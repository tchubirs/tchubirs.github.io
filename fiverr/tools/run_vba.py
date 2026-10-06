"""Run an Excel VBA macro on a copy of a workbook in LibreOffice and say what it changed. Used for the
Fiverr gig "Google Sheets and Excel": a macro is tested on the client's own file before delivery.

Usage: python3 run_vba.py book.xlsm MacroName                     a macro that is in the workbook
       python3 run_vba.py book.xlsx MacroName --code Module1.bas  a module I wrote or fixed, put into a copy
              [--sheet NAME] [--answer yes|no] [--input TEXT] [--timeout 60] [-o after.xlsx] [--report report.txt]
       (needs LibreOffice with its Python bridge, apt-get install python3-uno, and pip install openpyxl)

MacroName can be Module1.MacroName when two modules have one with that name. The file itself is never
touched: LibreOffice opens a copy with macros off to keep the state before, then opens it again with
macros on (so Workbook_Open runs, as in Excel), runs the macro and keeps the state after.

A macro starts on the sheet the client is on, which matters for a recorded one (Range("A1") with no sheet
acts on that sheet). The run starts on the sheet --sheet names; without it, on the sheet with the macro's
button (a form button, or a shape or picture the macro is assigned to), as when the client clicks it;
else on the sheet open when the file opens. The report says which.

LibreOffice runs most VBA written for Excel. Tried on the usual kinds of macro: ranges and cells, adding
sheets, copying sheets into one, sorting, AutoFilter, Find, formatting, arrays, Collection, string and date
functions, WorksheetFunction and formulas. It has no Range.RemoveDuplicates and no Scripting.Dictionary
(Excel on a Mac has no Dictionary either). What needs Windows or Excel itself (those two, Outlook and
other programs, Windows API calls, UserForms, pivot tables, other files) may not run here, so the report
names those lines. MsgBox and InputBox
do not wait for anyone: MsgBox gets --answer (yes by default; no means No or Cancel), InputBox gets
--input, and the report lists every question. When the macro stops on an error, the report gives the
module, the line as the VBA editor numbers it, and the error. Last comes every change: sheets added or
removed, formulas, typed values and each result that moved, recalculated (audit_sheet.py --compare).
Colors, fonts and widths are not compared; -o keeps the workbook as the macro left it, to look at them.
"""
import argparse
import html
import os
import pathlib
import posixpath
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import xml.etree.ElementTree as ET
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_sheet  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402

NORMAL, CLASS = 1, 2                     # com.sun.star.script.ModuleType
HEADER = re.compile(r"(?i)^(?:rem attribute vba_moduletype=|option vbasupport 1|option classmodule)")
PROC_START = re.compile(r"(?i)^(?:(?:public|private|friend)\s+)?(?:static\s+)?"
                        r"(?:sub|function|property\s+(?:get|let|set))\s+(\w+)")
PROC_END = re.compile(r"(?i)^end\s+(?:sub|function|property)\b")
# Lines that cannot have a statement put in front of them.
UNTOUCHED = re.compile(r"(?i)^(?:case\b|else\b|elseif\b|end\b|next\b|loop\b|wend\b|#|'|rem\b|\d+\s|[a-z_]\w*:(?!=))")
OUTSIDE = [(r"(?i)\bCreateObject\s*\(\s*\"Outlook", "Outlook"),
           (r"(?i)\bCreateObject\s*\(\s*\"Scripting\.Dictionary", "Scripting.Dictionary, which only Excel on Windows "
            "has (a Mac does not)"),
           (r"(?i)\bCreateObject\s*\(\s*\"Scripting\.FileSystemObject", "FileSystemObject, which only Windows has"),
           (r"(?i)\b(?:CreateObject|GetObject)\s*\(", "another program (CreateObject or GetObject)"),
           (r"(?i)\.RemoveDuplicates\b", "Range.RemoveDuplicates, which LibreOffice does not have"),
           (r"(?i)^\s*(?:public\s+|private\s+)?declare\b", "a Windows API call (Declare)"),
           (r"(?i)\.Show\b|\bUserForm", "a UserForm"),
           (r"(?i)\bShell\s*\(|\bSendKeys\b", "another program (Shell or SendKeys)"),
           (r"(?i)\bPivot(?:Tables?|Caches?|Fields?)\b", "pivot tables"),
           (r"(?i)\bApplication\.OnTime\b", "Application.OnTime"),
           (r"(?i)\bWorkbooks\.(?:Open|Add)\b|\.SaveAs\b|\bSaveCopyAs\b", "other files (open or save)"),
           (r"(?i)^\s*(?:Kill|MkDir|RmDir|ChDir)\s", "files and folders on the computer")]
# Excel 365 records .Formula2 and .Formula2R1C1, which LibreOffice lacks; .Formula and .FormulaR1C1 give the
# same result unless the formula spills over more cells, so the run uses those.
FORMULA2 = re.compile(r"(?i)\.Formula2(R1C1Local|R1C1|Local)?\b")
# Code that acts on whatever sheet is open: ActiveSheet, Selection, or Range and Cells with no sheet before them.
OPEN_SHEET = re.compile(r"(?i)\b(?:ActiveSheet|ActiveCell|Selection)\b|(?<![.\w])(?:Range|Cells|Rows|Columns)\s*\(")
# The macro a button runs: macro="[0]!Module1.Name" on a form button in the sheet or on a shape in the drawing,
# <x:FmlaMacro>[0]!Name</x:FmlaMacro> in the VML drawing of a form button.
BUTTON_MACRO = re.compile(r'\bmacro="([^"]+)"|<(?:\w+:)?FmlaMacro>([^<]+)<')
FORMATTING = re.compile(r"(?i)\.(?:Interior|Font|NumberFormat|ColumnWidth|RowHeight|AutoFit|Borders|"
                        r"HorizontalAlignment|WrapText|Style)\b")
# LibreOffice reports some errors with its own names; these are the errors Excel gives for the same lines.
UNO_ERRORS = {"NoSuchElementException": ("9", "Subscript out of range: no sheet or item has that name"),
              "IndexOutOfBoundsException": ("9", "Subscript out of range: the number is past the last one"),
              "IllegalArgumentException": ("5", "Invalid procedure call or argument")}

SUPPORT = '''Option VBASupport 1
Public FiverrAt As String
Public FiverrAsked As String
Public FiverrSays As String
Public FiverrTyped As String

Public Function FiverrAnswer(Prompt, Optional Buttons) As Integer
    Dim kind As Integer
    If Not IsMissing(Buttons) Then kind = Buttons And 7
    FiverrAsked = FiverrAsked & "MsgBox: " & Prompt & Chr(2)
    If kind = 3 Or kind = 4 Then
        FiverrAnswer = IIf(FiverrSays = "no", 7, 6)
    ElseIf kind = 1 Or kind = 5 Then
        FiverrAnswer = IIf(FiverrSays = "no", 2, IIf(kind = 1, 1, 4))
    ElseIf kind = 2 Then
        FiverrAnswer = IIf(FiverrSays = "no", 3, 4)
    Else
        FiverrAnswer = 1
    End If
End Function

Public Function FiverrInput(Prompt) As String
    FiverrAsked = FiverrAsked & "InputBox: " & Prompt & Chr(2)
    FiverrInput = FiverrTyped
End Function
'''
STANDINS = {"MsgBox": '''
Private Function MsgBox(Prompt, Optional Buttons, Optional Title, Optional HelpFile, Optional Context) As Integer
    MsgBox = FiverrAnswer(Prompt, Buttons)
End Function
''', "InputBox": '''
Private Function InputBox(Prompt, Optional Title, Optional Default, Optional XPos, Optional YPos, Optional HelpFile, _
                          Optional Context) As String
    InputBox = FiverrInput(Prompt)
End Function
'''}
READY = '''
Public Function FiverrReady() As String
    FiverrReady = "ready"
End Function
'''
MAIN = '''
Public Function FiverrMain() As String
    FiverrSays = "{says}"
    FiverrTyped = "{typed}"
    On Error GoTo FiverrFailed
    {macro}
    FiverrMain = "ok" & Chr(1) & FiverrAsked
    Exit Function
FiverrFailed:
    FiverrMain = "error" & Chr(1) & FiverrAsked & Chr(1) & FiverrAt & Chr(1) & Err & Chr(1) & Error$
End Function
'''


class Unrunnable(Exception):
    """Something that keeps the macro from being run at all."""


def prop(name, value):
    from com.sun.star.beans import PropertyValue
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


class Office:
    """A LibreOffice of its own for this run, with no screen and a fresh profile, driven through its Python
    bridge. Documents are opened with a view: without one, much of Excel's VBA (End(xlUp), Worksheets.Add)
    stops with "No ViewShell available"."""

    def __enter__(self):
        try:
            import uno
        except ImportError:
            raise Unrunnable("LibreOffice's Python bridge is missing: apt-get install python3-uno")
        if not shutil.which("soffice"):
            raise Unrunnable("LibreOffice is not installed (the soffice command)")
        self.uno, self.profile = uno, tempfile.mkdtemp(prefix="run-vba-")
        pipe = f"run_vba_{os.getpid()}_{time.monotonic_ns()}"
        self.proc = subprocess.Popen(
            ["soffice", "--headless", "--invisible", "--norestore", "--nologo", "--nodefault",
             f"-env:UserInstallation={pathlib.Path(self.profile, 'profile').as_uri()}",
             f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=dict(os.environ, HOME=self.profile),
            start_new_session=True)
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        ctx = None
        for _ in range(240):
            try:
                ctx = resolver.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except Exception:       # NoConnectException until LibreOffice is listening
                if self.proc.poll() is not None:
                    break
                time.sleep(0.25)
        if ctx is None:
            self.close()
            raise Unrunnable("LibreOffice did not start")
        self.desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        return self

    def __exit__(self, *exc):
        self.close()

    def close(self):
        """Stops LibreOffice. soffice starts soffice.bin under it, so the whole process group goes: killing
        soffice alone leaves soffice.bin running, and after a macro that never ends, running at full speed."""
        try:
            os.killpg(self.proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        self.proc.wait(timeout=30)
        for _ in range(100):            # a large process takes a moment to finish exiting
            if not running(self.proc.pid):
                break
            time.sleep(0.1)
        shutil.rmtree(self.profile, ignore_errors=True)

    def open(self, path, macros):
        doc = self.desktop.loadComponentFromURL(self.uno.systemPathToFileUrl(os.path.abspath(path)), "_blank", 0,
                                                (prop("MacroExecutionMode", 4 if macros else 0),))
        if doc is None:
            raise Unrunnable(f"LibreOffice could not open {os.path.basename(path)}")
        return doc

    def save(self, doc, path):
        doc.storeToURL(self.uno.systemPathToFileUrl(os.path.abspath(path)),
                       (prop("FilterName", "Calc MS Excel 2007 XML"),))


def running(group):
    """Whether a process of this process group is still running (one that has ended but was not yet
    collected by its parent does not count)."""
    for pid in filter(str.isdigit, os.listdir("/proc")):
        try:
            stat = pathlib.Path(f"/proc/{pid}/stat").read_text()
        except OSError:
            continue
        state, _, pgrp = stat.rsplit(")", 1)[1].split()[:3]
        if int(pgrp) == group and state != "Z":
            return True
    return False


def read_bas(path):
    """A module exported from the VBA editor (.bas): its name and its code as the editor shows it, without
    the Attribute lines, the module's and those Excel writes under a recorded macro with a shortcut key.
    Excel writes these files in the Windows code page."""
    raw = pathlib.Path(path).read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
    name = re.search(r'(?m)^Attribute VB_Name = "([^"]+)"', text)
    code = "\n".join(line for line in text.splitlines() if not line.startswith("Attribute "))
    return (name.group(1) if name else pathlib.Path(path).stem), code


def header(source):
    """How many lines LibreOffice keeps above the code the VBA editor shows (Rem Attribute, Option VBASupport)."""
    n = 0
    for line in source.splitlines():
        if not HEADER.match(line.strip()):
            break
        n += 1
    return n


def traced(module, source):
    """The module with each statement inside a procedure first noting where it is, so that an error can be
    placed: FiverrAt = "Module1 12": the line. Lines that cannot take a statement in front stay as they
    are: Case, Else, End, Next, Loop, labels, comments and the rest of a line continued with _."""
    skip, out, inside, continued = header(source), [], False, False
    for i, line in enumerate(source.splitlines()):
        text = line.strip()
        if i >= skip and inside and text and not continued and not PROC_END.match(text) and not UNTOUCHED.match(text):
            line = f'FiverrAt = "{module} {i + 1 - skip}": {line}'
        if not continued and PROC_START.match(text):
            inside = True
        elif PROC_END.match(text):
            inside = False
        continued = text.endswith(" _")
        out.append(line)
    return "\n".join(out)


def parameters(text):
    """What is inside the parentheses that open the parameter list, strings skipped, so that a comment after
    them, Sub Main() ' (Ctrl+M), is not taken for parameters; "" when there are none."""
    if not text.startswith("("):
        return ""
    depth, quoted = 0, False
    for i, ch in enumerate(text):
        if ch == '"':
            quoted = not quoted
        elif not quoted and ch in "()":
            depth += 1 if ch == "(" else -1
            if depth == 0:
                return text[1:i].strip()
    return text[1:].strip()             # the list goes on to the next line, after a _


def procedures(source):
    """The Sub and Function names of a module, with their parameters, and whether each is Private."""
    found = {}
    for line in source.splitlines():
        text = line.strip()
        m = PROC_START.match(text)
        if m:
            found[m.group(1).lower()] = (m.group(1), parameters(text[m.end():].lstrip()),
                                         text.lower().startswith("private"))
    return found


def outside(modules):
    """The lines that need Windows or Excel itself, as (module, line, what)."""
    found = []
    for name, source in modules.items():
        skip = header(source)
        for i, line in enumerate(source.splitlines()[skip:], start=1):
            code = line.split("'", 1)[0]
            for rx, what in OUTSIDE:
                if re.search(rx, code):
                    found.append((name, i, what))
                    break
    return found


def basic_text(value):
    return value.replace('"', '""').replace("\r", " ").replace("\n", " ")


def run(book, macro, code=(), answer="yes", typed="", timeout=60, keep=None, changes=True, sheet=None):
    """Run the macro on a copy and return what happened, as a dict the report is written from. Without
    changes, the comparison of before and after (the slow part) is left out. sheet is where it starts
    (see the top of this file)."""
    work = tempfile.mkdtemp(prefix="run-vba-work-")
    try:
        copy = os.path.join(work, "book" + os.path.splitext(book)[1].lower())
        shutil.copyfile(book, copy)
        before, after = os.path.join(work, "before.xlsx"), os.path.join(work, "after.xlsx")
        result = {"macro": macro, "status": None}
        with Office() as office:
            doc = office.open(copy, macros=False)
            office.save(doc, before)
            doc.close(True)

            doc = office.open(copy, macros=True)
            sheets, opened = list(doc.Sheets.getElementNames()), doc.CurrentController.ActiveSheet.Name
            libs = doc.BasicLibraries
            libs.VBACompatibilityMode = True
            # The workbook's own modules, then the .bas files, which replace a module of the same name.
            modules, kinds = {}, {}
            for library in libs.getElementNames():
                lib = libs.getByName(library)
                for name in lib.getElementNames():
                    modules[name] = lib.getByName(name)
                    try:
                        kinds[name] = lib.getModuleInfo(name).ModuleType
                    except Exception:   # a library without module information holds normal modules
                        kinds[name] = NORMAL
            for path in code:
                name, text = read_bas(path)
                modules[name], kinds[name] = "Option VBASupport 1\n" + text, NORMAL
            result["outside"] = outside(modules)
            result["formula2"] = sum(len(FORMULA2.findall(modules[name])) for name in modules
                                     if kinds[name] in (NORMAL, CLASS))

            where, _, wanted = macro.rpartition(".")
            homes = [name for name, source in modules.items() if kinds[name] == NORMAL and wanted.lower()
                     in procedures(source) and (not where or name.lower() == where.lower())]
            if len(homes) != 1:
                listed = sorted(f"{name}.{p[0]}" for name, source in modules.items() if kinds[name] == NORMAL
                                for p in procedures(source).values() if not p[1])
                why = "is in more than one module, name it as Module.Macro" if homes else "is not in the workbook"
                raise Unrunnable(f"{macro} {why}. Macros that can run: {', '.join(listed) or 'none'}")
            home = homes[0]
            name_, params, _ = procedures(modules[home])[wanted.lower()]
            if params and not all(p.strip().lower().startswith("optional") for p in params.split(",")):
                raise Unrunnable(f"{home}.{name_} needs arguments ({params}); run the macro that calls it")
            on = [title for module, title in buttons(book).get(name_.lower(), ())
                  if module in ("", home.lower()) and title in sheets]
            if sheet:
                start = next((title for title in sheets if title.lower() == sheet.lower()), None)
                if start is None:
                    raise Unrunnable(f"There is no sheet {sheet} in the workbook. Sheets: {', '.join(sheets)}")
                how, on = "chosen", []
            elif on:
                start, how, on = on[0], "button", on[1:]
            else:
                start, how = opened, "open"
            if len(sheets) > 1:
                result["start"] = (start, how, on)
                result["open_sheet"] = any(OPEN_SHEET.search(line.split("'", 1)[0])
                                           for line in body(modules[home], name_))

            main = MAIN.format(says=answer, typed=basic_text(typed), macro=name_)
            ordinary = [name for name in modules if kinds[name] == NORMAL]
            classes = [name for name in modules if kinds[name] == CLASS]
            attempts = iter(range(1, 100))

            def put(trace, skip=()):
                """A new library holding the test's own module, the class modules and the normal modules with
                the test's code, traced or as they are; True when LibreOffice compiles it. A new library each
                time: changing one that LibreOffice failed to compile can crash it."""
                name = f"FiverrRun{next(attempts)}"
                libs.createLibrary(name)
                lib = libs.getByName(name)
                lib.insertByName("FiverrRun", SUPPORT)
                for module in classes:
                    lib.insertByName(module, as_libreoffice(modules[module]))
                for module in ordinary:
                    if module not in skip:
                        source, extra = as_libreoffice(modules[module]), standins(modules[module]) + READY
                        extra += main if module == home else ""
                        lib.insertByName(module, (traced(module, source) if trace else source) + "\n" + extra)
                result["library"] = name
                return ask(doc, name, home, "FiverrReady") == "ready"

            if not put(True):
                # LibreOffice compiles a whole library at once: find the modules it cannot compile.
                broken = {name: uncompiled(doc, libs, modules[name]) for name in ordinary
                          if not compiles(doc, libs, modules[name])}
                if home in broken:
                    result.update(status="compile", module=home, procedures=broken[home])
                else:
                    result["left_out"] = broken     # Excel too compiles a module only when it is used
                    if not put(True, broken):
                        result["untraced"] = True
                        if not put(False, broken):
                            result.update(status="compile", module=home, procedures=[])
            if result["status"] is None:
                doc.CurrentController.setActiveSheet(doc.Sheets.getByName(start))
                result.update(call(doc, result["library"], home, timeout))
                if result["status"] == "timeout":
                    office.close()
                    return result
                office.save(doc, after)
                doc.close(True)
        if os.path.exists(after) and changes:
            result["changes"] = audit_sheet.compare(before, after)
            result["layout"] = layout(before, after)
            if keep:
                shutil.copyfile(after, keep)
        result["sources"] = modules
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


def body(source, name):
    """The lines of one procedure of a module, from its Sub or Function line to its End line."""
    lines, inside = [], False
    for line in source.splitlines():
        m = PROC_START.match(line.strip())
        if m and not inside:
            inside = m.group(1).lower() == name.lower()
        if inside:
            lines.append(line)
            if PROC_END.match(line.strip()):
                break
    return lines


def buttons(path):
    """The sheets with a button that runs each macro, read from an .xlsm or .xlsx: {macro name in lower case:
    [(its module in lower case, or "" when the button does not say, sheet)]}, in the order of the sheets. A
    button is a form button, or a shape or picture the macro is assigned to. {} for other files."""
    found = {}
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())

            def related(part):
                folder, base = posixpath.split(part)
                rels = posixpath.join(folder, "_rels", base + ".rels")
                if rels not in names:
                    return {}
                return {r.get("Id"): posixpath.normpath(posixpath.join(folder, r.get("Target", ""))).lstrip("/")
                        for r in ET.fromstring(z.read(rels)) if r.get("TargetMode") != "External"}

            if "xl/workbook.xml" not in names:
                return found
            parts = related("xl/workbook.xml")
            for sheet in ET.fromstring(z.read("xl/workbook.xml")).iter():
                if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                    continue
                part = parts.get(next((v for k, v in sheet.attrib.items() if k.endswith("}id")), None))
                if part not in names:
                    continue
                drawn = [p for p in related(part).values() if p in names and p.startswith("xl/drawings/")]
                for p in [part] + drawn:
                    for pair in BUTTON_MACRO.findall(z.read(p).decode("utf-8", "replace")):
                        module, _, name = html.unescape("".join(pair)).rpartition("!")[2].strip().rpartition(".")
                        entry = (module.lower(), sheet.get("name"))
                        if name and entry not in found.setdefault(name.lower(), []):
                            found[name.lower()].append(entry)
    except (zipfile.BadZipFile, OSError, ET.ParseError):      # not a file Excel saves as zip (.xls), or damaged
        return {}
    return found


def as_libreoffice(source):
    """The module with .Formula2 and .Formula2R1C1, of Excel 365, as .Formula and .FormulaR1C1, on the same lines."""
    return FORMULA2.sub(lambda m: ".Formula" + (m.group(1) or ""), source)


def hidden(path):
    """The hidden rows and columns of each sheet of a saved workbook, as {sheet: (rows, column numbers)}."""
    out = {}
    for ws in load_workbook(path).worksheets:
        rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        cols = {c for key, d in ws.column_dimensions.items() if d.hidden
                for c in range(d.min or column_index_from_string(key), (d.max or column_index_from_string(key)) + 1)}
        out[ws.title] = (rows, cols)
    return out


def spans(numbers, name=str):
    """3, 4, 5 and 9 as "3 to 5, 9"; name turns a number into what is shown (a column letter)."""
    out, numbers = [], sorted(numbers)
    while numbers:
        first = last = numbers.pop(0)
        while numbers and numbers[0] == last + 1:
            last = numbers.pop(0)
        out.append(name(first) if first == last else f"{name(first)} to {name(last)}")
    return ", ".join(out)


def layout(before, after):
    """The rows and columns the macro hid or showed, as lines like "Travel: column C hidden"."""
    old, new = hidden(before), hidden(after)
    lines = []
    for sheet, (rows, cols) in new.items():
        was_rows, was_cols = old.get(sheet, (set(), set()))
        for what, items, name in (("row", rows - was_rows, str), ("column", cols - was_cols, get_column_letter)):
            if items:
                lines.append(f"{sheet}: {what}{'s' if len(items) > 1 else ''} {spans(items, name)} hidden")
        for what, items, name in (("row", was_rows - rows, str), ("column", was_cols - cols, get_column_letter)):
            if items:
                lines.append(f"{sheet}: {what}{'s' if len(items) > 1 else ''} {spans(items, name)} shown again")
    return lines


def ask(doc, library, module, function):
    """The value a function in a document module returns, or None when it could not run."""
    try:
        script = doc.getScriptProvider().getScript(
            f"vnd.sun.star.script:{library}.{module}.{function}?language=Basic&location=document")
        return script.invoke((), (), ())[0]
    except Exception:       # the bridge's own errors: the function is not there or did not compile
        return None


def standins(source):
    """MsgBox and InputBox of the test's own, for a module that uses them. They have to be in the same
    module: from another one, LibreOffice's own MsgBox is called."""
    return "".join(text for word, text in STANDINS.items() if re.search(rf"(?i)(?<![.\w]){word}\b", source)
                   and not re.search(rf"(?im)^\s*(?:public\s+|private\s+)?function\s+{word}\b", source))


def compiles(doc, libs, source):
    """Whether LibreOffice compiles this code, tried in a new library of its own (see put in run)."""
    name = next(f"FiverrCheck{n}" for n in range(1, 10000) if not libs.hasByName(f"FiverrCheck{n}"))
    libs.createLibrary(name)
    libs.getByName(name).insertByName("Trial", source + READY)
    return ask(doc, name, "Trial", "FiverrReady") == "ready"


def uncompiled(doc, libs, source):
    """The procedures of a module that LibreOffice cannot compile, each tried with the module's declarations."""
    skip = header(source)
    declarations, blocks, current = [], [], None
    for line in source.splitlines()[skip:]:
        text = line.strip()
        if current is None and PROC_START.match(text):
            current = [line]
        elif current is not None:
            current.append(line)
            if PROC_END.match(text):
                blocks.append(current)
                current = None
        else:
            declarations.append(line)
    top = "Option VBASupport 1\n" + "\n".join(declarations)
    if not compiles(doc, libs, top):
        return ["the declarations above the procedures"]
    return [PROC_START.match(block[0].strip()).group(1) for block in blocks
            if not compiles(doc, libs, top + "\n" + "\n".join(block))]


def call(doc, library, home, timeout):
    """Runs FiverrMain in its own thread, so that a macro that never ends can be stopped."""
    done = {}

    def go():
        done["value"] = ask(doc, library, home, "FiverrMain")
    started = time.monotonic()
    thread = threading.Thread(target=go, daemon=True)
    thread.start()
    thread.join(timeout)
    seconds = round(time.monotonic() - started, 2)
    if thread.is_alive():
        return {"status": "timeout", "seconds": seconds}
    value = done.get("value")
    if not value:
        return {"status": "ended", "seconds": seconds}
    parts = value.split("\x01")
    asked = [" ".join(q.split()) for q in parts[1].split("\x02") if q.strip()]    # a prompt on several lines is one
    if parts[0] == "ok":
        return {"status": "ok", "seconds": seconds, "asked": asked}
    at, number, message = (parts + ["", "", ""])[2:5]
    for name, (excel, text) in UNO_ERRORS.items():
        if name in message:
            number, message = excel, f"{text} (LibreOffice: {name})"
    module, _, line = at.partition(" ")
    return {"status": "error", "seconds": seconds, "asked": asked, "at": (module, int(line)) if line else None,
            "number": number, "message": " ".join(message.split())}


def report(result, answer):
    """The report as lines of text."""
    macro, lines = result["macro"], []
    status = result["status"]
    if status == "ok":
        lines.append(f"{macro} ran to the end in {result['seconds']:g} s, with no error.")
    elif status == "error":
        place = result["at"]
        where = f" at {place[0]} line {place[1]}" if place else ""
        lines.append(f"{macro} stopped{where} with error {result['number']}: {result['message']}")
        if place:
            source = result["sources"].get(place[0], "")
            code = source.splitlines()[header(source):]
            if 0 < place[1] <= len(code):
                lines.append(f"  {place[1]}: {code[place[1] - 1].strip()}")
    elif status == "timeout":
        lines.append(f"{macro} did not finish in {result['seconds']:g} s: a loop that never ends, or it is very slow. "
                     "Nothing was saved.")
    elif status == "ended":
        lines.append(f"{macro} ended without coming back to the test: an End statement, or an error LibreOffice "
                     "kept to itself.")
    elif status == "compile":
        names = ", ".join(result["procedures"]) or "a line outside the procedures"
        lines.append(f"LibreOffice cannot compile {result['module']} ({names}), so {macro} was not run. Some VBA only "
                     "Excel knows; check those procedures in Excel.")
    start = result.get("start")
    if start and status != "compile":
        sheet, how, more = start
        if how == "chosen":
            lines.append(f"Started on the sheet {sheet}, as asked with --sheet.")
        elif how == "button":
            lines.append(f"Started on the sheet {sheet}, where its button is"
                         + (f" (it has one on {', '.join(more)} too)." if more else "."))
        else:
            lines.append(f"Started on the sheet {sheet}, the one open when the file opens"
                         + ("; the macro works on the open sheet (ActiveSheet, Selection or a Range with no "
                            "sheet), so --sheet NAME starts it on another." if result.get("open_sheet") else "."))
    if result.get("left_out"):
        lines.append("Left out of the run, LibreOffice cannot compile them (Excel compiles a module only when it is "
                     "used, so the macro can still work there): " + ", ".join(
                         f"{name} ({', '.join(procs) or 'unknown'})" for name, procs in result["left_out"].items()))
    if result.get("untraced"):
        lines.append("  The lines of an error are not known in this run (the code runs without line marks).")
    if result.get("asked"):
        lines.append(f"Questions it asked, answered {answer}:")
        lines += [f"  {q}" for q in result["asked"]]
    if result.get("outside"):
        lines.append("Lines that need Windows or Excel itself and may not run here:")
        lines += [f"  {module} line {n}: {what}" for module, n, what in result["outside"]]
    if result.get("formula2"):
        lines.append("Formula2 and Formula2R1C1 (Excel 365) in the code ran as Formula and FormulaR1C1: the same "
                     "result unless a formula spills over more cells.")
    if "changes" in result:
        changes = audit_sheet.describe_changes(result["changes"])
        lines.append("What changed in the workbook:" if status == "ok" else "What changed before it stopped:")
        lines += changes[1:]
        lines += [f"  {line}" for line in result.get("layout", ())]
        if any(FORMATTING.search(source) for source in result["sources"].values()):
            lines.append("  Colors, fonts, number formats and widths are not compared: look at them in the file "
                         "saved with -o.")
    return lines


def main(argv=None):
    ap = argparse.ArgumentParser(description="Run a VBA macro on a copy of a workbook and list what it changed.")
    ap.add_argument("book", help=".xlsm, .xlsx, .xls or .ods")
    ap.add_argument("macro", help="the Sub to run, or Module1.Sub")
    ap.add_argument("--code", nargs="+", default=[], help=".bas modules to add to the copy (or to replace there)")
    ap.add_argument("--sheet", help="the sheet the macro starts on (else the one with its button, or the one open)")
    ap.add_argument("--answer", choices=["yes", "no"], default="yes", help="how MsgBox questions are answered")
    ap.add_argument("--input", default="", help="what InputBox gets")
    ap.add_argument("--timeout", type=float, default=60, help="seconds before the macro is stopped")
    ap.add_argument("-o", "--out", help="keep the workbook as the macro left it (.xlsx)")
    ap.add_argument("--report", help="write the report to this file too")
    a = ap.parse_args(argv)
    try:
        result = run(a.book, a.macro, a.code, a.answer, a.input, a.timeout, a.out, sheet=a.sheet)
    except Unrunnable as e:
        print(e)
        return 2
    text = "\n".join(report(result, a.answer))
    print(text)
    if a.report:
        pathlib.Path(a.report).write_text(text + "\n", encoding="utf-8")
    return 0 if result["status"] == "ok" else 1


if __name__ == "__main__":
    sys.exit(main())
