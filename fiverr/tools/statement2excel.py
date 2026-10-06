"""Bank statement PDF -> clean Excel (Transactions + Monthly summary + checks).

Used to fulfil the Fiverr gig "convert bank statement PDF to Excel". Works on PDFs
downloaded from online banking and on scans. Every bank lays pages out differently,
so this is a strong first pass that is then checked against the statement's own totals.

Usage: python3 statement2excel.py statement.pdf [more.pdf ...] -o out.xlsx   (or out.csv)
       [--dates dmy|mdy] (default: guess from the data)
       [--categories [rules.json]]  a Category column and a By category sheet; categories.json is the
                                    default list of keywords, copied and edited when a client wants others
       [--date-format dd/mm/yyyy] [--sep ";" --decimal ","]  for the client's Excel
       [--ocr]  read every page as a scan, also pages with text (a scan whose text layer is poor)
       [--lang fra]  the languages of the scan (default: English, French, Portuguese and Spanish)
       [--whole]  also amounts without cents (12.990, 1.250.000), as banks in Chile print them
       [--password 1503]  for a statement locked with a password (JPG and PNG pictures are read as scans)
       [--for quickbooks] [--for xero]  also the CSV file to upload to QuickBooks Online or Xero, next to
                                         the output: out-quickbooks.csv, out-xero.csv

A page that is only a picture is read with Tesseract (apt-get install tesseract-ocr tesseract-ocr-fra
tesseract-ocr-por tesseract-ocr-spa): it is turned upright if it was scanned sideways or upside down,
straightened, and read as one block, so each line of the table is one row. The console then lists the
rows whose date or amount the OCR was unsure of, and every failed balance check names its two rows.

Separate debit and credit columns are read from the heading line (Debit/Credit, Withdrawals/Deposits,
Paid out/Paid in, Débit/Crédit, Cargos/Abonos and so on): a number under the debit heading is money out.
The opening and closing balance lines are not transactions; they go to the Checks sheet, where the
opening balance plus every movement must give the closing balance. Dates without a year take it from
the dates the statement prints above its transactions ("1 December to 1 January 2026" puts 3 December
in 2025), and move to the next year when the statement crosses 31 December. Once a table with money
columns has been read, the heading of any other table (cheques, daily balances, a loan) ends it, and a
second transaction table after that, usually another account, is counted in the Checks sheet instead
of being mixed in. When that table starts with an opening balance equal to where the first one ended,
it is the next month of the same account joined in the same PDF, and it is read.
A table of cheques (number, date, amount, as US banks print them) is read as money out, and a table
of daily balances (Balance by Date) ends the transactions. The rows of each file come out in date
order, as the deposits, withdrawals and cheques of one statement are often in separate tables.

Each PDF is read on its own, then the files are put in the order of their first transaction, so
statements sent out of order come out in date order. A file with the same transactions as another is
read once. With several files, the Checks sheet gives the rows of each one and whether each starts at
the closing balance of the one before: where it does not, a statement may be missing.

Amounts without cents are read only with --whole, because without cents a document number or the 3 of
"cuota 3 de 12" looks like an amount too. They are read under the money columns of a heading or, when
the statement has a single amount column, at the end of a row (the amount, then the balance). When
nothing was read and the lines that start with a date hold such amounts, the console says to run again
with it.
"""
import argparse
import collections
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Month names and abbreviations in English, French, Portuguese and Spanish, without accents.
MONTHS = {word: month for month, words in enumerate([
    "jan january janv janvier ene enero janeiro", "feb february fev fevr fevrier febrero fevereiro",
    "mar march mars marzo marco", "apr april avr avril abr abril", "may mai mayo maio",
    "jun june juin junio junho", "jul july juil juillet julio julho", "aug august aou aout ago agosto",
    "sep sept september septembre septiembre set setembro", "oct october octobre octubre out outubro",
    "nov november novembre noviembre novembro", "dec december decembre dic diciembre dez dezembro"], start=1)
    for word in words.split()}

DATE_RE = re.compile(
    r"^(?P<d>\d{1,2})[./-](?P<m>\d{1,2})(?:[./-](?P<y>\d{2,4}))?\b"         # 12/09/2026, 12.09.26, 12/09
    r"|^(?P<iy>\d{4})-(?P<im>\d{2})-(?P<id>\d{2})\b"                          # 2026-09-12
    r"|^(?P<td>\d{1,2})(?:st|nd|rd|th|er)?\s*(?P<tm>[^\W\d_]{3,10})\.?(?:,?\s+(?P<ty>\d{4}))?\b"   # 1 February 2026, 02FEB
    r"|^(?P<mm>[^\W\d_]{3,10})\.?\s*(?:[\u2013-]\s*)?(?P<md>\d{1,2})(?:st|nd|rd|th)?"
    r"(?:,?\s+(?P<my>\d{4}))?\b")                                                # Nov 01, Nov - 01, Nov 1, 2019
TOTALS = re.compile(r"(?i)^(sub-?)?totals?\b(\s+(money|amount|debits?|credits?|withdrawals|deposits|paid|in|out|"
                    r"for|of|checks|cheques)\b|\s*:?\s*$)|^(totaux|sous-total)\b")
RULES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "categories.json")
LEADING_DATE = re.compile(r"^\d{1,2}[./-]\d{1,2}(?:[./-]\d{2,4})?\s+")   # the value date after the operation date

def phrase(pattern):
    """A pattern of words that may be one or two spaces apart in a line (see apart), case aside."""
    return re.compile("(?i)" + pattern.replace(" ", r"\s+"))


OPENING = phrase(r"solde pr[ée]c[ée]dent|ancien solde|solde (initial|d'ouverture)|previous (statement )?balance|"
                 r"opening balance|beginning balance|starting balance|balance (brought )?forward|brought forward|"
                 r"saldo anterior|saldo inicial")
CLOSING = phrase(r"nouveau solde|solde final|closing balance|new balance|ending balance|current statement balance|"
                 r"saldo final|saldo atual|saldo actual")
# The title of a table of daily balances, which ends the table of transactions.
DAILY = phrase(r"daily (ending |ledger )?balances?|balances? by date|daily balance summary")
# A cheque paid, in the table of cheques of a US statement: number, date, amount, often two or three a line.
CHECK_ROW = re.compile(r"(?<![\w/.,])(?P<no>\d{1,8})\s*[*^]?\s+(?P<date>\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?)\s+"
                       r"\$?\s?(?P<amount>\d{1,3}(?:,\d{3})*\.\d{2}|\d+\.\d{2})(?![\w.,])")
CHECK_WORDS = {"check", "checks", "cheque", "cheques", "serial", "chk", "no", "number", "#"}
EITHER = phrase(r"solde (cr[ée]diteur|d[ée]biteur|au)\b")       # opening before the transactions, closing after
CARRIED = phrase(r"carried forward|[àa] reporter|report de la page|suma y sigue|a transportar")
PAGE_TOTAL = phrase(r"total des op[ée]rations|totaux")
FOOTER = phrase(r"p[aá]g(e|ina) \d|balance|solde|saldo|total|^\s*\d+\s*(of|de|sur|/)\s*\d+\s*$")   # not a description
HEADS = {"debit": "out", "debits": "out", "withdrawal": "out", "withdrawals": "out", "out": "out",
         "debito": "out", "debitos": "out", "retiros": "out", "giros": "out", "cargo": "out", "cargos": "out",
         "saidas": "out", "levantamentos": "out", "depenses": "out",
         "credit": "in", "credits": "in", "deposit": "in", "deposits": "in", "in": "in", "credito": "in",
         "creditos": "in", "abono": "in", "abonos": "in", "ingresos": "in", "entradas": "in", "depositos": "in",
         "recettes": "in",
         "balance": "balance", "solde": "balance", "saldo": "balance"}
DATE_WORDS = {"date", "dates", "fecha", "data", "datum"}
WHEN = {"yyyy-mm-dd": "%Y-%m-%d", "dd/mm/yyyy": "%d/%m/%Y", "mm/dd/yyyy": "%m/%d/%Y"}
# A date with its year anywhere in a line: the period of a statement, the day it was made.
FULL_DATE = re.compile(
    r"(?<![\d/.-])(?P<a>\d{1,2})[./-](?P<b>\d{1,2})[./-](?P<y>20\d{2}|\d{2})(?!\d)"
    r"|(?<!\d)(?P<iy>20\d{2})-(?P<im>\d{2})-(?P<id>\d{2})(?!\d)"
    r"|(?<!\w)(?P<td>\d{1,2})(?:st|nd|rd|th|er)?\s+(?P<tm>[^\W\d_]{3,10})\.?,?\s+(?P<ty>20\d{2})(?!\d)"
    r"|(?<!\w)(?P<mm>[^\W\d_]{3,10})\.?\s+(?P<md>\d{1,2})(?:st|nd|rd|th)?,?\s+(?P<my>20\d{2})(?!\d)")
Read = collections.namedtuple("Read", "tx order opening closings extra scanned plain files twice")
IMPORTS = {"quickbooks": "QuickBooks Online", "xero": "Xero"}
QUICKBOOKS_LINES = 1000          # lines in one upload to QuickBooks Online, which also takes 350 KB at most
OCR_DPI = 300
OCR_LANGS = ("eng", "fra", "por", "spa")
UNSURE = 80                       # Tesseract's confidence, 0 to 100, under which a word is worth a look
DIGITS = str.maketrans("OolI|SB", "0011158")
NUMBER = re.compile(r"[-(]?[$€£]?\d[\d.,/]*\d\)?-?")
THOUSANDS = re.compile(r"[-(]?[$€£]?\d{1,3}(?:[:;]\d{3})+[.,]\d{2}\)?")    # 1:127,12 read for 1.127,12
AMOUNT_RE = re.compile(
    r"(?<![\w.,])(?P<neg>[-−(])?\s?(?P<cur>[$€£])?\s?"
    # One kind of thousands separator per number, so "5,000 505,491.59" stays two numbers.
    r"(?P<num>\d{1,3}(?:(?P<sep>[ ,.\u202f\u00a0])\d{3}(?:(?P=sep)\d{3})*)?[.,]\d{2}|\d+[.,]\d{2}|\.\d{2})"
    r"\)?\s?(?P<sign>CR|DR|Cr|Dr|-)?(?![\w]|[.,]\d)")          # 15.03 in the date 15.03.2026 is not one
# With --whole, also amounts without cents (12.990, 1.250.000, 948), as banks in Chile print them. A
# document number (0045217, 452173), a date, a time or a RUT (12.345.678-9) is none of them.
WHOLE_RE = re.compile(
    r"(?<![\w.,/:-])(?P<neg>[-−(])?\s?(?P<cur>[$€£])?\s?"
    r"(?P<num>\d{1,3}(?:(?P<sep>[ ,.\u202f\u00a0])\d{3}(?:(?P=sep)\d{3})*)?[.,]\d{2}|\d+[.,]\d{2}|\.\d{2}"
    r"|\d{1,3}(?:(?P<group>[.,])\d{3}(?:(?P=group)\d{3})*)|[1-9]\d{0,2}|0)"
    r"\)?\s?(?P<sign>CR|DR|Cr|Dr|-)?(?![\w/:]|[.,]\d|-\w)")
GAP = re.compile(r"[\s$€£+|]*")                  # what can stand between the amount and the balance


def parse_amount(m):
    num = m.group("num").replace(" ", "").replace(" ", "").replace(" ", "")
    if re.search(r"[.,]\d{2}$", num):
        whole, dec = num[:-3], num[-2:]
        whole = whole.replace(",", "").replace(".", "")
        value = float(f"{whole}.{dec}")
    else:
        value = float(num.replace(",", "").replace(".", ""))
    neg = m.group("neg") in ("-", "−", "(") or (m.group("sign") or "").upper() in ("DR", "-")
    return -value if neg else value


def parse_date(m, order, year_hint):
    """The date a DATE_RE match stands for, or None when it is not a real date (2 PIZZAS, 31 February)."""
    try:
        if m.group("iy"):
            return dt.date(int(m.group("iy")), int(m.group("im")), int(m.group("id")))
        if m.group("td") or m.group("mm"):
            day, word, year = (m.group("td"), m.group("tm"), m.group("ty")) if m.group("td") else \
                (m.group("md"), m.group("mm"), m.group("my"))
            month = MONTHS.get(plain_word(word))
            return dt.date(int(year) if year else year_hint, month, int(day)) if month else None
        a, b = int(m.group("d")), int(m.group("m"))
        y = m.group("y")
        year = year_hint if not y else (int(y) + 2000 if len(y) == 2 else int(y))
        day, month = (a, b) if order == "dmy" else (b, a)
        return dt.date(year, month, day)
    except ValueError:
        return None


def guess_order(lines):
    """dmy unless a first number above 12 never appears but a second one does."""
    first_big = second_big = 0
    for line in lines:
        m = re.match(r"^(\d{1,2})[./-](\d{1,2})", line)
        if m:
            first_big += int(m.group(1)) > 12
            second_big += int(m.group(2)) > 12
    return "mdy" if second_big > first_big else "dmy"


def plain_word(word):
    return unicodedata.normalize("NFKD", word).encode("ascii", "ignore").decode().lower().strip(".:()/*")


def heading(words):
    """x positions of the money columns named on a heading line, or None when the line is not one.

    In a table the money columns sit well to the right of the date and apart from each other; that keeps
    a bank name and a card line ("CREDIT AGRICOLE", "Carte de debit") from passing for a heading."""
    names = [plain_word(w) for _, _, w, _, _ in words]
    dates = [x0 for (x0, *_), name in zip(words, names) if name in DATE_WORDS]
    if not dates:
        return None
    found = {}
    for (x0, x1, *_), name in zip(words, names):
        kind = HEADS.get(name)
        if kind and kind not in found:
            found[kind] = (x0, x1)
    if "out" not in found or "in" not in found:
        return None
    out_x, in_x = found["out"][0], found["in"][0]
    return found if min(out_x, in_x) > min(dates) + 100 and abs(out_x - in_x) > 20 else None


def foreign(words):
    """The heading of some other table: a date word among its first two words, no digits, no money columns."""
    names = [plain_word(w) for _, _, w, _, _ in words]
    return (len(names) >= 3 and bool(DATE_WORDS & set(names[:2]))
            and not any(c.isdigit() for _, _, w, _, _ in words for c in w))


def column(cols, words, m):
    """Which money column an amount sits in, or "text" when it is left of them all (part of the description)."""
    inside = [(x0, x1) for x0, x1, _, start, end in words if start < m.end() and end > m.start()]
    if not inside:
        return None
    x0, x1 = min(a for a, _ in inside), max(b for _, b in inside)
    if x1 < min(a for a, _ in cols.values()) - 20:
        return "text"
    return min(cols, key=lambda k: min(abs((cols[k][0] + cols[k][1]) / 2 - (x0 + x1) / 2), abs(cols[k][1] - x1)))


def signed(value, kind):
    return -abs(value) if kind == "out" else abs(value) if kind == "in" else value


class Unreadable(Exception):
    """A scanned page and no Tesseract to read it."""


def text_lines(page):
    """The lines of a page with text, as (text, words, unsure): words sharing a baseline, left to right,
    each with where it sits and where it is in the text. Text pages have no unsure words."""
    rows, seen = {}, {}
    for x0, y0, x1, y1, word, *_ in page.get_text("words"):
        # The same word drawn again on itself: bold made by printing twice, or a text layer repeated.
        if any(abs(x0 - a) < 1.5 and abs(y1 - b) < 1.5 for a, b in seen.get(word, ())):
            continue
        seen.setdefault(word, []).append((x0, y1))
        rows.setdefault(round(y1 / 3), []).append((x0, x1, word, y1 - y0))
    lines = []
    for key in sorted(rows):
        text, words, right = "", [], None
        height = max(h for *_, h in rows[key])
        for x0, x1, word, _ in sorted(rows[key]):
            text += apart(x0, right, height)
            words.append((x0, x1, word, len(text), len(text) + len(word)))
            text += word
            right = x1
        lines.append((text, words, []))
    return lines


def apart(left, right, height):
    """What goes between two words of a line: one space, or two where they are farther apart than half
    the height of the tallest word of the line, as between two columns. "Cheque #31" and "148.11" in the
    next column then cannot read as 31 148.11, while the space inside 1 234,56 stays one."""
    return "" if right is None else "  " if left - right > height / 2 else " "


def is_scan(page):
    """A page that is a picture: an image over at least half of it and no text, or two words at most (a
    scanner's stamp). A bank's own page with a background image and a few words, "Page 3 of 3", is not one."""
    if len(page.get_text("words")) >= 3:
        return False
    return any(abs(pymupdf.Rect(info["bbox"]) & page.rect) >= abs(page.rect) / 2 for info in page.get_image_info())


def tesseract(img, *args, dpi=OCR_DPI, must=False):
    """What Tesseract prints for a page image. When it fails: "" (orientation on a page with little text),
    or with must, Unreadable with Tesseract's own words (a language that is not installed)."""
    buf = io.BytesIO()
    img.save(buf, "PNG")
    done = subprocess.run(["tesseract", "stdin", "stdout", "--dpi", str(dpi), *args], input=buf.getvalue(),
                          capture_output=True)
    if done.returncode and must:
        said = done.stderr.decode("utf-8", "replace").strip().splitlines()
        raise Unreadable("Tesseract could not read the page: " + (said[0] if said else f"exit {done.returncode}"))
    return done.stdout.decode("utf-8", "replace") if done.returncode == 0 else ""


def ocr_langs():
    """English, French, Portuguese and Spanish, those of them installed for Tesseract."""
    listed = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True).stdout.split()
    return "+".join(lang for lang in OCR_LANGS if lang in listed) or "eng"


def upright(img):
    """The page turned the right way up, when Tesseract finds it was scanned sideways or upside down."""
    osd = tesseract(img.reduce(2), "--psm", "0", dpi=OCR_DPI // 2)     # half the size is enough, and faster
    turn = re.search(r"Rotate: (\d+)", osd)
    sure = re.search(r"Orientation confidence: ([\d.]+)", osd)
    if turn and int(turn.group(1)) and sure and float(sure.group(1)) >= 2:
        return img.rotate(-int(turn.group(1)), expand=True, fillcolor=255)
    return img


def skew(img):
    """The angle, up to 5 degrees either way, that makes the lines of text level: the one where the rows
    with ink and the white rows between them are most sharply apart."""
    import numpy as np
    small = img.copy()
    small.thumbnail((1200, 1200))
    ink = small.point(lambda v: 255 if v < 160 else 0)

    def sharpness(angle):
        rows = np.asarray(ink.rotate(angle), dtype=np.float64).sum(axis=1)
        return float(np.square(np.diff(rows)).sum())
    best = max((a / 2 for a in range(-10, 11)), key=sharpness)
    return max((best + a / 10 for a in range(-4, 5)), key=sharpness)


def ocr_word(word):
    """A word as Tesseract read it, with plain hyphens for dashes, digits for the letters it can take for
    digits inside a number (1,2O4.56 is 1,204.56) and a thousands separator for a colon (1:127,12)."""
    word = word.replace("\u2014", "-").replace("\u2013", "-")
    if THOUSANDS.fullmatch(word):
        word = re.sub(r"(?<=\d)[:;](?=\d{3})", "." if word.rstrip(")")[-3] == "," else ",", word)
    if sum(c.isdigit() for c in word) >= 2:
        fixed = word.translate(DIGITS)
        if fixed != word and NUMBER.fullmatch(fixed) and sum(a != b for a, b in zip(word, fixed)) <= 2:
            return fixed
    return word


def ocr_lines(page, lang):
    """The lines of a scanned page, read by Tesseract once the page is upright and level, as
    (text, words, unsure) like a text page's; unsure holds where in the text the words are that
    Tesseract was not sure of."""
    from PIL import Image
    pix = page.get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY)
    img = upright(Image.frombytes("L", (pix.width, pix.height), pix.samples))
    angle = skew(img)
    if abs(angle) >= 0.1:
        img = img.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=255)
    found = {}
    # One block of text: each line then runs across the whole table, so it holds one row.
    for row in tesseract(img, "--psm", "6", "-l", lang, "tsv", must=True).splitlines()[1:]:
        f = row.split("\t")
        if len(f) == 12 and f[0] == "5" and f[11].strip():
            line = tuple(int(n) for n in f[2:5])
            found.setdefault(line, []).append((int(f[6]), int(f[8]), ocr_word(f[11].strip()), float(f[10]),
                                               int(f[9])))
    lines, scale = [], 72 / OCR_DPI
    for key in sorted(found):
        text, words, unsure, right = "", [], [], None
        height = max(h for *_, h in found[key])
        for left, width, word, conf, _ in sorted(found[key]):
            text += apart(left, right, height)
            right = left + width
            words.append((left * scale, (left + width) * scale, word, len(text), len(text) + len(word)))
            if conf < UNSURE:
                unsure.append((len(text), len(text) + len(word)))
            text += word
        lines.append((text, words, unsure))
    return lines


def full_dates(text, order):
    """Every date with a year in a line, wherever it is in the line: 01/12/2025, 2025-12-01, 1 December 2025,
    December 1, 2025."""
    found = []
    for m in FULL_DATE.finditer(text):
        try:
            if m.group("y"):
                a, b = int(m.group("a")), int(m.group("b"))
                day, month = (a, b) if order == "dmy" else (b, a)
                year = int(m.group("y"))
                found.append(dt.date(year + 2000 if year < 100 else year, month, day))
            elif m.group("iy"):
                found.append(dt.date(int(m.group("iy")), int(m.group("im")), int(m.group("id"))))
            else:
                day, word, year = ((m.group("td"), m.group("tm"), m.group("ty")) if m.group("td") else
                                   (m.group("md"), m.group("mm"), m.group("my")))
                month = MONTHS.get(plain_word(word))
                if month:
                    found.append(dt.date(int(year), month, int(day)))
        except ValueError:
            pass
    return found


def year_for(day, stated):
    """The year of the first date a statement prints without one (day holds its day and month), from the
    dates with a year that it prints above its transactions: the year that puts it nearest one of them.
    "1 December to 1 January 2026" puts 3 December in 2025; "For Dec 1 to Dec 31, 2022", with "as of
    April 2, 2024" further down, puts Dec 30 in 2022."""
    options = []
    for year in range(stated[0].year - 1, stated[-1].year + 2):
        try:
            options.append(day.replace(year=year))
        except ValueError:                           # 29 February
            pass
    return min(options, key=lambda d: (min(abs((d - s).days) for s in stated), d)).year


def read_pdf(path, ocr, lang, passwords=()):
    """The lines of one PDF as (text, words), the words the OCR was unsure of by line, how many pages
    were read from a scan, and the OCR languages (looked up at the first scanned page). A picture (JPG,
    PNG) is read as a scanned page. A file that cannot be opened, is neither a PDF nor a picture, or is
    locked with a password none of passwords opens, raises Unreadable with what to ask the client."""
    name = os.path.basename(path)
    try:
        doc = pymupdf.open(path)
    except (RuntimeError, OSError) as e:             # broken, empty or missing
        raise Unreadable(f"{name} cannot be opened ({e}). Ask the client to download it again.") from e
    if not (doc.is_pdf or doc.metadata.get("format") == "Image"):
        raise Unreadable(f"{name} is not a PDF or a picture (it reads as {doc.metadata.get('format')}). Ask the "
                         "client to download it again.")
    if doc.needs_pass and not any(doc.authenticate(password) for password in passwords):
        raise Unreadable(f"{name} is locked with a password. " + ("None of the passwords given opens it."
                         if passwords else "Ask the client for it and run again with --password."))
    lines, unsure, scanned = [], {}, 0
    for number, page in enumerate(doc, start=1):
        if ocr or is_scan(page):
            if not shutil.which("tesseract"):
                raise Unreadable(f"{os.path.basename(path)} page {number} is a scanned image. To read it, install "
                                 "Tesseract: apt-get install tesseract-ocr tesseract-ocr-fra tesseract-ocr-por "
                                 "tesseract-ocr-spa")
            lang = lang or ocr_langs()
            read = ocr_lines(page, lang)
            scanned += 1
        else:
            read = text_lines(page)
        for text, words, doubt in read:
            if doubt:
                unsure[len(lines)] = doubt
            lines.append((text, words))
    return lines, unsure, scanned, lang


def parse(lines, unsure, order, whole):
    """The transactions of one statement file as {date, desc, values, kinds, unsure}, its opening balance,
    every closing balance it prints, how many other transaction tables were left out and how many lines
    that start with a date have an amount without cents (1.250.000)."""
    first = next((i for i, (text, _) in enumerate(lines) if DATE_RE.match(text) and
                  (AMOUNT_RE.search(text) or whole and WHOLE_RE.search(text))), len(lines))
    stated = sorted(d for text, _ in lines[:first + 1] for d in full_dates(text, order))
    year_hint, anchored = dt.date.today().year, False
    for text, _ in lines:
        y = re.search(r"\b(20\d{2})\b", text)
        if y:
            year_hint = int(y.group(1))
            break
    plain = sum(1 for text, _ in lines if DATE_RE.match(text)
                and any(re.search(r"[.,]\d{3}$", m.group("num")) for m in WHOLE_RE.finditer(text)))
    tx, cols, closed, extra, opening, closings, last, recent = [], None, False, 0, None, [], None, []
    cheques = False                                   # in a table of cheques: number, date, amount

    def dated(dm):
        """The date a DATE_RE match stands for. A date printed without a year takes it from the dates
        the statement prints above its transactions, then moves to the next year after December."""
        nonlocal year_hint, anchored
        date = parse_date(dm, order, year_hint)
        if not any(dm.group(g) for g in ("y", "iy", "ty", "my")):
            day = parse_date(dm, order, 2000)        # 2000 has a 29 February
            if day and stated and not anchored:
                year_hint, anchored = year_for(day, stated), True
                date = parse_date(dm, order, year_hint)
            elif date and last and (last - date).days > 180:
                year_hint += 1                       # December, then January
                date = parse_date(dm, order, year_hint)
        return date

    def money(k, under):
        """The amounts on line k. Without cents a document number or the 3 of "3 de 12" looks like an
        amount too, so then only the numbers under the money columns count or, with no money columns,
        the numbers at the end of the line."""
        text, words = lines[k]
        if not whole:
            return list(AMOUNT_RE.finditer(text))
        found = list(WHOLE_RE.finditer(text))
        if under:
            return [m for m in found if column(under, words, m) != "text"]
        end = found[-1:]
        for m in reversed(found[:-1]):
            if not GAP.fullmatch(text, m.end(), end[0].start()):
                break
            end.insert(0, m)
        return end

    def ahead(i):
        """The words of this line and of up to two lines after it, while none of them has an amount."""
        out = []
        for k in range(i, min(i + 3, len(lines))):
            if money(k, cols):
                break
            out += lines[k][1]
        return sorted(out)

    def continues(i, heads):
        """Whether the table under the heading on line i is the next statement of the same account, as in
        several months joined in one PDF: its first amount is an opening balance equal to where the
        transactions read so far end. Another account starts elsewhere."""
        reached = closings[-1] if closings else next(
            (v for t in reversed(tx) for v, k in zip(t["values"], t["kinds"]) if k == "balance"), None)
        for k in range(i + 1, min(i + 6, len(lines))):
            placed = [(parse_amount(a), column(heads, lines[k][1], a)) for a in money(k, heads)]
            placed = [p for p in placed if p[1] != "text"]
            if placed:
                text = lines[k][0]
                return (reached is not None and bool(OPENING.search(text) or EITHER.search(text))
                        and abs(signed(*placed[-1]) - reached) < 0.005)
        return False

    for i, (line, words) in enumerate(lines):
        names = {plain_word(w) for _, _, w, _, _ in words}
        if cheques and not closed:
            paid = [(m, DATE_RE.match(m.group("date"))) for m in CHECK_ROW.finditer(line)]
            paid = [(m, dated(dm)) for m, dm in paid if dm]
            if paid and all(date for _, date in paid):
                for m, date in paid:
                    tx.append({"date": date, "desc": f"Check {m.group('no')}", "unsure": False, "desc_x": None,
                               "values": [parse_amount(AMOUNT_RE.search(m.group("amount")))], "kinds": ["out"]})
                    last = max(last, date) if last else date
                continue
            cheques = not any(c.isalpha() for c in line)   # a title or a heading ends the table
        if names & CHECK_WORDS and {"date", "amount"} <= names and not names & set(HEADS):
            cheques = True                           # Serial Date Amount, Check No. Date Paid Amount
            continue
        amounts = money(i, cols)
        # A heading can be spread over two or three lines ("Money" above "out"); look at them together.
        recent = (recent + [words])[-3:] if not amounts else []
        heads = None if amounts else heading(words) or heading(sorted(w for ws in recent for w in ws))
        if heads:
            if closed and tx and not continues(i, heads):
                extra += 1
            else:
                cols, closed = heads, False
            recent = []
            continue
        # The first line of a heading spread over two lines can look like another table's heading. Before
        # any transaction, the money columns were those of a summary box: this table has none.
        if cols and not amounts and foreign(words) and not heading(ahead(i)):
            cols, closed = None, bool(tx)
            continue
        if not amounts and len(words) <= 6 and DAILY.search(line):
            cols, closed = None, bool(tx)            # Balance by Date, Daily ledger balances
            continue
        if closed:
            continue
        placed = [(a, parse_amount(a), column(cols, words, a) if cols else None) for a in amounts]
        placed = [p for p in placed if p[2] != "text"]
        if placed and PAGE_TOTAL.search(line):
            continue                                 # the totals of a page or of the statement
        if placed and (OPENING.search(line) or CLOSING.search(line) or EITHER.search(line) or CARRIED.search(line)):
            value = signed(*placed[-1][1:])
            if CARRIED.search(line) and tx:
                closings.append(value)               # the last one carried is the closing balance
            elif OPENING.search(line) or ((EITHER.search(line) or CARRIED.search(line)) and not tx):
                opening = value if not tx else opening   # the table's own line comes after any summary
            else:
                closings.append(value)
            continue
        dm = DATE_RE.match(line)
        date = dated(dm) if dm else None
        # Many banks print the date only on the first row of each day: with money columns known, a row
        # without a date takes the date of the row above.
        if placed and (date or (cols and tx)):
            start = dm.end() if date else 0
            desc = " ".join(LEADING_DATE.sub("", line[start:placed[0][0].start()].strip(" -|")).split())
            if cols and all(k == "balance" for _, _, k in placed) or not date and TOTALS.match(desc):
                continue                             # a balance on a row of its own, or a total
            after = [x0 for x0, _, _, begin, _ in words if begin >= start]
            # The OCR was unsure of a word in the date or in one of the amounts.
            spans = [(a.start(), a.end()) for a, _, _ in placed] + ([dm.span()] if date else [])
            doubt = any(s0 < e1 and e0 > s1 for s0, e0 in spans for s1, e1 in unsure.get(i, ()))
            tx.append({"date": date or tx[-1]["date"], "desc": desc, "values": [v for _, v, _ in placed],
                       "kinds": [k for _, _, k in placed], "desc_x": after[0] if after else None,
                       "unsure": doubt})
            last = date or last
        elif tx and line.strip() and not amounts and not date:
            # A wrapped description starts under the description; a page heading starts at the margin. A
            # line without a letter (1.0, the version of the form) is not part of a description.
            under = tx[-1]["desc_x"] is not None and words[0][0] >= tx[-1]["desc_x"] - 5 and \
                any(c.isalpha() for c in line)
            if under and len(tx[-1]["desc"]) < 120 and not FOOTER.search(line):
                tx[-1]["desc"] = " ".join((tx[-1]["desc"] + " " + line).split())
    tx.sort(key=lambda t: t["date"])                 # stable: rows of one day keep their order
    return tx, opening, closings, extra, plain


def extract(paths, order=None, ocr=False, lang=None, whole=False, passwords=()):
    """Every file read on its own, then put in the order of its first transaction, so statements sent out
    of order come out in date order, and a file with the same transactions as another is read once. With
    ocr, every page is read as a scan; with whole, amounts without cents are read too.

    Returns a Read: the transactions, the date order, the opening balance of the first file, the closing
    balances in file order, the other transaction tables left out, the pages read from a scan, the lines
    with amounts without cents, each file kept {name, rows, first, last, opening, closing} and each file
    left out as (its name, the name of the file it repeats)."""
    files = []
    for path in paths:
        lines, unsure, scanned, lang = read_pdf(path, ocr, lang, passwords)
        files.append((os.path.basename(path), lines, unsure, scanned))
    order = order or guess_order([text for _, lines, _, _ in files for text, _ in lines])
    parts = []
    for name, lines, unsure, scanned in files:
        tx, opening, closings, extra, plain = parse(lines, unsure, order, whole)
        parts.append({"name": name, "tx": tx, "opening": opening, "closings": closings, "extra": extra,
                      "plain": plain, "scanned": scanned})
    parts.sort(key=lambda part: part["tx"][0]["date"] if part["tx"] else dt.date.max)
    kept, seen, twice = [], {}, []
    for part in parts:
        same = tuple((t["date"], t["desc"], tuple(t["values"])) for t in part["tx"])
        if part["tx"] and same in seen:
            twice.append((part["name"], seen[same]))
            continue
        seen.setdefault(same, part["name"])
        kept.append(part)
    tx = [t for part in kept for t in part["tx"]]
    listed = [{"name": part["name"], "rows": len(part["tx"]), "first": part["tx"][0]["date"] if part["tx"] else None,
               "last": part["tx"][-1]["date"] if part["tx"] else None, "opening": part["opening"],
               "closing": part["closings"][-1] if part["closings"] else None} for part in kept]
    return Read(tx, order, kept[0]["opening"] if kept else None, [c for part in kept for c in part["closings"]],
                sum(part["extra"] for part in kept), sum(part["scanned"] for part in parts),
                sum(part["plain"] for part in kept), listed, twice)


def columns(tx, opening=None):
    """Split each row's amounts into signed amount and running balance.

    With debit and credit columns, the column gives the sign. Otherwise two numbers on a row are
    amount and balance, and the sign of the amount comes from the balance movement when possible.
    """
    rows = []
    prev_bal = opening
    for t in tx:
        v, kinds = t["values"], t.get("kinds") or [None] * len(t["values"])
        if any(kinds):
            moves = [signed(x, k) for x, k in zip(v, kinds) if k in ("in", "out")]
            bals = [x for x, k in zip(v, kinds) if k == "balance"]
            amount, balance = sum(moves), (bals[-1] if bals else None)
        else:
            amount, balance = (v[0], None) if len(v) == 1 else (v[-2], v[-1])
        if balance is not None and prev_bal is not None and abs(abs(balance - prev_bal) - abs(amount)) < 0.005:
            amount = balance - prev_bal
        rows.append((t["date"], t["desc"], round(amount, 2), balance))
        if balance is not None:
            prev_bal = balance
    return rows


def fold(text):
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()


def load_rules(path):
    """Categories in the order they are tried, each with one pattern made of its keywords. A keyword
    matches whole words, accents and capitals aside; one ending in * also matches longer words."""
    with open(path, encoding="utf-8") as f:
        rules = json.load(f)
    out = []
    for category, words in rules.items():
        parts = [re.escape(fold(w).strip()[:-1]) if w.strip().endswith("*")
                 else re.escape(fold(w).strip()) + r"(?![a-z0-9])" for w in words]
        out.append((category, re.compile(r"(?<![a-z0-9])(?:" + "|".join(parts) + ")")))
    return out


def categorise(rows, rules):
    """The first category whose keywords appear in each description, or "" when none does."""
    return [next((name for name, rx in rules if rx.search(fold(desc or ""))), "") for _, desc, _, _ in rows]


def checks(rows, got):
    """The Checks sheet as (label, value) pairs, how many balance checks failed, for each failure the two
    rows of the Transactions sheet between which the balance does not follow, and the files after which
    the next one does not start at the closing balance (a statement may be missing there)."""
    order, opening, closings = got.order, got.opening, got.closings
    bals = ([(-1, opening)] if opening is not None else []) + [(i, b) for i, (_, _, _, b) in enumerate(rows)
                                                                if b is not None]
    mismatches, where = 0, []
    for (i0, b0), (i1, b1) in zip(bals, bals[1:]):
        moved = sum(rows[k][2] for k in range(i0 + 1, i1 + 1))
        if abs((b1 - b0) - moved) > 0.01:
            mismatches += 1
            where.append((i0 + 2 if i0 >= 0 else None, i1 + 2))    # rows of the sheet; None: the opening balance
    closes, closing = "not checked", (closings[-1] if closings else None)
    if closings and bals:
        start, base = bals[-1]
        reached = base + sum(r[2] for r in rows[start + 1:])
        same = [c for c in closings if abs(reached - c) < 0.01]
        closing, closes = (same[0], "yes") if same else (closings[-1], "no")
        mismatches += closes == "no"
    pairs = [("Transactions", len(rows)), ("Date order used", order),
             ("Opening balance", opening if opening is not None else "not on the statement"),
             ("Closing balance", closing if closing is not None else "not on the statement"),
             ("Opening balance plus movements gives the closing balance", closes),
             ("Balance steps checked", max(0, len(bals) - 1)), ("Balance mismatches", mismatches)]
    if got.extra:
        pairs.append(("Other transaction tables left out", got.extra))
    if got.scanned:
        pairs.append(("Pages read from a scan (OCR)", got.scanned))
    breaks = []
    if len(got.files) > 1:
        start = 2
        for k, f in enumerate(got.files, start=1):
            span = (f"rows {start} to {start + f['rows'] - 1}, {f['first']:%Y-%m-%d} to {f['last']:%Y-%m-%d}"
                    if f["rows"] else "no transactions")
            pairs.append((f"File {k}: {f['name']}", span))
            start += f["rows"]
        steps = [(a, b) for a, b in zip(got.files, got.files[1:])
                 if a["closing"] is not None and b["opening"] is not None]
        breaks = [(a["name"], b["name"]) for a, b in steps if abs(a["closing"] - b["opening"]) > 0.005]
        pairs.append(("Each file starts at the closing balance of the one before",
                      "no: " + ", ".join(f"{a} to {b}" for a, b in breaks) if breaks else "yes" if steps
                      else "not checked"))
    for name, same in got.twice:
        pairs.append((f"Left out: {name}", f"the same transactions as {same}"))
    return pairs, mismatches, where, breaks


def write(rows, out, got, cats=None, date_format="yyyy-mm-dd"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    head = ["Date", "Description", "Money in", "Money out", "Balance"] + (["Category"] if cats else [])
    ws.append(head)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F6F78")
    for i, (d, desc, amount, bal) in enumerate(rows):
        ws.append([d, desc, amount if amount > 0 else None, -amount if amount < 0 else None, bal]
                  + ([cats[i] or "Other"] if cats else []))
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(head) + 1):
            ws.cell(r, c).font = Font(name="Arial")
        ws.cell(r, 1).number_format = date_format
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0.00"
    for i, w in enumerate([12, 40, 13, 13, 13, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"
    ws.page_setup.orientation = "landscape"  # prints on one page width
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    sm = wb.create_sheet("Monthly summary")
    sm.append(["Month", "Money in", "Money out", "Net"])
    months = sorted({(d.year, d.month) for d, *_ in rows})
    for y, m in months:
        inn = sum(a for d, _, a, _ in rows if (d.year, d.month) == (y, m) and a > 0)
        out_ = -sum(a for d, _, a, _ in rows if (d.year, d.month) == (y, m) and a < 0)
        sm.append([f"{y}-{m:02d}", round(inn, 2), round(out_, 2), round(inn - out_, 2)])
    for c in sm[1]:
        c.font = Font(bold=True)
    for i in (1, 2, 3, 4):
        sm.column_dimensions[get_column_letter(i)].width = 14

    if cats:
        # Formulas, so the totals follow when the client moves a row to another category.
        bc = wb.create_sheet("By category")
        bc.append(["Category", "Money in", "Money out", "Net", "Rows"])
        names = list(dict.fromkeys(c or "Other" for c in cats))
        last = len(rows) + 1
        for r, name in enumerate(names, start=2):
            bc.append([name, f"=SUMIFS(Transactions!$C$2:$C${last},Transactions!$F$2:$F${last},A{r})",
                       f"=SUMIFS(Transactions!$D$2:$D${last},Transactions!$F$2:$F${last},A{r})", f"=B{r}-C{r}",
                       f"=COUNTIFS(Transactions!$F$2:$F${last},A{r})"])
            for c in (2, 3, 4):
                bc.cell(r, c).number_format = "#,##0.00"
        for c in bc[1]:
            c.font = Font(bold=True)
        for i, w in enumerate([24, 14, 14, 14, 8], start=1):
            bc.column_dimensions[get_column_letter(i)].width = w

    ck = wb.create_sheet("Checks")
    pairs, mismatches, where, breaks = checks(rows, got)
    if cats:
        pairs.append(("Rows without a category, under Other", sum(1 for c in cats if not c)))
    for pair in pairs:
        ck.append(list(pair))
    ck.column_dimensions["A"].width = 52
    wb.save(out)
    return mismatches, where, breaks


def write_csv(rows, out, cats=None, date_format="yyyy-mm-dd", sep=",", decimal="."):
    """The Transactions sheet as CSV, with a byte order mark so Excel reads the accents."""
    when = WHEN[date_format]

    def money(v):
        return "" if v is None else f"{v:.2f}".replace(".", decimal)

    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=sep)
        w.writerow(["Date", "Description", "Money in", "Money out", "Balance"] + (["Category"] if cats else []))
        for i, (d, desc, amount, bal) in enumerate(rows):
            w.writerow([d.strftime(when), desc, money(amount) if amount > 0 else "",
                        money(-amount) if amount < 0 else "", money(bal)] + ([cats[i] or "Other"] if cats else []))



def no_symbols(text):
    """A description QuickBooks takes: accents dropped, and only letters, digits, spaces, hyphens and
    apostrophes. Its help pages say a special character in a description can stop the upload, and its
    list of characters that are safe in names leaves out $ % ( ) " : and the slash."""
    text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^A-Za-z0-9 '-]", " ", text).split())


def write_import(rows, out, target, date_format):
    """The CSV files to upload as a bank statement, next to out: QuickBooks Online (Date, Description,
    Amount) or Xero (Date, Amount, Payee, Description, Reference). One amount, money out below zero, with
    no currency sign or thousands separator, and every date in one format. QuickBooks takes up to 1,000
    lines a file and no amount of 0, so a longer statement is split and a row of 0 is left out. Returns
    the files written with their number of rows, and how many rows were left out."""
    when = WHEN[date_format]
    stem = os.path.splitext(out)[0] + "-" + target
    if target == "quickbooks":
        head = ["Date", "Description", "Amount"]
        lines = [[d.strftime(when), no_symbols(desc), f"{amount:.2f}"] for d, desc, amount, _ in rows if amount]
        parts = [lines[i:i + QUICKBOOKS_LINES] for i in range(0, len(lines), QUICKBOOKS_LINES)] or [[]]
    else:
        head = ["Date", "Amount", "Payee", "Description", "Reference"]
        lines = [[d.strftime(when), f"{amount:.2f}", desc or "", "", ""] for d, desc, amount, _ in rows]
        parts = [lines]
    files = []
    for n, part in enumerate(parts, start=1):
        path = f"{stem}.csv" if len(parts) == 1 else f"{stem}-{n}.csv"
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(head)
            w.writerows(part)
        files.append((path, len(part)))
    return files, len(rows) - len(lines)

def main(argv=None):
    ap = argparse.ArgumentParser(description="Bank statement PDFs to one Excel or CSV file.")
    ap.add_argument("pdfs", nargs="+")
    ap.add_argument("-o", "--out", default="statement.xlsx", help="ends in .xlsx or .csv")
    ap.add_argument("--dates", choices=["dmy", "mdy"], help="how the statement writes dates (default: guessed)")
    ap.add_argument("--categories", nargs="?", const=RULES, metavar="RULES.json",
                    help="add a Category column; without a file, the rules in categories.json")
    ap.add_argument("--date-format", choices=list(WHEN),
                    help="default yyyy-mm-dd; in the files of --for, the day and month order of the statement")
    ap.add_argument("--sep", default=",", help="CSV separator, for example ; for Excel in French or Portuguese")
    ap.add_argument("--decimal", default=".", help="CSV decimal mark")
    ap.add_argument("--ocr", action="store_true", help="read every page as a scan, also pages that have text")
    ap.add_argument("--lang", help="languages for the OCR, for example fra or eng+spa (default: eng, fra, por "
                                   "and spa, those installed)")
    ap.add_argument("--whole", action="store_true",
                    help="also amounts without cents (12.990), under the money columns or at the end of a row")
    ap.add_argument("--password", action="append", default=[],
                    help="the password of a locked PDF, as the client sends it (may be given more than once)")
    ap.add_argument("--for", dest="imports", action="append", choices=list(IMPORTS), default=[],
                    help="also the CSV file to upload to QuickBooks Online or Xero (may be given twice)")
    a = ap.parse_args(argv)
    try:
        got = extract(a.pdfs, a.dates, a.ocr, a.lang, a.whole, a.password)
    except Unreadable as e:
        print(e)
        return 2
    rows = columns(got.tx, got.opening)
    cats = categorise(rows, load_rules(a.categories)) if a.categories else None
    if a.out.lower().endswith(".csv"):
        write_csv(rows, a.out, cats, a.date_format or "yyyy-mm-dd", a.sep, a.decimal)
        pairs, bad, where, breaks = checks(rows, got)
        print("; ".join(f"{label}: {value}" for label, value in pairs))
    else:
        bad, where, breaks = write(rows, a.out, got, cats, a.date_format or "yyyy-mm-dd")
    print(f"{len(rows)} transactions, dates {got.order}, balance mismatches {bad} -> {a.out}")
    if len(got.files) > 1:
        print("Files in date order:", ", ".join(f"{f['name']} ({f['rows']} rows)" for f in got.files))
    for name, same in got.twice:
        print(f"Left out {name}: the same transactions as {same}")
    for first, second in breaks:
        print(f"{second} does not start at the closing balance of {first}: is a statement missing between them?")
    if where:
        print("The balance does not follow between these rows of Transactions:",
              ", ".join(f"{first or 'the opening balance'} and {second}" for first, second in where))
    if got.scanned:
        unsure = [str(i + 2) for i, t in enumerate(got.tx) if t["unsure"]]
        print(f"Pages read from a scan: {got.scanned}. Rows of Transactions with a date or amount the OCR was "
              "unsure of:", ", ".join(unsure) or "none")
    for target in dict.fromkeys(a.imports):
        when = a.date_format or ("mm/dd/yyyy" if got.order == "mdy" else "dd/mm/yyyy")
        files, left = write_import(rows, a.out, target, when)
        print(f"For {IMPORTS[target]}, dates {when}:", ", ".join(f"{path} ({n} rows)" for path, n in files))
        if left:
            print(f"{left} rows with an amount of 0 left out, as QuickBooks does not take them")
    if got.extra:
        print(f"Other transaction tables left out: {got.extra}. That is another account, or a statement of this "
              "account out of date order inside one PDF: then split the PDF, one file per statement, and run "
              "again.")
    if not a.whole and got.plain >= 3 and got.plain > len(rows):
        print(f"{got.plain} lines that start with a date have amounts without cents (like 12.990), which are read "
              "only with --whole. Run again with --whole.")
    return 0 if rows else 1


if __name__ == "__main__":
    sys.exit(main())
