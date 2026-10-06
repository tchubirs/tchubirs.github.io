"""Round-trip test: draw synthetic statements (EU, US, French, UK and Chilean layouts), extract, compare;
then scan them and read them back with OCR."""
import contextlib
import csv
import datetime as dt
import io
import json
import os
import random
import re
import shutil
import sys
import tempfile
from difflib import SequenceMatcher

import numpy as np
import pymupdf
from openpyxl import load_workbook
from PIL import Image, ImageFilter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_sheet  # noqa: E402
import statement2excel as s2e  # noqa: E402

SHOPS = ["CARREFOUR MARKET", "SNCF VOYAGES", "AMAZON EU SARL", "PHARMACIE CENTRALE", "SALAIRE ACME SAS",
         "LOYER AGENCE DU PORT", "BOULANGERIE PAUL", "FREE MOBILE", "VIREMENT MAMAN", "UBER EATS"]


def fmt(x, eu):
    s = f"{abs(x):,.2f}"
    if eu:
        s = s.replace(",", " ").replace(".", ",").replace(" ", ".")
    return ("-" if x < 0 else "") + s


def make(path, eu, n=45, seed=1, dots=False):
    rnd = random.Random(seed)
    bal = 1200.00
    truth = []
    doc = pymupdf.open()
    page, y = None, 0
    day = dt.date(2026, 7, 1)
    for i in range(n):
        if page is None or y > 760:
            page = doc.new_page()
            page.insert_text((50, 40), "RELEVE DE COMPTE" if eu else "ACCOUNT STATEMENT", fontsize=12)
            page.insert_text((50, 60), f"Page {doc.page_count}", fontsize=8)
            page.insert_text((50, 90), "Date      Libelle / Description                         Montant     Solde",
                             fontsize=8)
            y = 110
        day += dt.timedelta(days=rnd.randint(0, 2))
        shop = rnd.choice(SHOPS)
        amount = round(rnd.uniform(2100, 2400), 2) if shop.startswith("SALAIRE") else -round(rnd.uniform(3, 180), 2)
        bal = round(bal + amount, 2)
        truth.append((day, amount, bal))
        date_s = day.strftime("%d.%m.%Y" if dots else "%d/%m/%Y") if eu else day.strftime("%m/%d/%Y")
        page.insert_text((50, y), date_s, fontsize=9)
        page.insert_text((115, y), f"{'CB ' if amount < 0 else ''}{shop}", fontsize=9)
        page.insert_text((400, y), fmt(amount, eu), fontsize=9)
        page.insert_text((480, y), fmt(bal, eu), fontsize=9)
        y += 14
        if i % 7 == 3:  # a wrapped description line
            page.insert_text((115, y), "REF 8841-" + str(i), fontsize=8)
            y += 12
    doc.save(path)
    return truth


def check(eu, dots=False):
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "s.pdf"), os.path.join(tmp, "s.xlsx")
        truth = make(pdf, eu, seed=7 if eu else 11, dots=dots)
        assert s2e.main([pdf, "-o", out]) == 0
        wb = load_workbook(out)
        ws = wb["Transactions"]
        got = [(r[0].value.date() if hasattr(r[0].value, "date") else r[0].value,
                (r[2].value or 0) - (r[3].value or 0), r[4].value) for r in ws.iter_rows(min_row=2)]
        assert len(got) == len(truth), (len(got), len(truth))
        for (gd, ga, gb), (td, ta, tb) in zip(got, truth):
            assert gd == td, (gd, td)
            assert abs(ga - ta) < 0.005, (ga, ta)
            assert abs(gb - tb) < 0.005, (gb, tb)
        checks = {r[0].value: r[1].value for r in wb["Checks"].iter_rows()}
        assert checks["Balance mismatches"] == 0, checks
        descs = [r[1].value or "" for r in ws.iter_rows(min_row=2)]
        assert any("REF 8841" in d for d in descs), "wrapped line lost"
        assert all(any(s in d for s in SHOPS) for d in descs), [d for d in descs if not any(s in d for s in SHOPS)][:3]
        assert not [d for d in descs if "RELEVE" in d or "STATEMENT" in d], "a page heading went into a description"
        print(("EU" if eu else "US"), "layout" + (" with dates like 02.07.2026" if dots else "") + ":", len(got),
              "transactions match, balances reconcile")


def make_columns(path, us, n=40, seed=3, split=False, start=dt.date(2026, 12, 1), opening=1234.56):
    """Money out and money in in separate columns, without signs, between opening and closing balance lines.
    French: no balance column, dates without a year from December to January, a value date after each
    date and an amount inside one description. US: a balance column carried to the next page, a summary
    above the table whose previous balance is not the table's, then a daily balance table and a savings
    account with the same columns and its own balance brought forward, which must stay out. With `split`,
    the heading of the later pages is on two lines, the first of which looks like the heading of some
    other table."""
    rnd = random.Random(seed)
    bal = opening
    truth = []
    doc = pymupdf.open()
    out_x, in_x, bal_x = (360, 430, 500) if us else (400, 470, None)
    state = {"page": None, "y": 0}

    def put(x, text):
        state["page"].insert_text((x, state["y"]), text, fontsize=9)

    def new_page():
        if state["page"] is not None and us:
            put(115, "BALANCE CARRIED FORWARD")
            put(bal_x, fmt(bal, False))
        state["page"], state["y"] = doc.new_page(), 40
        put(50, "ACCOUNT STATEMENT 12/01/2026 to 01/31/2027" if us else "RELEVE DE COMPTE du 01/12/2026 au 31/01/2027")
        state["y"] = 70
        heads = ([(50, "Date"), (115, "Description"), (out_x, "Withdrawals"), (in_x, "Deposits"), (bal_x, "Balance")]
                 if us else [(50, "Date"), (95, "Valeur"), (140, "Libellé"), (out_x, "Débit"), (in_x, "Crédit")])
        if split and doc.page_count > 1:
            heads = heads[:2] + [(250, "Reference")]
            for x, word in [(out_x, "Withdrawals"), (in_x, "Deposits"), (bal_x, "Balance")]:
                state["y"] = 80
                put(x, word)
            state["y"] = 70
        for x, word in heads:
            put(x, word)
        state["y"] = 90
        if us:
            put(115, "BALANCE BROUGHT FORWARD")
            put(bal_x, fmt(bal, False))
            state["y"] += 14

    new_page()
    if not us:
        put(140, "SOLDE PRECEDENT AU 30/11/2026")
        put(in_x, fmt(opening, True))
        state["y"] += 14
    day = start
    for i in range(n):
        if state["y"] > 600:
            new_page()
        day += dt.timedelta(days=rnd.randint(0, 2))
        pay = day.day in (5, 6) and not any(t[0].month == day.month and t[1] > 0 for t in truth)
        shop = "VIR SALAIRE ACME" if pay else ("RETRAIT DAB 50,00" if i == 9 else "CB " + rnd.choice(SHOPS))
        amount = 2150.00 if pay else (-50.00 if i == 9 else -round(rnd.uniform(3, 180), 2))
        bal = round(bal + amount, 2)
        truth.append((day, amount, bal if us else None))
        if us:
            put(50, day.strftime("%m/%d/%Y"))
            put(115, shop)
        else:
            put(50, day.strftime("%d/%m"))
            put(95, (day + dt.timedelta(days=rnd.randint(0, 2))).strftime("%d/%m"))
            put(140, shop)
        put(in_x if amount > 0 else out_x, fmt(abs(amount), not us))
        if us:
            put(bal_x, fmt(bal, False))
        state["y"] += 14
    if not us:
        put(140, "TOTAL DES OPERATIONS")
        put(out_x, fmt(-sum(t[1] for t in truth if t[1] < 0), True))
        put(in_x, fmt(sum(t[1] for t in truth if t[1] > 0), True))
        state["y"] += 14
    put(115 if us else 140, "CLOSING BALANCE" if us else "NOUVEAU SOLDE AU 31/01/2027")
    put(bal_x if us else (in_x if bal > 0 else out_x), fmt(abs(bal), not us))
    if us:
        first = doc[0]
        first.insert_text((50, 50), "Previous balance $999.99", fontsize=9)
        first.insert_text((50, 60), f"Ending balance ${fmt(bal, False)}", fontsize=9)
        heads = [(50, "Date"), (115, "Description"), (out_x, "Withdrawals"), (in_x, "Deposits"), (bal_x, "Balance")]
        for x_words in ([(50, "DAILY BALANCE SUMMARY")],
                        [(50, "Date"), (150, "Amount"), (250, "Date"), (350, "Amount")],
                        [(50, "12/01/2026"), (150, "$1,100.00"), (250, "12/02/2026"), (350, "$900.00")],
                        [(50, "SAVINGS Account Number: 9999")], heads,
                        [(115, "BALANCE BROUGHT FORWARD"), (bal_x, "1,900.00")],
                        [(50, "12/15/2026"), (115, "TRANSFER"), (in_x, "100.00"), (bal_x, "2,000.00")]):
            state["y"] += 14
            for x, word in x_words:
                put(x, word)
    doc.save(path)
    return truth, opening, bal


def check_columns(us, split=False):
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "c.pdf"), os.path.join(tmp, "c.xlsx")
        truth, opening, closing = make_columns(pdf, us, split=split)
        assert s2e.main([pdf, "-o", out]) == 0
        wb = load_workbook(out)
        ws = wb["Transactions"]
        got = [(r[0].value.date(), (r[2].value or 0) - (r[3].value or 0), r[4].value, r[1].value)
               for r in ws.iter_rows(min_row=2)]
        assert len(got) == len(truth), (len(got), len(truth))
        for (gd, ga, gb, desc), (td, ta, tb) in zip(got, truth):
            assert (gd, round(ga, 2), gb) == (td, ta, tb), ((gd, ga, gb, desc), (td, ta, tb))
            assert not s2e.LEADING_DATE.match(desc or ""), desc
            assert not any(w in (desc or "") for w in ("RELEVE", "STATEMENT", "BALANCE", "Date")), desc
        assert got[-1][0].year == 2027 and any(g[3] == "RETRAIT DAB 50,00" for g in got) or us, got[-3:]
        checks = {r[0].value: r[1].value for r in wb["Checks"].iter_rows()}
        assert (checks["Opening balance"], checks["Closing balance"]) == (opening, closing), checks
        assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
        assert checks["Balance mismatches"] == 0, checks
        assert checks["Balance steps checked"] == (len(truth) if us else 0), checks
        assert checks.get("Other transaction tables left out") == (1 if us else None), checks
        print(("US withdrawals and deposits" if us else "French debit and credit") + " columns"
              + (", heading on two lines after page 1" if split else "") + ":", len(got),
              "transactions with the right sign, opening and closing balances agree")


def check_bank_header():
    """A bank name and a card line above the table ("CREDIT AGRICOLE", "Carte de debit", "Date d'arret")
    must not be taken for the money columns of a statement that has one amount column."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "s.pdf"), os.path.join(tmp, "s.xlsx")
        truth = make(pdf, eu=True, seed=7)
        doc = pymupdf.open(pdf)
        for page in doc:
            for y, text in ((18, "CREDIT AGRICOLE"), (26, "Carte de débit"), (34, "Date d'arrêté du relevé")):
                page.insert_text((50, y), text, fontsize=8)
        doc.saveIncr()
        with contextlib.redirect_stdout(io.StringIO()):
            assert s2e.main([pdf, "-o", out]) == 0
        got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4])
               for r in load_workbook(out)["Transactions"].iter_rows(min_row=2, values_only=True)]
        assert got == truth, [(g, t) for g, t in zip(got, truth) if g != t][:3]
    print("bank name above the table: not taken for debit and credit columns")


UK_SHOPS = ["TESCO STORES 2041", "CARD PAYMENT TO COSTA COFFEE", "DIRECT DEBIT EE LIMITED", "TFL TRAVEL CH",
            "AMAZON MKTPLACE", "2 FOR 1 PIZZA CO"]


def make_uk(path, n=36, seed=5, start=dt.date(2026, 2, 1), opening=2500.00):
    """A UK layout: the heading spread over three lines (Paid above out and in), dates like 3 February only
    on the first row of each day, and the balance brought forward at the top and carried forward at the end."""
    rnd = random.Random(seed)
    bal = opening
    truth = []
    doc = pymupdf.open()
    page = doc.new_page()

    def put(x, y, text):
        page.insert_text((x, y), text, fontsize=9)

    end = (start + dt.timedelta(days=31)).replace(day=1)
    put(50, 40, f"Your statement {start.day} {start:%B} to {end.day} {end:%B} {end.year}")
    for x, y, word in [(330, 70, "Paid"), (400, 70, "Paid"), (50, 80, "Date"), (110, 80, "Description"),
                       (330, 90, "out"), (400, 90, "in"), (470, 90, "Balance")]:
        put(x, y, word)
    y = 110
    put(110, y, "Balance brought forward")
    put(470, y, fmt(bal, False))
    day, shown = start, None
    for i in range(n):
        y += 14
        day += dt.timedelta(days=rnd.choice([0, 0, 1, 2]))
        amount = 2400.00 if i == 10 else -round(rnd.uniform(2, 120), 2)
        bal = round(bal + amount, 2)
        truth.append((day, amount, bal))
        if day != shown:
            put(50, y, f"{day.day} {day.strftime('%B')}")
            shown = day
        put(110, y, "SALARY ACME LTD" if i == 10 else rnd.choice(UK_SHOPS))
        put(330 if amount < 0 else 400, y, fmt(abs(amount), False))
        put(470, y, fmt(bal, False))
    put(110, y + 14, "Balance carried forward")
    put(470, y + 14, fmt(bal, False))
    doc.save(path)
    return truth, opening, bal


def check_uk():
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "uk.pdf"), os.path.join(tmp, "uk.xlsx")
        truth, opening, closing = make_uk(pdf)
        with contextlib.redirect_stdout(io.StringIO()):
            assert s2e.main([pdf, "-o", out]) == 0
        wb = load_workbook(out)
        got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4], r[1])
               for r in wb["Transactions"].iter_rows(min_row=2, values_only=True)]
        assert [g[:3] for g in got] == truth, [(g, t) for g, t in zip(got, truth) if g[:3] != t][:3]
        assert any(g[3] == "2 FOR 1 PIZZA CO" for g in got), "a description that starts like a date was lost"
        checks = {r[0].value: r[1].value for r in wb["Checks"].iter_rows()}
        assert (checks["Opening balance"], checks["Closing balance"]) == (opening, closing), checks
        assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
        assert checks["Balance mismatches"] == 0, checks
    shared = sum(1 for a, b in zip(truth, truth[1:]) if a[0] == b[0])
    print(f"UK layout: {len(got)} transactions, {shared} of them without a printed date, the heading on three "
          "lines, balances brought and carried forward")



CL_SHOPS = ["COMPRA LIDER EXPRESS", "COMPRA COPEC LAS CONDES", "PAGO SERVIPAG ENEL", "GIRO CAJERO AUTOMATICO",
            "COMPRA FARMACIA AHUMADA", "PAC SEGURO AUTO", "COMPRA UBER TRIP", "COMPRA JUMBO"]


def cl(x):
    """A whole amount as banks in Chile print it: 1.250.000."""
    return f"{x:,}".replace(",", ".")


def make_whole(path, n=28, seed=9, single=False):
    """A cartola from a bank in Chile: amounts without cents (12.990, 1.250.000) right-aligned under Cargos,
    Abonos and Saldo, document numbers just left of them, a RUT and a "cuota 3 de 12" in the descriptions,
    a fee under 1.000, a summary box above the table and a page footer. With single, one Monto column of
    amounts with a sign (-$ 12.990) instead of Cargos and Abonos, as digital accounts print them; the second line of a
    description then ends in a word, since one that ends in a number is left out."""
    rnd = random.Random(seed)
    opening = bal = 1234567
    truth, doc = [], pymupdf.open()
    page = doc.new_page()

    def put(x, y, text, right=False):
        page.insert_text((x - pymupdf.get_text_length(text, fontsize=9) if right else x, y), text, fontsize=9)

    put(40, 40, "CARTOLA DE CUENTA CORRIENTE N° 45")
    put(40, 52, "Fecha de emisión: 01/04/2026   Periodo: 01/03/2026 al 31/03/2026")
    for x, word in [(40, "Fecha"), (95, "Descripción"), (315, "N° Docto")]:
        put(x, 110, word)
    for x, word in [(465, "Monto"), (545, "Saldo")] if single else [(395, "Cargos"), (465, "Abonos"), (545, "Saldo")]:
        put(x, 110, word, right=True)
    y = 126
    put(95, y, "SALDO ANTERIOR")
    put(545, y, cl(opening), right=True)
    day = dt.date(2026, 3, 2)
    special = {4: ("ABONO REMUNERACION ACME SPA", 1250000), 8: ("TRANSF A 12.345.678-9", -350000),
               12: ("CUOTA 3 DE 12 CREDITO CONSUMO", -85430), 15: ("COMISION MANTENCION", -4990),
               16: ("IVA COMISION", -948)}
    for i in range(n):
        y += 13
        day += dt.timedelta(days=rnd.randint(0, 2))
        desc, amount = special.get(i) or (rnd.choice(CL_SHOPS), -rnd.randrange(99, 8999) * 10)
        bal += amount
        truth.append((day, amount, bal))
        put(40, y, day.strftime("%d/%m/%Y"))
        put(95, y, desc)
        if i not in (15, 16):
            put(315, y, f"{rnd.randrange(1, 10 ** 7):07d}")
        if single:
            put(465, y, ("-" if amount < 0 else "") + "$ " + cl(abs(amount)), right=True)
        else:
            put(395 if amount < 0 else 465, y, cl(abs(amount)), right=True)
        put(545, y, cl(bal), right=True)
        if i == 8:                                   # what the transfer was for, on a second line
            y += 13
            put(95, y, "ARRIENDO MARZO" if single else "ARRIENDO MARZO DEPTO 21")
    put(95, y + 13, "SALDO FINAL")
    put(545, y + 13, cl(bal), right=True)
    spent, got = -sum(t[1] for t in truth if t[1] < 0), sum(t[1] for t in truth if t[1] > 0)
    for x, label, value in [(40, "Saldo anterior", opening), (160, "Total cargos", spent),
                            (280, "Total abonos", got), (400, "Saldo final", bal)]:
        put(x, 70, label)
        put(x, 82, cl(value))
    put(270, 800, "Página 1 de 1")
    doc.save(path)
    return truth, opening, bal


def check_whole():
    """Both cartolas read with --whole: every amount and balance exact, the document numbers, the RUT, the
    3 of "cuota 3 de 12" and the page footer not taken for amounts, the summary box not taken for the
    money columns of the table. Without --whole nothing is read and the console says to use it. The
    other layouts read the same with --whole."""
    found = [m.group() for m in s2e.WHOLE_RE.finditer("RUT 15.432.876-K 12.990- 0045217 15.03.2026 10:45 948 4.990")]
    assert found == ["12.990-", "948", "4.990"], found
    words = [(40, 62, "Fecha", 0, 5), (300, 324, "Cargo", 6, 11), (380, 405, "Abono", 12, 17),
             (460, 482, "Saldo", 18, 23)]
    assert s2e.heading(words) == {"out": (300, 324), "in": (380, 405), "balance": (460, 482)}, s2e.heading(words)
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "cl.pdf"), os.path.join(tmp, "cl.xlsx")
        for single in (False, True):
            truth, opening, closing = make_whole(pdf, single=single)
            code, rows, checks, said = read(pdf, out, "--whole")
            got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4]) for r in rows]
            assert code == 0 and got == truth, (single, said, [(g, t) for g, t in zip(got, truth) if g != t][:3])
            assert (checks["Opening balance"], checks["Closing balance"]) == (opening, closing), checks
            assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
            assert checks["Balance mismatches"] == 0 and checks["Balance steps checked"] == len(truth), checks
            descs = [r[1] for r in rows]
            wrapped = "ARRIENDO MARZO" if single else "ARRIENDO MARZO DEPTO 21"
            assert any(d.startswith("TRANSF A 12.345.678-9") and d.endswith(wrapped) for d in descs), descs
            assert any(d.startswith("CUOTA 3 DE 12 CREDITO CONSUMO") for d in descs), descs
            assert not any("Página" in d for d in descs) and "--whole" not in said, (descs, said)

            code, rows, _, said = read(pdf, out)
            assert code == 1 and not rows and f"{len(truth)} lines that start with a date" in said, said
            assert "Run again with --whole" in said, said

        for name, build in [("EU", lambda p: make(p, eu=True, seed=7)), ("US", lambda p: make(p, eu=False, seed=11)),
                            ("US columns", lambda p: make_columns(p, us=True)),
                            ("French", lambda p: make_columns(p, us=False)), ("UK", make_uk)]:
            build(pdf)
            plain, whole = read(pdf, out), read(pdf, out, "--whole")
            assert plain[0] == 0 and plain[1:3] == whole[1:3] and "--whole" not in plain[3], (name, plain[3])
    print(f"Chilean cartolas with --whole, with Cargos and Abonos and with one Monto column: {len(truth)} amounts "
          "without cents and their balances exact, document numbers, a RUT and the page footer left out; without "
          "it nothing is read and the console says why; the EU, US, French and UK layouts read the same with --whole")


EXPECTED = {"CARREFOUR MARKET": "Groceries", "SNCF VOYAGES": "Transport", "AMAZON EU SARL": "Shopping",
            "PHARMACIE CENTRALE": "Health", "SALAIRE ACME SAS": "Salary and income",
            "LOYER AGENCE DU PORT": "Housing", "BOULANGERIE PAUL": "Groceries", "FREE MOBILE": "Phone and internet",
            "VIREMENT MAMAN": "Transfers", "UBER EATS": "Restaurants and takeaway"}


def check_rules():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "rules.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"Health": ["pharmac*"], "Going out": ["bar"], "Food": ["épicerie"]}, f)
        rules = s2e.load_rules(path)
        rows = [(None, d, 0, None) for d in ["CB PHARMACIE DU PORT", "BARCLAYS FEE", "BAR TABAC", "EPICERIE FINE", "X"]]
        assert s2e.categorise(rows, rules) == ["Health", "", "Going out", "Food", ""]
    print("category rules: whole words, a * for longer words, accents and capitals ignored")


def check_categories():
    """Categories on the EU statement, the By category formulas recalculated in LibreOffice, and CSV."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out, out_csv = (os.path.join(tmp, name) for name in ("s.pdf", "s.xlsx", "s.csv"))
        make(pdf, eu=True, seed=7)
        assert s2e.main([pdf, "-o", out, "--categories", "--date-format", "dd/mm/yyyy"]) == 0
        wb = load_workbook(out)
        ws = wb["Transactions"]
        assert [c.value for c in ws[1]][-1] == "Category" and ws["A2"].number_format == "dd/mm/yyyy"
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for r in rows:
            shop = next(k for k in EXPECTED if k in r[1])
            assert r[5] == EXPECTED[shop], (r[1], r[5])
        assert {r[0].value: r[1].value for r in wb["Checks"].iter_rows()}["Rows without a category, under Other"] == 0

        found = audit_sheet.audit_files([out])[0]
        assert not found["must_fix"] and not found["odd"], found
        calc = load_workbook(audit_sheet.recalculate([out], tempfile.mkdtemp(dir=tmp))[0], data_only=True)
        totals = {r[0]: r[1:] for r in calc["By category"].iter_rows(min_row=2, values_only=True)}
        assert set(totals) == set(EXPECTED.values()), totals
        for name, (inn, out_, net, count) in totals.items():
            mine = [r for r in rows if r[5] == name]
            assert abs(inn - sum(r[2] or 0 for r in mine)) < 0.005 and abs(out_ - sum(r[3] or 0 for r in mine)) < 0.005
            assert abs(net - (inn - out_)) < 0.005 and count == len(mine), (name, totals[name])

        assert s2e.main([pdf, "-o", out_csv, "--categories", "--date-format", "dd/mm/yyyy",
                         "--sep", ";", "--decimal", ","]) == 0
        with open(out_csv, encoding="utf-8-sig", newline="") as f:
            lines = list(csv.reader(f, delimiter=";"))
        assert lines[0] == ["Date", "Description", "Money in", "Money out", "Balance", "Category"], lines[0]
        assert len(lines) == len(rows) + 1
        for line, r in zip(lines[1:], rows):
            money = ["" if v is None else f"{v:.2f}".replace(".", ",") for v in r[2:5]]
            assert line == [r[0].strftime("%d/%m/%Y"), r[1], *money, r[5]], (line, r)
    print("categories:", len(rows), "rows sorted into", len(totals), "categories; the By category formulas give the "
          "same totals in LibreOffice; the CSV with ; and decimal commas matches the Excel file")


def csv_rows(path, encoding="utf-8"):
    with open(path, encoding=encoding, newline="") as f:
        return list(csv.reader(f))


def check_imports():
    """The files to upload to QuickBooks Online and Xero, in the formats their help pages give: for
    QuickBooks Date, Description and Amount, money out below zero, no currency sign or thousands separator,
    one date format, no symbol in a description, 1,000 lines a file at most and no amount of 0; for Xero
    Date and Amount with Payee, Description and Reference. The amounts are those of the Excel file."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "s.pdf"), os.path.join(tmp, "s.xlsx")
        for us in (False, True):
            make_columns(pdf, us=us)
            code, rows, _, said = read(pdf, out, "--for", "quickbooks", "--for", "xero")
            assert code == 0, said
            when = "%m/%d/%Y" if us else "%d/%m/%Y"
            moves = [(r[0].strftime(when), r[1] or "", f"{(r[2] or 0) - (r[3] or 0):.2f}") for r in rows]
            qb = csv_rows(os.path.join(tmp, "s-quickbooks.csv"), "ascii")
            assert qb[0] == ["Date", "Description", "Amount"] and len(qb) == len(rows) + 1, qb[:2]
            for line, (date, desc, amount) in zip(qb[1:], moves):
                assert line[0] == date and line[2] == amount, (line, date, amount)
                words = " ".join(re.split(r"\W+", desc)).strip()
                assert re.fullmatch(r"[A-Za-z0-9 '-]+", line[1]) and line[1] == words, (line[1], desc)
            assert "RETRAIT DAB 50 00" in [line[1] for line in qb], "the comma of RETRAIT DAB 50,00"
            xero = csv_rows(os.path.join(tmp, "s-xero.csv"))
            assert xero[0] == ["Date", "Amount", "Payee", "Description", "Reference"], xero[0]
            assert [(line[0], line[2], line[1]) for line in xero[1:]] == moves, "Xero rows differ from the Excel file"
            for name in ("s-quickbooks.csv", "s-xero.csv"):
                assert f"{name} ({len(rows)} rows)" in said, said
            assert f"dates {'mm/dd/yyyy' if us else 'dd/mm/yyyy'}" in said, said

        code, _, _, said = read(pdf, out, "--for", "quickbooks", "--date-format", "dd/mm/yyyy")
        first = dt.datetime.strptime(moves[0][0], "%m/%d/%Y").strftime("%d/%m/%Y")
        assert code == 0 and csv_rows(os.path.join(tmp, "s-quickbooks.csv"))[1][0] == first, "--date-format lost"

        rows = [(dt.date(2026, 1, 1) + dt.timedelta(days=i % 300), f"CAFÉ {i} & CO, LTD. (50%)",
                 0.0 if i in (5, 77) else round(-1.5 - i, 2), None) for i in range(2345)]
        files, left = s2e.write_import(rows, os.path.join(tmp, "big.xlsx"), "quickbooks", "dd/mm/yyyy")
        assert [(os.path.basename(path), n) for path, n in files] == [
            ("big-quickbooks-1.csv", 1000), ("big-quickbooks-2.csv", 1000), ("big-quickbooks-3.csv", 343)], files
        assert left == 2 and all(os.path.getsize(path) <= 350 * 1024 for path, _ in files)
        lines = [line for path, _ in files for line in csv_rows(path, "ascii")[1:]]
        assert len(lines) == 2343 and lines[0] == ["01/01/2026", "CAFE 0 CO LTD 50", "-1.50"], lines[0]
        assert "0.00" not in {line[2] for line in lines} and "-0.00" not in {line[2] for line in lines}
        assert abs(sum(float(line[2]) for line in lines) - sum(r[2] for r in rows)) < 0.005
    print("QuickBooks and Xero: the upload files hold every row of the Excel file with its sign, dates in the "
          "statement's day and month order unless --date-format says otherwise, descriptions without symbols "
          "for QuickBooks, and a statement of 2,345 rows split into files of 1,000 lines with the two rows of 0 "
          "left out")


def scan(src, dst, angle, dpi=200, turn=0, seed=1):
    """What a scanner makes of a printed PDF: each page a grey JPEG picture, tilted by `angle` degrees,
    grainy and a little blurred, and turned by `turn` degrees (90 is on its side, 180 upside down)."""
    grain = np.random.default_rng(seed)
    out = pymupdf.open()
    for page in pymupdf.open(src):
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        img = img.rotate(angle, resample=Image.BICUBIC, fillcolor=255).rotate(turn, expand=True)
        dots = np.asarray(img, dtype=np.float64) * 0.9 + 12 + grain.normal(0, 14, (img.height, img.width))
        img = Image.fromarray(np.clip(dots, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(0.6))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=60)
        w, h = (page.rect.width, page.rect.height)[::1 if turn % 180 == 0 else -1]
        sheet = out.new_page(width=w, height=h)
        sheet.insert_image(sheet.rect, stream=buf.getvalue())
    out.save(dst)


def read(pdf, out, *flags):
    """Exit code, Transactions rows, Checks sheet and what was printed; pdf may be a list of files."""
    with contextlib.redirect_stdout(io.StringIO()) as said:
        code = s2e.main([*(pdf if isinstance(pdf, list) else [pdf]), "-o", out, *flags])
    if code:
        return code, [], {}, said.getvalue()
    wb = load_workbook(out)
    return (code, list(wb["Transactions"].iter_rows(min_row=2, values_only=True)),
            {r[0]: r[1] for r in wb["Checks"].iter_rows(values_only=True)}, said.getvalue())


def make_canada(path, cibc=False):
    """Two Canadian chequing layouts, drawn after the sample statements RBC and CIBC publish, with made-up
    rows. The date only on the first row of each day and the balance only on its last row; interest
    printed as .15; cheque numbers in the description just left of the amount (Cheque #31 148.11, Cheque
    105 250.00); a form version (1.0) and a page footer (1 of 1) under the description column. With
    cibc, dates like Dec 30 under "For Dec 1 to Dec 31, 2022", an "as of April 2, 2024" above it in the
    header, and the balance carried forward at the top of the table instead of an opening balance."""
    doc = pymupdf.open()
    page = doc.new_page()

    def put(x, y, text, right=False):
        page.insert_text((x - pymupdf.get_text_length(text, fontsize=9) if right else x, y), text, fontsize=9)

    if cibc:
        opening, year = 33781.82, 2022
        days = [("Dec 30", [("ACCOUNT FEE", -4.00), ("BALANCE FEE WAIVER", 4.00), ("PAPER STMNT FEE", -3.50)]),
                ("Dec 31", [("E-TRANSFER FROM SAVINGS", 250.00)])]
        put(40, 40, "The names shown are based on our current records, as of April 2, 2024.")
        put(40, 52, "Account Statement")
        put(40, 64, "For Dec 1 to Dec 31, 2022")
    else:
        opening, year = 4247.14, 2026
        days = [("15 Mar", [("Transfer from savings", 85.00), ("Interest paid", 0.15), ("ATM withdrawal", -100.00),
                            ("Interac purchase - 1361 - Corner Grocer", -47.82)]),
                ("17 Mar", [("Overdraft interest", -0.93)]),
                ("18 Mar", [("Cheque #30", -40.00), ("Cheque #31", -148.11), ("Cheque 105", -250.00)]),
                ("22 Mar", [("ATM withdrawal", -20.00), ("Interac purchase - 1361 - Hardware Store", -125.13)])]
        put(40, 40, "Your personal chequing account statement")
        put(40, 52, "From March 12, 2026 to April 12, 2026")
        put(40, 80, f"Your opening balance on March 12, 2026 ${opening:,.2f}")
    for x, word in [(40, "Date"), (95, "Description")]:
        put(x, 120, word)
    for x, word in [(380, "Withdrawals ($)"), (460, "Deposits ($)"), (545, "Balance ($)")]:
        put(x, 120, word, right=True)
    y, bal, truth = 134, opening, []
    put(95, y, "Balance carried forward" if cibc else "Opening balance")
    put(545, y, f"{opening:,.2f}", right=True)
    for day, rows in days:
        date = dt.datetime.strptime(f"{day} {year}", "%b %d %Y" if cibc else "%d %b %Y").date()
        for k, (desc, amount) in enumerate(rows):
            y += 13
            bal = round(bal + amount, 2)
            truth.append((date, amount, bal if k == len(rows) - 1 else None, desc))
            if k == 0:
                put(40, y, day)
            put(95, y, desc)
            put(380 if amount < 0 else 460, y, ".15" if desc == "Interest paid" else f"{abs(amount):,.2f}", right=True)
            if k == len(rows) - 1:
                put(545, y, f"{bal:,.2f}", right=True)
            if desc == "PAPER STMNT FEE":
                y += 13
                put(95, y, "PER STATEMENT")
    put(300, y + 13, "1.0")
    put(95, y + 26, "Closing balance")
    put(545, y + 26, f"{bal:,.2f}", right=True)
    put(280, 800, "1 of 1")
    doc.save(path)
    return truth, opening, bal


def check_canada():
    """The two Canadian layouts: every row with its date, amount and end of day balance, nothing from the
    margin or the footer in a description, the right year, and the opening and closing balances agree.
    Before, .15 was not an amount, Cheque 105 250.00 read as 105,250.00 and Dec 30 went to 2024."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "ca.pdf"), os.path.join(tmp, "ca.xlsx")
        for cibc in (False, True):
            truth, opening, closing = make_canada(pdf, cibc)
            code, rows, checks, said = read(pdf, out)
            got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4], r[1]) for r in rows]
            want = [(date, amount, bal, desc + (" PER STATEMENT" if desc == "PAPER STMNT FEE" else ""))
                    for date, amount, bal, desc in truth]
            assert code == 0 and got == want, (cibc, [(g, w) for g, w in zip(got, want) if g != w][:3], len(got))
            assert (checks["Opening balance"], checks["Closing balance"]) == (opening, closing), checks
            assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
            assert checks["Balance mismatches"] == 0, checks
    print("Canadian layouts: interest of .15, cheque numbers next to the amount, end of day balances, Dec 30 "
          "in 2022 despite a date of 2024 in the header, a balance carried forward at the top taken as the "
          "opening balance, and the form version and the footer kept out of the descriptions")


def make_us_checks(path, checks_first):
    """A US statement laid out as the sample TD Bank publishes, with made-up rows: a summary with two-digit
    years (10/14/18), a table of cheques paid (number, date and amount, two to a line, a * after a number
    where the sequence skips), the other transactions with dates without a year, and a table of balances
    by date, which must stay out. checks_first puts the cheques above the other transactions, as TD does;
    else they come below, as on a Chase statement."""
    doc = pymupdf.open()
    page = doc.new_page()

    def put(x, y, text, right=False):
        page.insert_text((x - pymupdf.get_text_length(text, fontsize=9) if right else x, y), text, fontsize=9)

    opening = 300.00
    rows = [("10/17", "POS Debit 10/15", -63.00, "Sneaker World"), ("10/21", "Deposit", 43.25, None),
            ("10/25", "Withdrawal 10/25", -20.00, None), ("11/05", "Deposit", 43.25, None),
            ("11/12", "Deposit", 125.00, None)]
    cheques = [("101", "10/25", 17.75), ("102", "10/25", 9.45), ("104*", "11/06", 18.00)]
    closing = round(opening + sum(r[2] for r in rows) - sum(c[2] for c in cheques), 2)
    put(40, 40, "SAMPLE CHECKING ACCOUNT STATEMENT")
    put(40, 52, "Statement Date 11/15/18")
    for y, label, value in [(70, "Previous Statement Balance as of 10/14/18", opening),
                            (82, "Current Statement Balance as of 11/15/18", closing)]:
        put(40, y, label + " " + "." * 30)
        put(420, y, f"{value:,.2f}", right=True)
    truth, state = [], {"y": 110}

    def cheque_table():
        y = state["y"]
        put(40, y, "***Checking Transactions***")
        for x, word in [(40, "Serial"), (90, "Date"), (150, "Amount"), (240, "Serial"), (290, "Date"), (350, "Amount")]:
            put(x, y + 13, word)
        for k, (number, date, amount) in enumerate(cheques):
            row_y, left = y + 26 + 13 * (k // 2), 0 if k % 2 == 0 else 200
            put(40 + left, row_y, number)
            put(90 + left, row_y, date)
            put(185 + left, row_y, f"{amount:,.2f}", right=True)    # the second one next to the Debits column
            truth.append((date, -amount, "Check " + number.rstrip("*")))
        y += 26 + 13 * ((len(cheques) + 1) // 2)
        put(40, y, "Total Checks Paid")
        put(420, y, f"{sum(c[2] for c in cheques):,.2f}", right=True)
        state["y"] = y + 26

    def other_table():
        y = state["y"]
        put(40, y, "***Checking Account Transactions***")
        for x, word in [(40, "Date"), (90, "Description")]:
            put(x, y + 13, word)
        for x, word in [(420, "Debits"), (500, "Credits")]:
            put(x, y + 13, word, right=True)
        y += 26
        for date, desc, amount, wrapped in rows:
            put(40, y, date)
            put(90, y, desc)
            put(420 if amount < 0 else 500, y, f"{abs(amount):,.2f}", right=True)
            truth.append((date, amount, desc + (" " + wrapped if wrapped else "")))
            if wrapped:
                y += 13
                put(90, y, wrapped)
            y += 13
        state["y"] = y + 13

    for table in (cheque_table, other_table) if checks_first else (other_table, cheque_table):
        table()
    y = state["y"]
    put(40, y, "***Balance by Date***")
    for k, (date, value) in enumerate([("10/15", 300.00), ("11/06", 258.30), ("10/25", 233.05), ("11/14", closing)]):
        put(40 + 400 * (k % 2), y + 13 + 13 * (k // 2), date)        # the balances under Debits and Credits
        put(420 + 80 * (k % 2), y + 13 + 13 * (k // 2), f"{value:,.2f}", right=True)
    doc.save(path)
    truth = sorted(truth, key=lambda t: (t[0], truth.index(t)))
    return [(dt.date(2018, int(d[:2]), int(d[3:])), amount, desc) for d, amount, desc in truth], opening, closing


def check_us_checks():
    """The US layout with a table of cheques, above the other transactions and below them: every cheque
    read as money out, the years from 10/14/18 in the summary, the balances by date left out, the rows
    in date order, and the opening balance plus the movements gives the closing balance. Before, the
    cheques were missing, the balances by date were read as deposits and the dates went to this year."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "us.pdf"), os.path.join(tmp, "us.xlsx")
        for checks_first in (True, False):
            truth, opening, closing = make_us_checks(pdf, checks_first)
            code, rows, checks, said = read(pdf, out)
            got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[1]) for r in rows]
            assert code == 0 and got == truth, (checks_first, [(g, w) for g, w in zip(got, truth) if g != w][:3],
                                                len(got), len(truth))
            assert (checks["Opening balance"], checks["Closing balance"]) == (opening, closing), checks
            assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
            assert checks["Balance mismatches"] == 0, checks
    print("US cheques: the table of cheques read above and below the other transactions, two to a line, as money "
          "out; years from 10/14/18; balances by date left out; rows in date order; the balances agree")


def check_merge():
    """Several statements of one account in one file: three US months sent out of order, each with its
    daily balance table and savings account at the end, come out complete and in date order, and so do
    the three joined in one PDF; a month sent twice is read once; a missing month is named; UK statements
    of December and January, whose dates have no year, get the right years. Before, only the first
    statement was read and the checks still passed."""
    assert s2e.year_for(dt.date(2000, 12, 3), [dt.date(2026, 1, 1)]) == 2025     # "1 December to 1 January 2026"
    assert s2e.year_for(dt.date(2000, 1, 3), [dt.date(2025, 1, 1), dt.date(2025, 12, 31)]) == 2025
    assert s2e.year_for(dt.date(2000, 2, 29), [dt.date(2024, 2, 1), dt.date(2024, 3, 1)]) == 2024
    assert s2e.year_for(dt.date(2000, 12, 15), [dt.date(2025, 12, 1)]) == 2025      # only the first day printed
    assert s2e.year_for(dt.date(2000, 12, 30), [dt.date(2022, 12, 1), dt.date(2022, 12, 31),
                                               dt.date(2024, 4, 2)]) == 2022        # and "as of April 2, 2024"
    assert s2e.year_for(dt.date(2000, 12, 30), [dt.date(2022, 12, 31), dt.date(2024, 4, 2)]) == 2022
    assert s2e.full_dates("Period 01/12/2025 to 31.12.2025, made 2026-01-02, due January 15, 2026", "dmy") == [
        dt.date(2025, 12, 1), dt.date(2025, 12, 31), dt.date(2026, 1, 2), dt.date(2026, 1, 15)]
    with tempfile.TemporaryDirectory() as tmp:
        months, opening = [], 1234.56
        for k, month in enumerate((1, 2, 3)):
            path = os.path.join(tmp, f"2026-{month:02d}.pdf")
            truth, start_balance, closing = make_columns(path, us=True, n=20, seed=3 + k,
                                                         start=dt.date(2026, month, 1), opening=opening)
            months.append((path, truth, start_balance, closing))
            opening = closing
        twice = os.path.join(tmp, "2026-01 (1).pdf")
        shutil.copy(months[0][0], twice)
        out = os.path.join(tmp, "all.xlsx")
        code, rows, checks, said = read([months[2][0], months[0][0], twice, months[1][0]], out)
        got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4]) for r in rows]
        assert code == 0 and got == [row for m in months for row in m[1]], said
        assert (checks["Opening balance"], checks["Closing balance"]) == (months[0][2], months[2][3]), checks
        assert checks["Opening balance plus movements gives the closing balance"] == "yes", checks
        assert checks["Balance mismatches"] == 0 and checks["Other transaction tables left out"] == 3, checks
        assert checks["File 2: 2026-02.pdf"] == f"rows 22 to 41, {months[1][1][0][0]} to {months[1][1][-1][0]}", checks
        assert checks["Each file starts at the closing balance of the one before"] == "yes", checks
        assert checks["Left out: 2026-01 (1).pdf"] == "the same transactions as 2026-01.pdf", checks
        assert "Files in date order: 2026-01.pdf (20 rows), 2026-02.pdf (20 rows), 2026-03.pdf (20 rows)" in said, said

        joined = pymupdf.open()
        for month in months:
            joined.insert_pdf(pymupdf.open(month[0]))
        joined.save(os.path.join(tmp, "2026-q1.pdf"))
        code, rows, checks, said = read(os.path.join(tmp, "2026-q1.pdf"), out)
        assert [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4]) for r in rows] == [
            row for m in months for row in m[1]], "three months joined in one PDF"
        assert checks["Balance mismatches"] == 0 and checks["Other transaction tables left out"] == 3, checks
        assert "Other transaction tables left out: 3. That is another account" in said, said

        code, rows, checks, said = read([months[0][0], months[2][0]], out)
        assert checks["Each file starts at the closing balance of the one before"] == "no: 2026-01.pdf to 2026-03.pdf"
        assert checks["Balance mismatches"] >= 1 and "2026-03.pdf does not start at the closing balance of " \
            "2026-01.pdf: is a statement missing between them?" in said, said

        december, january = os.path.join(tmp, "dec.pdf"), os.path.join(tmp, "jan.pdf")
        dec_truth, dec_opening, dec_closing = make_uk(december, n=20, start=dt.date(2025, 12, 1))
        jan_truth, _, jan_closing = make_uk(january, n=20, seed=6, start=dt.date(2026, 1, 1), opening=dec_closing)
        code, rows, checks, said = read([january, december], out)
        got = [(r[0].date(), round((r[2] or 0) - (r[3] or 0), 2), r[4]) for r in rows]
        want = dec_truth + jan_truth
        assert code == 0 and got == want, [(g, w) for g, w in zip(got, want) if g != w][:2]
        assert (checks["Opening balance"], checks["Closing balance"]) == (dec_opening, jan_closing), checks
        assert checks["Balance mismatches"] == 0 and checks["Each file starts at the closing balance of the one "
                                                            "before"] == "yes", checks
    print("several statements: three US months sent out of order come out complete and in date order, and "
          "so do the three joined in one PDF; a month sent twice is read once, a missing month is named, and "
          "UK dates without a year get 2025 and 2026 from the period the statements print")


def check_overprint():
    """Text drawn more than once on itself, as bold made by printing twice a third of a point off, or a text
    layer repeated four times, as in the sample statement a council in Australia publishes: each word is
    read once, and the result is that of the plain PDF. Before, each amount was read two or four times."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, bold, out = (os.path.join(tmp, name) for name in ("s.pdf", "bold.pdf", "s.xlsx"))
        make_columns(pdf, us=True)
        _, want, want_checks, _ = read(pdf, out)
        for shifts in ([(0.3, 0.2)], [(0, 0)] * 3):
            src, doc = pymupdf.open(pdf), pymupdf.open()
            for k, page in enumerate(src):
                new = doc.new_page(width=page.rect.width, height=page.rect.height)
                for dx, dy in [(0, 0)] + shifts:
                    new.show_pdf_page(new.rect + (dx, dy, dx, dy), src, k)
            doc.save(bold)
            code, got, checks, said = read(bold, out)
            assert code == 0 and got == want and checks == want_checks, (shifts, said, len(got), len(want))
    print("text drawn twice for bold, or a text layer repeated four times: each word read once, the same "
          "rows and checks as the plain PDF")


def check_locked():
    """A statement locked with a password, as many banks send them: without --password the console asks for
    it, a wrong one is named, the right one gives the rows of the open PDF. A broken file, or a web page
    saved as .pdf, is named instead of stopping with an error. A statement saved as a picture (JPG) is read
    like a scan."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, locked, out = (os.path.join(tmp, name) for name in ("s.pdf", "locked.pdf", "s.xlsx"))
        make(pdf, eu=True, n=20, seed=7)
        _, want, want_checks, _ = read(pdf, out)
        pymupdf.open(pdf).save(locked, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1503", owner_pw="bank")
        code, _, _, said = read(locked, out)
        assert code == 2 and "locked.pdf is locked with a password" in said and "--password" in said, said
        code, _, _, said = read(locked, out, "--password", "0000")
        assert code == 2 and "None of the passwords given opens it" in said, said
        code, got, checks, _ = read(locked, out, "--password", "0000", "--password", "1503")
        assert code == 0 and got == want and checks == want_checks, (code, len(got), len(want))
        for name, data in (("broken.pdf", bytes(range(256)) * 4), ("page.pdf", b"<html><body>Not found</body></html>")):
            with open(os.path.join(tmp, name), "wb") as f:
                f.write(data)
            code, _, _, said = read(os.path.join(tmp, name), out)
            assert code == 2 and name in said and "download it again" in said, said
        if shutil.which("tesseract"):
            picture = os.path.join(tmp, "statement.jpg")
            pymupdf.open(pdf)[0].get_pixmap(dpi=200).save(picture, jpg_quality=80)
            code, got, _, said = read(picture, out)
            assert code == 0 and [(r[0], r[2:5]) for r in got] == [(r[0], r[2:5]) for r in want], said
    print("locked statements: asked for the password, a wrong one named, the right one opens it; a broken file "
          "and a web page named; a JPG picture read like a scan" + ("" if shutil.which("tesseract") else
                                                                   " (not checked: no Tesseract)"))


def check_footer_page():
    """A bank's own last page: a background image over the whole page and "Page 3 of 3". It is not a scan,
    so the statement converts even without Tesseract, and an image alone is still a scan."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf, out = os.path.join(tmp, "s.pdf"), os.path.join(tmp, "s.xlsx")
        truth = make(pdf, eu=True, seed=7)
        doc = pymupdf.open(pdf)
        grey = io.BytesIO()
        Image.new("L", (850, 1100), 235).save(grey, "PNG")
        for words in ("Page 3 of 3", ""):
            page = doc.new_page()
            page.insert_image(page.rect, stream=grey.getvalue())
            if words:
                page.insert_text((50, 800), words, fontsize=8)
        assert not s2e.is_scan(doc[-2]) and s2e.is_scan(doc[-1])
        doc.delete_page(-1)
        doc.saveIncr()
        which = shutil.which
        shutil.which = lambda name: None
        try:
            code, rows, checks, _ = read(pdf, out)
        finally:
            shutil.which = which
        assert code == 0 and len(rows) == len(truth) and "Pages read from a scan (OCR)" not in checks, (code, checks)
    print("a last page with a background image and a footer: read as text, no OCR needed; an image alone is a scan")


def check_scans():
    """Each layout printed, scanned and read back with OCR must give what its text PDF gives: tilted up to
    2.5 degrees either way, upside down, on its side, at 150 or 200 dpi, grainy and blurred by JPEG."""
    if not shutil.which("tesseract"):
        print("scans: not checked, Tesseract is not installed (apt-get install tesseract-ocr tesseract-ocr-fra "
              "tesseract-ocr-por tesseract-ocr-spa)")
        return
    cases = [("EU, tilted 1.5 degrees", lambda p: make(p, eu=True, seed=7), {"angle": 1.5}),
             ("US withdrawals and deposits, upside down", lambda p: make_columns(p, us=True),
              {"angle": -0.7, "turn": 180}),
             ("French debit and credit, on its side and tilted 2.5 degrees", lambda p: make_columns(p, us=False),
              {"angle": 2.5, "turn": 90}),
             ("Chilean cartola without cents, tilted 1 degree", make_whole, {"angle": 1.0}, "--whole"),
             ("Chilean cartola with one amount column, upside down at 150 dpi",
              lambda p: make_whole(p, single=True), {"angle": -1.0, "turn": 180, "dpi": 150}, "--whole"),
             ("UK at 150 dpi, tilted 2.5 degrees the other way", make_uk, {"angle": -2.5, "dpi": 150})]
    with tempfile.TemporaryDirectory() as tmp:
        pdf, scanned = os.path.join(tmp, "text.pdf"), os.path.join(tmp, "scan.pdf")
        for name, build, how, *flags in cases:
            build(pdf)
            scan(pdf, scanned, **how)
            _, want, want_checks, _ = read(pdf, os.path.join(tmp, "text.xlsx"), *flags)
            code, got, checks, said = read(scanned, os.path.join(tmp, "scan.xlsx"), *flags)
            assert code == 0 and checks.pop("Pages read from a scan (OCR)") == pymupdf.open(pdf).page_count, name
            assert checks == want_checks, (name, checks, want_checks)
            assert [(r[0], r[2:5]) for r in got] == [(r[0], r[2:5]) for r in want], \
                (name, [(a, b) for a, b in zip(got, want) if (a[0], a[2:5]) != (b[0], b[2:5])][:3])
            alike = [SequenceMatcher(None, a[1] or "", b[1] or "").ratio() for a, b in zip(got, want)]
            assert min(alike) > 0.8 and sum(alike) / len(alike) > 0.97, (name, min(alike), sum(alike) / len(alike))
            assert "Pages read from a scan: " in said, said
            print(f"scan, {name}: the {len(got)} dates, amounts and balances of the text PDF, and the same checks")

        # The last statement again: forced through the OCR although it has text, then without Tesseract.
        code, got, checks, _ = read(pdf, os.path.join(tmp, "forced.xlsx"), "--ocr")
        assert code == 0 and checks["Pages read from a scan (OCR)"] == 1, checks
        assert [(r[0], r[2:5]) for r in got] == [(r[0], r[2:5]) for r in want]
        which = shutil.which
        shutil.which = lambda name: None
        try:
            code, _, _, said = read(scanned, os.path.join(tmp, "none.xlsx"))
        finally:
            shutil.which = which
        assert code == 2 and "scan.pdf page 1 is a scanned image" in said and "apt-get install" in said, said
        assert not os.path.exists(os.path.join(tmp, "none.xlsx"))
        code, _, _, said = read(scanned, os.path.join(tmp, "none.xlsx"), "--lang", "xyz")
        assert code == 2 and "Tesseract could not read the page" in said and "xyz" in said, said
        assert not os.path.exists(os.path.join(tmp, "none.xlsx"))
    print("scan: --ocr reads a text PDF as a scan; without Tesseract, or with a language it does not have, the "
          "page is named and nothing is written")


if __name__ == "__main__":
    check(eu=True)
    check(eu=False)
    check(eu=True, dots=True)
    check_columns(us=False)
    check_columns(us=True)
    check_columns(us=True, split=True)
    check_uk()
    check_whole()
    check_merge()
    check_canada()
    check_us_checks()
    check_overprint()
    check_locked()
    check_bank_header()
    check_rules()
    check_categories()
    check_imports()
    check_footer_page()
    check_scans()
    print("all good")
