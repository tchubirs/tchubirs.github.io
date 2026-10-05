"""Public web pages -> Excel, CSV and JSON. Used to fulfil the Fiverr gig "web scraping".

A small JSON file describes the job: where to start, which block is one item, which fields to take
from it and how to reach the next page. Every site is different, so a new order usually needs a new
job file and nothing else.

Usage: python3 scrape.py job.json -o out        writes out.xlsx, out.csv and out.json
       (needs: pip install beautifulsoup4 requests openpyxl; example-job.json is a working job)

    {
      "start": "https://books.toscrape.com/",
      "item": "article.product_pod",
      "fields": {"title": "h3 a@title", "price": "p.price_color", "link": "h3 a@href"},
      "numbers": ["price"],
      "next": "li.next a@href",
      "max_pages": 50,
      "delay": 1
    }

A field is a CSS selector inside the item, and "@name" at the end takes that attribute instead of
the text. Links come out as full addresses. "start" can also be a list of pages. The fields named in
"numbers" become real numbers in Excel. The run stops at the page limit, when there is no next page,
or when the site's robots.txt does not allow the page. robots.txt is read the way RFC 9309 says
(wildcards, the longest rule wins), its Crawl-delay is kept when it is longer than "delay", and nothing
is fetched from a site whose robots.txt cannot be read.

Shops and sites that load their data from a JSON address (a Shopify store's /products.json, or an
address seen in the browser's Network tab) use "json" instead of "item" (example-shopify-job.json):
  "json": "products"
        where the list of items is in the answer ("" when the answer is the list). Each field is then a
        path in one item: "title", "variants.0.price" (the first variant), "images.*.src" (all of them,
        joined with commas)
  "rows": "variants"
        one row for each element of this list in an item; the fields that start with "variants.*." are
        read from the element, the others from the item
  "page_param": "page"
        asks for page 2, 3 and so on with this address parameter, until a page brings no items
  "next": "links.next"
        with "json", the path in the answer to the next page's address
  "plain": ["description"]
        fields that hold HTML come out as plain text

Optional keys (example-detail-job.json uses "detail", example-scroll-job.json uses "render" and "scroll"):
  "detail": {"link": "link", "fields": {"upc": "table tr td"}}
        opens the page in each row's "link" field and adds these fields, read from that whole page.
        A page that fails or that robots.txt keeps out leaves them empty, and the About sheet counts it.
  "render": true
        opens every page in headless Chromium first, for sites that build their content with
        JavaScript (needs: pip install playwright)
  "wait_for": "article.item"
        with "render", waits until this selector is on the page
  "scroll": 20
        with "render", scrolls to the bottom of each listing page up to 20 times, for pages that load more
        items as you scroll, and stops as soon as a scroll brings nothing new ("scroll_wait" seconds, 5)
  "click_more": "button.load-more"
        with "render", clicks this button instead of scrolling, up to "scroll" times (20 if not set),
        until it is gone or a click brings nothing new
  "key": "link"
        rows with the same value in this list field are duplicates (by default the whole row must match)
"""
import argparse
import csv
import datetime as dt
import json
import re
import sys
import time
import urllib.parse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

AGENT = "Mozilla/5.0 (compatible; scrape.py data export)"
BREAKS = re.compile(r"(?i)<(?:br|/p|/div|/li|/h[1-6]|/tr|/td|/th)\b[^>]*>")   # where HTML text gets a space
RETRY = (429, 500, 502, 503, 504)


def pick(node, rule, base):
    """The text, or with "@attr" the attribute, of the first match of a CSS selector inside node."""
    css, _, attr = rule.rpartition("@") if "@" in rule else (rule, "", "")
    el = node.select_one(css) if css.strip() else node
    if el is None:
        return ""
    if attr:
        value = el.get(attr) or ""
        if isinstance(value, list):     # class and rel come back as a list of words
            value = " ".join(value)
        value = value.strip()
        return urllib.parse.urljoin(base, value) if value and attr in ("href", "src") else value
    return " ".join(el.get_text(" ", strip=True).split())


def number(text):
    """'£1,234.50', '1.234,50 €' or '-3' as a float; None when there is no number in it."""
    m = re.search("[-−]?\\d[\\d.,\\s]*", text or "")
    if not m:
        return None
    raw = re.sub(r"\s", "", m.group())
    negative = raw[0] in "-−"
    raw = raw.lstrip("-−").rstrip(".,")
    if "," in raw and "." in raw:          # the last one is the decimal mark
        dec = "," if raw.rfind(",") > raw.rfind(".") else "."
        raw = raw.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in raw:                        # 1,234 is thousands, 12,50 is decimals
        whole, _, tail = raw.rpartition(",")
        raw = raw.replace(",", "") if len(tail) == 3 else whole.replace(",", "") + "." + tail
    try:
        value = float(raw)
    except ValueError:
        return None
    return -value if negative else value


def columns(job):
    extra = job.get("detail", {}).get("fields", {})
    return list(job["fields"]) + [name for name in extra if name not in job["fields"]]


class Robots:
    """robots.txt read the way RFC 9309 and Google read it: the groups that name this agent, or else the *
    group; the longest rule that matches the address wins and Allow wins a tie; * in a rule matches
    anything and a $ at its end ends the address. Python's own parser takes the first rule that matches
    and no wildcards, so after "Allow: /" it lets everything through."""

    def __init__(self, text, agent, refusal=""):
        groups, group, last = [], None, None
        for raw in text.splitlines():
            key, sep, value = raw.split("#", 1)[0].partition(":")
            key, value = key.strip().lower(), value.strip()
            if not sep:
                continue
            if key == "user-agent":
                if last != "user-agent":
                    group = {"agents": set(), "rules": [], "delay": None}
                    groups.append(group)
                group["agents"].add(value.lower())
            elif key in ("allow", "disallow") and group is not None and value:
                group["rules"].append((key == "allow", value))
            elif key == "crawl-delay" and group is not None:
                try:
                    group["delay"] = float(value)
                except ValueError:
                    pass
            last = key
        words = set(re.findall(r"[a-z0-9._-]+", agent.lower()))
        mine = [g for g in groups if g["agents"] & words] or [g for g in groups if "*" in g["agents"]]
        self.rules = [(len(rule), allow, self.pattern(rule)) for g in mine for allow, rule in g["rules"]]
        self.delay = max([g["delay"] for g in mine if g["delay"] is not None] or [0])
        self.refusal = refusal       # why nothing on this site is fetched, when robots.txt could not be read

    @staticmethod
    def pattern(rule):
        end = rule.endswith("$")
        return re.compile("".join(".*" if c == "*" else re.escape(c) for c in rule[:-1 if end else None])
                          + ("$" if end else ""))

    def allowed(self, url):
        parts = urllib.parse.urlsplit(url)
        target = (parts.path or "/") + ("?" + parts.query if parts.query else "")
        if self.refusal:
            return False
        if target == "/robots.txt":
            return True
        best = max(((length, allow) for length, allow, rx in self.rules if rx.match(target)), default=(0, True))
        return best[1]


class Fetcher:
    errors = (requests.RequestException,)   # what a page that cannot be read raises

    def __init__(self, delay=1.0, agent=AGENT):
        self.delay, self.session, self.robots, self.last = delay, requests.Session(), {}, 0.0
        self.session.headers["User-Agent"] = agent

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()

    def close(self):
        self.session.close()

    def site(self, url):
        """The robots.txt of the site of this address, read once. Without one, everything is allowed. When it
        cannot be read (no answer, a server error, 429) or answers 401 or 403, nothing there is fetched."""
        root = "{0.scheme}://{0.netloc}".format(urllib.parse.urlsplit(url))
        if root not in self.robots:
            agent = self.session.headers["User-Agent"]
            try:
                r = self.session.get(root + "/robots.txt", timeout=30)
                closed = r.status_code in (401, 403, 429) or r.status_code >= 500
                self.robots[root] = Robots("" if closed or r.status_code != 200 else r.text, agent,
                                           f"{root}/robots.txt answered {r.status_code}" if closed else "")
            except requests.RequestException as e:
                self.robots[root] = Robots("", agent, f"{root}/robots.txt could not be read ({type(e).__name__})")
        return self.robots[root]

    def allowed(self, url):
        return self.site(url).allowed(url)

    def refusal(self, url):
        """Why this address is not fetched, for the log."""
        return self.site(url).refusal or f"robots.txt does not allow {url}"

    def pause(self, url):
        """Waits the job's delay between pages, or longer when robots.txt asks for it (Crawl-delay)."""
        wait = self.last + max(self.delay, self.site(url).delay) - time.time()
        if wait > 0:
            time.sleep(wait)
        self.last = time.time()

    def get(self, url, listing=False):
        for attempt in range(4):
            self.pause(url)
            r = self.session.get(url, timeout=60)
            if r.status_code not in RETRY:
                break
            time.sleep(2 ** (attempt + 1))
        r.raise_for_status()
        return r.text


class Browser(Fetcher):
    """For sites that build their pages with JavaScript: the same as Fetcher, but every page is opened in
    headless Chromium and read once it has loaded. robots.txt is still read the plain way."""

    def __init__(self, delay=1.0, agent=AGENT, wait_for=None, scroll=0, click_more=None, scroll_wait=5, item=None):
        super().__init__(delay, agent)
        from playwright.sync_api import Error, TimeoutError, sync_playwright   # only needed here
        self.errors, self.timeout, self.wait_for = (requests.RequestException, Error), TimeoutError, wait_for
        self.loads = scroll or (20 if click_more else 0)
        self.click_more, self.scroll_wait, self.item = click_more, scroll_wait, item
        self.pw = sync_playwright().start()
        self.browser = self.pw.chromium.launch()
        self.page = self.browser.new_page(user_agent=agent)

    def close(self):
        self.browser.close()
        self.pw.stop()
        super().close()

    def get(self, url, listing=False):
        for attempt in range(4):
            self.pause(url)
            r = self.page.goto(url, wait_until="load", timeout=60000)
            status = r.status if r else 200
            if status not in RETRY:
                break
            time.sleep(2 ** (attempt + 1))
        if status >= 400:
            raise requests.HTTPError(f"{status} for {url}")
        if self.wait_for:
            self.page.wait_for_selector(self.wait_for, timeout=30000)
        else:
            try:
                self.page.wait_for_load_state("networkidle", timeout=10000)
            except self.timeout:
                pass    # some sites never go quiet; what has loaded by now is used
        if listing:
            self.load_more()
        return self.page.content()

    def load_more(self):
        """Scroll to the bottom, or click the "more" button, until no new item comes or the limit."""
        for _ in range(self.loads):
            # New items are the sign. The page height is not: a "loading" line makes it grow before anything
            # comes, and it stays put while the page is shorter than the window.
            before = [self.item, self.page.locator(self.item).count() if self.item else 0,
                      self.page.evaluate("document.body.scrollHeight")]
            if self.click_more:
                button = self.page.locator(self.click_more).first
                if not button.count() or not button.is_visible():
                    break
                button.click()
            else:
                self.page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            try:
                self.page.wait_for_function(
                    "([item, n, h]) => item ? document.querySelectorAll(item).length > n"
                    " : document.body.scrollHeight > h", arg=before, timeout=self.scroll_wait * 1000)
            except self.timeout:
                break   # nothing new came


def check_job(job):
    """The mistakes a job file can have, told before anything is fetched."""
    if "start" not in job or not job.get("fields"):
        raise ValueError('a job needs "start" and "fields"')
    if "json" not in job and not job.get("item"):
        raise ValueError('a job on HTML pages needs "item", the CSS selector of one item')
    unknown = [name for key in ("numbers", "plain") for name in job.get(key, []) if name not in columns(job)]
    if unknown:
        raise ValueError(f"numbers or plain name a field the job does not have: {', '.join(unknown)}")


def scrape(job, fetcher=None, log=print):
    check_job(job)
    if fetcher:
        return collect(job, fetcher, log)
    settings = job.get("delay", 1), job.get("user_agent", AGENT)
    more = job.get("wait_for"), job.get("scroll", 0), job.get("click_more"), job.get("scroll_wait", 5), job.get("item")
    with Browser(*settings, *more) if job.get("render") and "json" not in job else Fetcher(*settings) as fetcher:
        return collect(job, fetcher, log)


def found(data, rule):
    """The values at a dotted path in JSON, and whether the path can give more than one: "variants.0.price"
    is the price of the first variant, "*" takes every element of a list and "" is the whole document."""
    values, many = [data], False
    for key in rule.split(".") if rule else []:
        many = many or key == "*"
        step = []
        for v in values:
            if isinstance(v, list) and key == "*":
                step += v
            elif isinstance(v, list) and re.fullmatch(r"-?\d+", key) and -len(v) <= int(key) < len(v):
                step.append(v[int(key)])
            elif isinstance(v, dict) and key in v:
                step.append(v[key])
        values = step
    return values, many


def cell(data, rule):
    """One field from JSON: a list becomes its values joined with commas, an object stays JSON text."""
    values, many = found(data, rule)
    flat = [x for v in values for x in (v if isinstance(v, list) else [v])]
    if not many and len(values) == 1 and not isinstance(values[0], list):
        flat = values
    flat = [json.dumps(x, ensure_ascii=False) if isinstance(x, dict) else x for x in flat]
    if len(flat) == 1 and not isinstance(flat[0], str):
        return flat[0]                  # a number or true/false keeps its type
    return ", ".join(str(x) for x in flat if x not in (None, ""))


def json_rows(item, job):
    """The rows of one JSON item: one, or with "rows" one for each element of that list, where the fields
    that start with that list's path and ".*." are read from the element."""
    if not job.get("rows"):
        return [{name: cell(item, rule) for name, rule in job["fields"].items()}]
    inner = job["rows"] + ".*."
    children = found(item, job["rows"] + ".*")[0] or [None]     # an item without any still gives a row
    return [{name: (cell(child, rule[len(inner):]) if child is not None else "") if rule.startswith(inner)
             else cell(item, rule) for name, rule in job["fields"].items()} for child in children]


def page_after(url, param):
    """The same address with the page number one higher (no number is page 1)."""
    parts = urllib.parse.urlsplit(url)
    query = urllib.parse.parse_qsl(parts.query, keep_blank_values=True)
    number = int(dict(query).get(param) or 1)
    query = [(k, v) for k, v in query if k != param] + [(param, str(number + 1))]
    return urllib.parse.urlunsplit(parts._replace(query=urllib.parse.urlencode(query)))


def collect(job, fetcher, log):
    queue = [job["start"]] if isinstance(job["start"], str) else list(job["start"])
    seen, rows, pages, told = set(), [], 0, set()
    while queue and pages < job.get("max_pages", 50):
        url = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        if not fetcher.allowed(url):
            log(f"{fetcher.refusal(url)}, stopping there")
            break
        site = fetcher.site(url)
        if site.delay > fetcher.delay and site not in told:
            told.add(site)          # it changes how long the order takes
            log(f"robots.txt asks for {site.delay:g} seconds between pages, so the run waits that long")
        try:
            body = fetcher.get(url, listing=True)
        except fetcher.errors as e:
            log(f"could not read {url} ({e}), stopping there; the rows read so far are kept")
            break
        pages += 1
        if "json" in job:
            try:
                data = json.loads(body)
            except ValueError:
                log(f"{url} did not answer with JSON, stopping there")
                break
            items = found(data, job["json"])[0]
            items = items[0] if len(items) == 1 and isinstance(items[0], list) else items
            for item in items:
                rows += json_rows(item, job)
            if job.get("page_param"):
                nxt = page_after(url, job["page_param"]) if items else ""
            else:
                nxt = cell(data, job["next"]) if job.get("next") else ""
                nxt = urllib.parse.urljoin(url, nxt) if isinstance(nxt, str) and nxt else ""
        else:
            soup = BeautifulSoup(body, "html.parser")
            for node in soup.select(job["item"]):
                rows.append({name: pick(node, rule, url) for name, rule in job["fields"].items()})
            nxt = pick(soup, job["next"], url) if job.get("next") else ""
        if nxt and nxt not in seen:
            queue.insert(0, nxt)
    key = job.get("key")
    unique, keys = [], set()
    for row in rows:
        k = row[key] if key else tuple(row.values())
        if k not in keys:
            keys.add(k)
            unique.append(row)
    stats = {"pages": pages, "found": len(rows), "duplicates": len(rows) - len(unique)}
    if job.get("detail"):
        stats.update(details(job["detail"], unique, fetcher, log))
    for row in unique:
        for name in job.get("plain", []):
            text = str(row[name] or "")
            row[name] = " ".join(BeautifulSoup(BREAKS.sub(" ", text), "html.parser").get_text().split()) \
                if "<" in text else text
        for name in job.get("numbers", []):
            value = row[name]
            row[name] = value if isinstance(value, (int, float)) and not isinstance(value, bool) else \
                number(str(value or ""))
    return unique, stats


def details(detail, rows, fetcher, log):
    """Opens the page in each row's link field and adds the detail fields, read from that whole page."""
    read = 0
    for i, row in enumerate(rows, start=1):
        url, soup = row[detail["link"]], None
        if url and not fetcher.allowed(url):
            log(f"{fetcher.refusal(url)}, its detail fields stay empty")
        elif url:
            try:
                soup = BeautifulSoup(fetcher.get(url), "html.parser")
                read += 1
            except fetcher.errors as e:
                log(f"could not read {url} ({e}), its detail fields stay empty")
        for name, rule in detail["fields"].items():
            row[name] = pick(soup, rule, url) if soup else ""
        if i % 100 == 0:
            log(f"detail pages: {i} of {len(rows)}")
    return {"details": read, "details_missing": len(rows) - read}


def write(rows, stats, job, out):
    names = columns(job)
    missing = {c: sum(1 for r in rows if r[c] in ("", None)) for c in names}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(names)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F6F78")
    for row in rows:
        ws.append([row[c] for c in names])
    for i, name in enumerate(names, start=1):
        longest = max([len(str(name))] + [len(str(r[name] or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
        if name in job.get("numbers", []):
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{ws.max_row}"

    about = wb.create_sheet("About")
    start = job["start"] if isinstance(job["start"], str) else job["start"][0]
    lines = [("Source", start), ("Extracted", dt.date.today().isoformat()), ("Pages read", stats["pages"]),
             ("Rows", len(rows)), ("Duplicate rows removed", stats["duplicates"])]
    if "details" in stats:
        lines += [("Detail pages read", stats["details"]), ("Detail pages not read", stats["details_missing"])]
    for line in lines + [(f"Empty values in {c}", n) for c, n in missing.items()]:
        about.append(line)
    about.column_dimensions["A"].width = 32
    about.column_dimensions["B"].width = 60
    for (cell,) in about.iter_rows(max_col=1):
        cell.font = Font(bold=True)
    wb.save(f"{out}.xlsx")

    with open(f"{out}.csv", "w", newline="", encoding="utf-8-sig") as f:   # the BOM makes Excel read UTF-8
        w = csv.DictWriter(f, fieldnames=names)
        w.writeheader()
        w.writerows(rows)
    with open(f"{out}.json", "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
    return missing


def main(argv=None):
    ap = argparse.ArgumentParser(description="Public web pages to Excel, CSV and JSON.")
    ap.add_argument("job", help="the job file (JSON)")
    ap.add_argument("-o", "--out", default="data", help="output name without extension")
    a = ap.parse_args(argv)
    with open(a.job, encoding="utf-8") as f:
        job = json.load(f)
    try:
        rows, stats = scrape(job)
    except ValueError as e:
        print("job file:", e)
        return 2
    missing = write(rows, stats, job, a.out)
    dup = stats["duplicates"]
    print(f"{len(rows)} rows from {stats['pages']} pages, {dup} duplicate{'' if dup == 1 else 's'} removed")
    if "details" in stats:
        print(f"  {stats['details']} detail pages read, {stats['details_missing']} not read")
    for name, n in missing.items():
        if n:
            print(f"  {n} empty values in {name}")
    print(f"wrote {a.out}.xlsx, {a.out}.csv and {a.out}.json")
    return 0 if rows else 1


if __name__ == "__main__":
    sys.exit(main())
