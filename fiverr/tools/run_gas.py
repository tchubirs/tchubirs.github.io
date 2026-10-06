"""Run Google Apps Script code on a copy of a spreadsheet, here, and say what it did. Used for the Fiverr gig
"Google Sheets and Excel": a script is tested on the client's own data before delivery.

Usage: python3 run_gas.py book.xlsx Code.gs [more.gs] --run sendReport     a function of the script
       python3 run_gas.py book.xlsx Code.gs --edit "Tasks!C2=Done"         a cell typed in, then onEdit(e)
       python3 run_gas.py book.xlsx Code.gs --open                         onOpen, and the menus it adds
              [--select "Tasks!A2:D9"] [--answer yes|no] [--input TEXT] [--fetch URL=saved-answer.json]
              [--property KEY=VALUE]
              [--timezone Europe/Paris] [--timeout 60] [-o after.xlsx] [--report report.txt]
       (needs Node.js, LibreOffice and pip install openpyxl)

A Google Sheet is tested from its copy as Excel (File > Download > Microsoft Excel). LibreOffice works out
its formulas first, so the script reads the values Google would show (functions only Google has, such as
QUERY or IMPORTRANGE, give an error there; audit_sheet.py names them). gas_sandbox.js runs the script in
Node with SpreadsheetApp and the other services scripts use most, held in memory: emails are listed and
not sent, alerts and prompts get --answer and --input, triggers are listed, UrlFetchApp reads saved
answers given with --fetch, and properties start from --property. A service this test does not have
(DriveApp, CalendarApp and so on) stops the script with its name, and the report lists it. When the
script stops on an error, the report gives the file, the line and the error. Last comes every change in
the workbook, recalculated (audit_sheet.py --compare), with the rows and columns hidden or shown;
background color, bold, font color and number format are kept in the file saved with -o, other
formatting is listed as not kept.

A script run from a menu works on what the client has open: getActiveSheet and getActiveRange give the
sheet and the cells selected. The run starts on the sheet open in the file, with A1 selected, or with
the cells --select gives; an edit selects the edited cell, as in Google. The report says when the result
depends on it. Hidden rows and columns and frozen rows come from the file.

A macro recorded in Google Sheets (Extensions > Macros > Record macro) runs as it is: activate, the current
cell, range lists, R1C1 formulas, fill down, filters, pastes. Formulas pasted, filled or sorted move their
references as in Sheets. A formula the script writes has no result here (LibreOffice works it out at the
end), so a sort by such formulas, or a paste of their values, is listed to check in Google.
"""
import argparse
import datetime as dt
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_sheet  # noqa: E402

SANDBOX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "gas_sandbox.js")
EPOCH = dt.datetime(1899, 12, 30)       # where Google Sheets puts a time of day without a date


class Unrunnable(Exception):
    """Something that keeps the script from being run at all."""


def to_json(value):
    if isinstance(value, dt.datetime):
        if value.tzinfo:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return {"$date": value.isoformat() + "Z"}
    if isinstance(value, dt.date):
        return {"$date": dt.datetime(value.year, value.month, value.day).isoformat() + "Z"}
    if isinstance(value, dt.time):
        return {"$date": EPOCH.replace(hour=value.hour, minute=value.minute, second=value.second).isoformat() + "Z"}
    return "" if value is None else value


def from_json(value):
    if isinstance(value, dict) and "$date" in value:
        return dt.datetime.fromisoformat(value["$date"].replace("Z", "+00:00")).replace(tzinfo=None)
    return None if value == "" else value


def sheets_of(base):
    """Each sheet as the script sees it: values (formulas already worked out) and formulas."""
    formulas, values = load_workbook(base), load_workbook(base, data_only=True)
    out = []
    for ws in formulas.worksheets:
        vs = values[ws.title]
        grid_v, grid_f = [], []
        for row in ws.iter_rows():
            line_v, line_f = [], []
            for cell in row:
                f = cell.value
                f = getattr(f, "text", f)               # an array formula
                is_formula = isinstance(f, str) and f.startswith("=")
                line_f.append(f if is_formula else "")
                line_v.append(to_json(vs[cell.coordinate].value))
            grid_v.append(line_v)
            grid_f.append(line_f)
        rows, cols = audit_sheet.hidden_of(ws)
        out.append({"name": ws.title, "values": grid_v, "formulas": grid_f, "hiddenRows": sorted(rows),
                    "hiddenCols": sorted(cols), "frozen": frozen_of(ws)})
    return out, formulas.active.title


def frozen_of(ws):
    """The rows and columns frozen at the top and left of a sheet, as [rows, columns]."""
    if not ws.freeze_panes:
        return [0, 0]
    col, row = coordinate_from_string(ws.freeze_panes)
    return [row - 1, column_index_from_string(col) - 1]


def show_hide(ws, rows, cols):
    """Hides these rows and columns of a sheet (numbers) and shows the others that were hidden."""
    old_rows, old_cols = audit_sheet.hidden_of(ws)
    for r in rows ^ old_rows:
        ws.row_dimensions[r].hidden = r in rows
    for c in cols ^ old_cols:
        # A run of columns alike is one entry in the file (C to E, or every column after the last used one):
        # cut it around the column, so that the column is hidden or shown alone.
        for key, d in list(ws.column_dimensions.items()):
            first, last = d.min or column_index_from_string(key), d.max or column_index_from_string(key)
            if first <= c <= last and first != last:
                del ws.column_dimensions[key]
                for a, b in ((first, c - 1), (c, c), (c + 1, last)):
                    if a <= b:
                        part = copy.copy(d)
                        part.index, part.min, part.max = get_column_letter(a), a, b
                        ws.column_dimensions[part.index] = part
                break
        ws.column_dimensions[get_column_letter(c)].hidden = c in cols


def color(text):
    """#1f6f78 or #fff as an openpyxl color; None for a color name, which is left as it was."""
    m = re.fullmatch(r"#?([0-9a-fA-F]{6}|[0-9a-fA-F]{3})", (text or "").strip())
    if not m:
        return None
    code = m.group(1) if len(m.group(1)) == 6 else "".join(ch * 2 for ch in m.group(1))
    return "FF" + code.upper()


def apply(base, result, out):
    """The workbook as the script left it: its changes written into a copy of the file."""
    wb = load_workbook(base)
    kept = {s["origin"] for s in result["sheets"] if s["origin"]}
    for name in list(wb.sheetnames):
        if name not in kept:
            del wb[name]
    ordered = []
    for s in result["sheets"]:
        ws = wb[s["origin"]] if s["origin"] else wb.create_sheet(s["name"])
        if ws.title != s["name"]:
            ws.title = s["name"]
        for r, c, value, formula in s["changes"]:
            ws.cell(r, c).value = formula or from_json(value)
        for key, fmt in s["formats"].items():
            cell = ws.cell(*map(int, key.split(",")))
            if fmt.get("cleared"):
                cell.style = "Normal"
            if color(fmt.get("background")):
                cell.fill = PatternFill("solid", start_color=color(fmt["background"]))
            if "bold" in fmt or color(fmt.get("color")):
                font = cell.font.copy(bold=fmt.get("bold", cell.font.b))
                cell.font = font.copy(color=color(fmt["color"])) if color(fmt.get("color")) else font
            if fmt.get("numberFormat"):
                cell.number_format = fmt["numberFormat"]
        show_hide(ws, set(s["hiddenRows"]), set(s["hiddenCols"]))
        if s["frozen"] != frozen_of(ws):
            rows, cols = s["frozen"]
            ws.freeze_panes = f"{get_column_letter(cols + 1)}{rows + 1}" if rows or cols else None
        ws.sheet_state = "hidden" if s["hidden"] else "visible"
        ordered.append(ws)
    wb._sheets = ordered
    wb.active = next(i for i, ws in enumerate(ordered) if ws.title == result["active"])
    wb.save(out)


def typed(text):
    """A value typed into a cell, as Google Sheets reads it: a number, TRUE or FALSE, or text."""
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d*\.\d+", text):
        return float(text)
    return {"TRUE": True, "FALSE": False}.get(text.upper(), text)


def run(book, scripts, call=None, edit=None, opened=False, answer="yes", typed_in="", fetch=None,
        properties=None, timezone="Etc/GMT", timeout=60, keep=None, changes=True, select=None):
    """Run the script on a copy and return what happened, as a dict the report is written from. select is
    {"sheet": name or None for the open one, "cells": "A2:D9"}, the cells selected when it starts."""
    if not shutil.which("node"):
        raise Unrunnable("Node.js is not installed (the node command)")
    if call and not re.fullmatch(r"[A-Za-z_$][\w$]*", call):
        raise Unrunnable(f"{call} is not the name of a function")
    work = tempfile.mkdtemp(prefix="run-gas-")
    try:
        base = audit_sheet.recalculate([book], work)[0]
        if not base:
            raise Unrunnable(f"LibreOffice could not open {os.path.basename(book)}")
        sheets, active = sheets_of(base)
        if select and select.get("sheet"):
            names = [s["name"] for s in sheets]
            title = next((n for n in names if n.lower() == select["sheet"].lower()), None)
            if title is None:
                raise Unrunnable(f"There is no sheet {select['sheet']} in the workbook. Sheets: {', '.join(names)}")
            select = dict(select, sheet=title)
        job = {"sheets": sheets, "active": active, "timezone": timezone, "user": "client@example.com",
               "answer": answer, "input": typed_in, "fetch": fetch or {}, "properties": properties or {},
               "timeout": timeout, "run": call, "edit": edit, "open": opened, "select": select,
               "title": pathlib.Path(book).stem}
        job_path, result_path = os.path.join(work, "job.json"), os.path.join(work, "result.json")
        pathlib.Path(job_path).write_text(json.dumps(job), encoding="utf-8")
        try:
            done = subprocess.run(["node", SANDBOX, job_path, result_path, *map(os.path.abspath, scripts)],
                                  capture_output=True, text=True, timeout=timeout + 60, env=dict(os.environ, TZ="UTC"))
        except subprocess.TimeoutExpired:
            return {"status": "timeout", "seconds": timeout, "what": what(call, edit, opened)}
        if not os.path.exists(result_path):
            raise Unrunnable("the test itself failed: " + (done.stderr.strip().splitlines() or ["no message"])[-1])
        result = json.loads(pathlib.Path(result_path).read_text(encoding="utf-8"))
        result["what"] = what(call, edit, opened)
        result["selected"] = bool(select)
        result["sources"] = {os.path.basename(s): pathlib.Path(s).read_text(encoding="utf-8") for s in scripts}
        if result["status"] == "absent":
            raise Unrunnable(f"{result['what']} is not in the script. Its functions: "
                             f"{', '.join(sorted(result['functions'])) or 'none'}")
        if changes:
            before, after = os.path.join(work, "before.xlsx"), os.path.join(work, "after.xlsx")
            load_workbook(base).save(before)    # written the same way as after, so only the script's changes show
            apply(base, result, after)
            result["changes"] = audit_sheet.compare(before, after)
            result["layout"] = audit_sheet.layout(before, after)
            if keep:
                shutil.copyfile(after, keep)
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


def what(call, edit, opened):
    if edit:
        return f"onEdit after {edit['sheet']}!{edit['cell']} = {edit['value']!r}"
    return "onOpen" if opened else call


def report(result, answer):
    """The report as lines of text."""
    lines, status, name = [], result["status"], result["what"]
    if status == "ok":
        lines.append(f"{name} ran to the end in {result['seconds']:g} s, with no error.")
    elif status == "timeout":
        lines.append(f"{name} did not finish in {result['seconds']:g} s: a loop that never ends, or it is very slow.")
    else:
        error, place = result["error"], result["error"].get("at")
        where = f" at {place['file']} line {place['line']}" if place else ""
        lines.append(f"{name} stopped{where} with {error['name']}: {error['message']}")
        if place:
            code = result["sources"].get(place["file"], "").splitlines()
            if 0 < place["line"] <= len(code):
                lines.append(f"  {place['line']}: {code[place['line'] - 1].strip()}")
    start, uses = result.get("start") or {}, result.get("uses") or {}
    picked, edited = f"{start.get('sheet')}!{start.get('cells')}", name.startswith("onEdit")
    if result.get("selected"):
        lines.append(f"Started with {picked} selected, as asked with --select.")
    elif uses.get("selection") and not edited:
        lines.append(f"It works on the selected cells: the run had {picked} selected, so give the cells the client "
                     "selects with --select Sheet!A2:D9.")
    elif uses.get("sheet") and start.get("sheets", 1) > 1 and not edited:
        lines.append(f"It works on the open sheet: the run started on {start['sheet']}, the one open in the file; "
                     "--select Sheet!A1 starts it on another.")
    if result.get("emails"):
        lines.append("Emails it would send (none was sent):")
        lines += [f"  to {m['to']}: {m['subject']!r}, {len(m['body'])} characters"
                  + (f", {m['attachments']} attachments" if m["attachments"] else "") for m in result["emails"]]
    if result.get("ui"):
        lines.append(f"Messages and questions, answered {answer}:")
        lines += [f"  {line}" for line in result["ui"]]
    for key, title in (("menus", "Menus it adds:"), ("triggers", "Triggers it would create:"),
                       ("fetched", "Addresses it fetched (from saved answers):"),
                       ("missing", "Not in this test, check these parts in Google:")):
        if result.get(key):
            lines.append(title)
            lines += [f"  {line}" for line in result[key]]
    if result.get("log"):
        lines.append("Log:")
        lines += [f"  {line}" for line in result["log"][:20]]
        if len(result["log"]) > 20:
            lines.append(f"  and {len(result['log']) - 20} more lines")
    if "changes" in result:
        lines.append("What changed in the workbook:" if status == "ok" else "What changed before it stopped:")
        lines += audit_sheet.describe_changes(result["changes"])[1:]
        lines += [f"  {line}" for line in result.get("layout", ())]
        if result.get("notApplied"):
            lines.append(f"  Not kept in the saved file: {', '.join(result['notApplied'])}.")
    return lines


def main(argv=None):
    ap = argparse.ArgumentParser(description="Run Apps Script code on a copy of a spreadsheet and list what it did.")
    ap.add_argument("book", help="the spreadsheet as .xlsx (a Google Sheet: File > Download > Microsoft Excel)")
    ap.add_argument("scripts", nargs="+", help="the script files (.gs or .js), in the order of the Apps Script editor")
    how = ap.add_mutually_exclusive_group(required=True)
    how.add_argument("--run", help="the function to run")
    how.add_argument("--edit", help='a cell typed in, then onEdit(e) runs: "Sheet1!C2=Done"')
    how.add_argument("--open", action="store_true", help="run onOpen, as when the spreadsheet opens")
    ap.add_argument("--select", metavar="SHEET!A2:D9",
                    help="the cells selected when the script starts, as the client selects them before using a menu")
    ap.add_argument("--answer", choices=["yes", "no"], default="yes", help="how alerts and questions are answered")
    ap.add_argument("--input", default="", help="what a prompt gets")
    ap.add_argument("--fetch", action="append", default=[], metavar="URL=FILE", help="the saved answer for an address")
    ap.add_argument("--property", action="append", default=[], metavar="KEY=VALUE", help="a script property")
    ap.add_argument("--timezone", default="Etc/GMT", help="the spreadsheet's time zone, as Session gives it")
    ap.add_argument("--timeout", type=float, default=60, help="seconds before the script is stopped")
    ap.add_argument("-o", "--out", help="keep the workbook as the script left it (.xlsx)")
    ap.add_argument("--report", help="write the report to this file too")
    a = ap.parse_args(argv)
    edit = None
    if a.edit:
        m = re.fullmatch(r"(?:'([^']+)'|([^!]+))!([A-Za-z]+\d+)=(.*)", a.edit, re.DOTALL)
        if not m:
            ap.error('--edit is written as Sheet1!C2=value')
        edit = {"sheet": m.group(1) or m.group(2), "cell": m.group(3).upper(), "value": typed(m.group(4))}
    select = None
    if a.select:
        m = re.fullmatch(r"(?:(?:'([^']+)'|([^!]+))!)?([A-Za-z]+\d+(?::[A-Za-z]+\d+)?)", a.select.strip())
        if not m:
            ap.error("--select is written as Sheet1!A2:D9 or A2:D9")
        select = {"sheet": m.group(1) or m.group(2), "cells": m.group(3).upper()}
    fetch = {}
    for item in a.fetch:
        url, _, path = item.rpartition("=")           # the address can have = in it, the file name not
        fetch[url] = pathlib.Path(path).read_text(encoding="utf-8")
    properties = dict(item.partition("=")[::2] for item in a.property)
    try:
        result = run(a.book, a.scripts, a.run, edit, a.open, a.answer, a.input, fetch, properties, a.timezone,
                     a.timeout, a.out, select=select)
    except Unrunnable as e:
        print(e)
        return 2
    text = "\n".join(report(result, a.answer))
    print(text)
    if a.report:
        pathlib.Path(a.report).write_text(text + "\n", encoding="utf-8")
    return 0 if result["status"] == "ok" else 1


if __name__ == "__main__":
    sys.exit(main())
