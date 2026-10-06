"""Run Apps Script code written for Google Sheets on a test workbook and check that the report says what it did."""
import contextlib
import datetime as dt
import io
import json
import os
import shutil
import sys
import tempfile

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import run_gas  # noqa: E402

CODE = '''function onOpen() {
  SpreadsheetApp.getUi().createMenu("Tools")
    .addItem("Archive done tasks", "archiveDone")
    .addItem("Send reminders", "sendReminders")
    .addToUi();
}

function onEdit(e) {
  const sheet = e.range.getSheet();
  if (sheet.getName() !== "Tasks" || e.range.getColumn() !== 3) return;
  sheet.getRange(e.range.getRow(), 4).setValue(e.value === "Done" ? new Date(Date.UTC(2026, 9, 5)) : "");
}

function archiveDone() {
  const ss = SpreadsheetApp.getActiveSpreadsheet();
  const tasks = ss.getSheetByName("Tasks");
  const archive = ss.getSheetByName("Archive") || ss.insertSheet("Archive");
  const rows = tasks.getRange(1, 1, tasks.getLastRow(), 4).getValues();
  const header = rows.shift();
  if (archive.getLastRow() === 0) archive.appendRow(header);
  for (let i = rows.length - 1; i >= 0; i--) {
    if (rows[i][2] === "Done") {
      archive.appendRow(rows[i]);
      tasks.deleteRow(i + 2);
    }
  }
  archive.getRange("A1:D1").setBackground("#1f6f78").setFontWeight("bold").setBorder(true, true, true, true, false, false);
  Logger.log("Archived %s rows", archive.getLastRow() - 1);
}

function sendReminders() {
  const ui = SpreadsheetApp.getUi();
  if (ui.alert("Send the reminders?", ui.ButtonSet.YES_NO) !== ui.Button.YES) return;
  const rows = SpreadsheetApp.getActive().getSheetByName("Tasks").getDataRange().getValues().slice(1);
  rows.filter(r => r[2] !== "Done" && r[1]).forEach(r => MailApp.sendEmail(r[1], "Reminder: " + r[0], "Please finish " + r[0]));
}

function askOwner() {
  const answer = SpreadsheetApp.getUi().prompt("Who owns the new task?");
  SpreadsheetApp.getActive().getSheetByName("Tasks").appendRow(["New task", answer.getResponseText(), "Open"]);
}

function openCount() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Tasks");
  sheet.getRange("H1").setValue("Open tasks: " + sheet.getRange("G1").getValue());
}

function broken() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Tasks");
  sheet.getRange("A1:B2").setValues([[1, 2, 3]]);
}

function useDrive() {
  const files = DriveApp.getFilesByName("report.pdf");
}

function loop() {
  while (true) {}
}

function importRates() {
  const key = PropertiesService.getScriptProperties().getProperty("API_KEY");
  const data = JSON.parse(UrlFetchApp.fetch("https://api.example.com/rates?key=" + key).getContentText());
  SpreadsheetApp.getActive().getSheetByName("Tasks").getRange("J1:K1").setValues([["EUR", data.EUR]]);
}

function schedule() {
  ScriptApp.newTrigger("sendReminders").timeBased().everyHours(6).create();
}

function dueSoon() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Tasks");
  const values = sheet.getRange(2, 1, sheet.getLastRow() - 1, 5).getValues();
  const today = new Date(Date.UTC(2026, 9, 5));
  values.forEach((row, i) => {
    if (row[4] instanceof Date && (row[4] - today) / 86400000 <= 3) {
      sheet.getRange(i + 2, 6).setValue("Due " + Utilities.formatDate(row[4], Session.getScriptTimeZone(), "dd/MM/yyyy"));
    }
  });
  try {
    sheet.getRange("ZZ").getValue();
  } catch (e) {
    Logger.log("%s %s %s", e instanceof Error, values instanceof Array, e.message);
  }
}
'''


def book(root):
    path = os.path.join(root, "tasks.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(["Task", "Owner", "Status", "Done on", "Due"])
    for row in [("Invoice March", "ana@example.com", "Done", None, dt.date(2026, 10, 6)),
                ("Call supplier", "rui@example.com", "Open", None, dt.date(2026, 10, 20)),
                ("Pay rent", "ana@example.com", "Done", None, dt.date(2026, 10, 8)),
                ("Count stock", "rui@example.com", "Open", None, None)]:
        ws.append(row)
    ws["G1"] = '=COUNTIF(C2:C100,"Open")'
    wb.save(path)
    code = os.path.join(root, "Code.gs")
    with open(code, "w", encoding="utf-8") as f:
        f.write(CODE)
    return path, code


def line_of(code):
    return next(i for i, line in enumerate(CODE.splitlines(), start=1) if line.strip() == code)


def cli(path, code, root):
    out, said = os.path.join(root, "edited.xlsx"), io.StringIO()
    with contextlib.redirect_stdout(said):
        assert run_gas.main([path, code, "--edit", "Tasks!C3=Done", "-o", out]) == 0
    report = said.getvalue()
    for line in ["onEdit after Tasks!C3 = 'Done' ran to the end", "Tasks!C3  Open  ->  Done",
                 "Tasks!D3  (empty)  ->  2026-10-05", "Tasks!G1  2  ->  1"]:
        assert line in report, (line, report)
    assert load_workbook(out)["Tasks"]["D3"].value == dt.datetime(2026, 10, 5), "the date is not a date"
    with contextlib.redirect_stdout(io.StringIO()) as said:
        assert run_gas.main([path, code, "--open"]) == 0
    assert "Tools: Archive done tasks -> archiveDone, Send reminders -> sendReminders" in said.getvalue()
    print("onEdit and onOpen: the date written as a date, the formula that counts open tasks recalculated, the "
          "menu listed")


def fetch_cli(path, code, root):
    """--fetch takes an address with a query string: the address ends at the last =."""
    answer = os.path.join(root, "rates.json")
    with open(answer, "w", encoding="utf-8") as f:
        json.dump({"EUR": 0.92}, f)
    with contextlib.redirect_stdout(io.StringIO()) as said:
        code_ = run_gas.main([path, code, "--run", "importRates", "--property", "API_KEY=k123",
                              "--fetch", f"https://api.example.com/rates?key=k123={answer}"])
    assert code_ == 0 and "https://api.example.com/rates?key=k123" in said.getvalue(), said.getvalue()
    print("--fetch: an address with ?key= and its saved answer")


def big(root):
    """A sheet of 3000 rows sorted, and getLastRow asked 300 times: quick, as in Google."""
    path = os.path.join(root, "big.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(3000):
        ws.append([f"item {i}", (i * 7919) % 3001, i % 5, "x", "y", "z", i, i * 2])
    wb.save(path)
    script = os.path.join(root, "Big.gs")
    with open(script, "w", encoding="utf-8") as f:
        f.write("function tidy() {\n  const sheet = SpreadsheetApp.getActive().getSheetByName('Data');\n"
                "  sheet.getDataRange().sort({column: 2, ascending: false});\n"
                "  for (let i = 0; i < 300; i++) sheet.getLastRow();\n}\n")
    result = run_gas.run(path, [script], "tidy", changes=False, timeout=30)
    assert result["status"] == "ok" and result["seconds"] < 10, (result["status"], result["seconds"])
    top = sorted(((i * 7919) % 3001 for i in range(3000)), reverse=True)[0]
    assert [1, 2, top, ""] in result["sheets"][0]["changes"], "the largest value is not on top"
    print(f"big sheet: 3000 rows sorted and getLastRow asked 300 times in {result['seconds']:.1f} s")


def archive(path, code, root):
    out = os.path.join(root, "archived.xlsx")
    result = run_gas.run(path, [code], "archiveDone", keep=out)
    report = "\n".join(run_gas.report(result, "yes"))
    assert result["status"] == "ok" and "Sheets added: Archive" in report and "Archived 2 rows" in report, report
    assert "Not kept in the saved file: borders." in report, report
    wb = load_workbook(out)
    def tasks(name):
        return [r[0] for r in wb[name].iter_rows(values_only=True) if r[0] is not None]
    assert tasks("Tasks") == ["Task", "Call supplier", "Count stock"], tasks("Tasks")
    assert tasks("Archive") == ["Task", "Pay rent", "Invoice March"], tasks("Archive")
    head = wb["Archive"]["A1"]
    assert head.fill.start_color.rgb == "FF1F6F78" and head.font.b, (head.fill.start_color.rgb, head.font.b)
    print("archive: 2 rows moved to a new sheet, its heading colored and bold in the file, borders listed as not kept")


def services(path, code):
    def go(call, **how):
        result = run_gas.run(path, [code], call, changes=False, **how)
        return result, "\n".join(run_gas.report(result, how.get("answer", "yes")))

    result, report = go("sendReminders")
    assert [m["to"] for m in result["emails"]] == ["rui@example.com", "rui@example.com"], report
    assert "alert: Send the reminders?" in report and "none was sent" in report, report
    result, _ = go("sendReminders", answer="no")
    assert result["emails"] == [], result["emails"]

    result, report = go("askOwner", typed_in="lea@example.com")
    assert result["status"] == "ok" and "prompt: Who owns the new task?" in report, report
    assert [6, 2, "lea@example.com", ""] in result["sheets"][0]["changes"], result["sheets"][0]["changes"]

    result, _ = go("openCount")
    assert [1, 8, "Open tasks: 2", ""] in result["sheets"][0]["changes"], result["sheets"][0]["changes"]

    result, report = go("broken")
    line = line_of("sheet.getRange(\"A1:B2\").setValues([[1, 2, 3]]);")
    assert result["status"] == "error" and result["error"]["at"] == {"file": "Code.gs", "line": line}, report
    assert "The number of rows in the data does not match the number of rows in the range. The data has 1 but the " \
           "range has 2." in report, report

    result, report = go("useDrive")
    assert result["status"] == "error" and result["missing"] == ["DriveApp.getFilesByName"], report
    assert "check these parts in Google" in report, report

    result, report = go("loop", timeout=3)
    assert result["status"] == "timeout" and "did not finish in 3 s" in report, report

    url = "https://api.example.com/rates?key=k123"
    result, report = go("importRates", fetch={url: json.dumps({"EUR": 0.92})}, properties={"API_KEY": "k123"})
    assert result["status"] == "ok" and result["fetched"] == [url], report
    assert [1, 11, 0.92, ""] in result["sheets"][0]["changes"], result["sheets"][0]["changes"]
    result, report = go("importRates", properties={"API_KEY": "k123"})
    assert result["status"] == "error" and "with no saved answer" in report, report

    result, report = go("schedule")
    assert result["triggers"] == ["sendReminders, on a timer, every 6 hours"], report

    result, report = go("dueSoon")
    assert result["log"] == ["true true Range not found: ZZ"], report
    changes = result["sheets"][0]["changes"]
    assert [2, 6, "Due 06/10/2026", ""] in changes and [4, 6, "Due 08/10/2026", ""] in changes, changes
    assert not [c for c in changes if c[0] == 3], changes

    root = os.path.dirname(path)
    broken = os.path.join(root, "Broken.gs")
    with open(broken, "w", encoding="utf-8") as f:
        f.write("function later() {\n  foo bar\n}\n")
    result = run_gas.run(path, [code, broken], "openCount", changes=False)
    assert result["status"] == "error" and result["error"]["at"] == {"file": "Broken.gs", "line": 2}, result["error"]

    check = os.path.join(root, "Check.gs")
    with open(check, "w", encoding="utf-8") as f:
        f.write('function onEdit(e) {\n  e.range.offset(0, 1).setValue(e.value === "TRUE" ? "ticked" : "not ticked");\n'
                '  e.range.offset(0, 2).setValue(e.oldValue === undefined ? "was empty" : "was " + e.oldValue);\n}\n')
    result = run_gas.run(path, [check], edit={"sheet": "Tasks", "cell": "H2", "value": True}, changes=False)
    changes = result["sheets"][0]["changes"]
    assert [2, 8, True, ""] in changes and [2, 9, "ticked", ""] in changes and [2, 10, "was empty", ""] in changes, changes

    try:
        go("nothing")
        raise AssertionError("an unknown function ran")
    except run_gas.Unrunnable as e:
        assert "nothing is not in the script" in str(e) and "archiveDone" in str(e), e
    print("services: emails listed and none sent, alert answered yes then no, a prompt, a formula's value read, "
          "Google's own error placed on its line, DriveApp named, a loop stopped, a saved answer for UrlFetchApp "
          "and a script property, a trigger, cell dates as Date, an unknown function refused")


TOOLS = """function deleteBlankRows() {
  const sheet = SpreadsheetApp.getActiveSheet();
  const range = sheet.getActiveRange();
  const values = sheet.getRange(range.getRowIndex(), 1, range.getNumRows(), sheet.getLastColumn()).getValues();
  for (let i = values.length - 1; i >= 0; i--) {
    if (values[i].every(v => v === "")) sheet.deleteRow(range.getRowIndex() + i);
  }
}

function hideDone() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Orders");
  const first = sheet.getFrozenRows() + 1;
  const status = sheet.getRange(first, 4, sheet.getLastRow() - first + 1, 1).getValues();
  status.forEach((row, i) => { if (row[0] === "Done") sheet.hideRows(first + i); });
  if (sheet.isColumnHiddenByUser(3)) sheet.showColumns(3);
  sheet.hideColumns(6, 2);
  Logger.log("frozen %s, row 2 hidden %s", sheet.getFrozenRows(), sheet.isRowHiddenByUser(2));
}

function dedupe() {
  const left = SpreadsheetApp.getActive().getSheetByName("Orders").getRange("B2:D9").removeDuplicates([2]);
  Logger.log("left %s", left.getA1Notation());
}

function findAndFix() {
  const book = SpreadsheetApp.getActive();
  const orders = book.getSheetByName("Orders");
  const all = book.createTextFinder("lisbon").findAll().map(r => r.getSheet().getName() + "!" + r.getA1Notation());
  const exact = orders.createTextFinder("Porto").matchEntireCell(true).matchCase(true).findAll().length;
  const replaced = orders.createTextFinder("Lisbon").replaceAllWith("Lisboa");
  const finder = orders.createTextFinder("done");
  const rows = [];
  let cell;
  while ((cell = finder.findNext())) rows.push(cell.getRow());
  const after = orders.createTextFinder("done").startFrom(orders.getRange("D5")).findNext().getRow();
  Logger.log("%s | %s | %s | %s | %s", all.join(" "), exact, replaced, rows.join(" "), after);
}

function finderOptions() {
  const orders = SpreadsheetApp.getActive().getSheetByName("Orders");
  orders.getRange("H4").setValue("Porto Alegre");
  const whole = orders.createTextFinder("Porto").matchEntireCell(true).findAll().length;
  const pattern = orders.createTextFinder("^B[a-z]+o$").useRegularExpression(true).findAll().length;
  const finder = orders.createTextFinder("Porto");
  finder.findNext();
  finder.findNext();
  const back = finder.findPrevious().getA1Notation();
  const one = finder.replaceWith("Oporto");
  orders.getRange("H2").setFormula("=E2*2");
  const formulas = orders.createTextFinder("E2").matchFormulaText(true).findAll().length;
  orders.getRange("H3").setValue("\u00c9vora");
  const accents = orders.createTextFinder("evora").ignoreDiacritics(true).findAll().length;
  const strict = orders.createTextFinder("evora").findAll().length;
  Logger.log("%s %s %s %s %s %s %s", whole, pattern, back, one, formulas, accents, strict);
}

function lastRows() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Orders");
  const up = sheet.getRange(sheet.getMaxRows(), 1).getNextDataCell(SpreadsheetApp.Direction.UP).getRow();
  const down = sheet.getRange("A1").getNextDataCell(SpreadsheetApp.Direction.DOWN).getRow();
  const next = sheet.getRange("A1").getNextDataCell(SpreadsheetApp.Direction.NEXT).getColumn();
  const empty = sheet.getRange("H1").getNextDataCell(SpreadsheetApp.Direction.NEXT).getColumn();
  Logger.log("%s %s %s %s", up, down, next, empty);
}

function shiftCells() {
  const sheet = SpreadsheetApp.getActive().getSheetByName("Orders");
  sheet.getRange("A9").setBackground("#ff0000");
  sheet.getRange("E1").setBackground("#00ff00");
  sheet.deleteRow(4);
  sheet.insertColumnBefore(2);
}

function countRows() {
  const sheet = SpreadsheetApp.getActiveSheet();
  Logger.log("%s has %s rows", sheet.getName(), sheet.getLastRow());
}

function onEdit() {
  const cell = SpreadsheetApp.getActiveSheet().getActiveCell();
  if (cell.getColumn() === 4) cell.offset(0, 4).setValue("seen " + cell.getA1Notation());
}
"""


def tools_book(root):
    """Orders with a frozen heading, column C and row 8 hidden, two empty rows and repeated customers; a second
    sheet."""
    path = os.path.join(root, "orders.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    rows = [["Order", "Customer", "City", "Status", "Amount", "Note", "Ref"],
            [101, "Ana", "Lisbon", "Done", 10], [102, "Bruno", "Porto", "Open", 20], [],
            [103, "ana", "Faro", "Done", 30], [104, "Carla", "Lisbon", "Open", 40], [],
            [105, "Duarte", "porto", "Open", 50], [106, "Bruno", "Braga", "Done", 60]]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    ws.freeze_panes = "A2"
    ws.column_dimensions["C"].hidden = True
    ws.row_dimensions[8].hidden = True
    for col in "EFG":
        ws.column_dimensions[col].width = 15         # saved by LibreOffice as one entry, E to G
    wb.create_sheet("Notes")["A1"] = "Lisbon office"
    wb.save(path)
    script = os.path.join(root, "Tools.gs")
    with open(script, "w", encoding="utf-8") as f:
        f.write(TOOLS)
    return path, script


def sheet_tools(root):
    """The Sheets methods client scripts use that the test lacked: the selection (--select), hidden rows and
    columns, frozen rows from the file, removeDuplicates, createTextFinder and getNextDataCell."""
    path, script = tools_book(root)
    out = os.path.join(root, "selected.xlsx")
    with contextlib.redirect_stdout(io.StringIO()) as said:
        assert run_gas.main([path, script, "--run", "deleteBlankRows", "--select", "Orders!A2:A9", "-o", out]) == 0
    report = said.getvalue()
    assert "Started with Orders!A2:A9 selected, as asked with --select." in report, report
    ws = load_workbook(out)["Orders"]
    left = [r[0] for r in ws.iter_rows(values_only=True)]
    assert left == ["Order", 101, 102, 103, 104, 105, 106, None, None], left     # the two rows left empty at the end
    assert ws.row_dimensions[6].hidden and not ws.row_dimensions[8].hidden, "the hidden row did not move up"
    result = run_gas.run(path, [script], "deleteBlankRows", changes=False)
    report = "\n".join(run_gas.report(result, "yes"))
    assert result["sheets"][0]["changes"] == [], result["sheets"][0]["changes"]   # only A1 was selected
    assert "It works on the selected cells: the run had Orders!A1 selected" in report, report
    try:
        run_gas.run(path, [script], "deleteBlankRows", changes=False, select={"sheet": "Sales", "cells": "A1"})
        raise AssertionError("ran with a sheet that is not there")
    except run_gas.Unrunnable as e:
        assert "There is no sheet Sales in the workbook. Sheets: Orders, Notes" in str(e), e

    result = run_gas.run(path, [script], "countRows", changes=False)
    report = "\n".join(run_gas.report(result, "yes"))
    assert result["log"] == ["Orders has 9 rows"], result["log"]
    assert "It works on the open sheet: the run started on Orders, the one open in the file" in report, report
    result = run_gas.run(path, [script], "countRows", changes=False, select={"sheet": "notes", "cells": "A1"})
    report = "\n".join(run_gas.report(result, "yes"))
    assert result["log"] == ["Notes has 1 rows"] and "Started with Notes!A1 selected" in report, report

    out = os.path.join(root, "hidden.xlsx")
    result = run_gas.run(path, [script], "hideDone", keep=out)
    report = "\n".join(run_gas.report(result, "yes"))
    assert result["log"] == ["frozen 1, row 2 hidden true"], result["log"]
    for line in ["No change in formulas, typed values or results.", "Orders: rows 2, 5, 9 hidden",
                 "Orders: columns F to G hidden", "Orders: column C shown again"]:
        assert line in report, (line, report)
    ws = load_workbook(out)["Orders"]
    assert ws.row_dimensions[5].hidden and not ws.row_dimensions[3].hidden and ws.freeze_panes == "A2", "rows"
    assert ws.column_dimensions["F"].hidden and not ws.column_dimensions["C"].hidden, "columns"
    spans = sorted((d.min, d.max) for d in ws.column_dimensions.values())
    assert all(a[1] < b[0] for a, b in zip(spans, spans[1:])), spans          # no column in two entries
    assert ws.column_dimensions["E"].width == ws.column_dimensions["F"].width == 15, "a width was lost"

    out = os.path.join(root, "deduped.xlsx")
    result = run_gas.run(path, [script], "dedupe", keep=out)
    assert result["log"] == ["left B2:D6"], result["log"]
    got = [r[:4] for r in load_workbook(out)["Orders"].iter_rows(min_row=2, max_row=9, values_only=True)]
    assert got == [(101, "Ana", "Lisbon", "Done"), (102, "Bruno", "Porto", "Open"), (None, None, None, None),
                   (103, "Carla", "Lisbon", "Open"), (104, "Duarte", "porto", "Open"), (None, None, None, None),
                   (105, None, None, None), (106, None, None, None)], got

    result = run_gas.run(path, [script], "findAndFix", changes=False)
    assert result["log"] == ["Orders!C2 Orders!C6 Notes!A1 | 1 | 2 | 2 5 9 | 9"], result["log"]
    assert sorted(result["sheets"][0]["changes"]) == [[2, 3, "Lisboa", ""], [6, 3, "Lisboa", ""]], result["sheets"][0]

    out = os.path.join(root, "shifted.xlsx")
    run_gas.run(path, [script], "shiftCells", keep=out)
    ws = load_workbook(out)["Orders"]
    assert ws["A8"].fill.start_color.rgb == "FFFF0000" and ws["F1"].fill.start_color.rgb == "FF00FF00", "colors"
    assert ws["A9"].fill.fill_type is None and ws["E1"].fill.fill_type is None, "colors left behind"
    assert ws.row_dimensions[7].hidden and ws.column_dimensions["D"].hidden, "hidden row and column"
    assert not ws.row_dimensions[8].hidden and not ws.column_dimensions["C"].hidden, "hidden left behind"
    result = run_gas.run(path, [script], "finderOptions", changes=False)
    assert result["log"] == ["2 2 C3 1 1 1 0"], result["log"]
    assert [3, 3, "Oporto", ""] in result["sheets"][0]["changes"], result["sheets"][0]["changes"]
    result = run_gas.run(path, [script], "lastRows", changes=False)
    assert result["log"] == ["9 3 7 26"], result["log"]

    result = run_gas.run(path, [script], edit={"sheet": "Orders", "cell": "D3", "value": "Done"}, changes=False)
    assert [3, 8, "seen D3", ""] in result["sheets"][0]["changes"], result["sheets"][0]["changes"]
    print("sheet tools: the cells selected with --select, rows and columns hidden and shown (in the report and the "
          "file, and moving when rows and columns are deleted or added), the frozen heading read from the file, "
          "duplicates removed in a range, find and replace across "
          "sheets, the last row with getNextDataCell, and the edited cell as the active cell in onEdit")


if __name__ == "__main__":
    if not shutil.which("node"):
        print("Apps Script: not checked, Node.js is not installed")
        sys.exit(0)
    with tempfile.TemporaryDirectory() as tmp:
        path, code = book(tmp)
        cli(path, code, tmp)
        archive(path, code, tmp)
        services(path, code)
        fetch_cli(path, code, tmp)
        big(tmp)
        sheet_tools(tmp)
    print("all good")
