"""Excel file -> list of what is wrong in it. Used for the Fiverr gig "Google Sheets and Excel", on the
client's file before the work and on ours before delivery.

LibreOffice opens a copy of each file, recalculates every formula and saves the results; the original
is not touched. Reads .xlsx, .xlsm, .xls and .ods. A Google Sheet is checked from File > Download >
Microsoft Excel.

Usage: python3 audit_sheet.py book.xlsx [more files] [-o report.txt]
       python3 audit_sheet.py --compare before.xlsx after.xlsx [-o changes.txt]
              what our work changed: formulas, typed values, and every result that moved, recalculated
       (needs LibreOffice, the soffice command, and pip install openpyxl)

Must fix, and the exit code is 1 when there is any:
  formulas that end in an error, split into the cells that cause it and the cells that only repeat it;
  circular references; references to deleted cells in named ranges, validation, conditional formatting
  or charts.
Worth a look:
  a formula that differs from the matching ones on both sides of it, the way Excel marks them; numbers
  typed as text where formulas read them or in a column of numbers; links to other files; manual
  calculation.
Not checked here:
  functions this LibreOffice lacks (XLOOKUP, FILTER, LET and the rest of Excel 2021 and later),
  Google Sheets functions and custom functions from macros, with the cells that read them.
Last, what openpyxl would drop if the file were saved with it (macros, shapes, slicers and so on), so
the work is done in Excel or LibreOffice when that line shows up.
"""
import argparse
import bisect
import functools
import os
import pathlib
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import warnings
import zipfile
from collections import Counter, defaultdict

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.formulas import FORMULAE

# Newer than Excel 2019, which shows #NAME? for them. LibreOffice has some from 24.8 on.
NEWER = {"XLOOKUP", "XMATCH", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY", "LET", "LAMBDA",
         "MAP", "REDUCE", "SCAN", "BYROW", "BYCOL", "MAKEARRAY", "ISOMITTED", "TEXTSPLIT", "TEXTBEFORE",
         "TEXTAFTER", "VSTACK", "HSTACK", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS", "TOCOL", "TOROW",
         "WRAPCOLS", "WRAPROWS", "EXPAND", "ARRAYTOTEXT", "VALUETOTEXT", "GROUPBY", "PIVOTBY", "PERCENTOF",
         "TRIMRANGE", "REGEXTEST", "REGEXEXTRACT", "REGEXREPLACE", "IMAGE", "STOCKHISTORY", "ANCHORARRAY",
         "SINGLE"}
MEANING = {"#DIV/0!": "division by zero or by an empty cell",  # ia-ok
           "#N/A": "a lookup found no match",
           "#REF!": "a reference to a deleted cell",  # ia-ok
           "#NAME?": "an unknown function or name, often a typo",
           "#VALUE!": "a value of the wrong type, often text where a number should be",  # ia-ok
           "#NUM!": "a number out of range",  # ia-ok
           "#NULL!": "two ranges that do not meet",  # ia-ok
           "#SPILL!": "the result has no empty cells to spill into",  # ia-ok
           "#CALC!": "a calculation with no result, like an empty FILTER"}  # ia-ok
BROKEN = "#REF!"  # ia-ok
GOOGLE = re.compile(r'__xludf\.DUMMYFUNCTION\(\s*"\s*([A-Za-z][\w.]*)\s*\(', re.I)
SHEET_REF = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^'!\[\]]+)!)?(?P<addr>[^!]+)$")
NUMBER_TEXT = re.compile(r"^(?:[$€£¥]|R\$)?\s?[-+]?(?:\d{1,3}(?:[,.  ]\d{3})+|\d+)(?:[.,]\d+)?\s?(?:%|[$€£])?$")
OTHER_FILE = re.compile(r"'?\[\d+\]")   # [1]Sheet1!A1 or '[1]My Sheet'!A1: a link to another workbook
CELL = re.compile(r"(?<![\w.])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\w.(])")
COLUMNS = re.compile(r"^(\$?)([A-Z]{1,3}):(\$?)([A-Z]{1,3})$")
ROWS = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")
LAST_ROW = 1048576
SHAPE = re.compile(rb"<xdr:(?:sp|grpSp|cxnSp)[ >]")
# Parts openpyxl does not read, so saving the file with it loses them.
PARTS = [("xl/vbaProject.bin", "macros (keep_vba=True keeps them)"), ("xl/ctrlProps/", "form controls"),
         ("xl/activeX/", "ActiveX controls"), ("xl/threadedComments/", "threaded comments"),
         ("xl/slicers/", "slicers"), ("xl/timelines/", "timelines"), ("xl/model/", "a data model"),
         ("xl/connections.xml", "data connections"), ("xl/queryTables/", "query tables")]
RECALC = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">
<value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse">
<value>0</value></prop></item>
</oor:items>
"""


def recalculate(paths, work):
    """For each file, the path of a copy that LibreOffice recalculated and saved as .xlsx, or None.

    LibreOffice keeps the results stored in a file unless its profile says otherwise, so a fresh
    profile in `work` turns full recalculation on. Copies get numbered names, so two files with the
    same name in different folders do not overwrite each other.
    """
    profile = pathlib.Path(work, "profile")
    (profile / "user").mkdir(parents=True, exist_ok=True)
    (profile / "user" / "registrymodifications.xcu").write_text(RECALC, encoding="utf-8")
    copies = []
    for i, path in enumerate(paths):
        if os.path.isfile(path):                    # a missing file comes back as None, the others are done
            copies.append(os.path.join(work, f"in{i}{os.path.splitext(path)[1].lower()}"))
            shutil.copyfile(path, copies[-1])
    out = os.path.join(work, "out")
    proc = subprocess.Popen(["soffice", "--headless", "--norestore", f"-env:UserInstallation={profile.as_uri()}",
                             "--convert-to", "xlsx", "--outdir", out, *copies], env=dict(os.environ, HOME=work),
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    try:
        proc.wait(timeout=120 + 60 * len(paths))
    except subprocess.TimeoutExpired:
        # Whatever it finished is used; the rest is reported as not opened. soffice runs soffice.bin under
        # it, so the whole group is stopped, or soffice.bin would keep running.
        os.killpg(proc.pid, signal.SIGKILL)
        proc.wait()
    done = [os.path.join(out, f"in{i}.xlsx") for i in range(len(paths))]
    return [p if os.path.exists(p) else None for p in done]


def where(sheet, row, col):
    name = sheet if re.fullmatch(r"[A-Za-z_][\w.]*", sheet) else "'" + sheet.replace("'", "''") + "'"
    return f"{name}!{get_column_letter(col)}{row}"


@functools.lru_cache(maxsize=None)
def tokens(text):
    try:
        return Tokenizer(text).items
    except Exception:   # a formula the tokenizer cannot read is skipped, the rest still counts
        return []


def relative(ref, row, col):
    """A reference the R1C1 way, counted from the cell that holds it, so copies of a formula match."""
    sheet, bang, addr = ref.rpartition("!")
    addr = addr.upper()

    def r(fixed, n):
        return f"R{n}" if fixed else f"R[{int(n) - row}]"

    def c(fixed, letters):
        n = column_index_from_string(letters)
        return f"C{n}" if fixed else f"C[{n - col}]"

    if m := COLUMNS.match(addr):
        addr = f"{c(m[1], m[2])}:{c(m[3], m[4])}"
    elif m := ROWS.match(addr):
        addr = f"{r(m[1], m[2])}:{r(m[3], m[4])}"
    else:
        addr = CELL.sub(lambda m: r(m[3], m[4]) + c(m[1], m[2]), addr)
    return sheet + bang + addr


def signature(toks, row, col):
    """A formula without spaces, in capitals outside quotes and with R1C1 references, to compare two."""
    out = []
    for t in toks:
        if t.subtype == "RANGE":
            out.append(relative(t.value, row, col))
        elif t.type != "WHITE-SPACE":
            out.append(t.value if t.subtype == "TEXT" else t.value.upper())
    return "".join(out)


def copied(toks, rows, cols):
    """A formula copied `rows` down and `cols` right, as Excel would."""
    return "=" + "".join(Translator.translate_range(t.value, rows, cols) if t.subtype == "RANGE" else t.value
                         for t in toks)


def calls(toks, text, macros, defined):
    """(kind, name) for each function a formula calls: newer, later, google, custom, unknown or known."""
    out = {("google", m[1].upper()) for m in GOOGLE.finditer(text)}
    for t in toks:
        if t.type != "FUNC" or t.subtype != "OPEN":
            continue
        name = t.value[:-1].upper()
        base = name.replace("_XLFN.", "").replace("_XLWS.", "")
        if name.startswith("__XLUDF.DUMMYFUNCTION"):
            continue
        if name.startswith("_XLUDF."):
            out.add(("custom", name[7:]))
        elif base in NEWER or base in defined:   # a named formula called like a function is a LAMBDA
            out.add(("newer", base))
        elif name.startswith("_XLFN."):
            out.add(("later", base))              # Excel 2010 to 2019; LibreOffice has most of them
        elif base not in FORMULAE:
            out.add(("custom" if macros else "unknown", base))
        else:
            out.add(("known", base))
    return out


def reasons(kinds):
    """The functions that explain a #NAME? from LibreOffice: newer, custom or Google ones, else later ones."""
    strong = {n for k, n in kinds if k in ("newer", "custom", "google")}
    return strong or {n for k, n in kinds if k == "later"}


def boxes(toks, sheet, names):
    """The ranges a formula reads, as (sheet, min_col, min_row, max_col, max_row), None when open-ended.
    Named ranges count; table columns, other files and INDIRECT do not."""
    out = []
    for t in toks:
        if t.type != "OPERAND" or t.subtype != "RANGE" or OTHER_FILE.match(t.value):
            continue
        targets = names.get(t.value.upper())
        if targets is None:
            m = SHEET_REF.match(t.value)
            if not m:
                continue
            s = m["sheet"]
            targets = [(s[1:-1].replace("''", "'") if s and s.startswith("'") else s or sheet, m["addr"])]
        for s, addr in targets:
            try:
                out.append((s, *range_boundaries(addr.replace("$", ""))))
            except (ValueError, TypeError):
                pass
    return out


def index_of(cells):
    """{(sheet, row, col)} -> sheet -> col -> sorted rows, for hits()."""
    index = defaultdict(lambda: defaultdict(list))
    for s, r, c in sorted(cells):
        index[s][c].append(r)
    return index


def hits(index, box):
    """The cells of an index that sit inside a range."""
    sheet, c1, r1, c2, r2 = box
    for col, rows in index.get(sheet, {}).items():
        if (c1 is None or c1 <= col) and (c2 is None or col <= c2):
            for row in rows[bisect.bisect_left(rows, r1 or 1):bisect.bisect_right(rows, r2 or LAST_ROW)]:
                yield sheet, row, col


def loops(graph):
    """The cells on a circle in graph {cell: cells it reads}, by Tarjan's strongly connected components."""
    order, low, stack, on, found = {}, {}, [], set(), set()
    for start in graph:
        if start in order:
            continue
        order[start] = low[start] = len(order)
        stack.append(start)
        on.add(start)
        work = [(start, iter(graph[start]))]
        while work:
            node, rest = work[-1]
            for nxt in rest:
                if nxt not in order:
                    order[nxt] = low[nxt] = len(order)
                    stack.append(nxt)
                    on.add(nxt)
                    work.append((nxt, iter(graph.get(nxt, ()))))
                    break
                if nxt in on:
                    low[node] = min(low[node], order[nxt])
            else:
                work.pop()
                if work:
                    low[work[-1][0]] = min(low[work[-1][0]], low[node])
                if low[node] == order[node]:
                    group = []
                    while not group or group[-1] != node:
                        group.append(stack.pop())
                        on.discard(group[-1])
                    if len(group) > 1 or node in graph.get(node, ()):
                        found.update(group)
    return found


def odd_ones(formulas, toks):
    """Formulas that differ from a matching pair on both sides of them, the way Excel marks them,
    with what the pair suggests instead."""
    sig = {rc: signature(toks[rc], *rc) for rc in formulas}
    out = []
    for (r, c), text in sorted(formulas.items()):
        for a, b in (((r - 1, c), (r + 1, c)), ((r, c - 1), (r, c + 1))):
            if a in sig and b in sig and sig[a] == sig[b] != sig[(r, c)]:
                try:
                    want = copied(toks[a], r - a[0], c - a[1])
                except Exception:   # a reference that would move off the sheet
                    want = formulas[a]
                out.append(((r, c), text, want))
                break
    return out


def around(mine, want, size=60):
    """Both formulas from a little before the first place they differ, so the difference shows."""
    i = next((k for k, (x, y) in enumerate(zip(mine, want)) if x != y), min(len(mine), len(want)))
    start = max(0, i - 15)
    return [("..." if start else "") + s[start:start + size] + ("..." if len(s) > start + size else "")
            for s in (mine, want)]


def looks_numeric(text):
    s = text.strip()
    digits = re.sub(r"\D", "", s)
    return bool(NUMBER_TEXT.match(s)) and not re.match(r"^[-+]?0\d", s) and 0 < len(digits) <= 15


def dropped(path, caught):
    """What saving this file with openpyxl would lose: parts it does not read and the warnings it gave."""
    gone = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        gone += [label for part, label in PARTS if any(n.startswith(part) for n in names)]
        if any(n.startswith("xl/drawings/drawing") and SHAPE.search(z.read(n)) for n in names):
            gone.append("shapes or text boxes")
    for w in caught:
        text = str(w.message)
        if "will be removed" in text:
            gone.append(text.split(" is not supported")[0].lower())
    return sorted(set(gone), key=gone.index)


def chart_breaks(path):
    with zipfile.ZipFile(path) as z:
        return sum(BROKEN.encode() in z.read(n) for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))


def audit(path, recalculated):
    """Everything this file shows, as a dict for describe()."""
    found = {"name": path, "failed": not recalculated}
    if not recalculated:
        return found
    openxml = os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        book = load_workbook(path if openxml else recalculated)
        values = load_workbook(recalculated, data_only=True)
    if openxml:
        with zipfile.ZipFile(path) as z:
            macros = "xl/vbaProject.bin" in z.namelist()
    else:
        macros = False
    defined = {n.upper() for n in book.defined_names} | {n.upper() for ws in book.worksheets for n in ws.defined_names}
    names = {}
    for name, dn in book.defined_names.items():
        try:
            names[name.upper()] = list(dn.destinations)
        except Exception:   # a name that is a constant or a formula, not a range
            pass

    sheets, typed = {}, {}
    for ws in book.worksheets:
        sheets[ws.title], typed[ws.title] = {}, {}
        for (r, c), cell in ws._cells.items():   # the cells that exist; iter_rows would create empty ones
            if cell.data_type == "f":
                text = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)
                if text:
                    sheets[ws.title][(r, c)] = text if text.startswith("=") else "=" + text
            elif cell.value is not None:
                typed[ws.title][(r, c)] = cell

    toks, reads, kinds, errors, on_purpose = {}, {}, {}, {}, []
    for title, formulas in sheets.items():
        local = dict(names)
        if title in book.sheetnames:
            for name, dn in book[title].defined_names.items():
                try:
                    local[name.upper()] = list(dn.destinations)
                except Exception:
                    pass
        out = values[title] if title in values.sheetnames else None
        for (r, c), text in formulas.items():
            cell = (title, r, c)
            toks[cell] = tokens(text)
            reads[cell] = boxes(toks[cell], title, local)
            kinds[cell] = calls(toks[cell], text, macros, defined)
            v = out._cells.get((r, c)) if out else None
            if v is None or v.data_type != "e":
                continue
            if v.value == "#N/A" and ("known", "NA") in kinds[cell]:
                on_purpose.append(cell)    # NA() on purpose, often so a chart skips the point
            else:
                errors[cell] = v.value
        for (r, c), v in typed[title].items():
            if v.data_type == "e":
                errors[(title, r, c)] = v.value    # an error typed or pasted as a value

    error_index = index_of(errors)
    graph = {cell: {hit for box in reads.get(cell, ()) for hit in hits(error_index, box)} for cell in errors}
    circle = loops(graph)
    lacking = {cell for cell in errors
               if cell not in circle and errors[cell] == "#NAME?" and reasons(kinds.get(cell, ()))}
    readers = defaultdict(set)
    for cell, parents in graph.items():
        for p in parents:
            readers[p].add(cell)
    unchecked, todo = set(lacking), list(lacking)
    while todo:
        for cell in readers[todo.pop()]:
            if cell not in unchecked and cell not in circle:
                unchecked.add(cell)
                todo.append(cell)
    sources, repeats = [], []
    for cell in sorted(errors):
        if cell in circle or cell in unchecked:
            continue
        text = sheets[cell[0]].get(cell[1:])
        if text is None or BROKEN in text.upper() or not graph[cell] - {cell}:
            sources.append((where(*cell), text, errors[cell]))
        else:
            repeats.append(where(*cell))

    broken = [f"named range {n}" for n, dn in book.defined_names.items() if BROKEN in (dn.attr_text or "")]
    for ws in book.worksheets:
        broken += [f"named range {n} on {ws.title}" for n, dn in ws.defined_names.items()
                   if BROKEN in (dn.attr_text or "")]
        for dv in ws.data_validations.dataValidation:
            if any(BROKEN in (f or "") for f in (dv.formula1, dv.formula2)):
                broken.append(f"data validation on {ws.title}!{dv.sqref}")
        for cf in ws.conditional_formatting:
            if any(BROKEN in f for rule in cf.rules for f in (rule.formula or ())):
                broken.append(f"conditional formatting on {ws.title}!{cf.sqref}")
    charts = chart_breaks(path) if openxml else 0
    if charts:
        broken.append(f"{charts} chart{'s' if charts > 1 else ''}")

    textual, top = {}, {}
    for title, cells in typed.items():
        count = defaultdict(lambda: [0, 0, 0])   # col -> numbers, numbers typed as text, other text
        for (r, c), v in cells.items():
            top[(title, c)] = min(r, top.get((title, c), r))
        mine = []                                   # this sheet's numbers typed as text
        for (r, c), v in cells.items():
            if r == top[(title, c)]:
                continue                            # the first cell of a column is usually its heading
            if v.data_type == "n":
                count[c][0] += 1
            elif v.data_type == "s" and looks_numeric(v.value):
                count[c][1] += 1
                textual[(title, r, c)] = v.value
                mine.append((title, r, c))
            elif v.data_type == "s":
                count[c][2] += 1
        for cell in mine:
            n, t, o = count[cell[2]]
            if n >= 3 and n >= (n + t + o) / 2:
                textual[cell] = (textual[cell], True)
    text_index = index_of(textual)
    read = {hit for cell_reads in reads.values() for box in cell_reads for hit in hits(text_index, box)}
    text_numbers = [(where(*cell), v[0] if isinstance(v, tuple) else v) for cell, v in sorted(textual.items())
                    if isinstance(v, tuple) or cell in read]

    external = [(where(*cell), sheets[cell[0]][cell[1:]]) for cell in sorted(toks)
                if any(t.type == "OPERAND" and OTHER_FILE.match(t.value) for t in toks[cell])]
    files = []
    for link in getattr(book, "_external_links", []):
        target = getattr(getattr(link, "file_link", None), "Target", None)
        if target:
            files.append(re.split(r"[\\/]", target)[-1])

    odd, by_sheet = [], defaultdict(dict)
    for (t, r, c), tokens_ in toks.items():     # once, not once a sheet
        by_sheet[t][(r, c)] = tokens_
    for title, formulas in sheets.items():
        odd += [(where(title, r, c), text, want) for (r, c), text, want in odd_ones(formulas, by_sheet[title])]

    newer, google = Counter(), Counter()
    for cell, ks in kinds.items():
        newer.update({n for k, n in ks if k == "newer"})
        google.update({n for k, n in ks if k == "google"})
    missing = Counter(n for cell in lacking for n in reasons(kinds[cell]))

    found.update(sheets=len(sheets), formulas=len(reads), errors=len(errors), on_purpose=len(on_purpose),
                 sources=sources, repeats=repeats,
                 circle=[where(*c) for c in sorted(circle)], broken=broken, odd=odd, text_numbers=text_numbers,
                 external=external, files=files, manual=book.calculation.calcMode == "manual",
                 lacking=[where(*c) for c in sorted(lacking)],
                 unchecked=[where(*c) for c in sorted(unchecked - lacking)],
                 missing=missing, newer=newer, google=google, drops=dropped(path, caught) if openxml else [])
    found["must_fix"] = bool(sources or repeats or circle or broken)
    return found


def short(text, size=60):
    text = " ".join(str(text).split())
    return text if len(text) <= size else text[:size - 3] + "..."


def some(items, size=8):
    shown = ", ".join(items[:size])
    return shown + (f" and {len(items) - size} more" if len(items) > size else "")


def plural(n, word):
    return f"{n} {word}" + ("" if n == 1 else "s")


def describe(f):
    """The report for one file, as lines of text."""
    lines = [f["name"]]
    if f["failed"]:
        why = f.get("reason") or ("LibreOffice could not open it "
                                  "(a password, a damaged file or a format it does not read)")
        return lines + [f"  Could not be read: {why}."]
    ok = f["formulas"] - f["errors"] - f["on_purpose"] + len([s for s in f["sources"] if s[1] is None])
    lines.append(f"  {plural(f['sheets'], 'sheet')}, {plural(f['formulas'], 'formula')}, "
                 f"{ok} of them with a result and no error"
                 + (f", {f['on_purpose']} with #N/A on purpose from NA()" if f["on_purpose"] else ""))
    must = []
    if f["sources"]:
        must.append(f"Cause an error ({len(f['sources'])}):")
        for place, text, value in f["sources"][:15]:
            meaning = MEANING.get(value, "an error")
            must.append(f"  {place}  {short(text) if text else '(typed as a value)'}  {value}, {meaning}")
        if len(f["sources"]) > 15:
            must.append(f"  and {len(f['sources']) - 15} more")
    if f["repeats"]:
        must.append(f"Only repeat an error from the cells they read ({len(f['repeats'])}): {some(f['repeats'])}")
    if f["circle"]:
        must.append(f"Circular references, the formula reads its own result ({len(f['circle'])}): {some(f['circle'])}")
    if f["broken"]:
        must.append(f"References to deleted cells ({len(f['broken'])}): {some(f['broken'])}")

    look = []
    if f["odd"]:
        look.append(f"Differs from the formulas on both sides ({len(f['odd'])}):")
        for place, text, want in f["odd"][:5]:
            mine, theirs = around(text, want)
            look += [f"  {place}  {mine}", f"  {'they suggest':>{len(place)}}  {theirs}"]
        if len(f["odd"]) > 5:
            look.append(f"  and {len(f['odd']) - 5} more")
    if f["text_numbers"]:
        look.append(f"Numbers typed as text, which SUM and lookups skip ({len(f['text_numbers'])}): "
                    + some([f'{p} "{short(v, 20)}"' for p, v in f["text_numbers"]]))
    if f["external"] or f["files"]:
        look.append(f"Links to other files ({len(f['external'])} formulas): "
                    + some([f"{p} {short(t, 40)}" for p, t in f["external"]], 3)
                    + (f"; files: {some(f['files'], 5)}" if f["files"] else ""))
    if f["manual"]:
        look.append("Calculation is set to manual, so results only change after F9.")

    unchecked = []
    if f["lacking"]:
        names = ", ".join(f"{n} ({plural(k, 'cell')})" for n, k in f["missing"].most_common())
        n = len(f["unchecked"])
        more = f", and {plural(n, 'cell')} that read{'s' if n == 1 else ''} them" if n else ""
        unchecked.append(f"LibreOffice here lacks {names}{more}: {some(f['lacking'] + f['unchecked'], 5)}")
    if f["google"]:
        unchecked.append("Google Sheets functions, checked only in Google Sheets: "
                         + ", ".join(f"{n} ({plural(k, 'cell')})" for n, k in f["google"].most_common()))

    if not must and not look:
        lines.append("  Nothing to fix.")
    for title, items in (("Must fix", must), ("Worth a look", look), ("Not checked here", unchecked)):
        if items:
            lines.append(f"  {title}:")
            lines += [f"    {item}" for item in items]
    if f["newer"]:
        lines.append("  Needs Excel 2021 or later, or Microsoft 365: "
                     + ", ".join(f"{n} ({plural(k, 'cell')})" for n, k in f["newer"].most_common())
                     + ". Excel 2019 and older show #NAME? there.")
    if f["drops"]:
        lines.append(f"  Saving with openpyxl would lose: {', '.join(f['drops'])}. Edit it in Excel or LibreOffice.")
    return lines


def audit_files(paths):
    work = tempfile.mkdtemp(prefix="audit-")
    try:
        found = []
        for path, done in zip(paths, recalculate(paths, work)):
            if not os.path.isfile(path):
                found.append({"name": path, "failed": True, "reason": "the file was not found"})
                continue
            try:
                found.append(audit(path, done))
            except Exception as e:   # one unreadable file should not stop the others
                found.append({"name": path, "failed": True, "reason": f"{type(e).__name__}: {e}"})
        return found
    finally:
        shutil.rmtree(work, ignore_errors=True)


def contents(path, recalculated):
    """Formulas and typed values as saved, and the recalculated result of every formula, by cell."""
    openxml = os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = load_workbook(path if openxml else recalculated)
        values = load_workbook(recalculated, data_only=True)
    formulas, typed, results = {}, {}, {}
    for ws in book.worksheets:
        out = values[ws.title] if ws.title in values.sheetnames else None
        for (r, c), cell in ws._cells.items():
            if cell.data_type == "f":
                text = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)
                if text:
                    formulas[(ws.title, r, c)] = text if text.startswith("=") else "=" + text
                    v = out._cells.get((r, c)) if out else None
                    results[(ws.title, r, c)] = v.value if v is not None else None
            elif cell.value is not None:
                typed[(ws.title, r, c)] = cell.value
    return book.sheetnames, formulas, typed, results


def same(a, b):
    if a in ("", None) and b in ("", None):
        return True
    numbers = (int, float)
    if isinstance(a, numbers) and isinstance(b, numbers) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-9 * max(1, abs(a), abs(b))
    return a == b


def compare(before, after):
    """What changed from one version of a file to the other, both recalculated in LibreOffice."""
    work = tempfile.mkdtemp(prefix="compare-")
    try:
        missing = [p for p in (before, after) if not os.path.isfile(p)]
        if missing:
            return {"failed": True, "names": (before, after), "reason": f"{missing[0]} was not found"}
        done = recalculate([before, after], work)
        if not all(done):
            return {"failed": True, "names": (before, after)}
        (sheets0, f0, t0, r0), (sheets1, f1, t1, r1) = contents(before, done[0]), contents(after, done[1])
    finally:
        shutil.rmtree(work, ignore_errors=True)

    def key(cell):
        return sheets0.index(cell[0]) if cell[0] in sheets0 else len(sheets0), cell[1], cell[2]

    def norm(cell, text):
        return signature(tokens(text), cell[1], cell[2]) if text else None

    # A formula replaced by its value (paste values), or the other way: the side without a formula shows
    # the value typed there, so the cell does not look cleared.
    formulas = [(where(*k), f0.get(k, t0.get(k)), f1.get(k, t1.get(k))) for k in sorted(set(f0) | set(f1), key=key)
                if norm(k, f0.get(k)) != norm(k, f1.get(k))]
    typed = [(where(*k), t0.get(k), t1.get(k)) for k in sorted(set(t0) | set(t1), key=key)
             if not same(t0.get(k), t1.get(k)) and k not in f0 and k not in f1]
    results = [(where(*k), r0.get(k, t0.get(k)), r1.get(k, t1.get(k))) for k in sorted(set(r0) | set(r1), key=key)
               if not same(r0.get(k, t0.get(k)), r1.get(k, t1.get(k)))]
    return {"failed": False, "names": (before, after), "added": [s for s in sheets1 if s not in sheets0],
            "removed": [s for s in sheets0 if s not in sheets1], "formulas": formulas, "typed": typed,
            "results": results}


def describe_changes(c):
    before, after = c["names"]
    lines = [f"{before} -> {after}"]
    if c["failed"]:
        return lines + [f"  {c.get('reason') or 'LibreOffice could not open one of the two files'}."]
    if c["added"] or c["removed"]:
        lines.append(f"  Sheets added: {', '.join(c['added']) or 'none'}; removed: {', '.join(c['removed']) or 'none'}")

    def shown(v):
        # Ten significant digits: 12345.67 -> 12345.68 must not show as 12345.7 on both sides.
        return "(empty)" if v in (None, "") else short(v, 40) if isinstance(v, str) else f"{v:.10g}" \
            if isinstance(v, float) else short(v, 40)

    for title, items in (("Formulas changed", c["formulas"]), ("Typed values changed", c["typed"]),
                         ("Results that changed", c["results"])):
        if items:
            lines.append(f"  {title} ({len(items)}):")
            lines += [f"    {place}  {shown(a)}  ->  {shown(b)}" for place, a, b in items[:15]]
            if len(items) > 15:
                lines.append(f"    and {len(items) - 15} more")
    if len(lines) == 1:
        lines.append("  No change in formulas, typed values or results.")
    return lines


def main(argv=None):
    ap = argparse.ArgumentParser(description="Recalculate spreadsheets in LibreOffice and list what is wrong.")
    ap.add_argument("files", nargs="+", help=".xlsx, .xlsm, .xls or .ods files")
    ap.add_argument("-o", "--out", help="also write the report to this file")
    ap.add_argument("--compare", action="store_true", help="list what changed from the first file to the second")
    a = ap.parse_args(argv)
    if not shutil.which("soffice"):
        print("LibreOffice is needed: the soffice command was not found.")
        return 2
    if a.compare:
        if len(a.files) != 2:
            ap.error("--compare takes two files: before and after")
        changes = compare(*a.files)
        text = "\n".join(describe_changes(changes))
        print(text)
        if a.out:
            with open(a.out, "w", encoding="utf-8") as fh:
                fh.write(text + "\n")
        return 1 if changes["failed"] else 0
    found = audit_files(a.files)
    text = "\n\n".join("\n".join(describe(f)) for f in found)
    print(text)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
    return 1 if any(f["failed"] or f["must_fix"] for f in found) else 0


if __name__ == "__main__":
    sys.exit(main())
