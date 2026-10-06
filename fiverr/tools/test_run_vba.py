"""Run macros written for Excel in LibreOffice and check that the report says what they did.

testdata/vbaProject.bin is the VBA project from XlsxWriter's examples, made in Excel (license in
testdata/LICENSE-XlsxWriter.txt): a Module1 with say_hello, which shows a MsgBox. It stands for a client's .xlsm.
"""
import contextlib
import io
import os
import sys
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import run_vba  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
MODULE = '''Attribute VB_Name = "Module1"
Option Explicit

Sub FillTotals()
    Dim ws As Worksheet, last As Long, r As Long
    Set ws = Worksheets("Data")
    last = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row
    For r = 2 To last
        ws.Cells(r, 4).Value = ws.Cells(r, 2).Value * ws.Cells(r, 3).Value
    Next r
    Dim s As Worksheet
    Set s = Worksheets.Add(After:=Worksheets(Worksheets.Count))
    s.Name = "Summary"
    With s
        .Range("A1").Value = "Grand total"
        .Range("B1").Value = Application.WorksheetFunction.Sum(ws.Range("D2:D" & last))
        .Range("B2").Formula = "=SUM(Data!D2:D" & last & ")"
    End With
    If MsgBox("Add a note?", vbYesNo) = vbYes Then s.Range("A3").Value = "checked"
End Sub

Sub AskName()
    Worksheets("Data").Range("F1").Value = InputBox("Who checked the totals?")
End Sub

Sub Broken()
    Dim ws As Worksheet
    Set ws = Worksheets("Sales")
End Sub

Sub Forever()
    Dim n As Long
    Do
        n = n + 1
    Loop
End Sub

Sub UsesHelper()
    Helper.Mark "from Module1"
End Sub

Private Sub Hidden()
    Worksheets("Data").Range("G1").Value = "private ran"
End Sub

Sub NeedsArgs(target As String)
    Worksheets("Data").Range("G2").Value = target
End Sub

Sub MailIt()
    Dim app As Object
    Set app = CreateObject("Outlook.Application")
End Sub

Sub Report() ' run with (Ctrl+R)
    MsgBox "Done" & vbCrLf & "Rows: " & 5
End Sub

Sub CountItems()
    Dim seen As Object
    Set seen = CreateObject("Scripting.Dictionary")
End Sub

Sub Recorded365()
    Worksheets("Data").Range("E2:E4").Formula2R1C1 = "=RC[-3]*2"
    Worksheets("Data").Columns("C:C").Hidden = True
    Worksheets("Data").Rows("3:4").EntireRow.Hidden = True
End Sub

Sub StampHere()
    Range("A1").Value = "stamped"
    MsgBox ActiveSheet.Name
End Sub
'''
HELPER = '''Attribute VB_Name = "Helper"
Sub Mark(text As String)
    Worksheets("Data").Range("H1").Value = text
End Sub
'''
BAD = '''Attribute VB_Name = "Bad"
Sub Typo()
    Dim x As Long
    x = = 1
End Sub

Sub Fine()
    Worksheets("Data").Range("A1").Value = "x"
End Sub
'''


def leftover():
    """LibreOffice processes this tool started that are still running."""
    found = []
    for pid in filter(str.isdigit, os.listdir("/proc")):
        try:
            with open(f"/proc/{pid}/cmdline", "rb") as f:
                if b"run-vba-" in f.read():
                    found.append(pid)
        except OSError:
            pass
    return found


def files(root):
    book = os.path.join(root, "orders.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Total"])
    for row in [("Pen", 10, 2), ("Pad", 5, 4), ("Ink", 3, 7)]:
        ws.append(row)
    wb.save(book)
    modules = {}
    for name, text in (("Module1", MODULE), ("Helper", HELPER), ("Bad", BAD)):
        modules[name] = os.path.join(root, name + ".bas")
        with open(modules[name], "w", encoding="cp1252", newline="\r\n") as f:     # as the VBA editor exports
            f.write(text)
    return book, modules


def line_of(text, code):
    """The line number the VBA editor gives this code in the module (Attribute lines are hidden there)."""
    shown = [line for line in text.splitlines() if not line.startswith("Attribute")]
    return next(i for i, line in enumerate(shown, start=1) if line.strip() == code)


def reading(root):
    """What the module files say, read without LibreOffice: the parameters of a Sub with a comment in
    parentheses after it, and the Attribute lines Excel writes under a recorded macro with a shortcut key."""
    found = run_vba.procedures("Sub Main() ' run with (Ctrl+M)\nFunction Fee(rate As Double) As Double\n"
                               "Private Sub Two(a As Long, _\n    b As Long)\n")
    assert found == {"main": ("Main", "", False), "fee": ("Fee", "rate As Double", False),
                     "two": ("Two", "a As Long, _", True)}, found
    path = os.path.join(root, "Recorded.bas")
    with open(path, "w", encoding="cp1252", newline="\r\n") as f:
        f.write('Attribute VB_Name = "Recorded"\nSub Macro1()\nAttribute Macro1.VB_ProcData.VB_Invoke_Func = "m\\n14"\n'
                '    Range("A1").Value = 1\nEnd Sub\n')
    name, code = run_vba.read_bas(path)
    assert (name, code.splitlines()) == ("Recorded", ["Sub Macro1()", '    Range("A1").Value = 1', "End Sub"]), code
    print("reading: no parameters taken from a comment, the Attribute line of a recorded macro left out")


def cli(book, modules, root):
    out, said = os.path.join(root, "after.xlsx"), io.StringIO()
    with contextlib.redirect_stdout(said):
        code = run_vba.main([book, "FillTotals", "--code", modules["Module1"], "-o", out])
    report = said.getvalue()
    assert code == 0, report
    for line in ["FillTotals ran to the end", "MsgBox: Add a note?", "Sheets added: Summary",
                 "Data!D2  (empty)  ->  20", "Data!D4  (empty)  ->  21", "Summary!B1  (empty)  ->  61",
                 "Summary!B2  (empty)  ->  =SUM(Data!D2:D4)", "Summary!A3  (empty)  ->  checked"]:
        assert line in report, (line, report)
    assert "Started on the sheet" not in report, report         # one sheet: nothing to choose
    after = load_workbook(out)
    assert after.sheetnames == ["Data", "Summary"] and after["Data"]["D3"].value == 20, after.sheetnames
    with contextlib.redirect_stdout(io.StringIO()) as said:
        assert run_vba.main([book, "FillTotals", "--code", modules["Module1"], "--answer", "no"]) == 0
    assert "Summary!A3" not in said.getvalue() and "answered no" in said.getvalue(), said.getvalue()
    assert load_workbook(book).sheetnames == ["Data"], "the client's file was changed"
    print("macro: ran in LibreOffice; the new sheet, 6 values, the formula and its result reported; MsgBox "
          "answered yes, then no; the client's file untouched")

    # Recorded in Excel 365, which writes Formula2R1C1; it also hides a column and two rows.
    with contextlib.redirect_stdout(io.StringIO()) as said:
        assert run_vba.main([book, "Recorded365", "--code", modules["Module1"]]) == 0
    report = said.getvalue()
    for line in ["Recorded365 ran to the end", "ran as Formula and FormulaR1C1", "Data!E2  (empty)  ->  =B2*2",
                 "Data!E4  (empty)  ->  =B4*2", "Data: column C hidden", "Data: rows 3 to 4 hidden"]:
        assert line in report, (line, report)
    print("recorded in Excel 365: Formula2R1C1 run as FormulaR1C1, the formulas and the hidden column and rows "
          "reported")


def paths(book, modules):
    def go(macro, *more, **how):
        result = run_vba.run(book, macro, [modules["Module1"], *more], changes=False, **how)
        return result, "\n".join(run_vba.report(result, how.get("answer", "yes")))

    result, report = go("Report")
    assert result["status"] == "ok" and result["asked"] == ["MsgBox: Done Rows: 5"], report

    result, report = go("AskName", typed="Ana")
    assert result["status"] == "ok" and result["asked"] == ["InputBox: Who checked the totals?"], report

    result, report = go("Broken")
    line = line_of(MODULE, 'Set ws = Worksheets("Sales")')
    assert result["status"] == "error" and result["at"] == ("Module1", line), report
    assert f"stopped at Module1 line {line} with error 9: Subscript out of range" in report, report
    assert f'  {line}: Set ws = Worksheets("Sales")' in report, report

    result, report = go("Forever", timeout=5)
    assert result["status"] == "timeout" and "did not finish in 5" in report, report
    assert not leftover(), f"LibreOffice still running after the loop was stopped: {leftover()}"

    result, _ = go("UsesHelper", modules["Helper"])
    assert result["status"] == "ok", result
    result, report = go("Hidden")
    assert result["status"] == "ok", report

    result, report = go("MailIt")
    outlook = line_of(MODULE, 'Set app = CreateObject("Outlook.Application")')
    assert result["status"] == "error" and f"Module1 line {outlook}: Outlook" in report, report

    result, report = go("CountItems")
    dictionary = line_of(MODULE, 'Set seen = CreateObject("Scripting.Dictionary")')
    assert result["status"] == "error" and result["at"] == ("Module1", dictionary), report
    assert f"Module1 line {dictionary}: Scripting.Dictionary, which only Excel on Windows has" in report, report

    result, report = go("Fine", modules["Bad"])
    assert result["status"] == "compile" and result["procedures"] == ["Typo"], report
    assert "LibreOffice cannot compile Bad (Typo)" in report, report
    result, report = go("Hidden", modules["Bad"])          # Bad is not used by it: left out, as Excel would
    assert result["status"] == "ok" and result["left_out"] == {"Bad": ["Typo"]} and "Bad (Typo)" in report, report

    for macro, said in (("NeedsArgs", "needs arguments (target As String)"), ("Nowhere", "is not in the workbook"),
                        ("Module1.Nowhere", "Macros that can run: Helper.Mark")):
        try:
            go(macro, modules["Helper"])
            raise AssertionError(f"{macro} ran")
        except run_vba.Unrunnable as e:
            assert said in str(e) or macro == "Module1.Nowhere" and "Module1.FillTotals" in str(e), (macro, e)
    print("paths: InputBox, an error placed on its line, a loop stopped after 5 s, a call into another module, "
          "a Private Sub, Outlook and Scripting.Dictionary flagged, a module that does not compile named and left out when not used, "
          "arguments and unknown names refused")


def sheets(root, modules):
    """Where a macro that acts on the open sheet starts: on the sheet open when the file opens, on the one
    --sheet names, or on the sheet with its button, as when the client clicks it there."""
    two = os.path.join(root, "two.xlsx")
    wb = Workbook()
    wb.active.title = "Data"
    wb.create_sheet("Notes")
    wb.active = 1
    wb.save(two)
    result = run_vba.run(two, "StampHere", [modules["Module1"]], changes=False)
    report = "\n".join(run_vba.report(result, "yes"))
    assert result["asked"] == ["MsgBox: Notes"] and result["start"] == ("Notes", "open", []), report
    assert "Started on the sheet Notes, the one open when the file opens; the macro works on the open sheet" \
        in report, report

    with contextlib.redirect_stdout(io.StringIO()) as said:
        code = run_vba.main([two, "StampHere", "--code", modules["Module1"], "--sheet", "data"])
    report = said.getvalue()
    for line in ["Started on the sheet Data, as asked with --sheet.", "MsgBox: Data", "Data!A1  (empty)  ->  stamped"]:
        assert code == 0 and line in report, (line, report)
    assert "Notes!A1" not in report, report
    try:
        run_vba.run(two, "StampHere", [modules["Module1"]], changes=False, sheet="Sales")
        raise AssertionError("ran on a sheet that is not there")
    except run_vba.Unrunnable as e:
        assert "There is no sheet Sales in the workbook. Sheets: Data, Notes" in str(e), e

    try:
        import xlsxwriter
    except ImportError:
        print("sheets: the open sheet and --sheet checked; buttons not, pip install xlsxwriter builds the test file")
        return
    # A form button on Report and, on Archive, a shape the macro is assigned to (XlsxWriter cannot assign one,
    # so it is written into the drawing as Excel does).
    made, path = os.path.join(root, "made.xlsm"), os.path.join(root, "buttons.xlsm")
    wb = xlsxwriter.Workbook(made)
    for title in ("Data", "Report", "Archive"):
        wb.add_worksheet(title)
    wb.get_worksheet_by_name("Report").insert_button("B3", {"macro": "StampHere", "caption": "Stamp"})
    wb.get_worksheet_by_name("Archive").insert_textbox("B3", "Stamp")
    wb.add_vba_project(os.path.join(HERE, "testdata", "vbaProject.bin"))
    wb.close()
    with zipfile.ZipFile(made) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename.startswith("xl/drawings/drawing"):
                assert data.count(b'<xdr:sp macro=""') == 1, data
                data = data.replace(b'<xdr:sp macro=""', b'<xdr:sp macro="[0]!Module1.StampHere"')
            dst.writestr(item, data)
    assert run_vba.buttons(path) == {"stamphere": [("", "Report"), ("module1", "Archive")]}, run_vba.buttons(path)
    result = run_vba.run(path, "StampHere", [modules["Module1"]], changes=False)
    report = "\n".join(run_vba.report(result, "yes"))
    assert result["asked"] == ["MsgBox: Report"] and result["start"] == ("Report", "button", ["Archive"]), report
    assert "Started on the sheet Report, where its button is (it has one on Archive too)." in report, report
    print("sheets: started on the sheet open when the file opens, on the one --sheet names, and on the sheet "
          "with the macro's button (a form button, or a shape it is assigned to); a sheet that is not there refused")


def workbook_macro(root):
    """A macro inside an .xlsm made in Excel, as a client sends it."""
    try:
        import xlsxwriter
    except ImportError:
        print("xlsm: not checked, pip install xlsxwriter builds the test file")
        return
    path = os.path.join(root, "hello.xlsm")
    wb = xlsxwriter.Workbook(path)
    wb.add_worksheet("Data").write_row(0, 0, ["Item", "Qty"])
    wb.add_vba_project(os.path.join(HERE, "testdata", "vbaProject.bin"))
    wb.close()
    result = run_vba.run(path, "say_hello", changes=False)
    report = "\n".join(run_vba.report(result, "yes"))
    assert result["status"] == "ok" and result["asked"] == ["MsgBox: Hello from Python!"], report  # ia-ok
    print("xlsm: the macro saved in the workbook by Excel ran, and its MsgBox was answered")


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        book, modules = files(tmp)
        reading(tmp)
        cli(book, modules, tmp)
        paths(book, modules)
        sheets(tmp, modules)
        workbook_macro(tmp)
    assert not leftover(), f"LibreOffice still running: {leftover()}"
    print("no LibreOffice left running")
    print("all good")
