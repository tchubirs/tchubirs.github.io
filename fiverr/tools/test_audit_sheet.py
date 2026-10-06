"""Build spreadsheets with known problems, check them in LibreOffice and compare the report with what was planted."""
import contextlib
import io
import os
import sys
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_sheet  # noqa: E402


def helpers():
    cases = {"12,50": True, "1,234.50": True, "1.234,50": True, "$1,200": True, "45%": True, "-3": True,
             "2024": True, "007": False, "abc": False, "12.09.2026": False, "+33 6 12 34 56 78": False,
             "1234567890123456": False, "": False}
    for text, want in cases.items():
        assert audit_sheet.looks_numeric(text) == want, (text, want)
    refs = {("A1", 5, 3): "R[-4]C[-2]", ("$A$1:B5", 5, 3): "R1C1:R[0]C[-1]", ("A:A", 9, 2): "C[-1]:C[-1]",
            ("$3:4", 4, 1): "R3:R[0]", ("'Q1 2024'!B2", 2, 2): "'Q1 2024'!R[0]C[0]", ("Q1SALES", 3, 3): "Q1SALES"}
    for (ref, row, col), want in refs.items():
        assert audit_sheet.relative(ref, row, col) == want, (ref, audit_sheet.relative(ref, row, col), want)
    assert audit_sheet.loops({1: {2}, 2: {1}, 3: {1}, 4: {4}, 5: set()}) == {1, 2, 4}
    print("helpers: numbers typed as text, R1C1 references and circles read as expected")


def client(path):
    """A small order book with one of each problem the checker looks for."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Price", "Qty", "Total"])
    for row in [("Pen", 2, 10), ("Pad", 4, 5), ("Ink", 7, 3), ("Bag", 12, 2), ("Cap", "12,50", 1)]:
        ws.append(row)
    for r in range(2, 6):
        ws[f"D{r}"] = f"=B{r}*C{r}"
    ws["D4"] = "=B4*C3"                      # copied wrong, it reads the row above
    ws["B7"] = "=SUM(B2:B6)"                 # skips the price typed as text
    ws["E2"] = "=B2/F2"                      # F2 is empty
    ws["E3"] = "=E2*2"                       # only repeats E2
    ws["E4"] = "=SUM(#REF!)"                 # a deleted range
    ws["E6"] = "=IF(C6>5,C6,NA())"           # #N/A on purpose
    ws["G1"], ws["H1"] = "=H1+1", "=G1+1"    # each reads the other
    calc = wb.create_sheet("Calc")
    calc["A1"] = "=_xlfn.XLOOKUP(2,Data!C2:C6,Data!B2:B6)"
    calc["A2"] = "=A1*2"
    calc["A3"] = "=SUMM(1,2)"                # a typo
    calc["A4"] = '=IFERROR(__xludf.DUMMYFUNCTION("QUERY(Data!A1:D6,""select A"")"),"Item")'
    calc["A5"] = "=_xlfn.LAMBDA(_xlpm.x,_xlpm.x+1)(2)"   # 3 in Excel 365; LibreOffice gives #VALUE!
    wb.defined_names["Old"] = DefinedName("Old", attr_text="#REF!")  # ia-ok
    dv = DataValidation(type="list", formula1="#REF!")  # ia-ok
    ws.add_data_validation(dv)
    dv.add("C2:C6")
    ws.conditional_formatting.add("D2:D6", FormulaRule(formula=["#REF!>0"]))
    wb.calculation.calcMode = "manual"
    wb.save(path)


def clean(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Units", "Price", "Revenue"])
    for r, (month, units, price) in enumerate([("Jan", 10, 5), ("Feb", 12, 5), ("Mar", 9, 6)], start=2):
        ws.append([month, units, price, f"=B{r}*C{r}"])
    ws["D5"] = "=SUM(D2:D4)"
    wb.save(path)


def files(root):
    # The same file name in two folders: the copies LibreOffice opens must not overwrite each other.
    paths = [os.path.join(root, "a", "book.xlsx"), os.path.join(root, "b", "book.xlsx")]
    for path in paths:
        os.makedirs(os.path.dirname(path))
    client(paths[0])
    clean(paths[1])
    bad, good = audit_sheet.audit_files(paths)

    assert not bad["failed"] and bad["must_fix"], bad
    assert (bad["sheets"], bad["formulas"], bad["on_purpose"]) == (2, 16, 1), bad
    assert [(p, v) for p, _, v in bad["sources"]] == [("Calc!A3", "#NAME?"), ("Data!E2", "#DIV/0!"),
                                                      ("Data!E4", "#REF!")], bad["sources"]  # ia-ok
    assert bad["repeats"] == ["Data!E3"] and bad["circle"] == ["Data!G1", "Data!H1"], bad
    # LibreOffice before 24.8 has no XLOOKUP: then that cell and the one reading it are set apart. It has no
    # LAMBDA, whose #VALUE! there is set apart too, and not taken for an error of the client's; nor QUERY, to
    # which 26.8 gives #VALUE! inside the IFERROR of Google.
    assert "Calc!A5" in bad["lacking"] and set(bad["lacking"]) <= {"Calc!A1", "Calc!A4", "Calc!A5"}, bad
    assert bad["unchecked"] == (["Calc!A2"] if "Calc!A1" in bad["lacking"] else []), bad
    assert sorted(bad["broken"]) == ["conditional formatting on Data!D2:D6", "data validation on Data!C2:C6",
                                     "named range Old"], bad["broken"]
    assert bad["odd"] == [("Data!D4", "=B4*C3", "=B4*C4")], bad["odd"]
    assert bad["text_numbers"] == [("Data!B6", "12,50")], bad["text_numbers"]
    assert bad["manual"] and bad["newer"] == {"XLOOKUP": 1, "LAMBDA": 1} and bad["google"] == {"QUERY": 1}, bad
    assert bad["external"] == [] and bad["drops"] == [], bad

    report = "\n".join(audit_sheet.describe(bad))
    for line in ["Cause an error (3):", "Data!E2  =B2/F2  #DIV/0!, division by zero or by an empty cell",
                 "Only repeat an error from the cells they read (1): Data!E3",
                 "Circular references, the formula reads its own result (2): Data!G1, Data!H1",
                 "Data!D4  =B4*C3", "they suggest  =B4*C4", 'Data!B6 "12,50"', "Calculation is set to manual",
                 "1 with #N/A on purpose from NA()", "checked only in Google Sheets: QUERY (1 cell)",
                 "Needs Excel 2021 or later, or Microsoft 365: XLOOKUP (1 cell)"]:
        assert line in report, (line, report)
    print("client file: 3 errors and the cell that repeats one, the circle, 3 broken references, the odd "
          "formula, the number typed as text, manual calculation, XLOOKUP, LAMBDA and QUERY, all found")

    assert not good["must_fix"] and (good["formulas"], good["sources"], good["odd"]) == (4, [], []), good
    assert "  Nothing to fix." in audit_sheet.describe(good), audit_sheet.describe(good)
    print("clean file with the same name: nothing to fix")
    return paths


def cli(paths, root):
    out, shown = os.path.join(root, "report.txt"), io.StringIO()
    with contextlib.redirect_stdout(shown):
        assert audit_sheet.main([paths[0], "-o", out]) == 1
        assert audit_sheet.main([paths[1]]) == 0
    with open(out, encoding="utf-8") as f:
        written = f.read()
    assert "Data!E2" in written and written.strip() in shown.getvalue(), written
    print("command line: exit code 1 with problems, 0 without, report written")


def drops(paths, root):
    path = os.path.join(root, "with-parts.xlsx")
    with zipfile.ZipFile(paths[1]) as src, zipfile.ZipFile(path, "w") as dst:
        for name in src.namelist():
            dst.writestr(name, src.read(name))
        dst.writestr("xl/vbaProject.bin", b"")
        dst.writestr("xl/drawings/drawing1.xml", b'<xdr:wsDr><xdr:twoCellAnchor><xdr:sp macro=""/></xdr:twoCellAnchor>')

    class Said:
        message = "Data Validation extension is not supported and will be removed"
    assert audit_sheet.dropped(path, [Said()]) == ["macros (keep_vba=True keeps them)", "shapes or text boxes",
                                                   "data validation extension"]
    print("openpyxl limits: macros, shapes and the extension it warned about are listed")


def edges(root):
    """A file that is not there, a formula pasted as its value, a cent that changed and a link to another
    workbook whose sheet name has a space."""
    good = os.path.join(root, "b", "book.xlsx")
    found = audit_sheet.audit_files([good, os.path.join(root, "nowhere.xlsx")])
    assert not found[0]["failed"] and found[1] == {"name": os.path.join(root, "nowhere.xlsx"), "failed": True,
                                                    "reason": "the file was not found"}, found[1]
    c = audit_sheet.compare(good, os.path.join(root, "nowhere.xlsx"))
    assert c["failed"] and "was not found" in "\n".join(audit_sheet.describe_changes(c)), c

    before, after = os.path.join(root, "paste0.xlsx"), os.path.join(root, "paste1.xlsx")
    for path, b5, price in ((before, "=SUM(B2:B4)", 5), (after, 27, 5.01)):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        for row in (["Month", "Units"], ["Jan", 10], ["Feb", 12], ["Mar", 5], ["Total", b5],
                    ["Price", price * 2469.134]):
            ws.append(row)
        wb.save(path)
    c = audit_sheet.compare(before, after)
    assert c["formulas"] == [("Sales!B5", "=SUM(B2:B4)", 27)] and c["results"] == [], c
    report = "\n".join(audit_sheet.describe_changes(c))
    assert "Sales!B5  =SUM(B2:B4)  ->  27" in report and "Sales!B6  12345.67  ->  12370.36134" in report, report

    linked = os.path.join(root, "linked.xlsx")
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "='[1]My Sheet'!A1"
    wb.active["A2"] = 5
    wb.save(linked)
    found = audit_sheet.audit_files([linked])[0]
    assert [place for place, _ in found["external"]] == ["Data!A1"], found["external"]
    print("edges: a missing file reported without stopping the others, a formula pasted as its value shown with "
          "the value, a change in the cents shown, a link to '[1]My Sheet' found")


def changes(root):
    """Fix two things in the client file and check that the compare report names them and what moved."""
    before, after, out = (os.path.join(root, n) for n in ("before.xlsx", "after.xlsx", "changes.txt"))
    client(before)
    wb = load_workbook(before)
    wb["Data"]["D4"] = "=B4*C4"      # the formula copied wrong, fixed
    wb["Data"]["B6"] = 12.5          # the price typed as text, now a number
    wb.save(after)
    c = audit_sheet.compare(before, after)
    assert not c["failed"] and c["added"] == c["removed"] == [], c
    assert c["formulas"] == [("Data!D4", "=B4*C3", "=B4*C4")], c["formulas"]
    assert c["typed"] == [("Data!B6", "12,50", 12.5)], c["typed"]
    assert c["results"] == [("Data!D4", 35, 21), ("Data!B7", 25, 37.5)], c["results"]
    with contextlib.redirect_stdout(io.StringIO()):
        assert audit_sheet.main(["--compare", before, after, "-o", out]) == 0
    with open(out, encoding="utf-8") as f:
        report = f.read()
    for line in ["Formulas changed (1):", "Data!D4  =B4*C3  ->  =B4*C4", "Typed values changed (1):",
                 "Results that changed (2):", "Data!B7  25  ->  37.5"]:
        assert line in report, (line, report)
    same = audit_sheet.compare(before, before)
    assert (same["formulas"], same["typed"], same["results"]) == ([], [], []), same
    assert "No change" in "\n".join(audit_sheet.describe_changes(same))
    print("compare: the fixed formula, the number no longer typed as text and the 2 results that moved, "
          "nothing else; the same file twice shows no change")


if __name__ == "__main__":
    helpers()
    root = tempfile.mkdtemp()
    paths = files(root)
    cli(paths, root)
    drops(paths, root)
    edges(root)
    changes(root)
    print("all good")
