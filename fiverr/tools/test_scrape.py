"""Scrape a small site served on this machine and compare every output with what the pages hold."""
import csv
import functools
import http.server
import json
import os
import sys
import tempfile
import threading
import urllib.parse

import requests
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import scrape  # noqa: E402

PAGES = {
    # Like many shops: everything allowed first, then the exceptions, some with wildcards.
    "robots.txt": "User-agent: *\nAllow: /\nDisallow: /private/\nDisallow: /*?*sort=\n",
    "api/one.json": '{"data": {"items": [{"name": "A", "n": 1}, {"name": "B", "n": 2}]}, "links": {"next": "two.json"}}',
    "api/two.json": '{"data": {"items": [{"name": "C", "n": 3}]}, "links": {"next": null}}',
    "shop/page1.html": """<html><body>
<article class="p"><h3><a href="item/a.html" title="Alpha">Alpha</a></h3><p class="price">£1,234.50</p>
<p class="star-rating Three"></p></article>
<article class="p"><h3><a href="item/b.html" title="Beta">Beta</a></h3><p class="price">12,50 €</p></article>
<article class="p"><h3><a href="item/c.html" title="Gamma">Gamma</a></h3></article>
<ul><li class="next"><a href="page2.html">next</a></li></ul></body></html>""",
    "shop/page2.html": """<html><body>
<article class="p"><h3><a href="item/a.html" title="Alpha">Alpha</a></h3><p class="price">£1,234.50</p>
<p class="star-rating Three"></p></article>
<article class="p"><h3><a href="/shop/item/d.html" title="Delta">Delta
      Four</a></h3><p class="price">-3</p></article>
<ul><li class="next"><a href="page1.html">back to the first page</a></li></ul></body></html>""",
    "private/x.html": "<html><body><article class='p'><h3><a title='Hidden'>Hidden</a></h3></article></body></html>",
    "deep/list.html": """<html><body>
<div class="row"><a href="one.html">One</a></div>
<div class="row"><a href="missing.html">Missing</a></div>
<div class="row"><a href="/private/two.html">Two</a></div></body></html>""",
    "deep/one.html": """<html><body><p class="desc">The first
      item</p><table><tr><td class="stock">In stock (7 available)</td></tr></table></body></html>""",
    "private/two.html": "<html><body><p class='desc'>Hidden</p></body></html>",
    "js/list.html": """<html><body><div id="list"></div><script>
setTimeout(function () {
  document.getElementById("list").innerHTML =
    '<article class="p"><h3><a href="a.html" title="Built">Built</a></h3></article>'
}, 300)
</script></body></html>""",
    "js/scroll.html": """<html><body style="margin:0"><div id="list"></div><script>
var n = 0, busy = false;
function add(k, height) {
  for (var i = 0; i < k; i++) {
    n += 1;
    var a = document.createElement("article");
    a.className = "p";
    a.style.height = height + "px";
    a.innerHTML = '<h3><a href="item' + n + '.html" title="Item ' + n + '">'
      + 'Item ' + n + '</a></h3>';
    document.getElementById("list").appendChild(a);
  }
}
add(5, 400);
window.addEventListener("scroll", function () {
  if (busy || n >= 20 || window.innerHeight + window.scrollY < document.body.scrollHeight - 50) return;
  busy = true;
  setTimeout(function () { add(5, 400); busy = false; }, 200);
});
</script></body></html>""",
    "js/more.html": """<html><body><div id="list"></div><button id="more">More</button><script>
var n = 0;
function add(k) {
  for (var i = 0; i < k; i++) {
    n += 1;
    var a = document.createElement("article");
    a.className = "p";
    a.innerHTML = '<h3><a href="item' + n + '.html" title="Item ' + n + '">Item ' + n + '</a></h3>';
    document.getElementById("list").appendChild(a);
  }
}
add(5);
document.getElementById("more").onclick = function () {
  setTimeout(function () { add(5); if (n >= 20) document.getElementById("more").remove(); }, 200);
};
</script></body></html>""",
}


# A shop that answers like Shopify's /products.json, a page at a time.
SHOP = [{"id": n, "title": f"Runner {n}", "vendor": "Acme", "tags": ["shoes", f"size-{n}"],
         "body_html": f"<p>Light shoe, <strong>number {n}</strong>.</p><ul><li>Mesh</li><li>Rubber</li></ul>",
         "images": [{"src": f"https://cdn.test/{n}a.jpg"}, {"src": f"https://cdn.test/{n}b.jpg"}],
         "variants": [{"title": size, "sku": f"R{n}-{size}", "price": f"{50 + n}.50", "available": size != "44"}
                      for size in ["40", "42", "44"][:(1 + n % 3) if n < 6 else 0]]} for n in range(1, 7)]
REQUESTED = []


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_GET(self):
        REQUESTED.append(self.path)
        url = urllib.parse.urlsplit(self.path)
        if url.path != "/store/products.json":
            return super().do_GET()
        ask = dict(urllib.parse.parse_qsl(url.query))
        page, limit = int(ask.get("page", 1)), int(ask.get("limit", 30))
        body = json.dumps({"products": SHOP[(page - 1) * limit:page * limit]}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def numbers():
    cases = {"£51.77": 51.77, "1.234,50 €": 1234.5, "1 234,50": 1234.5, "In stock (22 available)": 22,
             "−7": -7, "1,234": 1234, "12,50": 12.5, "none": None, "": None}
    for text, want in cases.items():
        assert scrape.number(text) == want, (text, scrape.number(text), want)
    print("numbers: all", len(cases), "read as expected")


def serve():
    root = tempfile.mkdtemp()
    for path, body in PAGES.items():
        os.makedirs(os.path.dirname(os.path.join(root, path)) or root, exist_ok=True)
        with open(os.path.join(root, path), "w", encoding="utf-8") as f:
            f.write(body)
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(Quiet, directory=root))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return f"http://127.0.0.1:{server.server_address[1]}", root, server


def site(base, root):
    job = {"start": base + "/shop/page1.html", "item": "article.p", "delay": 0, "numbers": ["price"],
           "fields": {"name": "h3 a@title", "text": "h3 a", "price": "p.price", "link": "h3 a@href",
                      "rating": "p.star-rating@class"},
           "next": "li.next a@href"}

    rows, stats = scrape.scrape(job, log=lambda *a: None)
    assert stats == {"pages": 2, "found": 5, "duplicates": 1}, stats
    assert [r["name"] for r in rows] == ["Alpha", "Beta", "Gamma", "Delta"], rows
    assert [r["price"] for r in rows] == [1234.5, 12.5, None, -3], rows
    assert rows[3]["text"] == "Delta Four", rows[3]
    assert rows[0]["rating"] == "star-rating Three" and rows[1]["rating"] == "", rows
    assert rows[0]["link"] == base + "/shop/item/a.html" and rows[3]["link"] == base + "/shop/item/d.html", rows

    out = os.path.join(root, "result")
    job_file = os.path.join(root, "job.json")
    with open(job_file, "w") as f:
        json.dump(job, f)
    assert scrape.main([job_file, "-o", out]) == 0
    wb = load_workbook(out + ".xlsx")
    sheet = [[c.value for c in r] for r in wb["Data"].iter_rows()]
    assert sheet[0] == ["name", "text", "price", "link", "rating"], sheet[0]
    assert [r[2] for r in sheet[1:]] == [1234.5, 12.5, None, -3], sheet
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows()}
    assert about["Rows"] == 4 and about["Duplicate rows removed"] == 1 and about["Empty values in price"] == 1, about
    assert "Detail pages read" not in about, about
    with open(out + ".csv", encoding="utf-8-sig") as f:
        assert len(list(csv.reader(f))) == 5
    with open(out + ".json", encoding="utf-8") as f:
        assert [r["name"] for r in json.load(f)] == ["Alpha", "Beta", "Gamma", "Delta"]
    print("site: 2 pages, 4 rows, 1 duplicate removed, the loop back to page 1 stopped, files match")

    hidden, stats = scrape.scrape(dict(job, start=base + "/private/x.html"), log=lambda *a: None)
    assert hidden == [] and stats["pages"] == 0, (hidden, stats)
    sorted_out, stats = scrape.scrape(dict(job, start=base + "/shop/page1.html?sort=price"), log=lambda *a: None)
    assert sorted_out == [] and stats["pages"] == 0, (sorted_out, stats)
    assert not [p for p in REQUESTED if p.startswith("/private/x") or "sort=" in p], REQUESTED
    said = []
    kept, stats = scrape.scrape(dict(job, start=[base + "/shop/page1.html", base + "/gone.html"], next=None),
                                log=said.append)
    assert [r["name"] for r in kept] == ["Alpha", "Beta", "Gamma"] and stats["pages"] == 1, (kept, stats)
    assert said and "could not read" in said[0] and "gone.html" in said[0], said
    print("a page that cannot be read: the run stops there and keeps the 3 rows read before it")
    print("robots.txt: after Allow: / the disallowed page and the address a wildcard rule matches were not fetched")


def robots():
    text = """# Shop rules
User-agent: Googlebot
Disallow: /

User-agent: *
Allow: /
Disallow: /cart
Allow: /cart/public
Disallow: /*?*sort_by=
Disallow: /*.pdf$
Allow: /page
Disallow: /page
Crawl-delay: 2

User-agent: scrape.py
User-agent: another-bot
Disallow: /mine/
"""
    anyone = scrape.Robots(text, "Mozilla/5.0 (X11; Linux x86_64) Chrome/120.0")
    cases = {"/": True, "/cart": False, "/cart/x": False, "/cart/public/x": True, "/shop?a=1&sort_by=price": False,
             "/shop?sort=price": True, "/files/a.pdf": False, "/files/a.pdf?v=2": True, "/page": True,
             "/robots.txt": True}
    for path, want in cases.items():
        assert anyone.allowed("https://shop.test" + path) == want, (path, want)
    assert anyone.delay == 2
    named = scrape.Robots(text, scrape.AGENT)          # a group that names scrape.py replaces the * group
    assert named.allowed("https://shop.test/cart") and not named.allowed("https://shop.test/mine/a") and not named.delay
    root_only = scrape.Robots("User-agent: *\nAllow: /$\nDisallow: /\n", scrape.AGENT)
    assert root_only.allowed("https://x.test/") and not root_only.allowed("https://x.test/a")

    class Answer:
        def __init__(self, code, text=""):
            self.status_code, self.text = code, text

    def down(url, timeout):
        raise requests.ConnectionError("no answer")
    fetcher = scrape.Fetcher(delay=0)
    for answer, allowed, said in [(lambda url, timeout: Answer(404), True, ""),
                                  (lambda url, timeout: Answer(503), False, "robots.txt answered 503"),
                                  (lambda url, timeout: Answer(403), False, "robots.txt answered 403"),
                                  (down, False, "robots.txt could not be read (ConnectionError)")]:
        fetcher.robots.clear()
        fetcher.session.get = answer
        assert fetcher.allowed("https://x.test/a") == allowed and said in fetcher.refusal("https://x.test/a"), said

    class Quick(scrape.Fetcher):
        def get(self, url, listing=False):
            return "<html></html>"
    slow, said = Quick(delay=0), []
    slow.robots["http://x.test"] = scrape.Robots("User-agent: *\nCrawl-delay: 0.1\n", scrape.AGENT)
    scrape.scrape({"start": ["http://x.test/a", "http://x.test/b"], "item": "p", "fields": {"t": "p"}}, slow,
                  said.append)
    assert said == ["robots.txt asks for 0.1 seconds between pages, so the run waits that long"], said
    print(f"robots.txt rules: {len(cases)} addresses read as RFC 9309 says (longest rule, wildcards, $, Allow on a "
          "tie), Crawl-delay, the group naming scrape.py; missing allows all, 403, 503 or no answer allows nothing")


def shop(base, root):
    """A store's JSON, one row per variant over numbered pages, then one row per product, then a JSON API
    whose answer gives the address of the next page."""
    job = {"start": base + "/store/products.json?limit=2", "json": "products", "page_param": "page",
           "rows": "variants", "delay": 0, "numbers": ["price"], "plain": ["description"],
           "fields": {"product": "title", "vendor": "vendor", "tags": "tags", "description": "body_html",
                      "variant": "variants.*.title", "sku": "variants.*.sku", "price": "variants.*.price",
                      "in_stock": "variants.*.available", "image": "images.0.src", "images": "images.*.src"}}
    rows, stats = scrape.scrape(job, log=lambda *a: None)
    variants = [(p, v) for p in SHOP for v in p["variants"] or [None]]
    assert stats == {"pages": 4, "found": len(variants), "duplicates": 0}, stats      # page 4 is empty
    assert [r["sku"] for r in rows] == [v["sku"] if v else "" for _, v in variants], rows
    first = rows[0]
    assert first == {"product": "Runner 1", "vendor": "Acme", "tags": "shoes, size-1",
                     "description": "Light shoe, number 1. Mesh Rubber", "variant": "40", "sku": "R1-40", "price": 51.5,
                     "in_stock": True, "image": "https://cdn.test/1a.jpg",
                     "images": "https://cdn.test/1a.jpg, https://cdn.test/1b.jpg"}, first
    assert [r["in_stock"] for r in rows if r["variant"] == "44"] == [False, False], rows
    assert rows[-1]["product"] == "Runner 6" and rows[-1]["price"] is None, rows[-1]

    out = os.path.join(root, "shop")
    scrape.write(rows, stats, job, out)
    sheet = [[c.value for c in r] for r in load_workbook(out + ".xlsx")["Data"].iter_rows()]
    assert sheet[1][6] == 51.5 and sheet[1][7] is True and len(sheet) == len(variants) + 1, sheet[:2]

    try:
        scrape.scrape(dict(job, rows=None, fields={"product": "title"}), log=lambda *a: None)
        raise AssertionError("a job whose numbers name a missing field ran")
    except ValueError as e:
        assert "price, description" in str(e), e
    per_product = dict(job, rows=None, plain=[], fields={"product": "title", "price": "variants.0.price",
                                                         "skus": "variants.*.sku"})
    rows, stats = scrape.scrape(per_product, log=lambda *a: None)
    assert len(rows) == len(SHOP) and rows[1] == {"product": "Runner 2", "price": 52.5,
                                                   "skus": "R2-40, R2-42, R2-44"}, rows

    api = {"start": base + "/api/one.json", "json": "data.items", "next": "links.next", "delay": 0,
           "fields": {"name": "name", "n": "n"}}
    rows, stats = scrape.scrape(api, log=lambda *a: None)
    assert [(r["name"], r["n"]) for r in rows] == [("A", 1), ("B", 2), ("C", 3)] and stats["pages"] == 2, rows

    said = []
    rows, stats = scrape.scrape(dict(api, start=base + "/shop/page1.html"), log=said.append)
    assert rows == [] and "did not answer with JSON" in said[0], said
    print(f"JSON: {len(variants)} variant rows from a store over 4 pages, prices as numbers, HTML as text, "
          f"{len(SHOP)} product rows, a next link read from the answer, an HTML page refused")


def deep(base, root):
    job = {"start": base + "/deep/list.html", "item": "div.row", "delay": 0, "numbers": ["stock"],
           "fields": {"name": "a", "link": "a@href"},
           "detail": {"link": "link", "fields": {"desc": "p.desc", "stock": "td.stock"}}}
    said = []
    rows, stats = scrape.scrape(job, log=said.append)
    assert stats == {"pages": 1, "found": 3, "duplicates": 0, "details": 1, "details_missing": 2}, stats
    assert rows[0] == {"name": "One", "link": base + "/deep/one.html", "desc": "The first item", "stock": 7}, rows
    assert [(r["desc"], r["stock"]) for r in rows[1:]] == [("", None), ("", None)], rows
    assert len(said) == 2 and "missing.html" in said[0] and "robots.txt" in said[1], said

    out = os.path.join(root, "deep-result")
    scrape.write(rows, stats, job, out)
    wb = load_workbook(out + ".xlsx")
    assert [c.value for c in wb["Data"][1]] == ["name", "link", "desc", "stock"]
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows()}
    assert about["Detail pages read"] == 1 and about["Detail pages not read"] == 2, about
    print("detail pages: 1 read, the missing one and the one robots.txt keeps out left empty, numbers read")


def render(base):
    try:
        import playwright  # noqa: F401
    except ImportError:
        print("render: skipped, Playwright is not installed here")
        return
    job = {"start": base + "/js/list.html", "item": "article.p", "delay": 0,
           "fields": {"name": "h3 a@title", "link": "h3 a@href"}}
    want = [{"name": "Built", "link": base + "/js/a.html"}]
    plain, _ = scrape.scrape(job, log=lambda *a: None)
    assert plain == [], plain
    for extra in ({"render": True}, {"render": True, "wait_for": "article.p"}):
        built, stats = scrape.scrape(dict(job, **extra), log=lambda *a: None)
        assert built == want and stats["pages"] == 1, (extra, built, stats)
    print("render: the page built by JavaScript came out in Chromium, with and without wait_for")


def more(base):
    try:
        import playwright  # noqa: F401
    except ImportError:
        print("scroll and click_more: skipped, Playwright is not installed here")
        return
    job = {"start": base + "/js/scroll.html", "item": "article.p", "delay": 0, "render": True, "scroll_wait": 1,
           "fields": {"name": "h3 a@title"}}
    counts = [len(scrape.scrape(dict(job, **extra), log=lambda *a: None)[0])
              for extra in ({}, {"scroll": 1}, {"scroll": 10})]
    assert counts == [5, 10, 20], counts
    rows, _ = scrape.scrape(dict(job, start=base + "/js/more.html", click_more="#more"), log=lambda *a: None)
    assert [r["name"] for r in rows] == [f"Item {i}" for i in range(1, 21)], rows
    print("scroll and click_more: 5 items without, 10 after one scroll, all 20 by scrolling or clicking to the end")


if __name__ == "__main__":
    numbers()
    robots()
    base, root, server = serve()
    site(base, root)
    shop(base, root)
    deep(base, root)
    render(base)
    more(base)
    server.shutdown()
    print("all good")
