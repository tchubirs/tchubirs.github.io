"""Wedding Budget Planner (Excel + Google Sheets): budget split, vendors and
payments with due-date warnings, guest list with RSVPs, and a countdown.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
wedding = today + dt.timedelta(days=210)
# Typical split of a wedding budget; the couple can change every percentage.
CATS = [("Venue", 26), ("Catering", 24), ("Photo & video", 11), ("Attire & beauty", 8), ("Flowers & decor", 8),
        ("Music & entertainment", 6), ("Rings", 5), ("Stationery", 2), ("Transport", 2), ("Favors & gifts", 2),
        ("Other & buffer", 6)]
C0, C1 = 6, 6 + len(CATS) - 1
V0, V1 = 6, 105
G0, G1 = 6, 305

wb = Workbook()

# ---------------------------------------------------------------- Budget
bu = wb.active
bu.title = "Budget"
sheet_base(bu, "Budget", "Total budget and how it is split. Percentages are a common starting point: change them.",
           [3, 26, 12, 15, 15, 15, 15, 30])
bu["B4"] = "Total budget"
bu["B4"].font = font(12, True, TEAL_D)
style(bu["C4"], True, MONEY, True)
bu["C4"] = 25000
bu.merge_cells("C4:D4")
header(bu, 5, 2, ["Category", "Share %", "Planned", "Committed", "Paid", "Left to pay"])
V = lambda col: f"Vendors!${col}${V0}:${col}${V1}"
for k, (name, pct) in enumerate(CATS):
    r = C0 + k
    style(bu.cell(r, 2, name), True, bold=True)
    style(bu.cell(r, 3, pct), True, "0", align="center")
    style(bu.cell(r, 4, f"=$C$4*C{r}/100"), False, MONEY)
    style(bu.cell(r, 5, f'=SUMIFS({V("D")},{V("C")},B{r})'), False, MONEY)
    style(bu.cell(r, 6, f'=SUMIFS({V("E")},{V("C")},B{r})'), False, MONEY)
    style(bu.cell(r, 7, f"=E{r}-F{r}"), False, MONEY)
    c = style(bu.cell(r, 8, f'=IF(D{r}=0,"",{bar(f"E{r}/D{r}", 15)}&"  "&TEXT(E{r}/D{r},"0%"))'), False)
    c.font = Font(name=F, size=11, color=TEAL)
t = C1 + 1
bu.cell(t, 2, "Total").font = font(11, True)
for col in "CDEFG":
    style(bu[f"{col}{t}"], False, "0" if col == "C" else MONEY, True)
    bu[f"{col}{t}"] = f"=SUM({col}{C0}:{col}{C1})"
bu.conditional_formatting.add(f"C{t}", FormulaRule(formula=[f"$C${t}<>100"], fill=fill(BAD)))
bu.conditional_formatting.add(f"E{C0}:E{C1}", FormulaRule(formula=[f"$E{C0}>$D{C0}"], fill=fill(BAD),
                                                          font=Font(name=F, bold=True, color="9B1C1C")))
bu[f"B{t + 1}"] = "Shares should add up to 100 (the total turns red if not). Committed turns red above plan."
bu[f"B{t + 1}"].font = font(10, color=MUTED, italic=True)
CAT_RANGE = f"Budget!$B${C0}:$B${C1}"

# ---------------------------------------------------------------- Vendors
ve = wb.create_sheet("Vendors")
sheet_base(ve, "Vendors & payments", "One line per vendor: what they quoted, what you paid, when the rest is due.",
           [3, 24, 22, 13, 13, 14, 13, 28], rows=V1 + 2)
header(ve, 5, 2, ["Vendor", "Category", "Quoted", "Paid so far", "Next payment due", "Left", "Status"])
d = lambda k: today + dt.timedelta(days=k)
sample = [("Château venue", "Venue", 6200, 2000, d(12)), ("Chef & Co", "Catering", 5900, 1500, d(45)),
          ("Lens & Light", "Photo & video", 2600, 800, d(2)), ("Atelier Rose", "Attire & beauty", 1800, 1800, None),
          ("Bloom Studio", "Flowers & decor", 1700, 300, d(-3)), ("DJ Nova", "Music & entertainment", 1100, 300, d(80)),
          ("Maison Or", "Rings", 1300, 1300, None), ("Paper & Ink", "Stationery", 420, 200, d(20))]
for r in range(V0, V1 + 1):
    s = sample[r - V0] if r - V0 < len(sample) else (None,) * 5
    style(ve.cell(r, 2, s[0]), True)
    style(ve.cell(r, 3, s[1]), True)
    style(ve.cell(r, 4, s[2]), True, MONEY)
    style(ve.cell(r, 5, s[3]), True, MONEY)
    style(ve.cell(r, 6, s[4]), True, DATE)
    style(ve.cell(r, 7, f'=IF(B{r}="","",D{r}-E{r})'), False, MONEY)
    late, left = f"TODAY()-F{r}", f"F{r}-TODAY()"
    st = (f'=IF(B{r}="","",IF(G{r}<=0,"Paid in full",IF(F{r}="","Balance due, no date set",'
          f'IF(F{r}<TODAY(),"Overdue by "&{late}&IF({late}=1," day"," days"),IF(F{r}=TODAY(),"Due today",'
          f'IF({left}<=14,"Due in "&{left}&IF({left}=1," day"," days"),"Next payment in "&{left}&" days"))))))')
    style(ve.cell(r, 8, st), False)
dv = DataValidation(type="list", formula1=f"={CAT_RANGE}", allow_blank=True)
ve.add_data_validation(dv)
dv.add(f"C{V0}:C{V1}")
ve.freeze_panes = "B6"
rng = f"H{V0}:H{V1}"
ve.conditional_formatting.add(rng, FormulaRule(formula=[f'$H{V0}="Paid in full"'], fill=fill(OK)))
ve.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($H{V0},7)="Overdue"'], fill=fill(BAD),
                                               font=Font(name=F, bold=True, color="9B1C1C")))
ve.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT($H{V0},6)="Due in",$H{V0}="Due today")'],
                                               fill=fill(WARN)))

# ---------------------------------------------------------------- Guests
gu = wb.create_sheet("Guests")
sheet_base(gu, "Guest list", "One line per guest. RSVP and meal choices add up on the Dashboard.",
           [3, 26, 12, 12, 10, 16, 8, 26], rows=G1 + 2)
header(gu, 5, 2, ["Guest", "Side", "RSVP", "Plus one", "Meal", "Table", "Note"])
names = ["Ana & Luca", "Chloé Martin", "Tom Becker", "Inès Dupont", "Rafael Costa", "Mia Laurent", "Hugo Bernard",
         "Sara Klein", "Noah Petit", "Lea Moreau", "Jonas Weber", "Emma Roux"]
rsvp = ["Yes", "Yes", "Pending", "Yes", "No", "Yes", "Pending", "Yes", "Yes", "Pending", "Yes", "Yes"]
meals = ["Fish", "Veggie", "", "Meat", "", "Meat", "", "Fish", "Meat", "", "Veggie", "Meat"]
for r in range(G0, G1 + 1):
    k = r - G0
    have = k < len(names)
    style(gu.cell(r, 2, names[k] if have else None), True)
    style(gu.cell(r, 3, ["Both", "Partner A", "Partner B"][k % 3] if have else None), True, align="center")
    style(gu.cell(r, 4, rsvp[k] if have else None), True, align="center")
    style(gu.cell(r, 5, ("Yes" if k % 4 == 0 else "No") if have else None), True, align="center")
    style(gu.cell(r, 6, meals[k] if have else None), True, align="center")
    style(gu.cell(r, 7, (k // 4) + 1 if have and rsvp[k] == "Yes" else None), True, "0", align="center")
    style(gu.cell(r, 8), True)
for col, options in (("C", ["Both", "Partner A", "Partner B"]), ("D", ["Yes", "No", "Pending"]),
                     ("E", ["Yes", "No"]), ("F", ["Meat", "Fish", "Veggie", "Kids"])):
    v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    gu.add_data_validation(v)
    v.add(f"{col}{G0}:{col}{G1}")
gu.conditional_formatting.add(f"D{G0}:D{G1}", FormulaRule(formula=[f'$D{G0}="Yes"'], fill=fill(OK)))
gu.conditional_formatting.add(f"D{G0}:D{G1}", FormulaRule(formula=[f'$D{G0}="Pending"'], fill=fill(WARN)))
gu.conditional_formatting.add(f"D{G0}:D{G1}", FormulaRule(formula=[f'$D{G0}="No"'], fill=fill("E4E8E8")))
gu.freeze_panes = "B6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Wedding dashboard", "Set the date below. Everything else fills itself in.",
           [3, 24, 16, 16, 16, 16, 3, 40])
db["B4"] = "Wedding date"
db["B4"].font = font(12, True, TEAL_D)
style(db["C4"], True, DATE, True)
db["C4"] = wedding
G = lambda col: f"Guests!${col}${G0}:${col}${G1}"
tiles = [
    ("B", "Total budget", "=Budget!$C$4", TEAL),
    ("C", "Committed", f"=SUM({V('D')})", "B4541F"),
    ("D", "Paid so far", f"=SUM({V('E')})", TEAL),
    ("E", "Still to pay", f"=SUM({V('G')})", "B4541F"),
    ("F", "Unallocated", "=B7-C7", TEAL),
]
for col, label, formula, colour in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = MONEY
    db[f"{col}7"].font = font(17, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db["H6"] = "Days to go"
db["H7"] = '=IF($C$4="","",MAX(0,$C$4-TODAY()))'
for a, size in (("H6", 11), ("H7", 28)):
    db[a].font = font(size, True, "FFFFFF")
    db[a].fill = fill(TEAL_D)
    db[a].alignment = Alignment(horizontal="center", vertical="center")
db["H7"].number_format = "0"
db.row_dimensions[7].height = 44

header(db, 9, 2, ["Guests", "Invited", "Yes", "Pending", "No"])
style(db["B10"], False, bold=True)
db["B10"] = "People (incl. plus ones)"
db["C10"] = f'=COUNTIFS({G("B")},"<>")+COUNTIFS({G("B")},"<>",{G("E")},"Yes")'
db["D10"] = f'=COUNTIFS({G("D")},"Yes")+COUNTIFS({G("D")},"Yes",{G("E")},"Yes")'
db["E10"] = f'=COUNTIFS({G("D")},"Pending")+COUNTIFS({G("D")},"Pending",{G("E")},"Yes")'
db["F10"] = f'=COUNTIFS({G("D")},"No")+COUNTIFS({G("D")},"No",{G("E")},"Yes")'
for col in "CDEF":
    style(db[f"{col}10"], False, "0", align="center")
style(db["B11"], False, bold=True)
db["B11"] = "Meals (guests who said yes)"
for i, meal in enumerate(["Meat", "Fish", "Veggie", "Kids"]):
    c = db.cell(11, 3 + i, f'="{meal}: "&COUNTIFS({G("D")},"Yes",{G("F")},"{meal}")')
    style(c, False, align="center")

db["H9"] = "Heads-up"
db["H9"].font = font(11, True, "FFFFFF")
db["H9"].fill = fill(TEAL)
late_n = f'COUNTIFS({V("H")},"Overdue*")'
soon_n = f'(COUNTIFS({V("H")},"Due in*")+COUNTIFS({V("H")},"Due today"))'
soon_sum = f'(SUMIFS({V("G")},{V("H")},"Due in*")+SUMIFS({V("G")},{V("H")},"Due today"))'
wait_n = f'COUNTIFS({G("D")},"Pending")'
notes = [
    (f'=IF({late_n}=0,"No overdue payments",IF({late_n}=1,"1 payment is overdue",{late_n}&" payments are overdue"))', BAD),
    (f'=IF({soon_n}=0,"Nothing due in the next 14 days",IF({soon_n}=1,"1 payment",{soon_n}&" payments")'
     f'&" due in the next 14 days: "&FIXED({soon_sum},2))', WARN),
    (f'=IF({wait_n}=0,"Every guest has answered",IF({wait_n}=1,"1 guest has not answered yet",'
     f'{wait_n}&" guests have not answered yet"))', INFO),
    ('=IF(C7>B7,"Over budget by "&FIXED(C7-B7,2),"Within budget")', OK),
]
for k, (formula, colour) in enumerate(notes):
    c = db.cell(10 + k, 8, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.border = box
for cell, start in (("H10", "No"), ("H11", "Nothing"), ("H12", "Every")):
    db.conditional_formatting.add(cell, FormulaRule(formula=[f'LEFT({cell},{len(start)})="{start}"'], fill=fill(OK)))
db.conditional_formatting.add("H13", FormulaRule(formula=['LEFT(H13,4)="Over"'], fill=fill(BAD)))
db["H9"].alignment = Alignment(vertical="center")

header(db, 14, 2, ["Category", "Planned", "Committed", "Used"])
for k in range(len(CATS)):
    r = 15 + k
    style(db.cell(r, 2, f"=Budget!B{C0 + k}"), False, bold=True)
    style(db.cell(r, 3, f"=Budget!D{C0 + k}"), False, MONEY)
    style(db.cell(r, 4, f"=Budget!E{C0 + k}"), False, MONEY)
    c = style(db.cell(r, 5, f'=IF(C{r}=0,"",{bar(f"D{r}/C{r}", 15)}&"  "&TEXT(D{r}/C{r},"0%"))'), False)
    c.font = Font(name=F, size=11, color=TEAL)
db.conditional_formatting.add(f"D15:D{14 + len(CATS)}", FormulaRule(formula=["$D15>$C15"], fill=fill(BAD)))
db.column_dimensions["E"].width = 26

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 100])
steps = [
    ("1", "Dashboard: type your wedding date and the countdown starts."),
    ("2", "Budget: type your total. Change the percentages if your priorities differ (they add up to 100)."),
    ("3", "Vendors: add each vendor once you have a quote, with the amount, what you paid and the next due date."),
    ("4", "Guests: one line per guest. Update the RSVP and meal as answers come in."),
    ("5", "The Dashboard then shows what is booked, paid and still to pay, and your guest numbers."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
st["C17"] = "The example vendors and guests show how it works. Delete them and add your own."
st["C17"].font = font(11, color=MUTED, italic=True)
st["C19"] = "Works in Google Sheets and Microsoft Excel, in any currency."
st["C19"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Wedding-Budget-Planner.xlsx")
wb.save(out)
print("saved", out)
