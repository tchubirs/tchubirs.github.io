"""Shared look for the Etsy spreadsheets: calm paper, teal headers, yellow inputs."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
TEAL, TEAL_D, INK, MUTED = "1F6F78", "15525A", "1E2A2F", "6B7B80"
PAPER, INPUT, LINE = "F6F4EF", "FFF4C2", "D9D4C7"
OK, WARN, BAD, INFO = "D8F0DC", "FFE2B8", "F9C9C4", "E3E8FF"
MONEY = '#,##0.00'
DATE = "DD MMM YYYY"

thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)

def font(size=11, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)

def sheet_base(ws, title, subtitle, widths, rows=80):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(1, rows):
        for c in range(1, len(widths) + 2):
            ws.cell(r, c).fill = fill(PAPER)
    ws["B2"] = title
    ws["B2"].font = font(22, True, TEAL_D)
    ws["B3"] = subtitle
    ws["B3"].font = font(11, color=MUTED, italic=True)
    ws.row_dimensions[2].height = 32
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

def header(ws, row, col, labels):
    for i, label in enumerate(labels):
        c = ws.cell(row, col + i, label)
        c.font = font(11, True, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[row].height = 30

def style(c, input_cell=False, fmt=None, bold=False, align=None):
    c.fill = fill(INPUT if input_cell else "FFFFFF")
    c.border = box
    c.font = font(11, bold)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    return c

def bar(fraction_expr, width=20):
    """Text progress bar that works in Excel and Google Sheets."""
    n = f"MIN({width},MAX(0,ROUND(({fraction_expr})*{width},0)))"
    return f'REPT("█",{n})&REPT("░",{width}-{n})'
