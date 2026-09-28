"""Vegetable Garden Planner and Harvest Tracker (Excel + Google Sheets).

Crops with their own sowing, planting out and harvest months, a calendar to print, what was planted in each bed and
when, a warning when a bed gets the same crop family as last year, the harvest valued at shop prices, watering and
other jobs, and what the garden cost. Only functions both apps have. The months are the ones the gardener types:
nothing here promises a date.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
Y = today.year
C0, C1 = 6, 65         # crops
B0, B1 = 6, 35         # beds
L0, L1 = 6, 505        # plantings
H0, H1 = 6, 1505       # harvest
J0, J1 = 6, 2005       # jobs
K0, K1 = 6, 505        # costs
NCAL = 40              # crops on the calendar
NAME, YEAR, CUR, WATER = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6", "Settings!$C$7"
MONTHS = "Settings!$H$1:$H$12"
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December"]
FAMILIES = ["Nightshades", "Cucurbits", "Legumes", "Brassicas", "Roots", "Alliums", "Leafy greens", "Herbs", "Fruit",
            "Other"]
CR = lambda c: f"Crops!${c}${C0}:${c}${C1}"
CR1 = lambda c: f"Crops!${c}$1:${c}${C1}"
BD = lambda c: f"Beds!${c}${B0}:${c}${B1}"
PL = lambda c: f"Plantings!${c}${L0}:${c}${L1}"
PL1 = lambda c: f"Plantings!${c}$1:${c}${L1}"
HV = lambda c: f"Harvest!${c}${H0}:${c}${H1}"
JB = lambda c: f"Jobs!${c}${J0}:${c}${J1}"
CO = lambda c: f"Costs!${c}${K0}:${c}${K1}"
crop_of = lambda cell, c: f"INDEX({CR(c)},MATCH({cell},{CR('B')},0))"


def in_range(f, t, m):
    """1 when month m is between the months f and t (numbers, 0 when missing), also across the new year."""
    return f"IF(OR({f}=0,{t}=0),0,IF({f}<={t},IF(AND({m}>={f},{m}<={t}),1,0),IF(OR({m}>={f},{m}<={t}),1,0)))"


# ---------------------------------------------------------------- example data (a made-up back garden)
rng = random.Random(43)
D = lambda y, m, d: dt.date(y, m, d)
beds = [("Bed 1", "1.2 x 2.4 m", "Full sun", None), ("Bed 2", "1.2 x 2.4 m", "Full sun", None),
        ("Bed 3", "1.2 x 2.4 m", "Full sun", "By the fence"), ("Bed 4", "1.2 x 2.4 m", "Part shade", None),
        ("Greenhouse", "2 x 3 m", "Full sun", None), ("Patio pots", "8 pots", "Full sun", None),
        ("Herb bed", "1 x 1 m", "Part shade", "By the back door")]
crops = [  # crop, family, unit, shop price, days to harvest, sow, plant out, harvest, seeds left
    ("Tomatoes", "Nightshades", "kg", 5.5, 75, ("Mar", "Apr"), ("May", "Jun"), ("Jul", "Oct"), "Half a packet"),
    ("Peppers", "Nightshades", "piece", 1.2, 80, ("Feb", "Mar"), ("May", "Jun"), ("Aug", "Oct"), "1 packet"),
    ("Potatoes", "Nightshades", "kg", 1.8, 100, ("Mar", "Apr"), (None, None), ("Jul", "Sep"), None),
    ("Zucchini", "Cucurbits", "kg", 3.5, 55, ("Apr", "May"), ("May", "Jun"), ("Jul", "Sep"), "6 seeds"),
    ("Cucumbers", "Cucurbits", "piece", 1.0, 60, ("Apr", "May"), ("May", "Jun"), ("Jul", "Sep"), "1 packet"),
    ("Runner beans", "Legumes", "kg", 6.0, 70, ("Apr", "May"), ("Jun", "Jun"), ("Jul", "Oct"), "Half a packet"),
    ("Peas", "Legumes", "kg", 7.0, 65, ("Mar", "May"), (None, None), ("Jun", "Jul"), "None left"),
    ("Lettuce", "Leafy greens", "head", 1.5, 45, ("Mar", "Sep"), ("Apr", "Sep"), ("May", "Oct"), "2 packets"),
    ("Spinach", "Leafy greens", "bunch", 2.0, 40, ("Aug", "Apr"), (None, None), ("Oct", "Jun"), "1 packet"),
    ("Carrots", "Roots", "kg", 2.0, 75, ("Apr", "Jun"), (None, None), ("Jul", "Nov"), "1 packet"),
    ("Beetroot", "Roots", "bunch", 2.5, 60, ("Apr", "Jul"), (None, None), ("Jun", "Oct"), "Half a packet"),
    ("Onions", "Alliums", "kg", 2.2, 120, ("Mar", "Apr"), (None, None), ("Aug", "Sep"), None),
    ("Garlic", "Alliums", "head", 0.8, 240, ("Oct", "Nov"), (None, None), ("Jun", "Jul"), "2 bulbs"),
    ("Kale", "Brassicas", "bunch", 2.5, 60, ("Apr", "Jun"), ("May", "Jul"), ("Sep", "Feb"), "1 packet"),
    ("Basil", "Herbs", "bunch", 2.0, 30, ("Apr", "May"), ("May", "Jun"), ("Jun", "Sep"), "Half a packet"),
    ("Strawberries", "Fruit", "kg", 9.0, None, (None, None), ("Apr", "Apr"), ("Jun", "Jul"), None),
]
plantings = [  # date, crop, bed, how, how many, cleared, note (sorted by date below)
    (D(Y - 1, 3, 20), "Potatoes", "Bed 2", "Sown outside", 20, D(Y - 1, 8, 20), "Seed potatoes"),
    (D(Y - 1, 4, 2), "Onions", "Bed 4", "Sown outside", 60, D(Y - 1, 8, 28), "Sets"),
    (D(Y - 1, 5, 20), "Tomatoes", "Bed 1", "Planted out", 8, D(Y - 1, 10, 15), None),
    (D(Y - 1, 6, 1), "Runner beans", "Bed 3", "Planted out", 12, D(Y - 1, 10, 20), None),
    (D(Y - 1, 4, 25), "Carrots", "Bed 3", "Sown outside", 1, D(Y - 1, 11, 10), "One row"),
    (D(Y - 1, 10, 20), "Garlic", "Bed 3", "Sown outside", 24, D(Y, 7, 5), None),
    (D(Y, 3, 15), "Tomatoes", "Greenhouse", "Sown inside", 12, D(Y, 5, 18), "Seed trays, planted out in May"),
    (D(Y, 3, 22), "Peas", "Bed 1", "Sown outside", 2, D(Y, 7, 25), "Two rows"),
    (D(Y, 4, 5), "Potatoes", "Bed 4", "Sown outside", 20, D(Y, 8, 30), "Seed potatoes"),
    (D(Y, 4, 12), "Onions", "Bed 2", "Sown outside", 60, D(Y, 9, 5), "Sets"),
    (D(Y, 4, 20), "Carrots", "Bed 2", "Sown outside", 1, None, "One row"),
    (D(Y, 4, 26), "Lettuce", "Herb bed", "Sown outside", 1, D(Y, 7, 10), None),
    (D(Y, 5, 10), "Peppers", "Bed 2", "Planted out", 6, None, None),
    (D(Y, 5, 18), "Tomatoes", "Greenhouse", "Planted out", 8, None, None),
    (D(Y, 5, 24), "Zucchini", "Bed 3", "Planted out", 3, None, None),
    (D(Y, 5, 24), "Cucumbers", "Greenhouse", "Planted out", 4, None, None),
    (D(Y, 5, 30), "Basil", "Patio pots", "Planted out", 6, None, None),
    (D(Y, 6, 2), "Runner beans", "Bed 1", "Planted out", 12, None, None),
    (D(Y, 6, 14), "Beetroot", "Bed 4", "Sown outside", 1, None, "After the potatoes come out"),
    (D(Y, 6, 20), "Kale", "Bed 3", "Planted out", 6, None, None),
    (D(Y, 8, 10), "Lettuce", "Herb bed", "Sown outside", 1, None, "Second sowing"),
    (D(Y, 4, 30), "Strawberries", "Patio pots", "Bought plants", 6, None, None),
    (D(Y, 9, 12), "Spinach", "Bed 1", "Sown outside", 1, None, "Under the beans"),
]
plantings.sort(key=lambda p: p[0])
harvest = []   # date, crop, bed, amount, note
season = [  # crop, bed, first pick, last pick, days between picks, amount per pick
    ("Tomatoes", "Bed 1", D(Y - 1, 7, 20), D(Y - 1, 10, 5), 5, (0.6, 1.8)),
    ("Runner beans", "Bed 3", D(Y - 1, 7, 25), D(Y - 1, 10, 10), 6, (0.3, 0.9)),
    ("Potatoes", "Bed 2", D(Y - 1, 7, 15), D(Y - 1, 8, 20), 9, (2.0, 4.0)),
    ("Onions", "Bed 4", D(Y - 1, 8, 25), D(Y - 1, 8, 28), 3, (4.0, 5.0)),
    ("Carrots", "Bed 3", D(Y - 1, 7, 20), D(Y - 1, 11, 5), 14, (0.4, 1.0)),
    ("Garlic", "Bed 3", D(Y, 7, 1), D(Y, 7, 5), 4, (20, 24)),
    ("Peas", "Bed 1", D(Y, 6, 10), D(Y, 7, 20), 5, (0.2, 0.6)),
    ("Lettuce", "Herb bed", D(Y, 6, 5), D(Y, 7, 8), 6, (1, 3)),
    ("Potatoes", "Bed 4", D(Y, 7, 20), D(Y, 8, 30), 8, (2.0, 4.5)),
    ("Onions", "Bed 2", D(Y, 9, 1), D(Y, 9, 5), 4, (4.0, 6.0)),
    ("Tomatoes", "Greenhouse", D(Y, 7, 12), today - dt.timedelta(days=2), 4, (0.5, 2.2)),
    ("Zucchini", "Bed 3", D(Y, 7, 8), today - dt.timedelta(days=4), 4, (0.4, 1.6)),
    ("Cucumbers", "Greenhouse", D(Y, 7, 15), D(Y, 9, 15), 5, (1, 3)),
    ("Basil", "Patio pots", D(Y, 6, 25), today - dt.timedelta(days=6), 9, (1, 2)),
    ("Runner beans", "Bed 1", D(Y, 7, 28), today - dt.timedelta(days=1), 5, (0.3, 1.1)),
    ("Peppers", "Bed 2", D(Y, 8, 20), today - dt.timedelta(days=5), 7, (2, 5)),
    ("Carrots", "Bed 2", D(Y, 7, 25), today - dt.timedelta(days=8), 12, (0.4, 0.9)),
    ("Beetroot", "Bed 4", D(Y, 8, 20), today - dt.timedelta(days=3), 10, (1, 2)),
    ("Kale", "Bed 3", D(Y, 9, 5), today - dt.timedelta(days=2), 7, (1, 2)),
    ("Lettuce", "Herb bed", D(Y, 9, 20), today - dt.timedelta(days=1), 4, (1, 2)),
]
whole = {"piece", "head", "bunch"}
for crop, bed, first, last, every, (lo, hi) in season:
    unit = next(c[2] for c in crops if c[0] == crop)
    d = first
    while d <= last:
        amt = rng.randint(int(lo), int(hi)) if unit in whole else round(rng.uniform(lo, hi), 2)
        harvest.append([d, crop, bed, amt, None])
        d += dt.timedelta(days=every + rng.randint(-1, 1))
harvest.append([D(Y, 6, 28), "Strawberries", "Patio pots", 1.4, "The first big pick"])
harvest.sort(key=lambda h: h[0])
jobs = []      # date, bed, job, note
for yr in (Y - 1, Y):
    d = D(yr, 5, 1)
    end = D(yr, 9, 30) if yr < Y else today - dt.timedelta(days=5)
    while d <= end:
        jobs.append([d, "All beds", "Water", None])
        d += dt.timedelta(days=rng.choice([2, 3, 3, 4]))
    d = D(yr, 6, 10)
    while d <= min(end, D(yr, 9, 20)):
        jobs.append([d, "Greenhouse", "Feed", "Tomato feed"])
        d += dt.timedelta(days=7)
for d, bed, job, note in ((D(Y, 3, 8), "Bed 1", "Weed", None), (D(Y, 3, 9), "Bed 2", "Mulch", "Compost"),
                          (D(Y, 4, 18), "Bed 4", "Weed", None), (D(Y, 6, 6), "Bed 3", "Pest check", "Slugs"),
                          (D(Y, 7, 2), "Bed 1", "Weed", None), (D(Y, 8, 14), "Greenhouse", "Prune", "Side shoots"),
                          (today - dt.timedelta(days=3), "All beds", "Water", None),
                          (today - dt.timedelta(days=1), "Greenhouse", "Water", None),
                          (today - dt.timedelta(days=1), "Patio pots", "Water", None),
                          (today - dt.timedelta(days=1), "Bed 1", "Water", None),
                          (today - dt.timedelta(days=1), "Bed 3", "Water", None),
                          (today - dt.timedelta(days=1), "Herb bed", "Water", None),
                          (today - dt.timedelta(days=2), "Bed 4", "Weed", None)):
    jobs.append([d, bed, job, note])
jobs.sort(key=lambda j: j[0])
costs = [  # date, what, kind, amount
    (D(Y - 1, 2, 20), "Seed packets", "Seeds", 24.50), (D(Y - 1, 3, 10), "Compost, 6 bags", "Soil and compost", 36.00),
    (D(Y - 1, 3, 20), "Seed potatoes and onion sets", "Plants", 14.00), (D(Y - 1, 5, 15), "Tomato plants", "Plants", 12.00),
    (D(Y - 1, 6, 1), "Bean poles", "Tools", 9.50), (D(Y, 2, 14), "Seed packets", "Seeds", 31.20),
    (D(Y, 2, 28), "Seed trays and pots", "Tools", 18.40), (D(Y, 3, 7), "Compost, 8 bags", "Soil and compost", 48.00),
    (D(Y, 3, 20), "Seed potatoes and onion sets", "Plants", 16.50), (D(Y, 4, 30), "Strawberry plants", "Plants", 15.00),
    (D(Y, 5, 12), "Tomato feed", "Other", 8.99), (D(Y, 5, 20), "Netting", "Tools", 12.00),
    (D(Y, 7, 1), "Hose connector", "Tools", 6.50), (D(Y, 8, 3), "Water bill, garden share", "Water", 14.00),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def check_fill(ws, rng_, first):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}<>""'], fill=fill(BAD)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The garden's name, the year shown, and when a bed needs water again.",
           [3, 30, 18, 4, 58], rows=14)
for r, lab, value, fmt, note in (
        (4, "Garden", "Back garden", None, "On the Dashboard and the calendar."),
        (5, "Year", "=YEAR(TODAY())", "0", "The year on the Dashboard. Type another to look back."),
        (6, "Currency", "USD", None, "For the shop prices and the costs."),
        (7, "Water again after (days)", 3, "0", "A bed shows Water when it has not been watered for this long.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, r == 4, "center")
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
for k, m in enumerate(MON):
    se[f"H{k + 1}"] = m                                                                        # H: month names
se.column_dimensions["H"].hidden = True

# ---------------------------------------------------------------- Crops
cr = wb.create_sheet("Crops")
sheet_base(cr, "Crops", "Each crop once, with the months you plan from your seed packets and your own experience.",
           [3, 18, 13, 8, 10, 9, 7, 7, 7, 7, 7, 7, 16, 9, 10, 11, 20], rows=C1 + 2)
header(cr, 5, 2, ["Crop", "Family", "Unit", "Shop price", "Days to harvest", "Sow from", "Sow to", "Plant out from",
                  "Plant out to", "Harvest from", "Harvest to", "Seeds left", "Planted", "Picked", "Value", "Check"])
now = "MONTH(TODAY())"
for r in range(C0, C1 + 1):
    x = crops[r - C0] if r - C0 < len(crops) else (None, None, None, None, None, (None, None), (None, None),
                                                      (None, None), None)
    vals = [x[0], x[1], x[2], x[3], x[4], *x[5], *x[6], *x[7], x[8]]
    for c, v in zip(range(2, 14), vals):
        fmt = MONEY if c == 5 else ("0" if c == 6 else None)
        style(cr.cell(r, c, v), True, fmt, c == 2, "center" if c in (3, 4, 6, 7, 8, 9, 10, 11, 12) else None)
    style(cr.cell(r, 14, f'=IF(B{r}="","",SUMIFS({PL("S")},{PL("C")},B{r},{PL("N")},{YEAR}))'), False, "0",
          align="center")
    style(cr.cell(r, 15, f'=IF(B{r}="","",SUMIFS({HV("E")},{HV("C")},B{r},{HV("J")},{YEAR}))'), False, "#,##0.##",
          align="center")
    style(cr.cell(r, 16, f'=IF(B{r}="","",SUMIFS({HV("H")},{HV("C")},B{r},{HV("J")},{YEAR}))'), False, MONEY, True)
    style(cr.cell(r, 17, f'=IF(B{r}="","",IF(COUNTIF({CR("B")},B{r})>1,"Listed twice",'
                         f'IF(AND(E{r}<>"",NOT(ISNUMBER(E{r}))),"Price not a number",'
                         f'IF(AND(F{r}<>"",NOT(ISNUMBER(F{r}))),"Days not a number",'
                         f'IF(OR((G{r}="")<>(H{r}=""),(I{r}="")<>(J{r}=""),(K{r}="")<>(L{r}="")),"A month is missing",'
                         f'"")))))'), False, align="center")
    for k, c in enumerate("GHIJKL"):                                                           # R to W: months
        cr.cell(r, 18 + k).value = f"=IFERROR(MATCH({c}{r},{MONTHS},0),0)"
    cr[f"X{r}"] = f'=IF(AND(B{r}<>"",{in_range(f"R{r}", f"S{r}", now)}=1),ROW(),"")'           # X: sow now
    cr[f"Y{r}"] = f'=IF(AND(B{r}<>"",{in_range(f"T{r}", f"U{r}", now)}=1),ROW(),"")'           # Y: plant out now
    cr[f"Z{r}"] = f'=IF(AND(B{r}<>"",{in_range(f"V{r}", f"W{r}", now)}=1),ROW(),"")'           # Z: harvest now
    cr[f"AA{r}"] = f'=IF(OR(B{r}="",N(O{r})<=0),"",ROUND(N(P{r})*100,0)*1000+1000-ROW())'      # AA: top crops
for c in ["R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA"]:
    cr.column_dimensions[c].hidden = True
check_fill(cr, f"Q{C0}:Q{C1}", f"Q{C0}")
dropdown(cr, '"' + ",".join(FAMILIES) + '"', f"C{C0}:C{C1}", strict=False)
dropdown(cr, '"kg,lb,g,oz,piece,head,bunch"', f"D{C0}:D{C1}", strict=False)
dropdown(cr, f"={MONTHS}", f"G{C0}:L{C1}")
cr.freeze_panes = "C6"

# ---------------------------------------------------------------- Beds
bd = wb.create_sheet("Beds")
sheet_base(bd, "Beds", "Each bed, pot group or greenhouse once. What grows in it and the last watering come by "
           "themselves.", [3, 18, 14, 12, 22, 10, 16, 14, 8, 10, 12], rows=B1 + 2)
header(bd, 5, 2, ["Bed", "Size", "Sun", "Note", "Growing", "Latest planting", "Last watered", "Days", "Status",
                  "Picked this year"])
for r in range(B0, B1 + 1):
    x = beds[r - B0] if r - B0 < len(beds) else (None,) * 4
    for c, v in zip(range(2, 6), x):
        style(bd.cell(r, c, v), True, bold=c == 2, align="center" if c in (3, 4) else None)
    style(bd.cell(r, 6, f'=IF(B{r}="","",SUMIFS({PL("Q")},{PL("D")},B{r}))'), False, "0", align="center")
    style(bd.cell(r, 7, f'=IF(OR(B{r}="",N(F{r})=0),"",INDEX({PL1("C")},'
                        f'MOD(SUMPRODUCT(MAX(({PL("D")}=B{r})*{PL("R")})),1000)))'), False, align="center")
    last_w = f'MAX(SUMPRODUCT(MAX(({JB("C")}=B{r})*({JB("D")}="Water")*{JB("G")})),MAX({JB("H")}))'
    style(bd.cell(r, 8, f'=IF(B{r}="","",IF({last_w}=0,"",{last_w}))'), False, DATE, align="center")
    style(bd.cell(r, 9, f'=IF(OR(B{r}="",H{r}=""),"",TODAY()-H{r})'), False, "0", align="center")
    style(bd.cell(r, 10, f'=IF(B{r}="","",IF(N(F{r})=0,"Empty",IF(AND(N({WATER})>0,OR(H{r}="",'
                         f'N(I{r})>=N({WATER}))),"Water","OK")))'), False, bold=True, align="center")
    style(bd.cell(r, 11, f'=IF(B{r}="","",SUMIFS({HV("H")},{HV("D")},B{r},{HV("J")},{YEAR}))'), False, MONEY)
bd.conditional_formatting.add(f"J{B0}:J{B1}", FormulaRule(formula=[f'J{B0}="Water"'], fill=fill(WARN)))
bd.conditional_formatting.add(f"J{B0}:J{B1}", FormulaRule(formula=[f'J{B0}="OK"'], fill=fill(OK)))
dropdown(bd, '"Full sun,Part shade,Shade"', f"D{B0}:D{B1}", strict=False)
bd.freeze_panes = "C6"

# ---------------------------------------------------------------- Plantings
pl = wb.create_sheet("Plantings")
sheet_base(pl, "Plantings", "Each time you sow or plant something, with the bed. A date under Cleared once it is "
           "finished.", [3, 13, 16, 14, 14, 9, 13, 22, 13, 12, 9, 20, 26], rows=L1 + 2)
header(pl, 5, 2, ["Date", "Crop", "Bed", "How", "How many", "Cleared", "Note", "Ready about", "Status", "Days",
                  "Check", "Rotation"])
for r in range(L0, L1 + 1):
    x = plantings[r - L0] if r - L0 < len(plantings) else (None,) * 7
    for c, v, fmt in zip(range(2, 9), x, (DATE, None, None, None, "0", DATE, None)):
        style(pl.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 5, 6, 7) else None)
    days = f"N(IFERROR({crop_of(f'C{r}', 'F')},0))"
    dated = f'AND(C{r}<>"",ISNUMBER(B{r}))'
    style(pl.cell(r, 9, f'=IF(NOT({dated}),"",IF({days}>0,INT(B{r})+{days},""))'), False, DATE, align="center")
    style(pl.cell(r, 10, f'=IF(NOT({dated}),"",IF(ISNUMBER(G{r}),"Done",IF(INT(B{r})>TODAY(),"Planned",'
                         f'IF(COUNTIFS({HV("C")},C{r},{HV("D")},D{r}&"",{HV("L")},">="&INT(B{r}))>0,"Picking",'
                         f'"Growing"))))'), False, bold=True, align="center")
    style(pl.cell(r, 11, f'=IF(OR(NOT({dated}),J{r}="Planned"),"",IF(ISNUMBER(G{r}),INT(G{r}),TODAY())-INT(B{r}))'),
          False, "0", align="center")
    style(pl.cell(r, 12, f'=IF(AND(B{r}="",C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                         f'IF(C{r}="","Needs a crop",IF(COUNTIF({CR("B")},C{r})=0,"Not on Crops",'
                         f'IF(AND(D{r}<>"",COUNTIF({BD("B")},D{r})=0),"Not on Beds",'
                         f'IF(AND(ISNUMBER(G{r}),G{r}<B{r}),"Cleared before it was planted",""))))))'), False,
          align="center")
    style(pl.cell(r, 13, f'=IF(OR(L{r}<>"",D{r}="",O{r}=""),"",IF(COUNTIFS({PL("D")},D{r},{PL("O")},O{r},'
                         f'{PL("N")},N{r}-1)>0,"Same family here last year",""))'), False, align="center")
    pl[f"N{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r}),"")'                                         # N: year
    pl[f"O{r}"] = f'=IF(C{r}="","",IFERROR({crop_of(f"C{r}", "C")}&"",""))'                   # O: family
    pl[f"P{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                            # P: date as a number
    pl[f"Q{r}"] = f'=IF(AND(L{r}="",OR(J{r}="Growing",J{r}="Picking")),1,0)'                   # Q: in the ground
    pl[f"R{r}"] = f'=IF(AND(Q{r}=1,D{r}<>""),P{r}*1000+ROW(),0)'                                # R: latest in bed
    pl[f"S{r}"] = f'=IF(AND({dated},L{r}=""),1,0)'                                              # S: counts
for c in "NOPQRS":
    pl.column_dimensions[c].hidden = True
check_fill(pl, f"L{L0}:L{L1}", f"L{L0}")
pl.conditional_formatting.add(f"M{L0}:M{L1}", FormulaRule(formula=[f'M{L0}<>""'], fill=fill(WARN)))
for text, colour in (("Picking", OK), ("Growing", INFO)):
    pl.conditional_formatting.add(f"J{L0}:J{L1}", FormulaRule(formula=[f'J{L0}="{text}"'], fill=fill(colour)))
dropdown(pl, f"={CR('B')}", f"C{L0}:C{L1}", strict=False)
dropdown(pl, f"={BD('B')}", f"D{L0}:D{L1}", strict=False)
dropdown(pl, '"Sown inside,Sown outside,Planted out,Bought plants"', f"E{L0}:E{L1}", strict=False)
pl.freeze_panes = "D6"

# ---------------------------------------------------------------- Harvest
hv = wb.create_sheet("Harvest")
sheet_base(hv, "Harvest", "Each time you pick something. It is valued at the shop price on the Crops tab.",
           [3, 13, 16, 14, 10, 24, 8, 11, 20], rows=H1 + 2)
header(hv, 5, 2, ["Date", "Crop", "Bed", "Amount", "Note", "Unit", "Value", "Check"])
for r in range(H0, H1 + 1):
    x = harvest[r - H0] if r - H0 < len(harvest) else (None,) * 5
    for c, v, fmt in zip(range(2, 7), x, (DATE, None, None, "#,##0.##", None)):
        style(hv.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 5) else None)
    style(hv.cell(r, 7, f'=IF(C{r}="","",IFERROR({crop_of(f"C{r}", "D")}&"",""))'), False, align="center")
    style(hv.cell(r, 8, f'=IF(OR(C{r}="",NOT(ISNUMBER(E{r}))),"",E{r}*N(IFERROR({crop_of(f"C{r}", "E")},0)))'),
          False, MONEY)
    style(hv.cell(r, 9, f'=IF(AND(C{r}="",E{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",IF(C{r}="","Needs a crop",'
                        f'IF(COUNTIF({CR("B")},C{r})=0,"Not on Crops",IF(NOT(ISNUMBER(E{r})),"Needs the amount",'
                        f'IF(AND(D{r}<>"",COUNTIF({BD("B")},D{r})=0),"Not on Beds",""))))))'), False, align="center")
    hv[f"J{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r}),"")'                                         # J: year
    hv[f"K{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                         # K: month
    hv[f"L{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                            # L: date as a number
for c in "JKL":
    hv.column_dimensions[c].hidden = True
check_fill(hv, f"I{H0}:I{H1}", f"I{H0}")
dropdown(hv, f"={CR('B')}", f"C{H0}:C{H1}", strict=False)
dropdown(hv, f"={BD('B')}", f"D{H0}:D{H1}", strict=False)
hv.freeze_panes = "D6"

# ---------------------------------------------------------------- Jobs
jb = wb.create_sheet("Jobs")
sheet_base(jb, "Jobs", "Watering and the other jobs, with the bed. All beds, or an empty bed, means the whole "
           "garden.", [3, 13, 16, 14, 28, 18], rows=J1 + 2)
header(jb, 5, 2, ["Date", "Bed", "Job", "Note", "Check"])
for r in range(J0, J1 + 1):
    x = jobs[r - J0] if r - J0 < len(jobs) else (None,) * 4
    for c, v, fmt in zip(range(2, 6), x, (DATE, None, None, None)):
        style(jb.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4) else None)
    style(jb.cell(r, 6, f'=IF(AND(B{r}="",C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                        f'IF(D{r}="","Needs the job",IF(AND(C{r}<>"",C{r}<>"All beds",'
                        f'COUNTIF({BD("B")},C{r})=0),"Not on Beds",""))))'), False, align="center")
    jb[f"G{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                            # G: date as a number
    jb[f"H{r}"] = f'=IF(AND(G{r}>0,D{r}="Water",OR(C{r}="",C{r}="All beds")),G{r},0)'           # H: watered it all
for c in "GH":
    jb.column_dimensions[c].hidden = True
check_fill(jb, f"F{J0}:F{J1}", f"F{J0}")
dropdown(jb, f'={BD("B")}', f"C{J0}:C{J1}", strict=False)
dropdown(jb, '"Water,Feed,Weed,Mulch,Prune,Pest check,Other"', f"D{J0}:D{J1}", strict=False)
jb.freeze_panes = "C6"

# ---------------------------------------------------------------- Costs
co = wb.create_sheet("Costs")
sheet_base(co, "Costs", "What the garden costs: seeds, plants, compost, tools and water.", [3, 13, 30, 18, 12, 18],
           rows=K1 + 2)
header(co, 5, 2, ["Date", "What", "Kind", "Amount", "Check"])
for r in range(K0, K1 + 1):
    x = costs[r - K0] if r - K0 < len(costs) else (None,) * 4
    for c, v, fmt in zip(range(2, 6), x, (DATE, None, None, MONEY)):
        style(co.cell(r, c, v), True, fmt, align="center" if c in (2, 4) else None)
    style(co.cell(r, 6, f'=IF(AND(C{r}="",E{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                        f'IF(NOT(ISNUMBER(E{r})),"Needs the amount","")))'), False, align="center")
    co[f"G{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r}),"")'                                         # G: year
    co[f"H{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                         # H: month
for c in "GH":
    co.column_dimensions[c].hidden = True
check_fill(co, f"F{K0}:F{K1}", f"F{K0}")
dropdown(co, '"Seeds,Plants,Soil and compost,Tools,Water,Other"', f"D{K0}:D{K1}", strict=False)
co.freeze_panes = "C6"

# ---------------------------------------------------------------- Calendar (to print)
ca = wb.create_sheet("Calendar")
sheet_base(ca, "", "", [3, 18] + [6] * 12, rows=NCAL + 12)
ca["B2"] = f'=IF({NAME}="","Planting calendar",{NAME}&": planting calendar")'
ca["B2"].font = font(18, True, TEAL_D)
ca["B3"] = "The months you typed on the Crops tab. S sow, P plant out, H harvest."
ca["B3"].font = font(10, color=MUTED, italic=True)
header(ca, 5, 2, ["Crop"] + MON)
for k in range(NCAL):
    r = 6 + k
    name = f"INDEX({CR('B')},{k + 1})"
    style(ca.cell(r, 2, f'=IF({name}="","",{name})'), False, bold=True)
    for i, c in enumerate("RSTUVW"):                                                        # P to U: months
        ca.cell(r, 16 + i).value = f"=N(INDEX({CR(c)},{k + 1}))"
    for m in range(12):
        s = in_range(f"$P{r}", f"$Q{r}", m + 1)
        p = in_range(f"$R{r}", f"$S{r}", m + 1)
        h = in_range(f"$T{r}", f"$U{r}", m + 1)
        style(ca.cell(r, 3 + m, f'=IF(B{r}="","",TRIM(IF({s}=1,"S ","")&IF({p}=1,"P ","")&IF({h}=1,"H","")))'),
              False, align="center").font = font(9, True)
    ca.row_dimensions[r].height = 16
for c in "PQRSTU":
    ca.column_dimensions[c].hidden = True
grid = f"C6:N{5 + NCAL}"
for letter, colour in (("H", OK), ("P", INFO), ("S", "FFF4C2")):
    ca.conditional_formatting.add(grid, FormulaRule(formula=[f'ISNUMBER(FIND("{letter}",C6))'], fill=fill(colour),
                                                    stopIfTrue=True))
ca.conditional_formatting.add("C5:N5", FormulaRule(formula=["COLUMN()-2=MONTH(TODAY())"], fill=fill(TEAL_D)))
r = 7 + NCAL
for i, (label, colour) in enumerate((("S  sow", "FFF4C2"), ("P  plant out", INFO), ("H  harvest", OK))):
    c = ca.cell(r, 3 + i * 3, label)
    c.fill = fill(colour)
    c.font = font(9, True)
    ca.merge_cells(start_row=r, start_column=3 + i * 3, end_row=r, end_column=5 + i * 3)
ca.page_setup.orientation = "portrait"
ca.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "", "", [3, 18, 12, 12, 12, 3, 16, 16, 16, 3], rows=52)
db["B2"] = f'=IF({NAME}="","Vegetable garden",{NAME})'
db["B2"].font = font(22, True, TEAL_D)
db["B3"] = (f'={YEAR}&": the harvest at shop prices, what the garden cost, and what to do in "'
            f'&CHOOSE(MONTH(TODAY()),{",".join(chr(34) + m + chr(34) for m in MONTH_NAMES)})&"."')
db["B3"].font = font(11, color=MUTED, italic=True)
tile(db, "B", 4, "Harvest value", f"=SUMIFS({HV('H')},{HV('J')},{YEAR})", MONEY)
tile(db, "C", 4, "Spent", f"=SUMIFS({CO('E')},{CO('G')},{YEAR})", MONEY)
tile(db, "D", 4, "Value less costs", "=B5-C5", MONEY)
tile(db, "E", 4, "Crops picked", f'=COUNTIF({CR("O")},">0")', "0")
tile(db, "G", 4, "Plantings", f"=SUMIFS({PL('S')},{PL('N')},{YEAR})", "0")
tile(db, "H", 4, "In the ground", f"=SUM({PL('Q')})", "0")
tile(db, "I", 4, "Beds to water", f'=COUNTIF({BD("J")},"Water")', "0", "B86E1C")
header(db, 7, 2, ["Harvest this year", "Picked", "Unit", "Value"])
for n in range(10):
    r = 8 + n
    db[f"L{r}"] = f"=IFERROR(LARGE({CR('AA')},{n + 1}),\"\")"                                    # L: crop key
    c = lambda col: f"INDEX({CR1(col)},1000-MOD(L{r},1000))"
    style(db.cell(r, 2, f'=IF(L{r}="","",{c("B")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(L{r}="","",{c("O")})'), False, "#,##0.##", align="center")
    style(db.cell(r, 4, f'=IF(L{r}="","",IF({c("D")}="","",{c("D")}))'), False, align="center")
    style(db.cell(r, 5, f'=IF(L{r}="","",{c("P")})'), False, MONEY)
header(db, 7, 7, ["Sow now", "Plant out now", "Harvest now"])
for n in range(10):
    r = 8 + n
    for i, key in enumerate("XYZ"):
        col = "MNO"[i]
        db[f"{col}{r}"] = f"=IFERROR(SMALL({CR(key)},{n + 1}),\"\")"                             # M to O: crops
        style(db.cell(r, 7 + i, f'=IF({col}{r}="","",INDEX({CR1("B")},{col}{r}))'), False, align="center")
db["G18"] = '=IF(COUNT(Crops!$Z$6:$Z$65)>10,"And "&(COUNT(Crops!$Z$6:$Z$65)-10)&" more to harvest","")'
db["G18"].font = font(9, color=MUTED, italic=True)
header(db, 20, 2, ["Bed", "Growing", "Last watered", "Status"])
for n in range(8):
    r = 21 + n
    b = lambda col: f"INDEX({BD(col)},{n + 1})"
    style(db.cell(r, 2, f'=IF({b("B")}="","",{b("B")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",{b("F")})'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",IF({b("H")}="","",{b("H")}))'), False, "d mmm", align="center")
    style(db.cell(r, 5, f'=IF(B{r}="","",{b("J")})'), False, bold=True, align="center")
db.conditional_formatting.add("E21:E28", FormulaRule(formula=['E21="Water"'], fill=fill(WARN)))
header(db, 20, 7, ["Month", "Harvest value", "Spent"])
for k in range(12):
    r = 21 + k
    key = f"({YEAR}*100+{k + 1})"
    style(db.cell(r, 7, MON[k]), False, bold=True, align="center")
    style(db.cell(r, 8, f"=SUMIFS({HV('H')},{HV('K')},{key})"), False, MONEY)
    style(db.cell(r, 9, f"=SUMIFS({CO('E')},{CO('H')},{key})"), False, MONEY)
for c in "LMNO":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 18)) + list(range(21, 33)):
    db.row_dimensions[r].height = 18
chart = BarChart()
chart.type = "col"
chart.title = "Harvest and costs by month"
chart.height, chart.width = 7, 12
chart.add_data(Reference(db, min_col=8, max_col=9, min_row=20, max_row=32), titles_from_data=True)
chart.set_categories(Reference(db, min_col=7, min_row=21, max_row=32))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E39A4B"
chart.y_axis.number_format = "#,##0"
chart.legend.position = "b"
db.add_chart(chart, "B34")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: the garden's name, and after how many days a bed needs water again."),
    ("2", "Beds and Crops: each bed once, and each crop with its shop price and the months you plan to sow, plant "
          "out and harvest it."),
    ("3", "Plantings: each time you sow or plant something, with the bed. A date under Cleared once it is finished."),
    ("4", "Harvest and Jobs: each time you pick something, and each watering or other job."),
    ("5", "Costs: seeds, plants, compost and tools, to see what the garden gave back."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example garden and its months are made up. Your own months depend on where you live "
                          "and on the weather, so use your seed packets and local advice.",
                          "Calendar: print it to keep in the shed.",
                          "Works in Google Sheets and Microsoft Excel, in any currency and any unit."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Plantings", "Harvest", "Jobs", "Crops", "Beds", "Calendar", "Costs", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Vegetable-Garden-Planner.xlsx")
wb.save(out)
print("saved", out, len(plantings), "plantings", len(harvest), "picks", len(jobs), "jobs")
