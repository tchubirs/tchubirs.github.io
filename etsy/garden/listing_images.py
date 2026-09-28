"""Etsy listing photos for the Vegetable Garden Planner (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Vegetable-Garden-Planner.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in; the logs show their latest lines.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
crops_end, beds_end = last(wb["Crops"], 65), last(wb["Beds"], 35)
pl_end, hv_end, jb_end = last(wb["Plantings"], 505), last(wb["Harvest"], 1505), last(wb["Jobs"], 2005)
garden = wb["Settings"]["C4"].value
rows = {"Crops": range(crops_end + 1, 66), "Beds": range(beds_end + 1, 36),
        "Plantings": range(pl_end + 1, 506),
        "Harvest": list(range(6, hv_end - 21)) + list(range(hv_end + 1, 1506)),
        "Jobs": list(range(6, jb_end - 9)) + list(range(jb_end + 1, 2006)),
        "Calendar": range(6 + crops_end - 5, 46)}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Calendar"), hide_rows=rows, hide={"Plantings": ["L"], "Crops": ["Q"]})


def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


def stack(top, lower):
    lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                         Image.LANCZOS)
    both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
    both.paste(top, (0, 0))
    both.paste(lower, (0, top.height + 40))
    return both


img = canvas(dark=True)
title(img, "Vegetable garden planner", "What you sow and plant in each bed, the harvest at shop prices, the watering, "
      "and a calendar to print", dark=True, width=2300)
laptop(img, r.shot(garden), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-calendar.jpg"), "A planting calendar to print",
       "Made from the months you type for each crop: S sow, P plant out, H harvest. This month stands out.",
       r.shot(f"{garden}: planting"))

detail(os.path.join(img_dir, "03-plantings.jpg"), "Every sowing and planting, bed by bed",
       "When it should be ready by your own days to harvest, and a note when a bed gets the same crop family as last "
       "year", r.shot("Plantings"),
       [("#D8F0DC", "Picking"), ("#E3E8FF", "Growing"), ("#FFE2B8", "Same family as last year")])

detail(os.path.join(img_dir, "04-harvest.jpg"), "The harvest, valued at shop prices",
       "Each pick with its amount. The value comes from the price you would pay in the shop for each crop.",
       r.shot("Harvest"))

detail(os.path.join(img_dir, "05-beds-watering.jpg"), "Beds and watering",
       "What grows in each bed, the last watering, and the beds that need water now",
       stack(r.shot("Beds"), r.shot("Jobs")), [("#FFE2B8", "Water"), ("#D8F0DC", "Watered lately")])

detail(os.path.join(img_dir, "06-crops.jpg"), "Each crop with your own months",
       "Family, unit, shop price and days to harvest, the months you plan, the seeds left, and what it gave this year",
       r.shot("Crops"))

img = canvas()
y = title(img, "Inside the file", "Eight tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page(garden)), "Dashboard", "The year and this month"),
    (content(r.page("Plantings")), "Plantings", "What went in, and where"),
    (content(r.page("Harvest")), "Harvest", "Each pick and its value"),
    (content(r.page("Beds")), "Beds", "What grows, when watered"),
    (content(r.page(f"{garden}: planting")), "Calendar", "To print"),
    (content(r.page("Crops")), "Crops", "Your months and prices"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "07-whats-inside.jpg"))
