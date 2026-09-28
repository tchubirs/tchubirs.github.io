"""Moving House Planner (Excel + Google Sheets).

A countdown to moving day; a checklist whose dates count back from it; every box numbered, with the room it
leaves and the room it goes to, what is inside, fragile or open first, packed and arrived; movers' quotes side by
side; the costs with deposits and what is still to pay; everyone who needs the new address; and the meter
readings. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
T0, T1 = 6, 75         # tasks
X0, X1 = 6, 305        # boxes
C0, C1 = 6, 55         # costs
Q0, Q1 = 6, 15         # movers' quotes
A0, A1 = 6, 65         # who to tell
R0, R1 = 6, 15         # rooms on Settings
DAY = "Dashboard!$C$4"
TA = lambda col: f"Tasks!${col}${T0}:${col}${T1}"
BX = lambda col: f"Boxes!${col}${X0}:${col}${X1}"
CO = lambda col: f"Costs!${col}${C0}:${col}${C1}"
QU = lambda col: f"Movers!${col}${Q0}:${col}${Q1}"
AD = lambda col: f"'New address'!${col}${A0}:${col}${A1}"
ROOMS = f"Settings!$B${R0}:$B${R1}"
CATS = ["Movers", "Van and fuel", "Packing", "Cleaning", "Storage", "Deposits", "Utilities", "Furniture", "Travel", "Other"]
KINDS = ["Money", "Work", "Health", "Home", "Government", "Post", "Subscriptions", "People", "Other"]
SIZES = ["Small", "Medium", "Large"]

# ---------------------------------------------------------------- example data (made-up companies and places)
move = today + dt.timedelta(days=20 + (5 - (today + dt.timedelta(days=20)).weekday()) % 7)   # a Saturday, 20 to 26 days away
rooms = ["Kitchen", "Living room", "Bedroom", "Kids room", "Bathroom", "Office", "Hallway", "Garage"]
tasks = [
    # task, days before moving day, who, done
    ("Set a budget and the moving date", 56, "Both", True), ("Get three quotes from movers", 56, "Sam", True),
    ("Book the movers", 49, "Sam", True), ("Start clearing out, room by room", 49, "Both", True),
    ("Order boxes, tape and bubble wrap", 42, "Alex", True), ("Book the day off work", 42, "Both", True),
    ("Sell or give away what does not come", 35, "Alex", False), ("Measure the new home for the furniture", 35, "Sam", True),
    ("Give notice on the old flat", 30, "Sam", True), ("Pack the rooms you rarely use", 28, "Both", True),
    ("Set up mail forwarding", 21, "Alex", True), ("Book internet at the new home", 21, "Sam", False),
    ("Tell the bank, work and the doctor", 14, "Both", False), ("Book cleaning for the old flat", 14, "Alex", False),
    ("Confirm the movers and the time", 10, "Sam", False), ("Arrange parking for the van", 10, "Sam", False),
    ("Use up the frozen food", 7, "Both", False), ("Pack all but the essentials", 7, "Both", False),
    ("Pack the open-first boxes", 3, "Alex", False), ("Defrost the fridge", 2, "Alex", False),
    ("Take photos of the old flat", 1, "Sam", False), ("Take the meter readings", 0, "Sam", False),
    ("Check every room and cupboard", 0, "Both", False), ("Hand over the keys", 0, "Sam", False),
    ("Unpack the open-first boxes", -1, "Both", False), ("Register with a new doctor", -7, "Alex", False),
    ("Get the deposit back", -14, "Sam", False), ("Unpack the last boxes", -21, "Both", False),
]
boxes = []
contents = {
    "Kitchen": [("Plates and bowls", "Medium", True), ("Glasses", "Small", True), ("Pots and pans", "Large", False),
                ("Baking things", "Medium", False), ("Mugs and the kettle", "Small", True), ("Cutlery and knives", "Small", False),
                ("Food cupboard", "Medium", False), ("Small appliances", "Large", False), ("Tea towels and aprons", "Small", False)],
    "Living room": [("Books A to M", "Small", False), ("Books N to Z", "Small", False), ("Lamps", "Medium", True),
                    ("Cushions and throws", "Large", False), ("Board games", "Medium", False), ("Photo frames", "Small", True),
                    ("Speakers and cables", "Medium", True)],
    "Bedroom": [("Bedding", "Large", False), ("Winter clothes", "Large", False), ("Shoes", "Medium", False),
                ("Summer clothes", "Large", False), ("Jewellery box and watch", "Small", True), ("Towels", "Medium", False)],
    "Kids room": [("Toys", "Large", False), ("Kids books", "Small", False), ("Kids clothes", "Large", False),
                  ("Night light and bedding", "Medium", False)],
    "Bathroom": [("Bathroom things", "Small", False), ("Medicine box", "Small", False)],
    "Office": [("Files and papers", "Small", False), ("Monitor", "Large", True), ("Printer", "Medium", True),
               ("Stationery", "Small", False), ("Cables and chargers", "Small", False)],
    "Hallway": [("Coats", "Large", False), ("Umbrellas and bags", "Medium", False)],
    "Garage": [("Tools", "Medium", False), ("Garden things", "Large", False), ("Bike gear", "Medium", False),
               ("Christmas decorations", "Large", False)],
}
open_first = {"Mugs and the kettle", "Bedding", "Bathroom things", "Night light and bedding", "Cables and chargers"}
packed_rooms = {"Garage", "Office", "Living room"}
n = 1
for room in rooms:
    for what, size, fragile in contents[room]:
        packed = room in packed_rooms or what in ("Books A to M", "Baking things", "Winter clothes", "Summer clothes")
        if what in ("Speakers and cables", "Cables and chargers"):
            packed = False
        to = {"Office": "Study", "Hallway": "Hall"}.get(room, room)
        boxes.append((n, room, to, what, size, fragile, what in open_first, packed, False))
        n += 1
costs = [
    # item, category, company, estimate, final price, paid, pay by
    ("Removal, 2 movers and a truck", "Movers", "Blue Box Removals", 950, 950, 200, move),
    ("Boxes, tape and bubble wrap", "Packing", "Box shop", 90, 86.40, 86.40, None),
    ("End of tenancy cleaning", "Cleaning", "Sparkle Clean", 180, 180, 0, move + dt.timedelta(days=1)),
    ("Mail forwarding, 6 months", "Other", "Post office", 35, 34.50, 34.50, None),
    ("Internet set-up", "Utilities", "Fibre company", 49, None, 0, move - dt.timedelta(days=21)),
    ("Van parking permit", "Van and fuel", "Council", 25, None, 0, move - dt.timedelta(days=10)),
    ("Takeaway on moving day", "Other", None, 40, None, 0, None),
    ("Curtains for the new bedroom", "Furniture", None, 120, None, 0, None),
]
quotes = [
    # company, price, date free, packing included, insurance, notes, chosen
    ("Oak Lane Removals", 1150, move, "x", "x", "Pack the kitchen for us", None),
    ("City Van Movers", 890, move + dt.timedelta(days=7), None, None, "Only free the week after", None),
    ("Blue Box Removals", 950, move, None, "x", "Good reviews, can store for a week", "x"),
]
tell = [
    # who, kind, how, reference, done
    ("Bank", "Money", "Online", "Current account", True), ("Credit card", "Money", "App", None, True),
    ("Employer", "Work", "Email to HR", None, True), ("Doctor", "Health", "Form at the surgery", None, False),
    ("Dentist", "Health", "Phone", None, False), ("Car insurance", "Money", "Online", "Policy number", False),
    ("Home insurance", "Money", "Phone", "Policy number", False), ("Electricity", "Home", "Online", "Meter reading", False),
    ("Gas", "Home", "Online", "Meter reading", False), ("Water", "Home", "Phone", None, False),
    ("Internet", "Home", "Online", "Booked for the new home", True), ("Local council", "Government", "Online form", None, False),
    ("Tax office", "Government", "Online", None, False), ("Driving licence", "Government", "Online", None, False),
    ("Mail forwarding", "Post", "Post office", "6 months", True), ("Mobile phone", "Subscriptions", "App", None, True),
    ("Streaming and magazines", "Subscriptions", "Online", None, False), ("Gym", "Subscriptions", "At the desk", None, False),
    ("School", "Other", "Letter", None, True), ("Family and friends", "People", "Card with the new address", None, False),
    ("Vet", "Health", "Phone", None, False), ("Online shops", "Subscriptions", "Each account", None, False),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The rooms of the home you are leaving, for the boxes and the Dashboard.", [3, 22, 4, 60], rows=25)
header(se, 5, 2, ["Rooms"])
for k in range(R1 - R0 + 1):
    style(se.cell(R0 + k, 2, rooms[k] if k < len(rooms) else None), True, bold=True)
se["D6"] = "The Dashboard counts the boxes of each room on this list."
se["D6"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Tasks
ta = wb.create_sheet("Tasks")
sheet_base(ta, "Tasks", "The checklist, counted back from moving day. Change the days before and the dates follow. "
           "0 is moving day, and a minus number is after it.", [3, 40, 9, 16, 14, 8, 12, 30], rows=T1 + 2)
header(ta, 5, 2, ["Task", "Days before", "Date", "Who", "Done", "Status", "Notes"])
for r in range(T0, T1 + 1):
    t = tasks[r - T0] if r - T0 < len(tasks) else (None,) * 4
    name, before, who, done = t
    style(ta.cell(r, 2, name), True, bold=True)
    style(ta.cell(r, 3, before), True, "0", align="center")
    style(ta.cell(r, 4, f'=IF(OR(B{r}="",C{r}="",{DAY}=""),"",{DAY}-C{r})'), False, "ddd d mmm", align="center")
    style(ta.cell(r, 5, who), True)
    style(ta.cell(r, 6, "x" if done else None), True, align="center")
    style(ta.cell(r, 7, f'=IF(B{r}="","",IF(F{r}<>"","Done",IF(D{r}="","",IF(D{r}<TODAY(),"Late",'
                        f'IF(D{r}-TODAY()<=7,"This week","Coming")))))'), False, align="center")
    style(ta.cell(r, 8), True)
    ta.cell(r, 10, f'=IF(AND(B{r}<>"",F{r}="",D{r}<>""),D{r}+ROW()/10000000,"")')      # J: order of what is left
ta.column_dimensions["J"].hidden = True
for status, colour in (("Done", OK), ("Late", BAD), ("This week", WARN)):
    ta.conditional_formatting.add(f"G{T0}:G{T1}", FormulaRule(formula=[f'G{T0}="{status}"'], fill=fill(colour)))
ta.conditional_formatting.add(f"B{T0}:E{T1}", FormulaRule(formula=[f'$F{T0}<>""'], font=Font(name=F, color="9AA7A9")))
ta.freeze_panes = "C6"

# ---------------------------------------------------------------- Boxes
bx = wb.create_sheet("Boxes")
sheet_base(bx, "Boxes", "Write the number on the box and one line here: the room it leaves, the room it goes to and "
           "what is inside. Put an x as it is packed and when it arrives.",
           [3, 7, 15, 15, 34, 9, 8, 8, 8, 8, 12, 24], rows=X1 + 2)
header(bx, 5, 2, ["Box", "From", "To", "What is inside", "Size", "Fragile", "Open first", "Packed", "Arrived", "Status",
                  "Notes"])
for r in range(X0, X1 + 1):
    b = boxes[r - X0] if r - X0 < len(boxes) else (None,) * 9
    num, frm, to, what, size, fragile, first, packed, arrived = b
    style(bx.cell(r, 2, num), True, "0", True, "center")
    style(bx.cell(r, 3, frm), True)
    style(bx.cell(r, 4, to), True)
    style(bx.cell(r, 5, what), True)
    style(bx.cell(r, 6, size), True, align="center")
    for c, v in ((7, fragile), (8, first), (9, packed), (10, arrived)):
        style(bx.cell(r, c, "x" if v else None), True, align="center")
    style(bx.cell(r, 11, f'=IF(AND(B{r}="",E{r}=""),"",IF(J{r}<>"","Arrived",IF(I{r}<>"","Packed","To pack")))'), False,
          align="center")
    style(bx.cell(r, 12), True)
    bx.cell(r, 14, f'=IF(AND(H{r}<>"",B{r}<>""),B{r}+ROW()/10000000,"")')                  # N: open-first order
dropdown(bx, f"={ROOMS}", f"C{X0}:C{X1}", strict=False)
dropdown(bx, f'"{",".join(SIZES)}"', f"F{X0}:F{X1}")
bx.column_dimensions["N"].hidden = True
bx.conditional_formatting.add(f"K{X0}:K{X1}", FormulaRule(formula=[f'K{X0}="Arrived"'], fill=fill(OK)))
bx.conditional_formatting.add(f"K{X0}:K{X1}", FormulaRule(formula=[f'K{X0}="Packed"'], fill=fill("CFE8E4")))
bx.conditional_formatting.add(f"G{X0}:G{X1}", FormulaRule(formula=[f'G{X0}<>""'], fill=fill(BAD)))
bx.conditional_formatting.add(f"H{X0}:H{X1}", FormulaRule(formula=[f'H{X0}<>""'], fill=fill(WARN)))
bx.conditional_formatting.add(f"B{X0}:B{X1}", FormulaRule(formula=[f'AND(B{X0}<>"",COUNTIF({BX("B")},B{X0})>1)'],
                                                          fill=fill(BAD)))
bx.freeze_panes = "C6"

# ---------------------------------------------------------------- Movers
mv = wb.create_sheet("Movers")
sheet_base(mv, "Movers", "Quotes side by side. Put an x under Chosen for the one you book, and add it to Costs.",
           [3, 24, 12, 14, 10, 10, 34, 9], rows=25)
header(mv, 5, 2, ["Company", "Price", "Free on", "Packing", "Insurance", "Notes", "Chosen"])
for r in range(Q0, Q1 + 1):
    q = quotes[r - Q0] if r - Q0 < len(quotes) else (None,) * 7
    company, price, free, packing, insurance, notes, chosen = q
    style(mv.cell(r, 2, company), True, bold=True)
    style(mv.cell(r, 3, price), True, MONEY)
    style(mv.cell(r, 4, free), True, DATE)
    style(mv.cell(r, 5, packing), True, align="center")
    style(mv.cell(r, 6, insurance), True, align="center")
    style(mv.cell(r, 7, notes), True)
    style(mv.cell(r, 8, chosen), True, align="center")
mv.conditional_formatting.add(f"B{Q0}:H{Q1}", FormulaRule(formula=[f'$H{Q0}<>""'], fill=fill(OK)))
mv.conditional_formatting.add(f"C{Q0}:C{Q1}", FormulaRule(formula=[f'AND(ISNUMBER(C{Q0}),C{Q0}=MIN({QU("C")}))'],
                                                          font=Font(name=F, bold=True, color="2E7D32")))
mv.conditional_formatting.add(f"D{Q0}:D{Q1}", FormulaRule(formula=[f'AND(ISNUMBER(D{Q0}),{DAY}<>"",D{Q0}<>{DAY})'],
                                                          fill=fill(WARN)))
mv[f"B{Q1 + 2}"] = "The cheapest price is in green. A date in orange is not moving day."
mv[f"B{Q1 + 2}"].font = font(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Costs
co = wb.create_sheet("Costs")
sheet_base(co, "Costs", "Everything the move costs: what you booked, what you paid and when the rest is due.",
           [3, 30, 14, 18, 11, 11, 11, 11, 13, 12, 22], rows=C1 + 2)
header(co, 5, 2, ["Item", "Category", "Company or shop", "Estimate", "Final price", "Paid", "Still to pay", "Pay by",
                  "Status", "Notes"])
for r in range(C0, C1 + 1):
    c_ = costs[r - C0] if r - C0 < len(costs) else (None,) * 7
    name, cat, shop, est, final, paid, due = c_
    style(co.cell(r, 2, name), True, bold=True)
    style(co.cell(r, 3, cat), True)
    style(co.cell(r, 4, shop), True)
    style(co.cell(r, 5, est), True, MONEY)
    style(co.cell(r, 6, final), True, MONEY)
    style(co.cell(r, 7, paid), True, MONEY)
    style(co.cell(r, 8, f'=IF(B{r}="","",MAX(0,IF(F{r}<>"",F{r},N(E{r}))-N(G{r})))'), False, MONEY)
    style(co.cell(r, 9, due), True, DATE)
    style(co.cell(r, 10, f'=IF(B{r}="","",IF(H{r}<=0,"Paid",IF(I{r}="","To pay",IF(I{r}<TODAY(),"Late",'
                         f'IF(I{r}-TODAY()<=7,"Due soon","To pay")))))'), False, align="center")
    style(co.cell(r, 11), True)
    co.cell(r, 13, f'=IF(B{r}="",0,IF(F{r}<>"",F{r},N(E{r})))')        # M: the cost that counts
co.column_dimensions["M"].hidden = True
dropdown(co, f'"{",".join(CATS)}"', f"C{C0}:C{C1}", strict=False)
for status, colour in (("Paid", OK), ("Late", BAD), ("Due soon", WARN)):
    co.conditional_formatting.add(f"J{C0}:J{C1}", FormulaRule(formula=[f'J{C0}="{status}"'], fill=fill(colour)))
co[f"B{C1 + 1}"] = "The final price counts once you know it. Until then the estimate does."
co[f"B{C1 + 1}"].font = font(9, color=MUTED, italic=True)
co.freeze_panes = "C6"

# ---------------------------------------------------------------- New address
ad = wb.create_sheet("New address")
sheet_base(ad, "New address", "Everyone who needs your new address. Put an x under Done when you have told them.",
           [3, 26, 14, 22, 22, 7, 24, 3, 16, 14, 14, 12], rows=A1 + 2)
header(ad, 5, 2, ["Who", "Kind", "How", "Account or reference", "Done", "Notes"])
for r in range(A0, A1 + 1):
    w = tell[r - A0] if r - A0 < len(tell) else (None,) * 5
    who, kind, how, ref, done = w
    style(ad.cell(r, 2, who), True, bold=True)
    style(ad.cell(r, 3, kind), True)
    style(ad.cell(r, 4, how), True)
    style(ad.cell(r, 5, ref), True)
    style(ad.cell(r, 6, "x" if done else None), True, align="center")
    style(ad.cell(r, 7), True)
dropdown(ad, f'"{",".join(KINDS)}"', f"C{A0}:C{A1}", strict=False)
ad.conditional_formatting.add(f"B{A0}:G{A1}", FormulaRule(formula=[f'$F{A0}<>""'], font=Font(name=F, color="9AA7A9")))
ad.conditional_formatting.add(f"F{A0}:F{A1}", FormulaRule(formula=[f'F{A0}<>""'], fill=fill(OK)))
header(ad, 5, 9, ["Meter", "Old home", "New home", "Date"])
for k, meter in enumerate(["Electricity", "Gas", "Water", "Other"]):
    r = 6 + k
    style(ad.cell(r, 9, meter), True, bold=True)
    style(ad.cell(r, 10), True, "0")
    style(ad.cell(r, 11), True, "0")
    style(ad.cell(r, 12, move if meter != "Other" else None), True, DATE)
ad["I11"] = "Take the readings on moving day, with a photo."
ad["I11"].font = font(9, color=MUTED, italic=True)
ad.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Moving day", "The countdown, the checklist, the boxes, the money and who still needs your new address.",
           [3, 18, 14, 14, 14, 14, 14, 3, 18, 10, 10, 10, 10, 3], rows=45)
db["B2"] = '=IF(C5="","Moving day","Moving to "&C5)'
for r, label in ((4, "Moving day"), (5, "New home"), (6, "Budget")):
    db[f"B{r}"] = label
    db[f"B{r}"].font = font(11, True)
style(db["C4"], True, DATE, True, "center")
db["C4"] = move
style(db["C5"], True, bold=True)
db["C5"] = "Maple Road"
db.merge_cells("C5:E5")
style(db["C6"], True, MONEY, True, "center")
db["C6"] = 1800
planned = f'SUM({CO("M")})'
boxes_n = f'COUNTIF({BX("K")},"?*")'
tile(db, "B", 8, "Countdown", f'=IF($C$4="","",IF($C$4-TODAY()>1,($C$4-TODAY())&" days",IF($C$4-TODAY()=1,"Tomorrow",'
                              f'IF($C$4=TODAY(),"Today","Moved"))))', "@")
tile(db, "C", 8, "Tasks done", f'=COUNTIFS({TA("B")},"<>",{TA("F")},"<>")&" of "&COUNTIF({TA("B")},"?*")', "@")
tile(db, "D", 8, "Late tasks", f'=COUNTIF({TA("G")},"Late")', "0", "B4541F")
tile(db, "E", 8, "Boxes packed", f'=(COUNTIF({BX("K")},"Packed")+COUNTIF({BX("K")},"Arrived"))&" of "&{boxes_n}', "@")
tile(db, "F", 8, "Told new address", f'=COUNTIFS({AD("B")},"<>",{AD("F")},"<>")&" of "&COUNTIF({AD("B")},"?*")', "@")
tile(db, "G", 8, "Budget left", f'=IF(N($C$6)=0,"",$C$6-{planned})', MONEY)
db.conditional_formatting.add("G9", FormulaRule(formula=["AND(ISNUMBER(G9),G9<0)"], font=Font(name=F, size=16, bold=True,
                                                                                               color="B4541F")))

header(db, 11, 2, ["Next up", "", "", "Date", "Who", "Status"])
db.merge_cells("B11:D11")
for i in range(8):
    r = 12 + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({TA("J")},{i + 1}),"")'
    row = f'MATCH({key},{TA("J")},0)'
    none = 'IF($C$4="","Type moving day first","Nothing left to do")' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({TA("B")},{row}))'), False, bold=True)
    for c in (3, 4):
        style(db.cell(r, c), False)
    db.merge_cells(f"B{r}:D{r}")
    style(db.cell(r, 5, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 6, f'=IF({key}="","",INDEX({TA("E")},{row}))'), False)
    style(db.cell(r, 7, f'=IF({key}="","",INDEX({TA("G")},{row}))'), False, align="center")
for status, colour in (("Late", BAD), ("This week", WARN)):
    db.conditional_formatting.add("G12:G19", FormulaRule(formula=[f'G12="{status}"'], fill=fill(colour)))

header(db, 11, 9, ["Boxes by room", "Boxes", "Packed", "Fragile", "Arrived"])
for k in range(R1 - R0 + 1):
    r = 12 + k
    room = f"Settings!B{R0 + k}"
    style(db.cell(r, 9, f'=IF({room}="","",{room})'), False, bold=True)
    style(db.cell(r, 10, f'=IF(I{r}="","",COUNTIF({BX("C")},I{r}))'), False, "0;;", align="center")
    style(db.cell(r, 11, f'=IF(I{r}="","",COUNTIFS({BX("C")},I{r},{BX("K")},"Packed")+COUNTIFS({BX("C")},I{r},{BX("K")},"Arrived"))'),
          False, "0;;", align="center")
    style(db.cell(r, 12, f'=IF(I{r}="","",COUNTIFS({BX("C")},I{r},{BX("G")},"<>"))'), False, "0;;", align="center")
    style(db.cell(r, 13, f'=IF(I{r}="","",COUNTIFS({BX("C")},I{r},{BX("K")},"Arrived"))'), False, "0;;", align="center")
db.conditional_formatting.add("K12:K21", FormulaRule(formula=["AND(N(J12)>0,K12=J12)"], fill=fill(OK)))

header(db, 22, 2, ["Money", "Planned", "Paid", "Still to pay"])
for k, cat in enumerate(CATS):
    r = 23 + k
    style(db.cell(r, 2, cat), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({CO("M")},{CO("C")},B{r})'), False, "#,##0.00;;")
    style(db.cell(r, 4, f'=SUMIFS({CO("G")},{CO("C")},B{r})'), False, "#,##0.00;;")
    style(db.cell(r, 5, f'=SUMIFS({CO("H")},{CO("C")},B{r})'), False, "#,##0.00;;")
r = 23 + len(CATS)
style(db.cell(r, 2, "Total"), False, bold=True)
for c, col in ((3, "M"), (4, "G"), (5, "H")):        # every cost, also one without a category from the list
    style(db.cell(r, c, f"=SUM({CO(col)})"), False, MONEY, True)
db[f"B{r + 1}"] = f'=IF(N($C$6)=0,"",{bar(f"MIN(1,C{r}/$C$6)", 24)})'
db[f"B{r + 1}"].font = Font(name=F, size=12, color=TEAL)
db.merge_cells(f"B{r + 1}:D{r + 1}")
db[f"E{r + 1}"] = f'=IF(N($C$6)=0,"",IF(C{r}>$C$6,"Over by "&FIXED(C{r}-$C$6,2),FIXED(C{r}/$C$6*100,0)&"% of the budget"))'
db[f"E{r + 1}"].font = font(10, True, TEAL_D)

header(db, 22, 9, ["Open first", "Box", "To"])
db.merge_cells("K22:M22")
for i in range(8):
    r = 23 + i
    key = f"Q{r}"
    db[key] = f'=IFERROR(SMALL({BX("N")},{i + 1}),"")'
    row = f'MATCH({key},{BX("N")},0)'
    none = '"No box marked open first"' if i == 0 else '""'
    style(db.cell(r, 9, f'=IF({key}="",{none},INDEX({BX("E")},{row}))'), False, bold=True)
    style(db.cell(r, 10, f'=IF({key}="","",INT({key}))'), False, "0", align="center")
    style(db.cell(r, 11, f'=IF({key}="","",INDEX({BX("D")},{row}))'), False)
    for c in (12, 13):
        style(db.cell(r, c), False)
    db.merge_cells(f"K{r}:M{r}")
db["I31"] = f'="Next box number: "&(MAX({BX("B")})+1)'
db["I31"].font = font(10, True, TEAL_D)
for c in ("P", "Q"):
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Dashboard: moving day, the new home and your budget. Settings: the rooms of the home you leave."),
    ("2", "Tasks: a checklist is ready, counted back from moving day. Put an x under Done as you go."),
    ("3", "Boxes: write a number on each box and one line here. Put an x when it is packed and when it arrives."),
    ("4", "Movers and Costs: compare the quotes, then note what you book, what you paid and when the rest is due."),
    ("5", "New address: everyone who needs it, with the meter readings to take on moving day."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example move, its companies and its places are made up. Delete them and add your own.",
                          "Change moving day and every task date moves with it.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Tasks", "Boxes", "Movers", "Costs", "New address", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Moving-Planner.xlsx")
wb.save(out)
print("saved", out, "moving on", move, len(boxes), "boxes", len(tasks), "tasks")
