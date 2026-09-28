"""Landlord Rental Property Tracker (Excel + Google Sheets).

Units and leases, a log of rent received, a rent roll with one column per
month (paid, part paid or late), rent owed to date, expenses per property and
net income. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
NET = '#,##0.00;[Red]-#,##0.00'
U0, U1 = 6, 25       # units (the rent roll uses the same rows)
G0, G1 = 6, 1005     # rent log rows
X0, X1 = 6, 1005     # expense rows
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PROPS = ["12 Oak Street", "8 River Road"]
CATEGORIES = ["Mortgage interest", "Insurance", "Property tax", "Repairs", "Maintenance", "Utilities", "HOA fees",
              "Management", "Legal and accounting", "Other"]
YR = "Dashboard!$C$4"
LG = lambda col: f"'Rent log'!${col}${G0}:${col}${G1}"
XP = lambda col: f"Expenses!${col}${X0}:${col}${X1}"
RR = lambda col: f"'Rent roll'!${col}${U0}:${col}${U1}"
UN = lambda col: f"Units!${col}${U0}:${col}${U1}"
D = lambda y, m, d: dt.date(y, m, d)

# ---------------------------------------------------------------- example data
units = [("Oak Street, apt 1", PROPS[0], "Maria L.", 1150, 1, 5, D(2025, 6, 1), D(2027, 5, 31), 1150),
         ("Oak Street, apt 2", PROPS[0], "James T.", 1200, 1, 5, D(2025, 11, 1), today + dt.timedelta(days=33), 1200),
         ("Oak Street, apt 3", PROPS[0], "Kevin P.", 950, 1, 5, D(2025, 2, 1), D(Y, 7, 31), 950),
         ("River Road house", PROPS[1], "Priya S.", 1800, 1, 5, D(2024, 9, 1), None, 1800)]
payments = []
for m in range(1, today.month + 1):
    payments.append((D(Y, m, 2), units[0][0], 1150, None, "Bank transfer", ""))
    payments.append((D(Y, m, 1), units[1][0], 600 if m == today.month else 1200, None, "Bank transfer",
                     "Paid half, rest promised" if m == today.month else ""))
    if m <= 7:
        payments.append((D(Y, m, 3), units[2][0], 950, None, "Check", ""))
    payments.append((D(Y, m, 12 if m == 6 else 4), units[3][0], 1800, D(Y, m, 1), "App", "Paid late" if m == 6 else ""))
payments.sort(key=lambda p: p[0])
expenses = []
for m in range(1, today.month + 1):
    expenses.append((D(Y, m, 5), PROPS[0], "Mortgage interest", 1100, ""))
    expenses.append((D(Y, m, 5), PROPS[1], "Mortgage interest", 650, ""))
expenses += [(D(Y, 1, 14), PROPS[0], "Repairs", 480, "Boiler repair"),
             (D(Y, 2, 10), PROPS[0], "Insurance", 1200, "Building insurance, yearly"),
             (D(Y, 3, 10), PROPS[1], "Insurance", 900, "Home insurance, yearly"),
             (D(Y, 4, 15), PROPS[0], "Property tax", 2400, "First half"),
             (D(Y, 4, 22), PROPS[1], "Maintenance", 150, "Gutter cleaning"),
             (D(Y, 5, 15), PROPS[1], "Property tax", 1600, "Yearly"),
             (D(Y, 8, 18), PROPS[0], "Repairs", 650, "Paint apt 3 between tenants"),
             (D(Y, 8, 29), PROPS[0], "Legal and accounting", 120, "New lease template")]
expenses.sort(key=lambda e: e[0])

wb = Workbook()

# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your properties and expense categories. The lists in the other tabs follow them.",
           [3, 30, 4, 60], rows=40)
header(se, 5, 2, ["Properties (up to 6)"])
for k in range(6):
    style(se.cell(6 + k, 2, PROPS[k] if k < len(PROPS) else None), True, bold=True)
header(se, 14, 2, ["Expense categories"])
for k, cat in enumerate(CATEGORIES):
    style(se.cell(15 + k, 2, cat), True)
se["D6"] = "A property can have several units. List the units on the Units tab."
se["D6"].font = font(10, color=MUTED, italic=True)
PROP_LIST = "Settings!$B$6:$B$11"

# ---------------------------------------------------------------- Units
un = wb.create_sheet("Units")
sheet_base(un, "Units and leases", "One line per unit. The lease status updates by itself.",
           [3, 22, 16, 14, 11, 9, 9, 13, 13, 11, 20, 20], rows=U1 + 3)
header(un, 5, 2, ["Unit", "Property", "Tenant", "Rent", "Due day", "Grace days", "Lease start", "Lease end",
                  "Deposit", "Lease status", "Note"])
for r in range(U0, U1 + 1):
    u = units[r - U0] if r - U0 < len(units) else (None,) * 9
    for c, (v, fmt) in enumerate(zip(u, (None, None, None, MONEY, "0", "0", DATE, DATE, MONEY)), start=2):
        style(un.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (6, 7) else None)
    ends = f"(I{r}-TODAY())"
    style(un.cell(r, 11, f'=IF(B{r}="","",IF(D{r}="","Vacant",IF(AND(I{r}<>"",I{r}<TODAY()),"Lease ended",'
                         f'IF(AND(I{r}<>"",{ends}<=60),"Ends in "&{ends}&IF({ends}=1," day"," days"),"Active"))))'), False)
    style(un.cell(r, 12), True)
dv = DataValidation(type="list", formula1=f"={PROP_LIST}", allow_blank=True)
un.add_data_validation(dv)
dv.add(f"C{U0}:C{U1}")
ls = f"K{U0}:K{U1}"
un.conditional_formatting.add(ls, FormulaRule(formula=[f'$K{U0}="Active"'], fill=fill(OK)))
un.conditional_formatting.add(ls, FormulaRule(formula=[f'LEFT($K{U0},7)="Ends in"'], fill=fill(WARN)))
un.conditional_formatting.add(ls, FormulaRule(formula=[f'OR($K{U0}="Vacant",$K{U0}="Lease ended")'], fill=fill(INFO)))
un[f"B{U1 + 2}"] = "When a tenant leaves, keep the line and its lease end date so the past months stay right."
un[f"B{U1 + 2}"].font = font(10, color=MUTED, italic=True)
un.freeze_panes = "C6"

# ---------------------------------------------------------------- Rent log
lg = wb.create_sheet("Rent log")
sheet_base(lg, "Rent received", "One line per payment. Leave 'For month' empty when it is for the month it arrives in.",
           [3, 13, 22, 12, 13, 15, 28], rows=G1 + 2)
header(lg, 5, 2, ["Date received", "Unit", "Amount", "For month", "Method", "Note"])
for r in range(G0, G1 + 1):
    p = payments[r - G0] if r - G0 < len(payments) else (None,) * 6
    for c, (v, fmt) in enumerate(zip(p, (DATE, None, MONEY, DATE, None, None)), start=2):
        style(lg.cell(r, c, v), True, fmt)
    lg.cell(r, 8, f'=IF(B{r}="","",IF(E{r}="",B{r},E{r}))')   # the month this payment counts for
lg.column_dimensions["H"].hidden = True
for col, formula in (("C", f"=Units!$B${U0}:$B${U1}"), ("F", '"Bank transfer,Check,Cash,App,Other"')):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    lg.add_data_validation(dv)
    dv.add(f"{col}{G0}:{col}{G1}")
lg.freeze_panes = "B6"

# ---------------------------------------------------------------- Expenses
ex = wb.create_sheet("Expenses")
sheet_base(ex, "Expenses", "One line per cost, with the property it belongs to.", [3, 13, 18, 22, 12, 32], rows=X1 + 2)
header(ex, 5, 2, ["Date", "Property", "Category", "Amount", "Note"])
for r in range(X0, X1 + 1):
    e = expenses[r - X0] if r - X0 < len(expenses) else (None,) * 5
    for c, (v, fmt) in enumerate(zip(e, (DATE, None, None, MONEY, None)), start=2):
        style(ex.cell(r, c, v), True, fmt)
for col, formula in (("C", f"={PROP_LIST}"), ("D", "=Settings!$B$15:$B$24")):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ex.add_data_validation(dv)
    dv.add(f"{col}{X0}:{col}{X1}")
ex.freeze_panes = "B6"

# ---------------------------------------------------------------- Rent roll
rr = wb.create_sheet("Rent roll")
sheet_base(rr, "Rent roll", "Rent received per unit and month for the year on the Dashboard.",
           [3, 20, 12, 10] + [9] * 12 + [11, 11, 22], rows=U1 + 4)
header(rr, 5, 2, ["Unit", "Tenant", "Rent"] + MONTHS + ["Received", "Owed now", "This month"])
EXP0, OWE0 = 20, 32   # hidden: expected rent per month (T..AE), then rent owed per month (AF..AQ)
for r in range(U0, U1 + 1):
    style(rr.cell(r, 2, f'=IF(Units!B{r}="","",Units!B{r})'), False, bold=True)
    style(rr.cell(r, 3, f'=IF(B{r}="","",Units!D{r})'), False)
    style(rr.cell(r, 4, f'=IF(B{r}="","",Units!E{r})'), False, MONEY)
    for m in range(12):
        ms, me = f"DATE({YR},{m + 1},1)", f"EOMONTH(DATE({YR},{m + 1},1),0)"
        got, exp_ = L(5 + m), L(EXP0 + m)
        paid = f'SUMIFS({LG("D")},{LG("C")},$B{r},{LG("H")},">="&{ms},{LG("H")},"<="&{me})'
        style(rr.cell(r, 5 + m, f'=IF($B{r}="","",IF(AND({paid}=0,{ms}>TODAY()),"",{paid}))'), False, "#,##0")
        rr.cell(r, EXP0 + m, f'=IF($B{r}="",0,IF(AND(OR(Units!$H{r}="",Units!$H{r}<={me}),OR(Units!$I{r}="",Units!$I{r}>={ms})),'
                             f'N(Units!$E{r}),0))')
        due = f"DATE({YR},{m + 1},MIN(N(Units!$F{r})+(N(Units!$F{r})=0),DAY({me})))+N(Units!$G{r})"
        rr.cell(r, OWE0 + m, f"=IF(AND({exp_}{r}>0,{due}<TODAY()),MAX(0,{exp_}{r}-N({got}{r})),0)")
    style(rr.cell(r, 17, f'=IF(B{r}="","",SUM(E{r}:P{r}))'), False, MONEY, True)
    style(rr.cell(r, 18, f'=IF(B{r}="","",SUM({L(OWE0)}{r}:{L(OWE0 + 11)}{r}))'), False, MONEY, True)
    m0 = "MONTH(TODAY())"
    want = f"INDEX({L(EXP0)}{r}:{L(EXP0 + 11)}{r},{m0})"
    got = f"N(INDEX(E{r}:P{r},{m0}))"
    due_day = f"MIN(N(Units!$F{r})+(N(Units!$F{r})=0),DAY(EOMONTH(TODAY(),0)))"
    th = f'IF(AND({due_day}>=11,{due_day}<=13),"th",IF(MOD({due_day},10)=1,"st",IF(MOD({due_day},10)=2,"nd",IF(MOD({due_day},10)=3,"rd","th"))))'
    style(rr.cell(r, 19, f'=IF(B{r}="","",IF({YR}<>YEAR(TODAY()),"",IF({want}=0,"No rent due",IF({got}>={want},"Paid",'
                         f'IF(TODAY()>DATE(YEAR(TODAY()),{m0},{due_day})+N(Units!$G{r}),"Late: "&FIXED({want}-{got},2)&" owed",'
                         f'IF({got}>0,"Part paid: "&FIXED({want}-{got},2)&" left","Due on the "&{due_day}&{th}))))))'), False)
    rr.cell(r, OWE0 + 12, f'=IF(B{r}="","",Units!C{r})')   # property, for totals per property
for c in range(EXP0, OWE0 + 13):
    rr.column_dimensions[L(c)].hidden = True
grid = f"E{U0}:P{U1}"
rr.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{U0}<>"",{L(EXP0)}{U0}>0,N(E{U0})>={L(EXP0)}{U0})'], fill=fill(OK)))
rr.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{U0}<>"",{L(OWE0)}{U0}>0)'], fill=fill(BAD),
                                                font=Font(name=F, bold=True, color="9B1C1C")))
rr.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{U0}<>"",{L(EXP0)}{U0}>0,N(E{U0})>0,N(E{U0})<{L(EXP0)}{U0})'],
                                                fill=fill(WARN)))
rr.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{U0}<>"",{L(EXP0)}{U0}=0)'], fill=fill("EEF1F1"),
                                                font=Font(name=F, color="9AA7A9")))
sm = f"S{U0}:S{U1}"
rr.conditional_formatting.add(sm, FormulaRule(formula=[f'$S{U0}="Paid"'], fill=fill(OK)))
rr.conditional_formatting.add(sm, FormulaRule(formula=[f'LEFT($S{U0},4)="Late"'], fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
rr.conditional_formatting.add(sm, FormulaRule(formula=[f'LEFT($S{U0},4)="Part"'], fill=fill(WARN)))
rr.conditional_formatting.add(f"R{U0}:R{U1}", FormulaRule(formula=[f"N($R{U0})>0"], fill=fill(BAD)))
key = U1 + 2
rr.cell(key, 2, "Colours:").font = font(10, True, MUTED)
for k, (label, colour) in enumerate([("Paid in full", OK), ("Part paid", WARN), ("Late", BAD), ("No rent due", "EEF1F1")]):
    c = rr.cell(key, 3 + k * 2, label)
    c.fill = fill(colour)
    c.font = font(10)
    c.border = box
rr.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Landlord dashboard", "Pick a year. Rent owed counts only rents whose due date and grace days have passed.",
           [3, 16, 15, 15, 15, 15, 15, 3, 24, 13, 14, 14, 3], rows=60)
db["B4"] = "Year"
db["B4"].font = font(11, True)
style(db["C4"], True, "0", True, "center")
db["C4"] = Y
exp_year = f'SUMIFS({XP("E")},{XP("B")},">="&DATE($C$4,1,1),{XP("B")},"<="&DATE($C$4,12,31))'
occupied = f'COUNTIFS({UN("K")},"Active")+COUNTIFS({UN("K")},"Ends in*")'
all_units = f'COUNTIFS({UN("B")},"?*")'
tiles = [("B", "Rent received", f"=SUM({RR('Q')})", TEAL, MONEY),
         ("C", "Rent owed now", f"=SUM({RR('R')})", "B4541F", MONEY),
         ("D", "Expenses", f"={exp_year}", "B4541F", MONEY),
         ("E", "Net income", "=B7-D7", TEAL, NET),
         ("F", "Units let", f'={occupied}&" of "&{all_units}', TEAL, None),
         ("G", "Deposits held", f'=SUMIFS({UN("J")},{UN("K")},"Active")+SUMIFS({UN("J")},{UN("K")},"Ends in*")', TEAL, MONEY)]
for col, label, formula, colour, fmt in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    if fmt:
        db[f"{col}7"].number_format = fmt
    db[f"{col}7"].font = font(16, True, colour)
    db[f"{col}7"].alignment = Alignment(horizontal="right")
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db.row_dimensions[7].height = 34

header(db, 9, 2, ["Month", "Rent due", "Received", "Expenses", "Net"])
for m in range(12):
    r = 10 + m
    ms, me = f"DATE($C$4,{m + 1},1)", f"EOMONTH(DATE($C$4,{m + 1},1),0)"
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f"=SUM('Rent roll'!{L(EXP0 + m)}{U0}:{L(EXP0 + m)}{U1})"), False, MONEY)
    style(db.cell(r, 4, f"=SUM('Rent roll'!{L(5 + m)}{U0}:{L(5 + m)}{U1})"), False, MONEY)
    style(db.cell(r, 5, f'=SUMIFS({XP("E")},{XP("B")},">="&{ms},{XP("B")},"<="&{me})'), False, MONEY)
    style(db.cell(r, 6, f"=D{r}-E{r}"), False, NET, True)
style(db.cell(22, 2, "Year"), False, bold=True, align="center")
for col in "CDEF":
    style(db[f"{col}22"], False, NET if col == "F" else MONEY, True)
    db[f"{col}22"] = f"=SUM({col}10:{col}21)"
db["B23"] = "Net is the rent received minus expenses. Rent due counts every month with an active lease."
db["B23"].font = font(9, color=MUTED, italic=True)
db.conditional_formatting.add("B10:F21", FormulaRule(formula=["AND($C$4=YEAR(TODAY()),ROW()-9=MONTH(TODAY()))"], fill=fill("E6F2F2")))

late_n = f'COUNTIFS({RR("R")},">0")'
ending = f'COUNTIFS({UN("K")},"Ends in*")'
empty = f'COUNTIFS({UN("K")},"Vacant")+COUNTIFS({UN("K")},"Lease ended")'
part = f'COUNTIFS({RR("S")},"Part paid*")'
notes = [
    (f'=IF({late_n}=0,"No rent is owed",IF({late_n}=1,"1 unit owes",{late_n}&" units owe")&" rent: "&FIXED(SUM({RR("R")}),2))', BAD, "No"),
    (f'=IF({ending}=0,"No lease ends in the next 60 days",IF({ending}=1,"1 lease ends",{ending}&" leases end")&" in the next 60 days")', WARN, "No"),
    (f'=IF({empty}=0,"Every unit is let",IF({empty}=1,"1 unit is",{empty}&" units are")&" empty")', INFO, "Every"),
    (f'=IF({part}=0,"No part payments this month",IF({part}=1,"1 unit has",{part}&" units have")&" paid part of this month")', INFO, "No"),
]
db.merge_cells("I9:L9")
db["I9"] = "Needs attention"
db["I9"].font = font(11, True, "FFFFFF")
db["I9"].fill = fill(TEAL)
db["I9"].alignment = Alignment(vertical="center")
for k, (formula, colour, calm) in enumerate(notes):
    r = 10 + k
    db.merge_cells(f"I{r}:L{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "IJKL":
        db[f"{col}{r}"].border = box
    db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))

header(db, 15, 9, ["Property", "Units", "Rent received", "Net"])
db.row_dimensions[15].height = None
for k in range(6):
    r = 16 + k
    p = f"Settings!B{6 + k}"
    style(db.cell(r, 9, f'=IF({p}="","",{p})'), False, bold=True)
    style(db.cell(r, 10, f'=IF(I{r}="","",COUNTIFS({UN("C")},I{r}))'), False, "0", align="center")
    style(db.cell(r, 11, f"=IF(I{r}=\"\",\"\",SUMIFS({RR('Q')},'Rent roll'!${L(OWE0 + 12)}${U0}:${L(OWE0 + 12)}${U1},I{r}))"), False, MONEY)
    style(db.cell(r, 12, f'=IF(I{r}="","",K{r}-SUMIFS({XP("E")},{XP("C")},I{r},{XP("B")},">="&DATE($C$4,1,1),{XP("B")},"<="&DATE($C$4,12,31)))'),
          False, NET, True)

chart = BarChart()
chart.type = "col"
chart.title = "Rent received and expenses by month"
chart.height, chart.width = 7.5, 16
chart.add_data(Reference(db, min_col=4, max_col=5, min_row=9, max_row=21), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=10, max_row=21))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E0A21B"
db.add_chart(chart, "B25")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: type the names of your properties. A building with several flats is one property."),
    ("2", "Units: one line per unit, with the tenant, rent, due day, grace days, lease dates and deposit."),
    ("3", "Rent log: one line each time rent arrives. Fill in 'For month' only for rent paid early or late."),
    ("4", "Expenses: one line per cost, with its property and category."),
    ("5", "Rent roll and Dashboard: every month shows paid, part paid or late, with what is owed and your net income."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Rent counts as late once the due day and the grace days have passed.",
                          "The example units, payments and expenses show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency. It does not give tax advice."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Landlord-Rental-Tracker.xlsx")
wb.save(out)
print("saved", out, len(payments), "payments", len(expenses), "expenses")
