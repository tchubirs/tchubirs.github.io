"""Social Media Content Planner (Excel + Google Sheets).

Every post with its date, platform, type, topic and status, and its results once it is out; a month calendar to
print with the posts of each day; followers by month for each platform; and a Dashboard with what is coming up, what
is late, each platform against its weekly goal, the best posts of the month and the posts published each week. Only
functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MUTED, OK, TEAL, TEAL_D, WARN, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
P0, P1 = 6, 805        # posts
NPL = 8                # platforms
NS = 12                # months on Stats
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
MONTHS_LONG = ('"January","February","March","April","May","June","July","August","September","October",'
               '"November","December"')
PS = lambda c: f"Posts!${c}${P0}:${c}${P1}"
PS1 = lambda c: f"Posts!${c}$1:${c}${P1}"
PLAT = f"Settings!$B$10:$B${9 + NPL}"
CODE = f"Settings!$C$10:$C${9 + NPL}"
MONTH1 = "Settings!$C$5"          # first day of the month on the calendar
STATS1 = "Settings!$C$6"          # first month on Stats
BIG = 10 ** 7
STATUSES = ["Idea", "Draft", "Scheduled", "Published"]
TYPES = ["Post", "Carousel", "Reel", "Story", "Video", "Short", "Live", "Pin", "Email", "Thread"]


def add_months(d, n):
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    last = [31, 29 if y % 4 == 0 and (y % 100 != 0 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return dt.date(y, m, min(d.day, last[m - 1]))


def g(x):
    """An empty cell as empty text, not 0."""
    return f'IF({x}="","",{x})'


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


rng = random.Random(40)
monday = today - dt.timedelta(days=today.weekday())

# ---------------------------------------------------------------- example data (a made-up small candle shop)
platforms = [("Instagram", "IG", 4), ("TikTok", "TT", 3), ("YouTube", "YT", 1), ("Pinterest", "PIN", 3),
             ("Newsletter", "NL", 1)]
pillars = ["Behind the scenes", "How to", "New products", "Customer stories", "Tips"]
titles = {
    "Behind the scenes": ["Pouring a new batch", "Studio tidy up", "Packing orders", "Testing a new wick",
                          "A day at the market"],
    "How to": ["How to trim a wick", "How to get an even burn", "Styling a shelf", "Care for your candle"],
    "New products": ["Autumn scents are here", "Meet the new tin", "Gift box sneak peek", "Refill jars"],
    "Customer stories": ["Your photos this month", "A lovely review", "Candle in a new home"],
    "Tips": ["Three cosy evening ideas", "Scent for small rooms", "Reuse your jar"],
}
kinds = {"Instagram": ["Post", "Carousel", "Reel", "Story"], "TikTok": ["Short"], "YouTube": ["Video"],
         "Pinterest": ["Pin"], "Newsletter": ["Email"]}
reach0 = {"Instagram": 1800, "TikTok": 3200, "YouTube": 900, "Pinterest": 1200, "Newsletter": 640}
posts = []
for w in range(-7, 3):
    for name, code, per_week in platforms:
        for n in range(per_week):
            day = monday + dt.timedelta(days=7 * w + [1, 3, 5, 0][n % 4] + (1 if name == "TikTok" else 0))
            pillar = rng.choice(pillars)
            title = rng.choice(titles[pillar])
            t = dt.time(rng.choice([9, 12, 18, 19]), rng.choice([0, 30]))
            if day < today:
                status = "Published"
            elif day <= today + dt.timedelta(days=6):
                status = rng.choice(["Scheduled", "Scheduled", "Draft"])
            else:
                status = rng.choice(["Draft", "Idea"])
            row = [day, t, name, rng.choice(kinds[name]), pillar, title, status, None]
            if status == "Published":
                reach = round(reach0[name] * rng.uniform(0.6, 1.6))
                likes = round(reach * rng.uniform(0.02, 0.08))
                row += [reach, likes, round(likes * rng.uniform(0.05, 0.2)), round(likes * rng.uniform(0.02, 0.1)),
                        round(likes * rng.uniform(0.05, 0.3))]
            else:
                row += [None] * 5
            posts.append(row)
posts[-3][6] = "Idea"
late = [p for p in posts if p[6] == "Published" and p[0] >= today - dt.timedelta(days=3)][:2]
for p in late:                                                  # two posts that did not go out on time
    p[6] = "Scheduled"
    p[8:13] = [None] * 5
posts.sort(key=lambda p: (p[0], p[1]))
for title, pillar, name in (("Candle making kit video", "How to", "YouTube"), ("Holiday gift guide", "New products",
                            "Pinterest"), ("Ask me anything", "Behind the scenes", "Instagram"),
                            ("Scent quiz", "Tips", "TikTok")):
    posts.append([None, None, name, kinds[name][0], pillar, title, "Idea", None] + [None] * 5)
posts[3][7] = "Link in bio"
stats_start = add_months(today.replace(day=1), -11)
followers = {"Instagram": (4200, 6100), "TikTok": (1500, 4800), "YouTube": (380, 720), "Pinterest": (900, 1500),
             "Newsletter": (610, 980)}

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your platforms and their weekly goals, your topics, and the month the calendar shows.",
           [3, 26, 16, 16, 4, 60], rows=30)
for r, lab, value, fmt, typed, note in (
        (4, "Show another month", None, "mmm yyyy", True, "Leave it empty for this month. Type any day of another."),
        (5, "The calendar shows", '=IF(C4="",DATE(YEAR(TODAY()),MONTH(TODAY()),1),DATE(YEAR(C4),MONTH(C4),1))',
         "mmmm yyyy", False, "The month on the Calendar tab."),
        (6, "First month on Stats", dt.datetime.combine(stats_start, dt.time()), "mmm yyyy", True,
         "Stats has 12 months from this one.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], typed, fmt, True, "center")
    se[f"C{r}"] = value
    se.cell(r, 6, note).font = font(10, color=MUTED, italic=True)
header(se, 9, 2, ["Platform", "Short name", "Posts a week"])
for k in range(NPL):
    r = 10 + k
    x = platforms[k] if k < len(platforms) else (None, None, None)
    style(se.cell(r, 2, x[0]), True, bold=True)
    style(se.cell(r, 3, x[1]), True, align="center")
    style(se.cell(r, 4, x[2]), True, "0", align="center")
header(se, 20, 2, ["Topic"])
for k in range(8):
    style(se.cell(21 + k, 2, pillars[k] if k < len(pillars) else None), True)
se["F9"] = "The short name goes on the calendar, next to each post."
se["F9"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Posts
ps = wb.create_sheet("Posts")
sheet_base(ps, "Posts", "Every post, from the idea to the results. The status moves along as it goes.",
           [3, 12, 8, 12, 11, 17, 30, 11, 16, 10, 9, 10, 9, 9, 12, 9, 9], rows=P1 + 2)
header(ps, 5, 2, ["Date", "Time", "Platform", "Type", "Topic", "Post", "Status", "Link or note", "Reach", "Likes",
                  "Comments", "Shares", "Saves", "Interactions", "Rate", "Check"])
this_month = "(YEAR(TODAY())*100+MONTH(TODAY()))"
for r in range(P0, P1 + 1):
    x = posts[r - P0] if r - P0 < len(posts) else (None,) * 13
    fmts = (DATE, "h:mm", None, None, None, None, None, None, "#,##0", "#,##0", "#,##0", "#,##0", "#,##0")
    for c, v, fmt in zip(range(2, 15), x, fmts):
        style(ps.cell(r, c, v), True, fmt, c == 7, "center" if c in (2, 3, 5, 8) else None)
    style(ps.cell(r, 15, f'=IF(COUNT(K{r}:N{r})=0,"",SUM(K{r}:N{r}))'), False, "#,##0")
    style(ps.cell(r, 16, f'=IF(OR(O{r}="",N(J{r})<=0),"",O{r}/J{r})'), False, "0.0%", align="center")
    ps[f"V{r}"] = (f'=IF(AND(ISNUMBER(B{r}),OR(H{r}="Draft",H{r}="Scheduled")),IF(INT(B{r})<TODAY(),1,0),0)')  # V: late
    style(ps.cell(r, 17, f'=IF(V{r}=1,"Late","")'), False, bold=True, align="center")
    mins = f'IF(ISNUMBER(C{r}),ROUND(MOD(C{r},1)*1440,0),0)'
    ps[f"R{r}"] = f'=IF(ISNUMBER(B{r}),INT(B{r})*{BIG}+{mins}*1000+ROW(),"")'               # R: in date order
    ps[f"S{r}"] = f'=IF(ISNUMBER(B{r}),INT(B{r})-WEEKDAY(B{r},3),"")'                       # S: its Monday
    ps[f"T{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                      # T: its month
    ps[f"U{r}"] = f'=IF(AND(ISNUMBER(B{r}),H{r}<>"Published"),IF(INT(B{r})>=TODAY(),R{r},""),"")'  # U: coming up
    ps[f"W{r}"] = (f'=IF(AND(H{r}="Published",N(O{r})>0),IF(T{r}={this_month},'
                   f'ROUND(O{r},0)*1000+1000-ROW(),""),"")')                                # W: best this month
for c in "RSTUVW":
    ps.column_dimensions[c].hidden = True
dropdown(ps, f"={PLAT}", f"D{P0}:D{P1}")
dropdown(ps, '"' + ",".join(TYPES) + '"', f"E{P0}:E{P1}", strict=False)
dropdown(ps, "=Settings!$B$21:$B$28", f"F{P0}:F{P1}", strict=False)
dropdown(ps, '"' + ",".join(STATUSES) + '"', f"H{P0}:H{P1}")
for text, colour in (("Published", OK), ("Scheduled", "E3E8FF"), ("Draft", WARN)):
    ps.conditional_formatting.add(f"H{P0}:H{P1}", FormulaRule(formula=[f'H{P0}="{text}"'], fill=fill(colour)))
ps.conditional_formatting.add(f"Q{P0}:Q{P1}", FormulaRule(formula=[f'Q{P0}="Late"'], fill=fill(BAD)))
ps.freeze_panes = "D6"

# ---------------------------------------------------------------- Calendar (a month to print)
ca = wb.create_sheet("Calendar")
sheet_base(ca, "", "", [3] + [21] * 7, rows=13)
ca["B2"] = f"=CHOOSE(MONTH({MONTH1}),{MONTHS_LONG})&\" \"&YEAR({MONTH1})"
ca["B2"].font = font(22, True, TEAL_D)
ca["B3"] = "The posts planned for each day, with the short name of the platform. Settings shows another month."
ca["B3"].font = font(10, color=MUTED, italic=True)
header(ca, 5, 2, ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"])
grid0 = f"({MONTH1}-WEEKDAY({MONTH1},3))"
for w in range(6):
    r = 6 + w
    ca.row_dimensions[r].height = 74
    for d in range(7):
        date = f"({grid0}+{7 * w + d})"
        base = 10 + d * 4                                                                   # J.. : helpers
        first = f"COUNTIF({PS('R')},\"<\"&{date}*{BIG})"
        for s in range(3):
            key = f"SMALL({PS('R')},{first}+{s + 1})"
            ca.cell(r, base + s, f'=IFERROR(IF(INT({key}/{BIG})={date},MOD({key},1000),""),"")')
        ca.cell(r, base + 3, f'=COUNTIFS({PS("R")},">="&{date}*{BIG},{PS("R")},"<"&({date}+1)*{BIG})')
        lines = []
        for s in range(3):
            h = f"{col(base + s)}{r}"
            plat = f"INDEX({PS1('D')},{h})"
            code = f'IFERROR(INDEX({CODE},MATCH({plat},{PLAT},0)),"")'
            name = f'IF({code}="",{plat},{code})'
            what = f'IF(INDEX({PS1("G")},{h})="",INDEX({PS1("E")},{h}),INDEX({PS1("G")},{h}))'
            lines.append(f'IF({h}="","",CHAR(10)&{name}&": "&{what})')
        more = f'IF({col(base + 3)}{r}>3,CHAR(10)&"and "&({col(base + 3)}{r}-3)&" more","")'
        text = f'=IF(MONTH({date})<>MONTH({MONTH1}),"",DAY({date})&' + "&".join(lines) + f"&{more})"
        c = style(ca.cell(r, 2 + d, text), False)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = font(9)
    for cc in range(10, 38):
        ca.column_dimensions[col(cc)].hidden = True
ca["AL1"] = f"={grid0}"                                                                      # AL1: first day shown
ca.column_dimensions["AL"].hidden = True
ca.conditional_formatting.add("B6:H11", FormulaRule(
    formula=['AND(B6<>"",$AL$1+(ROW(B6)-6)*7+(COLUMN(B6)-2)=TODAY())'], fill=fill("E3F1EF")))
ca.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Stats
sx = wb.create_sheet("Stats")
sheet_base(sx, "Stats", "Followers on each platform at the end of each month.", [3, 18] + [10] * NS + [11, 10],
           rows=24)
header(sx, 5, 2, ["Platform"])
for k in range(NS):
    c = sx.cell(5, 3 + k, f"=EDATE({STATS1},{k})")
    c.number_format = "mmm yy"
    c.font = font(10, True, "FFFFFF")
    c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
header(sx, 5, 3 + NS, ["Growth", "Growth %"])
for i in range(NPL):
    r = 6 + i
    name = platforms[i][0] if i < len(platforms) else None
    style(sx.cell(r, 2, name), True, bold=True)
    for k in range(NS):
        v = None
        if name and stats_start <= add_months(stats_start, k) <= today:
            a, b = followers[name]
            v = round(a + (b - a) * (k / (NS - 1)) ** 1.3)
        style(sx.cell(r, 3 + k, v), True, "#,##0")
    rowr = f"C{r}:{col(2 + NS)}{r}"
    sx[f"R{r}"] = f"=SUMPRODUCT(MIN(ISNUMBER({rowr})*(COLUMN({rowr})-2)+(1-ISNUMBER({rowr}))*999))"  # R: first
    sx[f"S{r}"] = f"=SUMPRODUCT(MAX(ISNUMBER({rowr})*(COLUMN({rowr})-2)))"                       # S: last
    first, last = f"INDEX({rowr},R{r})", f"INDEX({rowr},S{r})"
    style(sx.cell(r, 3 + NS, f'=IF(COUNT({rowr})<2,"",{last}-{first})'), False, "#,##0")
    style(sx.cell(r, 4 + NS, f'=IF(OR({col(3 + NS)}{r}="",COUNT({rowr})<2),"",IF({first}<=0,"",'
                             f'{col(3 + NS)}{r}/{first}))'), False, "0%", align="center")
r = 6 + NPL + 1
sx.cell(r, 2, "All platforms").font = font(11, True, TEAL_D)
for k in range(NS):
    L = col(3 + k)
    style(sx.cell(r, 3 + k, f'=IF(COUNT({L}6:{L}{5 + NPL})=0,"",SUM({L}6:{L}{5 + NPL}))'), False, "#,##0", True)
    sx.cell(r + 1, 3 + k, f"=IF(COUNT({L}6:{L}{5 + NPL})>0,{k + 1},0)")                   # month in use
sx.row_dimensions[r + 1].hidden = True
TOTAL_ROW = r
for c in "RS":
    sx.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Content planner", "", [3, 12, 12, 30, 12, 3, 24, 12, 12, 12, 3], rows=50)
db["B3"] = f'="The week of "&{short("$M$1")}&", and "&CHOOSE(MONTH(TODAY()),{MONTHS_LONG})&" so far."'
db["B3"].font = font(11, color=MUTED, italic=True)
db["M1"] = "=TODAY()-WEEKDAY(TODAY(),3)"                                                      # M1: this Monday
db["M2"] = f"=MAX(Stats!C{TOTAL_ROW + 1}:{col(2 + NS)}{TOTAL_ROW + 1})"                     # M2: latest Stats month
wk = f'{PS("B")},">="&$M$1,{PS("B")},"<"&($M$1+7)'
tile(db, "B", 4, "This week", f'=COUNTIFS({wk})-COUNTIFS({wk},{PS("H")},"Idea")', "0")
tile(db, "C", 4, "Next 7 days",
     f'=COUNTIFS({PS("B")},">="&TODAY(),{PS("B")},"<"&(TODAY()+7),{PS("H")},"Scheduled")'
     f'+COUNTIFS({PS("B")},">="&TODAY(),{PS("B")},"<"&(TODAY()+7),{PS("H")},"Draft")', "0")
tile(db, "D", 4, "Late", f"=SUM({PS('V')})", "0", "B23B2E")
tile(db, "E", 4, "Ideas", f'=COUNTIF({PS("H")},"Idea")', "0")
tile(db, "G", 4, "Published this month", f'=COUNTIFS({PS("H")},"Published",{PS("T")},{this_month})', "0")
tile(db, "H", 4, "Reach this month", f'=SUMIFS({PS("J")},{PS("H")},"Published",{PS("T")},{this_month})', "#,##0")
tile(db, "I", 4, "Rate this month",
     f'=IF(H5=0,"",SUMIFS({PS("O")},{PS("H")},"Published",{PS("T")},{this_month},{PS("J")},">0")/H5)', "0.0%")
tile(db, "J", 4, "Followers", f'=IF($M$2=0,"",INDEX(Stats!$C${TOTAL_ROW}:${col(2 + NS)}${TOTAL_ROW},$M$2))', "#,##0")

header(db, 7, 2, ["Coming up", "Platform", "Post", "Status"])
for n in range(10):
    r = 8 + n
    db[f"N{r}"] = f"=IFERROR(SMALL({PS('U')},{n + 1}),\"\")"                                   # N: key
    a = lambda c: f"INDEX({PS1(c)},MOD(N{r},1000))"
    style(db.cell(r, 2, f'=IF(N{r}="","",INT(N{r}/{BIG}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 3, f'=IF(N{r}="","",{g(a("D"))})'), False, bold=True)
    style(db.cell(r, 4, f'=IF(N{r}="","",IF({a("G")}="",{g(a("E"))},{a("G")}))'), False)
    style(db.cell(r, 5, f'=IF(N{r}="","",{g(a("H"))})'), False, align="center")
db.conditional_formatting.add("E8:E17", FormulaRule(formula=['E8="Draft"'], fill=fill(WARN)))
db.conditional_formatting.add("E8:E17", FormulaRule(formula=['E8="Idea"'], fill=fill(BAD)))

header(db, 7, 7, ["Platform", "This week", "Goal a week", "Rate"])
for k in range(NPL):
    r = 8 + k
    p = f"Settings!$B${10 + k}"
    style(db.cell(r, 7, f'=IF({p}="","",{p})'), False, bold=True)
    pw = f'{PS("D")},G{r},{wk}'
    style(db.cell(r, 8, f'=IF(G{r}="","",COUNTIFS({pw})-COUNTIFS({pw},{PS("H")},"Idea"))'), False, "0",
          align="center")
    style(db.cell(r, 9, f'=IF(G{r}="","",N(Settings!$D${10 + k}))'), False, "0", align="center")
    reach = f'SUMIFS({PS("J")},{PS("D")},G{r},{PS("H")},"Published",{PS("T")},{this_month})'
    inter = f'SUMIFS({PS("O")},{PS("D")},G{r},{PS("H")},"Published",{PS("T")},{this_month},{PS("J")},">0")'
    style(db.cell(r, 10, f'=IF(G{r}="","",IF({reach}=0,"",{inter}/{reach}))'), False, "0.0%", align="center")
db.conditional_formatting.add("H8:H15", FormulaRule(formula=["AND(ISNUMBER(H8),H8<N(I8))"], fill=fill(WARN)))

header(db, 19, 2, ["Week of", "Published", "Reach"])
for n in range(8):
    r = 20 + n
    wk_n = f"($M$1-{7 * (7 - n)})"
    style(db.cell(r, 2, f"={wk_n}"), False, "d mmm", True, "center")
    ww = f'{PS("S")},B{r},{PS("H")},"Published"'
    style(db.cell(r, 3, f"=COUNTIFS({ww})"), False, "0", align="center")
    style(db.cell(r, 4, f"=SUMIFS({PS('J')},{ww})"), False, "#,##0", align="center")
header(db, 19, 7, ["Best this month", "Platform", "Interactions", "Rate"])
for n in range(5):
    r = 20 + n
    db[f"O{r}"] = f"=IFERROR(LARGE({PS('W')},{n + 1}),\"\")"                                   # O: key
    a = lambda c: f"INDEX({PS1(c)},1000-MOD(O{r},1000))"
    style(db.cell(r, 7, f'=IF(O{r}="","",IF({a("G")}="",{g(a("E"))},{a("G")}))'), False, bold=True)
    style(db.cell(r, 8, f'=IF(O{r}="","",{g(a("D"))})'), False, align="center")
    style(db.cell(r, 9, f'=IF(O{r}="","",{a("O")})'), False, "#,##0", align="center")
    style(db.cell(r, 10, f'=IF(O{r}="","",IF({a("P")}="","",{a("P")}))'), False, "0.0%", align="center")
for c in "MNO":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 18)) + list(range(20, 28)):
    db.row_dimensions[r].height = 18
    for c in range(2, 11):
        cell = db.cell(r, c)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="center")
chart = BarChart()
chart.type = "col"
chart.title = "Published each week"
chart.height, chart.width = 6.2, 11.5
chart.add_data(Reference(db, min_col=3, min_row=19, max_row=27), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=20, max_row=27))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.y_axis.number_format = "0"
db.add_chart(chart, "G26")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your platforms with a short name and how many posts a week you aim for, and your topics."),
    ("2", "Posts: each post with its date, platform, type, topic and status. Ideas can wait without a date."),
    ("3", "Once a post is out: its reach, likes, comments, shares and saves. The rate works itself out."),
    ("4", "Calendar: the month with the posts of each day, ready to print."),
    ("5", "Stats: your followers on each platform at the end of each month."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example shop, its posts and its numbers are made up. Delete them and add your own.",
                          "Rate is interactions divided by reach: likes, comments, shares and saves together.",
                          "Works in Google Sheets and Microsoft Excel."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Calendar", "Posts", "Stats", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Content-Planner.xlsx")
wb.save(out)
print("saved", out, len(posts), "posts")
