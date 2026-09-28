"""Etsy listing photos for the Content Planner (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Content-Planner.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups: the latest posts, and the Stats rows in use.
wb = load_workbook(xlsx)
ps = wb["Posts"]
last = max(r for r in range(6, 806) if any(ps.cell(r, c).value not in (None, "") for c in range(2, 9)))
import datetime as dt  # noqa: E402
week_ago = dt.datetime.now() - dt.timedelta(days=8)
start = min(r for r in range(6, last + 1) if ps.cell(r, 2).value is not None and ps.cell(r, 2).value >= week_ago)
stats_last = max(r for r in range(6, 14) if wb["Stats"].cell(r, 2).value not in (None, ""))
rows = {"Posts": list(range(6, start)) + list(range(start + 26, 806)),                  # the week before and after
        "Stats": range(stats_last + 1, 14)}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Calendar"))
close = Renders(xlsx, tmp, one_page=("Dashboard", "Calendar"), hide_rows=rows, hide={"Posts": ["I"]})

img = canvas(dark=True)
title(img, "Social media content planner", "Every post from the idea to the results, a calendar to print, and "
      "each platform against its weekly goal", dark=True, width=2300)
laptop(img, r.shot("Content planner"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

month = next(t for t in r.titles if t and t[0].isalpha() and t.split()[-1].isdigit())
detail(os.path.join(img_dir, "02-calendar.jpg"), "A month of posts to print",
       "The posts of each day with the short name of the platform. Today is shaded.", r.shot(month))

detail(os.path.join(img_dir, "03-posts.jpg"), "Every post, from the idea to the results",
       "Date, platform, type, topic and status, then reach, likes, comments, shares and saves once it is out",
       close.shot("Posts"), [("#D8F0DC", "Published"), ("#E3E8FF", "Scheduled"), ("#FFE2B8", "Draft"),
                             ("#F9C9C4", "Late")])

# The Dashboard from the weeks down.
page = r.doc[r._index("Content planner")]
words = page.get_text("words")
week = next(w for w in words if w[4] == "Week" and w[1] > 200)
full = r.page("Content planner")
cut = int((week[1] - 14) * 220 / 72)
detail(os.path.join(img_dir, "04-what-worked.jpg"), "What worked this month",
       "The posts published each week with their reach, and the best posts of the month",
       content(full.crop((0, cut, full.width, full.height))))

img = canvas()
y = title(img, "Inside the file", "Four tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Content planner")), "Dashboard", "The week and the month"),
    (content(r.page(month)), "Calendar", "A month to print"),
    (content(close.page("Posts")), "Posts", "Plans and results"),
    (content(close.page("Stats")), "Stats", "Followers by month"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "05-whats-inside.jpg"))
