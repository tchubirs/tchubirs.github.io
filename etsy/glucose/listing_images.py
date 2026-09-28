"""Etsy listing photos for the Blood Sugar and Medicine Log (run build.py first)."""
import datetime as dt
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Blood-Sugar-Log.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups: the last ten days of readings, and the medicines typed in.
wb = load_workbook(xlsx)
day = lambda v: v.date() if isinstance(v, dt.datetime) else v
rd = wb["Readings"]
rows = [r for r in range(6, 1506) if rd.cell(r, 2).value not in (None, "")]
recent = [r for r in rows if day(rd.cell(r, 2).value) >= dt.date.today() - dt.timedelta(days=9)]
meds_end = max(r for r in range(6, 26) if wb["Medicines"].cell(r, 2).value not in (None, ""))

r = Renders(xlsx, tmp, one_page=("Dashboard", "Report"))
close = Renders(xlsx, tmp, one_page=("Dashboard", "Report"),
                hide_rows={"Readings": [x for x in range(6, 1506) if x not in recent],
                           "Medicines": range(meds_end + 1, 26)})

img = canvas(dark=True)
title(img, "Blood sugar and medicine log", "Readings against the range you type, what is still to take today, and "
      "a one-page report for an appointment", dark=True, width=2300)
laptop(img, r.shot("Blood sugar and medicines"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-report.jpg"), "One page for your appointment",
       "A summary for each moment of the day, the medicines, and every reading day by day. Print it or save it as a PDF.",
       r.shot("Blood sugar report"), [("#F9C9C4", "Below your range"), ("#FFE2B8", "Above your range")])

detail(os.path.join(img_dir, "03-readings.jpg"), "Every reading in its place",
       "The date, time and moment of the day, with notes on food, exercise or how you felt",
       close.shot("Readings"), [("#F9C9C4", "Below"), ("#D8F0DC", "In range"), ("#FFE2B8", "Above")])

detail(os.path.join(img_dir, "04-medicines.jpg"), "Medicines and their times",
       "Up to four times a day. Log each dose and the Dashboard shows what is still to take today.",
       close.shot("Medicines"), [("#D8F0DC", "Taking"), ("#E3E8FF", "Stopped")])

# The Dashboard from the two weeks table down: the numbers and the chart.
page = r.doc[r._index("Blood sugar and medicines")]
cut = int((page.search_for("Last 14 days")[0].y0 - 14) * 220 / 72)
full = r.page("Blood sugar and medicines")
lower = Image.new("RGB", (full.width, full.height - cut + 60), "#FFFFFF")
lower.paste(full.crop((0, cut, full.width, full.height)), (0, 30))
detail(os.path.join(img_dir, "05-two-weeks.jpg"), "Two weeks at a time",
       "The average, the lowest and the highest reading of each day, as numbers and as a chart",
       content(lower))

img = canvas()
y = title(img, "Inside the file", "Five tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Blood sugar and medicines")), "Dashboard", "Today and this week"),
    (content(r.page("Readings")), "Readings", "Every check"),
    (content(r.page("Medicines")), "Medicines", "What and when"),
    (content(r.page("Doses")), "Doses", "Each one taken"),
    (content(r.page("Blood sugar report")), "Report", "One page to print"),
    (content(r.page("Settings")), "Settings", "Your own ranges"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
