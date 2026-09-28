"""Blood Sugar and Medicine Log (Excel + Google Sheets).

Readings with the date, time and moment of the day, each marked below, inside or above the range the person types
for that moment; medicines with their times, and a Doses log that shows what is still to take today; averages for
today, the week and the month; and a Report tab, one printed page with a summary by moment, the medicines and a
day-by-day grid to take to an appointment. A log only: it gives no medical advice. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
R0, R1 = 6, 1505       # readings
M0, M1 = 6, 25         # medicines
D0, D1 = 6, 2005       # doses
G0, G1 = 10, 18        # moments and their ranges (Settings)
MOMENTS = ["Fasting", "After breakfast", "Before lunch", "After lunch", "Before dinner", "After dinner", "Bedtime",
           "Night", "Other"]
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
FULL_LIST = ('"January","February","March","April","May","June","July","August","September","October","November",'
             '"December"')
RD = lambda col: f"Readings!${col}${R0}:${col}${R1}"
MD = lambda col: f"Medicines!${col}${M0}:${col}${M1}"
DS = lambda col: f"Doses!${col}${D0}:${col}${D1}"
NAME, UNITS, FROM, TO = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6", "Settings!$C$7"
MOMENT_RANGE = f"Settings!$B${G0}:$B${G1}"
RED, AMBER = "B4541F", "B07A00"


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{FULL_LIST})&" "&YEAR({d})'


def hhmm(t):
    return f'HOUR({t})&":"&RIGHT("0"&MINUTE({t}),2)'


# ---------------------------------------------------------------- example data (a made-up person)
targets = {"Fasting": (80, 130), "After breakfast": (80, 180), "Before lunch": (80, 130), "After lunch": (80, 180),
           "Before dinner": (80, 130), "After dinner": (80, 180), "Bedtime": (90, 150), "Night": (80, 130),
           "Other": (80, 180)}
rng = random.Random(33)
readings = []   # date, time, moment, reading, notes
first = today - dt.timedelta(days=44)
d = first
while d <= today:
    plan = [("Fasting", dt.time(7, rng.choice([0, 5, 10, 15])), 116, 11),
            ("After breakfast", dt.time(9, rng.choice([30, 40, 45])), 158, 22)]
    if rng.random() < 0.6:
        plan.append(("Before dinner", dt.time(18, rng.choice([0, 15, 30])), 116, 14))
    if rng.random() < 0.55:
        plan.append(("After dinner", dt.time(20, rng.choice([15, 30, 45])), 164, 26))
    if rng.random() < 0.45:
        plan.append(("Bedtime", dt.time(22, rng.choice([15, 30])), 138, 14))
    for moment, t, mean, sd in plan:
        if d == today and t.hour > 10:
            continue
        v = max(60, round(rng.gauss(mean, sd)))
        readings.append([d, t, moment, v, None])
    d += dt.timedelta(days=1)
notes = {3: "Walked 30 min after lunch", 9: "Pasta for dinner", 15: "Felt shaky, had juice", 22: "Birthday cake",
         30: "Busy day, late dinner", 37: "Swim in the morning"}
for k, text in notes.items():
    day = first + dt.timedelta(days=k)
    rows = [x for x in readings if x[0] == day]
    if rows:
        target = rows[-1] if "dinner" in text or "cake" in text else rows[0]
        target[4] = text
        if "shaky" in text:
            target[3] = 66
        if "cake" in text:
            target[3] = 232
readings.append([today - dt.timedelta(days=5), dt.time(12, 20), "Before lunch", 71, "After the gym"])
readings.append([today - dt.timedelta(days=2), dt.time(3, 10), "Night", 102, "Woke up, checked"])
readings.sort(key=lambda x: (x[0], x[1]))
medicines = [
    # medicine, dose, time 1..4, with food, started, stopped, notes
    ("Tablet A", "1 tablet", dt.time(8, 0), dt.time(20, 0), None, None, "Yes", dt.date(2025, 3, 1), None, None),
    ("Tablet B", "1 tablet", dt.time(8, 0), None, None, None, "Yes", dt.date(2026, 5, 12), None, None),
    ("Vitamin D", "1 capsule", dt.time(8, 0), None, None, None, None, dt.date(2026, 1, 10), None, "Winter months"),
    ("Tablet C", "1 tablet", dt.time(22, 0), None, None, None, None, dt.date(2025, 9, 1), None, None),
    ("Old tablet", "1 tablet", dt.time(12, 0), None, None, None, "Yes", dt.date(2025, 3, 1), today - dt.timedelta(days=14),
     "Stopped at the last check-up"),
]
doses = []      # date, time, medicine, notes
d = first
while d <= today:
    for name, dose, *times, food, started, stopped, note in medicines:
        for t in times:
            if t is None or started > d or (stopped is not None and stopped <= d):
                continue
            if d == today and t.hour > 9:
                continue
            if rng.random() < 0.04:
                continue                                        # a missed dose now and then
            taken = (dt.datetime.combine(d, t) + dt.timedelta(minutes=rng.choice([0, 5, 10, 20]))).time()
            doses.append([d, taken, name, None])
    d += dt.timedelta(days=1)
doses.sort(key=lambda x: (x[0], x[1]))

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def band_fills(ws, rng, first_cell, low, high):
    """Below the range in red, above it in amber (low and high are cells on the same sheet)."""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({first_cell}),{low}<>"",{first_cell}<{low})'],
                                                   fill=fill(BAD)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({first_cell}),{high}<>"",{first_cell}>{high})'],
                                                   fill=fill(WARN)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your name, the units, the dates of the report, and the range for each moment of the day.",
           [3, 30, 14, 14, 4, 64], rows=G1 + 3)
for r, label, value, fmt in ((4, "Name on the report", "Jamie", None), (5, "Units", "mg/dL", None),
                             (6, "Report from", "=TODAY()-29", DATE), (7, "Report to", "=TODAY()", DATE)):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
dropdown(se, '"mg/dL,mmol/L"', "C5", strict=False)
header(se, G0 - 1, 2, ["Moment", "Lowest", "Highest"])
for k, m in enumerate(MOMENTS):
    r = G0 + k
    style(se.cell(r, 2, m), False, bold=True)
    style(se.cell(r, 3, targets[m][0]), True, "General", align="center")
    style(se.cell(r, 4, targets[m][1]), True, "General", align="center")
for k, text in enumerate(["Type the range your doctor or nurse gave you for each moment.",
                          "The example numbers are made up. This file only keeps a record: it gives no advice.",
                          "Report from and to: the dates the Report tab covers. As typed, the last 30 days.",
                          "Units: only a label. Type every reading in the same units."]):
    se.cell(4 + k, 6, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Readings
rd = wb.create_sheet("Readings")
sheet_base(rd, "Readings", "One line per reading. Each is marked against the range for its moment on the Settings tab.",
           [3, 12, 8, 17, 10, 34, 12], rows=R1 + 2)
header(rd, 5, 2, ["Date", "Time", "Moment", "Reading", "Notes", "Range"])
for r in range(R0, R1 + 1):
    x = readings[r - R0] if r - R0 < len(readings) else (None,) * 5
    style(rd.cell(r, 2, x[0]), True, DATE)
    style(rd.cell(r, 3, x[1]), True, "hh:mm", align="center")
    style(rd.cell(r, 4, x[2]), True)
    style(rd.cell(r, 5, x[3]), True, "General", True, "center")
    style(rd.cell(r, 6, x[4]), True)
    at = lambda col: f'INDEX(Settings!${col}${G0}:${col}${G1},M{r})'
    rd[f"M{r}"] = f'=IF(D{r}="","",IFERROR(MATCH(D{r},{MOMENT_RANGE},0),""))'                  # M: which moment
    rd[f"J{r}"] = f'=IF(M{r}="","",IF({at("C")}="","",{at("C")}))'                            # J: lowest
    rd[f"K{r}"] = f'=IF(M{r}="","",IF({at("D")}="","",{at("D")}))'                            # K: highest
    style(rd.cell(r, 7, f'=IF(OR(B{r}="",E{r}="",J{r}="",K{r}=""),"",IF(E{r}<J{r},"Below",IF(E{r}>K{r},"Above",'
                        f'"In range")))'), False, align="center")
    rd[f"L{r}"] = f'=IF(OR(B{r}="",E{r}=""),"",B{r}+MOD(N(C{r}),1)+ROW()/10000000)'             # L: in time order
    rd[f"N{r}"] = f'=IF(OR(B{r}="",E{r}="",M{r}=""),"",IF(AND(B{r}>={FROM},B{r}<={TO}),M{r}*10000+E{r},""))'  # N: report
    rd[f"O{r}"] = f'=IF(OR(B{r}="",E{r}=""),"",B{r}*10000+E{r})'                                 # O: by day
    rd[f"P{r}"] = f'=IF(N{r}="","",E{r})'                                                        # P: in the report
for c in "JKLMNOP":
    rd.column_dimensions[c].hidden = True
dropdown(rd, f"={MOMENT_RANGE}", f"D{R0}:D{R1}")
for text, colour in (("Below", BAD), ("Above", WARN), ("In range", OK)):
    rd.conditional_formatting.add(f"G{R0}:G{R1}", FormulaRule(formula=[f'G{R0}="{text}"'], fill=fill(colour)))
rd.freeze_panes = "C6"

# ---------------------------------------------------------------- Medicines
md = wb.create_sheet("Medicines")
sheet_base(md, "Medicines", "What you take and when. Up to four times a day. Tick them off on the Doses tab.",
           [3, 18, 13, 8, 8, 8, 8, 10, 12, 12, 24, 12, 9, 9, 9], rows=M1 + 2)
header(md, 5, 2, ["Medicine", "Dose", "Time 1", "Time 2", "Time 3", "Time 4", "With food", "Started", "Stopped", "Notes",
                  "Status", "A day", "Today", "Left"])
for r in range(M0, M1 + 1):
    x = medicines[r - M0] if r - M0 < len(medicines) else (None,) * 10
    name, dose, t1, t2, t3, t4, food, started, stopped, note = x
    style(md.cell(r, 2, name), True, bold=True)
    style(md.cell(r, 3, dose), True)
    for c, t in zip((4, 5, 6, 7), (t1, t2, t3, t4)):
        style(md.cell(r, c, t), True, "hh:mm", align="center")
    style(md.cell(r, 8, food), True, align="center")
    style(md.cell(r, 9, started), True, DATE, align="center")
    style(md.cell(r, 10, stopped), True, DATE, align="center")
    style(md.cell(r, 11, note), True)
    style(md.cell(r, 12, f'=IF(B{r}="","",IF(AND(J{r}<>"",J{r}<=TODAY()),"Stopped",IF(AND(I{r}<>"",I{r}>TODAY()),'
                         f'"Not started","Taking")))'), False, align="center")
    style(md.cell(r, 13, f'=IF(L{r}<>"Taking","",COUNT(D{r}:G{r}))'), False, "0", align="center")
    style(md.cell(r, 14, f'=IF(B{r}="","",COUNTIFS({DS("B")},TODAY(),{DS("D")},B{r}))'), False, "0", align="center")
    style(md.cell(r, 15, f'=IF(L{r}<>"Taking","",MAX(0,M{r}-N{r}))'), False, "0;;", align="center")
    # P to S: each time of today in order; T to W: which dose of the day it is.
    for s, col in enumerate("DEFG"):
        md.cell(r, 16 + s, f'=IF(OR(L{r}<>"Taking",{col}{r}=""),"",MOD({col}{r},1)+({r * 10 + s + 1})/10000000)')
        rank = "+".join(f'({c}{r}<>"")*(N({c}{r})<=N({col}{r}))' for c in "DEFG")
        md.cell(r, 20 + s, f'=IF({col}{r}="","",{rank})')
    md[f"X{r}"] = f'=IF(L{r}="Taking",ROW(),"")'                                              # X: on the report
for c in "PQRSTUVWX":
    md.column_dimensions[c].hidden = True
dropdown(md, '"Yes,No"', f"H{M0}:H{M1}", strict=False)
for text, colour in (("Taking", OK), ("Stopped", INFO), ("Not started", WARN)):
    md.conditional_formatting.add(f"L{M0}:L{M1}", FormulaRule(formula=[f'L{M0}="{text}"'], fill=fill(colour)))
md.freeze_panes = "C6"

# ---------------------------------------------------------------- Doses
ds = wb.create_sheet("Doses")
sheet_base(ds, "Doses", "One line each time you take a medicine. The Dashboard shows what is still to take today.",
           [3, 12, 8, 18, 34], rows=D1 + 2)
header(ds, 5, 2, ["Date", "Time", "Medicine", "Notes"])
for r in range(D0, D1 + 1):
    x = doses[r - D0] if r - D0 < len(doses) else (None,) * 4
    style(ds.cell(r, 2, x[0]), True, DATE)
    style(ds.cell(r, 3, x[1]), True, "hh:mm", align="center")
    style(ds.cell(r, 4, x[2]), True, bold=True)
    style(ds.cell(r, 5, x[3]), True)
dropdown(ds, f"={MD('B')}", f"D{D0}:D{D1}")
ds.freeze_panes = "C6"

# ---------------------------------------------------------------- Report (one printed page)
rp = wb.create_sheet("Report")
sheet_base(rp, "Blood sugar report", "", [3, 16, 9, 9, 10, 10, 10, 10, 10, 10, 10, 3], rows=62)
rp["B3"] = f'={NAME}&IF({NAME}="","",", ")&{long_date(FROM)}&" to "&{long_date(TO)}&", in "&{UNITS}'
rp["B3"].font = font(11, color=MUTED, italic=True)
header(rp, 5, 2, ["Moment", "From", "To", "Readings", "Average", "Lowest", "Highest", "In range", "Below", "Above"])
for k, m in enumerate(MOMENTS):
    r = 6 + k
    n = k + 1
    inside = f'{RD("N")},">="&{n * 10000},{RD("N")},"<"&{(n + 1) * 10000}'
    low = f'SMALL({RD("N")},COUNTIF({RD("N")},"<"&{n * 10000})+1)-{n * 10000}'
    high = f'LARGE({RD("N")},COUNTIF({RD("N")},">="&{(n + 1) * 10000})+1)-{n * 10000}'
    cells = [(2, m, None), (3, f'=IF(Settings!C{G0 + k}="","",Settings!C{G0 + k})', "General"),
             (4, f'=IF(Settings!D{G0 + k}="","",Settings!D{G0 + k})', "General"),
             (5, f"=COUNTIFS({inside})", "0;;"),
             (6, f'=IF(E{r}=0,"",ROUND(SUMIFS({RD("E")},{inside})/E{r},1))', "General"),
             (7, f'=IF(E{r}=0,"",IFERROR({low},""))', "General"),
             (8, f'=IF(E{r}=0,"",IFERROR({high},""))', "General"),
             (9, f'=IF(E{r}=0,"",COUNTIFS({inside},{RD("G")},"In range")/E{r})', "0%"),
             (10, f'=IF(E{r}=0,"",COUNTIFS({inside},{RD("G")},"Below"))', "0"),
             (11, f'=IF(E{r}=0,"",COUNTIFS({inside},{RD("G")},"Above"))', "0")]
    for c, v, fmt in cells:
        style(rp.cell(r, c, v), False, fmt, c == 2, None if c == 2 else "center")
all_row = 6 + len(MOMENTS)
for c, v, fmt in [(2, "All", None), (3, None, None), (4, None, None), (5, f"=SUM(E6:E{all_row - 1})", "0"),
                  (6, f'=IF(E{all_row}=0,"",ROUND(SUM({RD("P")})/E{all_row},1))', "General"),
                  (7, f'=IF(E{all_row}=0,"",MIN({RD("P")}))', "General"),
                  (8, f'=IF(E{all_row}=0,"",MAX({RD("P")}))', "General"),
                  (9, f'=IF(E{all_row}=0,"",COUNTIFS({RD("N")},">0",{RD("G")},"In range")/E{all_row})', "0%"),
                  (10, f'=IF(E{all_row}=0,"",SUM(J6:J{all_row - 1}))', "0"),
                  (11, f'=IF(E{all_row}=0,"",SUM(K6:K{all_row - 1}))', "0")]:
    style(rp.cell(all_row, c, v), False, fmt, True, None if c == 2 else "center")
    rp.cell(all_row, c).fill = fill("E8F0EF")
band_fills(rp, f"F6:H{all_row - 1}", "F6", "$C6", "$D6")
med_head = all_row + 2
header(rp, med_head, 2, ["Medicine", "Dose", "", "Times", "", "", "With food", "Since", "", ""])
for a, b in (("C", "D"), ("E", "G"), ("I", "K")):
    rp.merge_cells(f"{a}{med_head}:{b}{med_head}")
for i in range(6):
    r = med_head + 1 + i
    key = f"N{r}"
    rp[key] = f'=IFERROR(SMALL({MD("X")},{i + 1}),"")'
    at = lambda col: f"INDEX({MD(col)},{key}-{M0 - 1})"
    t = lambda col: f'IF({at(col)}="","",", "&{hhmm(at(col))})'
    times = f'MID({t("D")}&{t("E")}&{t("F")}&{t("G")},3,60)'
    style(rp.cell(r, 2, f'=IF({key}="",IF({i}=0,"None",""),{at("B")})'), False, bold=True)
    style(rp.cell(r, 3, f'=IF({key}="","",IF({at("C")}="","",{at("C")}))'), False)
    style(rp.cell(r, 5, f'=IF({key}="","",{times})'), False)
    style(rp.cell(r, 8, f'=IF({key}="","",IF({at("H")}="","",{at("H")}))'), False, align="center")
    style(rp.cell(r, 9, f'=IF({key}="","",IF({at("I")}="","",{at("I")}))'), False, "d mmm yyyy", align="center")
    for a, b in (("C", "D"), ("E", "G"), ("I", "K")):
        rp.merge_cells(f"{a}{r}:{b}{r}")
grid_head = med_head + 8
start = f"MAX({FROM},{TO}-30)"
header(rp, grid_head, 2, ["Date"] + MOMENTS[:8] + [""])
rp.merge_cells(f"J{grid_head}:K{grid_head}")
rp[f"J{grid_head}"] = "Night"
rp[f"I{grid_head}"] = "Bedtime"
for i in range(31):
    r = grid_head + 1 + i
    day = f"({start}+{i})"
    style(rp.cell(r, 2, f'=IF({day}>{TO},"",{day})'), False, "ddd d mmm", True)
    for k, m in enumerate(MOMENTS[:8]):
        col = 3 + k if k < 7 else 10
        style(rp.cell(r, col, f'=IF($B{r}="","",IFERROR(ROUND(AVERAGEIFS({RD("E")},{RD("B")},$B{r},{RD("D")},'
                              f'"{m}"),1),""))'), False, "General", align="center")
    style(rp.cell(r, 11), False)
    rp.merge_cells(f"J{r}:K{r}")
    rp.row_dimensions[r].height = 15
for k in range(8):
    col = rp.cell(1, 3 + k if k < 7 else 10).column_letter
    band_fills(rp, f"{col}{grid_head + 1}:{col}{grid_head + 31}", f"{col}{grid_head + 1}", f"$C${6 + k}", f"$D${6 + k}")
note_row = grid_head + 33
rp[f"B{note_row}"] = (f'=IF({TO}-{FROM}>30,"The grid shows the last 31 days. The summary covers all the dates.","")')
rp[f"B{note_row}"].font = font(9, color=MUTED, italic=True)
rp[f"B{note_row + 1}"] = "Red: below the range typed for that moment. Amber: above it. A record only, with no advice."
rp[f"B{note_row + 1}"].font = font(9, color=MUTED, italic=True)
rp.column_dimensions["N"].hidden = True
rp.page_setup.orientation = "portrait"
rp.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Blood sugar and medicines", "Readings on the left, medicines on the right.",
           [3, 17, 14, 14, 14, 3, 13, 18, 13, 16, 3], rows=40)
last_key = f"MAX({RD('L')})"
last_row = f"MATCH({last_key},{RD('L')},0)"
db["O4"] = f'=IF(COUNT({RD("L")})=0,"",INDEX({RD("G")},{last_row}))'                          # O4: its range
tile(db, "B", 4, "Last reading", f'=IF(COUNT({RD("L")})=0,"",INDEX({RD("E")},{last_row}))', "General")
db["B6"] = f'=IF(COUNT({RD("L")})=0,"",{short(f"INT({last_key})")}&", "&{hhmm(last_key)}&IF(O4="",""," ("&LOWER(O4)&")"))'
db["B6"].font = font(9, color=MUTED, italic=True)
db.conditional_formatting.add("B5", FormulaRule(formula=['$O$4="Below"'], font=Font(name=F, size=16, bold=True, color=RED)))
db.conditional_formatting.add("B5", FormulaRule(formula=['$O$4="Above"'], font=Font(name=F, size=16, bold=True, color=AMBER)))
window = lambda days: f'{RD("B")},">="&(TODAY()-{days - 1}),{RD("B")},"<="&TODAY()'
tile(db, "C", 4, "Today's average", f'=IFERROR(ROUND(AVERAGEIFS({RD("E")},{window(1)}),1),"")', "General")
tile(db, "D", 4, "7-day average", f'=IFERROR(ROUND(AVERAGEIFS({RD("E")},{window(7)}),1),"")', "General")
tile(db, "E", 4, "30-day average", f'=IFERROR(ROUND(AVERAGEIFS({RD("E")},{window(30)}),1),"")', "General")
rated7 = "+".join(f'COUNTIFS({RD("G")},"{s}",{window(7)})' for s in ("In range", "Below", "Above"))
tile(db, "G", 4, "In range, 7 days", f'=IFERROR(COUNTIFS({RD("G")},"In range",{window(7)})/({rated7}),"")', "0%")
tile(db, "H", 4, "Below, 7 days", f'=COUNTIFS({RD("G")},"Below",{window(7)})', "0", RED)
tile(db, "I", 4, "Above, 7 days", f'=COUNTIFS({RD("G")},"Above",{window(7)})', "0", AMBER)
tile(db, "J", 4, "Doses left today", f"=SUM({MD('O')})", "0")

header(db, 8, 2, ["Last 7 days", "Readings", "Average", "In range"])
for k, m in enumerate(MOMENTS):
    r = 9 + k
    style(db.cell(r, 2, m), False, bold=True)
    of = f'{RD("D")},B{r},{window(7)}'
    style(db.cell(r, 3, f"=COUNTIFS({of},{RD('E')},\">0\")"), False, "0;;", align="center")
    style(db.cell(r, 4, f'=IFERROR(ROUND(AVERAGEIFS({RD("E")},{of}),1),"")'), False, "General", align="center")
    rated = "+".join(f'COUNTIFS({of},{RD("G")},"{s}")' for s in ("In range", "Below", "Above"))
    style(db.cell(r, 5, f'=IFERROR(COUNTIFS({of},{RD("G")},"In range")/({rated}),"")'), False, "0%", align="center")

header(db, 8, 7, ["Today", "Medicine", "Dose", "Status"])
slots = f"Medicines!$P${M0}:$S${M1}"
for i in range(10):
    r = 9 + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({slots},{i + 1}),"")'
    which = lambda col: f"MATCH({key},Medicines!${col}${M0}:${col}${M1},0)"
    row = (f'IFERROR({which("P")},IFERROR({which("Q")},IFERROR({which("R")},{which("S")})))')
    slot = (f'IF(ISNUMBER({which("P")}),1,IF(ISNUMBER({which("Q")}),2,IF(ISNUMBER({which("R")}),3,4)))')
    db[f"Q{r}"] = f'=IF({key}="","",{row})'                                                   # Q: the medicine's line
    db[f"R{r}"] = f'=IF({key}="","",{slot})'                                                  # R: which of its times
    at = lambda col: f"INDEX({MD(col)},Q{r})"
    rank = f"INDEX(Medicines!$T${M0}:$W${M1},Q{r},R{r})"
    none = '"Nothing to take today"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},ROUND({key}*1440,0)/1440)'), False, "hh:mm", True, "center")
    style(db.cell(r, 8, f'=IF({key}="","",{at("B")})'), False)
    style(db.cell(r, 9, f'=IF({key}="","",IF({at("C")}="","",{at("C")}))'), False)
    style(db.cell(r, 10, f'=IF({key}="","",IF({rank}<={at("N")},"Taken",IF(G{r}<=MOD(NOW(),1),"Not taken","Later")))'),
          False, align="center")
for text, colour in (("Taken", OK), ("Not taken", WARN), ("Later", INFO)):
    db.conditional_formatting.add("J9:J18", FormulaRule(formula=[f'J9="{text}"'], fill=fill(colour)))

header(db, 20, 2, ["Last 14 days", "Average", "Lowest", "Highest"])
for i in range(14):
    r = 21 + i
    day = f"(TODAY()-{13 - i})"
    style(db.cell(r, 2, f"={day}"), False, "ddd d mmm", True)
    low = f'SMALL({RD("O")},COUNTIF({RD("O")},"<"&B{r}*10000)+1)-B{r}*10000'
    high = f'LARGE({RD("O")},COUNTIF({RD("O")},">="&(B{r}+1)*10000)+1)-B{r}*10000'
    style(db.cell(r, 3, f'=IFERROR(ROUND(AVERAGEIFS({RD("E")},{RD("B")},B{r}),1),"")'), False, "General", align="center")
    style(db.cell(r, 4, f'=IF(C{r}="","",IFERROR({low},""))'), False, "General", align="center")
    style(db.cell(r, 5, f'=IF(C{r}="","",IFERROR({high},""))'), False, "General", align="center")
    # G to J, under the chart: the same, with gaps a chart can skip (white text, the chart covers them).
    db[f"G{r}"] = f'={short(f"B{r}")}'
    for c, src_col in zip("HIJ", "CDE"):
        db[f"{c}{r}"] = f'=IF({src_col}{r}="",NA(),{src_col}{r})'
    for c in "GHIJ":
        db[f"{c}{r}"].font = Font(name=F, size=8, color="FFFFFF")
chart = LineChart()
chart.title = "Last 14 days"
chart.height, chart.width = 7.2, 12.6
for c, name, colour in zip("HIJ", ("Average", "Lowest", "Highest"), (TEAL, RED, "F2A65A")):
    chart.add_data(Reference(db, min_col=ord(c) - 64, min_row=21, max_row=34), titles_from_data=False)
    s = chart.series[-1]
    s.tx = SeriesLabel(v=name)
    s.smooth = False
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = 22000
    s.marker.symbol = "circle"
    s.marker.size = 5
    s.marker.graphicalProperties.solidFill = colour
    s.marker.graphicalProperties.line.solidFill = colour
chart.set_categories(Reference(db, min_col=7, min_row=21, max_row=34))
chart.y_axis.number_format = "General"
db.add_chart(chart, "G20")
for c in "OPQR":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: your name, the units, and the range your doctor or nurse gave you for each moment of the day."),
    ("2", "Readings: one line per reading with the date, time and moment. Notes are for food, exercise, how you felt."),
    ("3", "Medicines: what you take, the dose and up to four times a day."),
    ("4", "Doses: one line each time you take one. The Dashboard shows what is still to take today."),
    ("5", "Report: set the dates on Settings, then print the Report tab or save it as a PDF for an appointment."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["This file keeps a record and gives no medical advice. Ask your doctor or nurse about your numbers.",
                          "The example person and every number in it are made up. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in mg/dL or mmol/L."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Readings", "Medicines", "Doses", "Report", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Blood-Sugar-Log.xlsx")
wb.save(out)
print("saved", out, len(readings), "readings", len(doses), "doses")
