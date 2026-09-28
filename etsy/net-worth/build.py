"""Net Worth and Dividend Tracker (Excel + Google Sheets).

Each account's balance once a month, assets and debts, with the net worth worked out month by month; income from
dividends, interest and rent as it comes in; each holding with what it paid over the last twelve months and its
yield; and a Dashboard with the net worth, the monthly passive income against a goal, and where the money is. Only
functions both apps have.
"""
import datetime as dt
import math
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
A0, A1 = 6, 45         # accounts on Balances
NM = 36                # months on Balances
M0 = 5                 # column of the first month (E)
H0, H1 = 6, 105        # holdings
I0, I1 = 6, 1005       # income lines
MLAST = col(M0 + NM - 1)
START, GOAL = "Settings!$C$5", "Settings!$C$6"
ASSETS = ["Cash", "Savings", "Investments", "Retirement", "Property", "Vehicle", "Other asset"]
DEBTS = ["Credit card", "Loan", "Mortgage", "Other debt"]
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
BL = f"Balances!${col(M0)}${A0}:${MLAST}${A1}"
ROW47, ROW48, ROW49 = (f"Balances!${col(M0)}${r}:${MLAST}${r}" for r in (47, 48, 49))
INC = lambda c: f"Income!${c}${I0}:${c}${I1}"
HD = lambda c: f"Holdings!${c}${H0}:${c}${H1}"
HD1 = lambda c: f"Holdings!${c}$1:${c}${H1}"


def add_months(d, n):
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    last = [31, 29 if y % 4 == 0 and (y % 100 != 0 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return dt.date(y, m, min(d.day, last[m - 1]))


def label(d):
    return f'CHOOSE(MONTH({d}),{MON_LIST})&" "&YEAR({d})'


rng = random.Random(39)
start = add_months(today.replace(day=1), -11)

# ---------------------------------------------------------------- example data (a made-up household)
accounts = [  # name, type, first balance, last balance, noise
    ("Everyday account", "Cash", 2100, 2400, 400),
    ("Emergency savings", "Savings", 8000, 11300, 0),
    ("Brokerage", "Investments", 24000, 30800, 900),
    ("Dividend portfolio", "Investments", 18200, 22600, 600),
    ("Retirement plan", "Retirement", 52000, 61400, 1200),
    ("Home", "Property", 320000, 330000, 0),
    ("Car", "Vehicle", 14000, 12400, 0),
    ("Credit card", "Credit card", 1200, 850, 350),
    ("Car loan", "Loan", 9000, 6600, 0),
    ("Mortgage", "Mortgage", 210000, 204800, 0),
]
balances = []
for name, kind, a, b, noise in accounts:
    row = []
    for k in range(12):
        v = a + (b - a) * k / 11 + (rng.uniform(-noise, noise) if 0 < k < 11 else 0)
        if kind == "Property":
            v = a if k < 6 else b                                  # a new valuation half way
        row.append(round(v, -1) if v > 1000 else round(v, 2))
    balances.append(row)
holdings = [  # name, account, value now
    ("Total market fund", "Brokerage", 16400), ("Dividend fund", "Dividend portfolio", 9800),
    ("Utility shares", "Dividend portfolio", 5200), ("Bank shares", "Dividend portfolio", 4100),
    ("Property fund", "Brokerage", 6300), ("Bond fund", "Brokerage", 8100),
    ("High yield savings", "Emergency savings", 11300), ("Spare room", None, None),
]
income = []
for k in range(12):
    m = add_months(start, k)
    q = k % 3 == 2
    income += [(m.replace(day=5), "Property fund", "Dividend", 42), (m.replace(day=28), "Bond fund", "Dividend", 55),
               (m.replace(day=1), "Spare room", "Rent", 300), (add_months(m, 1) - dt.timedelta(days=1),
                                                               "High yield savings", "Interest", round(33 + k * 0.4, 2))]
    if q:
        income += [(m.replace(day=20), "Dividend fund", "Dividend", round(170 + k * 2.5, 2)),
                   (m.replace(day=15), "Utility shares", "Dividend", 95),
                   (m.replace(day=24), "Total market fund", "Dividend", round(200 + k * 3, 2))]
    if k in (4, 10):
        income.append((m.replace(day=10), "Bank shares", "Dividend", 140))
income = [x for x in income if x[0] <= today]
income.sort(key=lambda x: x[0])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label_, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label_
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The currency, the first month on Balances, and your passive income goal.",
           [3, 30, 18, 4, 60], rows=12)
for r, lab, value, fmt, note in (
        (4, "Currency", "USD", None, "Amounts have no currency sign, so any currency works."),
        (5, "First month on Balances", dt.datetime.combine(start, dt.time()), "mmm yyyy",
         "Balances has 36 months from this one."),
        (6, "Passive income goal a month", 1000, MONEY, "Dividends, interest and rent together.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Balances
bl = wb.create_sheet("Balances")
sheet_base(bl, "Balances", "Each account once, with its balance at the end of each month. Debts as the amount "
           "owed.", [3, 22, 13, 8] + [11] * NM, rows=51)
header(bl, 5, 2, ["Account", "Type", "Kind"])
for k in range(NM):
    c = bl.cell(5, M0 + k, f"=EDATE({START},{k})")
    c.number_format = "mmm yy"
    c.font = font(10, True, "FFFFFF")
    c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
bl.row_dimensions[5].height = 30
for r in range(A0, A1 + 1):
    i = r - A0
    x = accounts[i] if i < len(accounts) else (None, None, None, None, None)
    style(bl.cell(r, 2, x[0]), True, bold=True)
    style(bl.cell(r, 3, x[1]), True, align="center")
    debt = "OR(" + ",".join(f'C{r}="{d}"' for d in DEBTS) + ")"
    style(bl.cell(r, 4, f'=IF(B{r}="","",IF({debt},"Debt","Asset"))'), False, align="center").font = font(10, color=MUTED)
    bl[f"AP{r}"] = f'=IF(B{r}="",0,IF({debt},-1,1))'                                        # AP: sign
    for k in range(NM):
        v = balances[i][k] if i < len(balances) and k < 12 else None
        style(bl.cell(r, M0 + k, v), True, MONEY)
bl.column_dimensions["AP"].hidden = True
dropdown(bl, '"' + ",".join(ASSETS + DEBTS) + '"', f"C{A0}:C{A1}")
for r, lab in ((47, "Assets"), (48, "Debts"), (49, "Net worth")):
    bl.cell(r, 2, lab).font = font(11, True, TEAL_D if r == 49 else MUTED)
    for k in range(NM):
        L = col(M0 + k)
        rng_ = f"{L}{A0}:{L}{A1}"
        if r == 47:
            f_ = f'=IF(COUNT({rng_})=0,"",SUMIFS({rng_},$AP${A0}:$AP${A1},1))'
        elif r == 48:
            f_ = f'=IF(COUNT({rng_})=0,"",SUMIFS({rng_},$AP${A0}:$AP${A1},-1))'
        else:
            f_ = f'=IF(COUNT({rng_})=0,"",{L}47-{L}48)'
        c = style(bl.cell(r, M0 + k, f_), False, MONEY, r == 49)
        bl.cell(50, M0 + k, f"=IF(COUNT({rng_})>0,{k + 1},0)")                               # row 50: month in use
bl.row_dimensions[50].hidden = True
bl.conditional_formatting.add(f"{col(M0)}49:{MLAST}49", FormulaRule(
    formula=[f"AND(ISNUMBER({col(M0)}49),{col(M0)}49<0)"], fill=fill(BAD)))
bl.freeze_panes = f"{col(M0)}6"

# ---------------------------------------------------------------- Holdings
hd = wb.create_sheet("Holdings")
sheet_base(hd, "Holdings", "What pays you: funds, shares, savings, a rented room. Its value now, for the yield.",
           [3, 24, 22, 13, 13, 13, 9, 12], rows=H1 + 2)
header(hd, 5, 2, ["Holding", "Account", "Value now", "Last 12 months", "This year", "Yield", "A month"])
since = f"DATE(YEAR(TODAY()),MONTH(TODAY())-11,1)"
for r in range(H0, H1 + 1):
    x = holdings[r - H0] if r - H0 < len(holdings) else (None, None, None)
    style(hd.cell(r, 2, x[0]), True, bold=True)
    style(hd.cell(r, 3, x[1]), True)
    style(hd.cell(r, 4, x[2]), True, MONEY)
    style(hd.cell(r, 5, f'=IF(B{r}="","",SUMIFS({INC("E")},{INC("C")},B{r},{INC("B")},">="&{since}))'), False,
          MONEY, True)
    style(hd.cell(r, 6, f'=IF(B{r}="","",SUMIFS({INC("E")},{INC("C")},B{r},{INC("B")},">="&DATE(YEAR(TODAY()),1,1)))'),
          False, MONEY)
    style(hd.cell(r, 7, f'=IF(OR(B{r}="",N(D{r})<=0),"",E{r}/D{r})'), False, "0.0%", align="center")
    style(hd.cell(r, 8, f'=IF(B{r}="","",E{r}/12)'), False, MONEY)
    hd[f"J{r}"] = f'=IF(OR(B{r}="",N(E{r})<=0),"",ROUND(E{r}*100,0)*1000+1000-ROW())'      # J: top payers
hd.column_dimensions["J"].hidden = True

# ---------------------------------------------------------------- Income
ic = wb.create_sheet("Income")
sheet_base(ic, "Income", "Each dividend, interest payment or rent as it comes in.", [3, 13, 24, 12, 12, 28],
           rows=I1 + 2)
header(ic, 5, 2, ["Date", "From", "Type", "Amount", "Note"])
for r in range(I0, I1 + 1):
    x = income[r - I0] if r - I0 < len(income) else (None, None, None, None)
    for c, v, fmt in zip(range(2, 6), x, (DATE, None, None, MONEY)):
        style(ic.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4) else None)
    style(ic.cell(r, 6), True)
    ic[f"G{r}"] = f'=IF(NOT(ISNUMBER(B{r})),"",YEAR(B{r})*100+MONTH(B{r}))'                 # G: month
ic.column_dimensions["G"].hidden = True
dropdown(ic, f"={HD('B')}", f"C{I0}:C{I1}", strict=False)
dropdown(ic, '"Dividend,Interest,Rent,Other"', f"D{I0}:D{I1}")
ic.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Net worth and passive income", "", [3, 14, 14, 14, 14, 3, 14, 14, 14, 14, 3], rows=60)
db["M1"] = f"=MAX(Balances!{col(M0)}50:{MLAST}50)"                                          # M1: latest month
db["M2"] = f'=IF(M1=0,"",EDATE({START},M1-1))'                                                # M2: its date
db["B3"] = f'=IF(M1=0,"Type the balances on Balances.","Balances at the end of "&{label("M2")}&"; income over the last 12 months.")'
db["B3"].font = font(11, color=MUTED, italic=True)
at = lambda row, k: f'IF(OR({k}<1,{k}>{NM}),"",INDEX({row},{k}))'
tile(db, "B", 4, "Net worth", f'=IF($M$1=0,"",{at(ROW49, "$M$1")})', MONEY)
tile(db, "C", 4, "On last month",
     f'=IF($M$1<2,"",IF(OR({at(ROW49, "$M$1-1")}=""),"",{at(ROW49, "$M$1")}-{at(ROW49, "$M$1-1")}))', MONEY)
tile(db, "D", 4, "Assets", f'=IF($M$1=0,"",{at(ROW47, "$M$1")})', MONEY)
tile(db, "E", 4, "Debts", f'=IF($M$1=0,"",{at(ROW48, "$M$1")})', MONEY, "B23B2E")
tile(db, "G", 4, "Income, 12 months", "=SUM(J9:J20)", MONEY)
tile(db, "H", 4, "A month on average", "=G5/12", MONEY)
tile(db, "I", 4, "Goal a month", f"=N({GOAL})", MONEY)
tile(db, "J", 4, "Of the goal", f'=IF(N({GOAL})<=0,"",H5/{GOAL})', "0%", "B86E1C")

header(db, 8, 2, ["Month", "Assets", "Debts", "Net worth"])
header(db, 8, 7, ["Month", "Dividends", "Other income", "Total"])
for j in range(12):
    r = 9 + j
    k = f"($M$1-{11 - j})"
    db[f"N{r}"] = f'=IF(OR($M$1=0,{k}<1),"",{k})'                                             # N: month on Balances
    style(db.cell(r, 2, f'=IF(N{r}="","",{label(f"EDATE({START},N{r}-1)")})'), False, bold=True)
    for c, row in ((3, ROW47), (4, ROW48), (5, ROW49)):
        style(db.cell(r, c, f'=IF(N{r}="","",{at(row, f"N{r}")})'), False, MONEY, c == 5)
    db[f"O{r}"] = f'=IF(ISNUMBER(E{r}),E{r},NA())'                                              # O: for the chart
    m = f"DATE(YEAR(TODAY()),MONTH(TODAY())-{11 - j},1)"
    db[f"P{r}"] = f"=YEAR({m})*100+MONTH({m})"                                                  # P: income month
    style(db.cell(r, 7, f"={label(m)}"), False, bold=True)
    style(db.cell(r, 8, f'=SUMIFS({INC("E")},{INC("D")},"Dividend",{INC("G")},P{r})'), False, MONEY)
    style(db.cell(r, 10, f'=SUMIFS({INC("E")},{INC("G")},P{r})'), False, MONEY, True)
    style(db.cell(r, 9, f"=J{r}-H{r}"), False, MONEY)
db.conditional_formatting.add("E9:E20", FormulaRule(formula=["AND(ISNUMBER(E9),E9<0)"], fill=fill(BAD)))
# The net worth chart reads visible cells behind it (white letters).
for j in range(12):
    r = 24 + j
    db.cell(r, 3, f"=B{9 + j}").font = Font(color="FFFFFF", size=8)
    db.cell(r, 4, f"=O{9 + j}").font = Font(color="FFFFFF", size=8)
db["D23"] = "Net worth"
db["D23"].font = Font(color="FFFFFF", size=8)
line = LineChart()
line.title = "Net worth"
line.height, line.width = 7.2, 12.5
line.add_data(Reference(db, min_col=4, min_row=23, max_row=35), titles_from_data=True)
line.set_categories(Reference(db, min_col=3, min_row=24, max_row=35))
line.series[0].graphicalProperties.line.solidFill = TEAL
line.series[0].graphicalProperties.line.width = 28000
line.series[0].smooth = False
line.legend = None
line.y_axis.number_format = "#,##0"
db.add_chart(line, "B22")
bars = BarChart()
bars.type = "col"
bars.grouping = "stacked"
bars.overlap = 100
bars.title = "Passive income by month"
bars.height, bars.width = 7.2, 12.5
bars.add_data(Reference(db, min_col=8, max_col=9, min_row=8, max_row=20), titles_from_data=True)
bars.set_categories(Reference(db, min_col=7, min_row=9, max_row=20))
for s, colour in zip(bars.series, (TEAL, "F2A65A")):
    s.graphicalProperties.solidFill = colour
bars.y_axis.number_format = "#,##0"
db.add_chart(bars, "G22")

header(db, 38, 2, ["Where it is", "Amount", "Share", ""])
latest = f"INDEX({BL},0,$M$1)"
for k, kind in enumerate(ASSETS):
    r = 39 + k
    style(db.cell(r, 2, kind), False, bold=True)
    style(db.cell(r, 3, f'=IF($M$1=0,"",SUMIFS({latest},Balances!$C${A0}:$C${A1},B{r}))'), False, MONEY)
    style(db.cell(r, 4, f'=IF(OR(C{r}="",N($D$5)<=0),"",C{r}/$D$5)'), False, "0%", align="center")
    style(db.cell(r, 5, f'=IF(D{r}="","",REPT("█",ROUND(D{r}*12,0)))'), False).font = Font(size=9, color=TEAL)
db.cell(38, 5).value = None
header(db, 38, 7, ["Top payers", "12 months", "Yield", "A month"])
for n in range(7):
    r = 39 + n
    db[f"Q{r}"] = f"=IFERROR(LARGE({HD('J')},{n + 1}),\"\")"                                   # Q: holding key
    h = lambda c: f"INDEX({HD1(c)},1000-MOD(Q{r},1000))"
    style(db.cell(r, 7, f'=IF(Q{r}="","",{h("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(Q{r}="","",{h("E")})'), False, MONEY)
    style(db.cell(r, 9, f'=IF(Q{r}="","",IF({h("G")}="","",{h("G")}))'), False, "0.0%", align="center")
    style(db.cell(r, 10, f'=IF(Q{r}="","",{h("H")})'), False, MONEY)
for c in "MNOPQ":
    db.column_dimensions[c].hidden = True
for r in list(range(9, 21)) + list(range(39, 46)):
    db.row_dimensions[r].height = 18
    for c in range(2, 11):
        cell = db.cell(r, c)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="center")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: the currency, the first month for Balances, and your passive income goal a month."),
    ("2", "Balances: each account once with its type, then its balance at the end of every month."),
    ("3", "Holdings: the funds, shares and accounts that pay you, with their value now."),
    ("4", "Income: each dividend, interest payment or rent as it comes in."),
    ("5", "Dashboard: your net worth month by month, and your passive income against the goal."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example accounts, holdings and amounts are made up. Delete them and add your own.",
                          "Debts go in as the amount owed. Net worth is everything you own less everything you owe.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Balances", "Holdings", "Income", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Net-Worth-Dividend-Tracker.xlsx")
wb.save(out)
print("saved", out, len(accounts), "accounts", len(income), "income lines")
