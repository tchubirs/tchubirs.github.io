"""Etsy listing photos for the Net Worth and Dividend Tracker (run build.py first)."""
import os
import re
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Net-Worth-Dividend-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in, and Balances shows only the months in use.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
bl = wb["Balances"]
used = max(c for c in range(5, 41) if any(bl.cell(r, c).value not in (None, "") for r in range(6, 46)))
rows = {"Holdings": range(last(wb["Holdings"], 105) + 1, 106), "Balances": range(last(bl, 45) + 1, 46),
        "Income": list(range(6, last(wb["Income"], 1005) - 17)) + list(range(last(wb["Income"], 1005) + 1, 1006))}
r = Renders(xlsx, tmp, one_page=("Dashboard",), hide_rows=rows,
            hide={"Balances": [get_column_letter(c) for c in range(used + 1, 41)]})


def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


def ln_level(sp):
    """Level text only: the chart labels are turned."""
    x0, y0, x1, y1 = sp["bbox"]
    return x1 - x0 > 1.5 * (y1 - y0)


def stack(top, lower):
    lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                         Image.LANCZOS)
    both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
    both.paste(top, (0, 0))
    both.paste(lower, (0, top.height + 40))
    return both


img = canvas(dark=True)
title(img, "Net worth and dividend tracker", "Your net worth month by month, your passive income against a goal, "
      "and what each holding pays", dark=True, width=2300)
laptop(img, r.shot("Net worth and passive income"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-balances.jpg"), "Every account, once a month",
       "Assets and debts at the end of each month. The net worth adds up by itself.", r.shot("Balances"),
       [("#F9C9C4", "Net worth below zero")])

detail(os.path.join(img_dir, "03-holdings.jpg"), "What each holding pays",
       "Dividends, interest and rent as they come in, with the last 12 months and the yield of each holding",
       stack(r.shot("Holdings"), r.shot("Income")))

# The Dashboard from the charts down: cut under the last month of the two tables.
page = r.doc[r._index("Net worth and passive income")]
words = page.get_text("words")
where = next(w for w in words if w[4] == "Where")
spans = [sp for b in page.get_text("dict")["blocks"] for ln in b.get("lines", []) for sp in ln["spans"]]
years = [sp["bbox"][3] for sp in spans if re.search(r"20\d\d$", sp["text"].strip()) and sp["bbox"][1] < where[1]
         and sp["color"] != 0xFFFFFF and ln_level(sp)]                            # the tables, not the charts
full = r.page("Net worth and passive income")
cut = int((max(years) + 6) * 220 / 72)
detail(os.path.join(img_dir, "04-where-it-is.jpg"), "Where your money is",
       "Net worth and passive income by month, your assets by type, and the holdings that pay you the most",
       content(full.crop((0, cut, full.width, full.height))))

img = canvas()
y = title(img, "Inside the file", "Four tabs, settings and a Start Here page. You type in the yellow cells and the rest "
          "is calculated.")
grid(img, [
    (content(r.page("Net worth and passive income")), "Dashboard", "Net worth and income"),
    (content(r.page("Balances")), "Balances", "Each account, each month"),
    (content(r.page("Holdings")), "Holdings", "What each one pays"),
    (content(r.page("Income")), "Income", "Dividends, interest, rent"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "05-whats-inside.jpg"))
