"""Car Maintenance Tracker (Excel + Google Sheets).

Up to five cars. Services by distance or by date, whichever comes first, with the day each one is likely due from
how far the car goes in a day; fill-ups with the consumption from one full tank to the next; every other cost;
insurance, tax and inspection with their end dates; and a Dashboard with each car's odometer, consumption and
cost per distance, what is due, and the year month by month. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
K0, K1 = 6, 10         # cars
V0, V1 = 6, 65         # services
U0, U1 = 6, 605        # fuel
C0, C1 = 6, 505        # costs
D0, D1 = 6, 45         # documents
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MON_LIST = ",".join(f'"{m}"' for m in MONTHS)
CA = lambda col: f"Cars!${col}${K0}:${col}${K1}"
SV = lambda col: f"Services!${col}${V0}:${col}${V1}"
FU = lambda col: f"Fuel!${col}${U0}:${col}${U1}"
CO = lambda col: f"Costs!${col}${C0}:${col}${C1}"
DO = lambda col: f"Documents!${col}${D0}:${col}${D1}"
DIST, VOL, SOON_D, SOON_T = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6", "Settings!$C$7"
CATS = ["Service", "Repair", "Tyres", "Insurance", "Tax", "Parking", "Tolls", "Cleaning", "Parts", "Other"]
DOCS = ["Insurance", "Road tax", "Inspection", "Registration", "Breakdown cover", "Warranty", "Other"]


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def ago(days):
    return today - dt.timedelta(days=days)


# ---------------------------------------------------------------- example data (made-up cars and garages)
cars = [
    # name, make and model, year, plate, fuel, odometer at start, start date
    ("Blue hatchback", "Hatchback, 1.2 petrol", 2019, "AB19 CDE", "Petrol", 41200, ago(300)),
    ("Grey van", "Small van, 1.6 diesel", 2016, "FG16 HJK", "Diesel", 118500, ago(300)),
]
rng = random.Random(28)
fuel = []
for name, per_day, per100, tank, price, partial_at in (("Blue hatchback", 38, 6.2, 40, 1.72, 5), ("Grey van", 55, 7.9, 55, 1.68, 3)):
    odo = next(c[5] for c in cars if c[0] == name)
    d = ago(300)
    k = 0
    while True:
        gap = rng.randint(9, 15)
        d += dt.timedelta(days=gap)
        if d > today:
            break
        dist = round(gap * per_day * rng.uniform(0.8, 1.2))
        odo += dist
        full = k != partial_at
        litres = round(dist * per100 * rng.uniform(0.95, 1.06) / 100 * (1 if full else 0.6), 2)
        if full and k == partial_at + 1:
            litres = round(litres + fuel[-1][3] * 0.67, 2)      # the tank topped up after a half fill
        paid = round(litres * price * rng.uniform(0.97, 1.03), 2)
        fuel.append((d, name, odo, litres, paid, "x" if full else None))
        k += 1
fuel.sort()
last_odo = {name: max(f[2] for f in fuel if f[1] == name) for name, *_ in cars}
services = [
    # car, service, every (distance), every (months), last done (odometer), last done (date)
    ("Blue hatchback", "Oil and filter", 15000, 12, last_odo["Blue hatchback"] - 13900, ago(250)),
    ("Blue hatchback", "Tyre rotation", 10000, 12, last_odo["Blue hatchback"] - 6200, ago(170)),
    ("Blue hatchback", "Brakes check", 20000, 24, last_odo["Blue hatchback"] - 8000, ago(400)),
    ("Blue hatchback", "Air filter", 30000, 24, last_odo["Blue hatchback"] - 21000, ago(700)),
    ("Blue hatchback", "Brake fluid", None, 24, None, ago(705)),
    ("Grey van", "Oil and filter", 15000, 12, last_odo["Grey van"] - 15400, ago(340)),
    ("Grey van", "Timing belt", 100000, 60, 100200, ago(900)),
    ("Grey van", "Tyres checked", 10000, 6, last_odo["Grey van"] - 4000, ago(120)),
    ("Grey van", "Coolant", None, 48, None, ago(1300)),
]
costs = [
    # date, car, category, what, odometer, amount
    (ago(250), "Blue hatchback", "Service", "Oil and filter, check-up", last_odo["Blue hatchback"] - 13900, 145.00),
    (ago(170), "Blue hatchback", "Tyres", "Rotation and balance", last_odo["Blue hatchback"] - 6200, 40.00),
    (ago(330), "Blue hatchback", "Insurance", "Yearly insurance", None, 486.00),
    (ago(200), "Blue hatchback", "Tax", "Road tax", None, 180.00),
    (ago(95), "Blue hatchback", "Repair", "New wiper blades", None, 32.50),
    (ago(60), "Blue hatchback", "Cleaning", "Valet", None, 45.00),
    (ago(340), "Grey van", "Service", "Oil and filter", last_odo["Grey van"] - 15400, 168.00),
    (ago(120), "Grey van", "Tyres", "Two new front tyres", last_odo["Grey van"] - 4000, 214.00),
    (ago(280), "Grey van", "Insurance", "Yearly insurance", None, 612.00),
    (ago(150), "Grey van", "Repair", "Rear light bulb", None, 18.00),
    (ago(40), "Grey van", "Parking", "Station parking, one month", None, 95.00),
    (ago(25), "Blue hatchback", "Tolls", "Bridge tolls", None, 12.60),
]
docs = [
    # car, document, ends, reference, cost
    ("Blue hatchback", "Insurance", ago(330 - 365), "Policy 55-9021", 486.00),
    ("Blue hatchback", "Road tax", ago(200 - 365), None, 180.00),
    ("Blue hatchback", "Inspection", today + dt.timedelta(days=150), None, 55.00),
    ("Grey van", "Insurance", ago(280 - 365), "Policy 71-3310", 612.00),
    ("Grey van", "Inspection", ago(12), None, 55.00),
    ("Grey van", "Breakdown cover", today + dt.timedelta(days=21), "Member 88213", 96.00),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Units, and how early you want to hear that something is due.", [3, 30, 12, 4, 60], rows=20)
for r, label, value, fmt in ((4, "Distance", "km", None), (5, "Fuel", "L", None), (6, "Warn this distance before", 1000, "#,##0"),
                             (7, "Warn this many days before", 30, "0")):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
dropdown(se, '"km,mi"', "C4")
dropdown(se, '"L,gal"', "C5")
for k, text in enumerate(["Miles and gallons work too: the consumption then reads in miles per gallon.",
                          "Services and documents turn orange this distance or this many days before they are due."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Fuel
fu = wb.create_sheet("Fuel")
sheet_base(fu, "Fuel", "One line per fill-up. Put an x under Full when you fill the tank: the consumption is worked "
           "out from one full tank to the next.", [3, 12, 16, 11, 9, 10, 6, 10, 10, 9, 11, 11], rows=U1 + 2)
header(fu, 5, 2, ["Date", "Car", "Odometer", "Fuel", "Paid", "Full", "Price", "Distance", "Used", "Per 100", "Distance per"])
fu["E5"] = f'="Fuel ("&{VOL}&")"'
fu["J5"] = f'="Used ("&{VOL}&")"'
fu["I5"] = f'="Distance ("&{DIST}&")"'
fu["K5"] = f'={VOL}&" per 100 "&{DIST}'
fu["L5"] = f'={DIST}&" per "&{VOL}'
for r in range(U0, U1 + 1):
    x = fuel[r - U0] if r - U0 < len(fuel) else (None,) * 6
    style(fu.cell(r, 2, x[0]), True, DATE)
    style(fu.cell(r, 3, x[1]), True, bold=True)
    style(fu.cell(r, 4, x[2]), True, "#,##0", align="center")
    style(fu.cell(r, 5, x[3]), True, "0.00", align="center")
    style(fu.cell(r, 6, x[4]), True, MONEY)
    style(fu.cell(r, 7, x[5]), True, align="center")
    style(fu.cell(r, 8, f'=IF(OR(N(E{r})=0,F{r}=""),"",F{r}/E{r})'), False, "0.000", align="center")
    style(fu.cell(r, 9, f'=IF(R{r}="","",D{r}-R{r})'), False, "#,##0", align="center")
    style(fu.cell(r, 10, f'=IF(R{r}="","",SUMIFS({FU("E")},{FU("C")},C{r},{FU("D")},">"&R{r},{FU("D")},"<="&D{r}))'),
          False, "0.00", align="center")
    style(fu.cell(r, 11, f'=IF(OR(I{r}="",N(I{r})=0),"",J{r}/I{r}*100)'), False, "0.0", True, "center")
    style(fu.cell(r, 12, f'=IF(OR(J{r}="",N(J{r})=0),"",I{r}/J{r})'), False, "0.0", True, "center")
    fu.cell(r, 14, f'=IF(OR(C{r}="",N(D{r})=0),"",D{r}+ROW()/10000000)')                               # N: order
    fu.cell(r, 15, f'=IF(N{r}="","",COUNTIFS({FU("C")},C{r},{FU("N")},">"&N{r}))')                      # O: 0 = newest
    fu.cell(r, 16, f'=IF(OR(N{r}="",G{r}=""),"",COUNTIFS({FU("C")},C{r},{FU("G")},"<>",{FU("N")},"<"&N{r}))')  # P: full no.
    fu.cell(r, 18, f'=IF(OR(P{r}="",N(P{r})=0),"",SUMIFS({FU("D")},{FU("C")},C{r},{FU("P")},P{r}-1))')  # R: last full
for c in "MNOPQR":
    fu.column_dimensions[c].hidden = True       # Q stays empty
dropdown(fu, f"={CA('B')}", f"C{U0}:C{U1}")
fu.conditional_formatting.add(f"B{U0}:L{U1}", FormulaRule(formula=[f'AND($B{U0}<>"",$G{U0}="")'],
                                                          font=Font(name=F, italic=True, color=MUTED)))
fu.freeze_panes = "C6"

# ---------------------------------------------------------------- Costs
co = wb.create_sheet("Costs")
sheet_base(co, "Costs", "Everything else the cars cost: services, repairs, tyres, insurance, tax, parking. The "
           "odometer is optional.", [3, 12, 16, 13, 34, 11, 11], rows=C1 + 2)
header(co, 5, 2, ["Date", "Car", "Category", "What", "Odometer", "Amount"])
for r in range(C0, C1 + 1):
    x = costs[r - C0] if r - C0 < len(costs) else (None,) * 6
    style(co.cell(r, 2, x[0]), True, DATE)
    style(co.cell(r, 3, x[1]), True, bold=True)
    style(co.cell(r, 4, x[2]), True)
    style(co.cell(r, 5, x[3]), True)
    style(co.cell(r, 6, x[4]), True, "#,##0", align="center")
    style(co.cell(r, 7, x[5]), True, MONEY)
    co.cell(r, 9, f'=IF(OR(C{r}="",N(F{r})=0),"",F{r}+ROW()/10000000)')                                # I: order
    co.cell(r, 10, f'=IF(I{r}="","",COUNTIFS({CO("C")},C{r},{CO("I")},">"&I{r}))')                      # J: 0 = newest
for c in "HIJ":
    co.column_dimensions[c].hidden = True
dropdown(co, f"={CA('B')}", f"C{C0}:C{C1}")
dropdown(co, f'"{",".join(CATS)}"', f"D{C0}:D{C1}", strict=False)
co.freeze_panes = "C6"

# ---------------------------------------------------------------- Cars
ka = wb.create_sheet("Cars")
sheet_base(ka, "Cars", "Each car once, with the odometer on the day you start the file. The rest comes from the other "
           "tabs.", [3, 18, 24, 7, 12, 10, 12, 13, 13, 12, 12], rows=20)
header(ka, 5, 2, ["Car", "Make and model", "Year", "Plate", "Fuel", "Odometer at start", "Start date", "Odometer now",
                  "Distance a day", "Per 100"])
ka["J5"] = f'={DIST}&" a day"'
ka["K5"] = f'={VOL}&" per 100 "&{DIST}'
for r in range(K0, K1 + 1):
    x = cars[r - K0] if r - K0 < len(cars) else (None,) * 7
    for c, v, fmt in ((2, x[0], None), (3, x[1], None), (4, x[2], "0"), (5, x[3], None), (6, x[4], None), (7, x[5], "#,##0"),
                      (8, x[6], DATE)):
        style(ka.cell(r, c, v), True, fmt, c == 2, "center" if c in (4, 6, 7) else None)
    fuel_now = f'SUMIFS({FU("D")},{FU("C")},$B{r},{FU("O")},0)'
    cost_now = f'SUMIFS({CO("F")},{CO("C")},$B{r},{CO("J")},0)'
    style(ka.cell(r, 9, f'=IF(B{r}="","",MAX(N(G{r}),{fuel_now},{cost_now}))'), False, "#,##0", True, "center")
    last_day = f'SUMIFS({FU("B")},{FU("C")},$B{r},{FU("O")},0)'
    style(ka.cell(r, 10, f'=IF(OR(B{r}="",N(G{r})=0,H{r}="",{last_day}<=H{r}),"",({fuel_now}-G{r})/({last_day}-H{r}))'),
          False, "0.0", align="center")
    style(ka.cell(r, 11, f'=IF(B{r}="","",IFERROR(SUMIFS({FU("J")},{FU("C")},$B{r})/SUMIFS({FU("I")},{FU("C")},$B{r})*100,""))'),
          False, "0.0", align="center")
dropdown(ka, '"Petrol,Diesel,Electric,Hybrid,Gas,Other"', f"F{K0}:F{K1}", strict=False)
ka["B12"] = "Odometer now is the highest reading in Fuel and Costs. Distance a day is the average since the start date."
ka["B12"].font = font(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Services
sv = wb.create_sheet("Services")
sheet_base(sv, "Services", "What each car needs, every so far or every so many months, whichever comes first. When it "
           "is done, change the last odometer and date.", [3, 16, 20, 10, 10, 11, 13, 11, 13, 11, 14, 12, 26],
           rows=V1 + 2)
header(sv, 5, 2, ["Car", "Service", "Every", "Every (months)", "Last odometer", "Last date", "Next odometer",
                  "Next date", "Left", "About", "Status", "Notes"])
sv["D5"] = f'="Every ("&{DIST}&")"'
sv["J5"] = f'="Left ("&{DIST}&")"'
now = lambda r: f"INDEX({CA('I')},MATCH($B{r},{CA('B')},0))"
per_day = lambda r: f"IFERROR(N(INDEX({CA('J')},MATCH($B{r},{CA('B')},0))),0)"
for r in range(V0, V1 + 1):
    x = services[r - V0] if r - V0 < len(services) else (None,) * 6
    car, what, every_d, every_m, last_odo_, last_date = x
    style(sv.cell(r, 2, car), True, bold=True)
    style(sv.cell(r, 3, what), True)
    style(sv.cell(r, 4, every_d), True, "#,##0", align="center")
    style(sv.cell(r, 5, every_m), True, "0", align="center")
    style(sv.cell(r, 6, last_odo_), True, "#,##0", align="center")
    style(sv.cell(r, 7, last_date), True, DATE)
    style(sv.cell(r, 8, f'=IF(OR(N(D{r})=0,F{r}=""),"",F{r}+D{r})'), False, "#,##0", align="center")
    style(sv.cell(r, 9, f'=IF(OR(N(E{r})=0,G{r}=""),"",EDATE(G{r},E{r}))'), False, DATE, align="center")
    style(sv.cell(r, 10, f'=IF(OR(H{r}="",ISNA(MATCH($B{r},{CA("B")},0))),"",H{r}-{now(r)})'), False, "#,##0;[Red]-#,##0",
          align="center")
    by_km = f'IF(OR(J{r}="",{per_day(r)}<=0),"",TODAY()+INT(J{r}/{per_day(r)}))'
    style(sv.cell(r, 11, f'=IF(AND(I{r}="",J{r}=""),"",IF({by_km}="",I{r},IF(I{r}="",{by_km},MIN(I{r},{by_km}))))'), False,
          "ddd d mmm yy", align="center")
    over = f'OR(K{r}<TODAY(),AND(ISNUMBER(J{r}),J{r}<0))'
    soon = f'OR(K{r}-TODAY()<=N({SOON_T}),AND(ISNUMBER(J{r}),J{r}<=N({SOON_D})))'
    style(sv.cell(r, 12, f'=IF(K{r}="","",IF({over},"Overdue",IF({soon},"Due soon","OK")))'), False, align="center")
    style(sv.cell(r, 13), True)
    sv.cell(r, 15, f'=IF(OR(L{r}="",L{r}="OK"),"",K{r}+ROW()/10000000)')                               # O: order
sv.column_dimensions["O"].hidden = True
dropdown(sv, f"={CA('B')}", f"B{V0}:B{V1}")
for status, colour in (("Overdue", BAD), ("Due soon", WARN), ("OK", OK)):
    sv.conditional_formatting.add(f"L{V0}:L{V1}", FormulaRule(formula=[f'L{V0}="{status}"'], fill=fill(colour)))
sv[f"B{V1 + 1}"] = "About: the earlier of the next date and the day the car should reach the next odometer at its usual pace."
sv[f"B{V1 + 1}"].font = font(9, color=MUTED, italic=True)
sv.freeze_panes = "C6"

# ---------------------------------------------------------------- Documents
do = wb.create_sheet("Documents")
sheet_base(do, "Documents", "Insurance, tax, inspection and the rest, with the day each one ends.",
           [3, 16, 18, 13, 22, 11, 16, 26], rows=D1 + 2)
header(do, 5, 2, ["Car", "Document", "Ends", "Reference", "Cost", "Status", "Notes"])
for r in range(D0, D1 + 1):
    x = docs[r - D0] if r - D0 < len(docs) else (None,) * 5
    style(do.cell(r, 2, x[0]), True, bold=True)
    style(do.cell(r, 3, x[1]), True)
    style(do.cell(r, 4, x[2]), True, DATE)
    style(do.cell(r, 5, x[3]), True)
    style(do.cell(r, 6, x[4]), True, MONEY)
    style(do.cell(r, 7, f'=IF(D{r}="","",IF(D{r}<TODAY(),"Ended",IF(D{r}-TODAY()<=N({SOON_T}),"Ends in "&(D{r}-TODAY())'
                        f'&IF(D{r}-TODAY()=1," day"," days"),"OK")))'), False, align="center")
    style(do.cell(r, 8), True)
    do.cell(r, 10, f'=IF(OR(G{r}="",G{r}="OK"),"",D{r}+ROW()/10000000)')                               # J: order
do.column_dimensions["J"].hidden = True
dropdown(do, f"={CA('B')}", f"B{D0}:B{D1}")
dropdown(do, f'"{",".join(DOCS)}"', f"C{D0}:C{D1}", strict=False)
for cond, colour in (('G{r}="Ended"', BAD), ('LEFT(G{r},4)="Ends"', WARN), ('G{r}="OK"', OK)):
    do.conditional_formatting.add(f"G{D0}:G{D1}", FormulaRule(formula=[cond.format(r=D0)], fill=fill(colour)))

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Car care", "Each car, what is due, and what the cars cost this year.",
           [3, 18, 20, 13, 12, 12, 13, 3, 16, 18, 13, 16, 3], rows=50)
this_year = f'">="&DATE(YEAR(TODAY()),1,1)'
fuel_year = f'SUMIFS({FU("F")},{FU("B")},{this_year})'
cost_year = f'SUMIFS({CO("G")},{CO("B")},{this_year})'
tile(db, "B", 5, "Cars", f'=COUNTIF({CA("B")},"?*")', "0")
tile(db, "C", 5, "Services due soon", f'=COUNTIF({SV("L")},"Due soon")', "0", "B07A00")
tile(db, "D", 5, "Overdue", f'=COUNTIF({SV("L")},"Overdue")', "0", "B4541F")
tile(db, "E", 5, "Papers to renew", f'=COUNTIF({DO("G")},"End*")', "0", "B4541F")
tile(db, "F", 5, "Fuel this year", f"={fuel_year}", MONEY)
tile(db, "G", 5, "All costs this year", f"={fuel_year}+{cost_year}", MONEY)

header(db, 9, 2, ["Car", "Odometer", "Per 100", "Distance per", "This year", "Cost per"])
db["D9"] = f'={VOL}&" per 100 "&{DIST}'
db["E9"] = f'={DIST}&" per "&{VOL}'
db["G9"] = f'="Cost per "&{DIST}'
for k in range(K1 - K0 + 1):
    r, kr = 10 + k, K0 + k
    style(db.cell(r, 2, f'=IF(Cars!B{kr}="","",Cars!B{kr})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",Cars!I{kr})'), False, "#,##0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",Cars!K{kr})'), False, "0.0", align="center")
    style(db.cell(r, 5, f'=IF(OR(B{r}="",N(D{r})=0),"",100/D{r})'), False, "0.0", align="center")
    style(db.cell(r, 6, f'=IF(B{r}="","",SUMIFS({FU("F")},{FU("C")},B{r},{FU("B")},{this_year})+SUMIFS({CO("G")},{CO("C")},B{r},'
                        f'{CO("B")},{this_year}))'), False, MONEY)
    driven = f"(N(Cars!I{kr})-N(Cars!G{kr}))"
    style(db.cell(r, 7, f'=IF(OR(B{r}="",{driven}<=0),"",(SUMIFS({FU("F")},{FU("C")},B{r})+SUMIFS({CO("G")},{CO("C")},B{r}))/{driven})'),
          False, "0.000", align="center")
db["B15"] = "Cost per distance: fuel and every cost, divided by the distance since the start of the file."
db["B15"].font = font(9, color=MUTED, italic=True)

header(db, 17, 2, ["Services", "", "About", "Left", "Status"])
db.merge_cells("B17:C17")
db["E17"] = f'="Left ("&{DIST}&")"'
for i in range(8):
    r = 18 + i
    key = f"N{r}"
    db[key] = f'=IFERROR(SMALL({SV("O")},{i + 1}),"")'
    row = f'MATCH({key},{SV("O")},0)'
    none = '"Nothing due soon"' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({SV("B")},{row}))'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",INDEX({SV("C")},{row}))'), False)
    style(db.cell(r, 4, f'=IF({key}="","",INT({key}))'), False, "d mmm yy", align="center")
    style(db.cell(r, 5, f'=IF({key}="","",INDEX({SV("J")},{row}))'), False, "#,##0;[Red]-#,##0", align="center")
    style(db.cell(r, 6, f'=IF({key}="","",INDEX({SV("L")},{row}))'), False, align="center")
db.conditional_formatting.add("F18:F25", FormulaRule(formula=['F18="Overdue"'], fill=fill(BAD)))
db.conditional_formatting.add("F18:F25", FormulaRule(formula=['F18="Due soon"'], fill=fill(WARN)))

header(db, 17, 9, ["Papers", "", "Ends", "Status"])
db.merge_cells("I17:J17")
for i in range(8):
    r = 18 + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({DO("J")},{i + 1}),"")'
    row = f'MATCH({key},{DO("J")},0)'
    none = '"Nothing to renew soon"' if i == 0 else '""'
    style(db.cell(r, 9, f'=IF({key}="",{none},INDEX({DO("B")},{row}))'), False, bold=True)
    style(db.cell(r, 10, f'=IF({key}="","",INDEX({DO("C")},{row}))'), False)
    style(db.cell(r, 11, f'=IF({key}="","",INT({key}))'), False, "d mmm yy", align="center")
    style(db.cell(r, 12, f'=IF({key}="","",INDEX({DO("G")},{row}))'), False, align="center")
db.conditional_formatting.add("L18:L25", FormulaRule(formula=['L18="Ended"'], fill=fill(BAD)))
db.conditional_formatting.add("L17:L24", FormulaRule(formula=['LEFT(L18,4)="Ends"'], fill=fill(WARN)))

header(db, 28, 2, ["Month", "Fuel", "Other costs", "Total"])
for m in range(12):
    r = 29 + m
    start = f"DATE(YEAR(TODAY()),{m + 1},1)"
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f'=SUMIFS({FU("F")},{FU("B")},">="&{start},{FU("B")},"<="&EOMONTH({start},0))'), False, "#,##0.00;;")
    style(db.cell(r, 4, f'=SUMIFS({CO("G")},{CO("B")},">="&{start},{CO("B")},"<="&EOMONTH({start},0))'), False, "#,##0.00;;")
    style(db.cell(r, 5, f"=C{r}+D{r}"), False, "#,##0.00;;", True)
chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.overlap = 100
chart.title = "This year by month"
chart.height, chart.width = 7.5, 15
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=28, max_row=40), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=29, max_row=40))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "F2A65A"
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "G28")
for c in "NO":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: km or miles, litres or gallons, and how early you want a warning."),
    ("2", "Cars: each car, with the odometer on the day you start and that date."),
    ("3", "Services: what each car needs, every so far or every so many months, and when it was last done."),
    ("4", "Fuel and Costs: one line per fill-up, with an x when you fill the tank, and one line per other cost."),
    ("5", "Documents and Dashboard: the end dates of insurance, tax and inspection, and everything that is due."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Service intervals differ from car to car: the handbook or the garage has yours.",
                          "The example cars and their numbers are made up. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Cars", "Services", "Fuel", "Costs", "Documents", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Car-Maintenance-Tracker.xlsx")
wb.save(out)
print("saved", out, len(cars), "cars", len(fuel), "fill-ups", len(costs), "costs")
