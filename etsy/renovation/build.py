"""Home Renovation Budget Planner (Excel + Google Sheets).

A budget per room with a contingency, one list for every quote, job and
purchase (quotes for the same job are compared automatically), payments
that are due, and a week-by-week timeline. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
T = lambda k: today + dt.timedelta(days=k)
B0, B1 = 6, 17       # rooms
C0, C1 = 6, 305      # quotes and costs
T0, T1 = 6, 35       # timeline tasks
WEEKS = 16
RED_NEG = '#,##0.00;[Red]-#,##0.00'
CS = lambda col: f"Costs!${col}${C0}:${col}${C1}"
TL = lambda col: f"Timeline!${col}${T0}:${col}${T1}"
ROOMS = f"Budget!$B${B0}:$B${B1}"


def days(expr, word="day"):
    """'3 days' or '1 day' for a number expression."""
    return f'{expr}&IF({expr}=1," {word}"," {word}s")'


def room_status(r, budget, committed):
    """Plain-word status for a room: over, near the limit, or share used."""
    used = f"ROUND({committed}{r}/{budget}{r}*100,0)"
    return (f'=IF(B{r}="","",IF(N({budget}{r})=0,IF(N({committed}{r})>0,"No budget set",""),'
            f'IF({committed}{r}>{budget}{r},"Over by "&FIXED({committed}{r}-{budget}{r},0),'
            f'IF({committed}{r}>=0.9*{budget}{r},{used}&"% used, near the limit",{used}&"% used"))))')


def status_colours(ws, rng, first):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${first},4)="Over"'], fill=fill(BAD),
                                                   font=Font(name=F, bold=True, color="9B1C1C")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("near",${first})),${first}="No budget set")'],
                                                   fill=fill(WARN)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${first}<>""'], fill=fill(OK)))


wb = Workbook()

# ---------------------------------------------------------------- Budget
bu = wb.active
bu.title = "Budget"
sheet_base(bu, "Budget by room", "Type a budget for each room or area. Committed and paid come from the Costs tab.",
           [3, 24, 14, 14, 14, 14, 15, 16, 26], rows=40)
bu["B4"] = "Contingency"
bu["B4"].font = font(11, True, TEAL_D)
style(bu["C4"], True, "0%", True, "center")
bu["C4"] = 0.10
bu["D4"] = "Money kept aside for surprises. Between 10 and 20% is usual."
bu["D4"].font = font(10, color=MUTED, italic=True)
header(bu, 5, 2, ["Room or area", "Budget", "Committed", "Paid", "Still to pay", "Left to commit", "Used", "Status"])
rooms = [("Kitchen", 18000), ("Bathroom", 9000), ("Living room", 4500), ("Primary bedroom", 3000),
         ("Whole house", 5000), ("Exterior", 3500)]
for r in range(B0, B1 + 1):
    name, budget = rooms[r - B0] if r - B0 < len(rooms) else (None, None)
    style(bu.cell(r, 2, name), True, bold=True)
    style(bu.cell(r, 3, budget), True, MONEY)
    style(bu.cell(r, 4, f'=IF(B{r}="","",SUMIFS({CS("E")},{CS("B")},B{r},{CS("F")},"Hired")'
                        f'+SUMIFS({CS("E")},{CS("B")},B{r},{CS("F")},"Bought"))'), False, MONEY)
    style(bu.cell(r, 5, f'=IF(B{r}="","",SUMIFS({CS("G")},{CS("B")},B{r},{CS("F")},"Hired")'
                        f'+SUMIFS({CS("G")},{CS("B")},B{r},{CS("F")},"Bought"))'), False, MONEY)
    style(bu.cell(r, 6, f'=IF(B{r}="","",D{r}-E{r})'), False, MONEY)
    style(bu.cell(r, 7, f'=IF(B{r}="","",C{r}-D{r})'), False, RED_NEG, True)
    c = style(bu.cell(r, 8, f'=IF(OR(B{r}="",N(C{r})=0),"",{bar(f"D{r}/C{r}", 12)})'), False)
    c.font = Font(name=F, size=10, color=TEAL)
    style(bu.cell(r, 9, room_status(r, "C", "D")), False)
    bu.cell(r, 10, f'=IF(B{r}="","",MAX(0,N(D{r})-N(C{r})))')  # over budget by, for the contingency
bu.column_dimensions["J"].hidden = True
status_colours(bu, f"I{B0}:I{B1}", f"I{B0}")
t = B1 + 1
bu.cell(t, 2, "Total").font = font(11, True)
for col in "CDEFG":
    style(bu[f"{col}{t}"], False, RED_NEG if col == "G" else MONEY, True)
    bu[f"{col}{t}"] = f"=SUM({col}{B0}:{col}{B1})"
committed_all = f'(SUMIFS({CS("E")},{CS("F")},"Hired")+SUMIFS({CS("E")},{CS("F")},"Bought"))'
summary = [("Room budgets", f"=C{t}", False),
           (f'="Contingency ("&ROUND(C4*100,0)&"%)"', f"=C{t + 2}*C4", False),
           ("Total budget", f"=C{t + 2}+C{t + 3}", True),
           ("Contingency used", f"=SUM(J{B0}:J{B1})+MAX(0,{committed_all}-D{t})", False),
           ("Contingency left", f"=C{t + 3}-C{t + 5}", True)]
for k, (label, formula, bold) in enumerate(summary):
    r = t + 2 + k
    bu.cell(r, 2, label).font = font(11, bold, TEAL_D if bold else "1E2A2F")
    style(bu.cell(r, 3, formula), False, RED_NEG, bold)
bu.cell(t + 8, 2, "Costs with a room that is not on this list also count against the contingency.").font = \
    font(10, color=MUTED, italic=True)
TOTAL_BUDGET, CONTINGENCY, CONT_LEFT, ROOM_TOTAL = f"Budget!$C${t + 4}", f"Budget!$C${t + 3}", f"Budget!$C${t + 6}", t
bu.freeze_panes = "C6"

# ---------------------------------------------------------------- Costs
co = wb.create_sheet("Costs")
sheet_base(co, "Quotes and costs", "One line per quote, job or purchase. Quotes for the same job need the same job name.",
           [3, 17, 26, 22, 13, 11, 13, 15, 13, 24, 26, 24], rows=C1 + 2)
header(co, 5, 2, ["Room", "Job or item", "Contractor or shop", "Amount", "Stage", "Paid so far", "Next payment due",
                  "Left to pay", "Payment status", "Quote check", "Note"])
costs = [
    ("Kitchen", "Cabinets and fitting", "Oakline Kitchens", 7900, "Quote", None, None),
    ("Kitchen", "Cabinets and fitting", "Riverside Cabinet Co", 7200, "Hired", 2160, T(12)),
    ("Kitchen", "Cabinets and fitting", "HomeFit Kitchens", 8450, "Quote", None, None),
    ("Kitchen", "Quartz countertops", "Stone Works", 2650, "Bought", 2650, None),
    ("Kitchen", "Appliances", "Appliance store", 3100, "Bought", 3100, None),
    ("Kitchen", "Kitchen electrics", "Spark Electrical", 1450, "Hired", 0, T(-3)),
    ("Bathroom", "Bathroom remodel", "Blue Tap Plumbing", 6800, "Hired", 2000, T(20)),
    ("Bathroom", "Bathroom remodel", "Clearwater Baths", 7450, "Quote", None, None),
    ("Bathroom", "Tiles", "Tile shop", 1380, "Bought", 1380, None),
    ("Bathroom", "Vanity, toilet and shower", "Bathroom store", 2150, "Bought", 1000, T(6)),
    ("Living room", "Oak flooring", "Floor Masters", 2900, "Hired", 0, T(35)),
    ("Living room", "Painting", "Fresh Coat Painters", 1100, "Quote", None, None),
    ("Living room", "Painting", "Local painter", 950, "Quote", None, None),
    ("Primary bedroom", "Built-in closet", "Riverside Cabinet Co", 2400, "Quote", None, None),
    ("Primary bedroom", "Built-in closet", "Closet Studio", 2150, "Quote", None, None),
    ("Whole house", "Building permit", "Building department", 450, "Bought", 450, None),
    ("Whole house", "Dumpster rental", "Waste hauler", 380, "Hired", 380, None),
    ("Whole house", "Partial rewire", "Spark Electrical", 3900, "Hired", 1950, T(9)),
    ("Exterior", "Fence repair", "Garden and Fence Co", 1250, "Quote", None, None),
    ("Exterior", "Fence repair", "Green Yard Services", 1480, "Quote", None, None),
]
J, A, S = CS("C"), CS("E"), CS("F")
for r in range(C0, C1 + 1):
    row = costs[r - C0] if r - C0 < len(costs) else (None,) * 7
    for col, v, fmt in ((2, row[0], None), (3, row[1], None), (4, row[2], None), (5, row[3], MONEY),
                        (6, row[4], None), (7, row[5], MONEY), (8, row[6], DATE)):
        style(co.cell(r, col, v), True, fmt, align="center" if col == 6 else None)
    style(co.cell(r, 9, f'=IF(OR(E{r}="",AND(F{r}<>"Hired",F{r}<>"Bought")),"",E{r}-N(G{r}))'), False, MONEY)
    late, left = f"(TODAY()-H{r})", f"(H{r}-TODAY())"
    pay = (f'=IF(C{r}&E{r}="","",IF(F{r}="","Choose a stage",IF(F{r}="Quote","",IF(E{r}="","Add the amount",'
           f'IF(I{r}<=0,"Paid",IF(H{r}="","Balance due, no date set",IF(H{r}<TODAY(),"Overdue by "&{days(late)},'
           f'IF(H{r}=TODAY(),"Due today",IF({left}<=14,"Due in "&{days(left)},"Next payment in "&{left}&" days")))))))))')
    style(co.cell(r, 10, pay), False)
    n_quotes = f'COUNTIFS({J},C{r},{A},"<>")'
    cheapest = f"SUMIFS({A},{J},C{r},{CS('M')},1)/COUNTIFS({J},C{r},{CS('M')},1)"
    style(co.cell(r, 11, f'=IF(OR(C{r}="",E{r}=""),"",IF({n_quotes}<2,"",IF(M{r}=1,"Cheapest of "&{n_quotes},'
                         f'FIXED(E{r}-{cheapest},0)&" more than the cheapest")))'), False)
    style(co.cell(r, 12), True)
    co.cell(r, 13, f'=IF(OR(C{r}="",E{r}=""),"",IF(COUNTIFS({J},C{r},{A},"<"&E{r})=0,1,0))')
    co.cell(r, 14, f'=IF(AND(C{r}<>"",F{r}="Quote",COUNTIFS({J},C{r},{S},"Hired")+COUNTIFS({J},C{r},{S},"Bought")=0),'
                   f'1/COUNTIFS({J},C{r},{S},"Quote"),"")')
    co.cell(r, 15, f'=IF(AND(OR(F{r}="Hired",F{r}="Bought"),N(I{r})>0,H{r}<>""),H{r}+ROW()/100000,"")')
for col in "MNO":
    co.column_dimensions[col].hidden = True
for rng, formula in ((f"B{C0}:B{C1}", f"={ROOMS}"), (f"F{C0}:F{C1}", '"Quote,Hired,Bought"')):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    co.add_data_validation(dv)
    dv.add(rng)
pr = f"J{C0}:J{C1}"
co.conditional_formatting.add(pr, FormulaRule(formula=[f'LEFT($J{C0},7)="Overdue"'], fill=fill(BAD),
                                              font=Font(name=F, bold=True, color="9B1C1C")))
co.conditional_formatting.add(pr, FormulaRule(formula=[f'OR(LEFT($J{C0},6)="Due in",$J{C0}="Due today")'], fill=fill(WARN)))
co.conditional_formatting.add(pr, FormulaRule(formula=[f'$J{C0}="Paid"'], fill=fill(OK)))
co.conditional_formatting.add(pr, FormulaRule(formula=[f'OR($J{C0}="Choose a stage",$J{C0}="Add the amount")'], fill=fill(INFO)))
co.conditional_formatting.add(f"K{C0}:K{C1}", FormulaRule(formula=[f'LEFT($K{C0},8)="Cheapest"'], fill=fill(OK)))
co.conditional_formatting.add(f"B{C0}:E{C1}", FormulaRule(formula=[f'$F{C0}="Quote"'], font=Font(name=F, color="7A8A8E")))
co.freeze_panes = "D6"

# ---------------------------------------------------------------- Timeline
tl = wb.create_sheet("Timeline")
sheet_base(tl, "Timeline", "One line per task. The coloured bars follow the start and end dates, one column per week.",
           [3, 26, 15, 20, 12, 12, 13] + [3.4] * WEEKS, rows=T1 + 8)
tl["B4"] = "Project start"
tl["B4"].font = font(11, True, TEAL_D)
start = T(-24) - dt.timedelta(days=T(-24).weekday())
style(tl["C4"], True, DATE, True)
tl["C4"] = start
header(tl, 5, 2, ["Task", "Room", "Who", "Start", "End", "Status"])
W0 = 8  # first week column (H)
for k in range(WEEKS):
    c = tl.cell(5, W0 + k, "=$C$4-WEEKDAY($C$4,3)" if k == 0 else f"={tl.cell(5, W0 + k - 1).coordinate}+7")
    c.number_format = "DD MMM"
    c.font = font(9, True, "FFFFFF")
    c.fill = fill(TEAL)
    c.border = box
    c.alignment = Alignment(textRotation=90, horizontal="center", vertical="center")
tl.row_dimensions[5].height = 58
tasks = [("Demolition", "Whole house", "Us", -24, -20, "Done"),
         ("Partial rewire", "Whole house", "Spark Electrical", -19, -10, "Done"),
         ("Kitchen electrics", "Kitchen", "Spark Electrical", -9, -5, "Done"),
         ("Bathroom remodel", "Bathroom", "Blue Tap Plumbing", -6, 12, "In progress"),
         ("Cabinets and fitting", "Kitchen", "Riverside Cabinet Co", -2, 9, "In progress"),
         ("Drywall repair in the hall", "Whole house", "Handyman", -4, 3, "Delayed"),
         ("Quartz countertops", "Kitchen", "Stone Works", 10, 16, "Planned"),
         ("Oak flooring", "Living room", "Floor Masters", 17, 24, "Planned"),
         ("Painting", "Living room", "Not chosen yet", 25, 32, "Planned"),
         ("Built-in closet", "Primary bedroom", "Not chosen yet", 35, 42, "Planned"),
         ("Fence repair", "Exterior", "Not chosen yet", 44, 47, "Planned"),
         ("Final clean and punch list", "Whole house", "Us", 50, 54, "Planned")]
last_col = tl.cell(5, W0 + WEEKS - 1).column_letter
for r in range(T0, T1 + 1):
    tk = tasks[r - T0] if r - T0 < len(tasks) else None
    style(tl.cell(r, 2, tk[0] if tk else None), True)
    style(tl.cell(r, 3, tk[1] if tk else None), True)
    style(tl.cell(r, 4, tk[2] if tk else None), True)
    style(tl.cell(r, 5, T(tk[3]) if tk else None), True, DATE)
    style(tl.cell(r, 6, T(tk[4]) if tk else None), True, DATE)
    style(tl.cell(r, 7, tk[5] if tk else None), True, align="center")
    for k in range(WEEKS):
        style(tl.cell(r, W0 + k), False)
grid = f"H{T0}:{last_col}{T1}"
on = f"$E{T0}<>\"\",$F{T0}<>\"\",H$5<=$F{T0},H$5+6>=$E{T0}"
TASK_COLOURS = [("Done", "9FD3A8"), ("In progress", TEAL), ("Delayed", "E58A7F"), ("Planned", "B9CCE8")]
for status, colour in TASK_COLOURS:
    cond = f'$G{T0}="{status}"' if status != "Planned" else f'OR($G{T0}="Planned",$G{T0}="")'
    tl.conditional_formatting.add(grid, FormulaRule(formula=[f"AND({on},{cond})"], fill=fill(colour)))
this_week = "AND(H$5<=TODAY(),H$5+6>=TODAY())"
tl.conditional_formatting.add(grid, FormulaRule(formula=[this_week], fill=fill("FFF1D6")))
tl.conditional_formatting.add(f"H5:{last_col}5", FormulaRule(formula=[this_week], fill=fill("F2C14E"),
                                                              font=Font(name=F, bold=True, color="1E2A2F")))
dv = DataValidation(type="list", formula1='"Planned,In progress,Done,Delayed"', allow_blank=True)
tl.add_data_validation(dv)
dv.add(f"G{T0}:G{T1}")
dv = DataValidation(type="list", formula1=f"={ROOMS}", allow_blank=True)
tl.add_data_validation(dv)
dv.add(f"C{T0}:C{T1}")
key = T1 + 2
tl.cell(key, 2, "Colours:").font = font(10, True, MUTED)
for k, (label, colour) in enumerate(TASK_COLOURS + [("This week", "FFF1D6")]):
    c = tl.cell(key, 3 + k, label)
    c.fill = fill(colour)
    c.font = font(10, color="FFFFFF" if colour == TEAL else "1E2A2F")
    c.alignment = Alignment(horizontal="center")
    c.border = box
tl.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Renovation dashboard", "Everything here comes from the Budget, Costs and Timeline tabs.",
           [3, 22, 15, 15, 15, 22, 24, 3, 14, 34, 15, 3], rows=40)
paid_all = f'(SUMIFS({CS("G")},{CS("F")},"Hired")+SUMIFS({CS("G")},{CS("F")},"Bought"))'
tiles = [("B", "Total budget", f"={TOTAL_BUDGET}", TEAL, MONEY),
         ("C", "Committed", f"={committed_all}", "B4541F", MONEY),
         ("D", "Paid", f"={paid_all}", TEAL, MONEY),
         ("E", "Still to pay", "=C7-D7", "B4541F", MONEY),
         ("F", "Contingency left", f"={CONT_LEFT}", TEAL, RED_NEG),
         ("G", "Budget committed", '=IF(B7=0,"",C7/B7)', TEAL, "0%")]
for col, label, formula, colour, fmt in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = fmt
    db[f"{col}7"].font = font(17, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db.merge_cells("I6:K6")
db.merge_cells("I7:K7")
db["I6"] = "Left to spend"
db["I7"] = "=B7-C7"
db["I7"].number_format = MONEY
for a, size in (("I6", 11), ("I7", 28)):
    db[a].font = font(size, True, "FFFFFF")
    db[a].fill = fill(TEAL_D)
    db[a].alignment = Alignment(horizontal="center", vertical="center")
db.conditional_formatting.add("I7", FormulaRule(formula=["$I$7<0"], fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
db.row_dimensions[7].height = 44

header(db, 9, 2, ["Room or area", "Budget", "Committed", "Paid", "Used", "Status"])
for k in range(B1 - B0 + 1):
    r, br = 10 + k, B0 + k
    style(db.cell(r, 2, f'=IF(Budget!B{br}="","",Budget!B{br})'), False, bold=True)
    for col, src in (("C", "C"), ("D", "D"), ("E", "E")):
        style(db[f"{col}{r}"], False, MONEY)
        db[f"{col}{r}"] = f'=IF(B{r}="","",Budget!{src}{br})'
    c = style(db.cell(r, 6, f'=IF(OR(B{r}="",N(C{r})=0),"",{bar(f"D{r}/C{r}", 14)})'), False)
    c.font = Font(name=F, size=10, color=TEAL)
    style(db.cell(r, 7, f'=IF(B{r}="","",Budget!I{br})'), False)
status_colours(db, "G10:G21", "G10")

# Hidden helpers: counts and dates used by the notes.
H = {"over": f'COUNTIFS(Budget!$I${B0}:$I${B1},"Over*")',
     "late": f'COUNTIFS({CS("J")},"Overdue*")', "late_sum": f'SUMIFS({CS("I")},{CS("J")},"Overdue*")',
     "soon": f'COUNTIFS({CS("J")},"Due in*")+COUNTIFS({CS("J")},"Due today")',
     "soon_sum": f'SUMIFS({CS("I")},{CS("J")},"Due in*")+SUMIFS({CS("I")},{CS("J")},"Due today")',
     "open": f'ROUND(SUM({CS("N")}),0)', "delayed": f'COUNTIFS({TL("G")},"Delayed")',
     "next": f'IFERROR(SMALL({TL("E")},COUNTIF({TL("E")},"<"&TODAY())+1),"")'}
ref = {}
for k, (name, formula) in enumerate(H.items()):
    db.cell(1 + k, 14, "=" + formula)
    ref[name] = f"$N${1 + k}"
db.column_dimensions["N"].hidden = True
nxt = ref["next"]
next_task = f'INDEX({TL("B")},MATCH({nxt},{TL("E")},0))'
notes = [
    (f'=IF({ref["over"]}=0,"No room is over budget",IF({ref["over"]}=1,"1 room is",{ref["over"]}&" rooms are")&" over budget")', BAD, "No"),
    (f'=IF({ref["late"]}=0,"No overdue payments",IF({ref["late"]}=1,"1 payment is",{ref["late"]}&" payments are")'
     f'&" overdue: "&FIXED({ref["late_sum"]},2))', BAD, "No"),
    (f'=IF({ref["soon"]}=0,"Nothing due in the next 14 days",IF({ref["soon"]}=1,"1 payment",{ref["soon"]}&" payments")'
     f'&" due in the next 14 days: "&FIXED({ref["soon_sum"]},2))', WARN, "Nothing"),
    (f'=IF({ref["open"]}=0,"No quotes waiting for a decision",IF({ref["open"]}=1,"1 job is",{ref["open"]}&" jobs are")'
     f'&" waiting for a decision between quotes")', INFO, "No"),
    (f'=IF({CONT_LEFT}>=0,"Contingency left: "&FIXED({CONT_LEFT},2)&" of "&FIXED({CONTINGENCY},2),'
     f'"Contingency used up, over by "&FIXED(-{CONT_LEFT},2))', OK, None),
    (f'=IF({ref["delayed"]}=0,"No delayed tasks",IF({ref["delayed"]}=1,"1 task is",{ref["delayed"]}&" tasks are")&" delayed")',
     BAD, "No"),
    (f'=IF({nxt}="","No tasks coming up","Next task: "&{next_task}&IF({nxt}=TODAY(),", starting today",'
     f'", starts in "&{days(f"({nxt}-TODAY())")}))', INFO, None),
]
db.merge_cells("I9:K9")
db["I9"] = "Needs attention"
db["I9"].font = font(11, True, "FFFFFF")
db["I9"].fill = fill(TEAL)
db["I9"].alignment = Alignment(vertical="center")
for k, (formula, colour, calm) in enumerate(notes):
    r = 10 + k
    db.merge_cells(f"I{r}:K{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "IJK":
        db[f"{col}{r}"].border = box
    if calm:
        db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))
db.conditional_formatting.add("I14", FormulaRule(formula=['LEFT(I14,16)="Contingency used"'], fill=fill(BAD)))

header(db, 18, 9, ["Next payments", "To", "Amount"])
for k in range(5):
    r = 19 + k
    db.cell(r, 13, f'=IFERROR(SMALL({CS("O")},{k + 1}),"")')
    hit = f'MATCH($M{r},{CS("O")},0)'
    style(db.cell(r, 9, f'=IF($M{r}="","",INT($M{r}))'), False, DATE, align="center")
    style(db.cell(r, 10, f'=IF($M{r}="","",INDEX({CS("D")},{hit})&", "&INDEX({CS("C")},{hit}))'), False)
    style(db.cell(r, 11, f'=IF($M{r}="","",INDEX({CS("I")},{hit}))'), False, MONEY)
db.column_dimensions["M"].hidden = True
db.conditional_formatting.add("I19:K23", FormulaRule(formula=['AND($I19<>"",$I19<TODAY())'], fill=fill(BAD)))
db.conditional_formatting.add("I19:K23", FormulaRule(formula=['AND($I19<>"",$I19-TODAY()<=14)'], fill=fill(WARN)))

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Budget: list your rooms or areas with a budget for each, and set the contingency (10 to 20% is usual)."),
    ("2", "Costs: add every quote, job and purchase. Quotes for the same job get the same job name."),
    ("3", "When you hire someone or buy something, change the stage from Quote to Hired or Bought."),
    ("4", "When you pay, update 'Paid so far' and the next payment date. Late payments turn red."),
    ("5", "Timeline: add each task with its start and end dates, and the bars fill in week by week."),
    ("6", "Dashboard: budget used per room, payments due soon, quotes to decide and the next task."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
st["C18"] = "Only lines marked Hired or Bought count as committed. Quotes stay grey until you choose one."
st["C18"].font = font(11, color=MUTED, italic=True)
st["C20"] = "The example rooms, quotes and tasks show how it works. Delete them and add your own."
st["C20"].font = font(11, color=MUTED, italic=True)
st["C22"] = "Works in Google Sheets and Microsoft Excel, in any currency."
st["C22"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Home-Renovation-Budget-Planner.xlsx")
wb.save(out)
print("saved", out)
