"""Travel Planner: itinerary, bookings, budget and packing list (Excel + Google Sheets).

A countdown, a day by day itinerary, every booking with its confirmation code
and what is still to pay, a budget per category that counts the bookings and
the spending in any currency, and a packing list. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
I0, I1 = 6, 305       # itinerary lines
B0, B1 = 6, 105       # bookings
X0, X1 = 6, 505       # expenses
K0, K1 = 6, 255       # packing list
CATS = ["Flights", "Lodging", "Transport", "Food", "Activities", "Shopping", "Other"]
TYPES = [("Flight", "Flights"), ("Hotel", "Lodging"), ("Train", "Transport"), ("Car", "Transport"), ("Tour", "Activities"),
         ("Restaurant", "Food"), ("Other", "Other")]
PACK_CATS = ["Documents", "Clothes", "Toiletries", "Electronics", "Health", "Other"]
IT = lambda col: f"Itinerary!${col}${I0}:${col}${I1}"
BK = lambda col: f"Bookings!${col}${B0}:${col}${B1}"
EX = lambda col: f"Expenses!${col}${X0}:${col}${X1}"
PK = lambda col: f"'Packing list'!${col}${K0}:${col}${K1}"
CURS, RATES = "Setup!$E$9:$E$18", "Setup!$F$9:$F$18"
START, END = "Setup!$C$6", "Setup!$C$7"

# ---------------------------------------------------------------- example data
S = dt.date(Y, 10, 20) if today < dt.date(Y, 10, 1) else today + dt.timedelta(days=22)
day = lambda n: S + dt.timedelta(days=n)
T = lambda h, m=0: dt.time(h, m)
bookings = [
    # type, what, starts, time, ends, confirmation, cost, currency, paid, note
    ("Flight", "Los Angeles to Tokyo Haneda", day(0), T(11, 5), day(1), "KX7Q2M", 1180, "USD", 1180, "Two seats, 23 kg bags"),
    ("Hotel", "Shinjuku hotel, 5 nights", day(1), T(15), day(6), "88213-04", 145000, "JPY", 0, "Pay at the hotel"),
    ("Tour", "teamLab Planets", day(2), T(15), day(2), "TLP-5521", 7600, "JPY", 7600, None),
    ("Tour", "Sushi cooking class", day(4), T(10), day(4), "SC-3310", 24000, "JPY", 24000, "Meet at Tsukiji station"),
    ("Train", "Rail pass, 7 days, two people", day(6), T(8), day(12), "JRP-77410", 100000, "JPY", 100000, "Exchange at Tokyo station"),
    ("Hotel", "Kyoto ryokan, 4 nights", day(6), T(15), day(10), "RK-2291", 168000, "JPY", 50000, "Dinner included twice"),
    ("Tour", "Tea ceremony in Kyoto", day(8), T(14), day(8), None, 8000, "JPY", 0, "Pay on the day"),
    ("Restaurant", "Kaiseki dinner", day(9), T(19), day(9), "KD-14", 36000, "JPY", 0, None),
    ("Hotel", "Osaka hotel, 1 night", day(10), T(15), day(11), "OS-5570", 22000, "JPY", 22000, None),
    ("Flight", "Osaka Kansai to Los Angeles", day(11), T(17, 20), day(11), "KX7Q2M", 1180, "USD", 1180, None),
]
itinerary = [
    (day(0), T(11, 5), "Flight to Tokyo", "Los Angeles airport", "KX7Q2M"),
    (day(1), T(15, 30), "Land at Haneda, train to Shinjuku", "Haneda airport", None),
    (day(1), T(17), "Check in", "Shinjuku hotel", "88213-04"),
    (day(1), T(19), "Ramen dinner", "Omoide Yokocho", None),
    (day(2), T(9), "Meiji Shrine", "Harajuku", None),
    (day(2), T(11), "Walk Omotesando", "Omotesando", None),
    (day(2), T(15), "teamLab Planets", "Toyosu", "TLP-5521"),
    (day(2), T(19), "Shibuya crossing at night", "Shibuya", None),
    (day(3), T(8), "Breakfast at the outer market", "Tsukiji", None),
    (day(3), T(11), "Senso-ji temple", "Asakusa", None),
    (day(3), T(16), "Tokyo Skytree", "Oshiage", None),
    (day(4), T(10), "Sushi cooking class", "Tsukiji", "SC-3310"),
    (day(4), T(15), "Akihabara", "Akihabara", None),
    (day(5), T(8), "Day trip to Nikko", "Nikko", None),
    (day(6), T(8), "Exchange rail pass, train to Kyoto", "Tokyo station", "JRP-77410"),
    (day(6), T(15), "Check in", "Kyoto ryokan", "RK-2291"),
    (day(6), T(17), "Fushimi Inari at sunset", "Fushimi", None),
    (day(7), T(7, 30), "Bamboo grove before the crowds", "Arashiyama", None),
    (day(7), T(13), "Golden Pavilion", "Kinkaku-ji", None),
    (day(8), T(10), "Nishiki market", "Central Kyoto", None),
    (day(8), T(14), "Tea ceremony", "Gion", None),
    (day(8), T(18), "Evening walk in Gion", "Gion", None),
    (day(9), T(9), "Day trip to Nara, deer park", "Nara", None),
    (day(9), T(19), "Kaiseki dinner", "Pontocho", "KD-14"),
    (day(10), T(11), "Train to Osaka, check in", "Osaka hotel", "OS-5570"),
    (day(10), T(18), "Street food in Dotonbori", "Dotonbori", None),
    (day(11), T(10), "Osaka castle", "Osaka", None),
    (day(11), T(17, 20), "Flight home", "Kansai airport", "KX7Q2M"),
]
expenses = [(today - dt.timedelta(days=30), "Travel insurance, two people", "Other", 180, "USD"),
            (today - dt.timedelta(days=21), "Guidebook", "Other", 25, "USD"),
            (today - dt.timedelta(days=12), "Pocket wifi rental", "Other", 6000, "JPY"),
            (today - dt.timedelta(days=5), "Cabin suitcase", "Shopping", 95, "USD")]
budget = [2400, 2200, 900, 1300, 800, 500, 400]
packing = {
    "Documents": ["Passports", "Printed bookings", "Travel insurance card", "Credit cards", "Some cash in yen", "Rail pass voucher"],
    "Clothes": ["T-shirts", "Long sleeve tops", "Jeans", "Light jacket", "Sweater", "Pajamas", "Underwear", "Socks",
                "Walking shoes", "Slip-on shoes for temples", "Scarf"],
    "Toiletries": ["Toothbrush and toothpaste", "Deodorant", "Sunscreen", "Shampoo, travel size", "Hairbrush", "Razor",
                   "Lip balm"],
    "Electronics": ["Phone and charger", "Power bank", "Plug adapter", "Headphones", "Camera"],
    "Health": ["Painkillers", "Plasters", "Motion sickness tablets", "Prescription medicine", "Hand sanitiser"],
    "Other": ["Day backpack", "Reusable water bottle", "Small towel", "Coin purse", "Folding umbrella", "Earplugs and eye mask"],
}
qty = {"T-shirts": 6, "Long sleeve tops": 3, "Jeans": 2, "Underwear": 12, "Socks": 12, "Passports": 2, "Credit cards": 2}
rng = random.Random(4)
packed = {i for c in ("Documents", "Electronics") for i in packing[c]} | {"Walking shoes", "Light jacket", "Plasters",
                                                                          "Plug adapter", "Folding umbrella"}
packed -= {"Some cash in yen", "Phone and charger"}

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL, span=1):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    first = ws[f"{col}{row}"].column
    for r in (row, row + 1):
        for c in range(first, first + span):
            ws.cell(r, c).border = box
        if span > 1:
            ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=first + span - 1)
    ws.row_dimensions[row + 1].height = 34


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


home = lambda amount, cur: f'IF(N({amount})=0,0,{amount}*IF({cur}="",1,IFERROR(INDEX({RATES},MATCH({cur},{CURS},0)),0)))'

# ---------------------------------------------------------------- Setup
se = wb.active
se.title = "Setup"
sheet_base(se, "Setup", "The trip, its dates, and the currencies you will pay in.", [3, 18, 16, 4, 12, 14, 4, 22, 14], rows=30)
for r, (label, value, fmt) in enumerate([("Trip", "Japan", None), ("Travelers", 2, "0"), ("First day", S, DATE),
                                         ("Last day", day(11), DATE), ("Home currency", "USD", None)], start=4):
    se.cell(r, 2, label).font = font(11, True)
    style(se.cell(r, 3, value), True, fmt, True, "center")
header(se, 8, 5, ["Currency", "1 unit is worth"])
style(se.cell(9, 5, "=$C$8"), False, bold=True, align="center")
style(se.cell(9, 6, 1), False, "0.0000", align="center")
for k, (cur, rate) in enumerate([("JPY", 0.0068), ("EUR", 1.08), ("GBP", 1.27)] + [(None, None)] * 6):
    style(se.cell(10 + k, 5, cur), True, bold=True, align="center")
    style(se.cell(10 + k, 6, rate), True, "0.0000", align="center")
header(se, 8, 8, ["Booking type", "Budget category"])
for k, (t, c) in enumerate(TYPES):
    style(se.cell(9 + k, 8, t), True, bold=True)
    style(se.cell(9 + k, 9, c), True)
dropdown(se, f'"{",".join(CATS)}"', "I9:I15")
se["B11"] = "Each currency's worth in your home currency, for example 1 JPY is worth 0.0068 USD."
se["B12"] = "Each kind of booking counts in one budget category."
for a in ("B11", "B12"):
    se[a].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Itinerary
it = wb.create_sheet("Itinerary")
sheet_base(it, "Itinerary", "What you plan to do, day by day. Put the lines in any order: the dashboard sorts them.",
           [3, 18, 8, 34, 22, 14, 26], rows=I1 + 2)
header(it, 5, 2, ["Date", "Time", "Plan", "Where", "Booking", "Note"])
for r in range(I0, I1 + 1):
    e = itinerary[r - I0] if r - I0 < len(itinerary) else (None,) * 5
    style(it.cell(r, 2, e[0]), True, "ddd d mmm yyyy", bold=True)
    style(it.cell(r, 3, e[1]), True, "h:mm", align="center")
    style(it.cell(r, 4, e[2]), True, bold=True)
    style(it.cell(r, 5, e[3]), True)
    style(it.cell(r, 6, e[4]), True, align="center")
    style(it.cell(r, 7), True)
    it.cell(r, 8, f'=IF(OR(B{r}="",D{r}=""),"",B{r}+N(C{r})+ROW()/100000000)')                     # H sort key
    it.cell(r, 9, f'=IF(B{r}="","","Day "&(B{r}-{START}+1))')                                     # I day number
it.column_dimensions["H"].hidden = True
it.column_dimensions["I"].hidden = True
it.conditional_formatting.add(f"B{I0}:G{I1}", FormulaRule(formula=[f'AND($B{I0}<>"",$B{I0}<>$B{I0 - 1},ROW()>{I0})'],
                                                          fill=fill("E6F2F2")))
it.conditional_formatting.add(f"B{I0}:G{I1}", FormulaRule(formula=[f'AND($B{I0}<>"",$B{I0}=TODAY())'], fill=fill("FCE9B8")))
it.freeze_panes = "B6"

# ---------------------------------------------------------------- Bookings
bk = wb.create_sheet("Bookings")
sheet_base(bk, "Bookings", "Every flight, hotel, train and ticket, with its confirmation code and what is still to pay.",
           [3, 12, 30, 13, 8, 13, 14, 12, 9, 12, 13, 13, 14, 24], rows=B1 + 2)
header(bk, 5, 2, ["Type", "What", "Starts", "Time", "Ends", "Confirmation", "Cost", "Currency", "Paid", "In home currency",
                  "Still to pay", "When", "Note"])
for r in range(B0, B1 + 1):
    b = bookings[r - B0] if r - B0 < len(bookings) else (None,) * 10
    typ, what, st_, tm, en, conf, cost, cur, paid, note = b
    style(bk.cell(r, 2, typ), True, bold=True)
    style(bk.cell(r, 3, what), True)
    style(bk.cell(r, 4, st_), True, DATE)
    style(bk.cell(r, 5, tm), True, "h:mm", align="center")
    style(bk.cell(r, 6, en), True, DATE)
    style(bk.cell(r, 7, conf), True, align="center")
    style(bk.cell(r, 8, cost), True, "#,##0.##")
    style(bk.cell(r, 9, cur), True, align="center")
    style(bk.cell(r, 10, paid), True, "#,##0.##")
    style(bk.cell(r, 11, f'=IF(N(H{r})=0,"",{home(f"H{r}", f"I{r}")})'), False, MONEY, True)
    left = f"{home(f'H{r}', f'I{r}')}-{home(f'J{r}', f'I{r}')}"
    style(bk.cell(r, 12, f'=IF(N(H{r})=0,"",IF({left}>0.005,{left},""))'), False, MONEY)
    days = lambda a: f'{a}&IF({a}=1," day"," days")'
    style(bk.cell(r, 13, f'=IF(D{r}="","",IF(D{r}>TODAY(),"In "&{days(f"(D{r}-TODAY())")},IF(D{r}=TODAY(),"Today",'
                         f'IF(OR(F{r}="",F{r}<TODAY()),"Done","Now"))))'), False, align="center")
    style(bk.cell(r, 14, note), True)
    bk.cell(r, 15, f'=IF(B{r}="","",IFERROR(INDEX(Setup!$I$9:$I$15,MATCH(B{r},Setup!$H$9:$H$15,0)),"Other"))')   # O category
    bk.cell(r, 16, f'=IF(N(H{r})=0,0,{home(f"J{r}", f"I{r}")})')                                    # P paid in home currency
    bk.cell(r, 17, f'=IF(OR(D{r}="",M{r}="Done"),"",D{r}+N(E{r})+ROW()/100000000)')                 # Q upcoming key
for c in "OPQ":
    bk.column_dimensions[c].hidden = True
dropdown(bk, "=Setup!$H$9:$H$15", f"B{B0}:B{B1}")
dropdown(bk, f"={CURS}", f"I{B0}:I{B1}")
bk.conditional_formatting.add(f"L{B0}:L{B1}", FormulaRule(formula=[f"N($L{B0})>0"], fill=fill(WARN)))
bk.conditional_formatting.add(f"M{B0}:M{B1}", FormulaRule(formula=[f'OR(M{B0}="Today",M{B0}="Now")'], fill=fill("FCE9B8")))
bk.conditional_formatting.add(f"M{B0}:M{B1}", FormulaRule(formula=[f'M{B0}="Done"'], fill=fill(OK)))
bk.freeze_panes = "C6"

# ---------------------------------------------------------------- Expenses
ex = wb.create_sheet("Expenses")
sheet_base(ex, "Expenses", "Everything you pay that is not a booking: food, tickets, souvenirs, taxis.",
           [3, 13, 30, 14, 12, 10, 14], rows=X1 + 2)
header(ex, 5, 2, ["Date", "What", "Category", "Amount", "Currency", "In home currency"])
for r in range(X0, X1 + 1):
    e = expenses[r - X0] if r - X0 < len(expenses) else (None,) * 5
    for c, (v, fmt) in enumerate(zip(e, (DATE, None, None, "#,##0.##", None)), start=2):
        style(ex.cell(r, c, v), True, fmt, align="center" if c == 6 else None)
    style(ex.cell(r, 7, f'=IF(N(E{r})=0,"",{home(f"E{r}", f"F{r}")})'), False, MONEY, True)
dropdown(ex, f'"{",".join(CATS)}"', f"D{X0}:D{X1}")
dropdown(ex, f"={CURS}", f"F{X0}:F{X1}")
ex.freeze_panes = "B6"

# ---------------------------------------------------------------- Packing list
pk = wb.create_sheet("Packing list")
sheet_base(pk, "Packing list", "Type x under Packed as things go in the bag. Add your own at the end.", [3, 30, 14, 7, 9],
           rows=K1 + 2)
header(pk, 5, 2, ["Item", "Category", "How many", "Packed"])
items = [(i, c) for c in PACK_CATS for i in packing[c]]
for r in range(K0, K1 + 1):
    i = items[r - K0] if r - K0 < len(items) else (None, None)
    style(pk.cell(r, 2, i[0]), True, bold=True)
    style(pk.cell(r, 3, i[1]), True)
    style(pk.cell(r, 4, qty.get(i[0], 1) if i[0] else None), True, "0", align="center")
    style(pk.cell(r, 5, "x" if i[0] in packed else None), True, align="center")
dropdown(pk, f'"{",".join(PACK_CATS)}"', f"C{K0}:C{K1}")
pk.conditional_formatting.add(f"B{K0}:E{K1}", FormulaRule(formula=[f'AND($B{K0}<>"",$E{K0}<>"")'], fill=fill(OK),
                                                          font=Font(name=F, color="5E6E72", strike=True)))
pk.freeze_panes = "B6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Trip planner", "The countdown, the budget, what is booked and what is next.",
           [3, 18, 13, 13, 13, 13, 13, 3, 12, 26, 13, 14, 3], rows=60)
db["B2"] = '=Setup!$C$4&" trip planner"'
MONTHS = ",".join(f'"{m}"' for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
nice = lambda d: f'DAY({d})&" "&CHOOSE(MONTH({d}),{MONTHS})&" "&YEAR({d})'
db["B4"] = f'=IF({START}="","",{nice(START)}&" to "&{nice(END)}&", "&Setup!$C$5&IF(Setup!$C$5=1," traveler"," travelers"))'
db["B4"].font = font(11, True, TEAL_D)
countdown = (f'IF({START}="","",IF(TODAY()<{START},{START}-TODAY(),IF(TODAY()<={END},"Day "&(TODAY()-{START}+1),'
             f'"Home")))')
tile(db, "B", 6, "Days to go", f"={countdown}", "0", "B07A00")
tile(db, "C", 6, "Nights", f'=IF({START}="","",{END}-{START})', "0")
planned = "SUM(C12:C18)"
tile(db, "D", 6, "Budget", f"={planned}", MONEY)
tile(db, "E", 6, "Booked and spent", "=SUM(D12:D18)", MONEY)
tile(db, "F", 6, "Left", f"={planned}-SUM(D12:D18)", MONEY, "B4541F")
tile(db, "G", 6, "Still to pay", f'=SUM({BK("L")})', MONEY, "B4541F")
packed_n = f'COUNTIFS({PK("B")},"?*",{PK("E")},"?*")'
tile(db, "I", 6, "Packed", f'=IF(COUNTIF({PK("B")},"?*")=0,"",{packed_n}/COUNTIF({PK("B")},"?*"))', "0%", span=2)

header(db, 11, 2, ["Budget", "Planned", "Booked and spent", "Paid so far", "Left", "Used"])
for k, cat in enumerate(CATS):
    r = 12 + k
    style(db.cell(r, 2, cat), False, bold=True)
    style(db.cell(r, 3, budget[k]), True, MONEY)
    committed = f'SUMIFS({BK("K")},{BK("O")},B{r})+SUMIFS({EX("G")},{EX("D")},B{r})'
    style(db.cell(r, 4, f"={committed}"), False, MONEY)
    style(db.cell(r, 5, f'=SUMIFS({BK("P")},{BK("O")},B{r})+SUMIFS({EX("G")},{EX("D")},B{r})'), False, MONEY)
    style(db.cell(r, 6, f"=C{r}-D{r}"), False, '#,##0.00;[Red]-#,##0.00', True)
    c = style(db.cell(r, 7, f'=IF(N(C{r})=0,"",{bar(f"D{r}/C{r}", 10)})'), False)
    c.font = Font(name=F, size=9, color=TEAL)
style(db.cell(19, 2, "Total"), False, bold=True)
for col in "CDEF":
    style(db[f"{col}19"], False, MONEY, True)
    db[f"{col}19"] = f"=SUM({col}12:{col}18)"
db.conditional_formatting.add("F12:F18", FormulaRule(formula=["F12<0"], fill=fill(BAD)))
db["B20"] = "Booked and spent counts the full cost of every booking, paid or not, plus the expenses."
db["B20"].font = font(9, color=MUTED, italic=True)

header(db, 11, 9, ["Next bookings", "", "Confirmation", "Still to pay"])
db.merge_cells("I11:J11")
for k in range(8):
    r = 12 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({BK("Q")},{k + 1}),"")'
    row = f'MATCH({key},{BK("Q")},0)'
    style(db.cell(r, 9, f'=IF({key}="","",INT({key}))'), False, "d mmm", align="center")
    style(db.cell(r, 10, f'=IF({key}="","",INDEX({BK("C")},{row}))'), False, bold=True)
    style(db.cell(r, 11, f'=IF({key}="","",INDEX({BK("G")},{row})&"")'), False, align="center")
    style(db.cell(r, 12, f'=IF({key}="","",INDEX({BK("L")},{row}))'), False, MONEY)
db.conditional_formatting.add("L12:L19", FormulaRule(formula=["N(L12)>0"], fill=fill(WARN)))

header(db, 22, 9, ["Itinerary", "", "", "Where"])
db.merge_cells("I22:K22")
for k in range(14):
    r = 23 + k
    key = f"P{r}"
    start_key = f"IF(TODAY()<{START},{START},TODAY())"
    db[key] = f'=IFERROR(SMALL({IT("H")},COUNTIF({IT("H")},"<"&{start_key})+{k + 1}),"")'
    row = f'MATCH({key},{IT("H")},0)'
    style(db.cell(r, 9, f'=IF({key}="","",IF(INT({key})=INT(N(P{r - 1})),"",INT({key})))' if k else
                  f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", True, "center")
    style(db.cell(r, 10, f'=IF({key}="","",HOUR({key})&":"&RIGHT("0"&MINUTE({key}),2)&"  "&INDEX({IT("D")},{row}))'), False)
    db.merge_cells(f"J{r}:K{r}")
    style(db.cell(r, 11), False)
    style(db.cell(r, 12, f'=IF({key}="","",INDEX({IT("E")},{row}))'), False)
for c in "OP":
    db.column_dimensions[c].hidden = True

chart = BarChart()
chart.type = "bar"
chart.title = "Budget and what is booked or spent"
chart.height, chart.width = 7, 12
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=11, max_row=18), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=12, max_row=18))
chart.series[0].graphicalProperties.solidFill = "B8D3CF"
chart.series[1].graphicalProperties.solidFill = TEAL
chart.x_axis.scaling.orientation = "maxMin"
chart.y_axis.number_format = "#,##0"
chart.legend.position = "b"
db.add_chart(chart, "B22")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Setup: the trip, how many travel, the first and last day, your home currency and the other currencies."),
    ("2", "Dashboard: your budget for each category."),
    ("3", "Bookings: every flight, hotel, train, tour and table, with its code, its cost and what you have paid."),
    ("4", "Itinerary: what you plan each day. Expenses: what you spend that is not a booking."),
    ("5", "Packing list: type x as things go in the bag. The list has room for your own items."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The budget counts the full cost of each booking as soon as you add it, so you see what is left.",
                          "The example trip shows how it works. Delete it and plan your own.",
                          "Works in Google Sheets and Microsoft Excel, with any currencies."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Itinerary", "Bookings", "Expenses", "Packing list", "Setup", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Travel-Planner.xlsx")
wb.save(out)
print("saved", out, len(itinerary), "plans", len(bookings), "bookings", len(items), "packing items")
