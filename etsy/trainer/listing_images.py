"""Etsy listing photos for the Personal Trainer Client Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Personal-Trainer-Client-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups: the lists down to the last line, the latest sessions with the ones booked ahead.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
s_end, c_end = last(wb["Sessions"], 3005), last(wb["Clients"], 105)
p_end, k_end = last(wb["Packages"], 505), last(wb["Check-ins"], 1505)
business = wb["Settings"]["C4"].value
import datetime as dt  # noqa: E402
now = dt.datetime.combine(dt.date.today(), dt.time())
ahead = min(r for r in range(6, s_end + 1) if wb["Sessions"].cell(r, 2).value >= now)
rows = {"Sessions": list(range(6, ahead - 18)) + list(range(ahead + 5, 3006)),
        "Clients": range(c_end + 1, 106),
        "Packages": list(range(6, p_end - 9)) + list(range(p_end + 1, 506)),
        "Check-ins": list(range(6, k_end - 11)) + list(range(k_end + 1, 1506))}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Report"), hide_rows=rows, hide={"Clients": ["C", "E"]})
report_page = [i for i, t in enumerate(r.titles) if t == business][-1]
pix = r.doc[report_page].get_pixmap(dpi=220)
report = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


def stack(top, lower):
    lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                         Image.LANCZOS)
    both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
    both.paste(top, (0, 0))
    both.paste(lower, (0, top.height + 40))
    return both


img = canvas(dark=True)
title(img, "Personal trainer client tracker", "Session packages, no-shows and late cancels, what each client owes, "
      "and a progress report to print", dark=True, width=2300)
laptop(img, r.shot(business), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-progress-report.jpg"), "A progress report for each client",
       "Pick a client: their measurements and marks from the first check-in to the latest, and their weight over "
       "time. Print it or save it as a PDF.", content(report))

detail(os.path.join(img_dir, "03-sessions.jpg"), "Every session, booked or done",
       "Attended, no-show, late cancel or cancelled. A no-show uses a session, and a late cancel too if you choose.",
       r.shot("Sessions"), [("#D8F0DC", "Attended"), ("#F9C9C4", "No-show"), ("#FFE2B8", "Late cancel")])

detail(os.path.join(img_dir, "04-clients.jpg"), "Each client, and who needs you",
       "Sessions attended and left, the next check-in, what each client owes, and what needs your attention",
       r.shot("Clients"), [("#FFE2B8", "Needs attention"), ("#F9C9C4", "Owes money")])

detail(os.path.join(img_dir, "05-checkins-packages.jpg"), "Check-ins and packages",
       "Measurements and your own three marks at each check-in, and each block of sessions with what was paid",
       stack(r.shot("Check-ins"), r.shot("Packages")))

img = canvas()
y = title(img, "Inside the file", "Six tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page(business)), "Dashboard", "The month and who needs you"),
    (content(r.page("Sessions")), "Sessions", "Booked and done"),
    (content(r.page("Clients")), "Clients", "Sessions left and owed"),
    (content(r.page("Packages")), "Packages", "Blocks of sessions"),
    (content(r.page("Check-ins")), "Check-ins", "Measurements and marks"),
    (content(report), "Report", "To print for a client"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
