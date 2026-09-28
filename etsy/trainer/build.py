"""Personal Trainer Client Tracker (Excel + Google Sheets).

Clients with their session packages, the sessions with no-shows and late cancellations, what each client owes,
check-ins with body measurements and three performance marks of the trainer's choice, a progress report for one
client to print or save as a PDF, and a Dashboard with the month and the clients who need attention. It records
numbers only: it gives no health or training advice. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
L0, L1 = 6, 105        # clients
S0, S1 = 6, 3005       # sessions
P0, P1 = 6, 505        # packages
C0, C1 = 6, 1505       # check-ins
BUSINESS, WUNIT, LUNIT, LOW, WEEKS, LATE = (f"Settings!$C${r}" for r in range(4, 10))
MARK = lambda i: f"Settings!$C${12 + i}"
BETTER = lambda i: f"Settings!$D${12 + i}"
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
CL = lambda c: f"Clients!${c}${L0}:${c}${L1}"
SS = lambda c: f"Sessions!${c}${S0}:${c}${S1}"
SS1 = lambda c: f"Sessions!${c}$1:${c}${S1}"
PK = lambda c: f"Packages!${c}${P0}:${c}${P1}"
CI = lambda c: f"'Check-ins'!${c}${C0}:${c}${C1}"
CI1 = lambda c: f"'Check-ins'!${c}$1:${c}${C1}"
THIS = "(YEAR(TODAY())*100+MONTH(TODAY()))"
MEASURES = ["Weight", "Body fat %", "Waist", "Hips", "Chest", "Arm", "Thigh"]


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{",".join(chr(34) + m + chr(34) for m in MON)})&" "&YEAR({d})'


# ---------------------------------------------------------------- example data (a made-up studio)
rng = random.Random(44)
D = lambda n: today + dt.timedelta(days=n)
clients = [  # name, contact, started, goal, status
    ("Maya Ortiz", "maya.o@mail.example", D(-240), "Stronger legs and back, run a 10 km in spring", "Active"),
    ("Ben Carter", "ben.c@mail.example", D(-200), "Lose 8 kg before the summer", "Active"),
    ("Priya Nair", "priya.n@mail.example", D(-180), "Back to training after a shoulder injury", "Active"),
    ("Tom Walsh", "tom.w@mail.example", D(-150), "First pull-up", "Active"),
    ("Leah Fischer", "leah.f@mail.example", D(-120), "Strength and energy for work", "Active"),
    ("Omar Haddad", "omar.h@mail.example", D(-100), "Row 1 km under 4 minutes", "Active"),
    ("Grace Kim", "grace.k@mail.example", D(-150), "Tone up for a wedding", "Paused"),
    ("Luis Moreno", "luis.m@mail.example", D(-60), "Build muscle", "Active"),
    ("Hannah Brooks", "hannah.b@mail.example", D(-300), "Stay fit and flexible", "Finished"),
    ("Sam Taylor", "sam.t@mail.example", D(-30), "Get into a routine", "Active"),
]
names = [c[0] for c in clients]
packages = []  # bought on, client, sessions, price, paid, expires, note
sessions = []  # date, time, client, type, status, note
checkins = []  # date, client, weight, fat, waist, hips, chest, arm, thigh, mark 1, 2, 3, note
profile = {  # weight, fat, waist, hips, chest, arm, thigh, squat, push-ups, row; change a week
    "Maya Ortiz": ((68, 27, 76, 98, 92, 29, 57, 50, 12, 4.9), (-0.08, -0.08, -0.12, -0.08, 0, 0.03, -0.05, 0.7, 0.25, -0.012)),
    "Ben Carter": ((96, 29, 102, 108, 110, 35, 62, 70, 15, 5.2), (-0.28, -0.12, -0.3, -0.15, -0.1, 0, -0.1, 0.6, 0.2, -0.012)),
    "Priya Nair": ((61, 25, 70, 95, 88, 27, 54, 35, 5, 5.4), (0, -0.05, -0.05, 0, 0, 0.03, 0, 0.6, 0.2, -0.01)),
    "Tom Walsh": ((78, 20, 84, 97, 100, 32, 56, 80, 25, 4.6), (-0.03, -0.05, -0.05, 0, 0.05, 0.03, 0, 0.5, 0.35, -0.008)),
    "Leah Fischer": ((72, 30, 80, 102, 94, 30, 59, 45, 8, 5.6), (-0.1, -0.1, -0.15, -0.1, -0.05, 0, -0.05, 0.7, 0.25, -0.015)),
    "Omar Haddad": ((84, 18, 86, 99, 104, 34, 58, 100, 30, 4.3), (0, -0.03, 0, 0, 0.05, 0.03, 0.03, 0.6, 0.25, -0.012)),
    "Grace Kim": ((64, 28, 72, 96, 89, 28, 55, 40, 6, 5.3), (-0.12, -0.1, -0.2, -0.15, -0.05, 0, -0.05, 0.5, 0.2, -0.012)),
    "Luis Moreno": ((70, 15, 78, 94, 96, 31, 55, 85, 28, 4.8), (0.1, 0, 0.03, 0, 0.15, 0.08, 0.05, 1.0, 0.3, -0.005)),
    "Hannah Brooks": ((58, 24, 68, 92, 86, 26, 52, 45, 15, 4.9), (0, -0.03, -0.05, 0, 0, 0, 0, 0.4, 0.15, -0.005)),
    "Sam Taylor": ((88, 26, 94, 104, 106, 33, 60, 55, 10, 5.5), (-0.2, -0.1, -0.2, -0.1, -0.05, 0, -0.05, 0.9, 0.35, -0.02)),
}
plans = {"Maya Ortiz": (10, 600), "Ben Carter": (10, 600), "Priya Nair": (5, 325), "Tom Walsh": (10, 600),
         "Leah Fischer": (10, 600), "Omar Haddad": (5, 325), "Grace Kim": (10, 600), "Luis Moreno": (10, 600),
         "Hannah Brooks": (5, 325), "Sam Taylor": (5, 325)}
slots = [dt.time(7, 0), dt.time(8, 0), dt.time(12, 30), dt.time(17, 30), dt.time(18, 30)]
for n, (name, contact, started, goal, status) in enumerate(clients):
    per_week = 2 if n % 3 else 1
    end = today if status == "Active" else started + dt.timedelta(days=90 if status == "Paused" else 200)
    slot = slots[n % len(slots)]
    d, count = started, 0
    while d < end:
        for k in range(per_week):
            day = d + dt.timedelta(days=k * 3)
            if day >= end:
                break
            r = rng.random()
            st = "Attended" if r < 0.86 else ("No-show" if r < 0.91 else ("Late cancel" if r < 0.95 else "Cancelled"))
            sessions.append([day, slot, name, "1 to 1" if n % 4 else "Duo", st, None])
            count += st != "Cancelled"
        d += dt.timedelta(days=7)
    size, price = plans[name]
    target = (2 if name == "Ben Carter" else rng.randint(3, 8)) if status == "Active" else 0
    bought, d = 0, started
    while bought + size <= count + target:
        paid = price
        if name == "Luis Moreno" and bought + 2 * size > count + target:
            paid = price / 2                                          # Luis still owes half of his last block
        packages.append([d, name, size, price, paid, d + dt.timedelta(days=120), None])
        bought += size
        d += dt.timedelta(days=max(14, int(size * 7 / per_week) - 10))
    if bought < count + target:
        extra = count + target - bought
        packages.append([min(d, today - dt.timedelta(days=3)), name, extra, round(price / size * extra), round(
            price / size * extra), min(d, today) + dt.timedelta(days=90), "Top-up"])
    if status == "Active":
        for k in (1, 3, 8):                                           # booked sessions ahead
            sessions.append([D(k + n % 3), slot, name, "1 to 1" if n % 4 else "Duo", None, None])
    base, step = profile[name]
    weeks = max(1, (end - started).days // 7)
    for w in range(0, weeks, 2 if n == 0 else 4):                   # the first client checks in every 2 weeks
        day = started + dt.timedelta(days=7 * w)
        vals = [round(b + s * w + rng.uniform(-0.3, 0.3) * (abs(s) + 0.1), 1) for b, s in zip(base, step)]
        vals[8] = int(round(vals[8]))
        vals[9] = round(vals[9], 2)
        if w and w % (4 if n == 0 else 8):
            vals[3:7] = [None] * 4                                    # hips, chest, arm and thigh now and then
        note = {0: "First check-in"}.get(w)
        checkins.append([day, name] + vals + [note])
sessions.sort(key=lambda s: (s[0], s[1]))
packages.sort(key=lambda p: p[0])
checkins.sort(key=lambda c: c[0])
next(c for c in reversed(checkins) if c[1] == clients[0][0])[12] = "Squat and row up again."
for s in sessions:
    if s[0] >= today:
        s[4] = None

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def check_fill(ws, rng_, first):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}<>""'], fill=fill(BAD)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your business name for the report, the units, and when a client needs attention.",
           [3, 34, 18, 12, 50], rows=18)
for r, lab, value, fmt, note in (
        (4, "Business name", "Northside Strength", None, "At the top of the progress report."),
        (5, "Weight unit", "kg", None, "Only a label: type the numbers in this unit."),
        (6, "Length unit", "cm", None, "For waist, hips, chest, arm and thigh."),
        (7, "Few sessions left, at most", 2, "0", "A client shows Few sessions left from this number down."),
        (8, "Check-in every (weeks)", 4, "0", "When the next check-in is due."),
        (9, "A late cancel uses a session", "Yes", None, "Yes or No. A no-show always uses one.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, r == 4, "center")
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
dropdown(se, '"Yes,No"', "C9")
se["B11"], se["C11"], se["D11"] = "Performance marks", "Name", "Better when"
for c in ("B11", "C11", "D11"):
    se[c].font = font(11, True, TEAL_D)
for i, (name, better) in enumerate((("Squat (kg)", "Higher"), ("Push-ups", "Higher"), ("1 km row (min)", "Lower"))):
    r = 12 + i
    se[f"B{r}"] = f"Mark {i + 1}"
    style(se[f"C{r}"], True, bold=True, align="center")
    se[f"C{r}"] = name
    style(se[f"D{r}"], True, align="center")
    se[f"D{r}"] = better
dropdown(se, '"Higher,Lower"', "D12:D14")
se["E12"] = "Any three tests you repeat, such as a lift, a count or a time."
se["E12"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Clients
cl = wb.create_sheet("Clients")
sheet_base(cl, "Clients", "Each client once. Sessions, packages and check-ins add up from the other tabs.",
           [3, 18, 24, 12, 30, 10, 9, 9, 8, 9, 8, 8, 12, 13, 11, 20], rows=L1 + 2)
header(cl, 5, 2, ["Client", "Phone or email", "Started", "Goal", "Status", "Attended", "This month", "Missed",
                  "Bought", "Used", "Left", "Last session", "Next check-in", "Owed", "Needs"])
for r in range(L0, L1 + 1):
    x = clients[r - L0] if r - L0 < len(clients) else (None,) * 5
    for c, v, fmt in zip(range(2, 7), x, (None, None, DATE, None, None)):
        style(cl.cell(r, c, v), True, fmt, c == 2, "center" if c in (4, 6) else None)
    style(cl.cell(r, 7, f'=IF(B{r}="","",SUMIFS({SS("L")},{SS("D")},B{r}))'), False, "0", align="center")
    style(cl.cell(r, 8, f'=IF(B{r}="","",SUMIFS({SS("L")},{SS("D")},B{r},{SS("J")},{THIS}))'), False, "0",
          align="center")
    style(cl.cell(r, 9, f'=IF(B{r}="","",SUMIFS({SS("M")},{SS("D")},B{r}))'), False, "0", align="center")
    style(cl.cell(r, 10, f'=IF(B{r}="","",SUMIFS({PK("D")},{PK("C")},B{r},{PK("M")},1))'), False, "0",
          align="center")
    style(cl.cell(r, 11, f'=IF(B{r}="","",SUMIFS({SS("K")},{SS("D")},B{r}))'), False, "0", align="center")
    style(cl.cell(r, 12, f'=IF(OR(B{r}="",J{r}=0),"",J{r}-K{r})'), False, "0", True, "center")
    last = f"SUMPRODUCT(MAX(({SS('D')}=B{r})*{SS('L')}*{SS('I')}))"
    style(cl.cell(r, 13, f'=IF(B{r}="","",IF({last}=0,"",{last}))'), False, DATE, align="center")
    style(cl.cell(r, 14, f'=IF(OR(B{r}="",F{r}="Finished"),"",IF(S{r}>0,S{r}+7*N({WEEKS}),'
                         f'IF(ISNUMBER(D{r}),INT(D{r})+7*N({WEEKS}),"")))'), False, DATE, align="center")
    style(cl.cell(r, 15, f'=IF(B{r}="","",SUMIFS({PK("J")},{PK("C")},B{r},{PK("M")},1))'), False, MONEY)
    style(cl.cell(r, 16, f'=IF(B{r}="","",IF(COUNTIF({CL("B")},B{r})>1,"Listed twice",IF(O{r}>0,"Owes money",'
                         f'IF(F{r}="Finished","",IF(F{r}="Paused","",IF(AND(J{r}>0,L{r}<=0),"No sessions left",'
                         f'IF(AND(J{r}>0,L{r}<=N({LOW})),"Few sessions left",IF(AND(Q{r}>0,L{r}>0,Q{r}<TODAY()),'
                         f'"Package expired",IF(AND(Q{r}>0,L{r}>0,Q{r}-TODAY()<=14),"Package ends soon",'
                         f'IF(AND(N{r}<>"",N{r}<=TODAY()),"Check-in due","")))))))))'), False, bold=True,
          align="center")
    cl[f"Q{r}"] = f'=IF(B{r}="",0,SUMPRODUCT(MAX(({PK("C")}=B{r})*{PK("N")})))'                # Q: package ends
    cl[f"R{r}"] = f'=IF(P{r}<>"",ROW(),"")'                                                    # R: needs attention
    cl[f"S{r}"] = f'=IF(B{r}="",0,SUMPRODUCT(MAX(({CI("C")}=B{r})*{CI("P")})))'                # S: last check-in
    active = f'B{r}<>"",F{r}<>"Paused",F{r}<>"Finished"'
    cl[f"T{r}"] = f'=IF(AND({active},J{r}>0,L{r}<=N({LOW})),1,0)'                               # T: few left
    cl[f"U{r}"] = f'=IF(AND({active},N{r}<>"",N{r}<=TODAY()),1,0)'                              # U: check-in due
    cl[f"V{r}"] = f'=IF(AND({active}),1,0)'                                                     # V: active
for c in "QRSTUV":
    cl.column_dimensions[c].hidden = True
cl.conditional_formatting.add(f"P{L0}:P{L1}", FormulaRule(formula=[f'OR(P{L0}="Owes money",P{L0}="No sessions left",'
                                                                   f'P{L0}="Listed twice")'], fill=fill(BAD)))
cl.conditional_formatting.add(f"P{L0}:P{L1}", FormulaRule(formula=[f'P{L0}<>""'], fill=fill(WARN)))
cl.conditional_formatting.add(f"O{L0}:O{L1}", FormulaRule(formula=[f'AND(ISNUMBER(O{L0}),O{L0}>0)'], fill=fill(BAD)))
dropdown(cl, '"Active,Paused,Finished"', f"F{L0}:F{L1}")
cl.freeze_panes = "C6"

# ---------------------------------------------------------------- Sessions
ss = wb.create_sheet("Sessions")
sheet_base(ss, "Sessions", "Each session, booked ahead or done. Leave the status empty until it has happened.",
           [3, 13, 8, 18, 12, 13, 26, 18], rows=S1 + 2)
header(ss, 5, 2, ["Date", "Time", "Client", "Type", "Status", "Note", "Check"])
late_used = f'AND(F{{r}}="Late cancel",{LATE}="Yes")'
for r in range(S0, S1 + 1):
    x = sessions[r - S0] if r - S0 < len(sessions) else (None,) * 6
    for c, v, fmt in zip(range(2, 8), x, (DATE, "hh:mm", None, None, None, None)):
        style(ss.cell(r, c, v), True, fmt, c == 4, "center" if c in (2, 3, 5, 6) else None)
    style(ss.cell(r, 8, f'=IF(AND(B{r}="",D{r}="",F{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                        f'IF(D{r}="","Needs a client",IF(COUNTIF({CL("B")},D{r})=0,"Not on Clients",'
                        f'IF(AND(F{r}="",INT(B{r})<TODAY()),"Needs a status","")))))'), False, align="center")
    ss[f"I{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                           # I: date as a number
    ss[f"J{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                         # J: month
    ss[f"K{r}"] = (f'=IF(AND(H{r}="",OR(F{r}="Attended",F{r}="No-show",'
                   f'{late_used.format(r=r)})),1,0)')                                           # K: uses a session
    ss[f"L{r}"] = f'=IF(AND(H{r}="",F{r}="Attended"),1,0)'                                     # L: attended
    ss[f"M{r}"] = f'=IF(AND(H{r}="",OR(F{r}="No-show",F{r}="Late cancel")),1,0)'                # M: missed
    ss[f"N{r}"] = (f'=IF(AND(H{r}="",F{r}="",D{r}<>"",I{r}>=TODAY()),'
                   f'ROUND((I{r}+IF(ISNUMBER(C{r}),MOD(C{r},1),0))*1440,0)*10000+ROW(),"")')     # N: coming up
for c in "IJKLMN":
    ss.column_dimensions[c].hidden = True
check_fill(ss, f"H{S0}:H{S1}", f"H{S0}")
for text, colour in (("Attended", OK), ("No-show", BAD), ("Late cancel", WARN)):
    ss.conditional_formatting.add(f"F{S0}:F{S1}", FormulaRule(formula=[f'F{S0}="{text}"'], fill=fill(colour)))
dropdown(ss, f"={CL('B')}", f"D{S0}:D{S1}", strict=False)
dropdown(ss, '"1 to 1,Duo,Small group,Online"', f"E{S0}:E{S1}", strict=False)
dropdown(ss, '"Attended,No-show,Late cancel,Cancelled"', f"F{S0}:F{S1}")
ss.freeze_panes = "E6"

# ---------------------------------------------------------------- Packages
pk = wb.create_sheet("Packages")
sheet_base(pk, "Packages", "Each block of sessions a client buys, with what they paid so far.",
           [3, 13, 18, 10, 11, 11, 13, 20, 12, 11, 22], rows=P1 + 2)
header(pk, 5, 2, ["Bought on", "Client", "Sessions", "Price", "Paid", "Use by", "Note", "Per session", "Owed",
                  "Check"])
for r in range(P0, P1 + 1):
    x = packages[r - P0] if r - P0 < len(packages) else (None,) * 7
    for c, v, fmt in zip(range(2, 9), x, (DATE, None, "0", MONEY, MONEY, DATE, None)):
        style(pk.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 7) else None)
    style(pk.cell(r, 9, f'=IF(OR(N(D{r})<=0,NOT(ISNUMBER(E{r}))),"",E{r}/D{r})'), False, MONEY)
    style(pk.cell(r, 10, f'=IF(OR(C{r}="",NOT(ISNUMBER(E{r}))),"",MAX(0,E{r}-N(F{r})))'), False, MONEY)
    style(pk.cell(r, 11, f'=IF(AND(C{r}="",D{r}="",E{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                         f'IF(C{r}="","Needs a client",IF(COUNTIF({CL("B")},C{r})=0,"Not on Clients",'
                         f'IF(NOT(ISNUMBER(D{r})),"Needs the sessions",IF(AND(E{r}<>"",NOT(ISNUMBER(E{r}))),'
                         f'"Price not a number",IF(AND(ISNUMBER(E{r}),N(F{r})>E{r}),"Paid more than the price",'
                         f'"")))))))'), False, align="center")
    pk[f"L{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                         # L: month
    pk[f"M{r}"] = f'=IF(AND(K{r}="",C{r}<>""),1,0)'                                            # M: counts
    pk[f"N{r}"] = f"=IF(AND(M{r}=1,ISNUMBER(G{r})),INT(G{r}),0)"                               # N: use by
for c in "LMN":
    pk.column_dimensions[c].hidden = True
check_fill(pk, f"K{P0}:K{P1}", f"K{P0}")
pk.conditional_formatting.add(f"J{P0}:J{P1}", FormulaRule(formula=[f'AND(ISNUMBER(J{P0}),J{P0}>0)'], fill=fill(BAD)))
dropdown(pk, f"={CL('B')}", f"C{P0}:C{P1}", strict=False)
pk.freeze_panes = "D6"

# ---------------------------------------------------------------- Check-ins
ci = wb.create_sheet("Check-ins")
sheet_base(ci, "Check-ins", "Measurements and marks, as often as you check. Leave empty what you did not measure.",
           [3, 13, 18, 9, 9, 9, 9, 9, 9, 9, 11, 11, 11, 26, 18], rows=C1 + 2)
header(ci, 5, 2, ["Date", "Client"] + [""] * 10 + ["Note", "Check"])
units = [WUNIT, None, LUNIT, LUNIT, LUNIT, LUNIT, LUNIT]
for i, (m, u) in enumerate(zip(MEASURES, units)):
    ci.cell(5, 4 + i).value = f'="{m}"&IF({u}="",""," ("&{u}&")")' if u else m
for i in range(3):
    ci.cell(5, 11 + i).value = f'=IF({MARK(i)}="","Mark {i + 1}",{MARK(i)})'
for r in range(C0, C1 + 1):
    x = checkins[r - C0] if r - C0 < len(checkins) else (None,) * 13
    for c, v in zip(range(2, 15), x):
        fmt = DATE if c == 2 else ("0.#" if 4 <= c <= 12 else ("0.##" if c == 13 else None))
        style(ci.cell(r, c, v), True, fmt, c == 3, "center" if c != 14 else None)
    style(ci.cell(r, 15, f'=IF(AND(B{r}="",C{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                         f'IF(C{r}="","Needs a client",IF(COUNTIF({CL("B")},C{r})=0,"Not on Clients",'
                         f'IF(COUNT(D{r}:M{r})=0,"Nothing measured","")))))'), False, align="center")
    ci[f"P{r}"] = f'=IF(AND(O{r}="",ISNUMBER(B{r})),INT(B{r}),0)'                              # P: date as a number
    ci[f"Q{r}"] = f'=IF(AND(P{r}>0,C{r}=Report!$C$5),P{r}*10000+ROW(),"")'                     # Q: on the report
    for i in range(10):                                                                        # R to AA: each value
        col = chr(68 + i)
        ci.cell(r, 18 + i).value = f'=IF(AND(Q{r}<>"",ISNUMBER({col}{r})),Q{r},"")'
    for i in range(3):                                                                         # AB to AD: marks
        col = chr(75 + i)
        ci.cell(r, 28 + i).value = f'=IF(AND(Q{r}<>"",ISNUMBER({col}{r})),{col}{r},"")'
for c in ["P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD"]:
    ci.column_dimensions[c].hidden = True
check_fill(ci, f"O{C0}:O{C1}", f"O{C0}")
dropdown(ci, f"={CL('B')}", f"C{C0}:C{C1}", strict=False)
ci.freeze_panes = "D6"

# ---------------------------------------------------------------- Report (to print)
rp = wb.create_sheet("Report")
sheet_base(rp, "", "", [3, 24, 14, 14, 14], rows=50)
rp["B2"] = f'=IF({BUSINESS}="","Progress report",{BUSINESS})'
rp["B2"].font = font(18, True, TEAL_D)
rp["B3"] = '="Progress report, "&' + long_date("TODAY()")
rp["B3"].font = font(10, color=MUTED, italic=True)
rp["B5"] = "Client"
rp["B5"].font = font(12, True)
style(rp["C5"], True, bold=True)
rp["C5"] = clients[0][0]
rp.merge_cells("C5:E5")
dropdown(rp, f"={CL('B')}", "C5", strict=False)
row_of = f"MATCH($C$5,{CL('B')},0)"
cv = lambda c: f"INDEX({CL(c)},{row_of})"
keys = f"'Check-ins'!$Q${C0}:$Q${C1}"
info = [("Goal", f'=IFERROR(IF({cv("E")}="","",{cv("E")}),"")', None),
        ("Started", f'=IFERROR(IF({cv("D")}="","",{cv("D")}),"")', DATE),
        ("Sessions attended", f'=IFERROR({cv("G")},"")', "0"),
        ("Attendance", f'=IFERROR(IF({cv("G")}+{cv("I")}=0,"",{cv("G")}/({cv("G")}+{cv("I")})),"")', "0%"),
        ("Sessions left", f'=IFERROR(IF({cv("J")}=0,"",{cv("L")}),"")', "0"),
        ("Check-ins", f'=IF(COUNT({keys})=0,"None yet",IF(COUNT({keys})=1,"1, on "&'
                      f'{long_date(f"INT(SMALL({keys},1)/10000)")},COUNT({keys})&", from "&'
                      f'{long_date(f"INT(SMALL({keys},1)/10000)")}&" to "&{long_date(f"INT(LARGE({keys},1)/10000)")}))',
         None)]
for k, (lab, f, fmt) in enumerate(info):
    r = 7 + k
    rp[f"B{r}"] = lab
    rp[f"B{r}"].font = font(11, True)
    rp[f"C{r}"] = f
    rp[f"C{r}"].font = font(11)
    rp[f"C{r}"].alignment = Alignment(horizontal="left")
    if fmt:
        rp[f"C{r}"].number_format = fmt
    rp.merge_cells(f"C{r}:E{r}")
header(rp, 14, 2, ["Measure", "First", "Latest", "Change"])
for i in range(7):
    r = 15 + i
    kcol = f"'Check-ins'!${chr(82 + i)}${C0}:${chr(82 + i)}${C1}"
    vcol = CI1(chr(68 + i))
    style(rp.cell(r, 2, f"='Check-ins'!{chr(68 + i)}5"), False, bold=True)
    style(rp.cell(r, 3, f'=IF(COUNT({kcol})=0,"",INDEX({vcol},MOD(SMALL({kcol},1),10000)))'), False, "0.#",
          align="center")
    style(rp.cell(r, 4, f'=IF(COUNT({kcol})=0,"",INDEX({vcol},MOD(LARGE({kcol},1),10000)))'), False, "0.#",
          align="center")
    style(rp.cell(r, 5, f'=IF(COUNT({kcol})<2,"",D{r}-C{r})'), False, '+0.#;-0.#;0', True, "center")
header(rp, 23, 2, ["Mark", "First", "Best", "Latest"])
for i in range(3):
    r = 24 + i
    kcol = f"'Check-ins'!${chr(89 + i)}${C0}:${chr(89 + i)}${C1}" if i < 2 else f"'Check-ins'!$AA${C0}:$AA${C1}"
    vcol = CI1(chr(75 + i))
    best = f"'Check-ins'!${['AB', 'AC', 'AD'][i]}${C0}:${['AB', 'AC', 'AD'][i]}${C1}"
    style(rp.cell(r, 2, f"='Check-ins'!{chr(75 + i)}5"), False, bold=True)
    style(rp.cell(r, 3, f'=IF(COUNT({kcol})=0,"",INDEX({vcol},MOD(SMALL({kcol},1),10000)))'), False, "0.##",
          align="center")
    style(rp.cell(r, 4, f'=IF(COUNT({best})=0,"",IF({BETTER(i)}="Lower",MIN({best}),MAX({best})))'), False,
          "0.##", True, "center")
    style(rp.cell(r, 5, f'=IF(COUNT({kcol})=0,"",INDEX({vcol},MOD(LARGE({kcol},1),10000)))'), False, "0.##",
          align="center")
rp["B28"] = f'=IF(COUNT({CI("R")})<2,"","Weight at each check-in")'
rp["B28"].font = font(11, True, TEAL_D)
for k in range(12):                                                         # rows 30 to 41: the chart's data
    r = 30 + k
    n = f"COUNT('Check-ins'!$R${C0}:$R${C1})"
    pos = f"(MAX(0,{n}-12)+{k + 1})"
    key = f"SMALL('Check-ins'!$R${C0}:$R${C1},{pos})"
    rp[f"B{r}"] = (f'=IF({pos}>{n},"",DAY(INT({key}/10000))&" "&CHOOSE(MONTH(INT({key}/10000)),'
                   + ",".join(chr(34) + m + chr(34) for m in MON) + "))")
    rp[f"C{r}"] = f'=IF({pos}>{n},NA(),INDEX({CI1("D")},MOD({key},10000)))'
    for c in ("B", "C"):
        rp[f"{c}{r}"].font = Font(color="FFFFFF", size=8)
chart = LineChart()
chart.title = None
chart.height, chart.width = 6.2, 13.5
chart.add_data(Reference(rp, min_col=3, min_row=30, max_row=41), titles_from_data=False)
chart.set_categories(Reference(rp, min_col=2, min_row=30, max_row=41))
chart.series[0].graphicalProperties.line.solidFill = TEAL
chart.series[0].graphicalProperties.line.width = 28000
chart.series[0].marker.symbol = "circle"
chart.series[0].marker.graphicalProperties.solidFill = TEAL
chart.series[0].marker.graphicalProperties.line.solidFill = TEAL
chart.series[0].smooth = False
chart.display_blanks = "gap"
chart.legend = None
rp.add_chart(chart, "B29")
rp["B43"] = "Latest note"
rp["B43"].font = font(11, True)
latest_note = f"INDEX({CI1('N')},MOD(LARGE({keys},1),10000))"
rp["B44"] = f'=IF(COUNT({keys})=0,"",IF({latest_note}="","",{latest_note}))'
rp["B44"].alignment = Alignment(wrap_text=True, vertical="top")
rp["B44"].font = font(11)
rp.merge_cells("B44:E46")
rp["B48"] = "These are the numbers recorded at each check-in. They are not medical advice."
rp["B48"].font = font(8, color=MUTED, italic=True)
rp.page_setup.orientation = "portrait"
rp.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "", "", [3, 20, 13, 13, 13, 3, 20, 13, 13, 3], rows=50)
db["B2"] = f'=IF({BUSINESS}="","Clients and sessions",{BUSINESS})'
db["B2"].font = font(22, True, TEAL_D)
db["B3"] = ('="This month, "&CHOOSE(MONTH(TODAY()),"January","February","March","April","May","June","July",'
            '"August","September","October","November","December")&": the sessions, what was sold, and who needs '
            'you."')
db["B3"].font = font(11, color=MUTED, italic=True)
tile(db, "B", 4, "Active clients", f"=SUM({CL('V')})", "0")
tile(db, "C", 4, "Sessions", f"=SUMIFS({SS('L')},{SS('J')},{THIS})", "0")
tile(db, "D", 4, "Missed", f"=SUMIFS({SS('M')},{SS('J')},{THIS})", "0", "B86E1C")
tile(db, "E", 4, "Sold", f"=SUMIFS({PK('E')},{PK('L')},{THIS},{PK('M')},1)", MONEY)
tile(db, "G", 4, "Owed to you", f"=SUM({CL('O')})", MONEY, "B23B2E")
tile(db, "H", 4, "Few sessions left", f"=SUM({CL('T')})", "0", "B86E1C")
tile(db, "I", 4, "Check-ins due", f"=SUM({CL('U')})", "0")
header(db, 7, 2, ["Needs attention", "Why", "", "Left"])
db.merge_cells("C7:D7")
for n in range(10):
    r = 8 + n
    db[f"L{r}"] = f"=IFERROR(SMALL({CL('R')},{n + 1}),\"\")"                                    # L: client row
    c = lambda col: f"INDEX(Clients!${col}$1:${col}${L1},L{r})"
    style(db.cell(r, 2, f'=IF(L{r}="","",{c("B")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(L{r}="","",{c("P")})'), False, align="center")
    db.merge_cells(f"C{r}:D{r}")
    style(db.cell(r, 5, f'=IF(L{r}="","",IF({c("J")}=0,"",{c("L")}))'), False, "0", align="center")
header(db, 7, 7, ["Coming up", "Date", "Time"])
for n in range(10):
    r = 8 + n
    db[f"M{r}"] = f"=IFERROR(SMALL({SS('N')},{n + 1}),\"\")"                                    # M: session row
    g = lambda col: f"INDEX({SS1(col)},MOD(M{r},10000))"
    style(db.cell(r, 7, f'=IF(M{r}="","",{g("D")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(M{r}="","",{g("B")})'), False, "ddd d mmm", align="center")
    style(db.cell(r, 9, f'=IF(M{r}="","",IF({g("C")}="","",{g("C")}))'), False, "hh:mm", align="center")
header(db, 19, 2, ["Month", "Attended", "Missed", "Sold"])
for k in range(12):
    r = 20 + k
    y = "YEAR(TODAY())"
    key = f"({y}*100+{k + 1})"
    style(db.cell(r, 2, MON[k]), False, bold=True)
    style(db.cell(r, 3, f"=SUMIFS({SS('L')},{SS('J')},{key})"), False, "0", align="center")
    style(db.cell(r, 4, f"=SUMIFS({SS('M')},{SS('J')},{key})"), False, "0", align="center")
    style(db.cell(r, 5, f"=SUMIFS({PK('E')},{PK('L')},{key},{PK('M')},1)"), False, MONEY)
for c in "LM":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 18)) + list(range(20, 32)):
    db.row_dimensions[r].height = 18
db.conditional_formatting.add("C8:C17", FormulaRule(formula=['OR(C8="Owes money",C8="No sessions left")'],
                                                    fill=fill(BAD)))
db.conditional_formatting.add("C8:C17", FormulaRule(formula=['C8<>""'], fill=fill(WARN)))
chart = BarChart()
chart.type = "col"
chart.title = "Sessions each month"
chart.grouping = "stacked"
chart.overlap = 100
chart.height, chart.width = 7.4, 11.5
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=19, max_row=31), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=20, max_row=31))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E39A4B"
chart.legend.position = "b"
db.add_chart(chart, "G19")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your business name, your units, and your three performance marks."),
    ("2", "Clients: each client once, with the day they started and their goal."),
    ("3", "Packages: each block of sessions a client buys, with what they have paid."),
    ("4", "Sessions: book them ahead, then pick Attended, No-show, Late cancel or Cancelled."),
    ("5", "Check-ins and Report: record the measurements, then print the report for the client."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example studio and its clients are made up. To remove them, select the yellow cells "
                          "and press Delete, rather than deleting whole rows.",
                          "The file records numbers. It gives no health or training advice.",
                          "Works in Google Sheets and Microsoft Excel, in any currency and any unit."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Sessions", "Clients", "Packages", "Check-ins", "Report", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Personal-Trainer-Client-Tracker.xlsx")
wb.save(out)
print("saved", out, len(clients), "clients", len(sessions), "sessions", len(packages), "packages",
      len(checkins), "check-ins")
