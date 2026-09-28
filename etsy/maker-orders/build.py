"""Handmade Order Tracker for makers who sell what they make (Excel + Google Sheets).

Orders with due dates, deposits and balances; what each product costs to make
from its materials; the materials to buy for open orders; fees per sales
channel; profit per order and per month; customers; and a month calendar of
what is due. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
D = lambda y, m, d: dt.date(y, m, d)
O0, O1 = 6, 1005      # orders
P0, P1 = 6, 105       # products
M0, M1 = 6, 205       # materials
R0, R1 = 6, 505       # product materials (bill of materials)
C0, C1 = 6, 505       # customers
H0, H1 = 6, 15        # sales channels
OR = lambda col: f"Orders!${col}${O0}:${col}${O1}"
PR = lambda col: f"Products!${col}${P0}:${col}${P1}"
MT = lambda col: f"Materials!${col}${M0}:${col}${M1}"
BM = lambda col: f"'Product materials'!${col}${R0}:${col}${R1}"
CU = lambda col: f"Customers!${col}${C0}:${col}${C1}"
CH = lambda col: f"Settings!${col}${H0}:${col}${H1}"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
STATUSES = ["New", "Making", "Ready", "Sent", "Picked up", "Cancelled"]
NET = '#,##0.00;[Red]-#,##0.00'
GREY = "EEF1F1"

# ---------------------------------------------------------------- example data
channels = [("Etsy", 0.095, 0.45), ("Own website", 0.035, 0.30), ("Instagram", 0.03, 0), ("Craft market", 0.026, 0),
            ("Friends and family", 0, 0)]
materials = [
    # material, unit, pack size, pack price, in stock
    ("Soy wax", "lb", 10, 32.00, 45), ("Cotton wick", "each", 100, 9.50, 200), ("Jar 8 oz", "each", 12, 18.00, 24),
    ("Jar 16 oz", "each", 12, 26.00, 4), ("Fragrance oil", "oz", 16, 21.00, 48), ("Label", "each", 250, 25.00, 250),
    ("Clamshell", "each", 50, 12.00, 30), ("Soap base", "lb", 5, 22.50, 6), ("Kraft wrap", "each", 100, 8.00, 70),
    ("Gift box", "each", 25, 30.00, 3), ("Tissue paper", "sheet", 100, 7.00, 40), ("Tin 4 oz", "each", 24, 19.00, 12),
]
products = [
    # product, price, minutes to make, [(material, amount per item)]
    ("Soy candle 8 oz", 22, 20, [("Soy wax", 0.5), ("Cotton wick", 1), ("Jar 8 oz", 1), ("Fragrance oil", 0.5),
                                 ("Label", 1)]),
    ("Soy candle 16 oz", 34, 25, [("Soy wax", 1), ("Cotton wick", 2), ("Jar 16 oz", 1), ("Fragrance oil", 1),
                                  ("Label", 1)]),
    ("Wax melts, 6 pack", 9, 10, [("Soy wax", 0.4), ("Fragrance oil", 0.4), ("Clamshell", 1), ("Label", 1)]),
    ("Soap bar", 8, 10, [("Soap base", 0.3), ("Fragrance oil", 0.1), ("Label", 1), ("Kraft wrap", 1)]),
    ("Gift box", 45, 35, [("Soy wax", 0.5), ("Cotton wick", 1), ("Jar 8 oz", 1), ("Fragrance oil", 0.7), ("Soap base", 0.6),
                          ("Label", 3), ("Gift box", 1), ("Tissue paper", 2)]),
    ("Wedding favours, 10 tins", 60, 60, [("Soy wax", 1.5), ("Cotton wick", 10), ("Tin 4 oz", 10), ("Fragrance oil", 1.5),
                                          ("Label", 10)]),
]
price = {p[0]: p[1] for p in products}
rng = random.Random(21)
FIRST = ["Ava", "Mia", "Lily", "Ella", "Sophie", "Ruby", "Alice", "Nora", "Hannah", "Jade", "Leah", "Megan", "Holly",
         "Amy", "Beth", "Clara", "Dina", "Eva", "Fiona", "Gina", "Helen", "Iris", "Julia", "Kara", "Lucy", "Nina",
         "Paula", "Rita", "Sara", "Tina", "Adam", "Chris", "Dan", "Eli", "Finn", "George", "Harry", "Ian", "Jon", "Kyle"]
LAST = ["Adams", "Baker", "Bell", "Brooks", "Cole", "Cook", "Cox", "Davies", "Ellis", "Fox", "Gray", "Hall", "Hill",
        "Hughes", "James", "Kelly", "King", "Lewis", "Mason", "Mills", "Morgan", "Murphy", "Owens", "Price", "Reed"]
customers = []
while len(customers) < 70:
    name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
    if name not in (c[0] for c in customers):
        customers.append((name, f"{name.split()[0].lower()}.{name.split()[1][0].lower()}@example.com",
                          rng.choice(["Portland", "Seattle", "Denver", "Austin", "Boston", "Chicago", "Local"]), None))
customers[0] = (customers[0][0], customers[0][1], "Local", "Wedding in October, 80 guests")
orders = []
day = D(Y, 1, 5)
while day <= today:
    for _ in range(rng.choice([0, 1, 1, 2])):
        product = rng.choices([p[0] for p in products], [35, 20, 15, 15, 10, 5])[0]
        channel = rng.choices([c[0] for c in channels], [45, 15, 15, 15, 10])[0]
        who = rng.choice(customers[:25]) if rng.random() < 0.5 else rng.choice(customers)
        qty = rng.choice([1, 1, 1, 2, 2, 3, 4]) if "favours" not in product else rng.choice([4, 6, 8])
        extras = 5 * qty if rng.random() < 0.1 and "Soy candle" in product else None
        ship = 6.95 if channel in ("Etsy", "Own website") else None
        due = day + dt.timedelta(days=rng.randint(4, 12) if "favours" not in product else rng.randint(25, 45))
        total = qty * price[product] + (extras or 0) + (ship or 0)
        if due < today:
            status = "Cancelled" if rng.random() < 0.02 else ("Sent" if ship else "Picked up")
            paid = total
        else:
            status = rng.choice(["New", "Making", "Ready"]) if due <= today + dt.timedelta(days=5) else "New"
            paid = total if channel in ("Etsy", "Own website") else (round(total / 2, 2) if total > 40 else None)
        postage = round(rng.uniform(5.2, 7.8), 2) if ship and status not in ("New", "Cancelled") else None
        note = "Custom label" if extras else None
        orders.append([day, who[0], product, qty, extras, ship, paid, due, status, channel, postage, note])
    day += dt.timedelta(days=1)
# A wedding order with a deposit, a late order and an unpaid pickup.
orders.append([today - dt.timedelta(days=20), customers[0][0], "Wedding favours, 10 tins", 8, 40, None, 250,
               today + dt.timedelta(days=16), "Making", "Friends and family", None, "Ivory labels with the date"])
late = next(o for o in reversed(orders) if o[8] in ("Sent", "Picked up") and o[7] < today - dt.timedelta(days=2))
late[8], late[10] = "Making", None
late[11] = "Waiting for the jars"
unpaid = next(o for o in reversed(orders) if o[8] == "Picked up" and o is not late)
unpaid[6], unpaid[11] = None, "Pays on Friday"
orders.sort(key=lambda o: o[0])

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def due_colours(ws, rng_, first):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT({first},4)="Late"'], fill=fill(BAD),
                                                    font=Font(name=F, bold=True, color="9B1C1C")))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'OR({first}="Due today",LEFT({first},6)="Due in")'],
                                                    fill=fill(WARN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="Done"'], fill=fill(OK)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="Cancelled"'], fill=fill(GREY),
                                                    font=Font(name=F, color="9AA7A9")))


# ---------------------------------------------------------------- Orders
od = wb.active
od.title = "Orders"
sheet_base(od, "Orders", "One line per order. The total, costs, fees and profit fill in from your products and channels.",
           [3, 12, 18, 22, 6, 9, 10, 11, 11, 11, 12, 11, 15, 11, 10, 10, 11, 16, 24], rows=O1 + 2)
header(od, 5, 2, ["Order date", "Customer", "Product", "Qty", "Extras", "Shipping charged", "Total", "Paid so far",
                  "Balance", "Due date", "Status", "Channel", "Materials", "Postage", "Fees", "Profit", "Due", "Note"])
for r in range(O0, O1 + 1):
    o = orders[r - O0] if r - O0 < len(orders) else (None,) * 12
    d_, who, product, qty, extras, ship, paid, due, status, channel, postage, note = o
    style(od.cell(r, 2, d_), True, DATE)
    style(od.cell(r, 3, who), True, bold=True)
    style(od.cell(r, 4, product), True)
    style(od.cell(r, 5, qty), True, "0", align="center")
    style(od.cell(r, 6, extras), True, MONEY)
    style(od.cell(r, 7, ship), True, MONEY)
    qty_ = f"IF(N(E{r})=0,1,E{r})"
    style(od.cell(r, 8, f'=IF(OR(B{r}="",D{r}=""),"",{qty_}*IFERROR(INDEX({PR("C")},MATCH(D{r},{PR("B")},0)),0)'
                        f'+N(F{r})+N(G{r}))'), False, MONEY, True)
    style(od.cell(r, 9, paid), True, MONEY)
    style(od.cell(r, 10, f'=IF(OR(H{r}="",L{r}="Cancelled"),"",IF(H{r}-N(I{r})>0.005,H{r}-N(I{r}),""))'), False, MONEY, True)
    style(od.cell(r, 11, due), True, DATE)
    style(od.cell(r, 12, status), True, align="center")
    style(od.cell(r, 13, channel), True)
    style(od.cell(r, 14, f'=IF(OR(H{r}="",L{r}="Cancelled"),"",{qty_}*IFERROR(INDEX({PR("E")},MATCH(D{r},{PR("B")},0)),0))'),
          False, MONEY)
    style(od.cell(r, 15, postage), True, MONEY)
    fee = f'IFERROR(H{r}*INDEX({CH("C")},MATCH(M{r},{CH("B")},0))+INDEX({CH("D")},MATCH(M{r},{CH("B")},0)),0)'
    style(od.cell(r, 16, f'=IF(OR(H{r}="",L{r}="Cancelled"),"",ROUND({fee},2))'), False, MONEY)
    style(od.cell(r, 17, f'=IF(OR(H{r}="",L{r}="Cancelled"),"",H{r}-N(N{r})-N(O{r})-N(P{r}))'), False, NET, True)
    days = lambda a, b: f'{a}&IF({a}=1," day"," days")'
    style(od.cell(r, 18, f'=IF(OR(B{r}="",K{r}=""),"",IF(OR(L{r}="Sent",L{r}="Picked up"),"Done",IF(L{r}="Cancelled",'
                         f'"Cancelled",IF(K{r}<TODAY(),"Late by "&{days(f"(TODAY()-K{r})", 0)},IF(K{r}=TODAY(),"Due today",'
                         f'"Due in "&{days(f"(K{r}-TODAY())", 0)})))))'), False, align="center")
    style(od.cell(r, 19, note), True)
    # Hidden helpers.
    open_ = f'AND(B{r}<>"",K{r}<>"",L{r}<>"Sent",L{r}<>"Picked up",L{r}<>"Cancelled")'
    od.cell(r, 20, f'=IF({open_},K{r}+ROW()/10000000,"")')                                        # T open by due date
    od.cell(r, 21, f'=IF(OR(B{r}="",C{r}=""),"",B{r}+ROW()/10000000)')                              # U order key
    od.cell(r, 22, f'=IF(U{r}="","",IF(COUNTIFS({OR("C")},C{r},{OR("U")},">"&U{r})=0,1,0))')        # V last order
    od.cell(r, 23, f'=IF({open_},{qty_},0)')                                                       # W open qty
    od.cell(r, 24, f'=IF(OR(B{r}="",D{r}=""),"",{qty_})')                                           # X quantity
for c in range(20, 25):
    od.column_dimensions[L(c)].hidden = True
dropdown(od, f"={CU('B')}", f"C{O0}:C{O1}")
dropdown(od, f"={PR('B')}", f"D{O0}:D{O1}")
dropdown(od, f'"{",".join(STATUSES)}"', f"L{O0}:L{O1}")
dropdown(od, f"={CH('B')}", f"M{O0}:M{O1}")
due_colours(od, f"R{O0}:R{O1}", f"$R{O0}")
od.conditional_formatting.add(f"J{O0}:J{O1}", FormulaRule(formula=[f"N($J{O0})>0"], fill=fill(WARN)))
od.freeze_panes = "D6"

# ---------------------------------------------------------------- Products
pd_ = wb.create_sheet("Products")
sheet_base(pd_, "Products", "What you sell and its price. The cost comes from the Product materials tab.",
           [3, 26, 10, 10, 12, 12, 10, 12, 13, 12], rows=P1 + 2)
header(pd_, 5, 2, ["Product", "Price", "Minutes to make", "Materials cost", "Profit per item", "Margin", "Sold this year",
                   "Sales this year", "In open orders"])
for r in range(P0, P1 + 1):
    p = products[r - P0][:3] if r - P0 < len(products) else (None,) * 3
    for c, (v, fmt) in enumerate(zip(p, (None, MONEY, "0")), start=2):
        style(pd_.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c == 4 else None)
    style(pd_.cell(r, 5, f'=IF(B{r}="","",SUMIFS({BM("F")},{BM("B")},B{r}))'), False, MONEY)
    style(pd_.cell(r, 6, f'=IF(B{r}="","",N(C{r})-E{r})'), False, NET, True)
    style(pd_.cell(r, 7, f'=IF(OR(B{r}="",N(C{r})=0),"",F{r}/C{r})'), False, "0%", align="center")
    done = f'{OR("B")},">="&DATE(YEAR(TODAY()),1,1)'
    style(pd_.cell(r, 8, f'=IF(B{r}="","",SUMIFS({OR("X")},{OR("D")},B{r},{OR("L")},"Sent",{done})'
                         f'+SUMIFS({OR("X")},{OR("D")},B{r},{OR("L")},"Picked up",{done}))'), False, "0", align="center")
    style(pd_.cell(r, 9, f'=IF(B{r}="","",SUMIFS({OR("H")},{OR("D")},B{r},{OR("L")},"<>Cancelled",{done}))'), False, MONEY)
    style(pd_.cell(r, 10, f'=IF(B{r}="","",SUMIFS({OR("W")},{OR("D")},B{r}))'), False, "0", align="center")
pd_.conditional_formatting.add(f"G{P0}:G{P1}", FormulaRule(formula=[f'AND(G{P0}<>"",G{P0}<0.5)'], fill=fill(WARN)))
pd_[f"B{P1 + 2}"] = "Margin is the profit per item before fees and postage, as a share of the price."
pd_[f"B{P1 + 2}"].font = font(10, color=MUTED, italic=True)
pd_.freeze_panes = "C6"

# ---------------------------------------------------------------- Product materials
bm = wb.create_sheet("Product materials")
sheet_base(bm, "Product materials", "One line per material in each product, with the amount one item uses.",
           [3, 26, 20, 10, 8, 11, 13], rows=R1 + 2)
header(bm, 5, 2, ["Product", "Material", "Amount per item", "Unit", "Cost per item", "For open orders"])
lines = [(p, m, a) for p, _, _, items in products for m, a in items]
for r in range(R0, R1 + 1):
    ln = lines[r - R0] if r - R0 < len(lines) else (None,) * 3
    style(bm.cell(r, 2, ln[0]), True)
    style(bm.cell(r, 3, ln[1]), True)
    style(bm.cell(r, 4, ln[2]), True, align="center")
    mrow = f"MATCH(C{r},{MT('B')},0)"
    style(bm.cell(r, 5, f'=IF(C{r}="","",IFERROR(INDEX({MT("C")},{mrow})&"","Add it to Materials"))'), False, align="center")
    style(bm.cell(r, 6, f'=IF(OR(C{r}="",E{r}="Add it to Materials"),"",N(D{r})*INDEX({MT("F")},{mrow}))'), False, MONEY)
    style(bm.cell(r, 7, f'=IF(OR(B{r}="",C{r}=""),"",N(D{r})*SUMIFS({OR("W")},{OR("D")},B{r}))'), False, "#,##0.##",
          align="center")
dropdown(bm, f"={PR('B')}", f"B{R0}:B{R1}")
dropdown(bm, f"={MT('B')}", f"C{R0}:C{R1}")
bm.conditional_formatting.add(f"B{R0}:B{R1}", FormulaRule(formula=[f'AND($B{R0}<>"",$B{R0}=$B{R0 - 1})'],
                                                          font=Font(name=F, color="A9B5B7")))
bm.conditional_formatting.add(f"E{R0}:E{R1}", FormulaRule(formula=[f'$E{R0}="Add it to Materials"'], fill=fill(WARN)))
bm.freeze_panes = "C6"

# ---------------------------------------------------------------- Materials
mt = wb.create_sheet("Materials")
sheet_base(mt, "Materials", "What you buy to make your products, what you have, and what the open orders still need.",
           [3, 20, 8, 10, 11, 11, 10, 13, 11, 24], rows=M1 + 2)
header(mt, 5, 2, ["Material", "Unit", "Pack size", "Pack price", "Cost per unit", "In stock", "For open orders",
                  "Short by", "Note"])
for r in range(M0, M1 + 1):
    m = materials[r - M0] if r - M0 < len(materials) else (None,) * 5
    for c, (v, fmt) in zip((2, 3, 4, 5, 7), zip(m, (None, None, "#,##0.##", MONEY, "#,##0.##"))):
        style(mt.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (3, 4, 7) else None)
    style(mt.cell(r, 6, f'=IF(B{r}="","",N(E{r})/IF(N(D{r})>0,D{r},1))'), False, "#,##0.000")
    style(mt.cell(r, 8, f'=IF(B{r}="","",SUMIFS({BM("G")},{BM("C")},B{r}))'), False, "#,##0.##", align="center")
    style(mt.cell(r, 9, f'=IF(B{r}="","",IF(H{r}-N(G{r})>0,H{r}-N(G{r}),""))'), False, "#,##0.##", True, "center")
    style(mt.cell(r, 10), True)
mt.conditional_formatting.add(f"I{M0}:I{M1}", FormulaRule(formula=[f"N($I{M0})>0"], fill=fill(BAD),
                                                          font=Font(name=F, bold=True, color="9B1C1C")))
mt.freeze_panes = "C6"

# ---------------------------------------------------------------- Customers
cu = wb.create_sheet("Customers")
sheet_base(cu, "Customers", "One line per customer. Orders, spending and what they owe come from the Orders tab.",
           [3, 20, 28, 14, 28, 8, 12, 12, 11], rows=C1 + 2)
header(cu, 5, 2, ["Customer", "Email or phone", "City", "Notes", "Orders", "Spent", "Last order", "Owes"])
for r in range(C0, C1 + 1):
    c_ = customers[r - C0] if r - C0 < len(customers) else (None,) * 4
    for c, v in enumerate(c_, start=2):
        style(cu.cell(r, c, v), True, bold=c == 2)
    style(cu.cell(r, 6, f'=IF(B{r}="","",COUNTIFS({OR("C")},B{r},{OR("L")},"<>Cancelled"))'), False, "0", align="center")
    style(cu.cell(r, 7, f'=IF(B{r}="","",SUMIFS({OR("H")},{OR("C")},B{r},{OR("L")},"<>Cancelled"))'), False, MONEY)
    last = f'SUMIFS({OR("U")},{OR("C")},B{r},{OR("V")},1)'
    style(cu.cell(r, 8, f'=IF(B{r}="","",IF({last}=0,"",INT({last})))'), False, DATE)
    owes = f'SUMIFS({OR("J")},{OR("C")},B{r})'
    style(cu.cell(r, 9, f'=IF(B{r}="","",IF({owes}=0,"",{owes}))'), False, MONEY, True)
cu.conditional_formatting.add(f"I{C0}:I{C1}", FormulaRule(formula=[f"N($I{C0})>0"], fill=fill(WARN)))
cu.freeze_panes = "C6"

# ---------------------------------------------------------------- Settings
se = wb.create_sheet("Settings")
sheet_base(se, "Settings", "Where you sell, and what each place takes from an order. Change the fees to match your own.",
           [3, 24, 12, 14, 4, 60], rows=H1 + 6)
header(se, 5, 2, ["Sales channel", "Fee (share)", "Fixed fee per order"])
for r in range(H0, H1 + 1):
    ch = channels[r - H0] if r - H0 < len(channels) else (None,) * 3
    style(se.cell(r, 2, ch[0]), True, bold=True)
    style(se.cell(r, 3, ch[1]), True, "0.0%", align="center")
    style(se.cell(r, 4, ch[2]), True, MONEY, align="center")
se["F6"] = "The fee share is taken from the order total, then the fixed fee is added."
se["F7"] = "Example: an Etsy order of 30.00 pays 9.5% of 30.00 plus 0.45, so 3.30."
se["F8"] = "The example fees are rough. Check what your own platforms and card readers charge."
for a in ("F6", "F7", "F8"):
    se[a].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Calendar
ca = wb.create_sheet("Calendar")
sheet_base(ca, "Due dates", "What is due each day of the month you pick.", [3] + [17] * 7, rows=32)
ca["B4"] = "Month"
ca["B4"].font = font(11, True)
style(ca["C4"], True, "0", True, "center")
ca["C4"] = "=MONTH(TODAY())"
ca["D4"] = "Year"
ca["D4"].font = font(11, True)
ca["D4"].alignment = Alignment(horizontal="right")
style(ca["E4"], True, "0", True, "center")
ca["E4"] = "=YEAR(TODAY())"
ca["F4"] = "Shows this month until you type another month and year."
ca["F4"].font = font(10, color=MUTED, italic=True)
header(ca, 6, 2, DAYS)
first = "DATE($E$4,$C$4,1)"
for w in range(6):
    for k in range(7):
        rd, rw = 7 + 2 * w, 8 + 2 * w
        day_ = f"({first}-WEEKDAY({first},3)+{w * 7 + k})"
        c = ca.cell(rd, 2 + k, f'=IF(MONTH({day_})=$C$4,DAY({day_}),"")')
        c.font = font(9, True, MUTED)
        c.alignment = Alignment(horizontal="right")
        c.border = box
        due_n = f'COUNTIFS({OR("K")},{day_},{OR("L")},"<>Cancelled")'
        open_n = f'COUNTIFS({OR("T")},">="&{day_},{OR("T")},"<"&({day_}+1))'
        c = ca.cell(rw, 2 + k, f'=IF(OR(MONTH({day_})<>$C$4,{due_n}=0),"",{due_n}&IF({due_n}=1," order"," orders")'
                               f'&IF({open_n}=0,""," ("&{open_n}&" open)"))')
        c.font = font(10, True, TEAL_D)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
        ca.cell(rw, 11 + k, f'=IF(MONTH({day_})<>$C$4,"",IF({open_n}>0,IF({day_}<TODAY(),"late","open"),'
                            f'IF({due_n}>0,"done","")))')                                         # K..Q state
    ca.row_dimensions[8 + 2 * w].height = 32
for c in range(11, 18):
    ca.column_dimensions[L(c)].hidden = True
for w in range(6):
    rng_ = f"B{7 + 2 * w}:H{8 + 2 * w}"
    ca.conditional_formatting.add(rng_, FormulaRule(formula=[f'K${8 + 2 * w}="late"'], fill=fill(BAD)))
    ca.conditional_formatting.add(rng_, FormulaRule(formula=[f'K${8 + 2 * w}="open"'], fill=fill(WARN)))
    ca.conditional_formatting.add(rng_, FormulaRule(formula=[f'K${8 + 2 * w}="done"'], fill=fill(OK)))
key = 20
ca.cell(key, 2, "Colours:").font = font(10, True, MUTED)
for k, (label, colour) in enumerate([("All sent or picked up", OK), ("Still to make or send", WARN), ("Late", BAD)]):
    c = ca.cell(key, 3 + k, label)
    c.fill = fill(colour)
    c.font = font(10)
    c.border = box

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Order dashboard", "What is due, what you are owed, and what each month earned.",
           [3, 13, 20, 22, 7, 12, 14, 3, 20, 10, 12, 12, 3], rows=60)
ms, me = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "EOMONTH(TODAY(),0)"
this_month = f'{OR("B")},">="&{ms},{OR("B")},"<="&{me},{OR("L")},"<>Cancelled"'
tile(db, "B", 5, "Open orders", f'=COUNT({OR("T")})', "0")
tile(db, "C", 5, "Due in 7 days", f'=COUNTIFS({OR("T")},">="&TODAY(),{OR("T")},"<"&(TODAY()+8))', "0", "B07A00")
tile(db, "D", 5, "Late", f'=COUNTIFS({OR("T")},"<"&TODAY())', "0", "B4541F")
tile(db, "F", 5, "Owed to you", f'=SUM({OR("J")})', MONEY, "B4541F")
db.merge_cells("F5:G5")
db.merge_cells("F6:G6")
tile(db, "I", 5, "Sales this month", f'=SUMIFS({OR("H")},{this_month})', MONEY)
tile(db, "J", 5, "Profit this month", f'=SUMIFS({OR("Q")},{this_month})', NET)
db.merge_cells("J5:K5")
db.merge_cells("J6:K6")
tile(db, "L", 5, "Orders this month", f'=COUNTIFS({this_month})', "0")
for a in ("G5", "G6", "K5", "K6"):
    db[a].border = box

header(db, 8, 2, ["Due", "Customer", "Product", "Qty", "Status", "When"])
for k in range(10):
    r = 9 + k
    key_ = f"O{r}"
    db[key_] = f'=IFERROR(SMALL({OR("T")},{k + 1}),"")'
    row = f'MATCH({key_},{OR("T")},0)'
    pick = lambda col: f'IF({key_}="","",INDEX({OR(col)},{row}))'
    style(db.cell(r, 2, f'=IF({key_}="","",INT({key_}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 3, f"={pick('C')}"), False, bold=True)
    style(db.cell(r, 4, f"={pick('D')}"), False)
    style(db.cell(r, 5, f"={pick('E')}"), False, "0", align="center")
    style(db.cell(r, 6, f"={pick('L')}"), False, align="center")
    style(db.cell(r, 7, f"={pick('R')}"), False, align="center")
db.column_dimensions["O"].hidden = True
due_colours(db, "G9:G18", "$G9")
db["B19"] = "Open orders, the soonest due first. Late orders come first."
db["B19"].font = font(9, color=MUTED, italic=True)

late_n = f'COUNTIFS({OR("T")},"<"&TODAY())'
owe_n = f'COUNTIFS({OR("J")},">0")'
short_n = f'COUNTIFS({MT("I")},">0")'
unpriced = f'COUNTIFS({BM("E")},"Add it to Materials")'
notes = [
    (f'=IF({late_n}=0,"No order is late",IF({late_n}=1,"1 order is",{late_n}&" orders are")&" past its due date")', BAD,
     "No order"),
    (f'=IF({owe_n}=0,"Every order is paid",IF({owe_n}=1,"1 order has",{owe_n}&" orders have")&" a balance to collect: "'
     f'&FIXED(SUM({OR("J")}),2))', WARN, "Every"),
    (f'=IF({short_n}=0,"You have the materials for every open order",IF({short_n}=1,"1 material is",'
     f'{short_n}&" materials are")&" short for the open orders")', BAD, "You have"),
    (f'=IF({unpriced}=0,"Every product has its cost",IF({unpriced}=1,"1 material line is",{unpriced}&" material lines are")'
     f'&" missing from Materials")', WARN, "Every"),
]
header(db, 8, 9, ["Needs attention", "", "", ""])
db.merge_cells("I8:L8")
for k, (formula, colour, calm) in enumerate(notes):
    r = 9 + k
    db.merge_cells(f"I{r}:L{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "IJKL":
        db[f"{col}{r}"].border = box
    db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))

header(db, 14, 9, ["Channel this year", "Orders", "Sales", "Profit"])
db.row_dimensions[14].height = None
for k in range(5):
    r = 15 + k
    ch = f"Settings!B{H0 + k}"
    year = f'{OR("M")},I{r},{OR("L")},"<>Cancelled",{OR("B")},">="&DATE(YEAR(TODAY()),1,1)'
    style(db.cell(r, 9, f'=IF({ch}="","",{ch})'), False, bold=True)
    style(db.cell(r, 10, f'=IF(I{r}="","",COUNTIFS({year}))'), False, "0", align="center")
    style(db.cell(r, 11, f'=IF(I{r}="","",SUMIFS({OR("H")},{year}))'), False, MONEY)
    style(db.cell(r, 12, f'=IF(I{r}="","",SUMIFS({OR("Q")},{year}))'), False, NET)

db["B21"] = "Year"
db["B21"].font = font(11, True)
style(db["C21"], True, "0", True, "center")
db["C21"] = Y
header(db, 23, 2, ["Month", "Orders", "Sales", "", "Materials", "Fees and postage", "Profit"])
db.merge_cells("D23:E23")
for m in range(12):
    r = 24 + m
    a, b = f"DATE($C$21,{m + 1},1)", f"EOMONTH(DATE($C$21,{m + 1},1),0)"
    within = f'{OR("B")},">="&{a},{OR("B")},"<="&{b},{OR("L")},"<>Cancelled"'
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f"=COUNTIFS({within})"), False, "0", align="center")
    style(db.cell(r, 4, f'=SUMIFS({OR("H")},{within})'), False, MONEY)
    style(db.cell(r, 5), False)
    db.merge_cells(f"D{r}:E{r}")
    style(db.cell(r, 6, f'=SUMIFS({OR("N")},{within})'), False, MONEY)
    style(db.cell(r, 7, f'=SUMIFS({OR("O")},{within})+SUMIFS({OR("P")},{within})'), False, MONEY)
    style(db.cell(r, 8, f"=D{r}-F{r}-G{r}"), False, NET, True)
db.column_dimensions["H"].width = 12
style(db.cell(36, 2, "Year"), False, bold=True, align="center")
for col, fmt in (("C", "#,##0"), ("D", MONEY), ("F", MONEY), ("G", MONEY), ("H", NET)):
    style(db[f"{col}36"], False, fmt, True, "center" if col == "C" else None)
    db[f"{col}36"] = f"=SUM({col}24:{col}35)"
style(db["E36"], False)
db.merge_cells("D36:E36")
db.conditional_formatting.add("B24:H35", FormulaRule(formula=["AND($C$21=YEAR(TODAY()),ROW()-23=MONTH(TODAY()))"],
                                                     fill=fill("E6F2F2")))
chart = BarChart()
chart.type = "col"
chart.title = "Sales and profit by month"
chart.height, chart.width = 8, 14
chart.add_data(Reference(db, min_col=4, min_row=23, max_row=35), titles_from_data=True)
chart.add_data(Reference(db, min_col=8, min_row=23, max_row=35), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=24, max_row=35))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E0A21B"
chart.y_axis.number_format = "#,##0"
chart.legend.position = "b"
db.add_chart(chart, "I22")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: where you sell, and the fee each place takes from an order."),
    ("2", "Materials: what you buy, with the pack size and price, and how much you have in stock."),
    ("3", "Products and Product materials: your prices, and how much of each material one item uses."),
    ("4", "Orders: one line per order, with the customer, product, quantity, due date and what was paid."),
    ("5", "Dashboard and Due dates: what to make next, what is late, what you are owed and what you earned."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["An order is open until its status is Sent, Picked up or Cancelled.",
                          "Profit is the total minus the materials, the postage you paid and the channel fees.",
                          "The example orders show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency. It does not give tax advice."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Handmade-Order-Tracker.xlsx")
wb.save(out)
print("saved", out, len(orders), "orders", len(customers), "customers", len(lines), "material lines")
