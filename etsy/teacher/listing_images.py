"""Etsy listing photos for the Teacher Planner and Gradebook (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Teacher-Planner-Gradebook.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in, and the gradebook shows only the assessments in use.
wb = load_workbook(xlsx)
last = lambda ws, end, first=6: max(r for r in range(first, end + 1) if ws.cell(r, 2).value not in (None, ""))
c1 = wb["Class 1"]
used = max(c for c in range(8, 28) if c1.cell(5, c).value not in (None, ""))
rows = {"Lessons": range(last(wb["Lessons"], 505) + 1, 506), "Absences": range(last(wb["Absences"], 405) + 1, 406),
        "Class 1": range(last(c1, 50, 11) + 1, 51)}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Week"))
close = Renders(xlsx, tmp, one_page=("Dashboard", "Week"), hide_rows=rows,
                hide={"Class 1": [get_column_letter(c) for c in range(used + 1, 28)], "Timetable": []})


def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


img = canvas(dark=True)
title(img, "Teacher planner and gradebook", "Lesson plans by the week, a gradebook for each class with weighted "
      "averages, and the students to watch", dark=True, width=2300)
laptop(img, r.shot("Teaching week"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-week-plan.jpg"), "A week plan to print",
       "The class and topic of every lesson from your timetable, with the homework set that week",
       r.shot("Week of"), [("#FFE2B8", "A lesson with no plan yet")])

detail(os.path.join(img_dir, "03-gradebook.jpg"), "A gradebook for each class",
       "Weighted averages, grades from your own scale, missing work and absences for every student",
       close.shot("7B Science"), [("#F9C9C4", "Below the pass mark"), ("#FFE2B8", "Missing work")])

detail(os.path.join(img_dir, "04-lessons.jpg"), "Every lesson by date and period",
       "Type the date, the period and the topic. The class comes from your timetable.",
       close.shot("Lessons"))

top, lower = close.shot("Timetable"), close.shot("Absences")
lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                     Image.LANCZOS)
both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
both.paste(top, (0, 0))
both.paste(lower, (0, top.height + 40))
detail(os.path.join(img_dir, "05-timetable.jpg"), "Your timetable and the absences",
       "Your usual week once, and only the students who were absent or late", both)

img = canvas()
y = title(img, "Inside the file", "Dashboard, week plan, lessons, timetable, six class gradebooks, absences and "
          "settings. You type in the yellow cells.")
grid(img, [
    (content(r.page("Teaching week")), "Dashboard", "Today, your classes, who to watch"),
    (content(r.page("Week of")), "Week plan", "Ready to print"),
    (content(close.page("7B Science")), "Gradebook", "One tab for each class"),
    (content(close.page("Lessons")), "Lessons", "By date and period"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
