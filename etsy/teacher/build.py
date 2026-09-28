"""Teacher Planner and Gradebook (Excel + Google Sheets).

A timetable of the week; lessons planned by date and period, with the class filled in from the timetable; a week
plan to print; a gradebook for each class with weighted averages, grades from your own scale and missing work; an
absence log; and a Dashboard with today's lessons, the classes and the students to watch. Only functions both apps
have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MUTED, OK, TEAL, TEAL_D, WARN, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
NC = 6                 # classes
NP = 10                # periods a day
S0, S1 = 11, 50        # students on each class tab
NA = 20                # assessments on each class tab
A0 = 8                 # column of the first assessment (H)
L0, L1 = 6, 505        # lessons
X0, X1 = 6, 405        # absences
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri"]
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
PASS, WEEK = "Settings!$C$4", "Settings!$C$6"
CLASS = [f"Settings!$C${9 + k}" for k in range(NC)]
MINS, LETTERS = "Settings!$B$17:$B$26", "Settings!$C$17:$C$26"
TT = "Timetable!$E$6:$I$15"
LS = lambda c: f"Lessons!${c}${L0}:${c}${L1}"
LS1 = lambda c: f"Lessons!${c}$1:${c}${L1}"
XS = lambda c: f"Absences!${c}${X0}:${c}${X1}"
TAB = [f"Class {k + 1}" for k in range(NC)]
QT = lambda k: f"'{TAB[k]}'"
AL, AR = col(A0), col(A0 + NA - 1)       # H .. AA


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def clock(t):
    return f'HOUR({t})&":"&RIGHT("0"&MINUTE({t}),2)'


D = lambda n: today + dt.timedelta(days=n)
monday = today - dt.timedelta(days=today.weekday())
rng = random.Random(37)

# ---------------------------------------------------------------- example data (a made-up science teacher)
classes = ["7B Science", "8A Science", "9C Biology", "10D Chemistry"]
starts = [dt.time(8, 45), dt.time(9, 40), dt.time(10, 50), dt.time(11, 45), dt.time(13, 30), dt.time(14, 25)]
timetable = [  # period by day, Mon..Fri
    ["7B Science", None, "9C Biology", "8A Science", None],
    ["7B Science", "10D Chemistry", None, "8A Science", "9C Biology"],
    [None, "8A Science", "7B Science", None, "10D Chemistry"],
    ["9C Biology", "8A Science", "7B Science", "10D Chemistry", None],
    ["10D Chemistry", None, "9C Biology", "7B Science", "8A Science"],
    [None, "9C Biology", "10D Chemistry", None, "7B Science"],
]
first_names = ["Amara", "Ben", "Chloe", "Daniel", "Ella", "Farah", "George", "Hana", "Isaac", "Jade", "Kofi", "Lena",
               "Mateo", "Nadia", "Oscar", "Priya", "Quinn", "Rosa", "Sami", "Tara", "Umar", "Vera", "Wes", "Yara",
               "Zane", "Aiden", "Bella", "Caleb", "Dina", "Eli", "Freya", "Gabe", "Holly", "Ivan", "Jonah", "Kira",
               "Liam", "Maya", "Noah", "Olive", "Pavel", "Ruby", "Seth", "Talia", "Vikram", "Willow", "Xavi", "Zoe"]
initials = "ABCDEFGHJKLMNOPRSTW"
sizes = [26, 24, 22, 18]
students = []
for k, n in enumerate(sizes):
    pool = rng.sample(first_names, n)
    students.append([f"{name} {rng.choice(initials)}." for name in pool])
assessments = [  # name, days ago, out of, weight
    [("Lab safety quiz", 24, 10, 1), ("Cells homework", 17, 20, 1), ("Microscope practical", 12, 25, 1),
     ("Unit 1 test", 5, 50, 2), ("Cell model", 1, 20, 1)],
    [("Particles quiz", 23, 10, 1), ("States of matter", 16, 20, 1), ("Diffusion practical", 10, 25, 1),
     ("Unit 1 test", 4, 60, 2)],
    [("Enzymes homework", 22, 20, 1), ("Food tests practical", 15, 30, 1), ("Enzymes test", 8, 50, 2),
     ("Digestion essay", 2, 25, 1)],
    [("Atoms quiz", 21, 10, 1), ("Bonding homework", 14, 20, 1), ("Titration practical", 9, 40, 1),
     ("Unit 1 test", 3, 80, 2), ("Rates homework", 0, 20, 1)],
]
scores = []
for k, names in enumerate(students):
    grid = []
    for s in names:
        skill = rng.uniform(0.58, 0.97)
        row = []
        for a, (_, ago, out, w) in enumerate(assessments[k]):
            if ago <= 1:
                row.append(None)                              # not marked yet
                continue
            r = rng.random()
            if r < 0.03:
                row.append("M")
            elif r < 0.04:
                row.append("Ex")
            else:
                row.append(max(0, min(out, round(out * min(1, rng.gauss(skill, 0.07))))))
        grid.append(row)
    scores.append(grid)
absences = []
for _ in range(22):
    k = rng.randrange(len(classes))
    day = monday - dt.timedelta(days=rng.choice([1, 2, 3, 4, 7, 8, 9, 10, 11, 14, 15, 16, 17, 18]) - 1)
    while day.weekday() > 4:
        day -= dt.timedelta(days=1)
    kind = rng.choices(["Absent", "Late", "Excused"], [7, 3, 1])[0]
    absences.append((day, classes[k], rng.choice(students[k]), kind, "Doctor's note" if kind == "Excused" else None))
absences += [(monday, "7B Science", students[0][3], "Late", None), (monday, "9C Biology", students[2][5], "Absent", None)]
absences.sort(key=lambda x: x[0])
topics = {
    "7B Science": ["Lab safety", "Using a microscope", "Plant and animal cells", "Specialised cells", "Unit 1 review",
                   "Cells test", "Unit 2: tissues and organs", "Organ systems", "Diffusion demo", "Organs quiz"],
    "8A Science": ["Particle model", "States of matter", "Changes of state", "Diffusion practical", "Gas pressure",
                   "Unit 1 test", "Pure substances", "Mixtures", "Filtration practical", "Chromatography"],
    "9C Biology": ["Enzymes", "Enzyme practical", "Food tests", "Digestive system", "Enzymes test", "Absorption",
                   "Essay planning", "Heart and blood", "Blood vessels", "Heart dissection"],
    "10D Chemistry": ["Atomic structure", "Isotopes", "Ionic bonding", "Covalent bonding", "Titration",
                      "Unit 1 test", "Rates of reaction", "Rates practical", "Catalysts", "Rates graphs"],
}
homework = {"Plant and animal cells": "Cells worksheet", "Changes of state": "Page 34, questions 1 to 6",
            "Food tests": "Write up the results", "Covalent bonding": "Dot and cross diagrams",
            "Organ systems": "Label the organ systems", "Mixtures": "Revise for the quiz",
            "Absorption": "Plan the essay", "Rates of reaction": "Rates homework sheet"}
aims = {"Using a microscope": "Name the parts and focus on a slide", "Plant and animal cells": "Compare the two",
        "Enzymes": "Explain the lock and key model", "Titration": "Carry out a titration safely",
        "Organ systems": "Link organs to their systems", "Mixtures": "Tell mixtures from compounds",
        "Heart and blood": "Trace the path of blood", "Rates of reaction": "Name four factors"}
lessons = []
count = {c: 0 for c in classes}
for w in (-1, 0):
    for d in range(5):
        day = monday + dt.timedelta(days=7 * w + d)
        for p in range(len(starts)):
            c = timetable[p][d]
            if not c:
                continue
            if w == 0 and (d, p) in ((3, 3), (4, 0), (4, 5), (3, 4)):
                continue                                       # not planned yet
            t = topics[c][count[c] % len(topics[c])]
            count[c] += 1
            done = "x" if day < today else None
            lessons.append([day, p + 1, None, t, aims.get(t), homework.get(t), done])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The pass mark, the week the plan shows, your classes and your grade scale.",
           [3, 30, 22, 4, 60], rows=30)
for r, label, value, fmt, typed, note in (
        (4, "Pass mark (percent)", 60, "0", True, "Students below it show in red."),
        (5, "Show another week", None, "ddd d mmm yyyy", True,
         "Leave it empty for the week we are in. Type any day to see and print its week."),
        (6, "The week starts on", '=IF(C5="",TODAY(),C5)-WEEKDAY(IF(C5="",TODAY(),C5),3)', "ddd d mmm yyyy", False,
         "The Monday of that week.")):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], typed, fmt, True, "center")
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
se["B8"] = "Classes"
se["B8"].font = font(12, True, TEAL_D)
se["E8"] = "One for each of the tabs Class 1 to Class 6. The timetable and the lessons use these names."
se["E8"].font = font(10, color=MUTED, italic=True)
for k in range(NC):
    se.cell(9 + k, 2, f"Class {k + 1}").font = font(10, color=MUTED)
    style(se.cell(9 + k, 3, classes[k] if k < len(classes) else None), True, bold=True)
se["B16"] = "Grade scale: from (percent)"
se["C16"] = "Grade"
for c in ("B16", "C16"):
    se[c].font = font(11, True, TEAL_D)
se["E16"] = "A grade applies from its percent up to the next one. Change them to your school's scale."
se["E16"].font = font(10, color=MUTED, italic=True)
scale = [(90, "A"), (80, "B"), (70, "C"), (60, "D"), (0, "F")]
for k in range(10):
    x = scale[k] if k < len(scale) else (None, None)
    style(se.cell(17 + k, 2, x[0]), True, "0", align="center")
    style(se.cell(17 + k, 3, x[1]), True, bold=True, align="center")

# ---------------------------------------------------------------- Timetable
tt = wb.create_sheet("Timetable")
sheet_base(tt, "Timetable", "Your usual week: the class in each period. Leave free periods empty.",
           [3, 8, 10, 10] + [18] * 5, rows=18)
header(tt, 5, 2, ["Period", "Starts", "Ends"] + ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
for p in range(NP):
    r = 6 + p
    tt.cell(r, 2, p + 1).font = font(12, True, TEAL_D)
    tt.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
    tt.cell(r, 2).border = box
    s_ = starts[p] if p < len(starts) else None
    e_ = (dt.datetime.combine(today, s_) + dt.timedelta(minutes=55)).time() if s_ else None
    style(tt.cell(r, 3, s_), True, "h:mm", align="center")
    style(tt.cell(r, 4, e_), True, "h:mm", align="center")
    for d in range(5):
        v = timetable[p][d] if p < len(timetable) else None
        style(tt.cell(r, 5 + d, v), True, bold=True, align="center")
    tt.row_dimensions[r].height = 22
dropdown(tt, "=Settings!$C$9:$C$14", "E6:I15")

# ---------------------------------------------------------------- Class tabs (the gradebook)
for k in range(NC):
    ws = wb.create_sheet(TAB[k])
    sheet_base(ws, "", "Scores as numbers. M for missing work (counts as 0), Ex for excused (left out), empty for "
               "not marked yet.", [3, 22, 9, 7, 8, 8, 7] + [6.5] * NA, rows=S1 + 2)
    ws["B2"] = f'=IF({CLASS[k]}="","{TAB[k]}",{CLASS[k]})'
    ws["B2"].font = font(22, True, TEAL_D)
    ws["AC2"] = f"={PASS}"                                                               # AC2: pass mark
    for r, label in ((5, "Assessment"), (6, "Date"), (7, "Out of"), (8, "Weight"), (9, "Class average")):
        ws.cell(r, 7, label).font = font(10, True, MUTED)
        ws.cell(r, 7).alignment = Alignment(horizontal="right", vertical="bottom" if r == 5 else "center")
    ws.row_dimensions[5].height = 92
    marks = assessments[k] if k < len(assessments) else []
    for a in range(NA):
        c = A0 + a
        L = col(c)
        x = marks[a] if a < len(marks) else (None, None, None, None)
        name, ago, out, w = x
        cell = style(ws.cell(5, c, name), True, bold=True)
        cell.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=True)
        style(ws.cell(6, c, D(-ago) if ago is not None else None), True, "d mmm", align="center")
        style(ws.cell(7, c, out), True, "0", align="center")
        style(ws.cell(8, c, w), True, "0.##", align="center")
        ws[f"{L}52"] = f'=IF(N({L}7)<=0,0,IF({L}8="",1,N({L}8)))'                       # row 52: weight used
        ws[f"{L}53"] = f"=IF({L}52=0,0,{L}52/{L}7)"                                      # row 53: weight a point
        ws[f"{L}54"] = (f'=IF(AND({L}5<>"",ISNUMBER({L}6)),IF(AND({L}6<=TODAY(),COUNT({L}{S0}:{L}{S1})'
                        f'+COUNTIF({L}{S0}:{L}{S1},"M")=0),1,0),0)')                    # row 54: to mark
        style(ws.cell(9, c, f'=IF(OR(N({L}7)<=0,COUNT({L}{S0}:{L}{S1})=0),"",AVERAGE({L}{S0}:{L}{S1})/{L}7)'),
              False, "0%", align="center").font = font(9, True, TEAL_D)
        ws.cell(10, c, a + 1)
    header(ws, 10, 2, ["Student", "Average", "Grade", "Missing", "Absent", "Late"])
    header(ws, 10, A0, list(range(1, NA + 1)))
    style(ws.cell(9, 3, f'=IF(COUNT(C{S0}:C{S1})=0,"",AVERAGE(C{S0}:C{S1}))'), False, "0.0%",
          align="center").font = font(10, True, TEAL_D)
    names = students[k] if k < len(students) else []
    for r in range(S0, S1 + 1):
        i = r - S0
        style(ws.cell(r, 2, names[i] if i < len(names) else None), True, bold=True)
        for a in range(NA):
            v = scores[k][i][a] if i < len(names) and a < len(marks) else None
            style(ws.cell(r, A0 + a, v), True, align="center")
        row = f"{AL}{r}:{AR}{r}"
        den = (f"SUMPRODUCT(--ISNUMBER({row}),${AL}$52:${AR}$52)"
               f'+SUMPRODUCT(--({row}="M"),${AL}$52:${AR}$52)')
        num = f"SUMPRODUCT({row},${AL}$53:${AR}$53)"
        style(ws.cell(r, 3, f'=IF(B{r}="","",IF(({den})=0,"",ROUND({num}/({den})*100,1)/100))'), False, "0.0%",
              True, "center")
        ws[f"AC{r}"] = (f'=IF(C{r}="","",SUMPRODUCT(MAX(ISNUMBER({MINS})*({MINS}<=ROUND(C{r}*100,1))'
                        f'*{MINS})))')                                                    # AC: grade reached
        style(ws.cell(r, 4, f'=IF(C{r}="","",IFERROR(INDEX({LETTERS},MATCH(AC{r},{MINS},0)),""))'), False,
              bold=True, align="center")
        style(ws.cell(r, 5, f'=IF(B{r}="","",COUNTIF({row},"M"))'), False, "0", align="center")
        for c, kind in ((6, "Absent"), (7, "Late")):
            style(ws.cell(r, c, f'=IF(B{r}="","",COUNTIFS({XS("C")},{CLASS[k]},{XS("D")},B{r},{XS("E")},'
                                f'"{kind}"))'), False, "0", align="center")
        ws[f"AD{r}"] = f'=IF(C{r}="","",ROUND(C{r}*1000,0)*1000+{(k + 1) * 100 + i + 1})'  # AD: to watch, in order
        ws[f"AE{r}"] = f'=IF(C{r}="",0,IF(ROUND(C{r}*100,1)<N($AC$2),1,0))'              # AE: below the pass mark
    for c in ("AC", "AD", "AE"):
        ws.column_dimensions[c].hidden = True
    for r in (52, 53, 54):
        ws.row_dimensions[r].hidden = True
    ws.conditional_formatting.add(f"C{S0}:D{S1}", FormulaRule(
        formula=[f"AND(ISNUMBER($C{S0}),ROUND($C{S0}*100,1)<N($AC$2))"], fill=fill(BAD)))
    ws.conditional_formatting.add(f"E{S0}:E{S1}", FormulaRule(formula=[f"N(E{S0})>0"], fill=fill(WARN)))
    grid = f"{AL}{S0}:{AR}{S1}"
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'{AL}{S0}="M"'], fill=fill(BAD)))
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f"AND(ISNUMBER({AL}{S0}),{AL}{S0}>N({AL}$7))"],
                                                    font=Font(color="C0392B", bold=True)))
    ws.freeze_panes = f"C{S0}"

# ---------------------------------------------------------------- Lessons
ls = wb.create_sheet("Lessons")
sheet_base(ls, "Lessons", "Each lesson by date and period. Leave the class empty and it comes from the timetable.",
           [3, 13, 8, 16, 30, 30, 28, 7, 16], rows=L1 + 2)
header(ls, 5, 2, ["Date", "Period", "Class", "Topic", "Aims", "Homework", "Done", "Class taught"])
for r in range(L0, L1 + 1):
    x = lessons[r - L0] if r - L0 < len(lessons) else (None,) * 7
    for c, v, fmt in zip(range(2, 9), x, (DATE, "0", None, None, None, None, None)):
        style(ls.cell(r, c, v), True, fmt, c == 5, "center" if c in (2, 3, 8) else None)
    for c in (6, 7):
        ls.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
    day = f"WEEKDAY(B{r},2)"
    slot = f"INDEX({TT},C{r},{day})"
    planned = (f'IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),IF(OR({day}>5,C{r}<1,C{r}>{NP}),"",'
               f'IF({slot}="","",{slot})),"")')
    style(ls.cell(r, 9, f'=IF(B{r}="","",IF(D{r}<>"",D{r},{planned}))'), False, bold=True, align="center")
    ls[f"K{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}*100+C{r},"")'              # K: date and period
    ls[f"L{r}"] = (f'=IF(OR(G{r}="",NOT(ISNUMBER(B{r}))),"",IF(AND(B{r}>={WEEK},B{r}<={WEEK}+6),'
                   f'B{r}*1000+ROW(),""))')                                              # L: homework this week
for c in "KL":
    ls.column_dimensions[c].hidden = True
dropdown(ls, "=Settings!$C$9:$C$14", f"D{L0}:D{L1}")
ls.freeze_panes = "C6"

# ---------------------------------------------------------------- Absences
ab = wb.create_sheet("Absences")
sheet_base(ab, "Absences", "Only the students who were absent or late. Everyone else counts as present.",
           [3, 13, 16, 22, 11, 26, 20], rows=X1 + 2)
header(ab, 5, 2, ["Date", "Class", "Student", "Type", "Note", "Check"])
for r in range(X0, X1 + 1):
    x = absences[r - X0] if r - X0 < len(absences) else (None,) * 5
    for c, v, fmt in zip(range(2, 7), x, (DATE, None, None, None, None)):
        style(ab.cell(r, c, v), True, fmt, c == 4, "center" if c in (2, 5) else None)
    found = ",".join(f"COUNTIF({QT(k)}!$B${S0}:$B${S1},D{r})" for k in range(NC))
    style(ab.cell(r, 7, f'=IF(OR(C{r}="",D{r}=""),"",IF(ISNA(MATCH(C{r},Settings!$C$9:$C$14,0)),'
                        f'"Class not on Settings",IF(CHOOSE(MATCH(C{r},Settings!$C$9:$C$14,0),{found})=0,'
                        f'"Not on the class list","")))'), False, align="center")
n = 0
for k in range(NC):
    for r in range(S0, S1 + 1):
        ab[f"J{X0 + n}"] = f'=IF({QT(k)}!B{r}="","",{QT(k)}!B{r})'                     # J: every student
        n += 1
ab.column_dimensions["J"].hidden = True
dropdown(ab, "=Settings!$C$9:$C$14", f"C{X0}:C{X1}")
dropdown(ab, f"=$J${X0}:$J${X0 + n - 1}", f"D{X0}:D{X1}", strict=False)
dropdown(ab, '"Absent,Late,Excused"', f"E{X0}:E{X1}")
ab.conditional_formatting.add(f"G{X0}:G{X1}", FormulaRule(formula=[f'G{X0}<>""'], fill=fill(BAD)))
ab.freeze_panes = "C6"

# ---------------------------------------------------------------- Week (the plan to print)
wk = wb.create_sheet("Week")
sheet_base(wk, "", "", [3, 8, 8] + [24] * 5, rows=30)
wk["B2"] = f'="Week of "&{short(WEEK)}&" to "&{short(f"({WEEK}+4)")}'
wk["B2"].font = font(20, True, TEAL_D)
wk["B3"] = "The class and topic of each lesson. Shaded: a lesson on the timetable with no plan yet."
wk["B3"].font = font(10, color=MUTED, italic=True)
header(wk, 5, 2, ["Period", "Time"] + [f'="{DAYS[d]} "&DAY({WEEK}+{d})' for d in range(5)])
for p in range(NP):
    r = 6 + p
    wk.cell(r, 2, p + 1).font = font(12, True, TEAL_D)
    wk.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
    wk.cell(r, 2).border = box
    style(wk.cell(r, 3, f'=IF(ISNUMBER(Timetable!C{r}),{clock(f"Timetable!C{r}")},"")'), False, align="center")
    for d in range(5):
        key = f"({WEEK}+{d})*100+{p + 1}"
        at = f"MATCH({key},{LS('K')},0)"
        slot = f"INDEX({TT},{p + 1},{d + 1})"
        fixed = f'IF({slot}="","",{slot})'
        wk.cell(r, 13 + d, f"=IFERROR({at},\"\")")                                        # M..Q: lesson row
        m = f"{col(13 + d)}{r}"
        taught = f'IF(ISNUMBER({m}),INDEX({LS("I")},{m}),{fixed})'
        topic = f'IF(ISNUMBER({m}),INDEX({LS("E")},{m}),"")'
        c = style(wk.cell(r, 4 + d, f'=IF({taught}="","",{taught}&IF({topic}="","",CHAR(10)&{topic}))'), False)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
        c.font = font(10)
        wk.cell(r, 18 + d, f'=IF(AND(NOT(ISNUMBER({m})),{fixed}<>""),1,0)')               # R..V: not planned
    wk.row_dimensions[r].height = 36
for c in "MNOPQRSTUV":
    wk.column_dimensions[c].hidden = True
wk.conditional_formatting.add("D6:H15", FormulaRule(formula=["R6=1"], fill=fill(WARN)))
wk["B17"] = "Homework set this week"
wk["B17"].font = font(12, True, TEAL_D)
header(wk, 18, 2, ["Day", "", "Class", "Homework", "", "", ""])
wk.merge_cells("B18:C18")
wk.merge_cells("E18:H18")
for n in range(8):
    r = 19 + n
    wk[f"W{r}"] = f"=IFERROR(SMALL({LS('L')},{n + 1}),\"\")"                               # W: lesson row
    at = lambda c: f"INDEX({LS1(c)},MOD(W{r},1000))"
    style(wk.cell(r, 2, f'=IF(W{r}="","",CHOOSE(WEEKDAY({at("B")},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")'
                        f'&" "&DAY({at("B")}))'), False, align="center")
    style(wk.cell(r, 4, f'=IF(W{r}="","",{at("I")})'), False, bold=True)
    c = style(wk.cell(r, 5, f'=IF(W{r}="","",{at("G")})'), False)
    for cc in (3, 6, 7, 8):
        style(wk.cell(r, cc), False)
    wk.merge_cells(f"B{r}:C{r}")
    wk.merge_cells(f"E{r}:H{r}")
wk.column_dimensions["W"].hidden = True
for r in (18,):
    wk.row_dimensions[r].height = 20
wk.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Teaching week", "Today's lessons, your classes and the students to watch.",
           [3, 9, 9, 18, 30, 3, 20, 12, 12, 12, 3], rows=45)
db["M1"] = f"={PASS}"                                                                        # M1: pass mark
slots = f'SUMPRODUCT(--({TT}<>""))'
tile(db, "B", 4, "Today", f'=IF(WEEKDAY(TODAY(),2)>5,0,SUMPRODUCT(({TT}<>"")*(COLUMN({TT})-4=WEEKDAY(TODAY(),2))))',
     "0")
tile(db, "D", 4, "Planned this week", f'=COUNTIFS({LS("B")},">="&{WEEK},{LS("B")},"<="&({WEEK}+4))&" of "&{slots}',
     "General")
tile(db, "E", 4, "To mark", "=" + "+".join(f"SUM({QT(k)}!{AL}54:{AR}54)" for k in range(NC)), "0")
tile(db, "G", 4, "Students", "=" + "+".join(f"COUNTA({QT(k)}!B{S0}:B{S1})" for k in range(NC)), "0")
tile(db, "H", 4, "Below pass", "=" + "+".join(f"SUM({QT(k)}!AE{S0}:AE{S1})" for k in range(NC)), "0",
     "B23B2E")
tile(db, "I", 4, "Missing work", "=" + "+".join(f"SUM({QT(k)}!E{S0}:E{S1})" for k in range(NC)), "0", "B86E1C")
tile(db, "J", 4, "Absent this week",
     f'=COUNTIFS({XS("B")},">="&{WEEK},{XS("B")},"<="&({WEEK}+6),{XS("E")},"Absent")', "0")
db.merge_cells("B4:C4")
db.merge_cells("B5:C5")

header(db, 7, 2, ["Today", "Time", "Class", "Topic"])
for p in range(NP):
    r = 8 + p
    key = f"TODAY()*100+{p + 1}"
    db[f"N{r}"] = f"=IFERROR(MATCH({key},{LS('K')},0),\"\")"                                 # N: lesson row
    slot = f"INDEX({TT},{p + 1},WEEKDAY(TODAY(),2))"
    fixed = f'IF(WEEKDAY(TODAY(),2)>5,"",IF({slot}="","",{slot}))'
    taught = f'IF(ISNUMBER(N{r}),INDEX({LS("I")},N{r}),{fixed})'
    style(db.cell(r, 2, p + 1), False, "0", True, "center")
    style(db.cell(r, 3, f'=IF(ISNUMBER(Timetable!C{6 + p}),{clock(f"Timetable!C{6 + p}")},"")'), False,
          align="center")
    style(db.cell(r, 4, f'=IF({taught}="","",{taught})'), False, bold=True)
    topic = f"INDEX({LS('E')},N{r})"
    style(db.cell(r, 5, f'=IF(ISNUMBER(N{r}),IF({topic}="","",{topic}),IF({fixed}="","","Not planned yet"))'), False)
db.conditional_formatting.add("E8:E17", FormulaRule(formula=['E8="Not planned yet"'], fill=fill(WARN)))

header(db, 7, 7, ["Class", "Students", "Average", "Below pass"])
for k in range(NC):
    r = 8 + k
    style(db.cell(r, 7, f'=IF({CLASS[k]}="","",{CLASS[k]})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(G{r}="","",COUNTA({QT(k)}!B{S0}:B{S1}))'), False, "0", align="center")
    style(db.cell(r, 9, f'=IF(G{r}="","",{QT(k)}!C9)'), False, "0%", align="center")
    style(db.cell(r, 10, f'=IF(G{r}="","",SUM({QT(k)}!AE{S0}:AE{S1}))'), False, "0", align="center")
db.conditional_formatting.add("J8:J13", FormulaRule(formula=["N(J8)>0"], fill=fill(BAD)))

header(db, 15, 7, ["Students to watch", "Class", "Average", "Absent"])
n = 0
for k in range(NC):
    for r in range(S0, S1 + 1):
        db[f"P{1 + n}"] = f"={QT(k)}!AD{r}"                                                   # P: every student
        n += 1
for m in range(8):
    r = 16 + m
    db[f"O{r}"] = f"=IFERROR(SMALL($P$1:$P${n},{m + 1}),\"\")"                               # O: key
    t = f"INT(MOD(O{r},1000)/100)"
    s = f"MOD(O{r},100)+{S0 - 1}"
    pick = lambda c: "CHOOSE(" + t + "," + ",".join(f"INDEX({QT(k)}!${c}$1:${c}${S1},{s})" for k in range(NC)) + ")"
    style(db.cell(r, 7, f'=IF(O{r}="","",{pick("B")})'), False, bold=True)
    which = f"CHOOSE({t}," + ",".join(CLASS) + ")"
    style(db.cell(r, 8, f'=IF(O{r}="","",IF({which}="","",{which}))'), False, align="center")
    style(db.cell(r, 9, f'=IF(O{r}="","",{pick("C")})'), False, "0.0%", align="center")
    style(db.cell(r, 10, f'=IF(O{r}="","",{pick("F")})'), False, "0", align="center")
db.conditional_formatting.add("G16:J23", FormulaRule(formula=["AND(ISNUMBER($I16),ROUND($I16*100,1)<N($M$1))"],
                                                     fill=fill(BAD)))
for r in range(8, 24):
    db.row_dimensions[r].height = 18
    for c in range(2, 11):
        cell = db.cell(r, c)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="center")
db.row_dimensions[15].height = 20
for c in "MNOP":
    db.column_dimensions[c].hidden = True

chart = BarChart()
chart.type = "bar"
chart.title = "Class averages"
chart.height, chart.width = 6.2, 12.5
chart.add_data(Reference(db, min_col=9, min_row=7, max_row=13), titles_from_data=True)
chart.set_categories(Reference(db, min_col=7, min_row=8, max_row=13))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.x_axis.scaling.orientation = "maxMin"
chart.y_axis.crosses = "max"
chart.y_axis.number_format = "0%"
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 1
db.add_chart(chart, "B19")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your classes, the pass mark and your grade scale."),
    ("2", "Timetable: the class in each period of your usual week, with the start times."),
    ("3", "Lessons: a line for each lesson, by date and period. The class fills in from the timetable."),
    ("4", "Week: the plan for the week, ready to print. Settings shows another week if you type a day."),
    ("5", "Class 1 to Class 6: the students, then each assessment with its date, marks out of and weight."),
    ("6", "Absences: only the students who were absent or late."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example classes and students are made up. Delete them and add your own.",
                          "Averages are weighted: a test with weight 2 counts twice as much as one with weight 1.",
                          "Works in Google Sheets and Microsoft Excel."]):
    st.cell(18 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Week", "Lessons", "Timetable"] + TAB + ["Absences", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Teacher-Planner-Gradebook.xlsx")
wb.save(out)
print("saved", out, sum(sizes), "students", len(lessons), "lessons", len(absences), "absences")
