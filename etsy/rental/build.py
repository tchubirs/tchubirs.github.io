"""Vacation Rental Host Tracker (Excel + Google Sheets).

One line per stay, from any booking platform. Nights are split across months,
so occupancy, nightly rate and revenue per month stay correct when a stay
crosses a month end. Revenue means the payout: what the guest paid plus the
cleaning fee, minus the platform fee. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
D = lambda k: today + dt.timedelta(days=k)
K0, K1 = 6, 505      # booking rows
E0, E1 = 6, 505      # expense rows
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
PLATFORMS = ["Airbnb", "Vrbo", "Booking.com", "Direct", "Other"]
CATEGORIES = ["Cleaning", "Supplies", "Utilities", "Repairs", "Furnishing", "Insurance", "Taxes and fees",
              "Marketing", "Mortgage or rent", "Other"]
PROPS = ["Beach apartment", "Mountain cabin"]
NET = '#,##0.00;[Red]-#,##0.00'
BK = lambda col: f"Bookings!${col}${K0}:${col}${K1}"
EX = lambda col: f"Expenses!${col}${E0}:${col}${E1}"
YR, SEL = "Dashboard!$C$4", "Dashboard!$E$4"
ALL = f'{SEL}="All properties"'

# ---------------------------------------------------------------- example data
rnd = random.Random(11)
FIRST = ["Emma", "Lucas", "Mia", "Noah", "Olivia", "Leo", "Ava", "Hugo", "Chloe", "Liam", "Sofia", "Ethan",
         "Nora", "Adam", "Lily", "Owen", "Zoe", "Max", "Ruby", "Sam"]
guest = lambda: f"{rnd.choice(FIRST)} {rnd.choice('ABCDEGHJKLMNPRSTW')}."


def rate(prop, day):
    m = day.month
    if prop == PROPS[0]:
        return 165 if m in (6, 7, 8) else 125 if m in (5, 9) else 95
    return 140 if m in (12, 1, 2, 3) else 110


def platform():
    p = rnd.choices(PLATFORMS[:4], weights=[55, 20, 15, 10])[0]
    return p, {"Airbnb": 0.03, "Vrbo": 0.05, "Booking.com": 0.15, "Direct": 0.0}[p]


stays = []
for prop in PROPS:
    d = dt.date(Y, 1, 2)
    while True:
        busy = d.month in (6, 7, 8) if prop == PROPS[0] else d.month in (1, 2, 12)
        ci = d + dt.timedelta(days=rnd.randint(0, 3) if busy else rnd.randint(2, 11))
        co = ci + dt.timedelta(days=rnd.randint(2, 7))
        if co > D(-12):
            break
        stays.append((prop, ci, co))
        d = co
# Recent and coming stays, fixed so the example always shows every status.
stays += [(PROPS[0], D(-9), D(-5)), (PROPS[0], D(-2), D(3)), (PROPS[0], D(5), D(9)), (PROPS[0], D(9), D(13)),
          (PROPS[1], D(-6), D(-1)), (PROPS[1], D(2), D(6)), (PROPS[1], D(11), D(15))]
stays.sort(key=lambda s: (s[1], s[0]))
bookings = []
for prop, ci, co in stays:
    nights = (co - ci).days
    total = sum(rate(prop, ci + dt.timedelta(days=k)) for k in range(nights))
    clean = 60 if prop == PROPS[0] else 80
    plat, pct = platform()
    paid = None if co > today else ("Yes" if co < D(-3) else "No")
    bookings.append((prop, guest(), plat, ci, co, rnd.randint(2, 5), total, clean, round((total + clean) * pct, 2), paid))

expenses = []
for m in range(1, today.month + 1):
    first = dt.date(Y, m, 1)
    last = dt.date(Y, m + 1, 1) - dt.timedelta(days=1) if m < 12 else dt.date(Y, 12, 31)
    if last > today:
        last = today
    for prop, per_stay, util in ((PROPS[0], 45, (70, 95)), (PROPS[1], 55, (110, 150))):
        n = sum(1 for b in bookings if b[0] == prop and first <= b[4] <= last)
        if n:
            expenses.append((last, prop, "Cleaning", n * per_stay, f"Cleaner, {n} stays"))
        expenses.append((first + dt.timedelta(days=9), prop, "Utilities", rnd.randint(*util), "Power, water and internet"))
        if m % 2 == 0:
            expenses.append((first + dt.timedelta(days=14), prop, "Supplies", rnd.randint(40, 90), "Coffee, soap, paper"))
expenses += [(dt.date(Y, 1, 20), "Shared", "Marketing", 250, "New listing photos"),
             (dt.date(Y, 2, 3), "Shared", "Insurance", 620, "Insurance for both homes"),
             (dt.date(Y, 3, 12), PROPS[1], "Repairs", 140, "Chimney sweep"),
             (dt.date(Y, 5, 8), PROPS[0], "Furnishing", 320, "New sofa bed"),
             (dt.date(Y, 7, 21), PROPS[0], "Repairs", 180, "Dishwasher repair"),
             (dt.date(Y, 7, 31), "Shared", "Taxes and fees", 410, "Tourist tax, first half")]
expenses.sort(key=lambda e: e[0])

wb = Workbook()

# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your properties and expense categories. The lists in the other tabs follow them.",
           [3, 30, 4, 60], rows=30)
header(se, 5, 2, ["Properties (up to 6)"])
for k in range(6):
    style(se.cell(6 + k, 2, PROPS[k] if k < len(PROPS) else None), True, bold=True)
header(se, 14, 2, ["Expense categories"])
for k, cat in enumerate(CATEGORIES):
    style(se.cell(15 + k, 2, cat), True)
se["D6"] = "Rename a property here and its bookings and expenses need the same new name."
se["D6"].font = font(10, color=MUTED, italic=True)
# Hidden lists for the dropdowns: "All properties" or "Shared" followed by the property names.
se["F6"], se["G6"] = "All properties", "Shared"
for k in range(6):
    se.cell(7 + k, 6, f'=IF(B{6 + k}="","",B{6 + k})')
    se.cell(7 + k, 7, f'=IF(B{6 + k}="","",B{6 + k})')
se.column_dimensions["F"].hidden = se.column_dimensions["G"].hidden = True
PROP_LIST, VIEW_LIST, EXP_PROP_LIST = "Settings!$B$6:$B$11", "Settings!$F$6:$F$12", "Settings!$G$6:$G$12"
UNITS = f'IF({ALL},COUNTA(Settings!$B$6:$B$11),1)'

# ---------------------------------------------------------------- Bookings
bk = wb.create_sheet("Bookings")
sheet_base(bk, "Bookings", "One line per stay, from any platform. Nights, payout and status are calculated.",
           [3, 18, 13, 13, 12, 12, 8, 8, 13, 12, 12, 13, 10, 20, 22], rows=K1 + 2)
header(bk, 5, 2, ["Property", "Guest", "Platform", "Check-in", "Check-out", "Nights", "Guests", "Booking total",
                  "Cleaning fee", "Platform fee", "Payout", "Paid out?", "Status", "Note"])
H0 = 17  # first hidden helper column (Q)
for r in range(K0, K1 + 1):
    b = bookings[r - K0] if r - K0 < len(bookings) else (None,) * 10
    for col, v, fmt in ((2, b[0], None), (3, b[1], None), (4, b[2], None), (5, b[3], DATE), (6, b[4], DATE),
                        (8, b[5], "0"), (9, b[6], MONEY), (10, b[7], MONEY), (11, b[8], MONEY), (13, b[9], None)):
        style(bk.cell(r, col, v), True, fmt, align="center" if col in (8, 13) else None)
    style(bk.cell(r, 7, f'=IF(OR(E{r}="",F{r}=""),"",F{r}-E{r})'), False, "0", align="center")
    style(bk.cell(r, 12, f'=IF(E{r}="","",N(I{r})+N(J{r})-N(K{r}))'), False, MONEY, True)
    style(bk.cell(r, 14, f'=IF(E{r}="","",IF(TODAY()<E{r},"Upcoming",IF(TODAY()<F{r},"Staying now",'
                         f'IF(M{r}="Yes","Checked out","Payout not received"))))'), False)
    style(bk.cell(r, 15), True)
    # Hidden helpers: payout and price per night, nights in each month of the dashboard year.
    bk.cell(r, H0, f"=IF(N(G{r})>0,L{r}/G{r},0)")
    bk.cell(r, H0 + 1, f"=IF(N(G{r})>0,N(I{r})/G{r},0)")
    for m in range(12):
        bk.cell(r, H0 + 2 + m, f'=IF(OR($E{r}="",$F{r}=""),0,MAX(0,MIN($F{r},DATE({YR},{m + 2},1))'
                               f'-MAX($E{r},DATE({YR},{m + 1},1))))')
    n1, n12 = L(H0 + 2), L(H0 + 13)
    bk.cell(r, H0 + 14, f"=SUM({n1}{r}:{n12}{r})")                                  # nights in the year
    bk.cell(r, H0 + 15, f"={L(H0)}{r}*{L(H0 + 14)}{r}")                             # payout in the year
    bk.cell(r, H0 + 16, f"={L(H0 + 1)}{r}*{L(H0 + 14)}{r}")                         # nightly price in the year
    bk.cell(r, H0 + 17, f"=SUMPRODUCT({n1}{r}:{n12}{r},--(COLUMN({n1}{r}:{n12}{r})-COLUMN({n1}{r})+1<=Dashboard!$N$1))")
    bk.cell(r, H0 + 18, f'=IF(OR(E{r}="",E{r}<TODAY()),"",E{r}+ROW()/100000)')    # next check-ins
    bk.cell(r, H0 + 19, f'=IF(OR(F{r}="",F{r}<TODAY()),"",F{r}+ROW()/100000)')    # next check-outs
for c in range(H0, H0 + 20):
    bk.column_dimensions[L(c)].hidden = True
PPN, NY, REVY, RATEY, NYTD, KIN, KOUT = (L(H0), L(H0 + 14), L(H0 + 15), L(H0 + 16), L(H0 + 17), L(H0 + 18), L(H0 + 19))
NM = lambda m: L(H0 + 2 + m)
for col, formula in (("B", f"={PROP_LIST}"), ("D", '"' + ",".join(PLATFORMS) + '"'), ("M", '"Yes,No"')):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    bk.add_data_validation(dv)
    dv.add(f"{col}{K0}:{col}{K1}")
st_rng = f"N{K0}:N{K1}"
bk.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$N{K0}="Upcoming"'], fill=fill(INFO)))
bk.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$N{K0}="Staying now"'], fill=fill(OK),
                                                  font=Font(name=F, bold=True)))
bk.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$N{K0}="Payout not received"'], fill=fill(WARN)))
bk.freeze_panes = "C6"

# ---------------------------------------------------------------- Expenses
ex = wb.create_sheet("Expenses")
sheet_base(ex, "Expenses", "One line per cost. Choose Shared for costs that are not for one property.",
           [3, 13, 18, 18, 12, 34], rows=E1 + 2)
header(ex, 5, 2, ["Date", "Property", "Category", "Amount", "Note"])
for r in range(E0, E1 + 1):
    e = expenses[r - E0] if r - E0 < len(expenses) else (None,) * 5
    for col, v, fmt in ((2, e[0], DATE), (3, e[1], None), (4, e[2], None), (5, e[3], MONEY), (6, e[4], None)):
        style(ex.cell(r, col, v), True, fmt)
for col, formula in (("C", f"={EXP_PROP_LIST}"), ("D", "=Settings!$B$15:$B$24")):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ex.add_data_validation(dv)
    dv.add(f"{col}{E0}:{col}{E1}")
ex.freeze_panes = "B6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Rental dashboard", "Pick a year and a property. Nights are counted in the month they fall in.",
           [3, 16, 14, 14, 15, 15, 15, 3, 15, 22, 20, 3], rows=60)
db["B4"] = "Year"
db["B4"].font = font(11, True)
style(db["C4"], True, "0", True, "center")
db["C4"] = Y
db["D4"] = "Property"
db["D4"].font = font(11, True)
db["D4"].alignment = Alignment(horizontal="right")
style(db["E4"], True, bold=True)
db["E4"] = "All properties"
db.merge_cells("E4:F4")
dv = DataValidation(type="list", formula1=f"={VIEW_LIST}", allow_blank=False)
db.add_data_validation(dv)
dv.add("E4")
# Hidden helpers: last month counted "to date", and the days in those months.
db["N1"] = "=IF($C$4<YEAR(TODAY()),12,IF($C$4>YEAR(TODAY()),0,MONTH(TODAY())))"
db["N2"] = "=IF($N$1=0,0,DATE($C$4,$N$1+1,1)-DATE($C$4,1,1))"
db.column_dimensions["M"].hidden = db.column_dimensions["N"].hidden = True
pick = lambda col: f"IF({ALL},SUM({BK(col)}),SUMIFS({BK(col)},{BK('B')},$E$4))"
ys, ye = "DATE($C$4,1,1)", "DATE($C$4,12,31)"
exp_year = (f'IF({ALL},SUMIFS({EX("E")},{EX("B")},">="&{ys},{EX("B")},"<="&{ye}),'
            f'SUMIFS({EX("E")},{EX("B")},">="&{ys},{EX("B")},"<="&{ye},{EX("C")},$E$4))')
tiles = [("B", "Revenue", f"={pick(REVY)}", TEAL, MONEY),
         ("C", "Expenses", f"={exp_year}", "B4541F", MONEY),
         ("D", "Profit", "=B7-C7", TEAL, NET),
         ("E", "Nights booked", f"={pick(NY)}", TEAL, "0"),
         ("F", "Occupancy to date", f'=IF(OR($N$2=0,{UNITS}=0),"",{pick(NYTD)}/({UNITS}*$N$2))', TEAL, "0%"),
         ("G", "Nightly rate", f'=IF(E7=0,"",{pick(RATEY)}/E7)', TEAL, MONEY)]
for col, label, formula, colour, fmt in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = fmt
    db[f"{col}7"].font = font(16, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db.row_dimensions[7].height = 34

header(db, 9, 2, ["Month", "Nights", "Occupancy", "Revenue", "Expenses", "Profit"])
for m in range(12):
    r = 10 + m
    ms, me = f"DATE($C$4,{m + 1},1)", f"EOMONTH(DATE($C$4,{m + 1},1),0)"
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f"=IF({ALL},SUM({BK(NM(m))}),SUMIFS({BK(NM(m))},{BK('B')},$E$4))"), False, "0", align="center")
    style(db.cell(r, 4, f'=IF($N{r}=0,"",C{r}/$N{r})'), False, "0%", align="center")
    style(db.cell(r, 5, f"=IF({ALL},SUMPRODUCT({BK(NM(m))},{BK(PPN)}),"
                        f"SUMPRODUCT({BK(NM(m))},{BK(PPN)},--({BK('B')}=$E$4)))"), False, MONEY)
    style(db.cell(r, 6, f'=IF({ALL},SUMIFS({EX("E")},{EX("B")},">="&{ms},{EX("B")},"<="&{me}),'
                        f'SUMIFS({EX("E")},{EX("B")},">="&{ms},{EX("B")},"<="&{me},{EX("C")},$E$4))'), False, MONEY)
    style(db.cell(r, 7, f"=E{r}-F{r}"), False, NET, True)
    db.cell(r, 14, f"=DAY(EOMONTH({ms},0))*{UNITS}")
style(db.cell(22, 2, "Year"), False, bold=True, align="center")
style(db.cell(22, 3, "=SUM(C10:C21)"), False, "0", True, "center")
style(db.cell(22, 4, '=IF(SUM(N10:N21)=0,"",C22/SUM(N10:N21))'), False, "0%", True, "center")
for col in "EFG":
    style(db[f"{col}22"], False, NET if col == "G" else MONEY, True)
    db[f"{col}22"] = f"=SUM({col}10:{col}21)"
db.conditional_formatting.add("B10:G21", FormulaRule(formula=["AND($C$4=YEAR(TODAY()),ROW()-9=MONTH(TODAY()))"],
                                                     fill=fill("E6F2F2")))

header(db, 24, 2, ["Property", "Nights", "Occupancy", "Revenue", "Expenses", "Profit"])
for k in range(6):
    r = 25 + k
    style(db.cell(r, 2, f'=IF(Settings!B{6 + k}="","",Settings!B{6 + k})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",SUMIFS({BK(NY)},{BK("B")},B{r}))'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(OR(B{r}="",$N$2=0),"",SUMIFS({BK(NYTD)},{BK("B")},B{r})/$N$2)'), False, "0%", align="center")
    style(db.cell(r, 5, f'=IF(B{r}="","",SUMIFS({BK(REVY)},{BK("B")},B{r}))'), False, MONEY)
    style(db.cell(r, 6, f'=IF(B{r}="","",SUMIFS({EX("E")},{EX("C")},B{r},{EX("B")},">="&{ys},{EX("B")},"<="&{ye}))'),
          False, MONEY)
    style(db.cell(r, 7, f'=IF(B{r}="","",E{r}-F{r})'), False, NET, True)
style(db.cell(31, 2, "Shared costs"), False, bold=True)
for col in "CDE":
    style(db[f"{col}31"], False)
style(db.cell(31, 6, f'=SUMIFS({EX("E")},{EX("C")},"Shared",{EX("B")},">="&{ys},{EX("B")},"<="&{ye})'), False, MONEY)
style(db.cell(31, 7, "=-F31"), False, NET, True)
db["B32"] = "Occupancy here counts the months up to today. Shared costs belong to no single property."
db["B32"].font = font(9, color=MUTED, italic=True)

# Right-hand column: what needs attention and what is coming up.
ins = f'COUNTIFS({BK("E")},">="&TODAY(),{BK("E")},"<="&TODAY()+7)'
now = f'COUNTIFS({BK("N")},"Staying now")'
owed = f'COUNTIFS({BK("N")},"Payout not received")'
owed_sum = f'SUMIFS({BK("L")},{BK("N")},"Payout not received")'
notes = [
    (f'=IF({ins}=0,"No check-ins in the next 7 days",IF({ins}=1,"1 check-in",{ins}&" check-ins")&" in the next 7 days")', INFO, None),
    (f'=IF({now}=0,"Nobody is staying right now",IF({now}=1,"1 booking is",{now}&" bookings are")&" staying right now")', OK, None),
    (f'=IF({owed}=0,"All payouts received",IF({owed}=1,"1 payout",{owed}&" payouts")&" not received yet: "&FIXED({owed_sum},2))',
     WARN, "All"),
    ('=IF(MAX(E10:E21)<=0,"No revenue yet this year","Best month so far: "&INDEX(B10:B21,MATCH(MAX(E10:E21),E10:E21,0)))',
     INFO, None),
]
db.merge_cells("I9:K9")
db["I9"] = "Heads-up"
db["I9"].font = font(11, True, "FFFFFF")
db["I9"].fill = fill(TEAL)
db["I9"].alignment = Alignment(vertical="center")
for k, (formula, colour, calm) in enumerate(notes):
    r = 10 + k
    db.merge_cells(f"I{r}:K{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "IJK":
        db[f"{col}{r}"].border = box
    if calm:
        db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))


def coming(top, title, keys, third_label, third):
    header(db, top, 9, [title, "Property", third_label])
    db.row_dimensions[top].height = None  # the row is shared with the tables on the left
    for k in range(5):
        r = top + 1 + k
        db.cell(r, 13, f'=IFERROR(SMALL({BK(keys)},{k + 1}),"")')
        hit = f"MATCH($M{r},{BK(keys)},0)"
        style(db.cell(r, 9, f'=IF($M{r}="","",INT($M{r}))'), False, DATE, align="center")
        style(db.cell(r, 10, f'=IF($M{r}="","",INDEX({BK("B")},{hit}))'), False)
        style(db.cell(r, 11, f'=IF($M{r}="","",{third(hit, r)})'), False, align="center")


coming(15, "Next check-ins", KIN, "Guest", lambda hit, r: f'INDEX({BK("C")},{hit})')
coming(25, "Next check-outs", KOUT, "Same-day arrival",
       lambda hit, r: f'IF(COUNTIFS({BK("B")},J{r},{BK("E")},I{r})>0,"Yes","No")')
db.conditional_formatting.add("K26:K30", FormulaRule(formula=['$K26="Yes"'], fill=fill(WARN), font=Font(name=F, bold=True)))

chart = BarChart()
chart.type = "col"
chart.title = "Revenue and expenses by month"
chart.height, chart.width = 7.5, 16
chart.add_data(Reference(db, min_col=5, max_col=6, min_row=9, max_row=21), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=10, max_row=21))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E0A21B"
db.add_chart(chart, "B34")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: type the names of your properties (up to 6). Change the expense categories if you like."),
    ("2", "Bookings: one line per stay, with the dates, what the guest paid, the cleaning fee and the platform fee."),
    ("3", "When the money reaches your account, choose Yes under 'Paid out?'."),
    ("4", "Expenses: one line per cost. Choose Shared for costs that are not for one property."),
    ("5", "Dashboard: pick the year and a property, or All properties, to see each month and each home."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
notes_st = ["Revenue here is your payout: what the guest paid plus the cleaning fee, minus the platform fee.",
            "A stay that crosses two months is split by nights, so each month shows its own nights and revenue.",
            "The example properties, bookings and expenses show how it works. Delete them and add your own.",
            "Works in Google Sheets and Microsoft Excel, in any currency, with bookings from any platform."]
for k, text in enumerate(notes_st):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Vacation-Rental-Host-Tracker.xlsx")
wb.save(out)
print("saved", out, len(bookings), "bookings", len(expenses), "expenses")
