"""Etsy listing photos for the Vacation Rental Host Tracker (run build.py first)."""
import os, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Vacation-Rental-Host-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
r = Renders(xlsx, tmp)
img_dir = os.path.join(HERE, "images")

img = canvas(dark=True)
title(img, "Vacation rental\nhost tracker", "Bookings from any platform, occupancy, nightly rate and profit per month",
      dark=True, width=2300)
laptop(img, r.shot("Rental dashboard", until="Occupancy here counts"), 600, 700, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

# The latest twelve bookings: older rows hidden, some columns hidden.
bk = load_workbook(xlsx)["Bookings"]
last = max(r_ for r_ in range(6, 506) if bk.cell(r_, 2).value)
close = Renders(xlsx, tmp, hide={"Bookings": ["H", "J", "K", "O"]}, hide_rows={"Bookings": list(range(6, last - 11)) + list(range(last + 1, 506))})
detail(os.path.join(img_dir, "02-bookings.jpg"), "Every stay, with its payout and status",
       "One line per booking. Nights, payout and status are worked out for you.",
       close.shot("Bookings", until=f"{bk.cell(last, 6).value:%d %b %Y}"),
       [("#E3E8FF", "Upcoming"), ("#D8F0DC", "Staying now"), ("#FFE2B8", "Payout not received")])

left = Renders(xlsx, tmp, hide={"Dashboard": ["I", "J", "K"]})
detail(os.path.join(img_dir, "03-months.jpg"), "Each month and each home",
       "Nights, occupancy, revenue, expenses and profit. Stays that cross a month end are split by nights.",
       left.shot("Rental dashboard", until="Shared costs"))

img = canvas()
y = title(img, "Inside the file", "Four tabs and a Start Here page. You type in the yellow cells and the rest is calculated.")
grid(img, [
    (content(r.page("Rental dashboard")), "Dashboard", "Year and property view, months, what is coming up"),
    (content(r.page("Bookings")), "Bookings", "Stays from any platform, payout, status"),
    (content(r.page("Expenses")), "Expenses", "Per property or shared, by category"),
    (content(r.page("Settings")), "Settings", "Up to 6 properties, your categories"),
], y + 80)
save(img, os.path.join(img_dir, "04-whats-inside.jpg"))
