"""Home Inventory for Insurance (Excel + Google Sheets).

Every item room by room, with what it cost, what it is worth now, the receipt, the photo, the serial number and the
warranty; what each room and category adds up to against the contents cover; what proof is missing; and a list to
print for one room or the whole home. It organises the list: it does not decide what an insurer pays. Only functions
both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
R0, R1 = 6, 35         # rooms
I0, I1 = 6, 1005       # items
K0, K1 = 14, 25        # categories on Settings
NREP = 40              # lines on the printed list
HOME, CUR, COVER, INSURER, POLICY, RENEWAL, SOON = (f"Settings!$C${r}" for r in range(4, 11))
RM = lambda c: f"Rooms!${c}${R0}:${c}${R1}"
IT = lambda c: f"Items!${c}${I0}:${c}${I1}"
IT1 = lambda c: f"Items!${c}$1:${c}${I1}"
CATS = f"Settings!$B${K0}:$B${K1}"
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{",".join(chr(34) + m + chr(34) for m in MON)})&" "&YEAR({d})'


# ---------------------------------------------------------------- example data (a made-up home)
rng = random.Random(45)
D = lambda y, m, d: dt.date(y, m, d)
Y = today.year
rooms = [("Living room", None), ("Kitchen", None), ("Main bedroom", None), ("Home office", None),
         ("Kids room", None), ("Bathroom", None), ("Hallway", "Coats and shoes"), ("Garage", "Bikes and tools")]
categories = ["Electronics", "Appliances", "Furniture", "Jewelry and watches", "Clothing", "Kitchenware", "Tools",
              "Sports", "Music", "Art and decor", "Toys", "Other"]
items = [  # room, item, category, model, bought, paid, worth now, receipt, photo, warranty years
    ("Living room", "55 inch TV", "Electronics", "Model KX-55", (Y - 1, 11, 20), 899, 750, "File", True, 2),
    ("Living room", "Sofa, three seats", "Furniture", "Linen, grey", (Y - 3, 4, 2), 1250, 900, "Paper", True, 0),
    ("Living room", "Sound bar", "Electronics", "SB-300", (Y - 2, 12, 1), 329, 250, "File", True, 1),
    ("Living room", "Game console", "Electronics", "GC-5 with two pads", (Y - 1, 12, 20), 549, 450, "File", True, 1),
    ("Living room", "Rug", "Art and decor", "Wool, 2 x 3 m", (Y - 4, 6, 10), 480, 300, "No", True, 0),
    ("Living room", "Painting of the harbour", "Art and decor", "Oil on canvas", (Y - 6, 8, 15), 900, 900, "Paper", True, 0),
    ("Living room", "Bookcase", "Furniture", "Oak, 5 shelves", (Y - 5, 2, 3), 380, 250, "No", False, 0),
    ("Kitchen", "Fridge freezer", "Appliances", "FF-360", (Y - 2, 3, 12), 1099, 850, "File", True, 5),
    ("Kitchen", "Dishwasher", "Appliances", "DW-45", (Y - 3, 9, 1), 549, 380, "File", True, 2),
    ("Kitchen", "Coffee machine", "Appliances", "Espresso EM-2", (Y, 1, 8), 429, 400, "File", True, 2),
    ("Kitchen", "Stand mixer", "Kitchenware", "SM-5, red", (Y - 2, 12, 25), 399, 300, "No", True, 0),
    ("Kitchen", "Knife set", "Kitchenware", "8 knives and block", (Y - 1, 7, 7), 229, 200, "Paper", False, 0),
    ("Kitchen", "Pots and pans", "Kitchenware", "10 pieces", (Y - 4, 11, 30), 350, 200, "No", False, 0),
    ("Main bedroom", "Bed and mattress", "Furniture", "King size", (Y - 2, 5, 20), 1600, 1300, "File", True, 10),
    ("Main bedroom", "Wardrobe", "Furniture", "Two doors, white", (Y - 5, 1, 12), 520, 350, "No", True, 0),
    ("Main bedroom", "Engagement ring", "Jewelry and watches", "Gold, one diamond", (Y - 7, 2, 14), 2400, 3100, "Paper", True, 0),
    ("Main bedroom", "Watch", "Jewelry and watches", "Steel, automatic", (Y - 3, 12, 24), 780, 700, "Paper", True, 2),
    ("Main bedroom", "Winter coats", "Clothing", "Two coats", (Y - 1, 11, 2), 460, 350, "No", False, 0),
    ("Main bedroom", "Clothes and shoes", "Clothing", "Everyday", None, None, 2500, "No", False, 0),
    ("Home office", "Laptop", "Electronics", "Pro 14, 16 GB", (Y, 2, 20), 1499, 1350, "File", True, 1),
    ("Home office", "Monitor", "Electronics", "27 inch, 4K", (Y - 1, 3, 3), 379, 320, "File", True, 3),
    ("Home office", "Desk chair", "Furniture", "Mesh, adjustable", (Y - 2, 9, 9), 449, 350, "File", True, 5),
    ("Home office", "Printer", "Electronics", "Colour laser", (Y - 3, 6, 6), 299, 150, "No", True, 1),
    ("Home office", "Camera", "Electronics", "Mirrorless with 2 lenses", (Y - 2, 7, 1), 1250, 950, "File", False, 2),
    ("Home office", "Acoustic guitar", "Music", "Dreadnought", (Y - 6, 12, 20), 650, 600, "No", True, 0),
    ("Kids room", "Bunk bed", "Furniture", "Pine, with drawers", (Y - 2, 8, 18), 690, 500, "File", True, 2),
    ("Kids room", "Tablet", "Electronics", "10 inch, 64 GB", (Y - 1, 12, 25), 349, 280, "File", True, 1),
    ("Kids room", "Toys and games", "Toys", "Lego, board games", None, None, 800, "No", True, 0),
    ("Kids room", "Keyboard piano", "Music", "61 keys", (Y - 1, 12, 25), 229, 180, "File", False, 1),
    ("Bathroom", "Hair dryer and straightener", "Appliances", "Two items", (Y - 1, 5, 5), 180, 140, "No", False, 0),
    ("Bathroom", "Electric toothbrushes", "Appliances", "Two, with chargers", (Y, 3, 3), 160, 150, "File", True, 2),
    ("Hallway", "Mirror", "Art and decor", "Round, brass frame", (Y - 3, 10, 10), 220, 180, "No", True, 0),
    ("Hallway", "Coats and boots", "Clothing", "Family", None, None, 900, "No", False, 0),
    ("Garage", "Road bike", "Sports", "Carbon, size 56", (Y - 2, 4, 25), 2100, 1700, "Paper", True, 2),
    ("Garage", "Kids bikes", "Sports", "Two bikes", (Y - 1, 6, 1), 520, 400, "File", True, 1),
    ("Garage", "Cordless drill set", "Tools", "18 V with 2 batteries", (Y - 1, 2, 10), 259, 220, "File", True, 3),
    ("Garage", "Lawn mower", "Tools", "Battery, 40 cm", (Y - 2, 5, 1), 399, 300, "No", True, 2),
    ("Garage", "Camping gear", "Sports", "Tent, sleeping bags", (Y - 3, 7, 1), 740, 500, "No", False, 0),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def check_fill(ws, rng_, first):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}<>""'], fill=fill(BAD)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your home, your policy, and the categories for your things.", [3, 34, 22, 4, 52],
           rows=K1 + 3)
for r, lab, value, fmt, note in (
        (4, "Home", "Our home", None, "At the top of the Dashboard and the printed list."),
        (5, "Currency", "USD", None, "Shown on the printed list."),
        (6, "Contents cover", 40000, MONEY, "The amount on your policy for the things inside your home."),
        (7, "Insurer", "Your insurer", None, "Where to call, so it is on hand."),
        (8, "Policy number", "HC-48213-07", None, None),
        (9, "Renews on", D(Y, 12, 1), DATE, "The Dashboard counts down to it."),
        (10, "Warranty ends soon (days)", 60, "0", "A warranty ending within this many days is shown.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, r == 4, "center")
    se[f"C{r}"] = value
    if note:
        se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
se[f"B{K0 - 1}"] = "Categories"
se[f"B{K0 - 1}"].font = font(12, True, TEAL_D)
for k in range(K1 - K0 + 1):
    style(se.cell(K0 + k, 2, categories[k] if k < len(categories) else None), True)

# ---------------------------------------------------------------- Rooms
rm = wb.create_sheet("Rooms")
sheet_base(rm, "Rooms", "Each room or place once, in the order you walk through your home.",
           [3, 20, 24, 9, 13, 12, 11], rows=R1 + 2)
header(rm, 5, 2, ["Room", "Note", "Items", "Worth", "No receipt", "No photo"])
for r in range(R0, R1 + 1):
    x = rooms[r - R0] if r - R0 < len(rooms) else (None, None)
    style(rm.cell(r, 2, x[0]), True, bold=True)
    style(rm.cell(r, 3, x[1]), True)
    style(rm.cell(r, 4, f'=IF(B{r}="","",SUMIFS({IT("S")},{IT("B")},B{r}))'), False, "0", align="center")
    style(rm.cell(r, 5, f'=IF(B{r}="","",SUMIFS({IT("T")},{IT("B")},B{r}))'), False, MONEY, True)
    style(rm.cell(r, 6, f'=IF(B{r}="","",SUMIFS({IT("W")},{IT("B")},B{r}))'), False, "0", align="center")
    style(rm.cell(r, 7, f'=IF(B{r}="","",SUMIFS({IT("X")},{IT("B")},B{r}))'), False, "0", align="center")
rm.freeze_panes = "C6"

# ---------------------------------------------------------------- Items
it = wb.create_sheet("Items")
sheet_base(it, "Items", "Everything worth listing, room by room. The photo column is where you keep the picture: a "
           "file name or a link.", [3, 15, 24, 17, 20, 15, 12, 11, 11, 12, 16, 13, 18, 11, 7, 12, 18, 20],
           rows=I1 + 2)
header(it, 5, 2, ["Room", "Item", "Category", "Brand and model", "Serial number", "Bought", "Paid", "Worth now",
                  "Receipt", "Photo", "Warranty until", "Note", "Value", "Age", "Warranty", "Missing", "Check"])
for r in range(I0, I1 + 1):
    if r - I0 < len(items):
        room, name, cat, model, bought, paid, worth, receipt, photo, wy = items[r - I0]
        bdate = D(*bought) if bought else None
        warranty = D(bdate.year + wy, bdate.month, min(bdate.day, 28)) if bdate and wy else None
        serial = f"SN{rng.randint(10000000, 99999999)}" if cat in ("Electronics", "Appliances", "Tools", "Sports") \
            else None
        pic = f"IMG_{rng.randint(1000, 9999)}.jpg" if photo else None
        x = [room, name, cat, model, serial, bdate, paid, worth, receipt, pic, warranty, None]
    else:
        x = [None] * 12
    for c, v, fmt in zip(range(2, 14), x, (None, None, None, None, None, DATE, MONEY, MONEY, None, None, DATE, None)):
        style(it.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 7, 10, 12) else None)
    style(it.cell(r, 14, f'=IF(C{r}="","",IF(ISNUMBER(I{r}),I{r},IF(ISNUMBER(H{r}),H{r},"")))'), False, MONEY, True)
    style(it.cell(r, 15, f'=IF(OR(C{r}="",NOT(ISNUMBER(G{r}))),"",(TODAY()-INT(G{r}))/365.25)'), False, "0.0",
          align="center")
    style(it.cell(r, 16, f'=IF(OR(C{r}="",NOT(ISNUMBER(L{r}))),"",IF(L{r}<TODAY(),"Ended",'
                         f'IF(L{r}-TODAY()<=N({SOON}),"Ends soon","Covered")))'), False, align="center")
    style(it.cell(r, 17, f'=IF(C{r}="","",MID(IF(N{r}="",", value","")&IF(OR(J{r}="",J{r}="No"),", receipt","")'
                         f'&IF(K{r}="",", photo",""),3,100))'), False, align="center")
    style(it.cell(r, 18, f'=IF(AND(B{r}="",C{r}=""),"",IF(C{r}="","Needs the item",IF(B{r}="","Needs a room",'
                         f'IF(COUNTIF({RM("B")},B{r})=0,"Not on Rooms",IF(AND(H{r}<>"",NOT(ISNUMBER(H{r}))),'
                         f'"Paid not a number",IF(AND(I{r}<>"",NOT(ISNUMBER(I{r}))),"Worth not a number",'
                         f'IF(AND(G{r}<>"",NOT(ISNUMBER(G{r}))),"Bought is not a date",'
                         f'IF(AND(L{r}<>"",NOT(ISNUMBER(L{r}))),"Warranty is not a date",""))))))))'), False,
          align="center")
    it[f"S{r}"] = f'=IF(AND(C{r}<>"",R{r}=""),1,0)'                                             # S: counts
    it[f"T{r}"] = f"=IF(S{r}=1,N(N{r}),0)"                                                      # T: value
    it[f"U{r}"] = f'=IF(AND(S{r}=1,T{r}>0),ROUND(T{r}*100,0)*10000+10000-ROW(),"")'            # U: most valuable
    it[f"V{r}"] = (f'=IF(AND(S{r}=1,OR(Report!$C$5="",B{r}=Report!$C$5)),'
                   f'MATCH(B{r},{RM("B")},0)*10000+ROW(),"")')                                  # V: printed list
    it[f"W{r}"] = f'=IF(AND(S{r}=1,OR(J{r}="",J{r}="No")),1,0)'                                # W: no receipt
    it[f"X{r}"] = f'=IF(AND(S{r}=1,K{r}=""),1,0)'                                              # X: no photo
    it[f"Y{r}"] = f'=IF(AND(S{r}=1,Q{r}<>""),ROUND(T{r}*100,0)*10000+10000-ROW(),"")'           # Y: proof to find
for c in "STUVWXY":
    it.column_dimensions[c].hidden = True
check_fill(it, f"R{I0}:R{I1}", f"R{I0}")
it.conditional_formatting.add(f"Q{I0}:Q{I1}", FormulaRule(formula=[f'Q{I0}<>""'], fill=fill(WARN)))
for text, colour in (("Ends soon", WARN), ("Covered", OK)):
    it.conditional_formatting.add(f"P{I0}:P{I1}", FormulaRule(formula=[f'P{I0}="{text}"'], fill=fill(colour)))
dropdown(it, f"={RM('B')}", f"B{I0}:B{I1}", strict=False)
dropdown(it, f"={CATS}", f"D{I0}:D{I1}", strict=False)
dropdown(it, '"Paper,File,No"', f"J{I0}:J{I1}", strict=False)
it.freeze_panes = "D6"

# ---------------------------------------------------------------- Report (to print)
rp = wb.create_sheet("Report")
sheet_base(rp, "", "", [3, 14, 22, 18, 14, 12, 12, 9, 13], rows=NREP + 20)
rp["B2"] = f'=IF({HOME}="","Home inventory",{HOME}&": home inventory")'
rp["B2"].font = font(18, True, TEAL_D)
rp["B3"] = '="Printed on "&' + long_date("TODAY()")
rp["B3"].font = font(10, color=MUTED, italic=True)
rp["B5"] = "Room"
rp["B5"].font = font(11, True)
style(rp["C5"], True, bold=True)
rp["C5"] = None
dropdown(rp, f"={RM('B')}", "C5", strict=False)
rp["D5"] = "Leave it empty for the whole home."
rp["D5"].font = font(9, color=MUTED, italic=True)
rp["B6"] = f'=IF({INSURER}="","","Insurer: "&{INSURER})&IF({POLICY}="","","   Policy: "&{POLICY})'
rp["B6"].font = font(10, color=MUTED)
header(rp, 8, 2, ["Room", "Item", "Brand and model", "Serial number", "Bought", "Value", "Receipt", "Photo"])
for k in range(NREP):
    r = 9 + k
    rp[f"K{r}"] = f"=IFERROR(SMALL({IT('V')},{k + 1}),\"\")"                                    # K: item row
    g = lambda c: f"INDEX({IT1(c)},MOD(K{r},10000))"
    for i, (c, fmt) in enumerate((("B", None), ("C", None), ("E", None), ("F", None), ("G", DATE), ("N", MONEY),
                                  ("J", None), ("K", None))):
        f = f'=IF(K{r}="","",IF({g(c)}="","",{g(c)}))'
        cell = style(rp.cell(r, 2 + i, f), False, fmt, i == 1, "center" if i in (3, 4, 6, 7) else None)
        cell.font = font(9, i == 1)
    rp.row_dimensions[r].height = 15
rp.column_dimensions["K"].hidden = True
r = 9 + NREP
rp[f"B{r}"] = f'=IF(COUNT({IT("V")})>{NREP},"And "&(COUNT({IT("V")})-{NREP})&" more items, in the total","")'
rp[f"B{r}"].font = font(9, color=MUTED, italic=True)
rp[f"B{r + 1}"] = f'=COUNT({IT("V")})&" items"'
rp[f"B{r + 1}"].font = font(11, True)
rp[f"F{r + 1}"] = "Total"
rp[f"F{r + 1}"].font = font(11, True)
rp[f"G{r + 1}"] = f'=SUMPRODUCT(--({IT("V")}<>""),{IT("T")})'
rp[f"G{r + 1}"].number_format = MONEY
rp[f"G{r + 1}"].font = font(11, True, TEAL_D)
rp[f"H{r + 1}"] = f'=IF({CUR}="","",{CUR})'
rp[f"H{r + 1}"].font = font(10, True, MUTED)
rp[f"B{r + 3}"] = "Keep a copy away from home. This list records what you own: it does not decide what an insurer pays."
rp[f"B{r + 3}"].font = font(8, color=MUTED, italic=True)
rp.page_setup.orientation = "portrait"
rp.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "", "", [3, 20, 9, 13, 9, 3, 24, 16, 13, 3], rows=48)
db["B2"] = f'=IF({HOME}="","Home inventory",{HOME})'
db["B2"].font = font(22, True, TEAL_D)
db["B3"] = '="What your things are worth, room by room, and the proof still to find."'
db["B3"].font = font(11, color=MUTED, italic=True)
tile(db, "B", 4, "Worth in all", f"=SUM({IT('T')})", MONEY)
tile(db, "C", 4, "Items", f"=SUM({IT('S')})", "0")
tile(db, "D", 4, "Contents cover", f'=IF(N({COVER})=0,"",N({COVER}))', MONEY)
tile(db, "E", 4, "No receipt", f"=SUM({IT('W')})", "0", "B86E1C")
tile(db, "G", 4, "Cover less the list", f'=IF(N({COVER})=0,"",N({COVER})-B5)', MONEY)
tile(db, "H", 4, "No photo", f"=SUM({IT('X')})", "0", "B86E1C")
tile(db, "I", 4, "Renews in (days)", f'=IF(NOT(ISNUMBER({RENEWAL})),"",INT({RENEWAL})-TODAY())', "0")
db.conditional_formatting.add("G5", FormulaRule(formula=['AND(ISNUMBER(G5),G5<0)'], fill=fill(BAD)))
header(db, 7, 2, ["Room", "Items", "Worth", "Share"])
for k in range(10):
    r = 8 + k
    rr = f"INDEX({RM('B')},{k + 1})"
    style(db.cell(r, 2, f'=IF({rr}="","",{rr})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",INDEX({RM("D")},{k + 1}))'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",INDEX({RM("E")},{k + 1}))'), False, MONEY)
    style(db.cell(r, 5, f'=IF(OR(B{r}="",N($B$5)=0),"",D{r}/$B$5)'), False, "0%", align="center")
header(db, 7, 7, ["Most valuable", "Room", "Worth"])
for n in range(10):
    r = 8 + n
    db[f"L{r}"] = f"=IFERROR(LARGE({IT('U')},{n + 1}),\"\")"                                     # L: item key
    g = lambda c: f"INDEX({IT1(c)},10000-MOD(L{r},10000))"
    style(db.cell(r, 7, f'=IF(L{r}="","",{g("C")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(L{r}="","",{g("B")})'), False, align="center")
    style(db.cell(r, 9, f'=IF(L{r}="","",{g("N")})'), False, MONEY)
header(db, 19, 2, ["Category", "Items", "Worth", ""])
for k in range(K1 - K0 + 1):
    r = 20 + k
    cat = f"Settings!$B${K0 + k}"
    style(db.cell(r, 2, f'=IF({cat}="","",{cat})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",SUMIFS({IT("S")},{IT("D")},B{r}))'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",SUMIFS({IT("T")},{IT("D")},B{r}))'), False, MONEY)
    style(db.cell(r, 5), False)
header(db, 19, 7, ["Proof to find first", "Missing", "Worth"])
for n in range(8):
    r = 20 + n
    db[f"M{r}"] = f"=IFERROR(LARGE({IT('Y')},{n + 1}),\"\")"                                     # M: item key
    g = lambda c: f"INDEX({IT1(c)},10000-MOD(M{r},10000))"
    style(db.cell(r, 7, f'=IF(M{r}="","",{g("C")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(M{r}="","",{g("Q")})'), False, align="center")
    style(db.cell(r, 9, f'=IF(M{r}="","",{g("N")})'), False, MONEY)
db["G29"] = f'=IF(COUNT({IT("Y")})>8,"And "&(COUNT({IT("Y")})-8)&" more with something missing","")'
db["G29"].font = font(9, color=MUTED, italic=True)
db["B33"] = f'=IF({INSURER}="","",{INSURER})&IF({POLICY}="","","   Policy "&{POLICY})'
db["B33"].font = font(10, color=MUTED)
for c in "LM":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 18)) + list(range(20, 32)):
    db.row_dimensions[r].height = 18
chart = BarChart()
chart.type = "bar"
chart.title = "Worth by room"
chart.height, chart.width = 7.5, 12
chart.add_data(Reference(db, min_col=4, min_row=7, max_row=17), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=8, max_row=17))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.y_axis.number_format = "#,##0"
chart.x_axis.scaling.orientation = "maxMin"
db.add_chart(chart, "G31")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your home, your contents cover and policy, and the categories you want."),
    ("2", "Rooms: each room or place once, in the order you walk through your home."),
    ("3", "Items: walk through each room and list what is worth listing, with the price, the receipt and a photo."),
    ("4", "Worth now is your own estimate. Leave it empty to use the price you paid."),
    ("5", "Report: pick a room, or none for the whole home, and print the list or save it as a PDF."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example home and its things are made up. To remove them, select the yellow cells and "
                          "press Delete, rather than deleting whole rows.",
                          "This list records what you own. What an insurer pays depends on your policy.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Items", "Rooms", "Report", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Home-Inventory.xlsx")
wb.save(out)
print("saved", out, len(items), "items")
