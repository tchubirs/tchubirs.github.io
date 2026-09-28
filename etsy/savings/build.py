"""Savings Goals and Sinking Funds Tracker (Excel + Google Sheets).

Several savings goals, each with a target, an optional date and a monthly plan;
a log of money put in and taken out; what each goal still needs per month and
per payday; a month by month grid; and two challenges (52 weeks and 100
envelopes). Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
D = lambda y, m, d: dt.date(y, m, d)
NET = '#,##0.00;[Red]-#,##0.00'
O0, O1 = 6, 20        # goals
L0, L1 = 6, 1005      # log lines
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
YR, FREQ = "Dashboard!$C$4", "Dashboard!$F$4"
PER_PAYDAY = f'IF({FREQ}="Weekly",52/12,IF({FREQ}="Every 2 weeks",26/12,IF({FREQ}="Twice a month",2,1)))'
LG = lambda col: f"'Savings log'!${col}${L0}:${col}${L1}"
GO = lambda col: f"Goals!${col}${O0}:${col}${O1}"
GREY = "EEF1F1"

# ---------------------------------------------------------------- example data
goals = [
    # goal, target, target date, saved before, plan per month, note
    ("Emergency fund", 6000, D(Y + 1, 6, 30), 2400, 300, "Three months of rent and bills"),
    ("Summer trip", 3200, D(Y + 1, 6, 15), None, 250, "Flights and a week by the sea"),
    ("Christmas gifts", 800, D(Y, 12, 1), None, 100, None),
    ("Car insurance", 1140, D(Y + 1, 3, 1), None, 95, "Paid once a year"),
    ("New laptop", 1500, D(Y + 1, 1, 31), 200, 150, None),
    ("Home repairs", 2000, None, 500, 100, "Roof and gutters, no fixed date"),
    ("Concert tickets", 240, D(Y, 8, 15), None, 60, None),
]
log = []
for m in range(1, today.month + 1):
    log.append((D(Y, m, 1), "Emergency fund", 300, None, "Transfer on payday"))
    log.append((D(Y, m, 2), "Home repairs", 100, None, None))
    if m >= 3:
        log.append((D(Y, m, 2), "Car insurance", 95, None, None))
    if m >= 4:
        log.append((D(Y, m, 15), "New laptop", 150, None, None))
    if 4 <= m <= 7:
        log.append((D(Y, m, 15), "Concert tickets", 60, None, None))
    if m >= 6:
        log.append((D(Y, m, 16), "Summer trip", 250, None, None))
    if m >= 7:
        log.append((D(Y, m, 16), "Christmas gifts", 100, None, None))
log.append((D(Y, 5, 12), "Emergency fund", None, 450, "Car repair"))
log.sort(key=lambda e: e[0])
WEEK1 = D(Y, 1, 5)
weeks_done = [k for k in range(1, 53) if WEEK1 + dt.timedelta(days=7 * (k - 1)) < today and k not in (12, 25)]
rng = random.Random(7)
envelopes = rng.sample(range(1, 101), 58)
env_days = sorted(today - dt.timedelta(days=rng.randrange(1, 90)) for _ in envelopes)
MONTH_LIST = ",".join(f'"{m}"' for m in MONTHS)

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL, span=1):
    """A number with its label above it, over `span` columns."""
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    last = L(ws[f"{col}{row}"].column + span - 1)
    for r in (row, row + 1):
        for c in range(ws[f"{col}{r}"].column, ws[f"{col}{r}"].column + span):
            ws.cell(r, c).border = box
        if span > 1:
            ws.merge_cells(f"{col}{r}:{last}{r}")
    ws.row_dimensions[row + 1].height = 34


def status_colours(ws, rng_, first):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'OR({first}="Reached",{first}="On track")'], fill=fill(OK)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'OR(LEFT({first},5)="Short",{first}="Past the date")'],
                                                    fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'OR(LEFT({first},5)="Needs",LEFT({first},2)="No")'],
                                                    fill=fill(INFO)))


# ---------------------------------------------------------------- Goals
go = wb.active
go.title = "Goals"
sheet_base(go, "Savings goals", "One line per goal. Leave the date empty for a goal with no deadline.",
           [3, 22, 12, 13, 12, 12, 12, 12, 10, 9, 12, 11, 30, 30], rows=O1 + 4)
header(go, 5, 2, ["Goal", "Target", "Target date", "Saved before", "Plan per month", "Saved", "Left", "Progress",
                  "Months left", "Needed per month", "Per payday", "Status", "Note"])
for r in range(O0, O1 + 1):
    g = goals[r - O0] if r - O0 < len(goals) else (None,) * 6
    for c, (v, fmt) in zip((2, 3, 4, 5, 6, 14), zip(g, (None, MONEY, DATE, MONEY, MONEY, None))):
        style(go.cell(r, c, v), True, fmt, bold=c == 2)
    style(go.cell(r, 7, f'=IF(B{r}="","",N(E{r})+SUMIFS({LG("D")},{LG("C")},B{r})-SUMIFS({LG("E")},{LG("C")},B{r}))'),
          False, MONEY, True)
    style(go.cell(r, 8, f'=IF(B{r}="","",MAX(0,N(C{r})-G{r}))'), False, MONEY)
    style(go.cell(r, 9, f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,MAX(0,G{r}/C{r})))'), False, "0%", align="center")
    months_left = f"MAX(1,(YEAR(D{r})-YEAR(TODAY()))*12+MONTH(D{r})-MONTH(TODAY()))"
    style(go.cell(r, 10, f'=IF(OR(B{r}="",D{r}=""),"",IF(D{r}<TODAY(),0,{months_left}))'), False, "0", align="center")
    style(go.cell(r, 11, f'=IF(B{r}="","",IF(N(H{r})=0,0,IF(D{r}="","",IF(N(J{r})=0,H{r},H{r}/J{r}))))'), False, MONEY)
    style(go.cell(r, 12, f'=IF(OR(B{r}="",K{r}=""),"",K{r}/{PER_PAYDAY})'), False, MONEY)
    finish = f"ROUNDUP(H{r}/F{r},0)"
    style(go.cell(r, 13, f'=IF(B{r}="","",IF(N(C{r})=0,"No target",IF(G{r}>=C{r},"Reached",IF(D{r}="",'
                         f'IF(N(F{r})>0,"No date, about "&{finish}&IF({finish}=1," month"," months")&" at your plan","No date"),'
                         f'IF(D{r}<TODAY(),"Past the date",IF(N(F{r})=0,"Needs "&FIXED(K{r},2)&" a month",'
                         f'IF(F{r}>=K{r}-0.005,"On track","Short by "&FIXED(K{r}-F{r},2)&" a month")))))))'), False).alignment = \
        Alignment(horizontal="left", indent=1)
    go.cell(r, 15, f'=IF(AND(B{r}<>"",D{r}<>"",N(H{r})>0),IF(D{r}>=TODAY(),D{r},""),"")')   # next target dates
go.column_dimensions["O"].hidden = True
status_colours(go, f"M{O0}:M{O1}", f"$M{O0}")
go[f"B{O1 + 2}"] = ("Months left counts the months after this one, up to the month of the target date. "
                    "Per payday follows how often you are paid, on the Dashboard.")
go[f"B{O1 + 2}"].font = font(10, color=MUTED, italic=True)
go.freeze_panes = "C6"

# ---------------------------------------------------------------- Savings log
lg = wb.create_sheet("Savings log")
sheet_base(lg, "Savings log", "One line each time you put money in or take it out.", [3, 13, 22, 12, 12, 32], rows=L1 + 2)
header(lg, 5, 2, ["Date", "Goal", "Put in", "Taken out", "Note"])
for r in range(L0, L1 + 1):
    e = log[r - L0] if r - L0 < len(log) else (None,) * 5
    for c, (v, fmt) in enumerate(zip(e, (DATE, None, MONEY, MONEY, None)), start=2):
        style(lg.cell(r, c, v), True, fmt)
dv = DataValidation(type="list", formula1=f"={GO('B')}", allow_blank=True)
lg.add_data_validation(dv)
dv.add(f"C{L0}:C{L1}")
lg.freeze_panes = "B6"

# ---------------------------------------------------------------- Month by month
mm = wb.create_sheet("Month by month")
sheet_base(mm, "Month by month", "What went into each goal each month of the year on the Dashboard, after money taken out.",
           [3, 22, 11] + [8] * 12 + [11], rows=O1 + 6)
header(mm, 5, 2, ["Goal", "Plan per month"] + MONTHS + ["Year"])
for r in range(O0, O1 + 1):
    style(mm.cell(r, 2, f'=IF(Goals!B{r}="","",Goals!B{r})'), False, bold=True)
    style(mm.cell(r, 3, f'=IF(OR(B{r}="",N(Goals!F{r})=0),"",Goals!F{r})'), False, "#,##0")
    for m in range(12):
        ms, me = f"DATE({YR},{m + 1},1)", f"EOMONTH(DATE({YR},{m + 1},1),0)"
        within = f'{LG("C")},$B{r},{LG("B")},">="&{ms},{LG("B")},"<="&{me}'
        style(mm.cell(r, 4 + m, f'=IF($B{r}="","",IF({ms}>TODAY(),"",SUMIFS({LG("D")},{within})-SUMIFS({LG("E")},{within})))'),
              False, "#,##0;[Red]-#,##0")
    style(mm.cell(r, 16, f'=IF(B{r}="","",SUM(D{r}:O{r}))'), False, NET, True)
tot = O1 + 1
style(mm.cell(tot, 2, "All goals"), False, bold=True)
style(mm.cell(tot, 3), False)
for m in range(13):
    col = L(4 + m)
    first = f"DATE({YR},{m + 1},1)>TODAY()" if m < 12 else "FALSE"
    style(mm.cell(tot, 4 + m, f'=IF({first},"",SUM({col}{O0}:{col}{O1}))'), False, "#,##0;[Red]-#,##0", True)
grid = f"D{O0}:O{O1}"
mm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{O0}<>"",D{O0}<>"",N(D{O0})<0)'], fill=fill(BAD)))
mm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{O0}<>"",N($C{O0})>0,D{O0}<>"",N(D{O0})>=$C{O0})'], fill=fill(OK)))
mm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{O0}<>"",N($C{O0})>0,D{O0}<>"",N(D{O0})>0)'], fill=fill(WARN)))
mm.conditional_formatting.add(grid, FormulaRule(formula=[f'AND($B{O0}<>"",D{O0}<>"",N(D{O0})=0)'], fill=fill(GREY),
                                                font=Font(name=F, color="9AA7A9")))
key = tot + 2
mm.cell(key, 2, "Colours:").font = font(10, True, MUTED)
for k, (label, colour) in enumerate([("Plan met", OK), ("Part of the plan", WARN), ("Nothing put in", GREY),
                                     ("More taken out", BAD)]):
    col = 4 + k * 3
    for c in (col, col + 1):
        mm.cell(key, c).fill = fill(colour)
        mm.cell(key, c).border = box
    mm.cell(key, col, label).font = font(10)
    mm.merge_cells(start_row=key, start_column=col, end_row=key, end_column=col + 1)
mm.freeze_panes = "C6"

# ---------------------------------------------------------------- 52 weeks
wk = wb.create_sheet("52 weeks")
sheet_base(wk, "52 week challenge", "Save a little more each week. Type x under Done for each week you save.",
           [3, 8, 13, 11, 8, 4, 8, 13, 11, 8], rows=42)
wk["B4"] = "First week"
wk["B4"].font = font(11, True)
style(wk["D4"], True, DATE, True, "center")
wk["D4"] = WEEK1
wk["B5"] = "Week 1 amount"
wk["B5"].font = font(11, True)
style(wk["D5"], True, MONEY, True, "center")
wk["D5"] = 1
wk["G4"] = "Order"
wk["G4"].font = font(11, True)
style(wk["H4"], True, None, True, "center")
wk["H4"] = "Low to high"
wk["G5"] = "Week 2 saves twice the week 1 amount, week 52 saves 52 times."
wk["G5"].font = font(10, color=MUTED, italic=True)
dv = DataValidation(type="list", formula1='"Low to high,High to low"', allow_blank=False)
wk.add_data_validation(dv)
dv.add("H4")
done = 'SUMIFS(D12:D37,E12:E37,"<>")+SUMIFS(I12:I37,J12:J37,"<>")'
tile(wk, "B", 7, "Saved", f"={done}", MONEY, span=2)
tile(wk, "D", 7, "Left", "=D5*1378-B8", MONEY, "B4541F", span=2)
tile(wk, "G", 7, "Weeks done", '=COUNTIF(E12:E37,"<>")+COUNTIF(J12:J37,"<>")', "0", span=2)
tile(wk, "I", 7, "Total", "=D5*1378", MONEY, span=2)
for block, col0 in ((0, 2), (1, 7)):
    header(wk, 11, col0, ["Week", "Date", "Amount", "Done"])
    for k in range(26):
        r, week = 12 + k, block * 26 + k + 1
        style(wk.cell(r, col0, week), False, "0", True, "center")
        style(wk.cell(r, col0 + 1, f"=$D$4+7*({L(col0)}{r}-1)"), False, DATE, align="center")
        style(wk.cell(r, col0 + 2, f'=$D$5*IF($H$4="High to low",53-{L(col0)}{r},{L(col0)}{r})'), False, MONEY)
        style(wk.cell(r, col0 + 3, "x" if week in weeks_done else None), True, align="center")
    cells = f"{L(col0)}12:{L(col0 + 3)}37"
    d_, done_ = L(col0 + 1), L(col0 + 3)
    wk.conditional_formatting.add(cells, FormulaRule(formula=[f'${done_}12<>""'], fill=fill(OK)))
    wk.conditional_formatting.add(cells, FormulaRule(formula=[f'AND(${done_}12="",${d_}12+6<TODAY())'], fill=fill(WARN)))
wk["B39"] = "Orange means the week has passed without an x. Low to high makes the last weeks the biggest."
wk["B39"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- 100 envelopes
ev = wb.create_sheet("100 envelopes")
sheet_base(ev, "100 envelope challenge", "Fill one envelope at a time, in any order. Type its number in the list on the right.",
           [3] + [7] * 10 + [3, 13, 11], rows=112)
ev["B4"] = "Envelopes"
ev["B4"].font = font(11, True)
style(ev["E4"], True, "0", True, "center")
ev["E4"] = 100
ev["B5"] = "Envelope 1 holds"
ev["B5"].font = font(11, True)
style(ev["E5"], True, MONEY, True, "center")
ev["E5"] = 1
ev["G4"] = "Envelope 7 holds 7 times envelope 1."
ev["G4"].font = font(10, color=MUTED, italic=True)
ev["G5"] = "With 100 envelopes of 1, the total is 5,050."
ev["G5"].font = font(10, color=MUTED, italic=True)
tile(ev, "B", 7, "Saved", "=SUM(P11:Y20)", MONEY, span=3)
tile(ev, "E", 7, "Left", "=E5*E4*(E4+1)/2-B8", MONEY, "B4541F", span=3)
tile(ev, "H", 7, "Filled", '=COUNTIF(P11:Y20,">0")&" of "&E4', "General", span=4)
for i in range(10):
    r = 11 + i
    ev.row_dimensions[r].height = 26
    for j in range(10):
        n = i * 10 + j + 1
        c = ev.cell(r, 2 + j, f'=IF({n}<=$E$4,{n},"")')
        c.font = font(12, True, TEAL_D)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
        ev.cell(r, 16 + j, f'=IF({L(2 + j)}{r}="","",IF(COUNTIF($N$11:$N$110,{L(2 + j)}{r})>0,{L(2 + j)}{r}*$E$5,0))')
ev.conditional_formatting.add("B11:K20", FormulaRule(formula=['AND(B11<>"",COUNTIF($N$11:$N$110,B11)>0)'],
                                                     fill=fill(OK), font=Font(name=F, bold=True, color="9AA7A9")))
header(ev, 10, 13, ["Date", "Envelope"])
for k in range(100):
    r = 11 + k
    style(ev.cell(r, 13, env_days[k] if k < len(envelopes) else None), True, DATE, align="center")
    style(ev.cell(r, 14, envelopes[k] if k < len(envelopes) else None), True, "0", align="center")
for c in range(16, 26):
    ev.column_dimensions[L(c)].hidden = True
ev["B22"] = "Green envelopes are filled. A number typed twice counts once."
ev["B22"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Savings dashboard", "Pick the year for the months, and how often you are paid for the amounts per payday.",
           [3, 22, 13, 13, 24, 8, 30, 3, 14, 12, 12, 12, 3], rows=60)
db["B4"] = "Year"
db["B4"].font = font(11, True)
style(db["C4"], True, "0", True, "center")
db["C4"] = Y
db["E4"] = "I am paid"
db["E4"].font = font(11, True)
db["E4"].alignment = Alignment(horizontal="right")
style(db["F4"], True, None, True, "center")
db["F4"] = "Every 2 weeks"
db.merge_cells("F4:G4")
dv = DataValidation(type="list", formula1='"Monthly,Twice a month,Every 2 weeks,Weekly"', allow_blank=False)
db.add_data_validation(dv)
dv.add("F4")
year_net = (f'SUMIFS({LG("D")},{LG("B")},">="&DATE($C$4,1,1),{LG("B")},"<="&DATE($C$4,12,31))'
            f'-SUMIFS({LG("E")},{LG("B")},">="&DATE($C$4,1,1),{LG("B")},"<="&DATE($C$4,12,31))')
tile(db, "B", 6, "Saved so far", f"=SUM({GO('G')})", MONEY)
tile(db, "C", 6, "Target", f"=SUM({GO('C')})", MONEY)
tile(db, "D", 6, "Left to save", f"=SUM({GO('H')})", MONEY, "B4541F")
tile(db, "E", 6, "Progress", "=IF(C7=0,0,MIN(1,B7/C7))", "0%")
tile(db, "F", 6, "Put in this year", f"={year_net}", MONEY, span=2)
tile(db, "I", 6, "Needed a month", f"=SUM({GO('K')})", MONEY, span=2)
tile(db, "K", 6, "Per payday", f"=I7/{PER_PAYDAY.replace(FREQ, '$F$4')}", MONEY, span=2)

header(db, 9, 2, ["Goal", "Saved", "Target", "Progress", "", "Status"])
db.merge_cells("E9:F9")
for k in range(O1 - O0 + 1):
    r, gr = 10 + k, O0 + k
    style(db.cell(r, 2, f'=IF(Goals!B{gr}="","",Goals!B{gr})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",Goals!G{gr})'), False, MONEY)
    style(db.cell(r, 4, f'=IF(B{r}="","",Goals!C{gr})'), False, MONEY)
    c = style(db.cell(r, 5, f'=IF(OR(B{r}="",Goals!I{gr}=""),"",{bar(f"Goals!I{gr}", 14)})'), False)
    c.font = Font(name=F, size=10, color=TEAL)
    style(db.cell(r, 6, f'=IF(OR(B{r}="",Goals!I{gr}=""),"",Goals!I{gr})'), False, "0%", align="center")
    style(db.cell(r, 7, f'=IF(B{r}="","",Goals!M{gr})'), False).alignment = Alignment(horizontal="left", indent=1)
status_colours(db, f"G10:G{10 + O1 - O0}", "$G10")

behind = f'COUNTIFS({GO("M")},"Short by*")+COUNTIFS({GO("M")},"Past the date")'
reached = f'COUNTIFS({GO("M")},"Reached")'
nxt = f'MIN({GO("O")})'
nxt_goal = f'INDEX({GO("B")},MATCH({nxt},{GO("O")},0))'
nice = lambda d: f'DAY({d})&" "&CHOOSE(MONTH({d}),{MONTH_LIST})&" "&YEAR({d})'
weeks = '(COUNTIF(\'52 weeks\'!E12:E37,"<>")+COUNTIF(\'52 weeks\'!J12:J37,"<>"))'
filled = "COUNTIF('100 envelopes'!P11:Y20,\">0\")"
notes = [
    (f'=IF({behind}=0,"Every goal with a date is on track",IF({behind}=1,"1 goal is behind its plan",'
     f'{behind}&" goals are behind their plan"))', BAD, "Every"),
    (f'=IF({reached}=0,"No goal reached yet",IF({reached}=1,"1 goal reached",{reached}&" goals reached"))', OK, None),
    (f'=IF({nxt}=0,"No target date coming up","Next target date: "&{nxt_goal}&", "&{nice(nxt)})', INFO, None),
    (f'="52 week challenge: "&{weeks}&" of 52 weeks done"', INFO, None),
    (f'="100 envelopes: "&{filled}&" of "&\'100 envelopes\'!E4&" filled"', INFO, None),
]
db.merge_cells("I9:L9")
db["I9"] = "Needs attention"
db["I9"].font = font(11, True, "FFFFFF")
db["I9"].fill = fill(TEAL)
db["I9"].alignment = Alignment(vertical="center")
for k, (formula, colour, calm) in enumerate(notes):
    r = 10 + k
    db.merge_cells(f"I{r}:L{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "IJKL":
        db[f"{col}{r}"].border = box
    if calm:
        db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))

header(db, 16, 9, ["Month", "Put in", "Taken out", "Net"])
db.row_dimensions[16].height = None
for m in range(12):
    r = 17 + m
    ms, me = f"DATE($C$4,{m + 1},1)", f"EOMONTH(DATE($C$4,{m + 1},1),0)"
    when = f'{LG("B")},">="&{ms},{LG("B")},"<="&{me}'
    style(db.cell(r, 9, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 10, f'=SUMIFS({LG("D")},{when})'), False, MONEY)
    style(db.cell(r, 11, f'=SUMIFS({LG("E")},{when})'), False, MONEY)
    style(db.cell(r, 12, f"=J{r}-K{r}"), False, NET, True)
style(db.cell(29, 9, "Year"), False, bold=True, align="center")
for col in "JKL":
    style(db[f"{col}29"], False, NET if col == "L" else MONEY, True)
    db[f"{col}29"] = f"=SUM({col}17:{col}28)"
db.conditional_formatting.add("I17:L28", FormulaRule(formula=["AND($C$4=YEAR(TODAY()),ROW()-16=MONTH(TODAY()))"],
                                                     fill=fill("E6F2F2")))

chart = BarChart()
chart.type = "col"
chart.title = "Put in and taken out"
chart.height, chart.width = 7.5, 17
chart.add_data(Reference(db, min_col=10, max_col=11, min_row=16, max_row=28), titles_from_data=True)
chart.set_categories(Reference(db, min_col=9, min_row=17, max_row=28))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E0A21B"
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "B27")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Dashboard: pick the year to show and how often you are paid."),
    ("2", "Goals: one line per goal, with the target, the date if there is one, what you had saved before, and your plan per month."),
    ("3", "Savings log: one line each time you put money in or take it out, with its goal."),
    ("4", "Goals and Dashboard: what is saved, what is left, and how much each goal needs per month and per payday."),
    ("5", "52 weeks and 100 envelopes: two savings challenges, if you like them. Type x for a week, or an envelope's number."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["A goal is on track when your plan per month covers what it still needs before its date.",
                          "Money for a challenge that goes into a goal can go in the Savings log too.",
                          "The example goals and payments show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Savings-Goals-Tracker.xlsx")
wb.save(out)
print("saved", out, len(goals), "goals", len(log), "log lines", len(weeks_done), "weeks", len(envelopes), "envelopes")
