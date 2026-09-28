"""Subscription Tracker (Excel + Google Sheets): find forgotten subscriptions,
catch free trials before they charge, see the yearly total.

Same compatibility rules as the budget: SUMIFS, COUNTIFS, EDATE, DATEDIF,
ROUNDUP, REPT, TEXT, IF, AND, OR, TODAY. No macros, no dynamic arrays.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
CATEGORIES = ["Streaming", "Music", "Apps", "Gaming", "Fitness", "News", "Cloud", "Shopping", "Other"]
BILLING = ["Weekly", "Monthly", "Quarterly", "Yearly"]
FIRST, LAST = 6, 55  # subscription rows

wb = Workbook()

# ---------------------------------------------------------------- Subscriptions
sub = wb.active
sub.title = "Subscriptions"
sheet_base(sub, "Your subscriptions", "Type in the yellow cells. Dates and totals work themselves out.",
           [3, 22, 13, 11, 12, 14, 9, 14, 10, 12, 12, 34])
header(sub, 5, 2, ["Name", "Category", "Cost", "Billing", "Start or last renewal", "Free trial?",
                   "Last time I used it", "Keep?", "Per month", "Per year", "What's next"])
d = lambda k: today + dt.timedelta(days=k)
sample = [
    ("Video streaming", "Streaming", 15.49, "Monthly", d(-28), "No", d(-1), "Keep"),
    ("Second video service", "Streaming", 9.99, "Monthly", d(-12), "No", d(-41), "Decide"),
    ("Music app", "Music", 10.99, "Monthly", d(-3), "No", d(0), "Keep"),
    ("Cloud storage 2 TB", "Cloud", 99.99, "Yearly", d(-360), "No", d(-2), "Keep"),
    ("Language app", "Apps", 12.99, "Monthly", d(-20), "No", d(-63), "Cancel"),
    ("Meal kit", "Shopping", 59.9, "Weekly", d(-5), "Yes", d(-5), "Decide"),
    ("Photo editor", "Apps", 4.99, "Monthly", d(-28), "Yes", d(-10), "Decide"),
    ("Gym", "Fitness", 29.9, "Monthly", d(-9), "No", d(-19), "Keep"),
    ("News site", "News", 36, "Quarterly", d(-80), "No", d(-33), "Decide"),
    ("Game pass", "Gaming", 14.99, "Monthly", d(-15), "No", d(-4), "Keep"),
]
for i in range(FIRST, LAST + 1):
    s = sample[i - FIRST] if i - FIRST < len(sample) else (None,) * 8
    style(sub.cell(i, 2, s[0]), True)
    style(sub.cell(i, 3, s[1]), True)
    style(sub.cell(i, 4, s[2]), True, MONEY)
    style(sub.cell(i, 5, s[3]), True, align="center")
    style(sub.cell(i, 6, s[4]), True, DATE)
    style(sub.cell(i, 7, s[5]), True, align="center")
    style(sub.cell(i, 8, s[6]), True, DATE)
    style(sub.cell(i, 9, s[7]), True, align="center")
    factor = f'IF(E{i}="Weekly",52/12,IF(E{i}="Monthly",1,IF(E{i}="Quarterly",1/3,IF(E{i}="Yearly",1/12,0))))'
    style(sub.cell(i, 10, f'=IF(B{i}="","",D{i}*{factor})'), False, MONEY)
    style(sub.cell(i, 11, f'=IF(B{i}="","",J{i}*12)'), False, MONEY)
    # Next charge after today, from the start date and the billing period.
    months = f'IF(E{i}="Monthly",1,IF(E{i}="Quarterly",3,12))'
    nxt = (f'IF(F{i}>=TODAY(),F{i},IF(E{i}="Weekly",F{i}+7*ROUNDUP((TODAY()-F{i})/7,0),'
           f'EDATE(F{i},{months}*ROUNDUP((DATEDIF(F{i},TODAY(),"m")+1)/{months},0))))')
    days = f'({nxt}-TODAY())'
    n_days = f'{days}&IF({days}=1," day"," days")'
    status = (f'=IF(OR(B{i}="",F{i}=""),"",IF(I{i}="Cancel","Cancel within "&{n_days},'
              f'IF(G{i}="Yes","Trial ends in "&{n_days},'
              f'IF(AND(H{i}<>"",TODAY()-H{i}>30),"Not used for "&(TODAY()-H{i})&" days",'
              f'"Renews in "&{n_days}))))')
    style(sub.cell(i, 12, status), False)
sub.freeze_panes = "C6"
for col, options in (("C", CATEGORIES), ("E", BILLING), ("G", ["Yes", "No"]), ("I", ["Keep", "Cancel", "Decide"])):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    sub.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")
rng = f"L{FIRST}:L{LAST}"
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},6)="Cancel"'], fill=fill(INFO)))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},5)="Trial"'], fill=fill(BAD),
                                                font=Font(name=F, bold=True, color="9B1C1C")))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},8)="Not used"'], fill=fill(WARN)))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},6)="Renews"'], fill=fill(OK)))

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "What your subscriptions cost", "Updated every time you open it.", [3, 26, 16, 16, 16, 34, 3, 40])
R = lambda col: f"Subscriptions!${col}${FIRST}:${col}${LAST}"
tiles = [
    ("B", "Per month", f"=SUM({R('J')})", TEAL),
    ("C", "Per year", f"=SUM({R('K')})", "B4541F"),
    ("D", "Subscriptions", f'=COUNTIFS({R("B")},"<>")', TEAL),
    ("E", "Marked to cancel", f'=SUMIFS({R("K")},{R("I")},"Cancel")', "2F6B3A"),
]
for col, label, formula, colour in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = "0" if col == "D" else MONEY
    db[f"{col}7"].font = font(18, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db["E8"] = "saved per year"
db["E8"].font = font(9, color=MUTED, italic=True)
db.row_dimensions[7].height = 30
db["F6"] = "Unused for 30+ days"
db["F6"].font = font(11, True, "FFFFFF")
db["F6"].fill = fill(TEAL_D)
db["F6"].alignment = Alignment(horizontal="center")
db["F7"] = f'=SUMIFS({R("K")},{R("B")},"<>",{R("H")},"<"&(TODAY()-30))'
db["F7"].number_format = MONEY
db["F7"].font = font(24, True, "FFFFFF")
db["F7"].fill = fill(TEAL_D)
db["F7"].alignment = Alignment(horizontal="center", vertical="center")
db["F8"] = "per year on things you are not using"
db["F8"].font = font(9, color=MUTED, italic=True)
db["F8"].alignment = Alignment(horizontal="center")
db.row_dimensions[7].height = 40

header(db, 10, 2, ["Category", "Per month", "Per year", "Count", "Share of total"])
for k, cat in enumerate(CATEGORIES):
    r = 11 + k
    style(db.cell(r, 2, cat), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({R("J")},{R("C")},B{r})'), False, MONEY)
    style(db.cell(r, 4, f"=C{r}*12"), False, MONEY)
    style(db.cell(r, 5, f'=COUNTIFS({R("C")},B{r})'), False, "0", align="center")
    share = f"IF($B$7=0,0,C{r}/$B$7)"
    c = style(db.cell(r, 6, f'={bar(share)}&"  "&TEXT({share},"0%")'), False)
    c.font = Font(name=F, size=11, color=TEAL)
    db.row_dimensions[r].height = 22

db["H10"] = "Heads-up"
db["H10"].alignment = Alignment(vertical="center")
db["H10"].font = font(11, True, "FFFFFF")
db["H10"].fill = fill(TEAL)
trial_n = f'COUNTIFS({R("L")},"Trial*")'
soon_n = "(" + "+".join(f'COUNTIFS({R("L")},"Renews in {k} day*")' for k in range(4)) + ")"
idle_n = f'COUNTIFS({R("L")},"Not used*")'
open_n = f'COUNTIFS({R("I")},"Decide")'
notes = [
    (f'=IF({trial_n}=0,"No free trial is about to charge",IF({trial_n}=1,"1 free trial will start charging soon",'
     f'{trial_n}&" free trials will start charging soon"))', BAD),
    (f'=IF({soon_n}=0,"No renewals in the next 3 days",IF({soon_n}=1,"1 renewal",{soon_n}&" renewals")&" in the next 3 days")', WARN),
    (f'=IF({idle_n}=0,"Everything was used in the last month",IF({idle_n}=1,"1 subscription has",{idle_n}&" subscriptions have")'
     f'&" not been used for over a month")', WARN),
    (f'=IF({open_n}=0,"Nothing left to decide",IF({open_n}=1,"1 subscription",{open_n}&" subscriptions")'
     f'&" still to decide: keep or cancel")', INFO),
]
for k, (formula, colour) in enumerate(notes):
    c = db.cell(11 + k * 2, 8, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.border = box
    c.alignment = Alignment(vertical="center")
for cell, prefix in (('H11', 'No'), ('H13', 'No'), ('H15', 'Everything'), ('H17', 'Nothing')):
    db.conditional_formatting.add(cell, FormulaRule(formula=[f'LEFT({cell},{len(prefix)})="{prefix}"'], fill=fill(OK)))

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 100])
steps = [
    ("1", "Subscriptions tab: add every subscription. Your bank statement and phone app store help you find them."),
    ("2", "Start or last renewal: the date you were last charged. The next charge is worked out for you."),
    ("3", "Free trial: choose Yes and the status turns red, with the days left before the first charge."),
    ("4", "Last time I used it: update it now and then. After 30 days without use the status turns orange."),
    ("5", "Keep: choose Keep, Cancel or Decide. The Dashboard shows what cancelling saves per year."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
    st.row_dimensions[r].height = 24
st["C17"] = "The example lines show how it works. Delete them and add your own."
st["C17"].font = font(11, color=MUTED, italic=True)
st["C19"] = "Works in Google Sheets (upload it to Drive, then File > Save as Google Sheets) and in Microsoft Excel."
st["C19"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Subscription-Tracker.xlsx")
wb.save(out)
print("saved", out)
