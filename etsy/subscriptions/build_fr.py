"""Suivi des abonnements (version française de build.py; Excel + Google Sheets): find forgotten subscriptions,
catch free trials before they charge, see the yearly total.

Same compatibility rules as the budget: SUMIFS, COUNTIFS, EDATE, DATEDIF,
ROUNDUP, REPT, TEXT, IF, AND, OR, TODAY. No macros, no dynamic arrays.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
CATEGORIES = ["Streaming", "Musique", "Applis", "Jeux", "Sport", "Presse", "Cloud", "Shopping", "Autre"]
BILLING = ["Hebdo", "Mensuel", "Trimestriel", "Annuel"]
FIRST, LAST = 6, 55  # subscription rows

wb = Workbook()

# ---------------------------------------------------------------- Subscriptions
sub = wb.active
sub.title = "Abonnements"
sheet_base(sub, "Tous vos abonnements, au même endroit", "Tapez dans les cases jaunes. Dates et totaux se calculent tout seuls.",
           [3, 22, 13, 11, 12, 14, 9, 14, 10, 12, 12, 34])
header(sub, 5, 2, ["Nom", "Catégorie", "Prix", "Facturation", "Début ou dernier prélèvement", "Essai gratuit ?",
                   "Dernière utilisation", "Garder ?", "Par mois", "Par an", "Et ensuite"])
d = lambda k: today + dt.timedelta(days=k)
sample = [
    ("Streaming vidéo", "Streaming", 15.49, "Mensuel", d(-28), False, d(-1), "Garder"),
    ("2e service vidéo", "Streaming", 9.99, "Mensuel", d(-12), False, d(-41), "Décider"),
    ("Appli musique", "Musique", 10.99, "Mensuel", d(-3), False, d(0), "Garder"),
    ("Stockage cloud 2 To", "Cloud", 99.99, "Annuel", d(-360), False, d(-2), "Garder"),
    ("Appli de langues", "Applis", 12.99, "Mensuel", d(-20), False, d(-63), "Annuler"),
    ("Box repas", "Shopping", 59.9, "Hebdo", d(-5), True, d(-5), "Décider"),
    ("Retouche photo", "Applis", 4.99, "Mensuel", d(-28), True, d(-10), "Décider"),
    ("Salle de sport", "Sport", 29.9, "Mensuel", d(-9), False, d(-19), "Garder"),
    ("Site d’info", "Presse", 36, "Trimestriel", d(-80), False, d(-33), "Décider"),
    ("Abonnement jeux", "Jeux", 14.99, "Mensuel", d(-15), False, d(-4), "Garder"),
]
for i in range(FIRST, LAST + 1):
    s = sample[i - FIRST] if i - FIRST < len(sample) else (None,) * 8
    style(sub.cell(i, 2, s[0]), True)
    style(sub.cell(i, 3, s[1]), True)
    style(sub.cell(i, 4, s[2]), True, MONEY)
    style(sub.cell(i, 5, s[3]), True, align="center")
    style(sub.cell(i, 6, s[4]), True, DATE)
    style(sub.cell(i, 7, s[5]), True, align="center")
    style(sub.cell(i, 8, s[6]), True, DATE)
    style(sub.cell(i, 9, s[7]), True, align="center")
    factor = f'IF(E{i}="Hebdo",52/12,IF(E{i}="Mensuel",1,IF(E{i}="Trimestriel",1/3,IF(E{i}="Annuel",1/12,0))))'
    style(sub.cell(i, 10, f'=IF(B{i}="","",D{i}*{factor})'), False, MONEY)
    style(sub.cell(i, 11, f'=IF(B{i}="","",J{i}*12)'), False, MONEY)
    # Next charge after today, from the start date and the billing period.
    months = f'IF(E{i}="Mensuel",1,IF(E{i}="Trimestriel",3,12))'
    nxt = (f'IF(F{i}>=TODAY(),F{i},IF(E{i}="Hebdo",F{i}+7*ROUNDUP((TODAY()-F{i})/7,0),'
           f'EDATE(F{i},{months}*ROUNDUP((DATEDIF(F{i},TODAY(),"m")+1)/{months},0))))')
    days = f'({nxt}-TODAY())'
    status = (f'=IF(OR(B{i}="",F{i}=""),"",IF(I{i}="Annuler","✗ À annuler sous "&{days}&" jour(s)",'
              f'IF(G{i}=TRUE,"⚠ Fin de l’essai dans "&{days}&" jour(s)",'
              f'IF(AND(H{i}<>"",TODAY()-H{i}>30),"💤 Pas utilisé depuis "&(TODAY()-H{i})&" jours : ça vaut encore le coup ?",'
              f'"Renouvellement dans "&{days}&" jour(s)"))))')
    style(sub.cell(i, 12, status), False)
sub.freeze_panes = "C6"
for col, options in (("C", CATEGORIES), ("E", BILLING), ("G", ["TRUE", "FALSE"]), ("I", ["Garder", "Annuler", "Décider"])):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    sub.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")
rng = f"L{FIRST}:L{LAST}"
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},1)="✗"'], fill=fill(INFO)))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},1)="⚠"'], fill=fill(BAD),
                                                font=Font(name=F, bold=True, color="9B1C1C")))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},2)="💤"'], fill=fill(WARN)))
sub.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($L{FIRST},6)="Renouv"'], fill=fill(OK)))

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Tableau de bord", 0)
sheet_base(db, "Où part l’argent chaque mois", "Mis à jour à chaque ouverture.", [3, 26, 16, 16, 16, 34, 3, 40])
R = lambda col: f"Abonnements!${col}${FIRST}:${col}${LAST}"
tiles = [
    ("B", "Par mois", f"=SUM({R('J')})", TEAL),
    ("C", "Par an", f"=SUM({R('K')})", "B4541F"),
    ("D", "Abonnements", f'=COUNTIFS({R("B")},"<>")', TEAL),
    ("E", "À annuler", f'=SUMIFS({R("K")},{R("I")},"Annuler")', "2F6B3A"),
]
for col, label, formula, colour in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = "0" if col == "D" else MONEY
    db[f"{col}7"].font = font(18, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db["E8"] = "économisés par an"
db["E8"].font = font(9, color=MUTED, italic=True)
db.row_dimensions[7].height = 30
db["F6"] = "Inutilisés depuis 30+ jours"
db["F6"].font = font(11, True, "FFFFFF")
db["F6"].fill = fill(TEAL_D)
db["F6"].alignment = Alignment(horizontal="center")
db["F7"] = f'=SUMIFS({R("K")},{R("B")},"<>",{R("H")},"<"&(TODAY()-30))'
db["F7"].number_format = MONEY
db["F7"].font = font(24, True, "FFFFFF")
db["F7"].fill = fill(TEAL_D)
db["F7"].alignment = Alignment(horizontal="center", vertical="center")
db["F8"] = "par an pour des choses que vous n’utilisez pas"
db["F8"].font = font(9, color=MUTED, italic=True)
db["F8"].alignment = Alignment(horizontal="center")
db.row_dimensions[7].height = 40

header(db, 10, 2, ["Catégorie", "Par mois", "Par an", "Nombre", "Part du total"])
for k, cat in enumerate(CATEGORIES):
    r = 11 + k
    style(db.cell(r, 2, cat), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({R("J")},{R("C")},B{r})'), False, MONEY)
    style(db.cell(r, 4, f"=C{r}*12"), False, MONEY)
    style(db.cell(r, 5, f'=COUNTIFS({R("C")},B{r})'), False, "0", align="center")
    share = f"IF($B$7=0,0,C{r}/$B$7)"
    c = style(db.cell(r, 6, f'={bar(share)}&"  "&TEXT({share},"0%")'), False)
    c.font = Font(name=F, size=11, color=TEAL)
    db.row_dimensions[r].height = 22

db["H10"] = "À surveiller"
db["H10"].font = font(11, True, "FFFFFF")
db["H10"].fill = fill(TEAL)
notes = [
    (f'=COUNTIFS({R("L")},"⚠*")&" essai(s) gratuit(s) bientôt payant(s)"', BAD),
    (f'=COUNTIFS({R("L")},"Renouvellement dans 0 jour*")+COUNTIFS({R("L")},"Renouvellement dans 1 jour*")+COUNTIFS({R("L")},"Renouvellement dans 2 jour*")'
     f'+COUNTIFS({R("L")},"Renouvellement dans 3 jour*")&" renouvellement(s) dans les 3 jours"', WARN),
    (f'=COUNTIFS({R("L")},"💤*")&" abonnement(s) pas utilisé(s) depuis un mois"', WARN),
    (f'=COUNTIFS({R("I")},"Décider")&" encore à décider : garder ou annuler ?"', INFO),
]
for k, (formula, colour) in enumerate(notes):
    c = db.cell(11 + k * 2, 8, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.border = box
    c.alignment = Alignment(vertical="center")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Commencer ici")
sheet_base(st, "Commencer ici : 2 minutes", "Tapez uniquement dans les cases JAUNES.", [3, 6, 100])
steps = [
    ("1", "Onglet Abonnements : ajoutez chaque abonnement (relevé bancaire + magasin d’applis du téléphone pour les retrouver)."),
    ("2", "Début ou dernier prélèvement : la date du dernier débit. Le prochain est calculé pour vous."),
    ("3", "Essai gratuit ? Mettez TRUE : alerte rouge avec les jours restants avant le premier débit."),
    ("4", "Dernière utilisation : mettez-la à jour de temps en temps. Plus de 30 jours = orange."),
    ("5", "Garder ? Choisissez Garder, Annuler ou Décider. Le Tableau de bord montre l’économie par an."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
    st.row_dimensions[r].height = 24
st["C17"] = "Les lignes d’exemple montrent comment ça marche. Supprimez-les et ajoutez les vôtres."
st["C17"].font = font(11, color=MUTED, italic=True)
st["C19"] = "Fonctionne dans Google Sheets (importez dans Drive, puis Fichier → Enregistrer au format Google Sheets) et dans Microsoft Excel."
st["C19"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Abonnements-Suivi.xlsx")
wb.save(out)
print("saved", out)
