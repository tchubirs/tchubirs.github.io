"""Party and Event Planner (Excel + Google Sheets).

A countdown to the day; a checklist whose dates count back from the party; the guest list with answers, adults
and kids, dietary needs, gifts and thank-you notes; a food and drink calculator that says how much to buy for the
guests who may come; the budget with deposits and what is still to pay; and the plan for the day itself.
Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
G0, G1 = 6, 305        # guests
T0, T1 = 6, 65         # tasks
F0, F1 = 8, 47         # food and drinks
B0, B1 = 6, 55         # budget
S0, S1 = 6, 40         # the day
MON_LIST = ",".join(f'"{m}"' for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
EV = "Dashboard!$C$5"
GU = lambda col: f"Guests!${col}${G0}:${col}${G1}"
TA = lambda col: f"Tasks!${col}${T0}:${col}${T1}"
FO = lambda col: f"'Food and drinks'!${col}${F0}:${col}${F1}"
BU = lambda col: f"Budget!${col}${B0}:${col}${B1}"
GROUPS = ["Family", "Friends", "School", "Work", "Neighbours", "Other"]
ANSWERS = ["Yes", "Maybe", "No"]
DIETS = ["Vegetarian", "Vegan", "Gluten free", "Dairy free", "Nut allergy", "Halal", "Kosher", "Other"]
CATS = ["Venue", "Cake", "Decorations", "Entertainment", "Rentals", "Invitations", "Favours", "Photos", "Outfits", "Other"]
KINDS = ["Food", "Drinks", "Supplies"]


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


# ---------------------------------------------------------------- example data (a made-up party)
party = today + dt.timedelta(days=24 + (5 - (today + dt.timedelta(days=24)).weekday()) % 7)   # a Saturday, 24 to 30 days away
guests = [
    # name, group, answer, adults, kids, diet, gift, thank-you sent, note
    ("Grandma Rosa and Grandpa Joe", "Family", "Yes", 2, 0, None, None, None, "Bringing the photo album"),
    ("Aunt Clara and Ben", "Family", "Yes", 2, 2, "Vegetarian", None, None, None),
    ("Uncle Dan", "Family", "No", 1, 0, None, "Book and card by post", "x", "Away that weekend"),
    ("The Okafor family", "School", "Yes", 2, 2, None, None, None, None),
    ("The Lindqvist family", "School", "Yes", 1, 1, "Nut allergy", None, None, "Ask about the cake"),
    ("The Brandt family", "School", "Maybe", 2, 1, None, None, None, None),
    ("The Varga family", "School", "Yes", 2, 1, None, None, None, None),
    ("The Castell family", "School", "Yes", 1, 2, "Gluten free", None, None, None),
    ("The Holloway family", "School", "No", 2, 1, None, "Gift card", None, None),
    ("The Quill family", "School", None, 2, 1, None, None, None, None),
    ("The Rook family", "School", "Yes", 2, 1, None, None, None, None),
    ("The Ashby family", "School", None, 1, 1, None, None, None, None),
    ("Nadia and Sam", "Friends", "Yes", 2, 0, "Vegan", None, None, None),
    ("Iris", "Friends", "Yes", 1, 0, None, None, None, "Can help with the games"),
    ("Tomas and Ines", "Friends", "Maybe", 2, 1, None, None, None, None),
    ("The Marlowe family", "Neighbours", "Yes", 2, 2, None, None, None, None),
    ("Mr and Mrs Thorne", "Neighbours", None, 2, 0, None, None, None, None),
    ("Coach Wren", "Other", "Yes", 1, 0, "Dairy free", None, None, None),
    ("The Fennimore family", "School", "Yes", 2, 1, None, None, None, None),
    ("Jude", "Friends", None, 1, 0, None, None, None, None),
]
tasks = [
    # task, days before, who, done
    ("Set the date, the time and the budget", 56, "Mum", True), ("Book the venue", 56, "Mum", True),
    ("Make the guest list", 49, "Mum and Dad", True), ("Send the invitations", 42, "Dad", True),
    ("Book the entertainment", 42, "Dad", True), ("Order the cake", 35, "Mum", True),
    ("Choose a theme and decorations", 35, "Mia and Mum", True), ("Plan the menu", 30, "Dad", True),
    ("Order the party bags", 35, "Mum", False), ("Book a face painter", 28, "Mum", True),
    ("Rent extra tables and chairs", 21, "Dad", False), ("Follow up on the answers", 21, "Mum", False),
    ("Plan the games and the music", 18, "Iris", False), ("Buy balloons and a banner", 14, "Dad", False),
    ("Confirm the entertainment", 14, "Dad", False), ("Make the plan for the day", 12, "Mum", False),
    ("Final guest count for the food", 10, "Mum", False), ("Shop for drinks and snacks", 7, "Dad", False),
    ("Confirm the venue and the cake", 7, "Mum", False), ("Ask helpers for a job each", 5, "Mum", False),
    ("Charge the camera and the speaker", 2, "Dad", False), ("Buy fresh food and ice", 1, "Dad", False),
    ("Pick up the cake", 1, "Mum", False), ("Set up the decorations", 0, "Everyone", False),
    ("Send thank-you notes", -3, "Mia and Mum", False), ("Return the rentals", -2, "Dad", False),
    ("Share the photos", -4, "Mum", False),
]
food = [
    # item, kind, per adult, per kid, unit, pack size, price per pack, bought
    ("Pizza", "Food", 2, 2, "slices", 8, 11.00, False), ("Wraps and sandwiches", "Food", 1, 0.5, "pieces", 12, 12.00, False),
    ("Crisps and snacks", "Food", 0.5, 0.5, "bags", 6, 3.20, True), ("Fruit platter", "Food", 0.5, 1, "portions", 12, 10.00, False),
    ("Veggies and dip", "Food", 0.3, 0.2, "portions", 10, 9.50, False),
    ("Soft drinks", "Drinks", 2, 1, "cans", 24, 11.00, True), ("Juice boxes", "Drinks", 0, 2, "boxes", 10, 4.50, True),
    ("Water", "Drinks", 1, 1, "bottles", 24, 5.00, False), ("Coffee and tea", "Drinks", 1, 0, "cups", 40, 7.50, False),
    ("Ice", "Drinks", 0.05, 0.05, "bags", 1, 2.50, False),
    ("Plates", "Supplies", 2, 2, "plates", 50, 4.00, True), ("Cups", "Supplies", 3, 3, "cups", 50, 3.50, True),
    ("Napkins", "Supplies", 3, 3, "napkins", 100, 2.50, True),
]
budget = [
    # item, category, shop or vendor, estimated, actual, paid, due date for the rest
    ("Community hall, 3 hours", "Venue", "Rowan Street hall", 180, 180, 50, party - dt.timedelta(days=7)),
    ("Unicorn cake, 24 slices", "Cake", "Corner bakery", 60, 55, 20, party - dt.timedelta(days=1)),
    ("Balloons and a banner", "Decorations", "Party shop", 45, 38.60, 38.60, None),
    ("Magician, 45 minutes", "Entertainment", "Magic Max", 150, 150, 0, party),
    ("Face painting", "Entertainment", "Painted Faces", 90, 90, 20, party),
    ("Extra tables and chairs", "Rentals", "Hire shop", 35, None, 0, party - dt.timedelta(days=2)),
    ("Printed invitations", "Invitations", "Print shop", 18, 16.50, 16.50, None),
    ("Party bags", "Favours", None, 40, None, 0, party - dt.timedelta(days=21)),
]
day = [
    (dt.time(12, 30), "Set up the room, balloons and tables", "Dad and Iris"),
    (dt.time(13, 30), "Food and drinks out, music on", "Dad"),
    (dt.time(14, 0), "Guests arrive, face painting starts", "Mum"),
    (dt.time(14, 45), "Games in the hall", "Iris"), (dt.time(15, 30), "Magician", "Magic Max"),
    (dt.time(16, 15), "Pizza and snacks", "Everyone"), (dt.time(16, 45), "Cake and singing", "Mum"),
    (dt.time(17, 15), "Party bags and goodbyes", "Mia"), (dt.time(17, 30), "Tidy up and return the hall key", "Dad"),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Guests
gu = wb.active
gu.title = "Guests"
sheet_base(gu, "Guests", "One line per guest or family. The answers, head counts and dietary needs add up on the "
           "Dashboard.", [3, 28, 12, 10, 8, 8, 14, 22, 11, 13, 26], rows=G1 + 2)
header(gu, 5, 2, ["Name", "Group", "Answer", "Adults", "Kids", "Dietary needs", "Gift", "Card sent", "Thank you",
                  "Notes"])
for r in range(G0, G1 + 1):
    g = guests[r - G0] if r - G0 < len(guests) else (None,) * 9
    name, group, answer, adults, kids, diet, gift, thanked, note = g
    style(gu.cell(r, 2, name), True, bold=True)
    style(gu.cell(r, 3, group), True)
    style(gu.cell(r, 4, answer), True, align="center")
    style(gu.cell(r, 5, adults), True, "0", align="center")
    style(gu.cell(r, 6, kids), True, "0", align="center")
    style(gu.cell(r, 7, diet), True)
    style(gu.cell(r, 8, gift), True)
    style(gu.cell(r, 9, thanked), True, align="center")
    style(gu.cell(r, 10, f'=IF(H{r}="","",IF(I{r}="","To send","Sent"))'), False, align="center")
    style(gu.cell(r, 11, note), True)
    gu.cell(r, 13, f'=IF(B{r}="",0,IF(E{r}="",1,N(E{r})))')          # M: adults, 1 when left empty
    gu.cell(r, 14, f'=IF(B{r}="",0,N(F{r}))')                          # N: kids
    gu.cell(r, 15, f'=IF(B{r}="","",IF(D{r}="","No answer",D{r}))')    # O: answer, with no answer spelled out
for c in "LMNO":
    gu.column_dimensions[c].hidden = True
dropdown(gu, f'"{",".join(GROUPS)}"', f"C{G0}:C{G1}", strict=False)
dropdown(gu, f'"{",".join(ANSWERS)}"', f"D{G0}:D{G1}")
dropdown(gu, f'"{",".join(DIETS)}"', f"G{G0}:G{G1}", strict=False)
for answer, colour in (("Yes", OK), ("Maybe", WARN), ("No", "EEF1F1")):
    gu.conditional_formatting.add(f"D{G0}:D{G1}", FormulaRule(formula=[f'D{G0}="{answer}"'], fill=fill(colour)))
gu.conditional_formatting.add(f"B{G0}:K{G1}", FormulaRule(formula=[f'$D{G0}="No"'], font=Font(name=F, color="9AA7A9")))
gu.conditional_formatting.add(f"J{G0}:J{G1}", FormulaRule(formula=[f'J{G0}="To send"'], fill=fill(WARN)))
gu.freeze_panes = "C6"

# ---------------------------------------------------------------- Tasks
ta = wb.create_sheet("Tasks")
sheet_base(ta, "Tasks", "The checklist, counted back from the day of the party. Change the days before and the dates "
           "follow. 0 is the day itself, and a minus number is after it.", [3, 38, 9, 16, 16, 8, 12, 30], rows=T1 + 2)
header(ta, 5, 2, ["Task", "Days before", "Date", "Who", "Done", "Status", "Notes"])
for r in range(T0, T1 + 1):
    t = tasks[r - T0] if r - T0 < len(tasks) else (None,) * 4
    name, before, who, done = t
    style(ta.cell(r, 2, name), True, bold=True)
    style(ta.cell(r, 3, before), True, "0", align="center")
    style(ta.cell(r, 4, f'=IF(OR(B{r}="",C{r}="",{EV}=""),"",{EV}-C{r})'), False, "ddd d mmm", align="center")
    style(ta.cell(r, 5, who), True)
    style(ta.cell(r, 6, "x" if done else None), True, align="center")
    style(ta.cell(r, 7, f'=IF(B{r}="","",IF(F{r}<>"","Done",IF(D{r}="","",IF(D{r}<TODAY(),"Late",'
                        f'IF(D{r}-TODAY()<=7,"This week","Coming")))))'), False, align="center")
    style(ta.cell(r, 8), True)
    ta.cell(r, 10, f'=IF(AND(B{r}<>"",F{r}="",D{r}<>""),D{r}+ROW()/10000000,"")')      # J: order of what is left
ta.column_dimensions["J"].hidden = True
for status, colour in (("Done", OK), ("Late", BAD), ("This week", WARN)):
    ta.conditional_formatting.add(f"G{T0}:G{T1}", FormulaRule(formula=[f'G{T0}="{status}"'], fill=fill(colour)))
ta.conditional_formatting.add(f"B{T0}:E{T1}", FormulaRule(formula=[f'$F{T0}<>""'], font=Font(name=F, color="9AA7A9")))
ta.freeze_panes = "C6"

# ---------------------------------------------------------------- Food and drinks
fo = wb.create_sheet("Food and drinks")
sheet_base(fo, "Food and drinks", "How much to buy, for everyone who may come: the yes, the maybe and those who have not "
           "answered. Amounts per person are for the whole party.", [3, 24, 10, 10, 10, 10, 10, 10, 10, 11, 11, 8, 20],
           rows=F1 + 3)
fo["B4"] = "Adults"
fo["B5"] = "Kids"
fo["E4"] = "Extra, to be safe"
for c in ("B4", "B5", "E4"):
    fo[c].font = font(11, True)
style(fo["C4"], False, "0", True, "center")
fo["C4"] = f'=SUMIFS({GU("M")},{GU("D")},"<>No")'
style(fo["C5"], False, "0", True, "center")
fo["C5"] = f'=SUMIFS({GU("N")},{GU("D")},"<>No")'
style(fo["G4"], True, "0%", True, "center")
fo["G4"] = 0.1
header(fo, 7, 2, ["Item", "Kind", "Per adult", "Per kid", "Unit", "Needed", "Pack of", "Packs", "Price a pack", "Cost",
                  "Bought", "Where"])
for r in range(F0, F1 + 1):
    it = food[r - F0] if r - F0 < len(food) else (None,) * 8
    name, kind, adult, kid, unit, pack, price, bought = it
    style(fo.cell(r, 2, name), True, bold=True)
    style(fo.cell(r, 3, kind), True)
    style(fo.cell(r, 4, adult), True, "0.##", align="center")
    style(fo.cell(r, 5, kid), True, "0.##", align="center")
    style(fo.cell(r, 6, unit), True, align="center")
    style(fo.cell(r, 7, f'=IF(B{r}="","",ROUNDUP(ROUND((N(D{r})*$C$4+N(E{r})*$C$5)*(1+N($G$4)),6),0))'), False, "0", True, "center")
    style(fo.cell(r, 8, pack), True, "0", align="center")
    style(fo.cell(r, 9, f'=IF(OR(G{r}="",N(H{r})=0),"",ROUNDUP(ROUND(G{r}/H{r},6),0))'), False, "0", True, "center")
    style(fo.cell(r, 10, price), True, MONEY)
    style(fo.cell(r, 11, f'=IF(OR(I{r}="",J{r}=""),"",I{r}*J{r})'), False, MONEY)
    style(fo.cell(r, 12, "x" if bought else None), True, align="center")
    style(fo.cell(r, 13), True)
dropdown(fo, f'"{",".join(KINDS)}"', f"C{F0}:C{F1}", strict=False)
fo.conditional_formatting.add(f"B{F0}:K{F1}", FormulaRule(formula=[f'$L{F0}<>""'], font=Font(name=F, color="9AA7A9")))
fo.conditional_formatting.add(f"L{F0}:L{F1}", FormulaRule(formula=[f'L{F0}<>""'], fill=fill(OK)))
fo[f"B{F1 + 1}"] = "Total"
fo[f"B{F1 + 1}"].font = font(11, True)
style(fo.cell(F1 + 1, 11, f"=SUM(K{F0}:K{F1})"), False, MONEY, True)
fo["I4"] = "A common rule for drinks: one per guest for each hour, and one more."
fo["I4"].font = font(9, color=MUTED, italic=True)
fo["I5"] = "Needed and packs round up, so there is always enough."
fo["I5"].font = font(9, color=MUTED, italic=True)
fo.freeze_panes = "C8"

# ---------------------------------------------------------------- Budget
bu = wb.create_sheet("Budget")
sheet_base(bu, "Budget", "Everything you book or buy besides the food: what it costs, what you paid and when the rest "
           "is due.", [3, 28, 14, 18, 11, 11, 11, 11, 13, 12, 22], rows=B1 + 2)
header(bu, 5, 2, ["Item", "Category", "Shop or vendor", "Estimate", "Final price", "Paid", "Still to pay", "Pay by",
                  "Status", "Notes"])
for r in range(B0, B1 + 1):
    b = budget[r - B0] if r - B0 < len(budget) else (None,) * 7
    name, cat, shop, est, final, paid, due = b
    style(bu.cell(r, 2, name), True, bold=True)
    style(bu.cell(r, 3, cat), True)
    style(bu.cell(r, 4, shop), True)
    style(bu.cell(r, 5, est), True, MONEY)
    style(bu.cell(r, 6, final), True, MONEY)
    style(bu.cell(r, 7, paid), True, MONEY)
    style(bu.cell(r, 8, f'=IF(B{r}="","",MAX(0,IF(F{r}<>"",F{r},N(E{r}))-N(G{r})))'), False, MONEY)
    style(bu.cell(r, 9, due), True, DATE)
    style(bu.cell(r, 10, f'=IF(B{r}="","",IF(H{r}<=0,"Paid",IF(I{r}="","To pay",IF(I{r}<TODAY(),"Late",'
                         f'IF(I{r}-TODAY()<=7,"Due soon","To pay")))))'), False, align="center")
    style(bu.cell(r, 11), True)
    bu.cell(r, 13, f'=IF(B{r}="",0,IF(F{r}<>"",F{r},N(E{r})))')        # M: the cost that counts
bu.column_dimensions["M"].hidden = True
dropdown(bu, f'"{",".join(CATS)}"', f"C{B0}:C{B1}", strict=False)
for status, colour in (("Paid", OK), ("Late", BAD), ("Due soon", WARN)):
    bu.conditional_formatting.add(f"J{B0}:J{B1}", FormulaRule(formula=[f'J{B0}="{status}"'], fill=fill(colour)))
bu[f"B{B1 + 1}"] = "The final price counts once you know it. Until then the estimate does."
bu[f"B{B1 + 1}"].font = font(9, color=MUTED, italic=True)
bu.freeze_panes = "C6"

# ---------------------------------------------------------------- The day
td = wb.create_sheet("The day")
sheet_base(td, "The day", "The plan for the party itself, hour by hour, and who looks after what.", [3, 10, 44, 20, 34],
           rows=S1 + 2)
header(td, 5, 2, ["Time", "What happens", "Who", "Notes"])
for r in range(S0, S1 + 1):
    d = day[r - S0] if r - S0 < len(day) else (None,) * 3
    style(td.cell(r, 2, d[0]), True, "h:mm", align="center")
    style(td.cell(r, 3, d[1]), True, bold=True)
    style(td.cell(r, 4, d[2]), True)
    style(td.cell(r, 5), True)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Party planner", "The countdown, the guests, the checklist and the money for the party.",
           [3, 18, 14, 14, 14, 14, 14, 3, 26, 11, 11, 11, 3], rows=45)
db["B2"] = '=IF(C4="","Party planner",C4)'
for r, label in ((4, "Party"), (5, "Date"), (6, "Budget")):
    db[f"B{r}"] = label
    db[f"B{r}"].font = font(11, True)
style(db["C4"], True, bold=True)
db["C4"] = "Mia turns 7"
db.merge_cells("C4:E4")
style(db["C5"], True, DATE, True, "center")
db["C5"] = party
db["D5"] = "Starts"
db["D5"].font = font(11, True)
db["D5"].alignment = Alignment(horizontal="right")
style(db["E5"], True, "h:mm", True, "center")
db["E5"] = dt.time(14, 0)
style(db["C6"], True, MONEY, True, "center")
db["C6"] = 1100
planned = f'SUM({BU("M")})+SUM({FO("K")})'
paid = f'SUM({BU("G")})+SUMIFS({FO("K")},{FO("L")},"<>")'
coming = f'SUMIFS({GU("M")},{GU("D")},"Yes")+SUMIFS({GU("N")},{GU("D")},"Yes")'
waiting = f'SUMIFS({GU("M")},{GU("O")},"No answer")+SUMIFS({GU("N")},{GU("O")},"No answer")'
tile(db, "B", 8, "Countdown", f'=IF($C$5="","",IF($C$5-TODAY()>1,($C$5-TODAY())&" days",IF($C$5-TODAY()=1,"Tomorrow",'
                              f'IF($C$5=TODAY(),"Today","Done"))))', "@")
tile(db, "C", 8, "Coming", f"={coming}", "0")
tile(db, "D", 8, "No answer yet", f"={waiting}", "0", "B4541F")
tile(db, "E", 8, "Tasks done", f'=COUNTIFS({TA("B")},"<>",{TA("F")},"<>")&" of "&COUNTIF({TA("B")},"?*")', "@")
tile(db, "F", 8, "Late tasks", f'=COUNTIF({TA("G")},"Late")', "0", "B4541F")
tile(db, "G", 8, "Budget left", f"=IF(N($C$6)=0,\"\",$C$6-({planned}))", MONEY)
db.conditional_formatting.add("G9", FormulaRule(formula=["AND(ISNUMBER(G9),G9<0)"], font=Font(name=F, size=16, bold=True,
                                                                                               color="B4541F")))

header(db, 11, 2, ["Next up", "", "", "Date", "Who", "Status"])
db.merge_cells("B11:D11")
for i in range(8):
    r = 12 + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({TA("J")},{i + 1}),"")'
    row = f'MATCH({key},{TA("J")},0)'
    none = 'IF($C$5="","Type the date of the party first","Nothing left to do")' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({TA("B")},{row}))'), False, bold=True)
    for c in (3, 4):
        style(db.cell(r, c), False)
    db.merge_cells(f"B{r}:D{r}")
    style(db.cell(r, 5, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 6, f'=IF({key}="","",INDEX({TA("E")},{row}))'), False)
    style(db.cell(r, 7, f'=IF({key}="","",INDEX({TA("G")},{row}))'), False, align="center")
for status, colour in (("Late", BAD), ("This week", WARN)):
    db.conditional_formatting.add("G12:G19", FormulaRule(formula=[f'G12="{status}"'], fill=fill(colour)))

header(db, 11, 9, ["Guests", "Adults", "Kids", "Total"])
for k, (label, answer) in enumerate((("Coming", "Yes"), ("Maybe", "Maybe"), ("No answer yet", "No answer"),
                                     ("Can't come", "No"))):
    r = 12 + k
    style(db.cell(r, 9, label), False, bold=True)
    style(db.cell(r, 10, f'=SUMIFS({GU("M")},{GU("O")},"{answer}")'), False, "0", align="center")
    style(db.cell(r, 11, f'=SUMIFS({GU("N")},{GU("O")},"{answer}")'), False, "0", align="center")
    style(db.cell(r, 12, f"=J{r}+K{r}"), False, "0", True, "center")
style(db.cell(16, 9, "Invited"), False, bold=True)
for c in (10, 11, 12):
    style(db.cell(16, c, f"=SUM({L(c)}12:{L(c)}15)"), False, "0", True, "center")
db.conditional_formatting.add("I12:L12", FormulaRule(formula=["TRUE"], fill=fill(OK)))
db["I17"] = f'="Thank-you notes to send: "&COUNTIF({GU("J")},"To send")'
db["I17"].font = font(10, True, TEAL_D)

header(db, 21, 2, ["Money", "Planned", "Paid", "Still to pay"])
for k, cat in enumerate(CATS):
    r = 22 + k
    style(db.cell(r, 2, cat), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({BU("M")},{BU("C")},B{r})'), False, "#,##0.00;;")
    style(db.cell(r, 4, f'=SUMIFS({BU("G")},{BU("C")},B{r})'), False, "#,##0.00;;")
    style(db.cell(r, 5, f'=SUMIFS({BU("H")},{BU("C")},B{r})'), False, "#,##0.00;;")
r = 22 + len(CATS)
style(db.cell(r, 2, "Food and drinks"), False, bold=True)
style(db.cell(r, 3, f'=SUM({FO("K")})'), False, "#,##0.00;;")
style(db.cell(r, 4, f'=SUMIFS({FO("K")},{FO("L")},"<>")'), False, "#,##0.00;;")
style(db.cell(r, 5, f"=C{r}-D{r}"), False, "#,##0.00;;")
style(db.cell(r + 1, 2, "Total"), False, bold=True)
for c in (3, 4, 5):
    style(db.cell(r + 1, c, f"=SUM({L(c)}22:{L(c)}{r})"), False, MONEY, True)
db[f"B{r + 2}"] = f'=IF(N($C$6)=0,"",{bar(f"MIN(1,C{r + 1}/$C$6)", 24)})'
db[f"B{r + 2}"].font = Font(name=F, size=12, color=TEAL)
db.merge_cells(f"B{r + 2}:D{r + 2}")
db[f"E{r + 2}"] = f'=IF(N($C$6)=0,"",IF(C{r + 1}>$C$6,"Over by "&FIXED(C{r + 1}-$C$6,2),FIXED(C{r + 1}/$C$6*100,0)&"% of the budget"))'
db[f"E{r + 2}"].font = font(10, True, TEAL_D)

header(db, 21, 9, ["Dietary needs of guests coming", "Guests"])
for k, diet in enumerate(DIETS):
    r2 = 22 + k
    style(db.cell(r2, 9, diet), False, bold=True)
    style(db.cell(r2, 10, f'=COUNTIFS({GU("G")},I{r2},{GU("D")},"Yes")'), False, "0;;", align="center")
db.column_dimensions["O"].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Dashboard: the name of the party, the date, the time it starts and your budget."),
    ("2", "Guests: one line per guest or family, with their answer and how many adults and kids."),
    ("3", "Tasks: a checklist is ready, counted back from the date. Put an x under Done as you go."),
    ("4", "Food and drinks: how much each person eats or drinks, and the size and price of a pack. It says what to buy."),
    ("5", "Budget and The day: what you book or buy, what you paid, and the plan for the party hour by hour."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example party, its guests and its shops are made up. Delete them and add your own.",
                          "Change the date and every task date moves with it.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Guests", "Tasks", "Food and drinks", "Budget", "The day", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Party-Planner.xlsx")
wb.save(out)
print("saved", out, "party on", party, len(guests), "guests", len(tasks), "tasks")
