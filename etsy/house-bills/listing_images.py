"""Etsy listing photos for the Roommate Bill Splitter (run build.py first)."""
import os, sys
from datetime import timedelta
from openpyxl import load_workbook
from PIL import Image
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Roommate-Bill-Splitter.xlsx")
tmp = os.path.join(HERE, "out")
r = Renders(xlsx, tmp)
# Payments close up: the last three months, from the first line of the month two months back.
pay = load_workbook(xlsx)["Payments"]
dates = [(row, pay.cell(row, 2).value) for row in range(6, 606) if pay.cell(row, 2).value]
last = max(d for _, d in dates)
start = last.replace(day=1)
for _ in range(2):
    start = (start - timedelta(days=1)).replace(day=1)
first_row = min(row for row, d in dates if d >= start)
close = Renders(xlsx, tmp, hide={"Bills": ["M", "N", "O"], "Payments": ["K", "L", "M"], "Chores": ["I", "J"]},
                hide_rows={"Bills": range(15, 28), "Chores": range(11, 16),
                           "Payments": list(range(6, first_row)) + list(range(dates[-1][0] + 1, 606))})
img_dir = os.path.join(HERE, "images")

img = canvas(dark=True)
title(img, "Roommate bill splitter", "Who paid what, who owes whom, the bills due this month and a chore rota",
      dark=True, width=2300)
laptop(img, r.shot("House bills"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

# The dashboard in two parts, cut in the blank row between the balances and the bills of the month.
page = r.doc[r._index("House bills")]
cut = int((page.search_for("Balance counts everything")[0].y1 + page.search_for("Category")[0].y0) / 2 * 220 / 72)
full = r.page("House bills")
top = Image.new("RGB", (full.width, cut + 120), "#FFFFFF")
top.paste(full.crop((0, 0, full.width, cut)), (0, 0))
detail(os.path.join(img_dir, "02-who-owes-whom.jpg"), "Who owes whom",
       "Everyone's share and balance, and the fewest payments to settle up",
       content(top), [("#D8F0DC", "Gets money back"), ("#FFE2B8", "Owes")])

detail(os.path.join(img_dir, "03-payments.jpg"), "One line per payment",
       "Leave the shares empty to use the bill's split, or put an x under the people who share it",
       close.shot("Payments"), [("#CFE8E4", "Shared only by the people marked")])

detail(os.path.join(img_dir, "04-split.jpg"), "Split rent by room",
       "Each bill split equally, by shares like the size of each room, or only between some of you",
       close.shot("Bills"), [("#CFE8E4", "Shares and marks"), ("#FFE2B8", "Due within 3 days"), ("#F9C9C4", "Late")])

detail(os.path.join(img_dir, "05-chores.jpg"), "A chore rota\nthat moves every week",
       "Everyone who lives in the house takes a turn",
       close.shot("Chores"))

img = canvas()
y = title(img, "Inside the file", "Six tabs and a Start Here page. You type in the yellow cells and the rest is "
          "calculated.")
grid(img, [
    (content(r.page("House bills")), "Dashboard", "Pick a month"),
    (content(r.page("Payments")), "Payments", "One line per bill paid"),
    (content(r.page("Bills")), "Bills", "Due dates and splits"),
    (content(r.page("Chores")), "Chores", "A weekly rota"),
    (content(r.page("Year")), "Year", "Months and categories"),
    (content(r.page("Housemates")), "Housemates", "Moving in and out"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
