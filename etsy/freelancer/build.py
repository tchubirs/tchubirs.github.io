"""Freelancer Income, Expenses & Tax Set-Aside (Excel + Google Sheets).

One line per invoice, one line per expense. The Dashboard shows profit, how
much to put aside for tax (your own rate), what is already put aside, unpaid
and overdue invoices, and a quarter-by-quarter view for estimated taxes.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
EXP_CATS = ["Software & tools", "Equipment", "Travel", "Phone & internet", "Office & coworking",
            "Training", "Marketing", "Bank & fees", "Other"]
I0, I1 = 6, 305   # income rows
E0, E1 = 6, 505   # expense rows
T0, T1 = 6, 105   # tax transfer rows

wb = Workbook()

# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Set these once a year.", [3, 40, 16, 4, 60])
rows = [(6, "Year", Y, "0"), (8, "Share of profit to put aside for tax (%)", 25, "0"),
        (10, "Invoice is overdue after (days)", 30, "0")]
for r, label, value, fmt in rows:
    se.cell(r, 2, label).font = font(12, True, TEAL_D)
    style(se.cell(r, 3, value), True, fmt, bold=True)
se["E8"] = "Ask your accountant or tax office for your rate; 25-30% is a common cautious choice."
se["E8"].font = font(10, color=MUTED, italic=True)
se["B13"] = "Expense categories"
se["B13"].font = font(12, True, TEAL_D)
for k, c in enumerate(EXP_CATS):
    style(se.cell(14 + k, 2, c), True)
CAT_RANGE = f"Settings!$B$14:$B${14 + len(EXP_CATS) + 5}"
for k in range(len(EXP_CATS), len(EXP_CATS) + 6):
    style(se.cell(14 + k, 2), True)
YEAR, RATE, DUE_DAYS = "Settings!$C$6", "Settings!$C$8/100", "Settings!$C$10"

# ---------------------------------------------------------------- Income
inc = wb.create_sheet("Income")
sheet_base(inc, "Income: one line per invoice", "Leave 'Paid on' empty until the money arrives.",
           [3, 13, 24, 12, 13, 13, 30], rows=I1 + 2)
header(inc, 5, 2, ["Invoice date", "Client", "Invoice #", "Amount", "Paid on", "Status"])
d = lambda k: today + dt.timedelta(days=k)
sample_inc = [(d(-200), "Studio North", "2026-009", 1600, d(-185)), (d(-150), "Acme Co", "2026-011", 2200, d(-128)),
              (d(-80), "Studio North", "2026-014", 1800, d(-70)), (d(-62), "Acme Co", "2026-015", 950, d(-40)),
              (d(-45), "Lumen Agency", "2026-016", 2400, d(-20)), (d(-38), "Studio North", "2026-017", 1200, None),
              (d(-12), "Pixel & Co", "2026-018", 640, None), (d(-3), "Acme Co", "2026-019", 1500, None)]
for r in range(I0, I1 + 1):
    s = sample_inc[r - I0] if r - I0 < len(sample_inc) else (None,) * 5
    style(inc.cell(r, 2, s[0]), True, DATE)
    style(inc.cell(r, 3, s[1]), True)
    style(inc.cell(r, 4, s[2]), True, align="center")
    style(inc.cell(r, 5, s[3]), True, MONEY)
    style(inc.cell(r, 6, s[4]), True, DATE)
    st = (f'=IF(OR(B{r}="",E{r}=""),"",IF(F{r}<>"","✓ Paid",IF(TODAY()-B{r}>{DUE_DAYS},'
          f'"⚠ Overdue "&(TODAY()-B{r}-{DUE_DAYS})&" day(s)","Waiting "&(TODAY()-B{r})&" day(s)")))')
    style(inc.cell(r, 7, st), False)
inc.freeze_panes = "B6"
inc.conditional_formatting.add(f"G{I0}:G{I1}", FormulaRule(formula=[f'LEFT($G{I0},1)="✓"'], fill=fill(OK)))
inc.conditional_formatting.add(f"G{I0}:G{I1}", FormulaRule(formula=[f'LEFT($G{I0},1)="⚠"'], fill=fill(BAD),
                                                            font=Font(name=F, bold=True, color="9B1C1C")))
inc.conditional_formatting.add(f"G{I0}:G{I1}", FormulaRule(formula=[f'LEFT($G{I0},7)="Waiting"'], fill=fill(WARN)))

# ---------------------------------------------------------------- Expenses
exp = wb.create_sheet("Expenses")
sheet_base(exp, "Expenses: one line per purchase", "Keep the receipt: tick 'Receipt?' when you have it.",
           [3, 13, 24, 22, 13, 11, 30], rows=E1 + 2)
header(exp, 5, 2, ["Date", "What", "Category", "Amount", "Receipt?", "Note"])
sample_exp = [(d(-190), "Laptop repair", "Equipment", 180, True, ""),
              (d(-140), "Accounting software", "Software & tools", 96, True, "yearly"),
              (d(-75), "Design app yearly plan", "Software & tools", 239, True, ""),
              (d(-60), "Train to client workshop", "Travel", 86.4, True, ""),
              (d(-41), "Phone plan", "Phone & internet", 19.99, True, "monthly"),
              (d(-33), "Monitor", "Equipment", 329, False, "find the invoice!"),
              (d(-20), "Coworking day passes", "Office & coworking", 75, True, ""),
              (d(-9), "Online course", "Training", 120, True, "")]
for r in range(E0, E1 + 1):
    s = sample_exp[r - E0] if r - E0 < len(sample_exp) else (None,) * 6
    style(exp.cell(r, 2, s[0]), True, DATE)
    style(exp.cell(r, 3, s[1]), True)
    style(exp.cell(r, 4, s[2]), True)
    style(exp.cell(r, 5, s[3]), True, MONEY)
    style(exp.cell(r, 6, s[4]), True, align="center")
    style(exp.cell(r, 7, s[5]), True)
dv_cat = DataValidation(type="list", formula1=f"={CAT_RANGE}", allow_blank=True)
dv_rec = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
exp.add_data_validation(dv_cat)
exp.add_data_validation(dv_rec)
dv_cat.add(f"D{E0}:D{E1}")
dv_rec.add(f"F{E0}:F{E1}")
exp.freeze_panes = "B6"
exp.conditional_formatting.add(f"F{E0}:F{E1}", FormulaRule(formula=[f'AND($B{E0}<>"",$F{E0}<>TRUE)'], fill=fill(WARN)))

# ---------------------------------------------------------------- Tax savings
tax = wb.create_sheet("Tax Savings")
sheet_base(tax, "Money put aside for tax", "Each time you move money to your tax savings, add a line.",
           [3, 13, 14, 40], rows=T1 + 2)
header(tax, 5, 2, ["Date", "Amount", "Note"])
for r in range(T0, T1 + 1):
    s = [(d(-180), 400, "Q1 set-aside"), (d(-120), 500, "Q2 set-aside"), (d(-35), 600, "from Acme + Lumen")][r - T0] if r - T0 < 3 else (None,) * 3
    style(tax.cell(r, 2, s[0]), True, DATE)
    style(tax.cell(r, 3, s[1]), True, MONEY)
    style(tax.cell(r, 4, s[2]), True)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Your freelance year at a glance", "Only the Settings year counts. Paid income only.",
           [3, 22, 16, 16, 16, 16, 3, 40])
IN = lambda col: f"Income!${col}${I0}:${col}${I1}"
EX = lambda col: f"Expenses!${col}${E0}:${col}${E1}"
TX = lambda col: f"'Tax Savings'!${col}${T0}:${col}${T1}"
start, end = f"DATE({YEAR},1,1)", f"DATE({YEAR},12,31)"
paid = f'SUMIFS({IN("E")},{IN("F")},">="&{start},{IN("F")},"<="&{end})'
spent = f'SUMIFS({EX("E")},{EX("B")},">="&{start},{EX("B")},"<="&{end})'
tiles = [
    ("B", "Paid income", f"={paid}", TEAL),
    ("C", "Expenses", f"={spent}", "B4541F"),
    ("D", "Profit", "=B7-C7", TEAL),
    ("E", "To put aside for tax", f"=MAX(0,D7)*{RATE}", "B4541F"),
    ("F", "Already put aside", f'=SUMIFS({TX("C")},{TX("B")},">="&{start},{TX("B")},"<="&{end})', TEAL),
]
for col, label, formula, colour in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = MONEY
    db[f"{col}7"].font = font(17, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db.row_dimensions[7].height = 30
db["H6"] = "Still to put aside"
db["H7"] = "=MAX(0,E7-F7)"
db["H7"].number_format = MONEY
for a, size in (("H6", 11), ("H7", 24)):
    db[a].font = font(size, True, "FFFFFF")
    db[a].fill = fill(TEAL_D)
    db[a].alignment = Alignment(horizontal="center", vertical="center")
db.row_dimensions[7].height = 40

header(db, 9, 2, ["Quarter", "Paid income", "Expenses", "Profit", "Tax share"])
for q in range(4):
    r = 10 + q
    qs, qe = f"DATE({YEAR},{3 * q + 1},1)", f"EOMONTH(DATE({YEAR},{3 * q + 3},1),0)"
    style(db.cell(r, 2, f"Q{q + 1}"), False, bold=True, align="center")
    style(db.cell(r, 3, f'=SUMIFS({IN("E")},{IN("F")},">="&{qs},{IN("F")},"<="&{qe})'), False, MONEY)
    style(db.cell(r, 4, f'=SUMIFS({EX("E")},{EX("B")},">="&{qs},{EX("B")},"<="&{qe})'), False, MONEY)
    style(db.cell(r, 5, f"=C{r}-D{r}"), False, MONEY, True)
    style(db.cell(r, 6, f"=MAX(0,E{r})*{RATE}"), False, MONEY)

db["H9"] = "Heads-up"
db["H9"].font = font(11, True, "FFFFFF")
db["H9"].fill = fill(TEAL)
notes = [
    (f'=COUNTIFS({IN("G")},"⚠*")&" overdue invoice(s): "&FIXED(SUMIFS({IN("E")},{IN("G")},"⚠*"),2)', BAD),
    (f'=COUNTIFS({IN("G")},"Waiting*")&" invoice(s) waiting: "&FIXED(SUMIFS({IN("E")},{IN("G")},"Waiting*"),2)', WARN),
    (f'=COUNTIFS({EX("B")},"<>",{EX("F")},"<>TRUE")&" expense(s) without a receipt"', WARN),
    ('="Tax set-aside progress: "&IF(E7=0,"—",FIXED(MIN(1,F7/E7)*100,0)&"%")', INFO),
]
for k, (formula, colour) in enumerate(notes):
    c = db.cell(10 + k, 8, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.border = box

header(db, 16, 2, ["Expense category", "This year", "Share"])
for k, cat in enumerate(EXP_CATS):
    r = 17 + k
    style(db.cell(r, 2, f"=Settings!B{14 + k}"), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({EX("E")},{EX("D")},B{r},{EX("B")},">="&{start},{EX("B")},"<="&{end})'), False, MONEY)
    share = f"IF($C$7=0,0,C{r}/$C$7)"
    c = style(db.cell(r, 4, f'=REPT("█",ROUND({share}*15,0))&"  "&TEXT({share},"0%")'), False)
    c.font = Font(name=F, size=11, color=TEAL)

chart = BarChart()
chart.type = "col"
chart.title = "Profit by quarter"
chart.height, chart.width = 7, 13
chart.add_data(Reference(db, min_col=5, min_row=9, max_row=13), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=10, max_row=13))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
db.add_chart(chart, "H16")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here: 5 minutes", "Only type in YELLOW cells.", [3, 6, 100])
steps = [
    ("1", "Settings: the year, and the share of profit you put aside for tax (ask your accountant; 25-30% is cautious)."),
    ("2", "Income: one line per invoice. Fill 'Paid on' when the money arrives. Late invoices turn red."),
    ("3", "Expenses: one line per business purchase. Tick 'Receipt?' when you have the receipt."),
    ("4", "Tax Savings: each time you move money aside for tax, add a line."),
    ("★", "Dashboard: profit, what to put aside, what is still missing, quarter by quarter."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
st["C17"] = "This is a tracking tool, not tax advice. Tax rules differ by country: check your rate with an accountant."
st["C17"].font = font(11, color=MUTED, italic=True)
st["C19"] = "Works in Google Sheets and Microsoft Excel. Example lines show how it works: delete them and add yours."
st["C19"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Freelancer-Tax-Tracker.xlsx")
wb.save(out)
print("saved", out)
