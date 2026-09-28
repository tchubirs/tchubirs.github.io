"""Etsy listing photos for the Christmas Planner (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Christmas-Planner.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, first, end: max(r for r in range(first, end + 1) if ws.cell(r, 2).value not in (None, ""))
gifts_end = last(wb["Gifts"], 6, 305)
people_end = last(wb["People"], 6, 65)
budget_end = last(wb["Budget"], 7, 20)

r = Renders(xlsx, tmp, one_page=("Dashboard", "Dinner"))
close = Renders(xlsx, tmp, one_page=("Dashboard", "Dinner"),
                hide_rows={"Gifts": range(gifts_end + 1, 306), "People": range(people_end + 1, 66),
                           "Budget": range(budget_end + 1, 22)})

img = canvas(dark=True)
title(img, "Christmas planner", "Gifts person by person, the budget, cards, the dinner and December, with the days "
      "left", dark=True, width=2300)
laptop(img, r.shot("Christmas 2026"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-gifts.jpg"), "Every gift, from idea to wrapped",
       "An x when you buy it and the day an order should arrive. A parcel that is late turns red.",
       close.shot("Gifts"), [("#E3E8FF", "On the way"), ("#F9C9C4", "Late"), ("#FFE2B8", "To wrap"), ("#D8F0DC", "Wrapped")])

# The dinner tab down to the guests: the menu and the plan for the day.
page = r.doc[r._index("Christmas dinner")]
cut = int((page.search_for("Diet or allergies")[0].y0 - 14) * 220 / 72)
full = r.page("Christmas dinner")
top = Image.new("RGB", (full.width, cut + 120), "#FFFFFF")
top.paste(full.crop((0, 0, full.width, cut)), (0, 0))
detail(os.path.join(img_dir, "03-dinner.jpg"), "Dinner, all ready at once",
       "Type how long each dish takes and when you sit down. The plan says when to start each one.", content(top))

detail(os.path.join(img_dir, "04-december.jpg"), "December on one page",
       "Parties, school events and travel go on a calendar by themselves, ready to print",
       r.shot("December 2026"))

detail(os.path.join(img_dir, "05-budget.jpg"), "The whole Christmas budget",
       "Gifts, food, decorations, travel and the rest: what is spent, what is still to pay and what is left",
       close.shot("Budget"))

detail(os.path.join(img_dir, "06-people.jpg"), "A budget for each person",
       "How many gifts, what they cost and who you still have to buy for",
       close.shot("People"), [("#FFE2B8", "To buy"), ("#E3E8FF", "All bought"), ("#D8F0DC", "All wrapped")])

img = canvas()
y = title(img, "Inside the file", "Nine tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Christmas 2026")), "Dashboard", "Days left, gifts, money"),
    (content(r.page("Gifts")), "Gifts", "Idea to wrapped"),
    (content(r.page("Budget")), "Budget", "Every category"),
    (content(r.page("Christmas dinner")), "Dinner", "Menu and cooking plan"),
    (content(r.page("December 2026")), "December", "The month on a calendar"),
    (content(r.page("Cards")), "Cards", "Who gets one"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "07-whats-inside.jpg"))
