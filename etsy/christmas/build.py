"""Christmas Planner (Excel + Google Sheets).

Gifts person by person with a budget each, orders that should have arrived and what is wrapped; the rest of the
Christmas money by category; the card list; the dinner menu with a cooking plan worked back from the time you sit
down; December on a calendar; and a Dashboard with the days left. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
YEAR = today.year if today <= dt.date(today.year, 12, 25) else today.year + 1
P0, P1 = 6, 65         # people
G0, G1 = 6, 305        # gifts
C0, C1 = 6, 305        # costs
B0, B1 = 7, 20         # budget categories (row 6 is the gifts)
K0, K1 = 6, 305        # cards
E0, E1 = 6, 155        # events
D0, D1 = 9, 28         # dishes
U0, U1 = 32, 51        # guests
L0, L1 = 9, 48         # steps of the cooking plan
PE = lambda col: f"People!${col}${P0}:${col}${P1}"
GF = lambda col: f"Gifts!${col}${G0}:${col}${G1}"
CO = lambda col: f"Costs!${col}${C0}:${col}${C1}"
CA = lambda col: f"Cards!${col}${K0}:${col}${K1}"
EV = lambda col: f"Plans!${col}${E0}:${col}${E1}"
YR, CUR = "Settings!$C$4", "Settings!$C$5"
XMAS = f"DATE({YR},12,25)"
RED, SOFT_RED = "B4541F", "F6D5D1"


def d(month, day, year=YEAR):
    return dt.date(year, month, day)


# ---------------------------------------------------------------- example data (made-up people)
people = [
    # name, group, budget
    ("Mum", "Family", 120), ("Dad", "Family", 100), ("Grandma", "Family", 60), ("Sam", "Family", 80),
    ("Lily", "Kids", 70), ("Noah", "Kids", 70), ("Aunt Rose", "Family", 40), ("Priya", "Friends", 40),
    ("Tom", "Friends", 40), ("Secret Santa", "Work", 25), ("Lily's teacher", "Kids", 15), ("Neighbours", "Neighbours", 20),
    ("Mia", "Friends", 30),
]
T = lambda n: today + dt.timedelta(days=n)
gifts = [
    # for, gift, shop, price, paid, bought, arrives by, here, wrapped
    ("Mum", "Cashmere scarf", "Department store", 65, 59, "x", None, None, "x"),
    ("Mum", "Cookbook", "Local bookshop", 28, 28, "x", None, None, None),
    ("Mum", "Bath set", "Online", 25, None, None, None, None, None),
    ("Dad", "Wireless headphones", "Online", 90, 84, "x", T(5), None, None),
    ("Dad", "Coffee sampler", "Craft market", 18, None, None, None, None, None),
    ("Grandma", "Photo calendar", "Online", 35, 32, "x", T(-2), None, None),
    ("Grandma", "Warm slippers", "Department store", 25, None, None, None, None, None),
    ("Sam", "Board game", "Toy shop", 45, 42, "x", None, "x", "x"),
    ("Sam", "Beanie hat", None, 20, None, None, None, None, None),
    ("Lily", "Art set", "Craft shop", 35, 35, "x", None, None, "x"),
    ("Lily", "Storybook", "Local bookshop", 15, 15, "x", None, None, None),
    ("Lily", "Pyjamas", None, 20, None, None, None, None, None),
    ("Noah", "Building blocks", "Toy shop", 55, 49, "x", T(12), None, None),
    ("Noah", "Football", "Sports shop", 15, None, None, None, None, None),
    ("Aunt Rose", "Scented candle", "Craft market", 22, 22, "x", None, None, None),
    ("Priya", "Plant and pot", "Garden centre", 30, None, None, None, None, None),
    ("Tom", "Hot sauce set", "Online", 28, 26, "x", T(3), None, None),
    ("Secret Santa", "Mug and cocoa", None, 15, None, None, None, None, None),
    ("Lily's teacher", "Card and chocolates", None, 12, None, None, None, None, None),
    ("Neighbours", "Mince pies", "Bakery", 10, None, None, None, None, None),
]
categories = [("Food and drink", 350), ("Decorations", 120), ("Cards and wrapping", 80), ("Travel", 150),
              ("Outings", 150), ("Clothes", 60), ("Charity", 50)]
costs = [
    # date, category, what, planned, paid
    (T(-8), "Decorations", "Advent calendars", 24, 24), (T(-3), "Cards and wrapping", "Box of 30 cards", 15, 15),
    (T(-3), "Cards and wrapping", "Wrapping paper and tags", 20, 18), (d(12, 1), "Cards and wrapping", "Stamps", 25, None),
    (d(12, 5), "Decorations", "Real tree", 60, None), (d(12, 5), "Decorations", "New lights", 30, None),
    (T(-12), "Travel", "Train tickets to Grandma's", 90, 86), (d(12, 20), "Travel", "Fuel", 40, None),
    (T(-5), "Outings", "Pantomime tickets", 96, 96), (d(12, 19), "Outings", "Christmas market", 40, None),
    (d(12, 22), "Food and drink", "Turkey", 60, None), (d(12, 23), "Food and drink", "Big food shop", 180, None),
    (T(-1), "Food and drink", "Drinks", 70, 64), (d(12, 10), "Clothes", "Christmas jumpers", 45, None),
    (d(12, 15), "Charity", "Food bank donation", 50, None),
]
cards = [
    ("Aunt Rose", "3 Ivy Lane", "Riverton", "x", None), ("Uncle Ben and Jo", "8 Holly Road", "Eastbridge", "x", "x"),
    ("The Parkers", "21 Mill Street", "Riverton", "x", None), ("Grandma", "5 Rose Court", "Westford", "x", None),
    ("Priya and Dev", None, None, "x", None), ("Tom", None, None, None, None), ("Mia", None, None, None, None),
    ("The Lees", "40 Oak Avenue", "Riverton", None, "x"), ("Mr and Mrs Grant", "2 Station Road", "Northam", None, None),
    ("Lily's teacher", None, None, None, None), ("Noah's coach", None, None, None, None),
    ("Cousin Ella", "17 Bay View", "Seaford", None, None), ("Cousin Max", None, None, None, None),
    ("Next door at 12", None, None, None, None), ("Next door at 16", None, None, None, None),
    ("Old friends from school", None, None, None, None), ("Work team", None, None, None, None),
    ("Sam and Kate", "9 Elm Close", "Eastbridge", None, None),
]
events = [
    # date, time, what, where
    (d(11, 14), None, "Order the turkey", "Butcher"), (d(11, 27), dt.time(10, 0), "Gift shopping day", "Town centre"),
    (d(12, 1), None, "Advent calendars start", "Home"), (d(12, 5), dt.time(14, 0), "Tree up and decorating", "Home"),
    (d(12, 8), dt.time(19, 0), "Office party", "The Old Mill"), (d(12, 11), dt.time(14, 30), "School concert", "School hall"),
    (d(12, 12), dt.time(10, 0), "Post the cards", "Post office"), (d(12, 13), dt.time(15, 0), "Pantomime", "Theatre"),
    (d(12, 19), dt.time(11, 0), "Christmas market", "Market square"), (d(12, 19), dt.time(18, 0), "Drinks with Priya and Tom", None),
    (d(12, 20), dt.time(17, 0), "Carol singing", "Church"), (d(12, 22), dt.time(12, 0), "Pick up the turkey", "Butcher"),
    (d(12, 23), dt.time(9, 0), "Big food shop", "Supermarket"), (d(12, 24), dt.time(18, 0), "Christmas Eve at Grandma's", "Westford"),
    (d(12, 25), dt.time(15, 0), "Christmas dinner", "Home"), (d(12, 26), dt.time(11, 0), "Boxing Day walk", "The river"),
    (d(12, 31), dt.time(20, 0), "New Year's Eve party", "Sam and Kate's"),
]
dishes = [
    # dish, course, who makes it, ready by, prep, cook, rest
    ("Prawn cocktail", "Starter", "Priya", None, 20, None, None),
    ("Roast turkey", "Main", "Me", None, 30, 210, 45), ("Roast potatoes", "Side", "Me", None, 20, 60, None),
    ("Stuffing", "Side", "Me", None, 15, 45, None), ("Glazed carrots", "Side", "Sam", None, 10, 30, None),
    ("Green beans", "Side", "Me", None, 10, 10, None), ("Gravy", "Side", "Me", None, 5, 15, None),
    ("Cranberry sauce", "Side", "Grandma", None, None, None, None),
    ("Christmas pudding", "Dessert", "Grandma", dt.time(16, 0), None, 120, None),
    ("Trifle", "Dessert", "Me", dt.datetime(YEAR, 12, 24, 20, 0), 40, None, None),
    ("Mulled wine", "Drinks", "Tom", dt.time(14, 30), 5, 20, None),
]
guests = [("Me", "x", None), ("Sam", "x", "Vegetarian"), ("Kate", "x", None), ("Grandma", "x", "Soft food"),
          ("Mum", "x", None), ("Dad", "x", "No nuts"), ("Priya", "x", "Gluten free"), ("Tom", "x", None),
          ("Aunt Rose", None, "Not sure yet")]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def status_fills(ws, rng, first, pairs):
    for text, colour in pairs:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{text}"'], fill=fill(colour)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The year and the currency. Everything else is on its own tab.", [3, 22, 16, 4, 60], rows=12)
for r, label, value, fmt in ((4, "Christmas of", YEAR, "0"), (5, "Currency", "USD", None)):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
for k, text in enumerate(["Next year, change the year: the countdown, the calendar and the dinner day follow it.",
                          "Amounts have no currency sign, so any currency works."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- People
pe = wb.create_sheet("People")
sheet_base(pe, "People", "Everyone you buy for, with a budget each. The rest comes from the Gifts tab.",
           [3, 20, 14, 11, 9, 9, 11, 11, 11, 16], rows=P1 + 2)
header(pe, 5, 2, ["Name", "Group", "Budget", "Gifts", "Bought", "Spent", "To pay", "Left", "Status"])
for r in range(P0, P1 + 1):
    x = people[r - P0] if r - P0 < len(people) else (None,) * 3
    style(pe.cell(r, 2, x[0]), True, bold=True)
    style(pe.cell(r, 3, x[1]), True)
    style(pe.cell(r, 4, x[2]), True, MONEY)
    style(pe.cell(r, 5, f'=IF(B{r}="","",COUNTIFS({GF("B")},B{r}))'), False, "0", align="center")
    style(pe.cell(r, 6, f'=IF(B{r}="","",SUMIFS({GF("M")},{GF("B")},B{r}))'), False, "0", align="center")
    style(pe.cell(r, 7, f'=IF(B{r}="","",SUMIFS({GF("N")},{GF("B")},B{r}))'), False, MONEY)
    style(pe.cell(r, 8, f'=IF(B{r}="","",SUMIFS({GF("O")},{GF("B")},B{r}))'), False, MONEY)
    style(pe.cell(r, 9, f'=IF(OR(B{r}="",D{r}=""),"",D{r}-G{r}-H{r})'), False, MONEY)
    style(pe.cell(r, 10, f'=IF(B{r}="","",IF(E{r}=0,"No gift yet",IF(F{r}<E{r},(E{r}-F{r})&" to buy",'
                         f'IF(COUNTIFS({GF("B")},B{r},{GF("K")},"Wrapped")=E{r},"All wrapped","All bought"))))'),
          False, align="center")
    pe.cell(r, 12, f'=IF(B{r}="","",IF(OR(J{r}="All bought",J{r}="All wrapped"),"",ROW()))')   # L: on the Dashboard
pe.column_dimensions["L"].hidden = True
dropdown(pe, '"Family,Friends,Kids,Work,Neighbours,Other"', f"C{P0}:C{P1}", strict=False)
pe.conditional_formatting.add(f"I{P0}:I{P1}", FormulaRule(formula=[f'AND(ISNUMBER(I{P0}),I{P0}<0)'], fill=fill(BAD)))
status_fills(pe, f"J{P0}:J{P1}", f"J{P0}", (("No gift yet", WARN), ("All bought", INFO), ("All wrapped", OK)))
pe.conditional_formatting.add(f"J{P0}:J{P1}", FormulaRule(formula=[f'RIGHT(J{P0},6)="to buy"'], fill=fill(WARN)))
pe.freeze_panes = "C6"

# ---------------------------------------------------------------- Gifts
gi = wb.create_sheet("Gifts")
sheet_base(gi, "Gifts", "One line per gift. An x when you buy it, the day an order should arrive, an x when it is here "
           "and when it is wrapped.", [3, 16, 24, 20, 10, 10, 8, 12, 7, 9, 12], rows=G1 + 2)
header(gi, 5, 2, ["For", "Gift", "Shop or link", "Price", "Paid", "Bought", "Arrives by", "Here", "Wrapped", "Status"])
for r in range(G0, G1 + 1):
    x = gifts[r - G0] if r - G0 < len(gifts) else (None,) * 9
    for c, (v, fmt, align) in enumerate(zip(x, (None, None, None, MONEY, MONEY, None, DATE, None, None),
                                            (None, None, None, None, None, "center", "center", "center", "center"))):
        style(gi.cell(r, 2 + c, v), True, fmt, c == 0, align)
    empty = f'AND(B{r}="",C{r}="")'
    style(gi.cell(r, 11, f'=IF({empty},"",IF(J{r}<>"","Wrapped",IF(I{r}<>"","To wrap",IF(AND(G{r}="",F{r}=""),"Idea",'
                         f'IF(H{r}="","To wrap",IF(H{r}<TODAY(),"Late","On the way"))))))'), False, align="center")
    # M: bought or not, N: spent, O: still to pay, P: orders on the way (for the Dashboard).
    gi.cell(r, 13, f'=IF({empty},"",IF(OR(F{r}<>"",G{r}<>"",I{r}<>"",J{r}<>""),1,0))')
    gi.cell(r, 14, f'=IF(M{r}="","",IF(M{r}=1,IF(F{r}<>"",N(F{r}),N(E{r})),0))')
    gi.cell(r, 15, f'=IF(M{r}="","",IF(M{r}=1,0,N(E{r})))')
    gi.cell(r, 16, f'=IF(OR(K{r}="On the way",K{r}="Late"),H{r}+ROW()/10000000,"")')
for c in "MNOP":
    gi.column_dimensions[c].hidden = True
dropdown(gi, f"={PE('B')}", f"B{G0}:B{G1}")
status_fills(gi, f"K{G0}:K{G1}", f"K{G0}", (("On the way", INFO), ("Late", BAD), ("To wrap", WARN), ("Wrapped", OK)))
gi.freeze_panes = "C6"

# ---------------------------------------------------------------- Costs
co = wb.create_sheet("Costs")
sheet_base(co, "Costs", "Everything else Christmas costs: what you plan to spend, and what you paid.",
           [3, 12, 20, 30, 11, 11], rows=C1 + 2)
header(co, 5, 2, ["Date", "Category", "What", "Planned", "Paid"])
for r in range(C0, C1 + 1):
    x = costs[r - C0] if r - C0 < len(costs) else (None,) * 5
    for c, (v, fmt) in enumerate(zip(x, (DATE, None, None, MONEY, MONEY))):
        style(co.cell(r, 2 + c, v), True, fmt)
    co.cell(r, 8, f"=N(F{r})")                                   # H: spent
    co.cell(r, 9, f'=IF(F{r}="",N(E{r}),0)')                     # I: still to pay
for c in "HI":
    co.column_dimensions[c].hidden = True
dropdown(co, f"=Budget!$B${B0}:$B${B1}", f"C{C0}:C{C1}", strict=False)
co.freeze_panes = "C6"

# ---------------------------------------------------------------- Budget
bu = wb.create_sheet("Budget")
sheet_base(bu, "Budget", "The gifts add up the people's budgets. Type a budget for everything else. Left takes off "
           "what is still to pay.", [3, 22, 12, 12, 12, 12, 24], rows=26)
header(bu, 5, 2, ["Category", "Budget", "Spent", "To pay", "Left", "Used"])
gift_costs = lambda col: f'SUMIFS({CO(col)},{CO("C")},"Gifts")'
style(bu.cell(6, 2, "Gifts"), False, bold=True)
style(bu.cell(6, 3, f"=SUM({PE('D')})"), False, MONEY)
style(bu.cell(6, 4, f"=SUM({GF('N')})+{gift_costs('H')}"), False, MONEY)
style(bu.cell(6, 5, f"=SUM({GF('O')})+{gift_costs('I')}"), False, MONEY)
for r in range(B0, B1 + 1):
    x = categories[r - B0] if r - B0 < len(categories) else (None, None)
    style(bu.cell(r, 2, x[0]), True, bold=True)
    style(bu.cell(r, 3, x[1]), True, MONEY)
    style(bu.cell(r, 4, f'=IF(B{r}="","",SUMIFS({CO("H")},{CO("C")},B{r}))'), False, MONEY)
    style(bu.cell(r, 5, f'=IF(B{r}="","",SUMIFS({CO("I")},{CO("C")},B{r}))'), False, MONEY)
other = B1 + 1
style(bu.cell(other, 4, f"=SUM({CO('H')})-{gift_costs('H')}-SUM(D{B0}:D{B1})"), False, "#,##0.00;-#,##0.00;;")
style(bu.cell(other, 5, f"=SUM({CO('I')})-{gift_costs('I')}-SUM(E{B0}:E{B1})"), False, "#,##0.00;-#,##0.00;;")
style(bu.cell(other, 2, f'=IF(AND(D{other}=0,E{other}=0),"","Not on this list")'), False, bold=True)
style(bu.cell(other, 3), False)
for r in range(6, other + 1):
    style(bu.cell(r, 6, f'=IF(OR(B{r}="",C{r}=""),"",N(C{r})-D{r}-E{r})'), False, MONEY)
    style(bu.cell(r, 7, f'=IF(OR(B{r}="",N(C{r})=0),"",{bar(f"(D{r}+E{r})/C{r}", 16)})'), False)
    bu.cell(r, 7).font = Font(name=F, size=9, color=TEAL)
    bu.cell(r, 9, f'=IF(AND(B{r}<>"",OR(N(C{r})>0,N(D{r})+N(E{r})>0)),ROW(),"")')       # I: on the Dashboard
tot = other + 1
style(bu.cell(tot, 2, "Total"), False, bold=True)
for c, formula in (("C", f"=SUM(C6:C{B1})"), ("D", f"=SUM(D6:D{other})"), ("E", f"=SUM(E6:E{other})"),
                   ("F", f"=C{tot}-D{tot}-E{tot}")):
    style(bu[f"{c}{tot}"], False, MONEY, True)
    bu[f"{c}{tot}"] = formula
style(bu.cell(tot, 7, f'=IF(N(C{tot})=0,"",{bar(f"(D{tot}+E{tot})/C{tot}", 16)})'), False)
bu.cell(tot, 7).font = Font(name=F, size=9, color=TEAL)
for c in "BCDEFG":
    bu[f"{c}{tot}"].fill = fill("E8F0EF")
bu.column_dimensions["I"].hidden = True
bu.conditional_formatting.add(f"F6:F{tot}", FormulaRule(formula=['AND(ISNUMBER(F6),F6<0)'], fill=fill(BAD)))

# ---------------------------------------------------------------- Cards
ca = wb.create_sheet("Cards")
sheet_base(ca, "Cards", "Who gets a card, the address, an x when it is posted and when one comes from them.",
           [3, 24, 26, 18, 8, 11, 26, 11], rows=K1 + 2)
header(ca, 5, 2, ["Name", "Address", "Town and postcode", "Sent", "From them", "Notes", "Status"])
for r in range(K0, K1 + 1):
    x = cards[r - K0] if r - K0 < len(cards) else (None,) * 5
    for c, v in enumerate(x):
        style(ca.cell(r, 2 + c, v), True, bold=c == 0, align="center" if c >= 3 else None)
    style(ca.cell(r, 7), True)
    style(ca.cell(r, 8, f'=IF(B{r}="","",IF(E{r}<>"","Sent","To send"))'), False, align="center")
status_fills(ca, f"H{K0}:H{K1}", f"H{K0}", (("Sent", OK), ("To send", WARN)))
ca.freeze_panes = "C6"

# ---------------------------------------------------------------- Plans (the dates) and December (the calendar)
pl = wb.create_sheet("Plans")
sheet_base(pl, "Plans", "Parties, school events, travel and the days to get things done. The December tab puts them "
           "on a calendar.", [3, 12, 8, 28, 20, 24], rows=E1 + 2)
header(pl, 5, 2, ["Date", "Time", "What", "Where", "Notes"])
for r in range(E0, E1 + 1):
    x = events[r - E0] if r - E0 < len(events) else (None,) * 4
    style(pl.cell(r, 2, x[0]), True, "ddd d mmm")
    style(pl.cell(r, 3, x[1]), True, "hh:mm", align="center")
    style(pl.cell(r, 4, x[2]), True, bold=True)
    style(pl.cell(r, 5, x[3]), True)
    style(pl.cell(r, 6), True)
    pl.cell(r, 8, f'=IF(B{r}="","",B{r}+MOD(N(C{r}),1)+ROW()/10000000)')                 # H: in date order
pl.column_dimensions["H"].hidden = True
pl.conditional_formatting.add(f"B{E0}:F{E1}", FormulaRule(formula=[f'AND($B{E0}<>"",$B{E0}<TODAY())'],
                                                          font=Font(name=F, italic=True, color=MUTED)))
pl.freeze_panes = "C6"

de = wb.create_sheet("December")
sheet_base(de, "", "Everything on the Plans tab, day by day. Up to two plans a day are shown.",
           [3] + [17] * 7 + [3], rows=20)
de["B2"] = f'="December "&{YR}'
first_monday = f"(DATE({YR},12,1)-WEEKDAY(DATE({YR},12,1),3))"
header(de, 5, 2, ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"])
keys = EV("H")
for w in range(6):
    top, low = 6 + 2 * w, 7 + 2 * w
    for k in range(7):
        col = 2 + k
        day = f"({first_monday}+{7 * w + k})"
        c = de.cell(top, col, f'=IF(MONTH({day})=12,DAY({day}),"")')
        c.font = font(11, True, TEAL)
        c.alignment = Alignment(horizontal="left", vertical="top")
        c.border = box
        before = f'COUNTIF({keys},"<"&{day})'
        nth = lambda n: f'SMALL({keys},{before}+{n})'
        what = lambda n: f'INDEX({EV("D")},MATCH({nth(n)},{keys},0))'
        count = f'COUNTIFS({keys},">="&{day},{keys},"<"&({day}+1))'
        c = de.cell(low, col, f'=IF(OR(MONTH({day})<>12,{count}=0),"",{what(1)}&IF({count}>1,CHAR(10)&{what(2)},"")'
                              f'&IF({count}>2,CHAR(10)&"+"&({count}-2)&" more",""))')
        c.font = font(9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = box
    de.row_dimensions[top].height = 18
    de.row_dimensions[low].height = 50
days = " ".join(f"B{6 + 2 * w}:H{6 + 2 * w}" for w in range(6))
texts = " ".join(f"B{7 + 2 * w}:H{7 + 2 * w}" for w in range(6))
de.conditional_formatting.add(days, FormulaRule(formula=["B6=25"], fill=fill(SOFT_RED)))
de.conditional_formatting.add(texts, FormulaRule(formula=["B6=25"], fill=fill(SOFT_RED)))
de.conditional_formatting.add(days, FormulaRule(formula=[f"AND(ISNUMBER(B6),DATE({YR},12,B6)=TODAY())"], fill=fill(INFO)))
de.conditional_formatting.add(texts, FormulaRule(formula=[f"AND(ISNUMBER(B6),DATE({YR},12,B6)=TODAY())"], fill=fill(INFO)))
de.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dinner (menu, guests and when to start each dish)
di = wb.create_sheet("Dinner")
sheet_base(di, "Christmas dinner", "The menu, who makes what, and when to start each dish so it is all ready together.",
           [3, 20, 13, 14, 14, 8, 8, 8, 13, 7, 3, 13, 36], rows=U1 + 2)
for r, label, value, fmt in ((4, "Dinner on", f"=DATE({YR},12,25)", "ddd d mmm"), (5, "Sit down at", dt.time(15, 0), "hh:mm")):
    di[f"B{r}"] = label
    di[f"B{r}"].font = font(11, True)
    style(di[f"C{r}"], True, fmt, True, "center")
    di[f"C{r}"] = value
di["B6"] = "Guests coming"
di["B6"].font = font(11, True)
style(di["C6"], False, "0", True, "center")
di["C6"] = f'=SUMPRODUCT(($B${U0}:$B${U1}<>"")*($C${U0}:$C${U1}<>""))'
di["E4"] = "Ready by: leave it empty for the time you sit down, or type a time, or a date and time."
di["E5"] = "Minutes: prep before it cooks, cooking, and resting after."
for c in ("E4", "E5"):
    di[c].font = font(9, color=MUTED, italic=True)
header(di, 8, 2, ["Dish", "Course", "Who makes it", "Ready by", "Prep (min)", "Cook (min)", "Rest (min)", "Start",
                  "Done"])
sit = "$C$4+$C$5"
for r in range(D0, D1 + 1):
    x = dishes[r - D0] if r - D0 < len(dishes) else (None,) * 7
    name, course, who, ready, prep, cook, rest = x
    style(di.cell(r, 2, name), True, bold=True)
    style(di.cell(r, 3, course), True)
    style(di.cell(r, 4, who), True)
    style(di.cell(r, 5, ready), True, "ddd hh:mm" if isinstance(ready, dt.datetime) else "hh:mm", align="center")
    for c, v in ((6, prep), (7, cook), (8, rest)):
        style(di.cell(r, c, v), True, "0", align="center")
    style(di.cell(r, 10), True, align="center")
    ready_at = f'IF(E{r}="",{sit},IF(E{r}>=1,E{r},$C$4+E{r}))'
    total = f"(N(F{r})+N(G{r})+N(H{r}))"
    style(di.cell(r, 9, f'=IF(OR(B{r}="",{total}=0),"",ROUND(({ready_at}-{total}/1440)*1440,0)/1440)'), False,
          "ddd hh:mm", align="center")
    # R, S, T: when to prepare it, start cooking it and take it out to rest.
    di.cell(r, 18, f'=IF(OR(B{r}="",N(F{r})=0),"",{ready_at}-{total}/1440+(100+ROW())/10000000)')
    di.cell(r, 19, f'=IF(OR(B{r}="",N(G{r})=0),"",{ready_at}-(N(G{r})+N(H{r}))/1440+(200+ROW())/10000000)')
    di.cell(r, 20, f'=IF(OR(B{r}="",N(H{r})=0),"",{ready_at}-N(H{r})/1440+(300+ROW())/10000000)')
dropdown(di, '"Starter,Main,Side,Dessert,Drinks,Other"', f"C{D0}:C{D1}", strict=False)
di.conditional_formatting.add(f"B{D0}:I{D1}", FormulaRule(formula=[f'$J{D0}<>""'], font=Font(name=F, italic=True,
                                                                                             color=MUTED)))
header(di, U0 - 1, 2, ["Guest", "Coming", "Diet or allergies"])
for r in range(U0, U1 + 1):
    x = guests[r - U0] if r - U0 < len(guests) else (None,) * 3
    style(di.cell(r, 2, x[0]), True, bold=True)
    style(di.cell(r, 3, x[1]), True, align="center")
    style(di.cell(r, 4, x[2]), True)
header(di, 8, 12, ["When", "What to do"])
# The time you sit down is one more step (R under the last dish), so everything sorts together.
di.cell(D1 + 1, 18, f"={sit}+999/10000000")
steps = f"$R${D0}:$T${D1 + 1}"
n_steps = f"COUNT({steps})"
for i in range(L1 - L0 + 1):
    r = L0 + i
    key = f"O{r}"
    di[key] = f'=IF({i + 1}<={n_steps},SMALL({steps},{i + 1}),"")'
    which = lambda col: f'MATCH({key},${col}${D0}:${col}${D1},0)'
    dish = lambda col: f'INDEX($B${D0}:$B${D1},{which(col)})'
    style(di.cell(r, 12, f'=IF({key}="","",ROUND({key}*1440,0)/1440)'), False, "ddd hh:mm", align="center")
    style(di.cell(r, 13, f'=IF({key}="","",IF({key}=$R${D1 + 1},"Sit down to eat",IF(ISNUMBER({which("R")}),'
                         f'"Prepare: "&{dish("R")},IF(ISNUMBER({which("S")}),"Cook: "&{dish("S")},"Rest: "&{dish("T")}))))'),
          False)
di[f"L{L1 + 1}"] = f'=IF({n_steps}>{L1 - L0 + 1},"More steps than fit here.","")'
di[f"L{L1 + 1}"].font = font(9, color=RED, italic=True)
di.conditional_formatting.add(f"L{L0}:M{L1}", FormulaRule(formula=[f'$M{L0}="Sit down to eat"'], fill=fill(OK)))
for c in "ORST":
    di.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "", "Gifts and plans on the left, the money on the right.",
           [3, 18, 14, 14, 14, 3, 18, 14, 14, 14, 3], rows=48)
db["B2"] = f'="Christmas "&{YR}'
bought, total_gifts = f"SUM({GF('M')})", f"COUNT({GF('M')})"
tile(db, "B", 4, "Days to Christmas", f'=IF(TODAY()={XMAS},"Today",MAX(0,{XMAS}-TODAY()))', "0")
tile(db, "C", 4, "Gifts bought", f'={bought}&" of "&{total_gifts}', "General")
tile(db, "D", 4, "Wrapped", f'=COUNTIF({GF("K")},"Wrapped")&" of "&{total_gifts}', "General")
tile(db, "E", 4, "Cards sent", f'=COUNTIF({CA("H")},"Sent")&" of "&(COUNTIF({CA("H")},"Sent")+COUNTIF({CA("H")},"To send"))',
     "General")
tile(db, "G", 4, "Budget", f"=Budget!C{tot}", MONEY)
tile(db, "H", 4, "Spent", f"=Budget!D{tot}", MONEY, "B07A00")
tile(db, "I", 4, "Still to pay", f"=Budget!E{tot}", MONEY, "B07A00")
tile(db, "J", 4, "Left", f"=Budget!F{tot}", MONEY)
db.conditional_formatting.add("J5", FormulaRule(formula=["J5<0"], font=Font(name=F, size=16, bold=True, color=RED)))

header(db, 8, 2, ["Still to buy for", "Status", "Budget", "Left"])
for k in range(8):
    r = 9 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL(People!$L${P0}:$L${P1},{k + 1}),"")'
    at = lambda col: f"INDEX({PE(col)},{key}-{P0 - 1})"
    none = '"Every gift is bought"' if k == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},{at("B")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",{at("J")})'), False, align="center")
    style(db.cell(r, 4, f'=IF({key}="","",IF({at("D")}="","",{at("D")}))'), False, MONEY)
    style(db.cell(r, 5, f'=IF({key}="","",{at("I")})'), False, MONEY)
db.conditional_formatting.add("C9:C16", FormulaRule(formula=['C9="No gift yet"'], fill=fill(WARN)))
db.conditional_formatting.add("E9:E16", FormulaRule(formula=['AND(ISNUMBER(E9),E9<0)'], fill=fill(BAD)))

header(db, 8, 7, ["Category", "Budget", "Spent", "To pay"])
for k in range(8):
    r = 9 + k
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL(Budget!$I$6:$I${other},{k + 1}),"")'
    at = lambda col: f"INDEX(Budget!${col}$6:${col}${other},{key}-5)"
    style(db.cell(r, 7, f'=IF({key}="","",{at("B")})'), False, bold=True)
    for c, col in ((8, "C"), (9, "D"), (10, "E")):
        style(db.cell(r, c, f'=IF({key}="","",N({at(col)}))'), False, "#,##0.00;-#,##0.00;;")

header(db, 19, 2, ["Orders on the way", "For", "Arrives by", "Status"])
for i in range(8):
    r = 20 + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({GF("P")},{i + 1}),"")'
    row = f'MATCH({key},{GF("P")},0)'
    none = '"No orders on the way"' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({GF("C")},{row}))'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",INDEX({GF("B")},{row}))'), False)
    style(db.cell(r, 4, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 5, f'=IF({key}="","",INDEX({GF("K")},{row}))'), False, align="center")
status_fills(db, "E20:E27", "E20", (("Late", BAD), ("On the way", INFO)))

header(db, 19, 7, ["Coming up", "Time", "What", ""])
db.merge_cells("I19:J19")
for i in range(8):
    r = 20 + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({EV("H")},COUNTIF({EV("H")},"<"&TODAY())+{i + 1}),"")'
    row = f'MATCH({key},{EV("H")},0)'
    none = '"Nothing planned yet"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},INT({key}))'), False, "ddd d mmm", bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",IF(INDEX({EV("C")},{row})="","",INDEX({EV("C")},{row})))'), False, "hh:mm",
          align="center")
    style(db.cell(r, 9, f'=IF({key}="","",INDEX({EV("D")},{row})&IF(INDEX({EV("E")},{row})="","",", "&'
                        f'INDEX({EV("E")},{row})))'), False)
    style(db.cell(r, 10), False)
    db.merge_cells(f"I{r}:J{r}")
db.conditional_formatting.add("G20:J27", FormulaRule(formula=["AND(ISNUMBER($P20),INT($P20)=TODAY())"], fill=fill(INFO)))

chart = BarChart()
chart.type = "col"
chart.title = "Budget by category"
chart.height, chart.width = 6.5, 24
chart.add_data(Reference(db, min_col=8, max_col=10, min_row=8, max_row=16), titles_from_data=True)
chart.set_categories(Reference(db, min_col=7, min_row=9, max_row=16))
for s, colour in zip(chart.series, (TEAL, "F2A65A", "B8D3CF")):
    s.graphicalProperties.solidFill = colour
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "B30")
for c in "OP":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps_text = [
    ("1", "People: everyone you buy for, with a budget each. Settings has the year and the currency."),
    ("2", "Gifts: one line per gift with its price. An x when you buy it, the day an order should arrive, then Here and Wrapped."),
    ("3", "Budget and Costs: a budget for food, decorations, travel and the rest, and one line per cost."),
    ("4", "Cards: who gets one, and an x when it is posted."),
    ("5", "Dinner, Plans and December: the menu with a cooking plan, the guests, and your dates on a calendar."),
]
for k, (n_, text) in enumerate(steps_text):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example family, their gifts and plans are made up. Delete them and add your own.",
                          "An order turns Late when its day has passed and it is not marked Here.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Gifts", "People", "Budget", "Costs", "Cards", "Dinner", "Plans", "December", "Settings",
         "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Christmas-Planner.xlsx")
wb.save(out)
print("saved", out, len(gifts), "gifts", len(people), "people", len(events), "events", len(dishes), "dishes")
