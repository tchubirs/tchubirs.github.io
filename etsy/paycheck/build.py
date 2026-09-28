"""Paycheck Budget Planner (Excel + Google Sheets).

The paydays for twelve months are listed from the pay frequency and the first
payday. Each bill falls on the paycheck that comes before its due date, so
every paycheck shows the bills it has to cover, the savings, what is left to
spend and how much per day until the next payday. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
NET = '#,##0.00;[Red]-#,##0.00'
K0, K1 = 6, 58       # paycheck rows (up to 53 weekly paydays)
B0, B1 = 6, 55       # bill rows
FREQS = ["Weekly", "Every 2 weeks", "Twice a month", "Monthly"]
FREQ, FIRST, DAY2, USUAL, SAVE, TIGHT = (f"Settings!$C${5 + k}" for k in range(6))
PK = lambda col: f"Paychecks!${col}${K0}:${col}${K1}"


def payday(k):
    """The k-th payday (k is a formula) for every pay frequency."""
    ym = lambda months: f"DATE(YEAR({FIRST}),MONTH({FIRST})+{months},1)"
    clamp = lambda months, day: f"DATE(YEAR({FIRST}),MONTH({FIRST})+{months},MIN({day},DAY(EOMONTH({ym(months)},0))))"
    half = f"INT(({k})/2)"
    return (f'IF({FREQ}="Weekly",{FIRST}+7*({k}),IF({FREQ}="Every 2 weeks",{FIRST}+14*({k}),'
            f'IF({FREQ}="Twice a month",{clamp(half, f"IF(MOD({k},2)=0,DAY({FIRST}),{DAY2})")},'
            f'{clamp(f"({k})", f"DAY({FIRST})")})))')


wb = Workbook()

# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "How you are paid. The Paychecks tab lists your paydays from these.", [3, 34, 16, 4, 60], rows=30)
jan_friday = dt.date(today.year, 1, 1) + dt.timedelta(days=(4 - dt.date(today.year, 1, 1).weekday()) % 7)
settings = [("How often you are paid", "Every 2 weeks", None, "Weekly, every 2 weeks, twice a month or monthly"),
            ("First payday", jan_friday, DATE, "The first payday of the 12 months you want to plan"),
            ("Second payday of the month", 15, "0", "Only for twice a month: the day of the second payday"),
            ("Usual paycheck", 1850, MONEY, "What you usually receive, after tax"),
            ("Savings from each paycheck", 150, MONEY, "Moved to savings before you spend"),
            ("Tight when less than this per day", 20, MONEY, "A paycheck turns orange below this amount per day")]
for k, (label, value, fmt, note) in enumerate(settings):
    r = 5 + k
    style(se.cell(r, 2, label), False, bold=True)
    style(se.cell(r, 3, value), True, fmt, True, "center")
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
dv = DataValidation(type="list", formula1='"' + ",".join(FREQS) + '"', allow_blank=False)
se.add_data_validation(dv)
dv.add("C5")

# ---------------------------------------------------------------- Bills
bi = wb.create_sheet("Bills")
sheet_base(bi, "Bills", "Each bill once, with the day of the month it is due. Each one goes on the paycheck before it.",
           [3, 24, 12, 9, 10, 18, 10, 26], rows=B1 + 2)
header(bi, 5, 2, ["Bill", "Amount", "Due day", "Only in month", "Category", "Autopay?", "Note"])
bills = [("Rent", 1200, 1, None, "Housing", "No", ""), ("Car payment", 285, 5, None, "Transport", "Yes", ""),
         ("Phone", 45, 12, None, "Utilities", "Yes", ""), ("Water", 40, 15, None, "Utilities", "No", ""),
         ("Internet", 60, 18, None, "Utilities", "Yes", ""), ("Electricity", 95, 20, None, "Utilities", "No", ""),
         ("Car insurance", 110, 22, None, "Insurance", "Yes", ""), ("Streaming", 16, 25, None, "Fun", "Yes", ""),
         ("Gym", 35, 27, None, "Health", "Yes", ""), ("Credit card minimum", 75, 28, None, "Debt", "No", ""),
         ("Renters insurance", 180, 10, 3, "Insurance", "No", "Once a year, in March")]
DUE0, AMT0 = 10, 23   # hidden helper columns: 13 due dates (J..V), then 13 amounts (W..AI)
for r in range(B0, B1 + 1):
    b = bills[r - B0] if r - B0 < len(bills) else (None,) * 7
    for c, (v, fmt) in enumerate(zip(b, (None, MONEY, "0", "0", None, None, None)), start=2):
        style(bi.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (4, 5, 7) else None)
    for j in range(13):
        month1 = f"DATE(YEAR({FIRST}),MONTH({FIRST})+{j},1)"
        bi.cell(r, DUE0 + j, f'=IF(OR(B{r}="",D{r}=""),"",DATE(YEAR({FIRST}),MONTH({FIRST})+{j},'
                             f'MIN(D{r},DAY(EOMONTH({month1},0)))))')
        bi.cell(r, AMT0 + j, f'=IF(OR(B{r}="",D{r}=""),0,IF(OR(E{r}="",E{r}=MONTH({month1})),N(C{r}),0))')
    dues, amts = f"{L(DUE0)}{r}:{L(DUE0 + 12)}{r}", f"{L(AMT0)}{r}:{L(AMT0 + 12)}{r}"
    bi.cell(r, 36, f'=IF(B{r}="","",SUMIFS({dues},{dues},">="&Dashboard!$C$4,{dues},"<"&Dashboard!$N$1,{amts},">0"))')
    bi.cell(r, 37, f'=IF(N(AJ{r})=0,"",AJ{r}+ROW()/100000)')
for c in range(DUE0, 38):
    bi.column_dimensions[L(c)].hidden = True
DUE2D = f"Bills!${L(DUE0)}${B0}:${L(DUE0 + 12)}${B1}"
AMT2D = f"Bills!${L(AMT0)}${B0}:${L(AMT0 + 12)}${B1}"
BL = lambda col: f"Bills!${col}${B0}:${col}${B1}"
for col, formula in (("G", '"Yes,No"'),):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    bi.add_data_validation(dv)
    dv.add(f"{col}{B0}:{col}{B1}")
bi.freeze_panes = "C6"

# ---------------------------------------------------------------- Paychecks
pc = wb.create_sheet("Paychecks")
sheet_base(pc, "Paychecks", "Your paydays for 12 months. Type an amount only when a paycheck is different.",
           [3, 14, 13, 13, 14, 13, 12, 13, 8, 11, 20], rows=K1 + 2)
header(pc, 5, 2, ["Payday", "Amount", "Different amount", "Next payday", "Bills due", "Savings", "Left to spend",
                  "Days", "Per day", "Status"])
end = f"EDATE({FIRST},12)"
for r in range(K0, K1 + 1):
    this, nxt = payday(f"ROW()-{K0}"), payday(f"ROW()-{K0 - 1}")
    style(pc.cell(r, 2, f'=IF({this}<{end},{this},"")'), False, DATE, True, "center")
    style(pc.cell(r, 3, f'=IF(B{r}="","",IF(ISNUMBER(D{r}),D{r},{USUAL}))'), False, MONEY)
    style(pc.cell(r, 4), True, MONEY)
    style(pc.cell(r, 5, f'=IF(B{r}="","",{nxt})'), False, DATE, align="center")
    style(pc.cell(r, 6, f'=IF(B{r}="","",SUMIFS({AMT2D},{DUE2D},">="&B{r},{DUE2D},"<"&E{r}))'), False, MONEY)
    style(pc.cell(r, 7, f'=IF(B{r}="","",{SAVE})'), False, MONEY)
    style(pc.cell(r, 8, f'=IF(B{r}="","",C{r}-F{r}-G{r})'), False, NET, True)
    style(pc.cell(r, 9, f'=IF(B{r}="","",E{r}-B{r})'), False, "0", align="center")
    style(pc.cell(r, 10, f'=IF(B{r}="","",H{r}/I{r})'), False, NET)
    style(pc.cell(r, 11, f'=IF(B{r}="","",IF(H{r}<0,"Short by "&FIXED(-H{r},2),IF(J{r}<{TIGHT},"Tight","OK")))'), False)
st_rng = f"K{K0}:K{K1}"
pc.conditional_formatting.add(st_rng, FormulaRule(formula=[f'LEFT($K{K0},5)="Short"'], fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
pc.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$K{K0}="Tight"'], fill=fill(WARN)))
pc.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$K{K0}="OK"'], fill=fill(OK)))
pc.conditional_formatting.add(f"B{K0}:J{K1}", FormulaRule(formula=[f'AND($B{K0}<>"",$B{K0}<=TODAY(),$E{K0}>TODAY())'],
                                                         fill=fill("E6F2F2"), font=Font(name=F, bold=True)))
pc.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Paycheck budget", "Pick a paycheck to see the bills it covers and what is left to spend.",
           [3, 14, 24, 13, 11, 14, 14, 3, 30, 15, 3], rows=60)
db["B4"] = "Paycheck of"
db["B4"].font = font(11, True)
count_past = f'COUNTIF({PK("B")},"<="&TODAY())'
db["C4"] = f'=IF({count_past}=0,INDEX({PK("B")},1),INDEX({PK("B")},{count_past}))'
style(db["C4"], True, DATE, True, "center")
dv = DataValidation(type="list", formula1=f"={PK('B')}", allow_blank=False)
db.add_data_validation(dv)
dv.add("C4")
db["D4"] = "The paycheck you are on today is picked for you. Choose another date to plan ahead."
db["D4"].font = font(9, color=MUTED, italic=True)
db["N1"] = f'=IFERROR(INDEX({PK("E")},MATCH($C$4,{PK("B")},0)),"")'
db["N2"] = f'=MATCH($C$4,{PK("B")},0)'
db.column_dimensions["M"].hidden = db.column_dimensions["N"].hidden = True
pick = lambda col: f'=IFERROR(INDEX({PK(col)},$N$2),"")'
tiles = [("B", "Paycheck", pick("C"), TEAL, MONEY), ("C", "Bills to pay from it", pick("F"), "B4541F", MONEY),
         ("D", "Savings", pick("G"), TEAL, MONEY), ("E", "Days", pick("I"), TEAL, "0"),
         ("F", "Left to spend", pick("H"), TEAL, NET), ("G", "Per day", pick("J"), TEAL, NET)]
for col, label, formula, colour, fmt in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = fmt
    db[f"{col}7"].font = font(16, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db.row_dimensions[7].height = 34
db.merge_cells("I6:J6")
db.merge_cells("I7:J7")
db["I6"] = "This paycheck"
db["I7"] = f'=IFERROR(INDEX({PK("K")},$N$2),"")'
for a, size in (("I6", 11), ("I7", 18)):
    db[a].font = font(size, True, "FFFFFF")
    db[a].fill = fill(TEAL_D)
    db[a].alignment = Alignment(horizontal="center", vertical="center")
db.conditional_formatting.add("I7", FormulaRule(formula=['LEFT($I$7,5)="Short"'], fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
db.conditional_formatting.add("I7", FormulaRule(formula=['$I$7="Tight"'], fill=fill(WARN), font=Font(name=F, bold=True, color="1E2A2F")))

header(db, 9, 2, ["Due", "Bill", "Amount", "Autopay", "Category"])
for k in range(12):
    r = 10 + k
    db.cell(r, 13, f'=IFERROR(SMALL({BL("AK")},{k + 1}),"")')
    hit = f"MATCH($M{r},{BL('AK')},0)"
    style(db.cell(r, 2, f'=IF($M{r}="","",INT($M{r}))'), False, DATE, align="center")
    style(db.cell(r, 3, f'=IF($M{r}="","",INDEX({BL("B")},{hit}))'), False)
    style(db.cell(r, 4, f'=IF($M{r}="","",INDEX({BL("C")},{hit}))'), False, MONEY)
    style(db.cell(r, 5, f'=IF($M{r}="","",INDEX({BL("G")},{hit}))'), False, align="center")
    style(db.cell(r, 6, f'=IF($M{r}="","",INDEX({BL("F")},{hit}))'), False)
db["B22"] = '=IF($N$1="","","The next payday comes "&E7&" days after this one.")'
db["B22"].font = font(9, color=MUTED, italic=True)

header(db, 9, 9, ["The 12 months", "Total"])
year = [("Paychecks", f'=COUNTIFS({PK("B")},">0")', "0"),
        ("Income", f'=SUM({PK("C")})', MONEY),
        ("Bills", f'=SUM({PK("F")})', MONEY),
        ("Savings", f'=SUM({PK("G")})', MONEY),
        ("Left to spend", f'=SUM({PK("H")})', NET),
        ("Tight paychecks", f'=COUNTIFS({PK("K")},"Tight")', "0"),
        ("Short paychecks", f'=COUNTIFS({PK("K")},"Short*")', "0")]
for k, (label, formula, fmt) in enumerate(year):
    r = 10 + k
    style(db.cell(r, 9, label), False, bold=True)
    style(db.cell(r, 10, formula), False, fmt)
db.conditional_formatting.add("J16", FormulaRule(formula=["$J$16>0"], fill=fill(BAD)))
db.conditional_formatting.add("J15", FormulaRule(formula=["$J$15>0"], fill=fill(WARN)))
nxt_tight = f'COUNTIFS({PK("B")},">"&TODAY(),{PK("K")},"Tight")+COUNTIFS({PK("B")},">"&TODAY(),{PK("K")},"Short*")'
db["I18"] = (f'=IF({nxt_tight}=0,"No tight or short paychecks ahead",{nxt_tight}&IF({nxt_tight}=1," paycheck ahead is"," paychecks ahead are")'
             f'&" tight or short. The Paychecks tab shows which.")')
db.merge_cells("I18:J19")
db["I18"].alignment = Alignment(wrap_text=True, vertical="center")
db["I18"].font = font(11)
db["I18"].fill = fill(INFO)
db.conditional_formatting.add("I18", FormulaRule(formula=['LEFT($I$18,2)="No"'], fill=fill(OK)))

chart = BarChart()
chart.type = "col"
chart.title = "Left to spend from each paycheck"
chart.height, chart.width = 7.5, 22
chart.add_data(Reference(pc, min_col=8, min_row=K0 - 1, max_row=K1), titles_from_data=True)
chart.set_categories(Reference(pc, min_col=2, min_row=K0, max_row=K1))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.x_axis.number_format = "DD MMM"
db.add_chart(chart, "B24")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: how often you are paid, your first payday, your usual paycheck and what you save from each one."),
    ("2", "Bills: each bill once, with the amount and the day of the month it is due. For a yearly bill, add its month."),
    ("3", "Paychecks: your paydays for 12 months are listed. Type an amount only when a paycheck is different."),
    ("4", "Dashboard: pick a paycheck to see the bills it covers, what is left and how much per day."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Each bill goes on the last paycheck before its due date, so that paycheck keeps the money for it.",
                          "Orange means less than your daily minimum is left. Red means the bills are more than the paycheck.",
                          "The example bills show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(14 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Paycheck-Budget-Planner.xlsx")
wb.save(out)
print("saved", out)
