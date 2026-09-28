"""Fee and Profit Calculator for Etsy sellers (Excel + Google Sheets).

Every fee on an order (listing, transaction, payment processing, Offsite Ads,
regulatory fee, currency conversion, VAT on fees) from an editable table of
rates per country, then profit, margin and the price that reaches a target
margin. Only functions both apps have.
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

NET = '#,##0.00;[Red]-#,##0.00'
PCT = "0.0%"
P0, P1 = 8, 207      # product rows
# Starting rates per country: currency, listing fee (0.20 US dollars in that currency, rounded),
# transaction %, payment processing % and fixed part, regulatory fee %, VAT charged on fees %.
COUNTRIES = [("United States", "USD", 0.20, 0.065, 0.03, 0.25, 0.0, 0.0),
             ("United Kingdom", "GBP", 0.16, 0.065, 0.04, 0.20, 0.0032, 0.20),
             ("France", "EUR", 0.18, 0.065, 0.04, 0.30, 0.0047, 0.20),
             ("Germany", "EUR", 0.18, 0.065, 0.04, 0.30, 0.0, 0.19),
             ("Italy", "EUR", 0.18, 0.065, 0.04, 0.30, 0.0032, 0.22),
             ("Spain", "EUR", 0.18, 0.065, 0.04, 0.30, 0.0072, 0.21),
             ("Canada", "CAD", 0.27, 0.065, 0.03, 0.25, 0.0, 0.0),
             ("Australia", "AUD", 0.30, 0.065, 0.03, 0.25, 0.0, 0.10)]
T0, T1 = 8, 8 + len(COUNTRIES) - 1

wb = Workbook()

# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Pick your country. Every calculation uses the rates on its line.",
           [3, 20, 11, 12, 13, 14, 14, 13, 13], rows=40)
se["B4"] = "Your country"
se["B4"].font = font(11, True, TEAL_D)
style(se["C4"], True, bold=True)
se["C4"] = "United States"
se.merge_cells("C4:D4")
dv = DataValidation(type="list", formula1=f"=$B${T0}:$B${T1}", allow_blank=False)
se.add_data_validation(dv)
dv.add("C4")
header(se, 7, 2, ["Country", "Currency", "Listing fee", "Transaction fee", "Payment processing",
                  "Processing fixed part", "Regulatory fee", "VAT on Etsy fees"])
for k, row in enumerate(COUNTRIES):
    r = T0 + k
    for c, (v, fmt) in enumerate(zip(row, (None, None, MONEY, PCT, PCT, MONEY, "0.00%", "0%")), start=2):
        style(se.cell(r, c, v), True, fmt, bold=c == 2, align=None if c == 2 else "center")
se.conditional_formatting.add(f"B{T0}:I{T1}", FormulaRule(formula=[f"$B{T0}=$C$4"], font=Font(name=F, bold=True, color=TEAL_D)))
notes = ["Etsy changes its fees from time to time. These rates are a starting point: check the current ones at",
         "etsy.com/legal/fees and change any number in this table. Every calculation follows.",
         "The listing fee is 0.20 US dollars, shown here in each currency and rounded. VAT on fees applies when you",
         "are not registered for VAT. In Canada, add the GST or HST of your province. Payment processing has no VAT."]
for k, text in enumerate(notes):
    se.cell(T1 + 2 + k, 2, text).font = font(10, color=MUTED, italic=True)
extra = T1 + 7
header(se, extra, 2, ["Other settings", "Value"])
others = [("Offsite Ads fee", 0.15, "0%", "15% for most shops, 12% after 10,000 US dollars of sales in a year"),
          ("Offsite Ads cap per order", 100, MONEY, "Etsy caps this fee at 100 US dollars per order"),
          ("Listing in another currency?", "No", None, "Yes adds the 2.5% currency conversion fee"),
          ("Sales tax the buyer pays", 0, "0.0%", "Only if you want processing fees on the tax too"),
          ("Your hourly rate", 15, MONEY, "For the time you spend on each item"),
          ("Target profit margin", 0.30, "0%", "Used for status colours and suggested prices")]
for k, (label, value, fmt, note) in enumerate(others):
    r = extra + 1 + k
    style(se.cell(r, 2, label), False, bold=True)
    style(se.cell(r, 3, value), True, fmt, align="center")
    se.cell(r, 4, note).font = font(10, color=MUTED, italic=True)
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
se.add_data_validation(dv)
dv.add(f"C{extra + 3}")
OA, OA_CAP, CONV, TAX, HOURLY, TARGET = (f"Settings!$C${extra + 1 + k}" for k in range(6))
row = f"MATCH(Settings!$C$4,Settings!$B${T0}:$B${T1},0)"
RATE = lambda col: f"INDEX(Settings!${col}${T0}:${col}${T1},{row})"
LF, TR, PR, PF, RR, VAT = RATE("D"), RATE("E"), RATE("F"), RATE("G"), RATE("H"), RATE("I")
CONV_RATE = f'IF({CONV}="Yes",0.025,0)'

# ---------------------------------------------------------------- Calculator
ca = wb.create_sheet("Calculator", 0)
sheet_base(ca, "Fee and profit calculator", "Type one sale in the yellow cells. Rates come from the Settings tab.",
           [3, 30, 13, 4, 30, 13, 11, 3], rows=40)
ca["B4"] = '="Country: "&Settings!$C$4&"   Currency: "&INDEX(Settings!$C$8:$C$15,' + row + ')'
ca["B4"].font = font(10, True, MUTED)
header(ca, 6, 2, ["The sale", "Amount"])
inputs = [("Item price", 18, MONEY), ("Shipping you charge", 4.5, MONEY), ("Discount on the item", 0, "0%"),
          ("Materials", 3.2, MONEY), ("Packaging", 0.8, MONEY), ("Shipping label you pay", 4.1, MONEY),
          ("Minutes of work", 12, "0"), ("Sold through an Offsite Ad?", "No", None)]
for k, (label, value, fmt) in enumerate(inputs):
    r = 7 + k
    style(ca.cell(r, 2, label), False, bold=True)
    style(ca.cell(r, 3, value), True, fmt, align="center" if fmt in (None, "0", "0%") else None)
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
ca.add_data_validation(dv)
dv.add("C14")
ORDER = "(C7*(1-C9)+C8)"
header(ca, 6, 5, ["Etsy fees", "Amount", "Share"])
fees = [("Listing fee", f"={LF}"),
        ("Transaction fee", f"={TR}*{ORDER}"),
        ("Payment processing", f"={PR}*{ORDER}*(1+{TAX})+{PF}"),
        ("Offsite Ads", f'=IF(C14="Yes",MIN({OA}*{ORDER},{OA_CAP}),0)'),
        ("Regulatory fee", f"={RR}*{ORDER}"),
        ("Currency conversion", f"={CONV_RATE}*{ORDER}"),
        ("VAT on fees", f"={VAT}*(F7+F8+F10+F11)")]
for k, (label, formula) in enumerate(fees):
    r = 7 + k
    style(ca.cell(r, 5, label), False)
    style(ca.cell(r, 6, formula), False, MONEY)
    style(ca.cell(r, 7, f'=IF({ORDER}=0,"",F{r}/{ORDER})'), False, PCT, align="center")
style(ca.cell(14, 5, "All Etsy fees"), False, bold=True)
style(ca.cell(14, 6, "=SUM(F7:F13)"), False, MONEY, True)
style(ca.cell(14, 7, f'=IF({ORDER}=0,"",F14/{ORDER})'), False, PCT, True, "center")

header(ca, 17, 2, ["What you keep", "Amount"])
result = [("The buyer pays you", f"={ORDER}", MONEY, False),
          ("Minus Etsy fees", "=-F14", NET, False),
          ("Minus your costs", f"=-(C10+C11+C12+C13/60*{HOURLY})", NET, False),
          ("Profit", "=C18+C19+C20", NET, True),
          ("Profit margin", '=IF(C18=0,"",C21/C18)', "0%", True)]
for k, (label, formula, fmt, bold) in enumerate(result):
    r = 18 + k
    style(ca.cell(r, 2, label), False, bold=bold)
    style(ca.cell(r, 3, formula), False, fmt, bold)
ca["B23"] = f'="Your time is counted at "&FIXED({HOURLY},2)&" per hour (Settings)."'
ca["B23"].font = font(9, color=MUTED, italic=True)
ca["E17"] = "Verdict"
ca["E17"].font = font(11, True, "FFFFFF")
ca["E17"].fill = fill(TEAL)
ca.merge_cells("E17:G17")
ca["E17"].alignment = Alignment(vertical="center")
ca["E18"] = (f'=IF(C18=0,"",IF(C21<0,"This sale loses money",IF(C22<{TARGET},"Profit, but below your target of "'
             f'&ROUND({TARGET}*100,0)&"%","On target: "&ROUND(C22*100,0)&"% profit margin")))')
ca.merge_cells("E18:G19")
ca["E18"].font = font(12, True)
ca["E18"].alignment = Alignment(vertical="center", wrap_text=True)
for r in (18, 19):
    for col in "EFG":
        ca[f"{col}{r}"].border = box
ca.conditional_formatting.add("E18", FormulaRule(formula=['LEFT($E$18,4)="This"'], fill=fill(BAD)))
ca.conditional_formatting.add("E18", FormulaRule(formula=['LEFT($E$18,6)="Profit"'], fill=fill(WARN)))
ca.conditional_formatting.add("E18", FormulaRule(formula=['LEFT($E$18,2)="On"'], fill=fill(OK)))
# Price that reaches the target margin: fees grow with the order, so solve for the order total.
rates = f"({TR}+{PR}*(1+{TAX})+IF(C14=\"Yes\",{OA},0)+{RR}+{CONV_RATE}+{VAT}*({TR}+IF(C14=\"Yes\",{OA},0)+{RR}))"
fixed = f"({LF}*(1+{VAT})+{PF}+C10+C11+C12+C13/60*{HOURLY})"
ca["E21"] = "Price for your target margin"
ca["E21"].font = font(11, True, TEAL_D)
ca["E22"] = (f'=IF(1-{rates}-{TARGET}<=0,"Not reachable",MAX(0,({fixed}/(1-{rates}-{TARGET})-C8)/(1-C9)))')
ca.merge_cells("E22:G23")
ca["E22"].number_format = MONEY
ca["E22"].font = font(22, True, TEAL)
ca["E22"].alignment = Alignment(horizontal="left", vertical="center")
ca["E24"] = "Item price, with the same shipping and discount as above."
ca["E24"].font = font(9, color=MUTED, italic=True)
ca.conditional_formatting.add("C21", FormulaRule(formula=["$C$21<0"], fill=fill(BAD)))

# ---------------------------------------------------------------- Products
pr = wb.create_sheet("Products", 1)
sheet_base(pr, "Your products", "One line per product. Fees, profit and a suggested price are calculated.",
           [3, 24, 10, 10, 10, 10, 10, 9, 10, 11, 11, 11, 11, 9, 20, 12], rows=P1 + 2)
pr["B4"] = '="Country: "&Settings!$C$4&".  Your time at "&FIXED(' + HOURLY + ',2)&" per hour.  Target margin "&ROUND(' + TARGET + '*100,0)&"%."'
pr["B4"].font = font(10, True, MUTED)
header(pr, 7, 2, ["Product", "Price", "Shipping charged", "Materials", "Packaging", "Label you pay", "Minutes",
                  "Share from Offsite Ads", "Buyer pays", "Etsy fees", "Your costs", "Profit", "Margin", "Status",
                  "Price for target"])
products = [("Soy candle, 200 g", 18, 4.5, 3.2, 0.8, 4.1, 12, 0.1),
            ("Brass earrings", 22, 3.5, 2.9, 0.5, 1.2, 20, 0.1),
            ("Art print, A4", 15, 3, 1.4, 0.9, 1.5, 5, 0),
            ("Knitted hat", 38, 5, 9.5, 0.7, 4.6, 180, 0),
            ("Printable planner", 6.5, 0, 0, 0, 0, 2, 0),
            ("Ceramic mug", 28, 6, 5.5, 1.5, 7.8, 40, 0.1),
            ("Sticker sheet", 4.5, 1.2, 0.6, 0.2, 0.7, 3, 0),
            ("Leather wallet", 55, 5, 14, 1.2, 4.5, 75, 0.2)]
for r in range(P0, P1 + 1):
    p = products[r - P0] if r - P0 < len(products) else (None,) * 8
    for c, (v, fmt) in enumerate(zip(p, (None, MONEY, MONEY, MONEY, MONEY, MONEY, "0", "0%")), start=2):
        style(pr.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (8, 9) else None)
    order = f"(N(C{r})+N(D{r}))"
    oa_fee = f"N(I{r})*MIN({OA}*{order},{OA_CAP})"
    fee = (f"{LF}+{TR}*{order}+{PR}*{order}*(1+{TAX})+{PF}+{oa_fee}+{RR}*{order}+{CONV_RATE}*{order}"
           f"+{VAT}*({LF}+{TR}*{order}+{oa_fee}+{RR}*{order})")
    style(pr.cell(r, 10, f'=IF(B{r}="","",{order})'), False, MONEY)
    style(pr.cell(r, 11, f'=IF(B{r}="","",{fee})'), False, MONEY)
    style(pr.cell(r, 12, f'=IF(B{r}="","",N(E{r})+N(F{r})+N(G{r})+N(H{r})/60*{HOURLY})'), False, MONEY)
    style(pr.cell(r, 13, f'=IF(B{r}="","",J{r}-K{r}-L{r})'), False, NET, True)
    style(pr.cell(r, 14, f'=IF(OR(B{r}="",N(J{r})=0),"",M{r}/J{r})'), False, "0%", align="center")
    style(pr.cell(r, 15, f'=IF(B{r}="","",IF(M{r}<0,"Loss",IF(N{r}<{TARGET},"Below target","On target")))'), False)
    rt = f"({TR}+{PR}*(1+{TAX})+N(I{r})*{OA}+{RR}+{CONV_RATE}+{VAT}*({TR}+N(I{r})*{OA}+{RR}))"
    fx = f"({LF}*(1+{VAT})+{PF}+L{r})"
    style(pr.cell(r, 16, f'=IF(B{r}="","",IF(1-{rt}-{TARGET}<=0,"Not reachable",MAX(0,{fx}/(1-{rt}-{TARGET})-N(D{r}))))'),
          False, MONEY, True)
st_rng = f"O{P0}:O{P1}"
pr.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$O{P0}="Loss"'], fill=fill(BAD), font=Font(name=F, bold=True, color="9B1C1C")))
pr.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$O{P0}="Below target"'], fill=fill(WARN)))
pr.conditional_formatting.add(st_rng, FormulaRule(formula=[f'$O{P0}="On target"'], fill=fill(OK)))
pr.conditional_formatting.add(f"P{P0}:P{P1}", FormulaRule(formula=[f'AND(ISNUMBER($P{P0}),$P{P0}>$C{P0}+0.005)'],
                                                          font=Font(name=F, bold=True, color="B4541F")))
pr["B5"] = (f'=COUNTIFS(B{P0}:B{P1},"?*")&" products:  "&COUNTIFS(O{P0}:O{P1},"On target")&" on target,  "'
            f'&COUNTIFS(O{P0}:O{P1},"Below target")&" below target,  "&COUNTIFS(O{P0}:O{P1},"Loss")&" at a loss."')
pr["B5"].font = font(11, True, TEAL_D)
pr.freeze_panes = "C8"

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: pick your country. Its fee rates are filled in, and you can change any of them."),
    ("2", "Settings: set your hourly rate and the profit margin you aim for."),
    ("3", "Calculator: type one sale, with the price, shipping and your costs, to see every fee and your profit."),
    ("4", "Products: list your products to compare them. The last column suggests a price for your target margin."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Etsy changes its fees from time to time. Check the current rates at etsy.com/legal/fees.",
                          "The example products show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel. This file is not made or endorsed by Etsy."]):
    st.cell(14 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Seller-Fee-Profit-Calculator.xlsx")
wb.save(out)
print("saved", out)
