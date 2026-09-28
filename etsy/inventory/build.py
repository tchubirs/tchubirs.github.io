"""Inventory & Sales Tracker for small sellers (Excel + Google Sheets).

Products are listed once; every purchase, sale, return or adjustment is one
line in Stock Moves. Stock levels, reorder alerts, margins and monthly sales
are calculated from those lines.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
P0, P1 = 6, 205    # product rows
M0, M1 = 6, 2005   # stock move rows
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

wb = Workbook()

# ---------------------------------------------------------------- Products
pr = wb.active
pr.title = "Products"
sheet_base(pr, "Products", "One line per product. Stock and status update from Stock Moves.",
           [3, 11, 26, 16, 11, 11, 11, 11, 10, 10, 12, 10, 18], rows=P1 + 2)
header(pr, 5, 2, ["SKU", "Product", "Category", "Unit cost", "Price", "Start stock", "Reorder at",
                  "In", "Out", "Stock now", "Margin", "Status"])
products = [("CAN-01", "Soy candle, lavender", "Candles", 4.2, 18, 20, 8),
            ("CAN-02", "Soy candle, cedar", "Candles", 4.2, 18, 15, 8),
            ("MUG-01", "Hand-painted mug", "Ceramics", 7.5, 28, 12, 5),
            ("EAR-01", "Brass earrings", "Jewelry", 3.1, 22, 30, 6),
            ("CRD-01", "Greeting card set", "Paper", 1.4, 9, 40, 10),
            ("BAG-01", "Linen tote bag", "Textile", 6.8, 24, 10, 4)]
MV = lambda col: f"'Stock Moves'!${col}${M0}:${col}${M1}"
for r in range(P0, P1 + 1):
    p = products[r - P0] if r - P0 < len(products) else (None,) * 7
    for c, v, fmt in ((2, p[0], None), (3, p[1], None), (4, p[2], None), (5, p[3], MONEY), (6, p[4], MONEY),
                      (7, p[5], "0"), (8, p[6], "0")):
        style(pr.cell(r, c, v), True, fmt, align="center" if c in (2, 7, 8) else None)
    ins = (f'=IF(B{r}="","",SUMIFS({MV("E")},{MV("C")},B{r},{MV("D")},"Purchase")'
           f'+SUMIFS({MV("E")},{MV("C")},B{r},{MV("D")},"Return")+SUMIFS({MV("E")},{MV("C")},B{r},{MV("D")},"Adjustment"))')
    style(pr.cell(r, 9, ins), False, "0", align="center")
    style(pr.cell(r, 10, f'=IF(B{r}="","",SUMIFS({MV("E")},{MV("C")},B{r},{MV("D")},"Sale"))'), False, "0", align="center")
    style(pr.cell(r, 11, f'=IF(B{r}="","",G{r}+I{r}-J{r})'), False, "0", align="center", bold=True)
    style(pr.cell(r, 12, f'=IF(OR(B{r}="",N(F{r})=0),"",(F{r}-E{r})/F{r})'), False, "0%", align="center")
    st = f'=IF(B{r}="","",IF(K{r}<=0,"Out of stock",IF(K{r}<=H{r},"Reorder now","OK")))'
    style(pr.cell(r, 13, st), False)
pr.freeze_panes = "D6"
rng = f"M{P0}:M{P1}"
pr.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{P0}="OK"'], fill=fill(OK)))
pr.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{P0}="Reorder now"'], fill=fill(WARN), font=Font(name=F, bold=True)))
pr.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{P0}="Out of stock"'], fill=fill(BAD),
                                               font=Font(name=F, bold=True, color="9B1C1C")))
SKUS = f"Products!$B${P0}:$B${P1}"

# ---------------------------------------------------------------- Stock moves
mv = wb.create_sheet("Stock Moves")
sheet_base(mv, "Stock moves", "One line per purchase, sale, return or adjustment (negative number = stock removed).",
           [3, 13, 11, 13, 9, 11, 13, 13, 26], rows=M1 + 2)
header(mv, 5, 2, ["Date", "SKU", "Type", "Qty", "Unit price", "Sale total", "Cost of sale", "Note"])
d = lambda k: today + dt.timedelta(days=k)
moves = [(d(-70), "CAN-01", "Purchase", 30, 4.2), (d(-70), "CAN-02", "Purchase", 20, 4.2), (d(-65), "MUG-01", "Purchase", 10, 7.5),
         (d(-60), "CAN-01", "Sale", 12, 18), (d(-55), "EAR-01", "Sale", 9, 22), (d(-50), "CRD-01", "Sale", 14, 9),
         (d(-40), "MUG-01", "Sale", 8, 28), (d(-33), "CAN-02", "Sale", 19, 18), (d(-30), "BAG-01", "Sale", 3, 24),
         (d(-26), "CAN-01", "Sale", 16, 18), (d(-20), "EAR-01", "Sale", 12, 22), (d(-18), "MUG-01", "Sale", 7, 28),
         (d(-14), "CRD-01", "Sale", 11, 9), (d(-12), "CAN-02", "Sale", 11, 16), (d(-9), "EAR-01", "Return", 1, 22),
         (d(-6), "BAG-01", "Sale", 7, 24), (d(-4), "CAN-01", "Sale", 9, 18), (d(-2), "CRD-01", "Adjustment", -2, 0),
         (d(-1), "MUG-01", "Sale", 3, 28)]
cost_of = f"INDEX(Products!$E${P0}:$E${P1},MATCH(C{{r}},{SKUS},0))"
for r in range(M0, M1 + 1):
    m = moves[r - M0] if r - M0 < len(moves) else (None,) * 5
    style(mv.cell(r, 2, m[0]), True, DATE)
    style(mv.cell(r, 3, m[1]), True, align="center")
    style(mv.cell(r, 4, m[2]), True, align="center")
    style(mv.cell(r, 5, m[3]), True, "0", align="center")
    style(mv.cell(r, 6, m[4]), True, MONEY)
    style(mv.cell(r, 7, f'=IF(D{r}="Sale",E{r}*F{r},IF(D{r}="Return",-E{r}*F{r},""))'), False, MONEY)
    cost = cost_of.format(r=r)
    style(mv.cell(r, 8, f'=IFERROR(IF(D{r}="Sale",E{r}*{cost},IF(D{r}="Return",-E{r}*{cost},"")),"")'), False, MONEY)
    style(mv.cell(r, 9), True)
dv_sku = DataValidation(type="list", formula1=f"={SKUS}", allow_blank=True)
dv_type = DataValidation(type="list", formula1='"Purchase,Sale,Return,Adjustment"', allow_blank=True)
mv.add_data_validation(dv_sku)
mv.add_data_validation(dv_type)
dv_sku.add(f"C{M0}:C{M1}")
dv_type.add(f"D{M0}:D{M1}")
mv.freeze_panes = "B6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Shop dashboard", "Pick a month and a year. Stock figures are always as of today.",
           [3, 24, 16, 16, 16, 16, 3, 40])
db["B4"] = "Month"
db["B4"].font = font(11, True)
style(db["C4"], True, bold=True)
db["C4"] = MONTHS[today.month - 1]
db["D4"] = "Year"
db["D4"].font = font(11, True)
db["D4"].alignment = Alignment(horizontal="right")
style(db["E4"], True, "0", bold=True)
db["E4"] = Y
dvm = DataValidation(type="list", formula1='"' + ",".join(MONTHS) + '"', allow_blank=False)
db.add_data_validation(dvm)
dvm.add("C4")
db["J1"] = '=MATCH(C4,{"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"},0)'
db["J2"] = "=DATE(E4,J1,1)"
db["J3"] = "=EOMONTH(J2,0)"
db.column_dimensions["J"].hidden = True
PR = lambda col: f"Products!${col}${P0}:${col}${P1}"
in_month = f'{MV("B")},">="&$J$2,{MV("B")},"<="&$J$3'
tiles = [
    ("B", "Sales this month", f'=SUMIFS({MV("G")},{in_month})', TEAL, MONEY),
    ("C", "Gross profit", f'=B7-SUMIFS({MV("H")},{in_month})', TEAL, MONEY),
    ("D", "Units sold", f'=SUMIFS({MV("E")},{MV("D")},"Sale",{in_month})', TEAL, "0"),
    ("E", "Stock value (cost)", f"=SUMPRODUCT({PR('E')},{PR('K')})", "B4541F", MONEY),
    ("F", "Stock value (price)", f"=SUMPRODUCT({PR('F')},{PR('K')})", TEAL, MONEY),
]
for col, label, formula, colour, fmt in tiles:
    db[f"{col}6"] = label
    db[f"{col}6"].font = font(10, True, MUTED)
    db[f"{col}7"] = formula
    db[f"{col}7"].number_format = fmt
    db[f"{col}7"].font = font(17, True, colour)
    for r in (6, 7):
        db[f"{col}{r}"].fill = fill("FFFFFF")
        db[f"{col}{r}"].border = box
db["H6"] = "To reorder now"
db["H7"] = f'=COUNTIFS({PR("M")},"Reorder now")+COUNTIFS({PR("M")},"Out of stock")'
for a, size in (("H6", 11), ("H7", 28)):
    db[a].font = font(size, True, "FFFFFF")
    db[a].fill = fill(TEAL_D)
    db[a].alignment = Alignment(horizontal="center", vertical="center")
db["H7"].number_format = "0"
db.row_dimensions[7].height = 44

db["H9"] = "Heads-up"
db["H9"].font = font(11, True, "FFFFFF")
db["H9"].fill = fill(TEAL)
best = f'INDEX({PR("C")},MATCH(MAX({PR("J")}),{PR("J")},0))'
out_n, low_n = f'COUNTIFS({PR("M")},"Out of stock")', f'COUNTIFS({PR("M")},"Reorder now")'
notes = [
    (f'=IF({out_n}=0,"Nothing is out of stock",IF({out_n}=1,"1 product is out of stock",{out_n}&" products are out of stock"))', BAD),
    (f'=IF({low_n}=0,"No product is at its reorder level",IF({low_n}=1,"1 product is at its reorder level",'
     f'{low_n}&" products are at their reorder level"))', WARN),
    (f'="Best seller so far: "&IFERROR({best},"none yet")', OK),
    (f'="Average margin: "&IFERROR(FIXED(AVERAGE({PR("L")})*100,0)&"%","n/a")', INFO),
]
for k, (formula, colour) in enumerate(notes):
    c = db.cell(10 + k, 8, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.border = box
db.conditional_formatting.add("H10:H11", FormulaRule(formula=['LEFT(H10,2)="No"'], fill=fill(OK)))
db["H9"].alignment = Alignment(vertical="center")

header(db, 9, 2, ["Month", "Sales", "Units", "Gross profit"])
for i, m in enumerate(MONTHS):
    r = 10 + i
    ms, me = f"DATE($E$4,{i + 1},1)", f"EOMONTH(DATE($E$4,{i + 1},1),0)"
    rng_m = f'{MV("B")},">="&{ms},{MV("B")},"<="&{me}'
    style(db.cell(r, 2, m), False, bold=True, align="center")
    style(db.cell(r, 3, f'=SUMIFS({MV("G")},{rng_m})'), False, MONEY)
    style(db.cell(r, 4, f'=SUMIFS({MV("E")},{MV("D")},"Sale",{rng_m})'), False, "0", align="center")
    style(db.cell(r, 5, f'=C{r}-SUMIFS({MV("H")},{rng_m})'), False, MONEY)
db.conditional_formatting.add("B10:E21", FormulaRule(formula=["ROW()-9=$J$1"], fill=fill(OK)))

chart = BarChart()
chart.type = "col"
chart.title = "Sales by month"
chart.height, chart.width = 7, 13
chart.add_data(Reference(db, min_col=3, min_row=9, max_row=21), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=10, max_row=21))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
db.add_chart(chart, "H15")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 100])
steps = [
    ("1", "Products: add each product once, with its code, cost, price, stock today and reorder level."),
    ("2", "Stock Moves: add a line each time stock changes (Purchase, Sale, Return or Adjustment)."),
    ("3", "Broken, lost or gifted items: add an Adjustment with a negative quantity."),
    ("4", "Products then shows the stock you have. Orange means reorder, red means sold out."),
    ("5", "Dashboard: sales, profit and units for the month you pick, and the value of your stock."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
st["C17"] = "The example products and moves show how it works. Delete them and add your own."
st["C17"].font = font(11, color=MUTED, italic=True)
st["C19"] = "Works in Google Sheets and Microsoft Excel, in any currency."
st["C19"].font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Inventory-Sales-Tracker.xlsx")
wb.save(out)
print("saved", out)
