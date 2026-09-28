"""Food Batch Tracker for people who make food to sell (Excel + Google Sheets).

Ingredients with their pack price and allergens; recipes that give the cost of each item; batches with a code and a
best-before date from the shelf life, and what was sold, given away or thrown out; purchases of ingredients with their
lot numbers, so the stock of each ingredient is known; labels to print for a batch, with the ingredients in order of
amount and the allergens filled in; and a Dashboard with what to sell first, what to buy and the waste by month. Only
functions both apps have.
"""
import datetime as dt
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
P0, P1 = 6, 45         # products
I0, I1 = 6, 105        # ingredients
R0, R1 = 6, 405        # recipe lines
B0, B1 = 6, 505        # batches
U0, U1 = 6, 405        # purchases
NA = 15                # allergens
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
PR = lambda col: f"Products!${col}${P0}:${col}${P1}"
IG = lambda col: f"Ingredients!${col}${I0}:${col}${I1}"
RC = lambda col: f"Recipes!${col}${R0}:${col}${R1}"
BT = lambda col: f"Batches!${col}${B0}:${col}${B1}"
BT1 = lambda col: f"Batches!${col}$1:${col}${B1}"
PU = lambda col: f"Purchases!${col}${U0}:${col}${U1}"
PU1 = lambda col: f"Purchases!${col}$1:${col}${U1}"
IG1 = lambda col: f"Ingredients!${col}$1:${col}${I1}"
RC1 = lambda col: f"Recipes!${col}$1:${col}${R1}"
SOON, PANTRY, WORD = "Settings!$C$5", "Settings!$C$6", "Settings!$C$7"
BIZ, ADDR, NOTE = "Settings!$C$8", "Settings!$C$9", "Settings!$C$10"
ALG = [f"Settings!$C${13 + k}" for k in range(NA)]
col = get_column_letter
ALG_IN = [col(13 + k) for k in range(NA)]          # Ingredients M..AA: an x for each allergen
ALG_FLAG = [col(29 + k) for k in range(NA)]        # Ingredients AC..AQ (hidden): 1 or 0
RC_FLAG = [col(10 + k) for k in range(NA)]         # Recipes J..X (hidden)
PR_COUNT = [col(16 + k) for k in range(NA)]        # Products P..AD (hidden)
LABEL_PRODUCT = "Label!$H$3"
MAX_ING, MAX_CHARS = 20, 240     # the most a label holds


def num_at(values, key, keys):
    return f"IFERROR(N(INDEX({values},MATCH({key},{keys},0))),0)"


def text_at(values, key, keys):
    at = f"INDEX({values},MATCH({key},{keys},0))"
    return f'IFERROR(IF({at}="","",{at}),"")'


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})&" "&YEAR({d})'


D = lambda n: today + dt.timedelta(days=n)

# ---------------------------------------------------------------- example data (a made-up home bakery)
allergens = ["Wheat", "Gluten", "Milk", "Eggs", "Soy", "Peanuts", "Tree nuts", "Sesame", "Fish", "Shellfish",
             "Molluscs", "Celery", "Mustard", "Lupin", "Sulphites"]
ingredients = [
    # name, unit, pack size, pack price, buy more at, not on label, allergens
    ("Plain flour", "g", 2000, 2.40, 1500, None, ["Wheat"]),
    ("Bread flour", "g", 2000, 3.20, 1500, None, ["Wheat"]),
    ("Rye flour", "g", 1000, 2.80, 500, None, ["Gluten"]),
    ("Caster sugar", "g", 1000, 1.90, 800, None, []),
    ("Butter", "g", 250, 2.60, 500, None, ["Milk"]),
    ("Eggs", "g", 600, 4.20, 600, None, ["Eggs"]),
    ("Lemons", "each", 4, 2.00, 4, None, []),
    ("Dark chocolate chips", "g", 500, 5.50, 300, None, ["Milk", "Soy"]),
    ("Cocoa powder", "g", 250, 3.80, 100, None, []),
    ("Strawberries", "g", 1000, 6.00, None, None, []),
    ("Jam sugar", "g", 1000, 2.50, 500, None, []),
    ("Rolled oats", "g", 1000, 2.20, 500, None, ["Gluten"]),
    ("Honey", "g", 340, 4.50, 340, None, []),
    ("Almonds", "g", 200, 3.60, 200, None, ["Tree nuts"]),
    ("Cinnamon", "g", 45, 2.10, 20, None, []),
    ("Milk", "ml", 1000, 1.10, 1000, None, ["Milk"]),
    ("Dried yeast", "g", 100, 2.30, 20, None, []),
    ("Peanut butter", "g", 340, 3.20, 340, None, ["Peanuts"]),
    ("Salt", "g", 1000, 0.80, 200, None, []),
    ("Vanilla extract", "ml", 60, 4.80, 20, None, []),
    ("Jars and lids", "each", 12, 9.00, 12, "x", []),
    ("Cookie bags", "each", 50, 6.00, 20, "x", []),
    ("Cake boxes", "each", 25, 11.00, 10, "x", []),
]
products = [
    # name, code, items per batch, shelf life (days), price, net weight
    ("Lemon drizzle loaf", "LEM", 2, 5, 9.00, "450 g"),
    ("Chocolate chip cookies", "CCC", 4, 7, 5.00, "6 cookies, 240 g"),
    ("Sourdough loaf", "SDL", 2, 3, 7.50, "800 g"),
    ("Strawberry jam", "JAM", 6, 180, 6.50, "340 g"),
    ("Granola", "GRA", 5, 60, 8.00, "400 g"),
    ("Cinnamon rolls", "CIN", 12, 2, 3.00, "1 roll, 110 g"),
    ("Dog biscuits", "DOG", 10, 30, 5.50, "200 g"),
    ("Brownies", "BRW", 16, 5, 2.50, "1 piece, 70 g"),
]
recipes = {
    "Lemon drizzle loaf": [("Plain flour", 450), ("Caster sugar", 450), ("Butter", 450), ("Eggs", 400), ("Lemons", 3),
                           ("Cake boxes", 2)],
    "Chocolate chip cookies": [("Plain flour", 375), ("Dark chocolate chips", 340), ("Caster sugar", 300),
                               ("Butter", 225), ("Eggs", 100), ("Vanilla extract", 5), ("Salt", 3), ("Cookie bags", 4)],
    "Sourdough loaf": [("Bread flour", 900), ("Rye flour", 100), ("Salt", 20)],
    "Strawberry jam": [("Strawberries", 1500), ("Jam sugar", 1200), ("Lemons", 1), ("Jars and lids", 6)],
    "Granola": [("Rolled oats", 1200), ("Honey", 340), ("Almonds", 200), ("Butter", 100), ("Cinnamon", 10),
                ("Cookie bags", 5)],
    "Cinnamon rolls": [("Bread flour", 600), ("Milk", 250), ("Caster sugar", 150), ("Butter", 150), ("Cinnamon", 15),
                       ("Eggs", 100), ("Salt", 8), ("Dried yeast", 7), ("Cake boxes", 2)],
    "Dog biscuits": [("Rolled oats", 500), ("Peanut butter", 340), ("Eggs", 100), ("Cookie bags", 10)],
    "Brownies": [("Caster sugar", 300), ("Butter", 250), ("Dark chocolate chips", 200), ("Plain flour", 100),
                 ("Eggs", 200), ("Cocoa powder", 60), ("Salt", 2), ("Cake boxes", 2)],
}
# Batches: made on, product, items made, sold, given or used, wasted, notes
batches = []
plan = [("Chocolate chip cookies", 8, 0.9), ("Lemon drizzle loaf", 4, 0.8), ("Brownies", 16, 0.85),
        ("Sourdough loaf", 4, 0.9), ("Cinnamon rolls", 12, 0.75)]
for m in range(10, 0, -1):                       # a batch of each most weeks, month by month
    first = D(-30 * m - 3)
    for w, (name, made, rate) in enumerate(plan):
        if (m + w) % 3 == 0:
            continue
        day = first + dt.timedelta(days=w * 5)
        sold = round(made * (rate + (m % 3) * 0.04))
        sold = min(sold, made)
        given = 1 if (m + w) % 4 == 0 and sold < made else 0
        batches.append((day, name, made, sold, given, made - sold - given, None))
for m in (9, 6, 3):
    batches.append((D(-30 * m - 10), "Strawberry jam", 12, 12, 0, 0, None))
for m in (8, 5, 2):
    batches.append((D(-30 * m - 12), "Granola", 10, 9, 1, 0, None))
batches.append((D(-100), "Dog biscuits", 10, 10, 0, 0, None))
batches += [
    (D(-40), "Strawberry jam", 12, 7, 1, 0, "Strawberries lot B-2207"),
    (D(-35), "Dog biscuits", 10, 8, 0, 0, None),
    (D(-20), "Granola", 10, 6, 0, 0, None),
    (D(-8), "Brownies", 16, 13, 0, 3, None),
    (D(-4), "Brownies", 16, 10, 2, 0, None),
    (D(-3), "Lemon drizzle loaf", 4, 2, 0, 0, None),
    (D(-3), "Cinnamon rolls", 24, 20, 0, 4, "Double batch for the market"),
    (D(-2), "Chocolate chip cookies", 8, 5, 0, 0, None),
    (D(-2), "Cinnamon rolls", 12, 9, 0, 0, None),
    (D(-1), "Sourdough loaf", 4, 3, 0, 0, None),
]
batches.sort(key=lambda b: b[0])

# Ingredient purchases: enough to cover what the batches used, with a few now running low.
used = {name: 0.0 for name, *_ in ingredients}
yields = {p[0]: p[2] for p in products}
for day, name, made, *_ in batches:
    for ing, amount in recipes[name]:
        used[ing] += amount * made / yields[name]
low = {"Butter", "Cake boxes", "Dried yeast"}
keeps = {"Butter": 30, "Eggs": 21, "Milk": 7, "Lemons": 14, "Strawberries": 5, "Dried yeast": 365}
fresh = {"Butter": (-10, 12), "Eggs": (-6, 15), "Milk": (-5, 2), "Dried yeast": (-20, -3), "Strawberries": (-41, -36),
         "Lemons": (-4, 10)}
packaging = ("Jars and lids", "Cookie bags", "Cake boxes")
purchases = []
lot = 2400
for k, (name, unit, pack, price, reorder, _, _) in enumerate(ingredients):
    target = used[name] + ((reorder or 0) * 0.5 if name in low else (reorder or 0) * 1.6 + pack * 0.3)
    packs = max(1, math.ceil(target / pack))
    n = min(5, packs)                            # bought every two months or so; the last lot is still open
    for j in range(n):
        share = packs // n + (1 if j < packs % n else 0)
        if j < n - 1:
            bought = -300 + j * 60 + k * 2
            bb = bought + keeps.get(name, 150)
        else:
            bought, bb = fresh.get(name, (-25 + k, 200))
        purchases.append([D(bought), name, share, round(share * price * (1 + 0.02 * j), 2),
                          "Corner Market" if j == n - 1 else "Wholesale Foods", f"L{lot}",
                          None if name in packaging else D(bb),
                          "x" if j < n - 1 or name == "Strawberries" else None])
        lot += 7 + j
purchases.sort(key=lambda p: p[0])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def status_colours(ws, rng, first, cell):
    for text, colour in (("Past its date", BAD), ("Use soon", WARN)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{cell}{first}="{text}"'], fill=fill(colour)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{cell}{first}="Sold out"'], font=Font(color="8A9699")))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Dates, the words on the labels, and the allergens they list.", [3, 34, 30, 4, 60],
           rows=30)
for r, label, value, fmt, note in (
        (4, "Currency", "USD", None, "Amounts have no currency sign, so any currency works."),
        (5, "Use soon: days before the date", 2, "0", "A batch this close to its date shows as Use soon."),
        (6, "Pantry: days before the date", 14, "0", "An ingredient this close to its date shows on the Dashboard."),
        (7, "Date wording on the labels", "Best before", None, "For example Best before or Use by."),
        (8, "Business name", "Little Oven Bakes", None, "On each label."),
        (9, "Address", "4 Mill Lane, Riverton", None, "On each label. Leave empty to leave it off."),
        (10, "Line at the bottom of each label", "Made in a home kitchen.", None,
         "The wording your local rules ask for.")):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, r == 8, "center" if r in (4, 5, 6) else None)
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
se["B12"] = "Allergens on the labels"
se["B12"].font = font(12, True, TEAL_D)
se["E12"] = "Rename or clear any of them to match the rules where you sell."
se["E12"].font = font(10, color=MUTED, italic=True)
for k, name in enumerate(allergens):
    se.cell(13 + k, 2, k + 1).font = font(10, color=MUTED)
    se.cell(13 + k, 2).alignment = Alignment(horizontal="right")
    style(se.cell(13 + k, 3, name), True)

# ---------------------------------------------------------------- Ingredients
ig = wb.create_sheet("Ingredients")
sheet_base(ig, "Ingredients", "Each ingredient and packaging item once, with its pack price. An x under each "
           "allergen it contains.", [3, 24, 7, 10, 10, 10, 8, 10, 10, 10, 10, 11] + [4.3] * NA, rows=I1 + 2)
header(ig, 5, 2, ["Ingredient", "Unit", "Pack size", "Pack price", "Buy more at", "Not on label", "Cost per unit",
                  "Bought", "Used", "Left", "Status"] + [f"={a}" for a in ALG])
ig.row_dimensions[5].height = 66
for k in range(NA):
    ig.cell(5, 13 + k).alignment = Alignment(horizontal="center", vertical="bottom", text_rotation=90)
    ig.cell(5, 13 + k).font = font(9, True, "FFFFFF")
for r in range(I0, I1 + 1):
    x = ingredients[r - I0] if r - I0 < len(ingredients) else (None,) * 6 + ([],)
    name, unit, pack, price, reorder, off, alg = x
    for c, v, fmt in ((2, name, None), (3, unit, None), (4, pack, "#,##0.##"), (5, price, MONEY),
                      (6, reorder, "#,##0.##"), (7, off, None)):
        style(ig.cell(r, c, v), True, fmt, c == 2, "center" if c in (3, 7) else None)
    for k in range(NA):
        style(ig.cell(r, 13 + k, "x" if allergens[k] in alg else None), True, align="center")
        ig[f"{ALG_FLAG[k]}{r}"] = f'=IF({ALG_IN[k]}{r}="",0,1)'
    style(ig.cell(r, 8, f'=IF(OR(B{r}="",N(D{r})=0),"",N(E{r})/D{r})'), False, "#,##0.0000")
    style(ig.cell(r, 9, f'=IF(B{r}="","",SUMIFS({PU("J")},{PU("C")},B{r}))'), False, "#,##0.##")
    style(ig.cell(r, 10, f'=IF(B{r}="","",SUMIFS({RC("H")},{RC("C")},B{r}))'), False, "#,##0.##")
    style(ig.cell(r, 11, f'=IF(OR(B{r}="",N(I{r})=0),"",I{r}-J{r})'), False, "#,##0.##")
    style(ig.cell(r, 12, f'=IF(OR(B{r}="",F{r}="",NOT(ISNUMBER(K{r}))),"",IF(K{r}<=N(F{r}),"Buy more",""))'),
          False, align="center")
    ig[f"AR{r}"] = f'=IF(L{r}="Buy more",ROW(),"")'                                      # AR: to buy, in order
for c in ALG_FLAG + ["AR"]:
    ig.column_dimensions[c].hidden = True
ig.conditional_formatting.add(f"L{I0}:L{I1}", FormulaRule(formula=[f'L{I0}="Buy more"'], fill=fill(WARN)))
ig.conditional_formatting.add(f"K{I0}:K{I1}", FormulaRule(formula=[f"AND(ISNUMBER(K{I0}),K{I0}<0)"], fill=fill(BAD)))
ig.freeze_panes = "C6"

# ---------------------------------------------------------------- Products
pr = wb.create_sheet("Products")
sheet_base(pr, "Products", "Each thing you sell: how many one batch makes, how long it keeps, and its price.",
           [3, 24, 7, 9, 9, 9, 16, 11, 10, 10, 9, 30, 9, 9], rows=P1 + 2)
header(pr, 5, 2, ["Product", "Code", "Items a batch", "Keeps (days)", "Price", "Net weight", "Cost a batch",
                  "Cost an item", "Profit an item", "Margin", "Contains", "In stock", "Waste"])
for r in range(P0, P1 + 1):
    x = products[r - P0] if r - P0 < len(products) else (None,) * 6
    for c, v, fmt in zip(range(2, 8), x, (None, None, "0", "0", MONEY, None)):
        style(pr.cell(r, c, v), True, fmt, c == 2, "center" if c in (3, 4, 5) else None)
    pr.cell(r, 7).alignment = Alignment(horizontal="left", indent=1)
    style(pr.cell(r, 8, f'=IF(B{r}="","",SUMIFS({RC("F")},{RC("B")},B{r}))'), False, MONEY)
    style(pr.cell(r, 9, f'=IF(OR(B{r}="",N(D{r})=0),"",H{r}/D{r})'), False, MONEY)
    style(pr.cell(r, 10, f'=IF(OR(NOT(ISNUMBER(I{r})),N(F{r})=0),"",F{r}-I{r})'), False, MONEY)
    style(pr.cell(r, 11, f'=IF(NOT(ISNUMBER(J{r})),"",J{r}/F{r})'), False, "0%", align="center")
    for k in range(NA):
        pr[f"{PR_COUNT[k]}{r}"] = f'=IF(B{r}="",0,SUMIFS({RC(RC_FLAG[k])},{RC("B")},B{r}))'
    joined = "&".join(f'IF(AND({PR_COUNT[k]}{r}>0,{ALG[k]}<>""),{ALG[k]}&", ","")' for k in range(NA))
    pr[f"AE{r}"] = f"={joined}"                                                           # AE: allergens, joined
    style(pr.cell(r, 12, f'=IF(OR(B{r}="",AE{r}=""),"",LEFT(AE{r},LEN(AE{r})-2))'), False).alignment = Alignment(horizontal="left", indent=1)
    live = (f'SUMIFS({BT("K")},{BT("C")},B{r},{BT("M")},"In stock")'
            f'+SUMIFS({BT("K")},{BT("C")},B{r},{BT("M")},"Use soon")')
    style(pr.cell(r, 13, f'=IF(B{r}="","",{live})'), False, "0", align="center")
    made = f'SUMIFS({BT("D")},{BT("C")},B{r})'
    style(pr.cell(r, 14, f'=IF(B{r}="","",IF({made}=0,"",SUMIFS({BT("G")},{BT("C")},B{r})/{made}))'), False, "0%",
          align="center")
for c in PR_COUNT + ["AE"]:
    pr.column_dimensions[c].hidden = True
pr.conditional_formatting.add(f"J{P0}:K{P1}", FormulaRule(formula=[f"AND(ISNUMBER($J{P0}),$J{P0}<0)"], fill=fill(BAD)))

# ---------------------------------------------------------------- Recipes
rc = wb.create_sheet("Recipes")
sheet_base(rc, "Recipes", "What goes into one batch of each product. Packaging too, so the cost is right.",
           [3, 24, 24, 12, 18, 10], rows=R1 + 2)
header(rc, 5, 2, ["Product", "Ingredient", "Amount a batch", "Unit", "Cost"])
lines = [(p, ing, amt) for p, _, *_ in products for ing, amt in recipes[p]]
for r in range(R0, R1 + 1):
    x = lines[r - R0] if r - R0 < len(lines) else (None,) * 3
    for c, v, fmt in zip((2, 3, 4), x, (None, None, "#,##0.##")):
        style(rc.cell(r, c, v), True, fmt, c == 2)
    style(rc.cell(r, 5, f'=IF(C{r}="","",IF(ISNA(MATCH(C{r},{IG("B")},0)),"Not on Ingredients",'
                        f'{text_at(IG("C"), f"C{r}", IG("B"))}))'), False, align="center")
    style(rc.cell(r, 6, f'=IF(OR(B{r}="",C{r}=""),"",N(D{r})*{num_at(IG("H"), f"C{r}", IG("B"))})'), False, MONEY)
    per = num_at(PR("D"), f"B{r}", PR("B"))
    rc[f"G{r}"] = f"=IF(OR(B{r}=\"\",C{r}=\"\"),0,{per})"                                 # G: items a batch
    rc[f"H{r}"] = (f'=IF(N(G{r})=0,0,N(D{r})*SUMIFS({BT("D")},{BT("C")},B{r})/G{r})')    # H: used so far
    rc[f"I{r}"] = (f'=IF(OR(B{r}="",C{r}=""),"",IF(B{r}<>{LABEL_PRODUCT},"",'            # I: order on the label
                   f'IF({num_at(IG("AS"), f"C{r}", IG("B"))}=1,"",ROUND(N(D{r})*1000,0)*100000+100000-ROW())))')
    for k in range(NA):
        rc[f"{RC_FLAG[k]}{r}"] = f'=IF(OR(B{r}="",C{r}=""),0,{num_at(IG(ALG_FLAG[k]), f"C{r}", IG("B"))})'
for c in ["G", "H", "I"] + RC_FLAG:
    rc.column_dimensions[c].hidden = True
rc.conditional_formatting.add(f"E{R0}:E{R1}", FormulaRule(formula=[f'E{R0}="Not on Ingredients"'], fill=fill(BAD)))
dropdown(rc, f"={PR('B')}", f"B{R0}:B{R1}")
dropdown(rc, f"={IG('B')}", f"C{R0}:C{R1}")
rc.freeze_panes = "B6"
for r in range(I0, I1 + 1):
    ig[f"AS{r}"] = f'=IF(G{r}="",0,1)'                                                    # AS: left off the label
ig.column_dimensions["AS"].hidden = True

# ---------------------------------------------------------------- Batches
bt = wb.create_sheet("Batches")
sheet_base(bt, "Batches", "One line for each batch you make. Update what was sold, given away and thrown out.",
           [3, 13, 24, 8, 8, 9, 8, 26, 15, 13, 7, 8, 16, 10], rows=B1 + 2)
header(bt, 5, 2, ["Made on", "Product", "Made", "Sold", "Given or used", "Wasted", "Notes", "Batch code",
                  f"={WORD}", "Left", "Days left", "Status", "Cost"])
for r in range(B0, B1 + 1):
    x = batches[r - B0] if r - B0 < len(batches) else (None,) * 7
    for c, v, fmt in zip(range(2, 9), x, (DATE, None, "0", "0", "0", "0", None)):
        style(bt.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 5, 6, 7) else None)
    bt.cell(r, 8).alignment = Alignment(horizontal="left", indent=1)
    code = text_at(PR("C"), f"C{r}", PR("B"))
    prefix = f'UPPER(IF({code}="",LEFT(SUBSTITUTE(C{r}," ",""),3),{code}))'
    stamp = f'RIGHT(YEAR(B{r}),2)&RIGHT("0"&MONTH(B{r}),2)&RIGHT("0"&DAY(B{r}),2)'
    style(bt.cell(r, 9, f'=IF(OR(C{r}="",NOT(ISNUMBER(B{r}))),"",{prefix}&"-"&{stamp}&"-"&'
                        f'COUNTIFS(C${B0}:C{r},C{r},B${B0}:B{r},B{r}))'), False, bold=True, align="center")
    bt[f"P{r}"] = f'=IF(C{r}="",0,{num_at(PR("E"), f"C{r}", PR("B"))})'                  # P: keeps (days)
    style(bt.cell(r, 10, f'=IF(OR(C{r}="",NOT(ISNUMBER(B{r})),N(P{r})=0),"",B{r}+P{r})'), False, DATE,
          align="center")
    style(bt.cell(r, 11, f'=IF(C{r}="","",N(D{r})-N(E{r})-N(F{r})-N(G{r}))'), False, "0", True, "center")
    style(bt.cell(r, 12, f'=IF(OR(C{r}="",N(K{r})<=0,NOT(ISNUMBER(J{r}))),"",J{r}-TODAY())'), False, "0",
          align="center")
    style(bt.cell(r, 13, f'=IF(C{r}="","",IF(N(K{r})<=0,"Sold out",IF(NOT(ISNUMBER(J{r})),"In stock",'
                         f'IF(J{r}<TODAY(),"Past its date",IF(J{r}-TODAY()<={SOON},"Use soon","In stock")))))'),
          False, align="center")
    bt[f"Q{r}"] = f'=IF(C{r}="",0,{num_at(PR("I"), f"C{r}", PR("B"))})'                  # Q: cost an item
    style(bt.cell(r, 14, f'=IF(C{r}="","",N(D{r})*Q{r})'), False, MONEY)
    bt[f"R{r}"] = f'=IF(OR(C{r}="",NOT(ISNUMBER(B{r}))),"",YEAR(B{r})*100+MONTH(B{r}))'  # R: month made
    bt[f"S{r}"] = f"=N(G{r})*Q{r}"                                                         # S: cost of the waste
    bt[f"T{r}"] = f'=IF(OR(M{r}="",M{r}="Sold out",NOT(ISNUMBER(J{r}))),"",J{r}*1000+ROW())'  # T: sell first
for c in "PQRST":
    bt.column_dimensions[c].hidden = True
status_colours(bt, f"I{B0}:N{B1}", B0, "$M")
dropdown(bt, f"={PR('B')}", f"C{B0}:C{B1}")
bt.freeze_panes = "D6"

# ---------------------------------------------------------------- Purchases
pu = wb.create_sheet("Purchases")
sheet_base(pu, "Purchases", "Ingredients as you buy them, with the lot number and date on the pack. An x once a pack "
           "is finished.", [3, 13, 24, 8, 10, 18, 11, 13, 9, 11, 16], rows=U1 + 2)
header(pu, 5, 2, ["Bought on", "Ingredient", "Packs", "Paid", "Where", "Lot number", "Date on pack", "Finished",
                  "Amount", "Status"])
for r in range(U0, U1 + 1):
    x = purchases[r - U0] if r - U0 < len(purchases) else (None,) * 8
    for c, v, fmt in zip(range(2, 10), x, (DATE, None, "0", MONEY, None, None, DATE, None)):
        style(pu.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4, 7, 8, 9) else None)
    pu.cell(r, 6).alignment = Alignment(horizontal="left", indent=1)
    style(pu.cell(r, 10, f'=IF(C{r}="","",N(D{r})*{num_at(IG("D"), f"C{r}", IG("B"))})'), False, "#,##0.##")
    style(pu.cell(r, 11, f'=IF(OR(C{r}="",I{r}<>"",NOT(ISNUMBER(H{r}))),"",IF(H{r}<TODAY(),"Past its date",'
                         f'IF(H{r}-TODAY()<={PANTRY},"Use soon","")))'), False, align="center")
    pu[f"L{r}"] = f'=IF(K{r}="","",H{r}*1000+ROW())'                                       # L: going off, in order
pu.column_dimensions["L"].hidden = True
status_colours(pu, f"K{U0}:K{U1}", U0, "$K")
dropdown(pu, f"={IG('B')}", f"C{U0}:C{U1}")
pu.freeze_panes = "D6"

# ---------------------------------------------------------------- Label (six to print and cut)
lb = wb.create_sheet("Label")
sheet_base(lb, "", "", [2, 46, 3, 46, 2], rows=30)
lb["B2"] = "Labels for batch"
lb["B2"].font = font(15, True, TEAL_D)
lb["B2"].alignment = Alignment(horizontal="right", vertical="center")
first_code = None
style(lb["D2"], True, bold=True)
lb["D2"].font = font(15, True, TEAL_D)
lb["D2"].alignment = Alignment(horizontal="left", vertical="center")
lb.row_dimensions[2].height = 30
dropdown(lb, f"={BT('I')}", "D2")
lb["B3"] = "Pick a batch in the yellow box, then print and cut along the lines."
lb["B3"].font = font(10, color=MUTED, italic=True)
helpers = {
    2: f'=IFERROR(MATCH(D2,{BT("I")},0)+{B0 - 1},"")',                                    # row on Batches
    3: f'=IF(H2="","",INDEX({BT1("C")},H2))',                                              # product
    4: f'=IF(H2="","",INDEX({BT1("B")},H2))',                                              # made on
    5: f'=IF(H2="","",INDEX({BT1("J")},H2))',                                              # best before
    6: f'=IF(H3="","",{text_at(PR("G"), "H3", PR("B"))})',                                 # net weight
    7: f'=IF(H3="","",{text_at(PR("L"), "H3", PR("B"))})',                                 # contains
    8: f'=IF(H2="","",INDEX({BT1("I")},H2))',                                              # batch code
}
for r, f_ in helpers.items():
    lb[f"H{r}"] = f_
for n in range(MAX_ING):                                                                   # H9..H28: ingredients
    r = 9 + n
    lb[f"I{r}"] = f"=IFERROR(LARGE({RC('I')},{n + 1}),\"\")"
    lb[f"H{r}"] = f'=IF(I{r}="","",INDEX({RC1("C")},100000-MOD(I{r},100000)))'
lb["H29"] = "=H9" + "".join(f'&IF(H{r}="","",", "&H{r})' for r in range(10, 9 + MAX_ING))
lb["H30"] = f"=COUNT({RC('I')})"                                                           # H30: how many
lb["H31"] = (f'=IF(H30=0,"",IF(H30>{MAX_ING},"Too many ingredients for one label: "&H30&". Keep it to {MAX_ING}.",'
             f'IF(LEN(H29)>{MAX_CHARS},"The ingredient list is too long for this label. Shorten the names on '
             f'Ingredients.","Ingredients: "&H29&".")))')
for c in "HI":
    lb.column_dimensions[c].hidden = True
cut = Side(style="dashed", color="9AA5A6")
texts = [
    ('=IF($H$3="","",$H$3)', font(14, True, TEAL_D), 24),
    ('=$H$31', font(9), 64),
    ('=IF($H$7="","","Contains: "&$H$7&".")', font(9, True), 16),
    ('=IF($H$3="","",IF($H$6="","","Net weight: "&$H$6&"      ")&"Batch: "&$H$8)', font(9), 15),
    (f'=IF($H$3="","","Made: "&{long_date("$H$4")}&IF(ISNUMBER($H$5),"      "&{WORD}&": "&{long_date("$H$5")},""))',
     font(9, True), 15),
    (f'=IF({BIZ}="","",{BIZ}&IF({ADDR}="","",", "&{ADDR}))', font(8), 14),
    (f'=IF({NOTE}="","",{NOTE})', font(8, color=MUTED, italic=True), 14),
]
for block in range(3):
    top = 5 + block * 9
    for k, (formula, fnt, height) in enumerate(texts):
        r = top + k
        lb.row_dimensions[r].height = height
        for c in (2, 4):
            cell = lb.cell(r, c, formula)
            cell.font = fnt
            cell.alignment = Alignment(wrap_text=True, vertical="top" if k == 1 else "center", indent=1)
            cell.border = Border(left=cut, right=cut, top=cut if k == 0 else None,
                                 bottom=cut if k == len(texts) - 1 else None)
lb.page_setup.orientation = "portrait"
lb.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Batches and best before", "What to sell first, what to buy, and how much goes to waste.",
           [3, 22, 15, 10, 14, 3, 20, 13, 13, 14, 3], rows=40)
this = "(YEAR(TODAY())*100+MONTH(TODAY()))"
db["M1"] = f"={SOON}"                                                                      # M1: days for Use soon
tile(db, "B", 4, "In stock", f'=SUMIFS({BT("K")},{BT("M")},"In stock")+SUMIFS({BT("K")},{BT("M")},"Use soon")', "0")
tile(db, "C", 4, "Use soon", f'=SUMIFS({BT("K")},{BT("M")},"Use soon")', "0", "B86E1C")
tile(db, "D", 4, "Past its date", f'=SUMIFS({BT("K")},{BT("M")},"Past its date")', "0", "B23B2E")
tile(db, "E", 4, "Waste this month",
     f'=IF(SUMIFS({BT("D")},{BT("R")},{this})=0,"",SUMIFS({BT("G")},{BT("R")},{this})/SUMIFS({BT("D")},{BT("R")},{this}))',
     "0%")
tile(db, "G", 4, "Made this month", f'=SUMIFS({BT("D")},{BT("R")},{this})', "0")
tile(db, "H", 4, "Cost of waste", f'=SUMIFS({BT("S")},{BT("R")},{this})', MONEY)
tile(db, "I", 4, "To buy", f'=COUNTIF({IG("L")},"Buy more")', "0")
tile(db, "J", 4, "Pantry going off",
     f'=COUNTIF({PU("K")},"Past its date")+COUNTIF({PU("K")},"Use soon")', "0")

header(db, 7, 2, ["Sell first", "Batch", "Left", f"={WORD}"])
for n in range(10):
    r = 8 + n
    db[f"N{r}"] = f"=IFERROR(SMALL({BT('T')},{n + 1}),\"\")"                              # N: row on Batches
    at = lambda c: f"INDEX({BT1(c)},MOD(N{r},1000))"
    style(db.cell(r, 2, f'=IF(N{r}="","",{at("C")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(N{r}="","",{at("I")})'), False, align="center")
    style(db.cell(r, 4, f'=IF(N{r}="","",{at("K")})'), False, "0", align="center")
    style(db.cell(r, 5, f'=IF(N{r}="","",{at("J")})'), False, DATE, align="center")
db.conditional_formatting.add("B8:E17", FormulaRule(formula=["AND(ISNUMBER($E8),$E8<TODAY())"], fill=fill(BAD)))
db.conditional_formatting.add("B8:E17", FormulaRule(formula=["AND(ISNUMBER($E8),$E8-TODAY()<=$M$1)"], fill=fill(WARN)))

header(db, 7, 7, ["To buy", "Left", "Buy more at", "Unit"])
for n in range(4):
    r = 8 + n
    db[f"O{r}"] = f"=IFERROR(SMALL({IG('AR')},{n + 1}),\"\")"                              # O: row on Ingredients
    at = lambda c: f"INDEX({IG1(c)},O{r})"
    style(db.cell(r, 7, f'=IF(O{r}="","",{at("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(O{r}="","",{at("K")})'), False, "#,##0.##", align="center")
    style(db.cell(r, 9, f'=IF(O{r}="","",N({at("F")}))'), False, "#,##0.##", align="center")
    style(db.cell(r, 10, f'=IF(O{r}="","",{at("C")})'), False, align="center")
header(db, 13, 7, ["Pantry", "Lot", "Date on pack", "Status"])
for n in range(4):
    r = 14 + n
    db[f"O{r}"] = f"=IFERROR(SMALL({PU('L')},{n + 1}),\"\")"                               # O: row on Purchases
    at = lambda c: f"INDEX({PU1(c)},MOD(O{r},1000))"
    style(db.cell(r, 7, f'=IF(O{r}="","",{at("C")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(O{r}="","",{at("G")})'), False, align="center")
    style(db.cell(r, 9, f'=IF(O{r}="","",{at("H")})'), False, DATE, align="center")
    style(db.cell(r, 10, f'=IF(O{r}="","",{at("K")})'), False, align="center")
status_colours(db, "G14:J17", 14, "$J")

for r in range(8, 18):
    db.row_dimensions[r].height = 18
    for c in range(2, 11):
        cell = db.cell(r, c)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="center")
header(db, 19, 2, ["Month made", "Made", "Wasted", "Waste"])
for k in range(12):
    r = 20 + k
    start = f"DATE(YEAR(TODAY()),MONTH(TODAY())-{11 - k},1)"
    db[f"M{r}"] = f"=YEAR({start})*100+MONTH({start})"                                     # M: month key
    style(db.cell(r, 2, f'=CHOOSE(MONTH({start}),{MON_LIST})&" "&YEAR({start})'), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({BT("D")},{BT("R")},M{r})'), False, "0", align="center")
    style(db.cell(r, 4, f'=SUMIFS({BT("G")},{BT("R")},M{r})'), False, "0", align="center")
    style(db.cell(r, 5, f'=IF(C{r}=0,"",D{r}/C{r})'), False, "0%", align="center")
for c in "MNO":
    db.column_dimensions[c].hidden = True

chart = BarChart()
chart.type = "col"
chart.title = "Made and wasted"
chart.height, chart.width = 7.4, 12.8
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=19, max_row=31), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=20, max_row=31))
for s, colour in zip(chart.series, (TEAL, "E07A5F")):
    s.graphicalProperties.solidFill = colour
chart.y_axis.number_format = "0"
db.add_chart(chart, "G19")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your business name and address for the labels, the date wording, and the allergen names."),
    ("2", "Ingredients: each ingredient and packaging item with its pack size and price, and an x for its allergens."),
    ("3", "Products: what you sell, how many one batch makes, how many days it keeps, and the price."),
    ("4", "Recipes: what goes into one batch of each product. The cost of each item fills in by itself."),
    ("5", "Batches: a line each time you bake or cook. The code and best-before date appear by themselves."),
    ("6", "Label: pick a batch and print. Purchases: log what you buy to keep the ingredient stock."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example bakery, its products and its prices are made up. Delete them and add your own.",
                          "Check your labels against the food rules where you sell. The file puts together what you "
                          "type; it does not know the law.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(18 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Batches", "Label", "Products", "Recipes", "Ingredients", "Purchases", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
# The label shows this week's cookies from the example.
pick = next(b for b in batches if b[1] == "Chocolate chip cookies" and b[0] == D(-2))
same = sum(1 for b in batches if b[1] == pick[1] and b[0] == pick[0])
lb["D2"] = f"CCC-{pick[0]:%y%m%d}-{same}"
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Food-Batch-Tracker.xlsx")
wb.save(out)
print("saved", out, len(batches), "batches", len(purchases), "purchases", len(lines), "recipe lines")
