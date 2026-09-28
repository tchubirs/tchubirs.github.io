"""Weekly Meal Planner with a grocery list that builds itself (Excel + Google Sheets).

Recipes and their ingredients, a pantry with store sections, packs and prices,
a week plan with the number of people each day, and a grocery list in store
order: what the week needs, minus what is at home, rounded up to whole packs.
Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
FIRST_DAY = today - dt.timedelta(days=today.weekday())   # Monday of this week
R0, R1 = 6, 105      # recipes
I0, I1 = 6, 1005     # ingredient lines
P0, P1 = 6, 305      # pantry items
G0, G1 = 8, 87       # grocery list lines
DAYS = range(3, 10)  # week plan columns C..I
M0, M1 = 9, 12       # week plan meal rows
PLAN = f"'Week plan'!$C${M0}:$I${M1}"
PEOPLE = f"'Week plan'!$Z${M0}:$AF${M1}"   # hidden: people eating each planned meal
RC = lambda col: f"Recipes!${col}${R0}:${col}${R1}"
IN = lambda col: f"Ingredients!${col}${I0}:${col}${I1}"
PA = lambda col: f"Pantry!${col}${P0}:${col}${P1}"
SECTIONS = ["Fruit and vegetables", "Bread and bakery", "Meat and fish", "Dairy and eggs", "Pasta, rice and grains",
            "Cans and jars", "Sauces and oils", "Spices and baking", "Frozen", "Snacks and drinks", "Household", "Other"]
SEC = "Settings!$B$6:$B$17"

# ---------------------------------------------------------------- example data
FV, BB, MF, DE, PG, CJ, SO, SB, FR, SD, HH = SECTIONS[:11]
pantry = [
    # item, unit, store section, pack, units in a pack, pack price, have at home, extra packs
    ("Onion", "each", FV, "bag of 6", 6, 3.49, 1, None),
    ("Garlic", "clove", FV, "bulb", 10, 0.60, 4, None),
    ("Carrot", "each", FV, "2 lb bag", 12, 1.49, None, None),
    ("Celery", "stalk", FV, "bunch", 8, 1.99, None, None),
    ("Bell pepper", "each", FV, "each", 1, 1.00, None, None),
    ("Broccoli", "head", FV, "each", 1, 1.99, None, None),
    ("Baby potatoes", "lb", FV, "1.5 lb bag", 1.5, 3.49, None, None),
    ("Green beans", "lb", FV, "1 lb bag", 1, 2.29, None, None),
    ("Lemon", "each", FV, "each", 1, 0.69, None, None),
    ("Lime", "each", FV, "each", 1, 0.35, None, None),
    ("Avocado", "each", FV, "each", 1, 1.25, None, None),
    ("Tomato", "each", FV, "each", 1, 0.60, None, None),
    ("Cherry tomatoes", "oz", FV, "10 oz box", 10, 2.99, None, None),
    ("Cucumber", "each", FV, "each", 1, 0.79, None, None),
    ("Spinach", "oz", FV, "5 oz bag", 5, 2.99, None, None),
    ("Romaine hearts", "each", FV, "pack of 3", 3, 3.49, None, None),
    ("Apple", "each", FV, "bag of 6", 6, 4.49, 2, None),
    ("Banana", "each", FV, "bunch of 6", 6, 1.49, None, 1),
    ("Sliced bread", "slice", BB, "loaf", 20, 3.29, 6, None),
    ("Tortillas", "each", BB, "pack of 10", 10, 3.19, None, None),
    ("Pizza dough", "each", BB, "each", 1, 2.99, None, None),
    ("Chicken breast", "lb", MF, "2 lb pack", 2, 7.98, None, None),
    ("Ground beef", "lb", MF, "1 lb pack", 1, 5.49, None, None),
    ("Salmon fillet", "each", MF, "pack of 2", 2, 9.99, None, None),
    ("Pepperoni", "oz", MF, "6 oz pack", 6, 3.49, None, None),
    ("Eggs", "each", DE, "dozen", 12, 3.49, 4, None),
    ("Milk", "cup", DE, "half gallon", 8, 2.79, 2, 1),
    ("Greek yogurt", "oz", DE, "32 oz tub", 32, 5.49, None, None),
    ("Butter", "tbsp", DE, "1 lb box", 32, 4.99, 20, None),
    ("Cheddar", "oz", DE, "8 oz block", 8, 3.29, None, None),
    ("Parmesan", "oz", DE, "5 oz wedge", 5, 4.49, 1, None),
    ("Mozzarella", "oz", DE, "8 oz bag", 8, 3.49, None, None),
    ("Spaghetti", "oz", PG, "1 lb box", 16, 1.49, None, None),
    ("Penne", "oz", PG, "1 lb box", 16, 1.49, None, None),
    ("Rice", "cup", PG, "2 lb bag", 4.5, 2.99, 2, None),
    ("Rolled oats", "cup", PG, "18 oz canister", 6, 3.49, 1, None),
    ("Red lentils", "cup", PG, "1 lb bag", 2.25, 1.99, None, None),
    ("Canned tomatoes", "can", CJ, "can", 1, 1.19, 1, None),
    ("Black beans", "can", CJ, "can", 1, 0.99, None, None),
    ("Tuna", "can", CJ, "can", 1, 1.29, None, None),
    ("Vegetable stock", "cup", CJ, "32 oz carton", 4, 2.49, None, None),
    ("Pizza sauce", "jar", CJ, "jar", 1, 2.49, None, None),
    ("Peanut butter", "tbsp", CJ, "16 oz jar", 28, 2.99, 10, None),
    ("Olive oil", "tbsp", SO, "17 oz bottle", 33, 7.99, 20, None),
    ("Soy sauce", "tbsp", SO, "10 oz bottle", 20, 2.49, 6, None),
    ("Caesar dressing", "tbsp", SO, "12 oz bottle", 24, 3.29, None, None),
    ("Mayonnaise", "tbsp", SO, "15 oz jar", 30, 3.99, 12, None),
    ("Honey", "tbsp", SO, "12 oz bottle", 16, 4.49, 10, None),
    ("Flour", "cup", SB, "5 lb bag", 18, 3.49, 6, None),
    ("Ground cumin", "tsp", SB, "jar", 20, 2.99, 12, None),
    ("Frozen berries", "cup", FR, "16 oz bag", 3.5, 3.99, None, None),
    ("Coffee", "bag", SD, "12 oz bag", 1, 8.99, None, 1),
    ("Dish soap", "bottle", HH, "bottle", 1, 2.99, None, 1),
]
recipes = [
    # name, meal, serves, minutes, notes, ingredients in the pantry unit
    ("Overnight oats", "Breakfast", 1, 5, "Make it the night before",
     [("Rolled oats", 0.5), ("Milk", 0.5), ("Greek yogurt", 2), ("Frozen berries", 0.5), ("Honey", 1)]),
    ("Scrambled eggs on toast", "Breakfast", 2, 10, None,
     [("Eggs", 4), ("Sliced bread", 4), ("Butter", 1)]),
    ("Veggie omelette", "Breakfast", 1, 15, None,
     [("Eggs", 3), ("Bell pepper", 0.5), ("Spinach", 1), ("Cheddar", 1)]),
    ("Pancakes", "Breakfast", 4, 25, "Weekend breakfast",
     [("Flour", 1.5), ("Milk", 1.25), ("Eggs", 1), ("Butter", 2), ("Honey", 2)]),
    ("Chicken stir fry with rice", "Dinner", 4, 30, None,
     [("Chicken breast", 1.5), ("Bell pepper", 2), ("Broccoli", 1), ("Garlic", 3), ("Soy sauce", 4), ("Rice", 1.5),
      ("Olive oil", 1)]),
    ("Spaghetti bolognese", "Dinner", 4, 45, "Freezes well",
     [("Ground beef", 1), ("Spaghetti", 16), ("Canned tomatoes", 2), ("Onion", 1), ("Garlic", 2), ("Carrot", 1),
      ("Parmesan", 1.5), ("Olive oil", 1)]),
    ("Sheet pan salmon and potatoes", "Dinner", 2, 35, None,
     [("Salmon fillet", 2), ("Baby potatoes", 1), ("Green beans", 0.5), ("Lemon", 1), ("Olive oil", 2)]),
    ("Black bean tacos", "Dinner", 4, 20, None,
     [("Black beans", 2), ("Tortillas", 8), ("Avocado", 2), ("Tomato", 2), ("Cheddar", 3), ("Lime", 1),
      ("Ground cumin", 1)]),
    ("Homemade pizza", "Dinner", 4, 30, "Friday night",
     [("Pizza dough", 2), ("Pizza sauce", 1), ("Mozzarella", 8), ("Pepperoni", 3), ("Bell pepper", 1)]),
    ("Beef chili", "Dinner", 6, 60, "Double it and freeze half",
     [("Ground beef", 1), ("Black beans", 2), ("Canned tomatoes", 2), ("Onion", 1), ("Garlic", 2), ("Ground cumin", 2)]),
    ("Chicken Caesar wraps", "Lunch", 2, 15, None,
     [("Chicken breast", 0.5), ("Tortillas", 2), ("Romaine hearts", 1), ("Parmesan", 0.5), ("Caesar dressing", 3)]),
    ("Lentil soup", "Lunch", 6, 45, "Freezes well",
     [("Red lentils", 1.5), ("Onion", 1), ("Carrot", 2), ("Celery", 2), ("Canned tomatoes", 1), ("Vegetable stock", 6),
      ("Ground cumin", 2), ("Olive oil", 1)]),
    ("Tuna pasta salad", "Lunch", 4, 20, None,
     [("Penne", 12), ("Tuna", 2), ("Cucumber", 1), ("Cherry tomatoes", 10), ("Mayonnaise", 4)]),
    ("Yogurt and berries", "Snack", 1, 2, None,
     [("Greek yogurt", 5), ("Frozen berries", 0.5), ("Honey", 1)]),
    ("Apple and peanut butter", "Snack", 1, 2, None,
     [("Apple", 1), ("Peanut butter", 2)]),
]
MEALS = ["Breakfast", "Lunch", "Dinner", "Snack"]
people = [4, 4, 4, 4, 4, 2, 4]
plan = [
    ["Overnight oats", "Scrambled eggs on toast", "Overnight oats", "Scrambled eggs on toast", "Overnight oats",
     "Veggie omelette", "Scrambled eggs on toast"],
    ["Lentil soup", "Lentil soup", "Tuna pasta salad", "Chicken Caesar wraps", "Tuna pasta salad", "Leftovers",
     "Chicken Caesar wraps"],
    ["Chicken stir fry with rice", "Spaghetti bolognese", "Black bean tacos", "Sheet pan salmon and potatoes",
     "Homemade pizza", "Eating out", "Spaghetti bolognese"],
    ["Apple and peanut butter", "Yogurt and berries", "Apple and peanut butter", "Yogurt and berries", None, None,
     "Apple and peanut butter"],
]
prep = ["Cook the soup for two days", None, "Move the salmon to the fridge", None, None, None, "Plan next week"]
lines = [(name, item, qty) for name, _, _, _, _, items in recipes for item, qty in items]

wb = Workbook()
wrap = Alignment(wrap_text=True, vertical="center")

# ---------------------------------------------------------------- Week plan
wp = wb.active
wp.title = "Week plan"
sheet_base(wp, "Week plan", "Pick a recipe for each meal. The grocery list builds itself from this plan.",
           [3, 14] + [19] * 7, rows=42)
wp["B5"] = "First day"
wp["B5"].font = font(11, True)
style(wp["C5"], True, DATE, True, "center")
wp["C5"] = FIRST_DAY
wp["D5"] = "Type the date of the first day. The other days follow."
wp["D5"].font = font(10, color=MUTED, italic=True)
header(wp, 7, 2, ["Meal"] + [""] * 7)
for k, col in enumerate(DAYS):
    wp.cell(7, col, f"=$C$5+{k}").number_format = "ddd d mmm"
style(wp.cell(8, 2, "People"), False, bold=True).alignment = Alignment(vertical="center")
for k, col in enumerate(DAYS):
    style(wp.cell(8, col, people[k]), True, "0", True, "center")
for i, meal in enumerate(MEALS):
    r = M0 + i
    style(wp.cell(r, 2, meal), True, bold=True).alignment = Alignment(vertical="center")
    for k, col in enumerate(DAYS):
        style(wp.cell(r, col, plan[i][k]), True).alignment = wrap
        wp.cell(r, 26 + k, f'=IF({L(col)}{r}="",0,N({L(col)}$8))')
        wp.cell(r, 34 + k, f'=IF({L(col)}{r}="","",IF(COUNTIF({RC("B")},{L(col)}{r})>0,"recipe","other"))')
    wp.row_dimensions[r].height = 36
style(wp.cell(13, 2, "Prep notes"), False, bold=True).alignment = Alignment(vertical="center")
for k, col in enumerate(DAYS):
    c = style(wp.cell(13, col, prep[k]), True)
    c.alignment = wrap
    c.font = font(10, italic=True)
wp.row_dimensions[13].height = 36
dv = DataValidation(type="list", formula1=f"={RC('B')}", allow_blank=True, showErrorMessage=False)
wp.add_data_validation(dv)
dv.add(f"C{M0}:I{M1}")
wp.conditional_formatting.add(f"C{M0}:I{M1}", FormulaRule(formula=[f'AH{M0}="other"'],
                                                           font=Font(name=F, italic=True, color=MUTED)))
wp.conditional_formatting.add("C7:I7", FormulaRule(formula=["C7=TODAY()"], fill=fill(WARN),
                                                   font=Font(name=F, bold=True, color=TEAL_D)))

tiles = [("C", "Meals planned", f'=COUNTIF(C{M0}:I{M1},"?*")', TEAL, "0"),
         ("D", "Recipes to cook", f'=COUNTIFS({RC("G")},">0")', TEAL, "0"),
         ("E", "Items to buy", f'=COUNTIFS({PA("L")},">0")', TEAL, "0"),
         ("F", "Shopping cost", f'=SUM({PA("M")})', "B4541F", MONEY),
         ("G", "Per person per day", "=IF(SUM(C8:I8)=0,0,F16/SUM(C8:I8))", TEAL, MONEY)]
for col, label, formula, colour, fmt in tiles:
    wp[f"{col}15"] = label
    wp[f"{col}15"].font = font(10, True, MUTED)
    wp[f"{col}16"] = formula
    wp[f"{col}16"].number_format = fmt
    wp[f"{col}16"].font = font(16, True, colour)
    wp[f"{col}16"].alignment = Alignment(horizontal="right")
    for r in (15, 16):
        wp[f"{col}{r}"].border = box
wp.row_dimensions[16].height = 34

missing = f'COUNTIFS({IN("E")},"Add to Pantry")'
empty = f'COUNTIFS({RC("G")},">0",{RC("J")},0)'
no_price = f'COUNTIFS({PA("L")},">0",{PA("G")},"")'
checks = [
    f'=IF({missing}=0,"Every ingredient is in the Pantry",IF({missing}=1,"1 ingredient line is not in the Pantry yet, '
    f'so the list leaves it out",{missing}&" ingredient lines are not in the Pantry yet, so the list leaves them out"))',
    f'=IF({empty}=0,"Every planned recipe has its ingredients",IF({empty}=1,"1 planned recipe has",'
    f'{empty}&" planned recipes have")&" no ingredients yet")',
    f'=IF({no_price}=0,"Every item on the list has a price",IF({no_price}=1,"1 item on the list has",'
    f'{no_price}&" items on the list have")&" no price yet, so the cost is too low")',
]
wp["B18"] = "Checks"
wp["B18"].font = font(11, True)
for k, formula in enumerate(checks):
    r = 18 + k
    wp.merge_cells(f"C{r}:I{r}")
    c = wp.cell(r, 3, formula)
    c.fill = fill(WARN)
    c.font = font(11)
    c.alignment = Alignment(vertical="center")
    for col in "CDEFGHI":
        wp[f"{col}{r}"].border = box
    wp.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f'LEFT(C{r},5)="Every"'], fill=fill(OK)))

wp["C22"] = "Cooking this week"
wp["C22"].font = font(13, True, TEAL_D)
header(wp, 23, 3, ["Recipe", "", "Times", "Servings", "Minutes", "Cost per serving"])
wp.merge_cells("C23:D23")
for k in range(15):
    r = 24 + k
    key = f"AP{r}"
    wp[key] = f'=IFERROR(SMALL({RC("K")},{k + 1}),"")'
    wp.merge_cells(f"C{r}:D{r}")
    pick = lambda col: f"INDEX(Recipes!${col}$1:${col}${R1},{key})"
    style(wp.cell(r, 3, f'=IF({key}="","",{pick("B")})'), False, bold=True)
    style(wp.cell(r, 4), False)
    style(wp.cell(r, 5, f'=IF({key}="","",{pick("G")})'), False, "0", align="center")
    style(wp.cell(r, 6, f'=IF({key}="","",{pick("H")})'), False, "0", align="center")
    style(wp.cell(r, 7, f'=IF({key}="","",IF({pick("E")}="","",{pick("E")}))'), False, "0", align="center")
    style(wp.cell(r, 8, f'=IF({key}="","",{pick("I")})'), False, MONEY)
for c in range(26, 43):
    wp.column_dimensions[L(c)].hidden = True
wp.freeze_panes = "C8"

# ---------------------------------------------------------------- Grocery list
gl = wb.create_sheet("Grocery list")
sheet_base(gl, "Grocery list", "What the week needs, minus what you have at home, in the order of the store.",
           [3, 22, 26, 7, 16, 11, 8], rows=G1 + 3)
gl.page_setup.orientation = "portrait"
gl["B5"] = "Week of"
gl["B5"].font = font(11, True)
gl["C5"] = "='Week plan'!C5"
gl["C5"].number_format = DATE
gl["C5"].font = font(11, True)
gl["C5"].alignment = Alignment(horizontal="left")
gl["E5"] = "Total"
gl["E5"].font = font(11, True)
gl["E5"].alignment = Alignment(horizontal="right")
gl["F5"] = f"=SUM({PA('M')})"
gl["F5"].number_format = MONEY
gl["F5"].font = font(11, True, TEAL_D)
header(gl, 7, 2, ["Section", "Item", "Buy", "Pack", "Cost", "Got it"])
for k in range(G1 - G0 + 1):
    r = G0 + k
    gl[f"I{r}"] = f'=IFERROR(SMALL({PA("O")},{k + 1}),"")'
    gl[f"J{r}"] = f'=IF(I{r}="","",MATCH(I{r},{PA("O")},0))'
    gl[f"K{r}"] = f'=IF(J{r}="","",IF(INDEX({PA("D")},J{r})&""="","Other",INDEX({PA("D")},J{r})&""))'
    c = style(gl.cell(r, 2, f'=IF(J{r}="","",IF(K{r}=K{r - 1},"",K{r}))'), False)
    c.font = font(10, True, TEAL_D)
    style(gl.cell(r, 3, f'=IF(J{r}="","",INDEX({PA("B")},J{r}))'), False)
    style(gl.cell(r, 4, f'=IF(J{r}="","",INDEX({PA("L")},J{r}))'), False, "0", True, "center")
    style(gl.cell(r, 5, f'=IF(J{r}="","",INDEX({PA("E")},J{r})&"")'), False)
    style(gl.cell(r, 6, f'=IF(J{r}="","",INDEX({PA("M")},J{r}))'), False, MONEY)
    style(gl.cell(r, 7), True, align="center")
for col in "IJK":
    gl.column_dimensions[col].hidden = True
gl.conditional_formatting.add(f"B{G0}:B{G1}", FormulaRule(formula=[f'$B{G0}<>""'], fill=fill("E6F2F2")))
gl.conditional_formatting.add(f"C{G0}:F{G1}", FormulaRule(formula=[f'$G{G0}<>""'],
                                                          font=Font(name=F, strike=True, color=MUTED)))
gl[f"B{G1 + 2}"] = "Type x under Got it as you shop. Clear that column when you start a new week."
gl[f"B{G1 + 2}"].font = font(10, color=MUTED, italic=True)
gl.freeze_panes = "B8"

# ---------------------------------------------------------------- Recipes
rc = wb.create_sheet("Recipes")
sheet_base(rc, "Recipes", "One line per recipe. Its ingredients go on the Ingredients tab.",
           [3, 30, 11, 8, 9, 28, 10, 10, 11, 12], rows=R1 + 3)
header(rc, 5, 2, ["Recipe", "Meal", "Serves", "Minutes", "Link or notes", "Times this week", "Servings this week",
                  "Cost per serving", "Ingredients"])
for r in range(R0, R1 + 1):
    rec = recipes[r - R0][:5] if r - R0 < len(recipes) else (None,) * 5
    for c, (v, fmt) in enumerate(zip(rec, (None, None, "0", "0", None)), start=2):
        style(rc.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (4, 5) else None)
    style(rc.cell(r, 7, f'=IF(B{r}="","",COUNTIFS({PLAN},B{r}))'), False, "0", align="center")
    style(rc.cell(r, 8, f'=IF(B{r}="","",SUMIFS({PEOPLE},{PLAN},B{r}))'), False, "0", align="center")
    style(rc.cell(r, 9, f'=IF(B{r}="","",ROUND(SUMIFS({IN("F")},{IN("B")},B{r})/IF(N(D{r})>0,D{r},1),2))'), False, MONEY)
    style(rc.cell(r, 10, f'=IF(B{r}="","",COUNTIFS({IN("B")},B{r}))'), False, "0", align="center")
    rc.cell(r, 11, f'=IF(N(G{r})>0,ROW(),"")')
rc.column_dimensions["K"].hidden = True
dv = DataValidation(type="list", formula1='"Breakfast,Lunch,Dinner,Snack,Any"', allow_blank=True)
rc.add_data_validation(dv)
dv.add(f"C{R0}:C{R1}")
rc.conditional_formatting.add(f"G{R0}:H{R1}", FormulaRule(formula=[f'AND($B{R0}<>"",N($G{R0})>0)'], fill=fill(OK)))
rc.conditional_formatting.add(f"J{R0}:J{R1}", FormulaRule(formula=[f'AND($B{R0}<>"",N($J{R0})=0)'], fill=fill(WARN)))
rc.freeze_panes = "C6"

# ---------------------------------------------------------------- Ingredients
ig = wb.create_sheet("Ingredients")
sheet_base(ig, "Ingredients", "One line per ingredient of each recipe, in the unit the Pantry uses.",
           [3, 30, 22, 10, 10, 10, 11], rows=I1 + 2)
header(ig, 5, 2, ["Recipe", "Ingredient", "Quantity", "Unit", "Cost", "This week"])
match_p = lambda r: f"MATCH(C{r},{PA('B')},0)"
match_r = lambda r: f"MATCH(B{r},{RC('B')},0)"
for r in range(I0, I1 + 1):
    ln = lines[r - I0] if r - I0 < len(lines) else (None,) * 3
    style(ig.cell(r, 2, ln[0]), True)
    style(ig.cell(r, 3, ln[1]), True)
    style(ig.cell(r, 4, ln[2]), True, align="center")
    style(ig.cell(r, 5, f'=IF(C{r}="","",IFERROR(INDEX({PA("C")},{match_p(r)})&"","Add to Pantry"))'), False,
          align="center")
    style(ig.cell(r, 6, f'=IF(OR(C{r}="",E{r}="Add to Pantry"),"",ROUND(N(D{r})*INDEX({PA("N")},{match_p(r)}),2))'),
          False, MONEY)
    serves = f"INDEX({RC('D')},{match_r(r)})"
    style(ig.cell(r, 7, f'=IF(OR(B{r}="",C{r}=""),"",IFERROR(ROUND(N(D{r})*N(INDEX({RC("H")},{match_r(r)}))'
                        f'/IF(N({serves})>0,{serves},1),2),0))'), False, align="center")
for col, formula in (("B", f"={RC('B')}"), ("C", f"={PA('B')}")):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ig.add_data_validation(dv)
    dv.add(f"{col}{I0}:{col}{I1}")
ig.conditional_formatting.add(f"B{I0}:B{I1}", FormulaRule(formula=[f'AND($B{I0}<>"",$B{I0}=$B{I0 - 1})'],
                                                          font=Font(name=F, color="A9B5B7")))
ig.conditional_formatting.add(f"E{I0}:E{I1}", FormulaRule(formula=[f'$E{I0}="Add to Pantry"'], fill=fill(WARN)))
ig.conditional_formatting.add(f"G{I0}:G{I1}", FormulaRule(formula=[f"N($G{I0})>0"], fill=fill(OK)))
ig.freeze_panes = "C6"

# ---------------------------------------------------------------- Pantry
pa = wb.create_sheet("Pantry")
sheet_base(pa, "Pantry", "Every food you buy, with its unit, store section, pack and price, and what you have at home.",
           [3, 20, 8, 21, 15, 9, 9, 9, 8, 10, 9, 8, 10], rows=P1 + 3)
header(pa, 5, 2, ["Item", "Unit", "Store section", "Pack", "Units in a pack", "Pack price", "Have at home",
                  "Extra packs", "Needed this week", "Still needed", "Packs to buy", "Cost"])
for r in range(P0, P1 + 1):
    it = pantry[r - P0] if r - P0 < len(pantry) else (None,) * 8
    for c, (v, fmt) in enumerate(zip(it, (None, None, None, None, None, MONEY, None, "0")), start=2):
        style(pa.cell(r, c, v), True, fmt, bold=c == 2, align="center" if c in (3, 6, 8, 9) else None)
    style(pa.cell(r, 10, f'=IF(B{r}="","",ROUND(SUMIFS({IN("G")},{IN("C")},B{r}),2))'), False, align="center")
    style(pa.cell(r, 11, f'=IF(B{r}="","",MAX(0,ROUND(J{r}-N(H{r}),2)))'), False, align="center")
    style(pa.cell(r, 12, f'=IF(B{r}="","",IF(N(K{r})=0,0,ROUNDUP(ROUND(K{r}/IF(N(F{r})>0,F{r},1),6),0))+N(I{r}))'),
          False, "0", True, "center")
    style(pa.cell(r, 13, f'=IF(B{r}="","",L{r}*N(G{r}))'), False, MONEY)
    pa.cell(r, 14, f'=IF(B{r}="","",N(G{r})/IF(N(F{r})>0,F{r},1))')                  # price of one unit
    pa.cell(r, 15, f'=IF(OR(B{r}="",N(L{r})=0),"",IFERROR(MATCH(D{r},{SEC},0),13)*1000+ROW())')   # list order
for col in "NO":
    pa.column_dimensions[col].hidden = True
dv = DataValidation(type="list", formula1=f"={SEC}", allow_blank=True)
pa.add_data_validation(dv)
dv.add(f"D{P0}:D{P1}")
pa.conditional_formatting.add(f"L{P0}:M{P1}", FormulaRule(formula=[f"N($L{P0})>0"], fill=fill(OK)))
pa.freeze_panes = "C6"

# ---------------------------------------------------------------- Settings
se = wb.create_sheet("Settings")
sheet_base(se, "Settings", "Store sections in the order you walk through the store. The grocery list follows it.",
           [3, 30, 4, 60], rows=30)
header(se, 5, 2, ["Store sections, in order"])
for k, name in enumerate(SECTIONS):
    style(se.cell(6 + k, 2, name), True)
se["D6"] = "Rename them or change the order to match your store. Items with no section go last."
se["D6"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Pantry: the foods you buy, the unit you cook with, the store section, the pack you buy and its price."),
    ("2", "Recipes: one line per recipe, with how many people it serves."),
    ("3", "Ingredients: one line per ingredient of each recipe, in the unit the Pantry uses."),
    ("4", "Week plan: type the first day and the number of people each day, then pick a recipe for each meal."),
    ("5", "Before you shop, fill in what you have at home on the Pantry tab. The Grocery list does the rest."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Cooking once and eating it twice? Put the recipe on both days. The list buys enough for both.",
                          "Anything you type that is not a recipe, like Eating out, adds nothing to the list.",
                          "Extra packs on the Pantry tab add things you buy anyway, like coffee or dish soap.",
                          "The example recipes, pantry and plan show how it works. Replace them with your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency and any unit."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Weekly-Meal-Planner.xlsx")
wb.save(out)
print("saved", out, len(recipes), "recipes", len(lines), "ingredient lines", len(pantry), "pantry items")
