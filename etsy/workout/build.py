"""Workout Log and Progress Tracker (Excel + Google Sheets).

A weekly plan, a log with one line per exercise and up to five sets, new
personal bests found by themselves, sets per muscle group, weekly volume, a
month calendar of the days you trained, and body weight and measurements.
Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
MONDAY = today - dt.timedelta(days=today.weekday())
G0, G1 = 6, 1505      # log lines
X0, X1 = 6, 105       # exercises
B0, B1 = 6, 405       # body entries
LG = lambda col: f"Log!${col}${G0}:${col}${G1}"
EX = lambda col: f"Exercises!${col}${X0}:${col}${X1}"
BD = lambda col: f"Body!${col}${B0}:${col}${B1}"
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
MUSCLES = ["Chest", "Back", "Shoulders", "Arms", "Legs", "Glutes", "Core"]
NUM = "#,##0.##"
UNIT = "Dashboard!$C$4"

# ---------------------------------------------------------------- example data
exercises = [
    ("Bench press", "Chest", "Weights"), ("Incline dumbbell press", "Chest", "Weights"), ("Push-up", "Chest", "Bodyweight"),
    ("Deadlift", "Back", "Weights"), ("Pull-up", "Back", "Bodyweight"), ("Barbell row", "Back", "Weights"),
    ("Lat pulldown", "Back", "Weights"), ("Overhead press", "Shoulders", "Weights"),
    ("Lateral raise", "Shoulders", "Weights"), ("Biceps curl", "Arms", "Weights"), ("Triceps pushdown", "Arms", "Weights"),
    ("Squat", "Legs", "Weights"), ("Leg press", "Legs", "Weights"), ("Leg curl", "Legs", "Weights"),
    ("Calf raise", "Legs", "Weights"), ("Romanian deadlift", "Glutes", "Weights"), ("Hip thrust", "Glutes", "Weights"),
    ("Plank", "Core", "Bodyweight"), ("Hanging leg raise", "Core", "Bodyweight"),
]
plan = {  # workout: [(exercise, sets, reps)]
    "Push": [("Bench press", 4, "6 to 8"), ("Overhead press", 3, "8 to 10"), ("Incline dumbbell press", 3, "8 to 12"),
             ("Lateral raise", 3, "12 to 15"), ("Triceps pushdown", 3, "10 to 12")],
    "Pull": [("Deadlift", 3, "5"), ("Pull-up", 3, "6 to 10"), ("Barbell row", 3, "8 to 10"), ("Lat pulldown", 3, "10 to 12"),
             ("Biceps curl", 3, "10 to 12")],
    "Legs": [("Squat", 4, "5 to 6"), ("Romanian deadlift", 3, "8 to 10"), ("Leg press", 3, "10 to 12"),
             ("Leg curl", 3, "10 to 12"), ("Calf raise", 4, "12 to 15")],
    "Upper": [("Bench press", 3, "8 to 10"), ("Barbell row", 3, "8 to 10"), ("Overhead press", 3, "8 to 10"),
              ("Pull-up", 3, "6 to 10"), ("Biceps curl", 2, "10 to 12"), ("Triceps pushdown", 2, "10 to 12")],
    "Lower": [("Squat", 3, "8"), ("Romanian deadlift", 3, "8 to 10"), ("Leg press", 3, "12 to 15"), ("Calf raise", 3, "15")],
}
week_plan = ["Push", "Pull", "Legs", "Rest", "Upper", "Lower", "Rest"]
start_w = {"Bench press": 135, "Overhead press": 85, "Incline dumbbell press": 45, "Lateral raise": 15,
           "Triceps pushdown": 40, "Deadlift": 225, "Pull-up": 0, "Barbell row": 115, "Lat pulldown": 100,
           "Biceps curl": 25, "Squat": 185, "Romanian deadlift": 155, "Leg press": 270, "Leg curl": 70, "Calf raise": 90}
step = {"Bench press": 5, "Overhead press": 5, "Incline dumbbell press": 5, "Lateral raise": 2.5, "Triceps pushdown": 5,
        "Deadlift": 10, "Pull-up": 0, "Barbell row": 5, "Lat pulldown": 5, "Biceps curl": 2.5, "Squat": 10,
        "Romanian deadlift": 10, "Leg press": 20, "Leg curl": 5, "Calf raise": 10}
rng = random.Random(5)
start = MONDAY - dt.timedelta(weeks=12)
log = []
day = start
while day <= today:
    name = week_plan[day.weekday()]
    if name != "Rest" and (rng.random() > 0.06 or day > today - dt.timedelta(days=21)):
        week = (day - start).days // 7
        for ex, sets, reps in plan[name]:
            lo, hi = (int(x) for x in (reps.split(" to ") if " to " in reps else (reps, reps)))
            w = start_w[ex] + step[ex] * (week // 3)
            if name in ("Upper", "Lower") and w:
                w = round(w * 0.9 / 5) * 5
            row = []
            for s in range(sets):
                r = hi - s - rng.randint(0, 1) if ex != "Pull-up" else 6 + week // 3 - s - rng.randint(0, 1)
                row += [w if w else None, max(lo if ex != "Pull-up" else 3, min(hi if ex != "Pull-up" else 12, r))]
            note = rng.choice([None] * 14 + ["Felt strong", "Short on time", "Slept badly", "Easy, go up next time"])
            log.append((day, name, ex, row, note))
    day += dt.timedelta(days=1)
body = []
for k in range(13):
    d = start + dt.timedelta(weeks=k)
    wt = round(186.4 - 0.45 * k + rng.uniform(-0.4, 0.4), 1)
    meas = (round(36.5 - 0.11 * k, 1), round(42.0 + 0.03 * k, 1), round(40.0 - 0.05 * k, 1), round(14.5 + 0.03 * k, 1),
            round(23.5 - 0.02 * k, 1)) if k % 4 == 0 else (None,) * 5
    body.append((d, wt) + meas + ("Start" if k == 0 else None,))

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


# ---------------------------------------------------------------- Log
lg = wb.active
lg.title = "Log"
sheet_base(lg, "Workout log", "One line per exercise, with up to five sets. For a bodyweight exercise leave the weight empty.",
           [3, 12, 10, 22] + [7, 6] * 5 + [7, 11, 11, 9, 11, 22], rows=G1 + 2)
header(lg, 5, 2, ["Date", "Workout", "Exercise", "Set 1", "Reps", "Set 2", "Reps", "Set 3", "Reps", "Set 4", "Reps",
                  "Set 5", "Reps", "Sets", "Volume", "Best set", "Est. max", "New best", "Note"])
score = lambda w, r: f"IF(N({r})=0,0,IF(N({w})>0,{w}*(1+{r}/30),{r}))"
for r in range(G0, G1 + 1):
    e = log[r - G0] if r - G0 < len(log) else (None, None, None, [], None)
    d_, name, ex, sets, note = e
    style(lg.cell(r, 2, d_), True, DATE)
    style(lg.cell(r, 3, name), True, align="center")
    style(lg.cell(r, 4, ex), True, bold=True)
    for k in range(10):
        style(lg.cell(r, 5 + k, sets[k] if k < len(sets) else None), True, NUM, align="center")
    pairs = [(f"{L(5 + 2 * k)}{r}", f"{L(6 + 2 * k)}{r}") for k in range(5)]
    style(lg.cell(r, 15, f'=IF(D{r}="","",COUNT({",".join(p[1] for p in pairs)}))'), False, "0", align="center")
    style(lg.cell(r, 16, f'=IF(D{r}="","",{"+".join(f"N({w})*N({x})" for w, x in pairs)})'), False, "#,##0", align="right")
    best = f'IF(N(X{r})=0,"",' + "".join(
        f'IF({score(w, x)}=X{r},IF(N({w})>0,{w}&" x "&{x},{x}&" reps"),' for w, x in pairs) + '""' + ")" * 6
    style(lg.cell(r, 17, f"={best}"), False, align="center")
    style(lg.cell(r, 18, f'=IF(N(X{r})=0,"",ROUND(X{r},1))'), False, NUM, align="center")
    earlier = f'COUNTIFS({LG("D")},D{r},{LG("B")},"<"&B{r})'
    style(lg.cell(r, 19, f'=IF(N(X{r})=0,"",IF(AND({earlier}>0,COUNTIFS({LG("D")},D{r},{LG("B")},"<"&B{r},'
                         f'{LG("X")},">="&X{r})=0),"New best",""))'), False, align="center")
    style(lg.cell(r, 20, note), True)
    # Hidden helpers.
    lg.cell(r, 21, f'=IF(D{r}="","",IFERROR(INDEX({EX("C")},MATCH(D{r},{EX("B")},0)),"Other"))')   # U muscle
    lg.cell(r, 22, f'=IF(B{r}="","",IF(COUNTIF($B${G0}:B{r},B{r})=1,1,0))')                     # V first of day
    lg.cell(r, 23, f'=IF(B{r}="","",B{r}-WEEKDAY(B{r},3))')                                      # W week start
    lg.cell(r, 24, f'=IF(D{r}="","",MAX({",".join(score(w, x) for w, x in pairs)}))')             # X best score
    lg.cell(r, 25, f'=IF(N(X{r})=0,"",X{r}+ROW()/1000000000)')                                    # Y unique score
    lg.cell(r, 26, f'=IF(Y{r}="","",IF(COUNTIFS({LG("D")},D{r},{LG("Y")},">"&Y{r})=0,1,0))')       # Z best flag
    lg.cell(r, 27, f'=IF(OR(B{r}="",D{r}=""),"",B{r}+ROW()/10000000)')                            # AA date key
    lg.cell(r, 28, f'=IF(AA{r}="","",IF(COUNTIFS({LG("D")},D{r},{LG("AA")},">"&AA{r})=0,1,0))')    # AB last flag
    lg.cell(r, 29, f'=IF(S{r}="New best",AA{r},"")')                                              # AC best dates
    lg.cell(r, 30, f"=ROW()")                                                                    # AD row
for c in range(21, 31):
    lg.column_dimensions[L(c)].hidden = True
dropdown(lg, "=Plan!$B$42:$B$47", f"C{G0}:C{G1}", strict=False)
dropdown(lg, f"={EX('B')}", f"D{G0}:D{G1}")
lg.conditional_formatting.add(f"S{G0}:S{G1}", FormulaRule(formula=[f'$S{G0}="New best"'], fill=fill("FCE9B8"),
                                                          font=Font(name=F, bold=True, color="7A5200")))
lg.conditional_formatting.add(f"B{G0}:D{G1}", FormulaRule(formula=[f'AND($B{G0}<>"",$V{G0}=1,ROW()>{G0})'],
                                                          fill=fill("E6F2F2")))
lg.freeze_panes = "E6"

# ---------------------------------------------------------------- Exercises
ex_ = wb.create_sheet("Exercises")
sheet_base(ex_, "Exercises", "Your exercises with their muscle group. Bests and last times come from the log.",
           [3, 24, 13, 13, 9, 13, 13, 11, 13], rows=X1 + 2)
header(ex_, 5, 2, ["Exercise", "Muscle group", "Type", "Times done", "Last done", "Best set", "Est. max", "Best on"])
best_row = lambda r: f'SUMIFS({LG("AD")},{LG("D")},B{r},{LG("Z")},1)'
for r in range(X0, X1 + 1):
    e = exercises[r - X0] if r - X0 < len(exercises) else (None,) * 3
    for c, v in enumerate(e, start=2):
        style(ex_.cell(r, c, v), True, bold=c == 2)
    style(ex_.cell(r, 5, f'=IF(B{r}="","",COUNTIFS({LG("D")},B{r}))'), False, "0", align="center")
    style(ex_.cell(r, 6, f'=IF(N(E{r})=0,"",INT(SUMIFS({LG("AA")},{LG("D")},B{r},{LG("AB")},1)))'), False, DATE)
    style(ex_.cell(r, 7, f'=IF(OR(N(E{r})=0,{best_row(r)}=0),"",INDEX(Log!$Q:$Q,{best_row(r)}))'), False, align="center")
    style(ex_.cell(r, 8, f'=IF(OR(N(E{r})=0,{best_row(r)}=0),"",INDEX(Log!$R:$R,{best_row(r)}))'), False, NUM,
          align="center")
    style(ex_.cell(r, 9, f'=IF(OR(N(E{r})=0,{best_row(r)}=0),"",INDEX(Log!$B:$B,{best_row(r)}))'), False, DATE)
dropdown(ex_, f'"{",".join(MUSCLES)}"', f"C{X0}:C{X1}")
dropdown(ex_, '"Weights,Bodyweight"', f"D{X0}:D{X1}")
ex_[f"B{X1 + 2}"] = ("Est. max is the weight you could lift once, estimated from your best set (weight x (1 + reps / 30)). "
                     "For a bodyweight exercise it is the most reps in one set.")
ex_[f"B{X1 + 2}"].font = font(10, color=MUTED, italic=True)
ex_.freeze_panes = "C6"

# ---------------------------------------------------------------- Plan
pl = wb.create_sheet("Plan")
sheet_base(pl, "Weekly plan", "Which workout on which day, and what each workout holds.", [3, 24, 7, 11, 4] * 3, rows=48)
header(pl, 5, 2, ["Day", "Workout", ""])
pl.merge_cells("C5:D5")
for k in range(7):
    r = 6 + k
    style(pl.cell(r, 2, DAYS[k]), False, bold=True)
    style(pl.cell(r, 3, week_plan[k]), True, bold=True, align="center")
    style(pl.cell(r, 4), True)
    pl.merge_cells(f"C{r}:D{r}")
pl["G6"] = "Type a workout name for each day, or Rest."
pl["G7"] = "The workouts below go into the list on the Log tab."
for a in ("G6", "G7"):
    pl[a].font = font(10, color=MUTED, italic=True)
names = list(plan)
for b in range(6):
    top, col = 15 + (b // 3) * 12, 2 + (b % 3) * 5
    name = names[b] if b < len(names) else None
    c = style(pl.cell(top, col, name), True, bold=True)
    c.font = font(12, True, TEAL_D)
    header(pl, top + 1, col, ["Exercise", "Sets", "Reps"])
    items = plan[name] if name else []
    for k in range(8):
        it = items[k] if k < len(items) else (None, None, None)
        style(pl.cell(top + 2 + k, col, it[0]), True)
        style(pl.cell(top + 2 + k, col + 1, it[1]), True, "0", align="center")
        style(pl.cell(top + 2 + k, col + 2, it[2]), True, align="center")
    pl.cell(42 + b, 2, f'=IF({L(col)}{top}="","",{L(col)}{top})')     # the workout names for the log's list
    dropdown(pl, f"={EX('B')}", f"{L(col)}{top + 2}:{L(col)}{top + 9}")
for r in range(42, 48):
    pl.row_dimensions[r].hidden = True
pl["B13"] = "Workouts"
pl["B13"].font = font(13, True, TEAL_D)

# ---------------------------------------------------------------- Body
bo = wb.create_sheet("Body")
sheet_base(bo, "Body", "Weigh in once a week, same day and time. Measure every few weeks.",
           [3, 13, 10, 9, 9, 9, 9, 9, 22, 4, 13, 10, 10], rows=B1 + 2)
header(bo, 5, 2, ["Date", "Weight", "Waist", "Chest", "Hips", "Arm", "Thigh", "Note"])
for r in range(B0, B1 + 1):
    e = body[r - B0] if r - B0 < len(body) else (None,) * 8
    for c, (v, fmt) in enumerate(zip(e, (DATE,) + (NUM,) * 6 + (None,)), start=2):
        style(bo.cell(r, c, v), True, fmt, align="center" if 3 <= c <= 8 else None)
    bo.cell(r, 16, f'=IF(OR(B{r}="",C{r}=""),"",B{r}+ROW()/10000000)')                   # P weigh-in key
bo.column_dimensions["P"].hidden = True
header(bo, 5, 11, ["Last weigh-ins", "Weight", "Goal"])
for k in range(12):
    r = 6 + k
    n = f'COUNT({BD("P")})'
    key = f'IF({n}-11+{k}<1,"",SMALL({BD("P")},{n}-11+{k}))'
    style(bo.cell(r, 11, f'=IFERROR(INT({key}),"")'), False, "d mmm", align="center")
    style(bo.cell(r, 12, f'=IFERROR(INDEX({BD("C")},MATCH({key},{BD("P")},0)),"")'), False, NUM, align="center")
    style(bo.cell(r, 13, f'=IF(K{r}="","",IF(N(Dashboard!$F$4)=0,"",Dashboard!$F$4))'), False, NUM, align="center")
chart = LineChart()
chart.title = "Weight"
chart.height, chart.width = 7.5, 15
chart.add_data(Reference(bo, min_col=12, max_col=13, min_row=5, max_row=17), titles_from_data=True)
chart.set_categories(Reference(bo, min_col=11, min_row=6, max_row=17))
chart.series[0].graphicalProperties.line.solidFill = TEAL
chart.series[0].graphicalProperties.line.width = 28000
chart.series[1].graphicalProperties.line.solidFill = "E0A21B"
chart.series[1].graphicalProperties.line.dashStyle = "dash"
chart.y_axis.number_format = "0"
bo.add_chart(chart, "K20")
bo.freeze_panes = "C6"

# ---------------------------------------------------------------- Calendar
ca = wb.create_sheet("Calendar")
sheet_base(ca, "Calendar", "Every day you trained, for the month you pick.", [3] + [16] * 7, rows=30)
ca["B4"] = "Month"
ca["B4"].font = font(11, True)
style(ca["C4"], True, "0", True, "center")
ca["C4"] = "=MONTH(TODAY())"
ca["D4"] = "Year"
ca["D4"].font = font(11, True)
ca["D4"].alignment = Alignment(horizontal="right")
style(ca["E4"], True, "0", True, "center")
ca["E4"] = "=YEAR(TODAY())"
ca["F4"] = "Shows this month until you type another month and year."
ca["F4"].font = font(10, color=MUTED, italic=True)
header(ca, 6, 2, DAYS)
first = "DATE($E$4,$C$4,1)"
for w in range(6):
    for k in range(7):
        rd, rw = 7 + 2 * w, 8 + 2 * w
        day_ = f"({first}-WEEKDAY({first},3)+{w * 7 + k})"
        c = ca.cell(rd, 2 + k, f'=IF(MONTH({day_})=$C$4,DAY({day_}),"")')
        c.font = font(9, True, MUTED)
        c.alignment = Alignment(horizontal="right")
        c.border = box
        name = f'IFERROR(INDEX({LG("C")},MATCH({day_},{LG("B")},0)),"")'
        c = ca.cell(rw, 2 + k, f'=IF(MONTH({day_})<>$C$4,"",{name})')
        c.font = font(11, True, TEAL_D)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
        ca.cell(rw, 11 + k, f'=IF(MONTH({day_})<>$C$4,0,COUNTIF({LG("B")},{day_}))')          # K..Q exercises that day
    ca.row_dimensions[8 + 2 * w].height = 30
for c in range(11, 18):
    ca.column_dimensions[L(c)].hidden = True
for w in range(6):
    rng_ = f"B{7 + 2 * w}:H{8 + 2 * w}"
    ca.conditional_formatting.add(rng_, FormulaRule(formula=[f"K${8 + 2 * w}>0"], fill=fill(OK)))
    ca.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(B${7 + 2 * w}<>\"\",DATE($E$4,$C$4,B${7 + 2 * w})=TODAY())"],
                                                    fill=fill("FCE9B8")))
month_from, month_to = "DATE($E$4,$C$4,1)", "EOMONTH(DATE($E$4,$C$4,1),0)"
within = lambda rng_: f'{rng_},">="&{month_from},{rng_},"<="&{month_to}'
tile(ca, "B", 20, "Workouts", f'=COUNTIFS({LG("V")},1,{within(LG("B"))})', "0")
tile(ca, "C", 20, "Sets", f'=SUMIFS({LG("O")},{within(LG("B"))})', "0")
tile(ca, "D", 20, "Volume", f'=SUMIFS({LG("P")},{within(LG("B"))})', "#,##0")
tile(ca, "E", 20, "New bests", f'=COUNTIFS({LG("S")},"New best",{within(LG("B"))})', "0", "B07A00")

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Training dashboard", "The last 7 days, the last 12 weeks, your new bests and your weight.",
           [3, 15, 11, 12, 17, 14, 13, 3, 20, 13, 12, 11, 3], rows=60)
db["B4"] = "Units"
db["B4"].font = font(11, True)
style(db["C4"], True, None, True, "center")
db["C4"] = "lb"
dropdown(db, '"lb,kg"', "C4")
db["E4"] = "Goal weight"
db["E4"].font = font(11, True)
db["E4"].alignment = Alignment(horizontal="right")
style(db["F4"], True, NUM, True, "center")
db["F4"] = 178
week7 = f'{LG("B")},">"&TODAY()-7,{LG("B")},"<="&TODAY()'
latest = f'MAX({BD("P")})'
w_now = f'INDEX({BD("C")},MATCH({latest},{BD("P")},0))'
w_start = f'INDEX({BD("C")},MATCH(MIN({BD("P")}),{BD("P")},0))'
tile(db, "B", 6, "Workouts, 7 days", f'=COUNTIFS({LG("V")},1,{week7})', "0")
tile(db, "C", 6, "Sets, 7 days", f'=SUMIFS({LG("O")},{week7})', "0")
tile(db, "D", 6, "Volume, 7 days", f'=SUMIFS({LG("P")},{week7})', "#,##0")
tile(db, "E", 6, "New bests, 30 days", f'=COUNTIFS({LG("S")},"New best",{LG("B")},">"&TODAY()-30)', "0", "B07A00")
tile(db, "F", 6, "Weight now", f'=IFERROR(FIXED({w_now},1)&" "&$C$4,"")', "General")
tile(db, "G", 6, "Since start", f'=IFERROR(FIXED({w_now}-{w_start},1)&" "&$C$4,"")', "General")
tile(db, "I", 6, "To goal", f'=IFERROR(IF(N($F$4)=0,"",FIXED({w_now}-$F$4,1)&" "&$C$4),"")', "General")

header(db, 9, 2, ["Week of", "Workouts", "Sets", "Volume"])
for k in range(12):
    r = 10 + k
    style(db.cell(r, 2, f"=TODAY()-WEEKDAY(TODAY(),3)-{7 * (11 - k)}"), False, "d mmm", align="center")
    style(db.cell(r, 3, f'=COUNTIFS({LG("V")},1,{LG("W")},B{r})'), False, "0", align="center")
    style(db.cell(r, 4, f'=SUMIFS({LG("O")},{LG("W")},B{r})'), False, "0", align="center")
    style(db.cell(r, 5, f'=SUMIFS({LG("P")},{LG("W")},B{r})'), False, "#,##0")
db.conditional_formatting.add("B21:E21", FormulaRule(formula=["TRUE"], fill=fill("E6F2F2")))
chart = BarChart()
chart.type = "col"
chart.title = "Volume by week"
chart.height, chart.width = 7, 13
chart.add_data(Reference(db, min_col=5, min_row=9, max_row=21), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=10, max_row=21))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.y_axis.number_format = "#,##0"
chart.legend = None
db.add_chart(chart, "B24")

header(db, 9, 9, ["Muscle group", "Sets, 7 days", "Sets, week before", ""])
for k, m in enumerate(MUSCLES):
    r = 10 + k
    style(db.cell(r, 9, m), False, bold=True)
    style(db.cell(r, 10, f'=SUMIFS({LG("O")},{LG("U")},I{r},{week7})'), False, "0", align="center")
    style(db.cell(r, 11, f'=SUMIFS({LG("O")},{LG("U")},I{r},{LG("B")},">"&TODAY()-14,{LG("B")},"<="&TODAY()-7)'),
          False, "0", align="center")
    style(db.cell(r, 12, f'=IF(N(J{r})=0,"",IF(J{r}<10,"Low",IF(J{r}>20,"High","")))'), False, align="center")
db.conditional_formatting.add("L10:L16", FormulaRule(formula=['L10="Low"'], fill=fill(INFO)))
db.conditional_formatting.add("L10:L16", FormulaRule(formula=['L10="High"'], fill=fill(WARN)))
db["I17"] = "Many programs aim for 10 to 20 hard sets per muscle group each week."
db["I17"].font = font(9, color=MUTED, italic=True)

header(db, 19, 9, ["New bests", "Best set", "Est. max", "Date"])
db.row_dimensions[19].height = None
for k in range(6):
    r = 20 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(LARGE({LG("AC")},{k + 1}),"")'
    row = f'MATCH({key},{LG("AC")},0)'
    style(db.cell(r, 9, f'=IF({key}="","",INDEX({LG("D")},{row}))'), False, bold=True)
    style(db.cell(r, 10, f'=IF({key}="","",INDEX({LG("Q")},{row}))'), False, align="center")
    style(db.cell(r, 11, f'=IF({key}="","",INDEX({LG("R")},{row}))'), False, NUM, align="center")
    style(db.cell(r, 12, f'=IF({key}="","",INT({key}))'), False, "d mmm", align="center")
db.column_dimensions["O"].hidden = True
chart = LineChart()
chart.title = "Weight"
chart.height, chart.width = 7, 11
chart.add_data(Reference(bo, min_col=12, max_col=13, min_row=5, max_row=17), titles_from_data=True)
chart.set_categories(Reference(bo, min_col=11, min_row=6, max_row=17))
chart.series[0].graphicalProperties.line.solidFill = TEAL
chart.series[0].graphicalProperties.line.width = 28000
chart.series[1].graphicalProperties.line.solidFill = "E0A21B"
chart.series[1].graphicalProperties.line.dashStyle = "dash"
chart.y_axis.number_format = "0"
chart.legend.position = "b"
db.add_chart(chart, "I28")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Dashboard: pick lb or kg and type your goal weight."),
    ("2", "Exercises: the exercises you do, with their muscle group. Add your own under the list."),
    ("3", "Plan: the workout for each day of the week, and the exercises, sets and reps of each workout."),
    ("4", "Log: after each workout, one line per exercise with the weight and reps of each set."),
    ("5", "Body: your weight once a week and your measurements every few weeks."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["A new best is a set that beats every earlier set of that exercise, by estimated max.",
                          "Volume is the weight times the reps of every set, added up.",
                          "The example log shows how it works. Delete it and start with your own workouts.",
                          "Works in Google Sheets and Microsoft Excel. It is a training log and does not give medical advice."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Workout-Log-Tracker.xlsx")
wb.save(out)
print("saved", out, len(log), "log lines", len(body), "weigh-ins")
