"""Habit and Mood Tracker for a whole year (Excel + Google Sheets).

Up to 15 habits, one tab per month to tick them off, streaks that carry over
from one month to the next, a mood and sleep row for every day, a year in
pixels of your mood, and a dashboard. Only functions both apps have.
"""
import calendar
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import MUTED, TEAL, TEAL_D, F, bar, box, fill, font, header, sheet_base, style  # noqa: E402

today = dt.date.today()
Y = today.year
TABS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
NAMES = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November",
         "December"]
YEAR = "Habits!$C$4"
H0, H1 = 8, 22          # habit rows on the Habits tab
N = H1 - H0 + 1         # 15 habits
G0 = 7                  # first habit row on a month tab
MOOD, SLEEP = G0 + N + 1, G0 + N + 2          # rows 23 and 24
S0 = 40                 # first row of the hidden streak grid on a month tab
DAY1 = 3                # column C is day 1
DONE_FILL, DONE_TEXT = "3E8E7E", "FFFFFF"
MOODS = [(5, "Great", "2E7D6B", "FFFFFF"), (4, "Good", "8CC5A8", "1E2A2F"), (3, "Okay", "F3DE8A", "1E2A2F"),
         (2, "Low", "F2A65A", "1E2A2F"), (1, "Bad", "D9645B", "FFFFFF")]
NA = "E3E7E7"


def col(day):
    return L(DAY1 + day - 1)


def done(cell):
    return f"OR({cell}=TRUE,ISTEXT({cell}))"


# ---------------------------------------------------------------- example data
habits = ["Drink 8 glasses of water", "Walk 30 minutes", "Read 20 pages", "No phone after 10 pm", "Stretch",
          "Journal", "Take vitamins", "Meditate 10 minutes"]
chance = [0.85, 0.62, 0.55, 0.45, 0.5, 0.4, 0.9, 0.5]
rng = random.Random(3)
marks, moods, sleep = {}, {}, {}
d = dt.date(Y, 1, 1)
while d <= today:
    walked = False
    for h, p in enumerate(chance):
        if d == today:
            ok = h in (0, 6) and rng.random() < p
        else:
            ok = rng.random() < min(0.95, p + (0.08 if d.month >= 6 else 0))
        if ok:
            marks[(d, h)] = "x"
            walked |= h == 1
    if d < today or rng.random() < 0.5:
        base = 3.4 + (0.6 if walked else 0) + rng.uniform(-1.3, 1.3)
        moods[d] = max(1, min(5, round(base)))
        sleep[d] = round(rng.choice([6, 6.5, 7, 7, 7.5, 7.5, 8, 8.5]) - (0.5 if moods[d] <= 2 else 0), 1)
    d += dt.timedelta(days=1)

wb = Workbook()


def mood_colours(ws, rng_, first, solid=False):
    for value, _, colour, text in MOODS:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"{first}={value}"], fill=fill(colour),
                                                        font=Font(name=F, bold=True, color=colour if solid else text)))


# ---------------------------------------------------------------- Habits (settings)
hb = wb.active
hb.title = "Habits"
sheet_base(hb, "Habits", "Your year and up to 15 habits. The month tabs and the dashboard follow this list.",
           [3, 30, 10, 4, 64], rows=40)
hb["B4"] = "Year"
hb["B4"].font = font(11, True)
style(hb["C4"], True, "0", True, "center")
hb["C4"] = Y
header(hb, 7, 2, ["Habit", ""])
hb.merge_cells("B7:C7")
for k in range(N):
    r = H0 + k
    style(hb.cell(r, 2, habits[k] if k < len(habits) else None), True, bold=True)
    style(hb.cell(r, 3), True)
    hb.merge_cells(f"B{r}:C{r}")
for k, text in enumerate(["Keep a habit on the same line all year, so its streak carries on from month to month.",
                          "On a month tab, type x in a day when you did the habit. In Google Sheets you can use",
                          "checkboxes instead: select the days and choose Insert > Checkbox.",
                          "Mood: 5 great, 4 good, 3 okay, 2 low, 1 bad. Sleep: the hours you slept that night."]):
    hb.cell(8 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- the 12 month tabs
for m in range(1, 13):
    tab = TABS[m - 1]
    ws = wb.create_sheet(tab)
    sheet_base(ws, f"{NAMES[m - 1]}", "Type x on each day you did a habit. Mood from 1 to 5, sleep in hours.",
               [3, 26] + [3.6] * 31 + [7, 7, 8, 8], rows=S0 + N + 2)
    ws["B2"] = f'="{NAMES[m - 1]} "&{YEAR}'
    first = f"DATE({YEAR},{m},1)"
    length = f"DAY(EOMONTH({first},0))"
    ws.row_dimensions[5].height = 18
    for day in range(1, 32):
        c = ws.cell(5, DAY1 + day - 1, f'=IF({day}<={length},{day},"")')
        c.font = font(9, True, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
        c = ws.cell(6, DAY1 + day - 1, f'=IF({col(day)}$5="","",CHOOSE(WEEKDAY(DATE({YEAR},{m},{day}),2),'
                                      f'"M","T","W","T","F","S","S"))')
        c.font = font(8, color=MUTED)
        c.alignment = Alignment(horizontal="center")
    c = ws.cell(5, 2, "Habit")
    c.font = font(10, True, "FFFFFF")
    c.fill = fill(TEAL)
    c.border = box
    for k, label in enumerate(["Done", "Share", "Streak now", "Best streak"]):
        c = ws.cell(5, 34 + k, label)
        c.font = font(8, True, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[5].height = 30
    so_far = (f"IF(AND({YEAR}=YEAR(TODAY()),{m}=MONTH(TODAY())),DAY(TODAY()),"
              f"IF({first}>TODAY(),0,{length}))")
    last = f"DAY(EOMONTH({first},0))"
    for k in range(N):
        r, s = G0 + k, S0 + k
        style(ws.cell(r, 2, f'=IF(Habits!B{H0 + k}="","",Habits!B{H0 + k})'), False, bold=True)
        ws.cell(r, 2).font = font(10, True)
        for day in range(1, 32):
            date = dt.date(Y, m, day) if day <= calendar.monthrange(Y, m)[1] else None
            v = marks.get((date, k)) if date else None
            c = ws.cell(r, DAY1 + day - 1, v)
            c.alignment = Alignment(horizontal="center")
            c.font = font(9, True, DONE_TEXT)
            c.border = box
            # Hidden streak grid: days in a row up to this day, carried over from the month before.
            prev = (f"IFERROR(INDEX({TABS[m - 2]}!$C${s}:$AG${s},DAY(EOMONTH(DATE({YEAR},{m - 1},1),0))),0)"
                    if m > 1 else "0") if day == 1 else f"{col(day - 1)}{s}"
            ws.cell(s, DAY1 + day - 1, f'=IF({col(day)}$5="",0,IF({done(f"{col(day)}{r}")},{prev}+1,0))')
        rng_ = f"C{r}:AG{r}"
        count = f'COUNTIF($C{s}:$AG{s},">0")'          # the streak grid is above 0 on every day done
        style(ws.cell(r, 34, f'=IF(B{r}="","",{count})'), False, "0", align="center")
        style(ws.cell(r, 35, f'=IF(OR(B{r}="",{so_far}=0),"",MIN(1,({count})/{so_far}))'), False, "0%", align="center")
        today_col = f"INDEX($C{s}:$AG{s},DAY(TODAY()))"
        before = f"IF(DAY(TODAY())=1,IFERROR(INDEX({TABS[m - 2]}!$C${s}:$AG${s},{last.replace(first, f'DATE({YEAR},{m - 1},1)')}),0)," \
                 f"INDEX($C{s}:$AG{s},DAY(TODAY())-1))" if m > 1 else f"IF(DAY(TODAY())=1,0,INDEX($C{s}:$AG{s},DAY(TODAY())-1))"
        now = (f"IF(AND({YEAR}=YEAR(TODAY()),{m}=MONTH(TODAY())),IF({today_col}>0,{today_col},{before}),"
               f"IF({first}>TODAY(),\"\",INDEX($C{s}:$AG{s},{length})))")
        style(ws.cell(r, 36, f'=IF(B{r}="","",{now})'), False, "0", True, "center")
        style(ws.cell(r, 37, f'=IF(B{r}="","",MAX($C{s}:$AG{s}))'), False, "0", align="center")
        ws.cell(s, 2, f"=B{r}")
    for r, label in ((MOOD, "Mood (1 to 5)"), (SLEEP, "Sleep (hours)")):
        style(ws.cell(r, 2, label), False, bold=True)
        for day in range(1, 32):
            date = dt.date(Y, m, day) if day <= calendar.monthrange(Y, m)[1] else None
            v = (moods if r == MOOD else sleep).get(date) if date else None
            c = ws.cell(r, DAY1 + day - 1, v)
            c.alignment = Alignment(horizontal="center")
            c.font = font(8)
            c.border = box
            c.fill = fill("FFF4C2")
    style(ws.cell(MOOD, 34, f'=IFERROR(AVERAGE(C{MOOD}:AG{MOOD}),"")'), False, "0.0", True, "center")
    style(ws.cell(SLEEP, 34, f'=IFERROR(AVERAGE(C{SLEEP}:AG{SLEEP}),"")'), False, "0.0", True, "center")
    ws.merge_cells(start_row=MOOD, start_column=34, end_row=MOOD, end_column=35)
    ws.merge_cells(start_row=SLEEP, start_column=34, end_row=SLEEP, end_column=35)
    ws.cell(MOOD, 36, "average").font = font(8, color=MUTED, italic=True)
    ws.cell(SLEEP, 36, "average").font = font(8, color=MUTED, italic=True)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{MOOD}:AG{MOOD}")
    grid = f"C{G0}:AG{SLEEP}"
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'C$5=""'], fill=fill(NA)))
    ws.conditional_formatting.add(f"C{G0}:AG{G0 + N - 1}", FormulaRule(formula=[f'AND($B{G0}<>"",{done(f"C{G0}")})'],
                                                                       fill=fill(DONE_FILL)))
    ws.conditional_formatting.add(f"C{G0}:AG{G0 + N - 1}", FormulaRule(formula=[f'$B{G0}=""'], fill=fill("F6F8F8")))
    mood_colours(ws, f"C{MOOD}:AG{MOOD}", f"C{MOOD}")
    ws.conditional_formatting.add("C5:AG5", FormulaRule(formula=[f"AND(C$5<>\"\",DATE({YEAR},{m},C$5)=TODAY())"],
                                                        fill=fill("F2C14E"), font=Font(name=F, bold=True, color=TEAL_D)))
    ws.conditional_formatting.add(f"AI{G0}:AI{G0 + N - 1}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                           end_type="num", end_value=1, end_color="8CC5A8"))
    for r in range(S0 - 1, S0 + N):
        ws.row_dimensions[r].hidden = True
    ws.freeze_panes = "C7"
    ws.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Year in pixels
px = wb.create_sheet("Year in pixels")
sheet_base(px, "Year in pixels", "Your mood on every day of the year, one square per day, from the month tabs.",
           [3, 6] + [6] * 12 + [3, 12], rows=40)
px["B2"] = f'="Year in pixels "&{YEAR}'
for m in range(12):
    c = px.cell(5, 3 + m, TABS[m])
    c.font = font(10, True, "FFFFFF")
    c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="center")
    c.border = box
for day in range(1, 32):
    r = 5 + day
    c = px.cell(r, 2, day)
    c.font = font(9, True, MUTED)
    c.alignment = Alignment(horizontal="right")
    px.row_dimensions[r].height = 15
    for m in range(12):
        src = f"{TABS[m]}!{col(day)}{MOOD}"
        c = px.cell(r, 3 + m, f'=IF({TABS[m]}!{col(day)}$5="","x",IF({src}="","",{src}))')
        c.alignment = Alignment(horizontal="center")
        c.font = font(8, color="FFFFFF")
        c.border = box
mood_colours(px, "C6:N36", "C6", solid=True)
px.conditional_formatting.add("C6:N36", FormulaRule(formula=['C6="x"'], fill=fill("FFFFFF"), font=Font(name=F, color="FFFFFF")))
header(px, 5, 16, ["Mood"])
for k, (value, label, colour, text) in enumerate(MOODS):
    c = px.cell(6 + k, 16, f"{value}  {label}")
    c.fill = fill(colour)
    c.font = font(10, True, text)
    c.border = box
px["P12"] = "An empty square is a day with no mood yet."
px["P12"].font = font(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Habit dashboard", "This month at the top, the whole year below.",
           [3, 28, 8, 9, 8, 10, 10, 10, 3, 8, 8, 8, 10, 3], rows=60)
cur = f"IF({YEAR}=YEAR(TODAY()),MONTH(TODAY()),12)"
db["B4"] = f'="Showing "&CHOOSE({cur},{",".join(chr(34) + n + chr(34) for n in NAMES)})&" "&{YEAR}'
db["B4"].font = font(12, True, TEAL_D)
pick = lambda cellref: f"CHOOSE({cur},{','.join(f'{t}!{cellref}' for t in TABS)})"
header(db, 9, 2, ["Habit", "Done", "Days", "Share", "", "Streak now", "Best this year"])
db.merge_cells("E9:F9")
for k in range(N):
    r, src = 10 + k, G0 + k
    style(db.cell(r, 2, f'=IF(Habits!B{H0 + k}="","",Habits!B{H0 + k})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",{pick(f"AH{src}")})'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",IF({YEAR}=YEAR(TODAY()),DAY(TODAY()),31))'), False, "0", align="center")
    style(db.cell(r, 5, f'=IF(B{r}="","",{pick(f"AI{src}")})'), False, "0%", align="center")
    c = style(db.cell(r, 6, f'=IF(OR(B{r}="",E{r}=""),"",{bar(f"E{r}", 8)})'), False)
    c.font = Font(name=F, size=9, color=TEAL)
    style(db.cell(r, 7, f'=IF(B{r}="","",{pick(f"AJ{src}")})'), False, "0", True, "center")
    style(db.cell(r, 8, f'=IF(B{r}="","",MAX({",".join(f"{t}!AK{src}" for t in TABS)}))'), False, "0", align="center")
db.conditional_formatting.add("G10:G24", FormulaRule(formula=["N(G10)>=7"], fill=fill("D8F0DC")))
share = f'IFERROR(SUM(C10:C24)/SUMIFS(D10:D24,B10:B24,"?*"),0)'
mood_now = pick(f"AH{MOOD}")
tiles = [("B", "Habits done this month", f"={share}", "0%"), ("C", "Longest streak now", "=MAX(G10:G24)", "0"),
         ("E", "Mood this month", f'=IFERROR({mood_now}*1,"")', "0.0"),
         ("G", "Sleep this month", f'=IFERROR({pick(f"AH{SLEEP}")}*1,"")', "0.0")]
for colL, label, formula, fmt in tiles:
    db[f"{colL}6"] = label
    db[f"{colL}6"].font = font(10, True, MUTED)
    db[f"{colL}7"] = formula
    db[f"{colL}7"].number_format = fmt
    db[f"{colL}7"].font = font(16, True, TEAL)
    db[f"{colL}7"].alignment = Alignment(horizontal="right")
for a, b in (("C6", "D6"), ("C7", "D7"), ("E6", "F6"), ("E7", "F7"), ("G6", "H6"), ("G7", "H7")):
    db.merge_cells(f"{a}:{b}")
for r in (6, 7):
    for colL in "BCDEFGH":
        db[f"{colL}{r}"].border = box
db.row_dimensions[7].height = 34

header(db, 9, 10, ["Month", "Mood", "Sleep", "Habits"])
for m in range(12):
    r = 10 + m
    t = TABS[m]
    style(db.cell(r, 10, t), False, bold=True, align="center")
    style(db.cell(r, 11, f'=IFERROR({t}!AH{MOOD}*1,"")'), False, "0.0", align="center")
    style(db.cell(r, 12, f'=IFERROR({t}!AH{SLEEP}*1,"")'), False, "0.0", align="center")
    done_sum = f"SUM({t}!AH{G0}:AH{G0 + N - 1})"
    days = f"IF({YEAR}<YEAR(TODAY()),DAY(EOMONTH(DATE({YEAR},{m + 1},1),0)),IF({YEAR}>YEAR(TODAY()),0,IF({m + 1}<MONTH(TODAY())," \
           f"DAY(EOMONTH(DATE({YEAR},{m + 1},1),0)),IF({m + 1}=MONTH(TODAY()),DAY(TODAY()),0))))"
    style(db.cell(r, 13, f'=IF(OR({days}=0,COUNTIF(B$10:B$24,"?*")=0),"",{done_sum}/({days}*COUNTIF(B$10:B$24,"?*")))'),
          False, "0%", align="center")
mood_avg = lambda c_: [FormulaRule(formula=[f"AND({c_}<>\"\",{c_}>={v - 0.5},{c_}<{v + 0.5})"], fill=fill(colour),
                                   font=Font(name=F, bold=True, color=text)) for v, _, colour, text in MOODS]
for rule in mood_avg("K10"):
    db.conditional_formatting.add("K10:K21", rule)
db.conditional_formatting.add("M10:M21", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                        end_type="num", end_value=1, end_color="8CC5A8"))
chart = BarChart()
chart.type = "col"
chart.title = "Habits done by month"
chart.height, chart.width = 7, 12
chart.add_data(Reference(db, min_col=13, min_row=9, max_row=21), titles_from_data=True)
chart.set_categories(Reference(db, min_col=10, min_row=10, max_row=21))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.y_axis.number_format = "0%"
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 1
chart.legend = None
db.add_chart(chart, "J24")

header(db, 27, 2, ["Share of days"] + TABS)
for k in range(N):
    r = 28 + k
    style(db.cell(r, 2, f"=B{10 + k}"), False, bold=True)
    for m in range(12):
        style(db.cell(r, 3 + m, f'=IF($B{r}="","",{TABS[m]}!AI{G0 + k})'), False, "0%", align="center")
db.conditional_formatting.add("C28:N42", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                        end_type="num", end_value=1, end_color="3E8E7E"))
for c_ in range(3, 15):
    db.column_dimensions[L(c_)].width = max(db.column_dimensions[L(c_)].width or 0, 7)

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow and white day cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Habits: type the year and up to 15 habits. Keep each habit on its own line all year."),
    ("2", "Month tabs: each day, type x under the habits you did. In Google Sheets you can use checkboxes instead."),
    ("3", "On the same tab, give the day a mood from 1 (bad) to 5 (great) and the hours you slept."),
    ("4", "Dashboard: this month's habits with their streaks, and the year month by month."),
    ("5", "Year in pixels: every day of the year in the colour of your mood."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["A streak counts the days in a row, and it carries on from one month to the next.",
                          "Streak now counts today once it is ticked, and until then the days up to yesterday.",
                          "The example year shows how it works. Clear the day cells and start with your own.",
                          "Works in Google Sheets and Microsoft Excel. For a new year, make a copy and change the year."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Year in pixels"] + TABS + ["Habits", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Habit-Mood-Tracker.xlsx")
wb.save(out)
print("saved", out, wb.sheetnames, len(marks), "ticks", len(moods), "moods")
