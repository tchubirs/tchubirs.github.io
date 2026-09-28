"""Plant Care Tracker (Excel + Google Sheets).

Every plant with its spot, its light, how often it needs water in the growing season and in the resting months,
how often to feed it and when to repot it. A care log where each watering, feed or repot is one line; from it,
the next watering of every plant, what to water today and what is late, the feeds of the week, the repots coming
up, and a week ahead in a grid. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
P0, P1 = 6, 65         # plants
C0, C1 = 6, 1505       # care log
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MON_LIST = ",".join(f'"{m[:3]}"' for m in MONTHS)
ACTIONS = ["Water", "Feed", "Mist", "Repot", "Prune", "Pest check", "Other"]
LIGHTS = ["Low", "Medium", "Bright indirect", "Direct sun"]
PL = lambda col: f"Plants!${col}${P0}:${col}${P1}"
LG = lambda col: f"'Care log'!${col}${C0}:${col}${C1}"
GROWING = "Settings!$C$7"
WATER, FEED, BOTH = "DCEBFA", "D8F0DC", "CFE8E4"


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def days_text(n):
    """1 day, 3 days."""
    return f'{n}&IF({n}=1," day"," days")'


# ---------------------------------------------------------------- example data (made-up nicknames)
plants = [
    # name, kind, room, light, water every (growing), water every (resting), feed every, repot every (months),
    # potted on, last watered (days ago), last fed (days ago), note
    ("Monty", "Monstera deliciosa", "Living room", "Bright indirect", 7, 12, 30, 24, today - dt.timedelta(days=690), 4, 20, "New leaf coming"),
    ("Sansa", "Snake plant", "Bedroom", "Low", 14, 28, 60, 36, today - dt.timedelta(days=400), 5, 40, None),
    ("Goldie", "Golden pothos", "Kitchen shelf", "Medium", 7, 10, 30, 24, today - dt.timedelta(days=300), 7, 29, "Trim the long vines"),
    ("Fig", "Fiddle leaf fig", "Living room", "Bright indirect", 7, 10, 30, 18, today - dt.timedelta(days=610), 6, 12, "Turn it each week"),
    ("Zed", "ZZ plant", "Hallway", "Low", 14, 28, 60, 36, today - dt.timedelta(days=200), 10, 45, None),
    ("Lily", "Peace lily", "Bathroom", "Low", 5, 8, 30, 24, today - dt.timedelta(days=250), 3, 16, "Droops when thirsty"),
    ("Cleo", "Calathea orbifolia", "Bedroom", "Medium", 5, 7, 30, 24, today - dt.timedelta(days=150), 7, 10, "Filtered water only"),
    ("Pearl", "String of pearls", "Office window", "Bright indirect", 10, 21, 30, 24, today - dt.timedelta(days=330), 8, 25, None),
    ("Vera", "Aloe vera", "Kitchen window", "Direct sun", 14, 28, 60, 24, today - dt.timedelta(days=380), 3, 50, None),
    ("Fern", "Boston fern", "Bathroom", "Medium", 3, 5, 30, 24, today - dt.timedelta(days=180), 3, 31, "Likes the steam"),
    ("Ruby", "Rubber plant", "Office", "Bright indirect", 8, 12, 30, 24, today - dt.timedelta(days=500), 2, 5, None),
    ("Basil", "Sweet basil", "Kitchen window", "Direct sun", 2, 3, 14, None, None, 3, 15, "Pick the tips"),
    ("Spidey", "Spider plant", "Living room", "Medium", 6, 10, 30, 24, today - dt.timedelta(days=120), 1, 22, "Babies to pot up"),
    ("Echo", "Echeveria", "Office window", "Direct sun", 12, 25, 60, 36, today - dt.timedelta(days=90), 9, 55, None),
]
rng = random.Random(25)
log = []
start = today - dt.timedelta(days=110)
for name, kind, room, light, w_grow, w_rest, feed, repot, potted, last_w, last_f, note in plants:
    d = today - dt.timedelta(days=last_w)
    while d >= start:
        log.append((d, name, "Water", None))
        d -= dt.timedelta(days=max(1, w_grow + rng.choice([-1, 0, 0, 1])))
    d = today - dt.timedelta(days=last_f)
    while d >= start:
        log.append((d, name, "Feed", "Half strength" if feed <= 30 else None))
        d -= dt.timedelta(days=feed)
for d, name, action, note in [(today - dt.timedelta(days=12), "Fern", "Mist", None), (today - dt.timedelta(days=5), "Fern", "Mist", None),
                              (today - dt.timedelta(days=9), "Cleo", "Mist", None), (today - dt.timedelta(days=30), "Fig", "Pest check", "Leaves clean, no pests"),
                              (today - dt.timedelta(days=45), "Goldie", "Prune", "Cut two long vines, put them in water"),
                              (today - dt.timedelta(days=60), "Spidey", "Repot", "Into a 17 cm pot"),
                              (today - dt.timedelta(days=16), "Echo", "Repot", "Cactus soil")]:
    log.append((d, name, action, note))
log.sort(key=lambda x: (x[0], x[1]))

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The months your plants grow. Outside them they rest: less water and no feeding.",
           [3, 26, 14, 4, 70], rows=20)
for r, label, value in ((4, "Growing season from", "March"), (5, "to the end of", "September")):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, bold=True, align="center")
    se[f"C{r}"] = value
dropdown(se, '"' + ",".join(MONTHS) + '"', "C4:C5")
se["B7"] = "Today it is"
se["B7"].font = font(11, True)
m_from = "MATCH($C$4,{" + ",".join(f'"{m}"' for m in MONTHS) + "},0)"
m_to = "MATCH($C$5,{" + ",".join(f'"{m}"' for m in MONTHS) + "},0)"
now = "MONTH(TODAY())"
se["C7"] = f"=IF({m_from}<={m_to},IF(AND({now}>={m_from},{now}<={m_to}),1,0),IF(OR({now}>={m_from},{now}<={m_to}),1,0))"
se["C7"].number_format = '"growing season";;"resting season"'
se["C7"].font = font(11, True, TEAL_D)
se["C7"].alignment = Alignment(horizontal="center")
for k, text in enumerate(["North of the equator, most houseplants grow from March to September.",
                          "South of it, pick September to March: the season can run over the new year.",
                          "Watering follows the season of today, and feeds stop while plants rest."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Care log
lg = wb.create_sheet("Care log")
sheet_base(lg, "Care log", "One line each time you water, feed, mist or repot a plant. The newest line of each kind "
           "sets the next date.", [3, 12, 18, 13, 40], rows=C1 + 2)
header(lg, 5, 2, ["Date", "Plant", "What", "Notes"])
for r in range(C0, C1 + 1):
    e = log[r - C0] if r - C0 < len(log) else (None,) * 4
    style(lg.cell(r, 2, e[0]), True, DATE)
    style(lg.cell(r, 3, e[1]), True, bold=True)
    style(lg.cell(r, 4, e[2]), True)
    style(lg.cell(r, 5, e[3]), True)
    lg.cell(r, 7, f'=IF(OR(B{r}="",C{r}="",D{r}=""),"",B{r}+ROW()/10000000)')                          # G: order
    lg.cell(r, 8, f'=IF(G{r}="","",IF(COUNTIFS({LG("C")},C{r},{LG("D")},D{r},{LG("G")},">"&G{r})=0,1,0))')  # H: newest
for c in "FGH":
    lg.column_dimensions[c].hidden = True
dropdown(lg, f"={PL('B')}", f"C{C0}:C{C1}")
dropdown(lg, f'"{",".join(ACTIONS)}"', f"D{C0}:D{C1}", strict=False)
lg.conditional_formatting.add(f"D{C0}:D{C1}", FormulaRule(formula=[f'D{C0}="Water"'], fill=fill(WATER)))
lg.conditional_formatting.add(f"D{C0}:D{C1}", FormulaRule(formula=[f'D{C0}="Feed"'], fill=fill(FEED)))
lg.conditional_formatting.add(f"C{C0}:C{C1}", FormulaRule(formula=[f'AND(C{C0}<>"",COUNTIF({PL("B")},C{C0})=0)'],
                                                          fill=fill(BAD)))
lg.freeze_panes = "C6"

# ---------------------------------------------------------------- Plants
pl = wb.create_sheet("Plants")
sheet_base(pl, "Plants", "Every plant once. Water every: days between waterings, in the growing season and in the "
           "resting one. Leave the resting one empty to keep the same.",
           [3, 14, 20, 15, 15, 9, 9, 9, 10, 13, 13, 13, 14, 13, 13, 26], rows=P1 + 2)
header(pl, 5, 2, ["Plant", "Kind", "Spot", "Light", "Water every", "Resting", "Feed every", "Repot (months)",
                  "Potted on", "Last watered", "Next water", "Water", "Next feed", "Next repot", "Notes"])
newest = lambda r, action: f'SUMIFS({LG("B")},{LG("C")},$B{r},{LG("D")},"{action}",{LG("H")},1)'
for r in range(P0, P1 + 1):
    p = plants[r - P0] if r - P0 < len(plants) else (None,) * 12
    name, kind, room, light, w_grow, w_rest, feed, repot, potted, _, _, note = p
    style(pl.cell(r, 2, name), True, bold=True)
    style(pl.cell(r, 3, kind), True)
    style(pl.cell(r, 4, room), True)
    style(pl.cell(r, 5, light), True)
    style(pl.cell(r, 6, w_grow), True, "0", align="center")
    style(pl.cell(r, 7, w_rest), True, "0", align="center")
    style(pl.cell(r, 8, feed), True, "0", align="center")
    style(pl.cell(r, 9, repot), True, "0", align="center")
    style(pl.cell(r, 10, potted), True, DATE)
    style(pl.cell(r, 11, f'=IF(B{r}="","",IF({newest(r, "Water")}=0,"",{newest(r, "Water")}))'), False, DATE)
    style(pl.cell(r, 12, f'=IF(OR(K{r}="",N(S{r})=0),"",K{r}+S{r})'), False, "ddd d mmm", align="center")
    style(pl.cell(r, 13, f'=IF(B{r}="","",IF(L{r}="",IF(N(S{r})=0,"","Not logged yet"),IF(L{r}<TODAY(),"Late "&'
                         f'{days_text(f"(TODAY()-L{r})")},IF(L{r}=TODAY(),"Today","In "&{days_text(f"(L{r}-TODAY())")}))))'),
          False, align="center")
    feed_next = f"IF(T{r}=0,TODAY(),T{r}+H{r})"
    style(pl.cell(r, 14, f'=IF(OR(B{r}="",N(H{r})=0),"",IF({GROWING}=0,"Resting",{feed_next}))'), False, "ddd d mmm",
          align="center")
    style(pl.cell(r, 15, f'=IF(OR(B{r}="",N(I{r})=0,N(U{r})=0),"",EDATE(U{r},I{r}))'), False, "mmm yyyy", align="center")
    style(pl.cell(r, 16, note), True)
    pl.cell(r, 19, f'=IF(B{r}="",0,IF({GROWING}=1,N(F{r}),IF(G{r}="",N(F{r}),N(G{r}))))')          # S: days between now
    pl.cell(r, 20, f'=IF(B{r}="",0,{newest(r, "Feed")})')                                            # T: last fed
    pl.cell(r, 21, f'=IF(B{r}="",0,MAX(N(J{r}),{newest(r, "Repot")}))')                             # U: last potted
    pl.cell(r, 22, f'=IF(L{r}="","",L{r}+ROW()/10000000)')                                          # V: water order
    pl.cell(r, 23, f'=IF(ISNUMBER(N{r}),N{r}+ROW()/10000000,"")')                                  # W: feed order
    pl.cell(r, 24, f'=IF(O{r}="","",O{r}+ROW()/10000000)')                                          # X: repot order
for c in range(18, 25):
    pl.column_dimensions[L(c)].hidden = True
dropdown(pl, f'"{",".join(LIGHTS)}"', f"E{P0}:E{P1}", strict=False)
pl.conditional_formatting.add(f"M{P0}:M{P1}", FormulaRule(formula=[f'LEFT(M{P0},4)="Late"'], fill=fill(BAD)))
pl.conditional_formatting.add(f"M{P0}:M{P1}", FormulaRule(formula=[f'M{P0}="Today"'], fill=fill(WATER)))
pl.conditional_formatting.add(f"N{P0}:N{P1}", FormulaRule(formula=[f'AND(ISNUMBER(N{P0}),N{P0}<=TODAY())'], fill=fill(FEED)))
pl.conditional_formatting.add(f"O{P0}:O{P1}", FormulaRule(formula=[f'AND(ISNUMBER(O{P0}),O{P0}<=TODAY()+30)'], fill=fill(WARN)))
pl.freeze_panes = "C6"

# ---------------------------------------------------------------- This week
wk = wb.create_sheet("This week")
sheet_base(wk, "This week", "What each plant needs over the next seven days, from what the care log says.",
           [3, 16, 15] + [13] * 7, rows=P1 + 2)
header(wk, 5, 2, ["Plant", "Spot"] + [""] * 7)
for k in range(7):
    wk.cell(5, 4 + k, f'=CHOOSE(WEEKDAY(TODAY()+{k},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")&" "&{short(f"(TODAY()+{k})")}')
for r in range(P0, P1 + 1):
    style(wk.cell(r, 2, f'=IF(Plants!B{r}="","",Plants!B{r})'), False, bold=True)
    style(wk.cell(r, 3, f'=IF(Plants!B{r}="","",Plants!D{r})'), False)
    first_w = f"MAX(Plants!L{r},TODAY())"
    first_f = f"MAX(Plants!N{r},TODAY())"
    for k in range(7):
        d = f"(TODAY()+{k})"
        # MOD by at least 1: AND works out every part, and a 0 would turn the whole cell into an error.
        water = f'AND(ISNUMBER(Plants!L{r}),N(Plants!S{r})>0,{d}>={first_w},MOD({d}-{first_w},MAX(1,N(Plants!S{r})))=0)'
        feed = f'AND(ISNUMBER(Plants!N{r}),N(Plants!H{r})>0,{d}>={first_f},MOD({d}-{first_f},MAX(1,N(Plants!H{r})))=0)'
        style(wk.cell(r, 4 + k, f'=IF({water},IF({feed},"Water, feed","Water"),IF({feed},"Feed",""))'), False,
              align="center")
wk.conditional_formatting.add(f"D{P0}:J{P1}", FormulaRule(formula=[f'D{P0}="Water"'], fill=fill(WATER)))
wk.conditional_formatting.add(f"D{P0}:J{P1}", FormulaRule(formula=[f'D{P0}="Feed"'], fill=fill(FEED)))
wk.conditional_formatting.add(f"D{P0}:J{P1}", FormulaRule(formula=[f'D{P0}="Water, feed"'], fill=fill(BOTH)))
wk.freeze_panes = "D6"

# ---------------------------------------------------------------- Today (dashboard)
db = wb.create_sheet("Today", 0)
sheet_base(db, "Plant care", "What your plants need today and in the days ahead.",
           [3, 16, 16, 16, 3, 16, 14, 3, 16, 14, 3], rows=45)
db["B4"] = f'="Today, "&{short("TODAY()")}&": "&IF({GROWING}=1,"growing season","resting season")'
db["B4"].font = font(11, True, TEAL_D)
LIST = 12
SOON = {"water": "TODAY()+7", "feed": "TODAY()+7", "repot": "EDATE(TODAY(),2)"}
tile(db, "B", 6, "Plants", f'=COUNTIF({PL("B")},"?*")', "0")
tile(db, "C", 6, "Water today", f'=COUNTIFS({PL("L")},"<="&TODAY())', "0", "1F5F99")
tile(db, "D", 6, "Late", f'=COUNTIFS({PL("L")},"<"&TODAY())', "0", "B4541F")
tile(db, "F", 6, "Feed this week", f'=COUNTIFS({PL("W")},"<"&{SOON["feed"]})', "0", "2E7D32")
tile(db, "I", 6, "Repot soon", f'=COUNTIFS({PL("X")},"<"&{SOON["repot"]})', "0", "B07A00")


def listing(first, head, keys, soon, cells, fmts, empty, hide):
    """The plants whose date is before `soon`, soonest first, in LIST rows. Keys go in the hidden column `hide`."""
    header(db, 9, first, head)
    for i in range(LIST):
        r = 10 + i
        key = f"{hide}{r}"
        db[key] = f'=IFERROR(IF(SMALL({keys},{i + 1})<{soon},SMALL({keys},{i + 1}),""),"")'
        row = f"MATCH({key},{keys},0)"
        for k, (cell, fmt) in enumerate(zip(cells, fmts)):
            blank = f'"{empty}"' if i == 0 and k == 0 else '""'
            style(db.cell(r, first + k, f'=IF({key}="",{blank},{cell(key, row)})'), False, fmt, k == 0,
                  "left" if k == 0 else "center")
    db.column_dimensions[hide].hidden = True


day_name = lambda d: f'CHOOSE(WEEKDAY({d},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")&" "&{short(d)}'
listing(2, ["Water", "Spot", "When"], PL("V"), SOON["water"],
        [lambda key, row: f"INDEX({PL('B')},{row})", lambda key, row: f"INDEX({PL('D')},{row})",
         lambda key, row: f'IF(INT({key})<TODAY(),"Late "&{days_text(f"(TODAY()-INT({key}))")},'
                          f'IF(INT({key})=TODAY(),"Today",{day_name(f"INT({key})")}))'],
        [None, None, None], "Nothing to water this week", "T")
listing(6, ["Feed", "When"], PL("W"), SOON["feed"],
        [lambda key, row: f"INDEX({PL('B')},{row})",
         lambda key, row: f'IF(INT({key})<=TODAY(),"Today",{day_name(f"INT({key})")})'],
        [None, None], "No feeds this week", "U")
listing(9, ["Repot", "When"], PL("X"), SOON["repot"],
        [lambda key, row: f"INDEX({PL('B')},{row})",
         lambda key, row: f'IF(INT({key})<=TODAY(),"Now",CHOOSE(MONTH(INT({key})),{MON_LIST})&" "&YEAR(INT({key})))'],
        [None, None], "No repots in the next two months", "V")
db.conditional_formatting.add(f"D10:D{9 + LIST}", FormulaRule(formula=['LEFT(D10,4)="Late"'], fill=fill(BAD)))
db.conditional_formatting.add(f"D10:D{9 + LIST}", FormulaRule(formula=['D10="Today"'], fill=fill(WATER)))
db.conditional_formatting.add(f"G10:G{9 + LIST}", FormulaRule(formula=['G10="Today"'], fill=fill(FEED)))
db.conditional_formatting.add(f"J10:J{9 + LIST}", FormulaRule(formula=['J10="Now"'], fill=fill(WARN)))
db[f"B{11 + LIST}"] = ("Water: the plants due in the next seven days, the late ones first. Feeds stop while plants "
                       "rest. Repot: the next two months.")
db[f"B{11 + LIST}"].font = font(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: the months your plants grow. North of the equator that is often March to September."),
    ("2", "Plants: each plant once, with how many days between waterings in the growing and resting seasons."),
    ("3", "Care log: one line each time you water, feed, mist or repot. Pick the plant from the list."),
    ("4", "Today: what to water now, what is late, the feeds of the week and the repots coming up."),
    ("5", "This week: the next seven days in a grid, for each plant."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The days between waterings are a starting point. Check the soil and change them to suit "
                          "your home.",
                          "The example plants and their nicknames are made up. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Today", "This week", "Plants", "Care log", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Plant-Care-Tracker.xlsx")
wb.save(out)
print("saved", out, len(plants), "plants", len(log), "log lines")
