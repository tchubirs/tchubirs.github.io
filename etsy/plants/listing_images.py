"""Etsy listing photos for the Plant Care Tracker (run build.py first)."""
import os, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Plant-Care-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
r = Renders(xlsx, tmp)
# Care log close up: its last 16 lines.
log = load_workbook(xlsx)["Care log"]
last = max(row for row in range(6, 1506) if log.cell(row, 2).value)
close = Renders(xlsx, tmp, hide={"Plants": ["J", "P"]},
                hide_rows={"This week": range(20, 66), "Plants": range(20, 66),
                           "Care log": list(range(6, last - 15)) + list(range(last + 1, 1508))})
img_dir = os.path.join(HERE, "images")

img = canvas(dark=True)
title(img, "Plant care tracker", "What to water today, what is late, the feeds of the week and the repots coming up",
      dark=True, width=2300)
laptop(img, r.shot("Plant care"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-today.jpg"))

detail(os.path.join(img_dir, "02-this-week.jpg"), "Your week of plant care",
       "The next seven days for every plant, worked out from the last time you watered and fed it",
       close.shot("This week"), [("#DCEBFA", "Water"), ("#D8F0DC", "Feed"), ("#CFE8E4", "Both")])

detail(os.path.join(img_dir, "03-plants.jpg"), "Each plant, its own rhythm",
       "Days between waterings in the growing and the resting season, feeding and repotting",
       close.shot("Plants"), [("#DCEBFA", "Water today"), ("#F9C9C4", "Late"), ("#D8F0DC", "Feed now"),
                              ("#FFE2B8", "Repot soon")])

detail(os.path.join(img_dir, "04-care-log.jpg"), "One line per watering",
       "Water, feed, mist, repot or prune: the newest line sets the next date",
       close.shot("Care log"), [("#DCEBFA", "Water"), ("#D8F0DC", "Feed")])

img = canvas()
y = title(img, "Inside the file", "Four tabs, settings and a Start Here page. You type in the yellow cells and the rest "
          "is calculated.")
grid(img, [
    (content(r.page("Plant care")), "Today", "What to water now"),
    (content(r.page("This week")), "This week", "Seven days ahead"),
    (content(r.page("Plants")), "Plants", "Each plant's rhythm"),
    (content(r.page("Care log")), "Care log", "One line per care"),
    (content(r.page("Settings")), "Settings", "The growing season"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "05-whats-inside.jpg"))
