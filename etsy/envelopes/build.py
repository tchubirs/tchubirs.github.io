"""Cash Envelope Budget, the cash stuffing way (Excel + Google Sheets).

Envelopes for spending and for sinking funds, what goes into each one every payday, the cash spent from them and
moved between them, what is left in each (it carries over to the next month), the notes to ask for at the bank, and
cards to print for the envelopes. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, bar, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
E0, E1 = 6, 21         # envelopes
P0, P1 = 6, 205        # paydays
S0, S1 = 6, 1005       # spending
NE = E1 - E0 + 1
CUR, PAID, SHOWN, LOW, PER_YEAR = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6", "Settings!$C$7", "Settings!$H$5"
N0, N1 = 11, 18        # note values on Settings
NOTES = f"Settings!$C${N0}:$C${N1}"
ENV = f"Envelopes!$B${E0}:$B${E1}"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MON = [m[:3] for m in MONTHS]
PC = lambda k: get_column_letter(5 + k)                        # the Paydays column of envelope k
PCOL = lambda k: f"Paydays!${PC(k)}${P0}:${PC(k)}${P1}"
PW = f"Paydays!$W${P0}:$W${P1}"
SP = lambda c: f"Spending!${c}${S0}:${c}${S1}"
EN = lambda c, k: f"Envelopes!${c}${E0 + k}"


def short_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{",".join(chr(34) + m + chr(34) for m in MON)})'


def flows(k, cond=None):
    """Stuffed, spent, and moved in minus out for envelope k, each SUMIFS with an extra date condition."""
    name = EN("B", k)
    cp = cond(PW) if cond else ""
    cs = cond(SP("H")) if cond else ""
    stuffed = f'SUMIFS({PCOL(k)},{PW},">0"{cp})'
    spent = f'SUMIFS({SP("D")},{SP("C")},{name},{SP("I")},1{cs})'
    moved = (f'SUMIFS({SP("D")},{SP("F")},{name},{SP("J")},1{cs})'
             f'-SUMIFS({SP("D")},{SP("C")},{name},{SP("J")},1{cs})')
    return stuffed, spent, moved


# ---------------------------------------------------------------- example data (a made-up household)
rng = random.Random(42)
envelopes = [  # name, kind, plan per payday, largest note, target, by, cash at the start
    ("Groceries", "Spending", 180, 20, None, None, 40),
    ("Gas", "Spending", 60, 20, None, None, 15),
    ("Eating out", "Spending", 40, 10, None, None, 0),
    ("Fun money", "Spending", 30, 5, None, None, 10),
    ("Personal care", "Spending", 25, 5, None, None, 0),
    ("Household", "Spending", 30, 10, None, None, 5),
    ("Kids", "Spending", 35, 10, None, None, 0),
    ("Christmas", "Sinking fund", 25, 20, 600, dt.date(today.year, 12, 15), 350),
    ("Car repairs", "Sinking fund", 20, 20, 500, None, 120),
    ("Vacation", "Sinking fund", 40, 50, 1200, dt.date(today.year + 1, 7, 1), 0),
    ("Birthdays", "Sinking fund", 15, 10, 300, None, 40),
    ("Emergency fund", "Sinking fund", 50, 100, 1500, None, 200),
]
names = [e[0] for e in envelopes]
last_pay = today - dt.timedelta(days=10)
paydays = [last_pay - dt.timedelta(days=14 * k) for k in range(21, -1, -1)]
stuff = []     # date, paycheck, amounts
for n, d in enumerate(paydays):
    amounts = [e[2] for e in envelopes]
    pay = 1480
    if n == 9:                                                       # a bonus: more for the vacation and savings
        pay, amounts[9], amounts[11] = 1780, 140, 150
    if n == 14:                                                      # a tight payday
        amounts[2], amounts[3] = 20, 15
    if n == len(paydays) - 1:                                        # the last one needs small notes and coins
        amounts[1], amounts[3] = 62.5, 33
    stuff.append([d, pay, amounts])

bal = {e[0]: e[6] for e in envelopes}
spend = []     # date, envelope, amount, what, moved to
what = {"Groceries": ["Supermarket", "Farmers market", "Bakery", "Corner shop"], "Gas": ["Gas station"],
        "Eating out": ["Pizza night", "Lunch out", "Coffee and cake", "Tacos"],
        "Fun money": ["Cinema", "Book", "Bowling", "Game"], "Personal care": ["Haircut", "Pharmacy", "Shampoo"],
        "Household": ["Cleaning things", "Light bulbs", "Hardware store"],
        "Kids": ["School trip", "Shoes", "Art supplies", "Swimming"]}
for n, (d, pay, amounts) in enumerate(stuff):
    for name, a in zip(names, amounts):
        bal[name] = round(bal[name] + a, 2)
    final = n + 1 == len(stuff)
    end = today if final else stuff[n + 1][0]
    days = (end - d).days
    for name in names[:7]:
        plan = amounts[names.index(name)]
        if final:                                                    # halfway to the next payday
            keep = 30 if name == "Groceries" else round(plan * rng.uniform(0.35, 0.6), 2)
        else:
            keep = round(plan * rng.uniform(*((0.1, 0.35) if name in ("Eating out", "Fun money") else (0, 0.15))),
                         2)                                          # a little is left, and it carries over
        total = round(bal[name] - keep, 2)
        lines = {"Groceries": 4, "Gas": 2, "Eating out": 2, "Fun money": 1}.get(name, 1)
        if final:
            lines = max(1, lines - 1)
        cuts = sorted(rng.uniform(0.2, 0.8) for _ in range(lines - 1))
        parts = [round(total * (b2 - a2), 2) for a2, b2 in zip([0] + cuts, cuts + [1])]
        parts[-1] = round(total - sum(parts[:-1]), 2)
        for i, amt in enumerate(parts):
            if amt <= 0:
                continue
            day = d + dt.timedelta(days=1 + int((i + rng.random()) * (days - 1) / lines))
            spend.append([day, name, amt, rng.choice(what[name]), None])
            bal[name] = round(bal[name] - amt, 2)
    if n % 2 == 1 and not final:                                     # leftovers go to the emergency fund
        for name in ("Eating out", "Fun money"):
            if bal[name] >= 5:
                amt = round(bal[name], 2)
                spend.append([end - dt.timedelta(days=1), name, amt, "Leftover to savings", "Emergency fund"])
                bal[name] = 0
                bal["Emergency fund"] = round(bal["Emergency fund"] + amt, 2)
sinking = [(dt.date(today.year - 1, 12, 6), "Christmas", 180, "Presents"),
           (dt.date(today.year - 1, 12, 13), "Christmas", 85, "Christmas dinner"),
           (dt.date(today.year - 1, 12, 20), "Christmas", 55, "Tree and lights"),
           (dt.date(today.year, 3, 14), "Birthdays", 60, "Birthday present"),
           (dt.date(today.year, 5, 20), "Emergency fund", 150, "Vet visit"),
           (dt.date(today.year, 6, 10), "Car repairs", 380, "New tires"),
           (dt.date(today.year, 8, 2), "Vacation", 250, "Cabin deposit"),
           (dt.date(today.year, 9, 12), "Birthdays", 45, "Party supplies")]
for d, name, amt, note in sinking:
    spend.append([d, name, amt, note, None])
spend.append([today - dt.timedelta(days=3), "Fun money", 10, "Top up for groceries", "Groceries"])
spend.sort(key=lambda s: s[0])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "How often you are paid, the month shown, and the notes your bank gives out.",
           [3, 30, 16, 4, 58], rows=24)
for r, lab, value, fmt, note in (
        (4, "Currency", "USD", None, "Shown next to the cash to take out."),
        (5, "Paid", "Every 2 weeks", None, "Every week, every 2 weeks, twice a month or every month."),
        (6, "Month shown", "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "mmmm yyyy",
         "The month on the Dashboard. Type a date in another month to look back."),
        (7, "Low when less than", 0.25, "0%", "An envelope is low when less than this share of its plan is left.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, False, "center")
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
se["H5"] = f'=IF(C5="Every week",52,IF(C5="Twice a month",24,IF(C5="Every month",12,26)))'   # H: paydays a year
se.column_dimensions["H"].hidden = True
dropdown(se, '"Every week,Every 2 weeks,Twice a month,Every month"', "C5")
se["B9"] = "Notes you use"
se["B9"].font = font(12, True, TEAL_D)
se["B10"] = "Any order. Leave out the ones you do not want in your envelopes."
se["B10"].font = font(10, color=MUTED, italic=True)
for k, v in enumerate([100, 50, 20, 10, 5, 1, None, None]):
    r = N0 + k
    se[f"B{r}"] = f"Note {k + 1}"
    se[f"B{r}"].font = font(11)
    style(se[f"C{r}"], True, "#,##0.##", False, "center")
    se[f"C{r}"] = v
se["E11"] = "For euros: 200, 100, 50, 20, 10 and 5, and the coins 2 and 1 if you like."
se["E12"] = "For pounds: 50, 20, 10 and 5. Anything smaller is counted as coins."
for c in ("E11", "E12"):
    se[c].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Envelopes
en = wb.create_sheet("Envelopes")
sheet_base(en, "Envelopes", "Each envelope once, with what goes in it each payday. A sinking fund saves up for "
           "something that comes later.", [3, 20, 14, 12, 10, 12, 13, 12, 12, 12, 11, 12, 13, 12], rows=E1 + 4)
header(en, 5, 2, ["Envelope", "Kind", "Plan per payday", "Largest note", "Target", "By", "Cash at the start",
                  "Stuffed", "Spent", "Moved", "In it now", "Status", "Per payday to get there"])
for k in range(NE):
    r = E0 + k
    x = envelopes[k] if k < len(envelopes) else (None,) * 7
    for c, v, fmt in zip(range(2, 9), x, (None, None, MONEY, "#,##0.##", MONEY, DATE, MONEY)):
        style(en.cell(r, c, v), True, fmt, c == 2, "center" if c in (3, 5, 7) else None)
    stuffed, spent, moved = flows(k)
    style(en.cell(r, 9, f'=IF(B{r}="","",{stuffed})'), False, MONEY)
    style(en.cell(r, 10, f'=IF(B{r}="","",{spent})'), False, MONEY)
    style(en.cell(r, 11, f'=IF(B{r}="","",{moved})'), False, MONEY)
    style(en.cell(r, 12, f'=IF(B{r}="","",N(H{r})+I{r}-J{r}+K{r})'), False, MONEY, True)
    style(en.cell(r, 13, f'=IF(B{r}="","",IF(L{r}<0,"Overspent",IF(C{r}="Sinking fund",IF(AND(N(F{r})>0,'
                         f'L{r}>=N(F{r})),"Reached",IF(AND(ISNUMBER(G{r}),G{r}<TODAY()),"Past its date","Saving")),'
                         f'IF(AND(N(D{r})>0,L{r}<N(D{r})*N({LOW})),"Low","OK"))))'), False, bold=True,
          align="center")
    style(en.cell(r, 14, f'=IF(OR(B{r}="",C{r}<>"Sinking fund",N(F{r})<=0),"",IF(L{r}>=N(F{r}),0,'
                         f'IF(OR(NOT(ISNUMBER(G{r})),G{r}<TODAY()),"",'
                         f'(N(F{r})-L{r})/MAX(1,INT((G{r}-TODAY())*{PER_YEAR}/365)))))'),
          False, MONEY)
    en[f"P{r}"] = f'=IF(AND(B{r}<>"",C{r}="Sinking fund"),ROW(),"")'                          # P: sinking funds
    on = lambda col: f',{col},"<="&INT(N(Cash!$C$4))'
    s2, p2, m2 = flows(k, on)
    en[f"Q{r}"] = f'=IF(B{r}="","",N(H{r})+{s2}-{p2}+{m2})'                                    # Q: on the payday
for c in "PQ":
    en.column_dimensions[c].hidden = True
for text, colour in (("Low", WARN), ("Overspent", BAD), ("Reached", OK)):
    en.conditional_formatting.add(f"M{E0}:M{E1}", FormulaRule(formula=[f'M{E0}="{text}"'], fill=fill(colour)))
dropdown(en, '"Spending,Sinking fund"', f"C{E0}:C{E1}")
dropdown(en, f"={NOTES}", f"E{E0}:E{E1}", strict=False)
en.freeze_panes = "C6"

# ---------------------------------------------------------------- Paydays
pd_ = wb.create_sheet("Paydays")
last_env = get_column_letter(4 + NE)
sheet_base(pd_, "Paydays", "Each payday, what goes into each envelope. The plan row above the names shows your usual "
           "amounts.", [3, 13, 12, 12] + [11] * NE + [13, 24], rows=P1 + 2)
header(pd_, 5, 2, ["Payday", "Paycheck", "Stuffed"] + [""] * NE + ["Not in envelopes", "Check"])
pd_["D4"] = "Plan"
pd_["D4"].font = font(10, True, MUTED)
pd_["D4"].alignment = Alignment(horizontal="right")
for k in range(NE):
    col = PC(k)
    pd_[f"{col}5"] = f'=IF({EN("B", k)}="","",{EN("B", k)})'
    pd_[f"{col}4"] = f'=IF({EN("B", k)}="","",N({EN("D", k)}))'
    pd_[f"{col}4"].number_format = MONEY
    pd_[f"{col}4"].font = font(10, color=MUTED, italic=True)
    pd_[f"{col}4"].alignment = Alignment(horizontal="right")
for r in range(P0, P1 + 1):
    x = stuff[r - P0] if r - P0 < len(stuff) else (None, None, [])
    style(pd_.cell(r, 2, x[0]), True, DATE, align="center")
    style(pd_.cell(r, 3, x[1]), True, MONEY)
    row = f"E{r}:{last_env}{r}"
    typed = f'SUMPRODUCT(({row}<>"")*1)'                         # cells with something in them
    words = f'SUMPRODUCT(({row}<>"")*(1-ISNUMBER({row})))'       # of those, the ones that are not numbers
    style(pd_.cell(r, 4, f'=IF(AND(B{r}="",{typed}=0),"",SUMPRODUCT((E$5:{last_env}$5<>"")*1,{row}))'),
          False, MONEY, True)
    for k in range(NE):
        v = x[2][k] if k < len(x[2]) else None
        style(pd_.cell(r, 5 + k, v), True, MONEY)
    style(pd_.cell(r, 5 + NE, f'=IF(OR(NOT(ISNUMBER(C{r})),D{r}=""),"",C{r}-D{r})'), False, MONEY)
    style(pd_.cell(r, 6 + NE, f'=IF(AND(B{r}="",C{r}="",{typed}=0),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                              f'IF({words}>0,"Amount not a number",'
                              f'IF(SUMPRODUCT((E$5:{last_env}$5="")*({row}<>""))>0,"Amount under no envelope",'
                              f'IF(COUNTIF({row},"<0")>0,"Amount below zero",'
                              f'IF(AND(ISNUMBER(C{r}),N(D{r})>C{r}),"More than the paycheck",""))))))'),
          False, align="center")
    pd_[f"W{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                         # W: date as a number
pd_.column_dimensions["W"].hidden = True
pd_.conditional_formatting.add(f"V{P0}:V{P1}", FormulaRule(formula=[f'V{P0}<>""'], fill=fill(BAD)))
pd_.freeze_panes = "E6"

# ---------------------------------------------------------------- Spending
sp = wb.create_sheet("Spending")
sheet_base(sp, "Spending", "Each time you spend cash from an envelope. To move cash to another envelope, pick it "
           "under Moved to.", [3, 13, 18, 12, 26, 18, 30], rows=S1 + 2)
header(sp, 5, 2, ["Date", "Envelope", "Amount", "What", "Moved to", "Check"])
for r in range(S0, S1 + 1):
    x = spend[r - S0] if r - S0 < len(spend) else (None,) * 5
    for c, v, fmt in zip(range(2, 7), x, (DATE, None, MONEY, None, None)):
        style(sp.cell(r, c, v), True, fmt, c == 3, "center" if c == 2 else None)
    sp.cell(r, 5).alignment = Alignment(horizontal="left", indent=1, vertical="center")
    style(sp.cell(r, 7, f'=IF(AND(C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                        f'IF(NOT(ISNUMBER(D{r})),"Needs the amount",IF(C{r}="","Needs an envelope",'
                        f'IF(COUNTIF({ENV},C{r})=0,"Not on Envelopes",IF(F{r}="",IF(D{r}=0,"Needs the amount",""),'
                        f'IF(COUNTIF({ENV},F{r})=0,"Moved to: not on Envelopes",IF(D{r}<=0,"A move needs an amount",'
                        f'IF(F{r}=C{r},"Moved to the same envelope","")))))))))'), False, align="center")
    sp[f"H{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                           # H: date as a number
    ok = f"H{r}>0,ISNUMBER(D{r}),COUNTIF({ENV},C{r})>0"
    sp[f"I{r}"] = f'=IF(AND({ok},F{r}=""),1,0)'                                                # I: spent
    sp[f"J{r}"] = f'=IF(AND({ok},F{r}<>"",COUNTIF({ENV},F{r})>0),1,0)'                         # J: moved
    sp[f"K{r}"] = f'=IF(I{r}=1,H{r}*10000+ROW(),"")'                                           # K: latest
for c in "HIJK":
    sp.column_dimensions[c].hidden = True
sp.conditional_formatting.add(f"G{S0}:G{S1}", FormulaRule(formula=[f'G{S0}<>""'], fill=fill(BAD)))
dropdown(sp, f"={ENV}", f"C{S0}:C{S1}", strict=False)
dropdown(sp, f"={ENV}", f"F{S0}:F{S1}", strict=False)
sp.freeze_panes = "C6"

# ---------------------------------------------------------------- Cash to take out (to print)
ca = wb.create_sheet("Cash")
sheet_base(ca, "Cash to take out", "The notes to ask for at the bank, and how to split them between the envelopes.",
           [3, 20, 12, 12] + [8] * 8 + [10], rows=32)
ca["B4"], ca["B5"] = "Payday", "Take out"
for c in ("B4", "B5"):
    ca[c].font = font(11, True)
style(ca["C4"], True, DATE, True, "center")
ca["C4"] = f"=MAX(Paydays!$B${P0}:$B${P1})"
ca["C5"] = "=SUM(C9:C24)"
ca["C5"].number_format = MONEY
ca["C5"].font = font(14, True, TEAL_D)
ca["D5"] = f'=IF({CUR}="","",{CUR})'
ca["D5"].font = font(11, True, MUTED)
ca["N4"] = f'=IF(N(C4)<=0,0,IFERROR(MATCH(INT(N(C4)),{PW},0),0))'                                            # N4: payday row
ca.column_dimensions["N"].hidden = True
ca["E4"] = '=IF(N4=0,"No payday on this date. Type one from the Paydays tab.","")'
ca["E4"].font = font(10, True, "B23B2E")
for j in range(8):
    col = get_column_letter(5 + j)
    ca[f"{col}7"] = f"=IFERROR(LARGE({NOTES},{j + 1}),0)"                                     # row 7: the notes
header(ca, 8, 2, ["Envelope", "Amount", "Largest note"] + [""] * 8 + ["Coins"])
for j in range(8):
    col = get_column_letter(5 + j)
    ca[f"{col}8"] = f'=IF({col}7<=0,"",{col}7)'
    ca[f"{col}8"].number_format = "#,##0.##"
for k in range(NE):
    r = 9 + k
    cap = f"N({EN('E', k)})"
    style(ca.cell(r, 2, f'=IF({EN("B", k)}="","",{EN("B", k)})'), False, bold=True)
    style(ca.cell(r, 3, f'=IF(OR(B{r}="",$N$4=0),"",N(INDEX(Paydays!$E${P0}:${last_env}${P1},$N$4,{k + 1})))'),
          False, '#,##0.00;-#,##0.00;""')
    style(ca.cell(r, 4, f'=IF(B{r}="","",IF({cap}>0,{cap},"Any"))'), False, "#,##0.##", align="center")
    for j in range(8):
        col = get_column_letter(5 + j)
        prev = "0" if j == 0 else f"SUMPRODUCT($E{r}:{get_column_letter(4 + j)}{r},$E$7:{get_column_letter(4 + j)}$7)"
        style(ca.cell(r, 5 + j, f'=IF(OR(N(C{r})<=0,{col}$7<=0,AND({cap}>0,{col}$7>{cap})),0,'
                                f'INT(ROUND(ROUND(C{r}-{prev},2)/{col}$7,6)))'), False, '0;-0;""', align="center")
    style(ca.cell(r, 13, f'=IF(OR(C{r}="",N(C{r})<=0),"",ROUND(C{r}-SUMPRODUCT(E{r}:L{r},$E$7:$L$7),2))'), False,
          '#,##0.00;-#,##0.00;""', align="center")
ca.row_dimensions[7].hidden = True
ca["B25"] = "Notes to ask for"
ca["B25"].font = font(11, True, TEAL_D)
for j in range(9):
    col = get_column_letter(5 + j)
    style(ca[f"{col}25"], False, '0;-0;""' if j < 8 else '#,##0.00;-#,##0.00;""', True, "center")
    ca[f"{col}25"] = f"=SUM({col}9:{col}24)"
    ca[f"{col}25"].font = font(12, True, TEAL_D)
for j in range(8):                                                                             # row 28: the words
    col = get_column_letter(5 + j)
    ca[f"{col}28"] = f'=IF({col}25>0,", "&{col}25&" of "&{col}7,"")'
ca["M28"] = '=IF(M25>0,IF(SUM(E25:L25)>0,", and ",", ")&FIXED(M25,2)&" in coins","")'
ca.row_dimensions[28].hidden = True
ca["B27"] = '=IF(C5<=0,"","Ask the bank for "&MID(E28&F28&G28&H28&I28&J28&K28&L28&M28,3,400)&".")'
ca["B27"].font = font(12, True, TEAL_D)
ca.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Cards (to print)
cd = wb.create_sheet("Cards")
sheet_base(cd, "Envelope cards", "Print this tab, cut out the cards and keep one in each envelope.",
           [3, 10, 17, 10, 10, 3, 10, 17, 10, 10], rows=56)
cd["B4"] = "Cards"
cd["B4"].font = font(11, True)
style(cd["C4"], True, None, True, "center")
cd["C4"] = "1 to 8"
dropdown(cd, '"1 to 8,9 to 16"', "C4")
cd["D4"] = '=IF(N(Cash!$C$4)=0,"","For the payday of "&' + short_date("Cash!$C$4") + ')'
cd["D4"].font = font(10, color=MUTED, italic=True)
for n in range(8):
    top = 6 + (n // 2) * 12
    c0 = 2 if n % 2 == 0 else 7
    cols = [get_column_letter(c0 + i) for i in range(4)]
    k = f'({n + 1}+IF($C$4="9 to 16",8,0))'
    name = f"INDEX({ENV},{k})"
    held = f"INDEX(Envelopes!$Q${E0}:$Q${E1},{k})"
    cd[f"{cols[0]}{top}"] = f'=IF({name}="","",{name})'
    cd[f"{cols[0]}{top}"].font = font(13, True, TEAL_D)
    cd[f"{cols[0]}{top + 1}"] = (f'=IF({name}="","",IF(N(Cash!$C$4)=0,"","Cash on '
                                 f'"&{short_date("Cash!$C$4")}&": "&FIXED(N({held}),2)))')
    cd[f"{cols[0]}{top + 1}"].font = font(10, color=MUTED, italic=True)
    for i, lab in enumerate(["Date", "What", "Amount", "Left"]):
        c = cd[f"{cols[i]}{top + 2}"]
        c.value = lab
        c.font = font(9, True, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="center")
        c.border = box
    for line in range(8):
        for col in cols:
            cd[f"{col}{top + 3 + line}"].border = box
            cd.row_dimensions[top + 3 + line].height = 20
cd.page_setup.orientation = "portrait"
cd.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Cash envelopes", "", [3, 18, 13, 13, 13, 13, 13, 14, 3, 18, 11, 15, 13, 11, 12, 3], rows=48)
start, end = "Dashboard!$V$1", "Dashboard!$V$2"
db["V1"] = f"=DATE(YEAR({SHOWN}),MONTH({SHOWN}),1)"                                          # V1, V2: the month
db["V2"] = f"=DATE(YEAR({SHOWN}),MONTH({SHOWN})+1,1)-1"
db["B3"] = (f'=CHOOSE(MONTH(V1),{",".join(chr(34) + m + chr(34) for m in MONTHS)})&" "&YEAR(V1)'
            f'&": what went into each envelope, what came out, and what is left."')
db["B3"].font = font(11, color=MUTED, italic=True)
tile(db, "B", 4, "In your envelopes", f"=SUM(Envelopes!$L${E0}:$L${E1})", MONEY)
tile(db, "C", 4, "Stuffed", "=SUM(D8:D23)", MONEY)
tile(db, "D", 4, "Spent", "=SUM(E8:E23)", MONEY)
tile(db, "E", 4, "Left to spend", '=SUMIFS(G8:G23,Q8:Q23,"Spending")', MONEY)
tile(db, "F", 4, "In sinking funds", '=SUMIFS(G8:G23,Q8:Q23,"Sinking fund")', MONEY)
tile(db, "G", 4, "Low now", f'=COUNTIF(Envelopes!$M${E0}:$M${E1},"Low")', "0", "B86E1C")
tile(db, "J", 4, "Plan per payday", f"=SUM(Envelopes!$D${E0}:$D${E1})", MONEY)
tile(db, "K", 4, "Last payday", f'=IF(COUNT(Paydays!$B${P0}:$B${P1})=0,"",MAX(Paydays!$B${P0}:$B${P1}))', "d mmm")
header(db, 7, 2, ["Envelope", "Carried in", "Stuffed", "Spent", "Moved", "Left", "Used"])
before = lambda col: f',{col},"<"&{start}'
within = lambda col: f',{col},">="&{start},{col},"<="&{end}'
for k in range(NE):
    r = 8 + k
    name = EN("B", k)
    s0, p0, m0 = flows(k, before)
    s1, p1, m1 = flows(k, within)
    style(db.cell(r, 2, f'=IF({name}="","",{name})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",N({EN("H", k)})+{s0}-{p0}+{m0})'), False, MONEY)
    style(db.cell(r, 4, f'=IF(B{r}="","",{s1})'), False, MONEY)
    style(db.cell(r, 5, f'=IF(B{r}="","",{p1})'), False, MONEY)
    style(db.cell(r, 6, f'=IF(B{r}="","",{m1})'), False, '#,##0.00;-#,##0.00;""')
    style(db.cell(r, 7, f'=IF(B{r}="","",C{r}+D{r}-E{r}+F{r})'), False, MONEY, True)
    used = f"E{r}/(C{r}+D{r}+MAX(0,F{r}))"
    style(db.cell(r, 8, f'=IF(B{r}="","",IF(C{r}+D{r}+MAX(0,F{r})<=0,"",{bar(used, 10)}))'), False).font = \
        Font(size=9, color=TEAL)
    db[f"Q{r}"] = f'=IF(B{r}="","",IF({EN("C", k)}="Sinking fund","Sinking fund","Spending"))'  # Q: kind
    db[f"R{r}"] = f'=IF(B{r}="","",{EN("M", k)})'                                               # R: status now
db.conditional_formatting.add("G8:G23", FormulaRule(formula=['AND(ISNUMBER(G8),G8<0)'], fill=fill(BAD)))
db.conditional_formatting.add("G8:G23", FormulaRule(formula=['R8="Low"'], fill=fill(WARN)))

header(db, 7, 10, ["Sinking fund", "Target", "Saved", "", "By", "Per payday"])
for n in range(8):
    r = 8 + n
    db[f"S{r}"] = f"=IFERROR(SMALL(Envelopes!$P${E0}:$P${E1},{n + 1}),\"\")"                   # S: envelope row
    e = lambda c: f"INDEX(Envelopes!${c}$1:${c}${E1},S{r})"
    style(db.cell(r, 10, f'=IF(S{r}="","",{e("B")})'), False, bold=True)
    style(db.cell(r, 11, f'=IF(S{r}="","",IF(N({e("F")})=0,"",{e("F")}))'), False, MONEY)
    style(db.cell(r, 12, f'=IF(S{r}="","",{e("L")})'), False, MONEY, True)
    style(db.cell(r, 13, f'=IF(OR(S{r}="",K{r}=""),"",{bar(f"L{r}/K{r}", 10)})'), False).font = \
        Font(size=9, color=TEAL)
    style(db.cell(r, 14, f'=IF(S{r}="","",IF({e("G")}="","",{e("G")}))'), False, "d mmm yy", align="center")
    style(db.cell(r, 15, f'=IF(S{r}="","",IF({e("M")}="Reached","Reached",IF({e("N")}="","",{e("N")})))'),
          False, MONEY, align="right")
db.conditional_formatting.add("O8:O15", FormulaRule(formula=['O8="Reached"'], fill=fill(OK)))

header(db, 17, 10, ["Latest spending", "Date", "Envelope", "Amount"])
for n in range(8):
    r = 18 + n
    db[f"T{r}"] = f"=IFERROR(LARGE({SP('K')},{n + 1}),\"\")"                                   # T: spending row
    g = lambda c: f"INDEX(Spending!${c}$1:${c}${S1},MOD(T{r},10000))"
    style(db.cell(r, 10, f'=IF(T{r}="","",IF({g("E")}="","",{g("E")}))'), False)
    style(db.cell(r, 11, f'=IF(T{r}="","",{g("B")})'), False, "d mmm", align="center")
    style(db.cell(r, 12, f'=IF(T{r}="","",{g("C")})'), False, align="center")
    style(db.cell(r, 13, f'=IF(T{r}="","",{g("D")})'), False, MONEY)

header(db, 26, 2, ["Month", "Stuffed", "Spent"])
for k in range(12):
    r = 27 + k
    ms = f"DATE(YEAR({start}),{k + 1},1)"
    me = f"DATE(YEAR({start}),{k + 2},1)-1"
    style(db.cell(r, 2, MON[k]), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS(Paydays!$D${P0}:$D${P1},{PW},">="&{ms},{PW},"<="&{me})'), False, MONEY)
    style(db.cell(r, 4, f'=SUMIFS({SP("D")},{SP("I")},1,{SP("H")},">="&{ms},{SP("H")},"<="&{me})'), False, MONEY)
for c in "QRSTV":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 24)) + list(range(27, 39)):
    db.row_dimensions[r].height = 18
chart = BarChart()
chart.type = "col"
chart.title = "Stuffed and spent each month"
chart.height, chart.width = 7.4, 16
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=26, max_row=38), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=27, max_row=38))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E39A4B"
chart.y_axis.number_format = "#,##0"
chart.legend.position = "b"
db.add_chart(chart, "F26")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: how often you are paid, and the notes your bank gives out."),
    ("2", "Envelopes: each envelope once, with what goes in it each payday and the cash already in it. A sinking "
          "fund can have a target and a date."),
    ("3", "Paydays: each payday, the date and what goes into each envelope. The plan row shows your usual amounts."),
    ("4", "Cash: the notes to ask for at the bank for that payday, and how many go in each envelope."),
    ("5", "Spending: each time you spend cash from an envelope. What is left carries over to the next month."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example envelopes and amounts are made up. Delete them and add your own.",
                          "Cards: print a card for each envelope and note what you spend, then type it in later.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Paydays", "Spending", "Envelopes", "Cash", "Cards", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Cash-Envelope-Budget.xlsx")
wb.save(out)
print("saved", out, len(stuff), "paydays", len(spend), "spending lines")
