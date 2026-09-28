"""Etsy listing photos for the Cash Envelope Budget (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Cash-Envelope-Budget.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups: the envelopes in use, the last paydays and the latest spending.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
env_end = last(wb["Envelopes"], 21)
n_env = env_end - 5
p_end, s_end = last(wb["Paydays"], 205), last(wb["Spending"], 1005)
n_notes = sum(1 for r in range(11, 19) if isinstance(wb["Settings"].cell(r, 3).value, (int, float)))
rows = {"Paydays": list(range(6, p_end - 11)) + list(range(p_end + 1, 206)),
        "Spending": list(range(6, s_end - 19)) + list(range(s_end + 1, 1006)),
        "Envelopes": range(env_end + 1, 22),
        "Cash": range(9 + n_env, 25)}
cols = {"Paydays": [get_column_letter(5 + k) for k in range(n_env, 16)] + ["V"],
        "Envelopes": ["E", "H"],
        "Cash": [get_column_letter(5 + j) for j in range(n_notes, 8)]}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Cash", "Cards"), hide_rows=rows, hide=cols)

img = canvas(dark=True)
title(img, "Cash envelope budget", "Cash stuffing by payday: what goes into each envelope, what you spend, and what "
      "is left to carry over", dark=True, width=2300)
laptop(img, r.shot("Cash envelopes"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-cash-to-take-out.jpg"), "The notes to ask for at the bank",
       "Type what goes into each envelope on payday. Each amount is split into notes, up to the largest you want in "
       "that envelope.", r.shot("Cash to take out"))

detail(os.path.join(img_dir, "03-envelope-cards.jpg"), "A card for each envelope",
       "Print them, keep one in each envelope and note what you spend. Each card starts with the cash in the "
       "envelope on payday.", r.shot("Envelope cards"))

detail(os.path.join(img_dir, "04-paydays.jpg"), "Every payday, envelope by envelope",
       "The plan row shows your usual amounts. What does not go into envelopes stays in the bank for the bills.",
       r.shot("Paydays"))

detail(os.path.join(img_dir, "05-envelopes.jpg"), "Each envelope and what is in it now",
       "What went in, what was spent or moved, and the leftovers that carry over. A sinking fund shows what it needs "
       "each payday to reach its target.", r.shot("Envelopes"), [("#FFE2B8", "Low"), ("#D8F0DC", "Target reached")])

detail(os.path.join(img_dir, "06-spending.jpg"), "Spend from an envelope, or move cash between them",
       "One line each time you spend. Leftovers can go to another envelope, such as your savings.",
       r.shot("Spending"))

img = canvas()
y = title(img, "Inside the file", "Six tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Cash envelopes")), "Dashboard", "The month, envelope by envelope"),
    (content(r.page("Paydays")), "Paydays", "What goes in each one"),
    (content(r.page("Spending")), "Spending", "What comes out"),
    (content(r.page("Envelopes")), "Envelopes", "Plans and sinking funds"),
    (content(r.page("Cash to take out")), "Cash", "The notes for the bank"),
    (content(r.page("Envelope cards")), "Cards", "To print"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "07-whats-inside.jpg"))
