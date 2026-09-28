"""Reading Tracker with a yearly goal and a bookshelf (Excel + Google Sheets).

Every book with its author, genre, pages, dates and rating; a reading goal for
the year that says if you are ahead or behind; books and pages by month and by
genre; what you are reading now; and a bookshelf that fills up with the spines
of the books you finish. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (DATE, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
B0, B1 = 6, 1005      # books
G0, G1 = 6, 17        # genres on the Settings tab
BK = lambda col: f"Books!${col}${B0}:${col}${B1}"
GENRES = [("Fiction", "8CC5A8"), ("Fantasy", "B39DDB"), ("Mystery", "7E9BB8"), ("Romance", "F2A7B8"),
          ("Science fiction", "80CBC4"), ("Historical fiction", "D7B98E"), ("Non-fiction", "F3DE8A"),
          ("Biography", "F2A65A"), ("Self-help", "A5D6A7"), ("Poetry", "CE93D8"), ("Classics", "BCAAA4"), ("Other", "CFD8DC")]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
YR = "Dashboard!$C$4"
IN_YEAR = f'{BK("G")},"Finished",{BK("I")},">="&DATE({YR},1,1),{BK("I")},"<="&DATE({YR},12,31)'

# ---------------------------------------------------------------- example data (made-up books and authors)
rng = random.Random(12)
ADJ = ["Salt", "Quiet", "Paper", "Winter", "Glass", "Last", "Hollow", "Silver", "Crooked", "Burning", "Northern", "Borrowed",
       "Little", "Wild", "Hidden", "Painted", "Lantern", "Midnight", "Copper", "Drowned"]
NOUN = ["Orchard", "Harbour", "Garden", "Kingdom", "Lighthouse", "House", "Tide", "Map", "Clockmaker", "River",
        "Island", "Archive", "Summer", "Ferry", "Bookshop", "Mountain", "Machine", "Orchestra", "Daughter", "Promise"]
OF = ["Salt", "Glass", "Ash", "Winter", "Small Hours", "Lost Things", "Quiet Water", "Borrowed Time", "Other Rooms",
      "Northern Lights"]
FIRST = ["Mira", "Tomas", "Elena", "Rafe", "Iris", "Colm", "Nadia", "Jude", "Ines", "Otto", "Maren", "Silas", "Wren", "Anika"]
LAST = ["Holloway", "Varga", "Quill", "Ashby", "Lindqvist", "Okafor", "Marlowe", "Brandt", "Castell", "Fennimore", "Rook",
        "Albescu", "Thorne", "Delacroix-Park"]
titles, books = set(), []


def title_():
    while True:
        t = rng.choice([f"The {rng.choice(ADJ)} {rng.choice(NOUN)}", f"The {rng.choice(NOUN)} of {rng.choice(OF)}",
                        f"{rng.choice(ADJ)} {rng.choice(NOUN)}"])
        if t not in titles:
            titles.add(t)
            return t


author = lambda: f"{rng.choice(FIRST)} {rng.choice(LAST)}"
genre_w = [18, 14, 12, 9, 8, 8, 10, 6, 5, 3, 4, 3]
d = dt.date(Y, 1, 2)
while True:
    pages = rng.randint(190, 560)
    days = max(3, round(pages / rng.uniform(42, 78)))
    end = d + dt.timedelta(days=days)
    if end >= today - dt.timedelta(days=2):
        break
    g = rng.choices([x[0] for x in GENRES], genre_w)[0]
    rating = rng.choices([2, 3, 4, 5], [1, 4, 8, 5])[0]
    note = rng.choice([None] * 6 + ["Could not put it down", "Slow start, great ending", "Book club pick",
                                    "Read it on holiday", "Lent to a friend"])
    books.append([title_(), author(), g, rng.choice(["Paper", "Paper", "Ebook", "Audio"]), pages, "Finished", d, end,
                  rating, note])
    d = end + dt.timedelta(days=rng.randint(0, 2))
books[len(books) // 2][5:10] = ["Did not finish", books[len(books) // 2][6], None, None, "Not for me"]
for pages, started in ((412, today - dt.timedelta(days=6)), (268, today - dt.timedelta(days=2))):
    books.append([title_(), author(), rng.choice(["Fantasy", "Mystery"]), "Paper", pages, "Reading", started, None, None,
                  None])
for _ in range(6):
    books.append([title_(), author(), rng.choices([x[0] for x in GENRES], genre_w)[0], None, rng.randint(220, 480),
                  "Want to read", None, None, None, rng.choice([None, "Recommended by Sam", "On sale"])])

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL, span=1):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    first = ws[f"{col}{row}"].column
    for r in (row, row + 1):
        for c in range(first, first + span):
            ws.cell(r, c).border = box
        if span > 1:
            ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=first + span - 1)
    ws.row_dimensions[row + 1].height = 34


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your genres and their colours on the bookshelf.", [3, 22, 8, 4, 60], rows=30)
header(se, 5, 2, ["Genre", "Colour"])
for k, (g, colour) in enumerate(GENRES):
    style(se.cell(G0 + k, 2, g), True, bold=True)
    c = se.cell(G0 + k, 3)
    c.fill = fill(colour)
    c.border = box
se["E6"] = "Rename the genres to the ones you read. Each place on the list keeps its colour."
se["E7"] = "A book with a genre that is not on the list gets the last colour."
se["E7"].font = font(10, color=MUTED, italic=True)
se["E6"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Books
bk = wb.create_sheet("Books")
sheet_base(bk, "Books", "One line per book: the ones you read, are reading, and want to read.",
           [3, 30, 20, 16, 9, 8, 14, 13, 13, 8, 7, 9, 28], rows=B1 + 2)
header(bk, 5, 2, ["Title", "Author", "Genre", "Format", "Pages", "Status", "Started", "Finished", "Rating", "Days",
                  "Pages a day", "Note"])
for r in range(B0, B1 + 1):
    b = books[r - B0] if r - B0 < len(books) else (None,) * 10
    t_, a_, g_, f_, p_, s_, st_, en_, ra_, no_ = b
    style(bk.cell(r, 2, t_), True, bold=True)
    style(bk.cell(r, 3, a_), True)
    style(bk.cell(r, 4, g_), True)
    style(bk.cell(r, 5, f_), True, align="center")
    style(bk.cell(r, 6, p_), True, "0", align="center")
    style(bk.cell(r, 7, s_), True, align="center")
    style(bk.cell(r, 8, st_), True, DATE)
    style(bk.cell(r, 9, en_), True, DATE)
    style(bk.cell(r, 10, ra_), True, "0", align="center")
    style(bk.cell(r, 11, f'=IF(OR(H{r}="",I{r}=""),"",I{r}-H{r}+1)'), False, "0", align="center")
    style(bk.cell(r, 12, f'=IF(OR(K{r}="",N(F{r})=0),"",F{r}/K{r})'), False, "0", align="center")
    style(bk.cell(r, 13, no_), True)
    bk.cell(r, 14, f'=IF(AND(G{r}="Finished",I{r}<>""),IF(YEAR(I{r})={YR},I{r}+ROW()/10000000,""),"")')   # N shelf order
    bk.cell(r, 15, f'=IF(B{r}="","",IFERROR(MATCH(D{r},Settings!$B${G0}:$B${G1},0),{G1 - G0 + 1}))')        # O genre number
    bk.cell(r, 16, f'=IF(AND(G{r}="Reading",B{r}<>""),IF(H{r}="",TODAY(),H{r})+ROW()/10000000,"")')        # P reading now
    bk.cell(r, 17, f'=IF(AND(N{r}<>"",N(J{r})=5),N{r},"")')                                                # Q five stars
for c in "NOPQ":
    bk.column_dimensions[c].hidden = True
dropdown(bk, f"=Settings!$B${G0}:$B${G1}", f"D{B0}:D{B1}")
dropdown(bk, '"Paper,Ebook,Audio"', f"E{B0}:E{B1}")
dropdown(bk, '"Want to read,Reading,Finished,Did not finish"', f"G{B0}:G{B1}")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
bk.add_data_validation(dv)
dv.add(f"J{B0}:J{B1}")
bk.conditional_formatting.add(f"G{B0}:G{B1}", FormulaRule(formula=[f'G{B0}="Finished"'], fill=fill(OK)))
bk.conditional_formatting.add(f"G{B0}:G{B1}", FormulaRule(formula=[f'G{B0}="Reading"'], fill=fill("FCE9B8")))
bk.conditional_formatting.add(f"G{B0}:G{B1}", FormulaRule(formula=[f'G{B0}="Did not finish"'], fill=fill("EEF1F1"),
                                                          font=Font(name=F, color="9AA7A9")))
bk.conditional_formatting.add(f"J{B0}:J{B1}", FormulaRule(formula=[f"N(J{B0})=5"], fill=fill("F3DE8A"),
                                                          font=Font(name=F, bold=True)))
bk.freeze_panes = "C6"

# ---------------------------------------------------------------- Bookshelf
sh = wb.create_sheet("Bookshelf")
SPINES, SHELVES, HELP = 20, 5, 25      # 20 books a shelf; hidden genre numbers from column Y
sheet_base(sh, "Bookshelf", "Every book you finish in the year on the Dashboard, in the order you finished them and in "
           "the colour of its genre. The shelves hold 100 books.", [3] + [4.6] * SPINES + [3], rows=30)
sh["B2"] = f'="Bookshelf "&{YR}'
sh["B4"] = f'=COUNT({BK("N")})&IF(COUNT({BK("N")})=1," book"," books")&" in "&{YR}'
sh["B4"].font = font(13, True, TEAL_D)
sh["B4"].alignment = Alignment(vertical="center")
sh.row_dimensions[4].height = 26
for s in range(SHELVES):
    top = 5 + s * 4
    sh.row_dimensions[top].height = 132
    sh.row_dimensions[top + 1].height = 7
    sh.row_dimensions[top + 2].hidden = True
    sh.row_dimensions[top + 3].height = 10
    for k in range(SPINES):
        n = s * SPINES + k + 1
        col = L(2 + k)
        key = f"{col}{top + 2}"
        sh[key] = f'=IFERROR(SMALL({BK("N")},{n}),"")'
        c = sh.cell(top, 2 + k, f'=IF({key}="","",INDEX({BK("B")},MATCH({key},{BK("N")},0)))')
        c.font = font(8, True, "1E2A2F")
        c.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=False)
        c.border = box
        sh.cell(top + 2, HELP + k, f'=IF({key}="",0,INDEX({BK("O")},MATCH({key},{BK("N")},0)))')     # genre number
    for k in range(SPINES + 1):
        sh.cell(top + 1, 2 + k).fill = fill("8D6E63")
    rng_ = f"B{top}:{L(1 + SPINES)}{top}"
    for g, (name, colour) in enumerate(GENRES, start=1):
        sh.conditional_formatting.add(rng_, FormulaRule(formula=[f"{L(HELP)}{top + 2}={g}"], fill=fill(colour)))
for c in range(HELP, HELP + SPINES):
    sh.column_dimensions[L(c)].hidden = True

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Reading dashboard", "Your goal for the year, what you are reading, and your books by month and genre.",
           [3, 22, 14, 14, 14, 14, 14, 3, 30, 12, 10, 10, 3], rows=60)
db["B4"] = "Year"
db["B4"].font = font(11, True)
style(db["C4"], True, "0", True, "center")
db["C4"] = Y
db["E4"] = "Goal (books)"
db["E4"].font = font(11, True)
db["E4"].alignment = Alignment(horizontal="right")
style(db["F4"], True, "0", True, "center")
db["F4"] = 45
done = f"COUNTIFS({IN_YEAR})"
share = f"IF({YR}<YEAR(TODAY()),1,IF({YR}>YEAR(TODAY()),0,(TODAY()-DATE({YR},1,1)+1)/(DATE({YR},12,31)-DATE({YR},1,1)+1)))"
expected = f"ROUND($F$4*{share},0)"
tile(db, "B", 6, f'="Books in "&{YR}', f"={done}", "0")
tile(db, "C", 6, "Goal", '=IF(N($F$4)=0,"",$F$4)', "0", MUTED)
tile(db, "D", 6, "Pages", f'=SUMIFS({BK("F")},{IN_YEAR})', "#,##0")
tile(db, "E", 6, "Average rating", f'=IFERROR(AVERAGEIFS({BK("J")},{IN_YEAR}),"")', "0.0", "B07A00")
tile(db, "F", 6, "Reading now", f'=COUNTIF({BK("G")},"Reading")', "0")
tile(db, "G", 6, "Want to read", f'=COUNTIF({BK("G")},"Want to read")', "0")
db["B9"] = f'=IF(N($F$4)=0,"",{bar(f"MIN(1,{done}/$F$4)", 30)})'
db["B9"].font = Font(name=F, size=14, color=TEAL)
db.merge_cells("B9:E9")
db["F9"] = (f'=IF(N($F$4)=0,"",IF({done}>={expected}+1,"Ahead by "&({done}-{expected})&IF({done}-{expected}=1," book"," books"),'
            f'IF({done}<={expected}-1,"Behind by "&({expected}-{done})&IF({expected}-{done}=1," book"," books"),"On track")))')
db["F9"].font = font(12, True, TEAL_D)
db.merge_cells("F9:G9")
db.conditional_formatting.add("F9", FormulaRule(formula=['LEFT(F9,6)="Behind"'], fill=fill(WARN)))
db.conditional_formatting.add("F9", FormulaRule(formula=['OR(LEFT(F9,5)="Ahead",F9="On track")'], fill=fill(OK)))
db.row_dimensions[9].height = 26

header(db, 11, 2, ["Month", "Books", "Pages"])
for m in range(12):
    r = 12 + m
    within = (f'{BK("G")},"Finished",{BK("I")},">="&DATE({YR},{m + 1},1),{BK("I")},"<="&EOMONTH(DATE({YR},{m + 1},1),0)')
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f"=COUNTIFS({within})"), False, "0", align="center")
    style(db.cell(r, 4, f'=SUMIFS({BK("F")},{within})'), False, "#,##0", align="center")
chart = BarChart()
chart.type = "col"
chart.title = "Books by month"
chart.height, chart.width = 7.2, 8.4
chart.add_data(Reference(db, min_col=3, min_row=11, max_row=23), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=12, max_row=23))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
db.add_chart(chart, "E11")

header(db, 11, 9, ["Genre", "Books", "Pages", "Rating"])
for k in range(G1 - G0 + 1):
    r = 12 + k
    g = f"Settings!B{G0 + k}"
    style(db.cell(r, 9, f'=IF({g}="","",{g})'), False, bold=True)
    style(db.cell(r, 10, f'=IF(I{r}="","",COUNTIFS({IN_YEAR},{BK("D")},I{r}))'), False, "0", align="center")
    style(db.cell(r, 11, f'=IF(I{r}="","",SUMIFS({BK("F")},{IN_YEAR},{BK("D")},I{r}))'), False, "#,##0", align="center")
    style(db.cell(r, 12, f'=IF(OR(I{r}="",N(J{r})=0),"",AVERAGEIFS({BK("J")},{IN_YEAR},{BK("D")},I{r}))'), False, "0.0",
          align="center")
    colour = GENRES[k][1]
    db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f"N(J{r})>0"], fill=fill(colour)))

header(db, 26, 2, ["Reading now", "", "Author", "", "Started", "Days"])
db.merge_cells("B26:C26")
db.merge_cells("D26:E26")
for k in range(4):
    r = 27 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({BK("P")},{k + 1}),"")'
    row = f'MATCH({key},{BK("P")},0)'
    style(db.cell(r, 2, f'=IF({key}="","",INDEX({BK("B")},{row}))'), False, bold=True)
    style(db.cell(r, 3), False)
    db.merge_cells(f"B{r}:C{r}")
    style(db.cell(r, 4, f'=IF({key}="","",INDEX({BK("C")},{row}))'), False)
    style(db.cell(r, 5), False)
    db.merge_cells(f"D{r}:E{r}")
    style(db.cell(r, 6, f'=IF({key}="","",INT({key}))'), False, "d mmm", align="center")
    style(db.cell(r, 7, f'=IF({key}="","",TODAY()-INT({key})+1)'), False, "0", align="center")

header(db, 26, 9, ["Five stars this year", "Author"])
db.merge_cells("J26:L26")
for k in range(6):
    r = 27 + k
    key = f"P{r}"
    db[key] = f'=IFERROR(LARGE({BK("Q")},{k + 1}),"")'
    row = f'MATCH({key},{BK("Q")},0)'
    style(db.cell(r, 9, f'=IF({key}="","",INDEX({BK("B")},{row}))'), False, bold=True)
    style(db.cell(r, 10, f'=IF({key}="","",INDEX({BK("C")},{row}))'), False)
    for c in (11, 12):
        style(db.cell(r, c), False)
    db.merge_cells(f"J{r}:L{r}")
for c in "OP":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Dashboard: the year and how many books you want to read in it."),
    ("2", "Books: one line per book, with its author, genre, format and pages."),
    ("3", "When you start a book, set it to Reading with the date. When you finish, set Finished, the date and a rating."),
    ("4", "Books you want to read go in the same list, as Want to read."),
    ("5", "Dashboard and Bookshelf: your goal, your months and genres, and shelves for 100 books that fill up as you read."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["Ahead or behind compares your finished books with where the goal says you should be today.",
                          "The example books and authors are made up. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Books", "Bookshelf", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Reading-Tracker.xlsx")
wb.save(out)
print("saved", out, len(books), "books", sum(b[5] == "Finished" for b in books), "finished")
