"""Donation and Volunteer Tracker for a small charity, club or church (Excel + Google Sheets).

Donors with their first and last gift and what they gave this year; each gift with its campaign and whether it was
thanked; campaigns against their goals; volunteers and their hours; a statement of a donor's gifts for the year to
print; and a Dashboard with the year's giving, the campaigns, the thank-you notes still to send and the volunteer
hours. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, bar, box, fill, font, header,  # noqa: E402
                      sheet_base, style)

today = dt.date.today()
D0, D1 = 6, 305        # donors
G0, G1 = 6, 1005       # gifts
C0, C1 = 6, 25         # campaigns
V0, V1 = 6, 105        # volunteers
H0, H1 = 6, 1005       # hours
YEAR, CUR, LAPSE = "Settings!$C$7", "Settings!$C$8", "Settings!$C$9"
ORG, ADDR, NOTE = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6"
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
GF = lambda c: f"Gifts!${c}${G0}:${c}${G1}"
GF1 = lambda c: f"Gifts!${c}$1:${c}${G1}"
DN = lambda c: f"Donors!${c}${D0}:${c}${D1}"
DN1 = lambda c: f"Donors!${c}$1:${c}${D1}"
HR = lambda c: f"Hours!${c}${H0}:${c}${H1}"


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})&" "&YEAR({d})'


rng = random.Random(41)
D = lambda n: today + dt.timedelta(days=n)

# ---------------------------------------------------------------- example data (a made-up community pantry)
donors = [  # name, type, email, address
    ("Grace Holloway", "Person", "grace@mail.example", "12 Birch Road, Riverton"),
    ("Daniel Mercer", "Person", "dan.mercer@mail.example", "4 Quay Street, Riverton"),
    ("Riverton Hardware", "Business", "shop@rivertonhardware.example", "88 High Street, Riverton"),
    ("Amina Yusuf", "Person", "amina.y@mail.example", "31 Elm Close, Riverton"),
    ("St Luke's Youth Group", "Group", None, "Church Lane, Riverton"),
    ("Tom and Jess Carter", "Person", "carters@mail.example", "7 Mill Yard, Riverton"),
    ("Corner Bakery", "Business", "hello@cornerbakery.example", "2 Market Square, Riverton"),
    ("Helen Price", "Person", None, "19 Orchard Way, Riverton"),
    ("Ravi Patel", "Person", "ravi.patel@mail.example", "5 Station Road, Riverton"),
    ("Riverton Rotary", "Group", "rotary@mail.example", None),
    ("Sofia Rossi", "Person", "sofia.r@mail.example", "40 Hill Street, Riverton"),
    ("Old Mill Trust", "Foundation", "grants@oldmill.example", "The Old Mill, Riverton"),
]
campaigns = [  # name, goal, start, end
    ("General fund", None, None, None),
    ("Winter coat drive", 3000, D(-40), D(50)),
    ("New freezer", 5000, D(-150), D(20)),
    ("Back to school", 600, D(-120), D(-30)),
    ("Spring garden", 800, D(30), D(120)),
]
gifts = []  # date, donor, amount, how, campaign, goods, thanked, note
regular = {"Grace Holloway": 25, "Ravi Patel": 20, "Helen Price": 10, "Amina Yusuf": 15}
for m in range(15):
    first = today.replace(day=1) - dt.timedelta(days=30 * m)
    for name, amount in regular.items():
        if name == "Helen Price" and m < 13:
            continue                                            # Helen stopped giving over a year ago
        gifts.append([first.replace(day=5 + len(name) % 7), name, amount, "Bank transfer", "General fund", None, "x",
                      "Monthly"])
extra = [  # date, donor, amount, how, campaign, goods, note
    (D(-400), "Riverton Hardware", 400, "Check", "General fund", None, None),
    (D(-380), "Old Mill Trust", 2000, "Bank transfer", "General fund", None, "Yearly grant"),
    (D(-330), "Corner Bakery", 120, "Card", "General fund", None, None),
    (D(-300), "St Luke's Youth Group", 260, "Cash", "General fund", None, "Carol singing"),
    (D(-140), "Riverton Hardware", 500, "Check", "New freezer", None, None),
    (D(-120), "Old Mill Trust", 2500, "Bank transfer", "New freezer", None, "Yearly grant"),
    (D(-100), "Corner Bakery", 150, "Card", "Back to school", None, None),
    (D(-95), "St Luke's Youth Group", 320, "Cash", "Back to school", None, "Car wash"),
    (D(-90), "Tom and Jess Carter", 100, "Card", "Back to school", None, None),
    (D(-75), None, 64, "Cash", "General fund", None, "Collection box"),
    (D(-60), "Riverton Rotary", 750, "Check", "New freezer", None, None),
    (D(-45), "Daniel Mercer", 60, "Online", "Back to school", None, None),
    (D(-35), "Sofia Rossi", 40, "Online", "Winter coat drive", None, None),
    (D(-30), "Corner Bakery", 200, "Card", "Winter coat drive", None, None),
    (D(-21), "Tom and Jess Carter", 120, "Online", "Winter coat drive", None, None),
    (D(-18), "Riverton Hardware", None, "Goods", "Winter coat drive", "12 boxes of gloves", None),
    (D(-14), "Riverton Hardware", 300, "Check", "Winter coat drive", None, None),
    (D(-10), "Daniel Mercer", 80, "Online", "Winter coat drive", None, None),
    (D(-8), "Corner Bakery", None, "Goods", "General fund", "40 loaves of bread", None),
    (D(-6), "St Luke's Youth Group", 185, "Cash", "Winter coat drive", None, "Bake sale"),
    (D(-4), None, 48, "Cash", "General fund", None, "Collection box"),
    (D(-3), "Sofia Rossi", 50, "Online", "New freezer", None, None),
    (D(-2), "Riverton Rotary", 400, "Check", "Winter coat drive", None, None),
    (D(-1), "Lena Brooks", 30, "Online", "Winter coat drive", None, "First gift"),
]
for d, name, amount, how, camp, goods, note in extra:
    thanked = "x" if name and d < today - dt.timedelta(days=12) else None
    gifts.append([d, name, amount, how, camp, goods, thanked, note])
gifts.sort(key=lambda g: g[0])
volunteers = [("Mark Evans", "Driver"), ("Lucy Chen", "Sorting"), ("Peter Olsen", "Front desk"),
              ("Nadia Karim", "Sorting"), ("Joe Ward", "Driver"), ("Beth Morgan", "Front desk"),
              ("Sam Reid", "Sorting"), ("Olga Novak", "Kitchen")]
contact = lambda name: name.split()[0].lower() + "@mail.example"
hours = []
for w in range(40):
    day = today - dt.timedelta(days=7 * w + 2)
    for name, role in volunteers:
        if rng.random() < 0.35:
            hours.append([day, name, rng.choice([2, 2.5, 3, 4]), rng.choice(["Saturday market", "Sorting day",
                                                                               "Deliveries", "Front desk"]), None])
hours.sort(key=lambda h: h[0])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, c, row, label, formula, fmt, colour=TEAL):
    ws[f"{c}{row}"] = label
    ws[f"{c}{row}"].font = font(10, True, MUTED)
    ws[f"{c}{row + 1}"] = formula
    ws[f"{c}{row + 1}"].number_format = fmt
    ws[f"{c}{row + 1}"].font = font(16, True, colour)
    ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{c}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your organisation for the statements, the year shown, and when a donor counts as lapsed.",
           [3, 30, 44, 4, 50], rows=14)
for r, lab, value, fmt, note in (
        (4, "Organisation", "Riverton Community Pantry", None, "At the top of each statement."),
        (5, "Address", "10 Church Lane, Riverton", None, "Under the name."),
        (6, "Line at the end of each statement",
         "Thank you for your support. This statement lists the gifts we received from you. Please keep it for your "
         "records.", None, "Add your charity number or anything your rules ask for."),
        (7, "Year", "=YEAR(TODAY())", "0", "The year on the Dashboard and the statements. Type another to look back."),
        (8, "Currency", "USD", None, "Shown on the statements."),
        (9, "Lapsed after (months)", 12, "0", "A donor with no gift for this long counts as lapsed.")):
    se[f"B{r}"] = lab
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, r == 4, "center" if r in (7, 8, 9) else None)
    se[f"C{r}"] = value
    se.cell(r, 5, note).font = font(10, color=MUTED, italic=True)
    if r == 6:
        se[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        se[f"B{r}"].alignment = se[f"E{r}"].alignment = Alignment(vertical="top")
        se.row_dimensions[r].height = 48

# ---------------------------------------------------------------- Gifts
gf = wb.create_sheet("Gifts")
sheet_base(gf, "Gifts", "Every gift, with its campaign. For goods, leave the amount empty and say what was given. "
           "An x once the donor has been thanked.", [3, 13, 24, 11, 14, 20, 22, 9, 18, 20], rows=G1 + 2)
header(gf, 5, 2, ["Date", "Donor", "Amount", "How", "Campaign", "Goods given", "Thanked", "Note", "Check"])
for r in range(G0, G1 + 1):
    x = gifts[r - G0] if r - G0 < len(gifts) else (None,) * 8
    for c, v, fmt in zip(range(2, 10), x, (DATE, None, MONEY, None, None, None, None, None)):
        style(gf.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 5, 8) else None)
    style(gf.cell(r, 10, f'=IF(AND(C{r}="",D{r}="",G{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                         f'IF(AND(D{r}<>"",NOT(ISNUMBER(D{r}))),"Amount not a number",IF(N(D{r})<0,"Amount below zero",'
                         f'IF(AND(N(D{r})=0,G{r}=""),"Needs an amount or goods",IF(C{r}="","",'
                         f'IF(COUNTIF({DN("B")},C{r})=0,"Not on Donors","")))))))'), False, align="center")
    gf[f"K{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r}),"")'                                          # K: year
    gf[f"L{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                          # L: month
    gf[f"M{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                            # M: date as a number
    gf[f"N{r}"] = (f'=IF(OR(C{r}="",F{r}="",Q{r}=0),0,'
                   f'IF(COUNTIFS(C${G0}:C{r},C{r},F${G0}:F{r},F{r},Q${G0}:Q{r},1)=1,1,0))')  # N: first for the campaign
    gf[f"O{r}"] = f'=IF(OR(C{r}="",Q{r}=0,H{r}<>""),"",M{r}*10000+ROW())'                     # O: to thank
    gf[f"P{r}"] = (f'=IF(OR(C{r}="",Q{r}=0),"",IF(AND(C{r}=Statement!$C$4,K{r}=Statement!$C$5),'
                   f'M{r}*10000+ROW(),""))')                                                  # P: on the statement
    gf[f"Q{r}"] = f'=IF(AND(ISNUMBER(B{r}),OR(N(D{r})>0,G{r}<>"")),1,0)'                      # Q: a gift that counts
for c in "KLMNOPQ":
    gf.column_dimensions[c].hidden = True
gf.conditional_formatting.add(f"J{G0}:J{G1}", FormulaRule(formula=[f'J{G0}<>""'], fill=fill(BAD)))
dropdown(gf, f"={DN('B')}", f"C{G0}:C{G1}", strict=False)
dropdown(gf, '"Cash,Check,Card,Bank transfer,Online,Goods,Other"', f"E{G0}:E{G1}", strict=False)
dropdown(gf, f"=Campaigns!$B${C0}:$B${C1}", f"F{G0}:F{G1}", strict=False)
gf.freeze_panes = "D6"

# ---------------------------------------------------------------- Donors
dn = wb.create_sheet("Donors")
sheet_base(dn, "Donors", "Each donor once. Their gifts add up from the Gifts tab.",
           [3, 24, 11, 26, 28, 13, 13, 12, 12, 8, 12], rows=D1 + 2)
header(dn, 5, 2, ["Donor", "Type", "Email", "Address", "First gift", "Last gift", "This year", "All time", "Gifts",
                  "Status"])
for r in range(D0, D1 + 1):
    x = donors[r - D0] if r - D0 < len(donors) else (None,) * 4
    for c, v in zip(range(2, 6), x):
        style(dn.cell(r, c, v), True, bold=c == 2, align="center" if c == 3 else None)
    mine = f"({GF('C')}=B{r})*{GF('Q')}"
    style(dn.cell(r, 6, f'=IF(OR(B{r}="",N(J{r})=0),"",SUMPRODUCT(MIN({mine}*{GF("M")}+(1-{mine})*1E9)))'), False,
          DATE, align="center")
    style(dn.cell(r, 7, f'=IF(OR(B{r}="",N(J{r})=0),"",SUMPRODUCT(MAX({mine}*{GF("M")})))'), False, DATE,
          align="center")
    style(dn.cell(r, 8, f'=IF(B{r}="","",SUMIFS({GF("D")},{GF("C")},B{r},{GF("K")},{YEAR},{GF("D")},">0"))'), False, MONEY, True)
    style(dn.cell(r, 9, f'=IF(B{r}="","",SUMIFS({GF("D")},{GF("C")},B{r},{GF("D")},">0",{GF("M")},">0"))'), False, MONEY)
    style(dn.cell(r, 10, f'=IF(B{r}="","",SUMIFS({GF("Q")},{GF("C")},B{r}))'), False, "0", align="center")
    style(dn.cell(r, 11, f'=IF(B{r}="","",IF(N(J{r})=0,"No gifts yet",IF(AND(N({LAPSE})>0,'
                         f'G{r}<EDATE(TODAY(),-N({LAPSE}))),"Lapsed",IF(YEAR(F{r})=N({YEAR}),"New","Active"))))'),
          False, bold=True, align="center")
    dn[f"M{r}"] = f'=IF(OR(B{r}="",N(H{r})<=0),"",ROUND(H{r}*100,0)*1000+1000-ROW())'        # M: top donors
    dn[f"N{r}"] = f'=IF(B{r}="","",SUMIFS({GF("Q")},{GF("C")},B{r},{GF("K")},{YEAR}))'       # N: gifts this year
for c in "MN":
    dn.column_dimensions[c].hidden = True
for text, colour in (("New", OK), ("Lapsed", WARN)):
    dn.conditional_formatting.add(f"K{D0}:K{D1}", FormulaRule(formula=[f'K{D0}="{text}"'], fill=fill(colour)))
dropdown(dn, '"Person,Business,Group,Foundation,Other"', f"C{D0}:C{D1}", strict=False)
dn.freeze_panes = "C6"

# ---------------------------------------------------------------- Campaigns
cp = wb.create_sheet("Campaigns")
sheet_base(cp, "Campaigns", "What you raise money for, with a goal and dates if it has them.",
           [3, 24, 12, 13, 13, 12, 9, 16, 9, 10, 12], rows=C1 + 2)
header(cp, 5, 2, ["Campaign", "Goal", "Starts", "Ends", "Raised", "Of goal", "", "Donors", "Days left", "Status"])
for r in range(C0, C1 + 1):
    x = campaigns[r - C0] if r - C0 < len(campaigns) else (None,) * 4
    for c, v, fmt in zip(range(2, 6), x, (None, MONEY, DATE, DATE)):
        style(cp.cell(r, c, v), True, fmt, c == 2, "center" if c in (4, 5) else None)
    style(cp.cell(r, 6, f'=IF(B{r}="","",SUMIFS({GF("D")},{GF("F")},B{r},{GF("D")},">0",{GF("M")},">0"))'), False, MONEY, True)
    style(cp.cell(r, 7, f'=IF(OR(B{r}="",N(C{r})<=0),"",F{r}/C{r})'), False, "0%", align="center")
    style(cp.cell(r, 8, f'=IF(G{r}="","",{bar(f"G{r}", 12)})'), False).font = Font(size=9, color=TEAL)
    style(cp.cell(r, 9, f'=IF(B{r}="","",SUMIFS({GF("N")},{GF("F")},B{r}))'), False, "0", align="center")
    style(cp.cell(r, 10, f'=IF(OR(B{r}="",NOT(ISNUMBER(E{r}))),"",MAX(0,E{r}-TODAY()))'), False, "0",
          align="center")
    style(cp.cell(r, 11, f'=IF(B{r}="","",IF(AND(N(C{r})>0,F{r}>=N(C{r})),"Reached",'
                         f'IF(AND(ISNUMBER(E{r}),E{r}<TODAY()),"Ended",IF(AND(ISNUMBER(D{r}),D{r}>TODAY()),'
                         f'"Not started","Running"))))'), False, bold=True, align="center")
cp.cell(5, 8).value = None
cp.conditional_formatting.add(f"K{C0}:K{C1}", FormulaRule(formula=[f'K{C0}="Reached"'], fill=fill(OK)))

# ---------------------------------------------------------------- Volunteers and Hours
vo = wb.create_sheet("Volunteers")
sheet_base(vo, "Volunteers", "Each volunteer once. Their hours add up from the Hours tab.",
           [3, 22, 26, 16, 12, 12, 13, 10], rows=V1 + 2)
header(vo, 5, 2, ["Volunteer", "Phone or email", "Role", "This year", "All time", "Last time", "Times"])
for r in range(V0, V1 + 1):
    x = volunteers[r - V0] if r - V0 < len(volunteers) else (None, None)
    style(vo.cell(r, 2, x[0]), True, bold=True)
    style(vo.cell(r, 3, contact(x[0]) if x[0] else None), True)
    style(vo.cell(r, 4, x[1]), True)
    style(vo.cell(r, 5, f'=IF(B{r}="","",SUMIFS({HR("D")},{HR("C")},B{r},{HR("I")},{YEAR},{HR("D")},">0"))'), False, "0.#",
          True, "center")
    style(vo.cell(r, 6, f'=IF(B{r}="","",SUMIFS({HR("D")},{HR("C")},B{r},{HR("D")},">0",{HR("K")},">0"))'), False, "0.#", align="center")
    style(vo.cell(r, 7, f'=IF(OR(B{r}="",N(F{r})=0),"",SUMPRODUCT(MAX(({HR("C")}=B{r})*{HR("L")})))'), False,
          DATE, align="center")
    style(vo.cell(r, 8, f'=IF(B{r}="","",COUNTIFS({HR("C")},B{r},{HR("I")},{YEAR},{HR("D")},">0"))'), False, "0",
          align="center")
dropdown(vo, '"Driver,Sorting,Front desk,Kitchen,Other"', f"D{V0}:D{V1}", strict=False)
vo.freeze_panes = "C6"
hr = wb.create_sheet("Hours")
sheet_base(hr, "Hours", "Each time someone helps, with the hours as a number: 2.5 for two and a half.",
           [3, 13, 22, 9, 22, 24, 20], rows=H1 + 2)
header(hr, 5, 2, ["Date", "Volunteer", "Hours", "What", "Note", "Check"])
for r in range(H0, H1 + 1):
    x = hours[r - H0] if r - H0 < len(hours) else (None,) * 5
    for c, v, fmt in zip(range(2, 7), x, (DATE, None, "0.#", None, None)):
        style(hr.cell(r, c, v), True, fmt, c == 3, "center" if c in (2, 4) else None)
    style(hr.cell(r, 7, f'=IF(AND(C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(B{r})),"Needs a date",'
                        f'IF(AND(D{r}<>"",NOT(ISNUMBER(D{r}))),"Hours not a number",IF(N(D{r})<=0,"Needs the hours",'
                        f'IF(C{r}="","Needs a name",IF(COUNTIF(Volunteers!$B${V0}:$B${V1},C{r})=0,"Not on Volunteers",'
                        f'""))))))'), False, align="center")
    hr[f"I{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r}),"")'                                          # I: year
    hr[f"J{r}"] = f'=IF(ISNUMBER(B{r}),YEAR(B{r})*100+MONTH(B{r}),"")'                          # J: month
    hr[f"K{r}"] = f"=IF(ISNUMBER(B{r}),INT(B{r}),0)"                                            # K: date as a number
    hr[f"L{r}"] = f"=IF(AND(ISNUMBER(D{r}),N(D{r})>0),K{r},0)"                                  # L: date with hours
for c in "IJKL":
    hr.column_dimensions[c].hidden = True
hr.conditional_formatting.add(f"G{H0}:G{H1}", FormulaRule(formula=[f'G{H0}<>""'], fill=fill(BAD)))
dropdown(hr, f"=Volunteers!$B${V0}:$B${V1}", f"C{H0}:C{H1}", strict=False)
hr.freeze_panes = "C6"

# ---------------------------------------------------------------- Statement (to print)
stt = wb.create_sheet("Statement")
sheet_base(stt, "", "", [3, 14, 13, 14, 20, 28], rows=48)
stt["B2"] = f'=IF({ORG}="","",{ORG})'
stt["B2"].font = font(18, True, TEAL_D)
stt["B3"] = f'=IF({ADDR}="","",{ADDR})'
stt["B3"].font = font(10, color=MUTED)
stt["B4"] = "Donor"
stt["B5"] = "Year"
for c in ("B4", "B5"):
    stt[c].font = font(11, True)
style(stt["C4"], True, bold=True)
stt["C4"] = donors[0][0]
stt.merge_cells("C4:E4")
style(stt["C5"], True, "0", True, "left")
stt["C5"] = f"={YEAR}"
dropdown(stt, f"={DN('B')}", "C4", strict=False)
stt["B7"] = '=IF(C4="","","Statement of gifts for "&C5)'
stt["B7"].font = font(15, True, TEAL_D)
stt["B8"] = (f'=IF(C4="","",C4&IFERROR(IF(INDEX({DN("E")},MATCH(C4,{DN("B")},0))="","",'
             f'", "&INDEX({DN("E")},MATCH(C4,{DN("B")},0))),""))')
stt["B8"].font = font(11)
stt["B9"] = f'="Issued on "&{long_date("TODAY()")}'
stt["B9"].font = font(9, color=MUTED, italic=True)
header(stt, 11, 2, ["Date", "Amount", "How", "Campaign", "Goods given"])
NR = 24
for m in range(NR):
    r = 12 + m
    stt[f"H{r}"] = f"=IFERROR(SMALL({GF('P')},{m + 1}),\"\")"                                   # H: gift row
    g = lambda c: f"INDEX({GF1(c)},MOD(H{r},10000))"
    style(stt.cell(r, 2, f'=IF(H{r}="","",{g("B")})'), False, DATE, align="center")
    style(stt.cell(r, 3, f'=IF(H{r}="","",IF(N({g("D")})=0,"",{g("D")}))'), False, MONEY)
    for c, col in ((4, "E"), (5, "F"), (6, "G")):
        style(stt.cell(r, c, f'=IF(H{r}="","",IF({g(col)}="","",{g(col)}))'), False,
              align="center" if c == 4 else None)
stt.column_dimensions["H"].hidden = True
r = 12 + NR
stt[f"B{r}"] = f'=IF(COUNT({GF("P")})>{NR},"And "&(COUNT({GF("P")})-{NR})&" more gifts, in the total","")'
stt[f"B{r}"].font = font(9, color=MUTED, italic=True)
stt[f"B{r + 1}"] = "Total"
stt[f"B{r + 1}"].font = font(12, True)
stt[f"C{r + 1}"] = f'=IF(C4="","",SUMIFS({GF("D")},{GF("C")},C4,{GF("K")},C5,{GF("D")},">0"))'
stt[f"C{r + 1}"].number_format = MONEY
stt[f"C{r + 1}"].font = font(12, True, TEAL_D)
stt[f"D{r + 1}"] = f'=IF({CUR}="","",{CUR})'
stt[f"D{r + 1}"].font = font(11, True, MUTED)
stt[f"B{r + 3}"] = f'=IF({NOTE}="","",{NOTE})'
stt.merge_cells(f"B{r + 3}:F{r + 5}")
stt[f"B{r + 3}"].alignment = Alignment(wrap_text=True, vertical="top")
stt[f"B{r + 3}"].font = font(10)
stt.page_setup.orientation = "portrait"
stt.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Giving and volunteering", "", [3, 16, 14, 14, 14, 3, 22, 13, 13, 13, 3], rows=50)
db["B3"] = f'="The year "&{YEAR}&", with the campaigns and the thank-you notes still to send."'
db["B3"].font = font(11, color=MUTED, italic=True)
yr = f'{GF("K")},{YEAR}'
tile(db, "B", 4, "Given this year", f"=SUMIFS({GF('D')},{yr},{GF('D')},\">0\")", MONEY)
tile(db, "C", 4, "Gifts", f"=SUMIFS({GF('Q')},{yr})", "0")
tile(db, "D", 4, "Average gift", f'=IF(COUNTIFS({yr},{GF("D")},">0")=0,"",B5/COUNTIFS({yr},{GF("D")},">0"))',
     MONEY)
tile(db, "E", 4, "Donors", f'=COUNTIF({DN("N")},">0")', "0")
tile(db, "G", 4, "New donors", f'=COUNTIFS({DN("F")},">="&DATE(N({YEAR}),1,1),{DN("F")},"<"&DATE(N({YEAR})+1,1,1))',
     "0")
tile(db, "H", 4, "Lapsed", f'=COUNTIF({DN("K")},"Lapsed")', "0", "B86E1C")
tile(db, "I", 4, "To thank", f"=COUNT({GF('O')})", "0", "B23B2E")
tile(db, "J", 4, "Volunteer hours", f"=SUMIFS({HR('D')},{HR('I')},{YEAR},{HR('D')},\">0\")", "0.#")

header(db, 7, 2, ["Campaign", "Goal", "Raised", "Of goal"])
for k in range(6):
    r = 8 + k
    cr = C0 + k
    style(db.cell(r, 2, f'=IF(Campaigns!B{cr}="","",Campaigns!B{cr})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",IF(N(Campaigns!C{cr})=0,"",Campaigns!C{cr}))'), False, MONEY)
    style(db.cell(r, 4, f'=IF(B{r}="","",Campaigns!F{cr})'), False, MONEY)
    style(db.cell(r, 5, f'=IF(B{r}="","",Campaigns!G{cr})'), False, "0%", align="center")
db.conditional_formatting.add("E8:E13", FormulaRule(formula=["AND(ISNUMBER(E8),E8>=1)"], fill=fill(OK)))
header(db, 7, 7, ["To thank", "Date", "Amount", "Campaign"])
for n in range(6):
    r = 8 + n
    db[f"N{r}"] = f"=IFERROR(SMALL({GF('O')},{n + 1}),\"\")"                                    # N: gift row
    g = lambda c: f"INDEX({GF1(c)},MOD(N{r},10000))"
    style(db.cell(r, 7, f'=IF(N{r}="","",{g("C")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(N{r}="","",{g("B")})'), False, "d mmm", align="center")
    style(db.cell(r, 9, f'=IF(N{r}="","",IF(N({g("D")})=0,"Goods",{g("D")}))'), False, MONEY, align="right")
    style(db.cell(r, 10, f'=IF(N{r}="","",IF({g("F")}="","",{g("F")}))'), False, align="center")

header(db, 15, 2, ["Month", "Given", "Gifts", "Hours"])
for k in range(12):
    r = 16 + k
    key = f"({YEAR}*100+{k + 1})"
    style(db.cell(r, 2, MON_LIST.replace('"', "").split(",")[k]), False, bold=True)
    style(db.cell(r, 3, f"=SUMIFS({GF('D')},{GF('L')},{key},{GF('D')},\">0\")"), False, MONEY)
    style(db.cell(r, 4, f"=SUMIFS({GF('Q')},{GF('L')},{key})"), False, "0", align="center")
    style(db.cell(r, 5, f"=SUMIFS({HR('D')},{HR('J')},{key},{HR('D')},\">0\")"), False, "0.#", align="center")
header(db, 15, 7, ["Top donors", "This year", "Gifts", "Last gift"])
for n in range(8):
    r = 16 + n
    db[f"O{r}"] = f"=IFERROR(LARGE({DN('M')},{n + 1}),\"\")"                                    # O: donor key
    dd = lambda c: f"INDEX({DN1(c)},1000-MOD(O{r},1000))"
    style(db.cell(r, 7, f'=IF(O{r}="","",{dd("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF(O{r}="","",{dd("H")})'), False, MONEY)
    style(db.cell(r, 9, f'=IF(O{r}="","",{dd("N")})'), False, "0", align="center")
    style(db.cell(r, 10, f'=IF(O{r}="","",{dd("G")})'), False, "d mmm yy", align="center")
for c in "NO":
    db.column_dimensions[c].hidden = True
for r in list(range(8, 14)) + list(range(16, 28)):
    db.row_dimensions[r].height = 18
    for c in range(2, 11):
        cell = db.cell(r, c)
        cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="center")
chart = BarChart()
chart.type = "col"
chart.title = "Given each month"
chart.height, chart.width = 7, 14
chart.add_data(Reference(db, min_col=3, min_row=15, max_row=27), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=16, max_row=27))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "G25")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells.", [3, 6, 108])
steps = [
    ("1", "Settings: your organisation's name and address, and the line at the end of each statement."),
    ("2", "Donors and Campaigns: each donor once, and what you raise money for, with a goal if it has one."),
    ("3", "Gifts: every gift with its donor and campaign. For goods, leave the amount empty and say what was "
          "given. An x under Thanked once you have said thank you."),
    ("4", "Volunteers and Hours: each volunteer once, then each time they help, with the hours."),
    ("5", "Statement: pick a donor and a year to print the list of their gifts."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example pantry, its donors and volunteers are made up. Delete them and add your own.",
                          "The statement lists the gifts. Whether they can be claimed on taxes depends on the rules "
                          "where you are.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Gifts", "Donors", "Campaigns", "Statement", "Volunteers", "Hours", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Donation-Volunteer-Tracker.xlsx")
wb.save(out)
print("saved", out, len(gifts), "gifts", len(hours), "hours")
