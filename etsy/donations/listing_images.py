"""Etsy listing photos for the Donation and Volunteer Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Donation-Volunteer-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups: the latest gifts and hours, and the lists down to the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
g_end, h_end = last(wb["Gifts"], 1005), last(wb["Hours"], 1005)
rows = {"Gifts": list(range(6, g_end - 21)) + list(range(g_end + 1, 1006)),
        "Donors": range(last(wb["Donors"], 305) + 1, 306),
        "Campaigns": range(last(wb["Campaigns"], 25) + 1, 26),
        "Volunteers": range(last(wb["Volunteers"], 105) + 1, 106),
        "Hours": list(range(6, h_end - 21)) + list(range(h_end + 1, 1006))}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Statement"), hide_rows=rows)
close = Renders(xlsx, tmp, hide_rows={"Donors": rows["Donors"]}, hide={"Donors": ["D", "E"]})
ORG = wb["Settings"]["C4"].value


def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


def stack(top, lower):
    lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                         Image.LANCZOS)
    both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
    both.paste(top, (0, 0))
    both.paste(lower, (0, top.height + 40))
    return both


img = canvas(dark=True)
title(img, "Donation and volunteer tracker", "Donors, gifts and campaigns, a statement of each donor's gifts to "
      "print, and the hours your volunteers give", dark=True, width=2300)
laptop(img, r.shot("Giving and volunteering"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-statement.jpg"), "A statement of each donor's gifts",
       "Pick a donor and a year. It lists their gifts with the total and your closing line, ready to print or save "
       "as a PDF.", r.shot(ORG))

detail(os.path.join(img_dir, "03-gifts.jpg"), "Every gift, and who still needs a thank-you",
       "Money or goods, with the campaign. A line that needs a look is flagged, like a donor who is not on the "
       "list yet.", r.shot("Gifts"), [("#F9C9C4", "Check this line")])

detail(os.path.join(img_dir, "04-donors.jpg"), "Each donor, from the first gift to the last",
       "What they gave this year and in all, and who is new or has not given for a year",
       close.shot("Donors"), [("#D8F0DC", "New this year"), ("#FFE2B8", "Lapsed")])

detail(os.path.join(img_dir, "05-campaigns-volunteers.jpg"), "Campaigns and volunteers",
       "What each campaign has raised against its goal, and the hours each volunteer gave this year",
       stack(r.shot("Campaigns"), r.shot("Volunteers")), [("#D8F0DC", "Goal reached")])

img = canvas()
y = title(img, "Inside the file", "Seven tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Giving and volunteering")), "Dashboard", "The whole year"),
    (content(r.page("Gifts")), "Gifts", "Money and goods"),
    (content(r.page("Donors")), "Donors", "New, active, lapsed"),
    (content(r.page(ORG)), "Statement", "To print for a donor"),
    (content(r.page("Campaigns")), "Campaigns", "Raised against goal"),
    (content(r.page("Hours")), "Hours", "Each time someone helps"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
