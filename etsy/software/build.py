"""Software and Tools Tracker for a small business (Excel + Google Sheets).

Every tool the business pays for: the plan, who looks after it, the card it goes on, what it costs a month and a year,
seats bought and seats used. The next charge moves on by itself, a notice period gives the last day to cancel, and a
card that runs out before a renewal is flagged. The Access tab says who uses what, so when someone leaves, the tools to
remove them from are listed. A Year ahead tab spreads every charge over the next twelve months. Only functions both
apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
T0, T1 = 6, 105        # tools
A0, A1 = 6, 405        # access lines
M0, M1 = 6, 105        # team
C0, C1 = 6, 25         # cards
K0, K1 = 10, 21        # categories (Settings)
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
TL = lambda col: f"Tools!${col}${T0}:${col}${T1}"
AC = lambda col: f"Access!${col}${A0}:${col}${A1}"
TM = lambda col: f"Team!${col}${M0}:${col}${M1}"
CD = lambda col: f"Cards!${col}${C0}:${col}${C1}"
CUR, WARN_DAYS, CARD_DAYS = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6"
RED = "B4541F"


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


D = lambda n: today + dt.timedelta(days=n)

# ---------------------------------------------------------------- example data (a made-up small business)
categories = ["Office", "Communication", "Design", "Finance", "Sales", "Marketing", "Website", "Operations", "Security",
              "Other"]
team = [
    # name, team, started, left on
    ("Ana", "Management", dt.date(2019, 3, 1), None), ("Ben", "Design", dt.date(2020, 6, 15), None),
    ("Chloe", "Design", dt.date(2023, 2, 1), None), ("Dev", "Tech", dt.date(2021, 9, 1), None),
    ("Ella", "Marketing", dt.date(2022, 4, 4), None), ("Finn", "Sales", dt.date(2024, 1, 8), None),
    ("Grace", "Operations", dt.date(2020, 11, 2), None), ("Hugo", "Sales", dt.date(2023, 5, 15), D(-12)),
]
cards = [("Business card A", "Ana", dt.date(2028, 6, 30)), ("Business card B", "Grace", dt.date(2027, 1, 31)),
         ("Debit card C", "Dev", D(63))]
tools = [
    # tool, category, owner, billing, price, per seat, seats bought, paid with, renews on, notice, cancelled on
    ("Email and docs", "Office", "Ana", "Monthly", 12, "x", 8, "Business card A", D(5), None, None),
    ("Team chat", "Communication", "Grace", "Monthly", 8.75, "x", 10, "Business card A", D(-16), None, None),
    ("Video calls", "Communication", "Grace", "Yearly", 149.90, None, None, "Business card B", D(126), None, None),
    ("Design suite", "Design", "Ben", "Yearly", 659.88, "x", 3, "Business card B", D(22), 14, None),
    ("Stock photos", "Design", "Chloe", "Monthly", 29, None, None, "Business card A", D(10), None, None),
    ("Accounting", "Finance", "Ana", "Monthly", 55, None, None, "Business card A", D(3), None, None),
    ("Website builder", "Website", "Ella", "Yearly", 276, None, None, "Debit card C", D(168), 30, None),
    ("Domain names", "Website", "Dev", "Yearly", 38, None, None, "Debit card C", D(65), None, None),
    ("CRM", "Sales", "Finn", "Monthly", 49, "x", 3, "Business card A", D(17), None, None),
    ("Newsletter", "Marketing", "Ella", "Monthly", 45, None, None, "Business card A", D(11), None, None),
    ("Project board", "Operations", "Grace", "Monthly", 10, "x", 8, "Business card A", D(20), None, None),
    ("Password manager", "Security", "Dev", "Yearly", 48, "x", 8, "Business card B", D(215), None, None),
    ("Cloud storage", "Office", "Dev", "Yearly", 120, None, None, "Business card A", D(43), None, None),
    ("Scheduling", "Operations", "Grace", "Quarterly", 36, None, None, "Business card A", D(-80), None, None),
    ("E-signature", "Operations", "Ana", "Monthly", 25, None, None, "Business card A", D(-13), None, D(-13)),
    ("AI writing assistant", "Marketing", "Ella", "Monthly", 20, "x", 2, "Business card A", D(7), None, D(33)),
]
everyone = [t[0] for t in team]
uses = {"Email and docs": everyone, "Team chat": everyone, "Video calls": ["Ana", "Grace", "Finn"],
        "Design suite": ["Ben", "Chloe"], "Stock photos": ["Chloe"], "Accounting": ["Ana", "Grace"],
        "Website builder": ["Ella", "Dev"], "Domain names": ["Dev"], "CRM": ["Finn", "Ana", "Ella", "Hugo"],
        "Newsletter": ["Ella"], "Project board": ["Ana", "Ben", "Chloe", "Dev", "Ella", "Finn", "Grace"],
        "Password manager": everyone, "Cloud storage": ["Dev", "Ana"], "Scheduling": ["Grace", "Finn"],
        "E-signature": ["Ana"], "AI writing assistant": ["Ella", "Finn"]}
admins = {"Email and docs": "Ana", "Team chat": "Grace", "Video calls": "Grace", "Design suite": "Ben",
          "Password manager": "Dev", "CRM": "Finn", "Project board": "Grace", "Website builder": "Ella"}
access = []
for i, (tool, people) in enumerate(uses.items()):
    adopted = dt.date(2023 + i % 3, 1 + i * 5 % 12, 1 + i * 7 % 27)          # when the business took the tool on
    for p in people:
        started = next(t[2] for t in team if t[0] == p)
        added = max(started, adopted)
        removed = D(-11) if p == "Hugo" and tool in ("Email and docs", "CRM") else (
            D(-13) if tool == "E-signature" else None)
        access.append((p, tool, "Admin" if admins.get(tool) == p else "User", added, removed))

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def status_fills(ws, rng, first, pairs, how="="):
    for text, colour in pairs:
        rule = f'{first}="{text}"' if how == "=" else f'LEFT({first},{len(text)})="{text}"'
        ws.conditional_formatting.add(rng, FormulaRule(formula=[rule], fill=fill(colour)))


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The currency, how early you want a warning, and the categories.", [3, 34, 16, 4, 60],
           rows=K1 + 3)
for r, label, value, fmt in ((4, "Currency", "USD", None), (5, "Warn me before a yearly charge (days)", 30, "0"),
                             (6, "Warn me before a card runs out (days)", 60, "0")):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
for k, text in enumerate(["Amounts have no currency sign, so any currency works.",
                          "Monthly charges are not warned about: they come every month.",
                          "A card warns you this many days before it expires."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)
header(se, K0 - 1, 2, ["Categories"])
for r in range(K0, K1 + 1):
    style(se.cell(r, 2, categories[r - K0] if r - K0 < len(categories) else None), True)

# ---------------------------------------------------------------- Team
tm = wb.create_sheet("Team")
sheet_base(tm, "Team", "Everyone who uses the tools. When someone leaves, type the day: the Access tab lists what to "
           "remove them from.", [3, 18, 16, 13, 13, 9, 13, 24], rows=M1 + 2)
header(tm, 5, 2, ["Name", "Team", "Started", "Left on", "Tools", "Seats a month", "Status"])
for r in range(M0, M1 + 1):
    x = team[r - M0] if r - M0 < len(team) else (None,) * 4
    style(tm.cell(r, 2, x[0]), True, bold=True)
    style(tm.cell(r, 3, x[1]), True)
    style(tm.cell(r, 4, x[2]), True, DATE, align="center")
    style(tm.cell(r, 5, x[3]), True, DATE, align="center")
    style(tm.cell(r, 6, f'=IF(B{r}="","",SUMIFS({AC("G")},{AC("B")},B{r}))'), False, "0", align="center")
    style(tm.cell(r, 7, f'=IF(B{r}="","",SUMIFS({AC("J")},{AC("B")},B{r}))'), False, MONEY)
    style(tm.cell(r, 8, f'=IF(B{r}="","",IF(AND(E{r}<>"",E{r}<=TODAY()),IF(F{r}>0,"Left, still has access","Left"),'
                        f'"Active"))'), False, align="center")
status_fills(tm, f"H{M0}:H{M1}", f"H{M0}", (("Left, still has access", BAD), ("Left", INFO)))
tm.freeze_panes = "C6"

# ---------------------------------------------------------------- Cards
cd = wb.create_sheet("Cards")
sheet_base(cd, "Cards", "The cards and accounts the tools are paid with. Type the last day of the month it expires.",
           [3, 22, 16, 13, 9, 12, 22], rows=C1 + 2)
header(cd, 5, 2, ["Card or account", "Held by", "Expires", "Tools", "Per month", "Status"])
for r in range(C0, C1 + 1):
    x = cards[r - C0] if r - C0 < len(cards) else (None,) * 3
    style(cd.cell(r, 2, x[0]), True, bold=True)
    style(cd.cell(r, 3, x[1]), True)
    style(cd.cell(r, 4, x[2]), True, DATE, align="center")
    style(cd.cell(r, 5, f'=IF(B{r}="","",COUNTIFS({TL("I")},B{r},{TL("T")},"<>Cancelled"))'), False, "0", align="center")
    style(cd.cell(r, 6, f'=IF(B{r}="","",SUMIFS({TL("O")},{TL("I")},B{r}))'), False, MONEY)
    late = f'COUNTIFS({TL("I")},B{r},{TL("Q")},">"&D{r})'
    style(cd.cell(r, 7, f'=IF(B{r}="","",IF(D{r}="","No date",IF(D{r}<TODAY(),"Expired",IF({late}>0,"Update "&{late}&'
                        f'IF({late}=1," tool"," tools"),IF(D{r}-TODAY()<=N({CARD_DAYS}),"Expires soon","OK")))))'),
          False, align="center")
status_fills(cd, f"G{C0}:G{C1}", f"G{C0}", (("Expired", BAD), ("Update", BAD), ("Expires soon", WARN), ("OK", OK)),
             how="left")

# ---------------------------------------------------------------- Tools
tl = wb.create_sheet("Tools")
sheet_base(tl, "Tools", "Every tool the business pays for. The next charge moves on by itself after each renewal.",
           [3, 20, 14, 10, 11, 10, 7, 8, 17, 12, 8, 12, 7, 10, 10, 11, 12, 12, 7, 22], rows=T1 + 2)
header(tl, 5, 2, ["Tool", "Category", "Owner", "Billing", "Price", "Per seat", "Seats", "Paid with", "Renews on",
                  "Notice (days)", "Cancelled on", "Seats used", "Each charge", "Per month", "Per year", "Next charge",
                  "Cancel by", "Unused", "Status"])
for r in range(T0, T1 + 1):
    x = tools[r - T0] if r - T0 < len(tools) else (None,) * 11
    fmts = (None, None, None, None, MONEY, None, "0", None, DATE, "0", DATE)
    aligns = (None, None, None, "center", None, "center", "center", None, "center", "center", "center")
    for c, (v, fmt, align) in enumerate(zip(x, fmts, aligns)):
        style(tl.cell(r, 2 + c, v), True, fmt, c == 0, align)
    months = f"U{r}"
    tl[months] = f'=IF(E{r}="Quarterly",3,IF(E{r}="Yearly",12,1))'                              # U: months a charge
    ended = f'AND(L{r}<>"",L{r}<=TODAY())'
    style(tl.cell(r, 13, f'=IF(B{r}="","",SUMIFS({AC("G")},{AC("C")},B{r}))'), False, "0", align="center")
    style(tl.cell(r, 14, f'=IF(B{r}="","",IF(G{r}<>"",N(F{r})*IF(H{r}="",M{r},N(H{r})),N(F{r})))'), False, MONEY)
    style(tl.cell(r, 15, f'=IF(B{r}="","",IF({ended},0,N{r}/{months}))'), False, MONEY)
    style(tl.cell(r, 16, f'=IF(B{r}="","",O{r}*12)'), False, MONEY)
    n = f'INT(DATEDIF(J{r},TODAY(),"m")/{months})*{months}'
    roll = f'IF(J{r}>=TODAY(),J{r},IF(EDATE(J{r},{n})>=TODAY(),EDATE(J{r},{n}),EDATE(J{r},{n}+{months})))'
    style(tl.cell(r, 17, f'=IF(OR(B{r}="",J{r}=""),"",IF(AND(L{r}<>"",{roll}>=L{r}),"",{roll}))'), False, DATE,
          align="center")
    style(tl.cell(r, 18, f'=IF(Q{r}="","",Q{r}-N(K{r}))'), False, DATE, align="center")
    style(tl.cell(r, 19, f'=IF(OR(B{r}="",G{r}="",H{r}="",{ended}),"",MAX(0,N(H{r})-N(M{r})))'), False, "0;;",
          align="center")
    tl[f"Y{r}"] = f'=IF(N(S{r})=0,0,S{r}*N(F{r})*12/{months})'                                # Y: unused, a year
    card = f'IFERROR(IF(INDEX({CD("D")},MATCH(I{r},{CD("B")},0))="","",INDEX({CD("D")},MATCH(I{r},{CD("B")},0))),"")'
    tl[f"W{r}"] = f"={card}"                                                                     # W: card expiry
    big = f'E{r}<>"Monthly"'
    style(tl.cell(r, 20, f'=IF(B{r}="","",IF({ended},"Cancelled",IF(L{r}<>"","Ends "&{short(f"L{r}")},'
                         f'IF(AND(Q{r}<>"",W{r}<>"",W{r}<Q{r}),"Card expires",'
                         f'IF(AND(Q{r}<>"",{big},N(K{r})>0,R{r}>=TODAY(),R{r}-TODAY()<=N({WARN_DAYS})),'
                         f'"Cancel by "&{short(f"R{r}")},IF(AND(Q{r}<>"",{big},Q{r}-TODAY()<=N({WARN_DAYS})),'
                         f'"Renews "&{short(f"Q{r}")},IF(N(S{r})>0,S{r}&IF(S{r}=1," unused seat"," unused seats"),'
                         f'"Active")))))))'), False, align="center")
    tl[f"V{r}"] = f'=IF(Q{r}="","",Q{r}+ROW()/10000000)'                                       # V: coming up
    tl[f"X{r}"] = f'=IF(OR(T{r}="",T{r}="Active",T{r}="Cancelled"),"",ROW())'                  # X: needs a look
for c in "UVWXY":
    tl.column_dimensions[c].hidden = True
dropdown(tl, f"=Settings!$B${K0}:$B${K1}", f"C{T0}:C{T1}", strict=False)
dropdown(tl, f"={TM('B')}", f"D{T0}:D{T1}", strict=False)
dropdown(tl, '"Monthly,Quarterly,Yearly"', f"E{T0}:E{T1}")
dropdown(tl, f"={CD('B')}", f"I{T0}:I{T1}")
status_fills(tl, f"T{T0}:T{T1}", f"T{T0}", (("Cancelled", INFO), ("Ends", INFO), ("Card expires", BAD),
                                           ("Cancel by", BAD), ("Renews", WARN), ("Active", OK)), how="left")
tl.conditional_formatting.add(f"T{T0}:T{T1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("unused",T{T0}))'], fill=fill(WARN)))
tl.conditional_formatting.add(f"B{T0}:S{T1}", FormulaRule(formula=[f'$T{T0}="Cancelled"'],
                                                          font=Font(name=F, italic=True, color=MUTED)))
tl.freeze_panes = "C6"

# ---------------------------------------------------------------- Access (who uses what)
ac = wb.create_sheet("Access")
sheet_base(ac, "Access", "One line per person and tool. Each line is a seat until the day it is removed.",
           [3, 16, 20, 10, 13, 13, 3, 16], rows=A1 + 2)
header(ac, 5, 2, ["Person", "Tool", "Role", "Added on", "Removed on"])
ac["H5"] = "Status"
header(ac, 5, 8, ["Status"])
for r in range(A0, A1 + 1):
    x = access[r - A0] if r - A0 < len(access) else (None,) * 5
    style(ac.cell(r, 2, x[0]), True, bold=True)
    style(ac.cell(r, 3, x[1]), True)
    style(ac.cell(r, 4, x[2]), True, align="center")
    style(ac.cell(r, 5, x[3]), True, DATE, align="center")
    style(ac.cell(r, 6, x[4]), True, DATE, align="center")
    empty = f'OR(B{r}="",C{r}="")'
    ac[f"G{r}"] = f'=IF({empty},"",IF(AND(F{r}<>"",F{r}<=TODAY()),0,1))'                  # G: holds a seat
    left = f'IFERROR(IF(INDEX({TM("E")},MATCH(B{r},{TM("B")},0))="","",INDEX({TM("E")},MATCH(B{r},{TM("B")},0))),"")'
    ac[f"I{r}"] = f"={left}"                                                                  # I: the day they left
    style(ac.cell(r, 8, f'=IF({empty},"",IF(G{r}=0,"Removed",IF(AND(I{r}<>"",I{r}<=TODAY()),"Remove now","Active")))'),
          False, align="center")
    tool = f'MATCH(C{r},{TL("B")},0)'
    seat = (f'IFERROR(IF(INDEX({TL("G")},{tool})="",0,IF(N(INDEX({TL("O")},{tool}))=0,0,'
            f'N(INDEX({TL("F")},{tool}))/INDEX({TL("U")},{tool}))),0)')
    ac[f"J{r}"] = f'=IF(G{r}<>1,0,{seat})'                                                   # J: seat cost a month
    ac[f"K{r}"] = f'=IF(H{r}="Remove now",ROW(),"")'                                        # K: to remove
for c in "GIJK":
    ac.column_dimensions[c].hidden = True
dropdown(ac, f"={TM('B')}", f"B{A0}:B{A1}")
dropdown(ac, f"={TL('B')}", f"C{A0}:C{A1}")
dropdown(ac, '"Admin,User"', f"D{A0}:D{A1}", strict=False)
status_fills(ac, f"H{A0}:H{A1}", f"H{A0}", (("Remove now", BAD), ("Removed", INFO)))
ac.freeze_panes = "C6"

# ---------------------------------------------------------------- Year ahead (every charge, month by month)
ya = wb.create_sheet("Year ahead")
sheet_base(ya, "Year ahead", "What each tool will charge in each of the next twelve months, from today.",
           [3, 20] + [9] * 12 + [11], rows=T1 + 3)
month = lambda j: f"DATE(YEAR(TODAY()),MONTH(TODAY())+{j},1)"
header(ya, 5, 2, ["Tool"] + [""] * 12 + ["Total"])
for j in range(12):
    ya.cell(5, 3 + j, f'=CHOOSE(MONTH({month(j)}),{MON_LIST})&" "&RIGHT(YEAR({month(j)}),2)')
style(ya.cell(6, 2, "Total"), False, bold=True)
for j in range(13):
    col = 3 + j
    L = ya.cell(6, col).column_letter
    style(ya.cell(6, col, f"=SUM({L}7:{L}{T1 + 1})"), False, "#,##0;;", True)
for c in range(2, 16):
    ya.cell(6, c).fill = fill("E8F0EF")
for r in range(T0 + 1, T1 + 2):
    t = r - 1
    style(ya.cell(r, 2, f'=IF(Tools!B{t}="","",Tools!B{t})'), False, bold=True)
    for j in range(12):
        m = month(j)
        diff = f"((YEAR({m})-YEAR(Tools!$Q{t}))*12+MONTH({m})-MONTH(Tools!$Q{t}))"
        style(ya.cell(r, 3 + j, f'=IF(OR(Tools!$B{t}="",Tools!$Q{t}=""),"",IF(AND({diff}>=0,MOD({diff},Tools!$U{t})=0,'
                                f'OR(Tools!$L{t}="",EDATE(Tools!$Q{t},{diff})<Tools!$L{t})),Tools!$N{t},""))'), False,
              "#,##0;;")
    style(ya.cell(r, 15, f'=IF(B{r}="","",SUM(C{r}:N{r}))'), False, "#,##0;;", True)
ya.freeze_panes = "C7"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Software and tools", "What the tools cost on the left, what needs a look on the right.",
           [3, 20, 14, 14, 14, 3, 20, 20, 14, 12, 3], rows=52)
active_people = f'COUNTIF({TM("H")},"Active")'
tile(db, "B", 4, "Per month", f"=SUM({TL('O')})", MONEY)
tile(db, "C", 4, "Per year", f"=SUM({TL('P')})", MONEY)
tile(db, "D", 4, "Tools in use", f'=COUNTIFS({TL("B")},"<>",{TL("T")},"<>Cancelled")', "0")
tile(db, "E", 4, "Per person a month", f'=IF({active_people}=0,"",SUM({TL("O")})/{active_people})', MONEY)
due = f'SUMIFS({TL("N")},{TL("Q")},">="&TODAY(),{TL("Q")},"<="&(TODAY()+N({WARN_DAYS})))'
tile(db, "G", 4, "", f"={due}", MONEY, "B07A00")
db["G4"] = f'="Due in the next "&N({WARN_DAYS})&" days"'
tile(db, "H", 4, "Unused seats", f"=SUM({TL('S')})", "0", "B07A00")
tile(db, "I", 4, "Unused a year", f"=SUM({TL('Y')})", MONEY, "B07A00")
tile(db, "J", 4, "To remove", f'=COUNTIF({AC("H")},"Remove now")', "0", RED)

TOP, NTOP = 9, 10
LOW = TOP + NTOP + 2                  # header row of the lower tables
header(db, 8, 2, ["Coming up", "Charge on", "Amount", "Cancel by"])
for i in range(NTOP):
    r = TOP + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({TL("V")},{i + 1}),"")'
    row = f'MATCH({key},{TL("V")},0)'
    at = lambda col: f"INDEX({TL(col)},{row})"
    none = '"Nothing to pay"' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},{at("B")})'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 4, f'=IF({key}="","",{at("N")})'), False, MONEY)
    style(db.cell(r, 5, f'=IF({key}="","",IF(N({at("K")})=0,"",{at("R")}))'), False, "ddd d mmm", align="center")

header(db, 8, 7, ["Needs a look", "Status", "Owner", "Per month"])
for i in range(NTOP):
    r = TOP + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({TL("X")},{i + 1}),"")'
    at = lambda col: f"INDEX({TL(col)},{key}-{T0 - 1})"
    none = '"Nothing needs a look"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},{at("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",{at("T")})'), False, align="center")
    style(db.cell(r, 9, f'=IF({key}="","",IF({at("D")}="","",{at("D")}))'), False)
    style(db.cell(r, 10, f'=IF({key}="","",{at("O")})'), False, "#,##0.00;;")
looks = f"H{TOP}:H{TOP + NTOP - 1}"
status_fills(db, looks, f"H{TOP}", (("Card expires", BAD), ("Cancel by", BAD), ("Renews", WARN), ("Ends", INFO)),
             how="left")
db.conditional_formatting.add(looks, FormulaRule(formula=[f'ISNUMBER(SEARCH("unused",H{TOP}))'], fill=fill(WARN)))

header(db, LOW, 2, ["By category", "Tools", "Per month", "Per year"])
for k in range(K1 - K0 + 1):
    r = LOW + 1 + k
    cat = f"Settings!$B${K0 + k}"
    style(db.cell(r, 2, f'=IF({cat}="","",{cat})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",COUNTIFS({TL("C")},B{r},{TL("T")},"<>Cancelled"))'), False, "0;;", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",SUMIFS({TL("O")},{TL("C")},B{r}))'), False, "#,##0.00;;")
    style(db.cell(r, 5, f'=IF(B{r}="","",D{r}*12)'), False, "#,##0.00;;")
other = LOW + 1 + K1 - K0 + 1
style(db.cell(other, 4, f"=SUM({TL('O')})-SUM(D{LOW + 1}:D{other - 1})"), False, "#,##0.00;-#,##0.00;;")
style(db.cell(other, 5, f"=D{other}*12"), False, "#,##0.00;-#,##0.00;;")
style(db.cell(other, 2, f'=IF(ROUND(D{other},6)=0,"","Not on the list")'), False, bold=True)
style(db.cell(other, 3), False)

header(db, LOW, 7, ["Access to remove", "Tool", "Left on", "Role"])
for i in range(8):
    r = LOW + 1 + i
    key = f"Q{r}"
    db[key] = f'=IFERROR(SMALL({AC("K")},{i + 1}),"")'
    at = lambda col: f"INDEX({AC(col)},{key}-{A0 - 1})"
    none = '"Nobody to remove"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},{at("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",{at("C")})'), False)
    style(db.cell(r, 9, f'=IF({key}="","",{at("I")})'), False, "d mmm yyyy", align="center")
    style(db.cell(r, 10, f'=IF({key}="","",IF({at("D")}="","",{at("D")}))'), False, align="center")
db.conditional_formatting.add(f"G{LOW + 1}:J{LOW + 8}", FormulaRule(formula=[f'$Q{LOW + 1}<>""'], fill=fill("FBE3E0")))

chart = BarChart()
chart.type = "col"
chart.title = "Charges over the next 12 months"
chart.height, chart.width = 6.5, 24
chart.add_data(Reference(ya, min_col=3, max_col=14, min_row=6, max_row=6), from_rows=True, titles_from_data=False)
chart.set_categories(Reference(ya, min_col=3, max_col=14, min_row=5, max_row=5))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.legend = None
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, f"B{other + 3}")
for c in "OPQ":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: the currency and how early you want a warning. Cards: each card or account and when it expires."),
    ("2", "Team: everyone who uses the tools, and the day someone leaves."),
    ("3", "Tools: each tool with its price, billing, card and the day it renews. Per seat: an x and how many seats."),
    ("4", "Access: one line per person and tool. The seats used and the leavers to remove come from it."),
    ("5", "Dashboard and Year ahead: what it all costs, what renews soon, and every charge month by month."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example business, its team and tools are made up. Delete them and add your own.",
                          "Keep passwords out of this file: a password manager is the place for them.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Tools", "Access", "Team", "Cards", "Year ahead", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Software-Tools-Tracker.xlsx")
wb.save(out)
print("saved", out, len(tools), "tools", len(access), "access lines", len(team), "people")
