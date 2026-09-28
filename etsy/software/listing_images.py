"""Etsy listing photos for the Software and Tools Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Software-Tools-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
tools_end, team_end, cards_end = last(wb["Tools"], 105), last(wb["Team"], 105), last(wb["Cards"], 25)

r = Renders(xlsx, tmp, one_page=("Dashboard",))
close = Renders(xlsx, tmp, one_page=("Dashboard",), hide={"Tools": ["K", "L", "N", "P", "R"]},
                hide_rows={"Tools": range(tools_end + 1, 106), "Team": range(team_end + 1, 106),
                           "Cards": range(cards_end + 1, 26), "Year ahead": range(tools_end + 2, 107)})

img = canvas(dark=True)
title(img, "Software and tools tracker", "What every tool costs, who uses it, what renews soon, and what to remove "
      "when someone leaves", dark=True, width=2300)
laptop(img, r.shot("Software and tools"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-tools.jpg"), "Every tool the business pays for",
       "Monthly, quarterly or yearly, per seat or flat. The next charge moves on by itself, with the last day to cancel.",
       close.shot("Tools"), [("#F9C9C4", "Card or last day to cancel"), ("#FFE2B8", "Renews soon or unused seats"),
                             ("#D8F0DC", "Active")])

detail(os.path.join(img_dir, "03-leavers.jpg"), "When someone leaves",
       "Type the day they leave. Every tool they can still get into is listed, and their seats stop being counted once "
       "removed.", close.shot("Team"), [("#F9C9C4", "Left, still has access"), ("#E3E8FF", "Left")])

detail(os.path.join(img_dir, "04-year-ahead.jpg"), "Every charge for the next twelve months",
       "Monthly, quarterly and yearly plans spread over the months they are paid in", close.shot("Year ahead"))

detail(os.path.join(img_dir, "05-cards.jpg"), "Cards that run out before a renewal",
       "Each card with the tools on it and what they cost. A card that expires before a charge is flagged.",
       close.shot("Cards"), [("#F9C9C4", "Update the card"), ("#FFE2B8", "Expires soon"), ("#D8F0DC", "OK")])

img = canvas()
y = title(img, "Inside the file", "Six tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Software and tools")), "Dashboard", "Costs and what needs a look"),
    (content(r.page("Tools")), "Tools", "Plans, seats and renewals"),
    (content(r.page("Access")), "Access", "Who uses what"),
    (content(r.page("Team")), "Team", "Starters and leavers"),
    (content(r.page("Cards")), "Cards", "What each card pays"),
    (content(r.page("Year ahead")), "Year ahead", "Month by month"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
