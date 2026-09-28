"""Etsy listing photos for the Personal Loan Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Personal-Loan-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, end, col: max(r for r in range(6, end + 1) if ws.cell(r, col).value not in (None, ""))
rows = {"Loans": range(last(wb["Loans"], 305, 3) + 1, 306), "Payments": range(last(wb["Payments"], 805, 2) + 1, 806)}
r = Renders(xlsx, tmp, one_page=("Dashboard", "Statement"), hide_rows=rows)



def band(img):
    """Height in pixels of the first teal header row, to match the size of type of two sheets."""
    a = np.asarray(img).astype(int)
    on = (np.abs(a - np.array([0x1F, 0x6F, 0x78])).sum(axis=2) < 40).mean(axis=1) > 0.3
    start = int(np.argmax(on))
    return int(np.argmin(on[start:]))


img = canvas(dark=True)
title(img, "Money lent and borrowed", "Loans to and from friends and family: who owes what, the next payments "
      "and what is overdue", dark=True, width=2300)
laptop(img, r.shot("Money lent and borrowed"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-statement.jpg"), "A statement for any loan",
       "Print it or share it, with a friendly reminder ready to send", r.shot("Loan statement"))

# Loans over the payments, at the same size of type.
top, lower = r.shot("Loans"), r.shot("Payments")
lower = lower.resize((round(lower.width * band(top) / band(lower)), round(lower.height * band(top) / band(lower))),
                     Image.LANCZOS)
both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
both.paste(top, (0, 0))
both.paste(lower, (0, top.height + 40))
detail(os.path.join(img_dir, "03-loans.jpg"), "Every loan with its plan",
       "Payments every week, two weeks or month, simple interest if you want it, and each payment as it happens",
       both, [("#F9C9C4", "Overdue"), ("#FFE2B8", "Due soon")])

# The top of the Dashboard, down to the monthly table.
page = r.doc[r._index("Money lent and borrowed")]
month = next(w for w in page.get_text("words") if w[4] == "Month")            # the header, not "this month"
cut = int((month[1] - 14) * 220 / 72)
full = r.page("Money lent and borrowed")
detail(os.path.join(img_dir, "04-who-owes.jpg"), "Who owes you, and what you owe",
       "What is owed each way, what is overdue, the next payments and each person's balance",
       content(full.crop((0, 0, full.width, cut))), [("#F9C9C4", "Overdue"), ("#FFE2B8", "Due soon")])

img = canvas()
y = title(img, "Inside the file", "Four tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Money lent and borrowed")), "Dashboard", "Who owes what"),
    (content(r.page("Loan statement")), "Statement", "To print or share"),
    (content(r.page("Loans")), "Loans", "Each loan and its plan"),
    (content(r.page("Payments")), "Payments", "In and out"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "05-whats-inside.jpg"))
