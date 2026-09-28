"""Etsy listing photos for the Timesheet and Invoice Maker (run build.py first)."""
import datetime as dt
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Timesheet-Invoice-Maker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Which rows to show close up: the end of last month and the start of this one, the invoices there are,
# and the lines of the invoice on the Invoice tab.
wb = load_workbook(xlsx)
ts, iv = wb["Timesheet"], wb["Invoices"]
day = lambda v: v.date() if isinstance(v, dt.datetime) else v
dates = {r: day(ts.cell(r, 2).value) for r in range(6, 2006) if ts.cell(r, 2).value}
first = dt.date.today().replace(day=1)
keep = [r for r, d in dates.items() if first - dt.timedelta(days=10) <= d <= first + dt.timedelta(days=9)]
last_invoice = max(r for r in range(6, 306) if iv.cell(r, 2).value)
chosen = wb["Invoice"]["G4"].value
row = next(r for r in range(6, 306) if iv.cell(r, 2).value == chosen)
client, frm, to = iv.cell(row, 3).value, day(iv.cell(row, 5).value), day(iv.cell(row, 6).value)
on_it = sum(1 for r, d in dates.items() if ts.cell(r, 3).value == client and frm <= d <= to and not ts.cell(r, 10).value)

r = Renders(xlsx, tmp)
close = Renders(xlsx, tmp, hide={"Timesheet": ["L"]},
                hide_rows={"Timesheet": [x for x in range(6, 2006) if x not in keep],
                           "Invoices": range(last_invoice + 1, 306),
                           "Hours report": range(7 + on_it, 52)})

img = canvas(dark=True)
title(img, "Timesheet and invoice maker", "Hours by client and project, invoices that fill themselves, and what is "
      "still to be paid", dark=True, width=2300)
laptop(img, r.shot("Hours and invoices"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-invoice.jpg"), "The invoice\nfills itself",
       "Pick its number: the client, the hours of those dates, the total and the day to pay are all there. Print it "
       "or save it as a PDF.", r.shot("Rivera"))

detail(os.path.join(img_dir, "03-timesheet.jpg"), "One line per piece of work",
       "Start and end times, or the hours typed in, at the client's rate, the project's or one of its own",
       close.shot("Timesheet"), [("#FFE2B8", "To invoice"), ("#E3E8FF", "Invoiced"), ("#D8F0DC", "Paid")])

detail(os.path.join(img_dir, "04-invoices.jpg"), "Who has paid, and who is late",
       "Each invoice adds up the work of its client and dates by itself. Type the day it is paid.",
       close.shot("Invoices"), [("#D8F0DC", "Paid"), ("#FFE2B8", "Waiting"), ("#F9C9C4", "Overdue")])

detail(os.path.join(img_dir, "05-hours-report.jpg"), "Every hour, if the client asks",
       "The Hours report lists each line of the same invoice with its times, ready to send with it",
       close.shot("Hours for invoice"))

img = canvas()
y = title(img, "Inside the file", "Seven tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Hours and invoices")), "Dashboard", "The week, the month, the money"),
    (content(r.page("Timesheet")), "Timesheet", "Every piece of work"),
    (content(r.page("Invoices")), "Invoices", "Paid, waiting, overdue"),
    (content(r.page("Rivera")), "Invoice", "The page you send"),
    (content(r.page("Clients")), "Clients", "Rates and days to pay"),
    (content(r.page("Projects")), "Projects", "Own rates and budgets"),
], y + 80, cols=3)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
