"""Timesheet and Invoice Maker (Excel + Google Sheets).

Hours by client and project, from start and end times or typed in, at the client's rate, the project's or one of
the line's own. Each invoice is one line on the Invoices tab: its number, the client and the dates of the work. The
lines of work in those dates go on it by themselves; the Invoice tab shows it as a page to print or save as a PDF,
and the Hours report lists every line for a client who wants the detail. The Dashboard has the hours of the week and
the month, what is still to invoice, and what is waiting or overdue. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
K0, K1 = 6, 45         # clients
P0, P1 = 6, 65         # projects
T0, T1 = 6, 2005       # timesheet
V0, V1 = 6, 305        # invoices
I0, I1 = 17, 26        # lines on the invoice (one per kind of work)
H0, H1 = 7, 51         # lines on the hours report
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MON_LIST = ",".join(f'"{m[:3]}"' for m in MONTHS)
FULL_LIST = ",".join(f'"{m}"' for m in MONTHS)
CL = lambda col: f"Clients!${col}${K0}:${col}${K1}"
PR = lambda col: f"Projects!${col}${P0}:${col}${P1}"
TS = lambda col: f"Timesheet!${col}${T0}:${col}${T1}"
IV = lambda col: f"Invoices!${col}${V0}:${col}${V1}"
S = lambda row: f"Settings!$C${row}"
CUR, TAXNAME, TAX, TERMS = S(10), S(11), S(12), S(13)
CHOSEN = "Invoice!$G$4"


def long_date(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{FULL_LIST})&" "&YEAR({d})'


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def month_start(d, n):
    m = d.month - 1 + n
    return dt.date(d.year + m // 12, m % 12 + 1, 1)


# ---------------------------------------------------------------- example data (a made-up studio and clients)
me = [("Your name or business", "Rivera Design Studio"), ("Address", "4 Canal Lane"), ("Town and postcode", "Riverton RT2 9QP"),
      ("Email", "hello@riveradesign.example"), ("Phone", "+44 7700 900123"), ("Tax number", None), ("Currency", "GBP"),
      ("Tax name", "VAT"), ("Tax rate", 0), ("Days to pay", 14), ("Bank", "Account name: Rivera Design Studio"),
      ("Bank details", "Sort code 00-00-00, account 12345678"), ("Note on invoices", "Thank you for your business.")]
clients = [
    # client, contact, email, address, town, rate, days to pay
    ("Harbour Books", "Priya Shah", "accounts@harbourbooks.example", "18 Quay Street", "Riverton RT1 4BB", 75, 30),
    ("Greenleaf Studio", "Tom Ellis", "tom@greenleaf.example", "Unit 5, Mill Yard", "Eastbridge EB3 2LD", 55, None),
    ("Atlas Fitness", "Maria Lopez", "maria@atlasfitness.example", "221 High Road", "Westford WF9 1AA", 65, 14),
]
first_this = today.replace(day=1)
last_month = month_start(today, -1)
m0 = month_start(today, -5)
projects = [
    # project, client, own rate, budget hours, first day, last day
    ("Spring catalogue", "Harbour Books", None, 140, m0, month_start(today, -3) - dt.timedelta(days=1)),
    ("Website redesign", "Harbour Books", None, 260, month_start(today, -3), today),
    ("Book launch posters", "Harbour Books", 85, 18, last_month, last_month + dt.timedelta(days=24)),
    ("Monthly social posts", "Greenleaf Studio", None, None, m0, today),
    ("Brand refresh", "Atlas Fitness", 70, 60, last_month + dt.timedelta(days=9), today),
]
weights = {"Spring catalogue": 3, "Website redesign": 3, "Book launch posters": 2, "Monthly social posts": 1.3,
           "Brand refresh": 2.2}
tasks = {"Spring catalogue": ["Catalogue layout", "Cover design", "Product pages", "Proof corrections"],
         "Website redesign": ["Page layouts", "Home page design", "Mobile layouts", "Client call",
                              "Revisions after feedback", "Style guide"],
         "Book launch posters": ["Poster concepts", "Final artwork", "Print files"],
         "Monthly social posts": ["Post designs", "Captions and scheduling", "Monthly report"],
         "Brand refresh": ["Logo sketches", "Colour and type", "Brand guidelines", "Workshop with the team"]}
rng = random.Random(30)
lines = []   # date, client, project, what, start, end, break, hours typed, not billable, own rate, put on invoice
d = m0
while d <= today:
    if d.weekday() < 5:
        active = [p for p in projects if p[4] <= d <= p[5]]
        t = dt.datetime.combine(d, dt.time(9, rng.choice([0, 30])))
        for _ in range(rng.choice([1, 2, 2, 3])):
            p = rng.choices(active, [weights[x[0]] for x in active])[0]
            length = rng.choice([1, 1.5, 2, 2, 2.5, 3, 3.5])
            what = rng.choice(tasks[p[0]])
            if rng.random() < 0.2:          # hours typed in, without times
                lines.append([d, p[1], p[0], what, None, None, None, length, None, None, None])
                continue
            brk = 15 if length >= 3 and rng.random() < 0.4 else None
            end = t + dt.timedelta(hours=length, minutes=brk or 0)
            lines.append([d, p[1], p[0], what, t.time(), end.time(), brk, None, None, None, None])
            t = end + dt.timedelta(minutes=rng.choice([15, 30, 45, 60]))
    d += dt.timedelta(days=1)
saturday = last_month + dt.timedelta(days=(5 - last_month.weekday()) % 7 + 14)
lines += [
    [last_month + dt.timedelta(days=17), "Harbour Books", "Website redesign", "Rush fixes before launch", dt.time(18, 0),
     dt.time(20, 30), None, None, None, 95, None],
    [saturday, "Atlas Fitness", "Brand refresh", "Weekend print check", None, None, None, 1.5, None, 100, None],
    [today - dt.timedelta(days=6), "Greenleaf Studio", None, "Quote for a new logo", None, None, None, 0.5, "x", None,
     None],
]
lines.sort(key=lambda x: (x[0], x[4] or dt.time()))
# Each month is invoiced on the first of the next; last month's are still being paid.
invoice_log = []   # number, client, date, work from, work to, paid on
n = 1
for k in range(-5, 0):
    frm, to = month_start(today, k), month_start(today, k + 1) - dt.timedelta(days=1)
    dated = month_start(today, k + 1)
    for c in clients:
        if any(x[1] == c[0] and frm <= x[0] <= to and x[8] is None for x in lines):
            due = dated + dt.timedelta(days=c[6] or 14)
            paid = None if (k == -1 and c[0] != "Greenleaf Studio") else due - dt.timedelta(days=rng.choice([1, 2, 3, 5]))
            invoice_log.append((f"RDS-{dated.year}-{n:03d}", c[0], dated, frm, to, paid if paid and paid <= today else None))
            n += 1
shown = next(v[0] for v in invoice_log if v[1] == "Harbour Books" and v[2] == first_this)

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def plain(ws, ref, value, size=10, bold=False, colour=None, italic=False, fmt=None, align=None):
    ws[ref] = value
    ws[ref].font = font(size, bold, colour or "1E2A2F", italic)
    if fmt:
        ws[ref].number_format = fmt
    if align:
        ws[ref].alignment = Alignment(horizontal=align)
    return ws[ref]


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "You, as the invoices show you, and how they are paid.", [3, 22, 44, 4, 50], rows=25)
for k, (label, value) in enumerate(me):
    r = 4 + k
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    fmt = "0%" if label == "Tax rate" else ("0" if label == "Days to pay" else None)
    style(se[f"C{r}"], True, fmt, label == "Your name or business")
    se[f"C{r}"] = value
for k, text in enumerate(["The Invoice tab uses everything on this page.",
                          "Leave the tax rate at 0% if you do not charge tax.",
                          "Days to pay: for clients with none of their own."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Clients
cl = wb.create_sheet("Clients")
sheet_base(cl, "Clients", "Each client once, with the rate per hour and how many days they have to pay.",
           [3, 20, 16, 28, 22, 20, 10, 10], rows=K1 + 2)
header(cl, 5, 2, ["Client", "Contact", "Email", "Address", "Town and postcode", "Rate", "Days to pay"])
month0 = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
month1 = "EOMONTH(TODAY(),0)"
for r in range(K0, K1 + 1):
    x = clients[r - K0] if r - K0 < len(clients) else (None,) * 7
    for c, v in enumerate(x):
        style(cl.cell(r, 2 + c, v), True, MONEY if c == 5 else ("0" if c == 6 else None), c == 0,
              "center" if c == 6 else None)
    busy = (f'SUMIFS({TS("M")},{TS("C")},B{r},{TS("B")},">="&{month0},{TS("B")},"<="&{month1})'
            f'+SUMIFS({TS("O")},{TS("C")},B{r},{TS("Q")},"To invoice")+SUMIFS({TS("O")},{TS("C")},B{r},{TS("Q")},"Draft")'
            f'+SUMIFS({IV("J")},{IV("C")},B{r},{IV("M")},"Waiting")+SUMIFS({IV("J")},{IV("C")},B{r},{IV("M")},"Overdue*")')
    cl.cell(r, 10, f'=IF(B{r}="","",IF({busy}>0,ROW(),""))')                  # J: on the Dashboard
cl.column_dimensions["J"].hidden = True

# ---------------------------------------------------------------- Projects
pr = wb.create_sheet("Projects")
sheet_base(pr, "Projects", "Optional. A project can have its own rate, and a budget of hours to keep an eye on.",
           [3, 24, 20, 10, 10, 10, 10, 12], rows=P1 + 2)
header(pr, 5, 2, ["Project", "Client", "Own rate", "Budget (h)", "Hours", "Left", "Value"])
for r in range(P0, P1 + 1):
    x = projects[r - P0] if r - P0 < len(projects) else (None,) * 4
    style(pr.cell(r, 2, x[0]), True, bold=True)
    style(pr.cell(r, 3, x[1]), True)
    style(pr.cell(r, 4, x[2]), True, MONEY)
    style(pr.cell(r, 5, x[3]), True, "0", align="center")
    style(pr.cell(r, 6, f'=IF(B{r}="","",SUMIFS({TS("M")},{TS("D")},B{r}))'), False, "0.0", align="center")
    style(pr.cell(r, 7, f'=IF(OR(B{r}="",N(E{r})=0),"",E{r}-F{r})'), False, "0.0;[Red]-0.0", align="center")
    style(pr.cell(r, 8, f'=IF(B{r}="","",SUMIFS({TS("O")},{TS("D")},B{r}))'), False, MONEY)
dropdown(pr, f"={CL('B')}", f"C{P0}:C{P1}")
pr.conditional_formatting.add(f"G{P0}:G{P1}", FormulaRule(formula=[f'AND(ISNUMBER(G{P0}),G{P0}<0)'], fill=fill(BAD)))
pr.conditional_formatting.add(f"G{P0}:G{P1}", FormulaRule(formula=[f'AND(ISNUMBER(G{P0}),G{P0}>=0,G{P0}<=0.1*E{P0})'],
                                                          fill=fill(WARN)))

# ---------------------------------------------------------------- Timesheet
ts = wb.create_sheet("Timesheet")
sheet_base(ts, "Timesheet", "One line per piece of work: start and end times, or the hours typed in. Each line goes on "
           "the invoice for its client and dates by itself.",
           [3, 11, 17, 20, 26, 7, 7, 7, 7, 8, 8, 14, 7, 8, 10, 14, 12], rows=T1 + 2)
header(ts, 5, 2, ["Date", "Client", "Project", "What", "Start", "End", "Break (min)", "Hours typed", "Not billable",
                  "Own rate", "Put on invoice", "Hours", "Rate", "Amount", "Invoice", "Status"])
match_inv = f'{IV("C")},C{{r}},{IV("E")},"<="&B{{r}},{IV("F")},">="&B{{r}}'
for r in range(T0, T1 + 1):
    x = lines[r - T0] if r - T0 < len(lines) else [None] * 11
    dte, client, project, task, begin, end, brk, hours, nb, own, put = x
    style(ts.cell(r, 2, dte), True, DATE)
    style(ts.cell(r, 3, client), True, bold=True)
    style(ts.cell(r, 4, project), True)
    style(ts.cell(r, 5, task), True)
    style(ts.cell(r, 6, begin), True, "hh:mm", align="center")
    style(ts.cell(r, 7, end), True, "hh:mm", align="center")
    style(ts.cell(r, 8, brk), True, "0", align="center")
    style(ts.cell(r, 9, hours), True, "0.0#", align="center")
    style(ts.cell(r, 10, nb), True, align="center")
    style(ts.cell(r, 11, own), True, MONEY)
    style(ts.cell(r, 12, put), True)
    style(ts.cell(r, 13, f'=IF(B{r}="","",IF(AND(F{r}<>"",G{r}<>""),MAX(0,ROUND(MOD(G{r}-F{r},1)*24-N(H{r})/60,2)),'
                         f'N(I{r})))'), False, "0.00", align="center")
    project_rate = f'IFERROR(INDEX({PR("D")},MATCH(D{r},{PR("B")},0)),"")'
    client_rate = f'IFERROR(INDEX({CL("G")},MATCH(C{r},{CL("B")},0)),0)'
    style(ts.cell(r, 14, f'=IF(B{r}="","",IF(K{r}<>"",K{r},IF(N({project_rate})>0,{project_rate},{client_rate})))'),
          False, MONEY)
    style(ts.cell(r, 15, f'=IF(B{r}="","",IF(J{r}<>"",0,ROUND(M{r}*N{r},2)))'), False, MONEY)
    m = match_inv.format(r=r)
    style(ts.cell(r, 16, f'=IF(OR(B{r}="",J{r}<>""),"",IF(L{r}<>"",L{r},IF(COUNTIFS({m})=1,INDEX({IV("B")},'
                         f'SUMIFS({IV("O")},{m})),IF(COUNTIFS({m})>1,"Two invoices",""))))'), False,
          align="center")
    found = f'MATCH(P{r},{IV("B")},0)'
    style(ts.cell(r, 17, f'=IF(B{r}="","",IF(J{r}<>"","Not billable",IF(P{r}="","To invoice",IF(P{r}="Two invoices",'
                         f'"Check dates",IF(ISNA({found}),"Check number",IF(INDEX({IV("M")},{found})="Paid","Paid",'
                         f'IF(INDEX({IV("M")},{found})="Draft","Draft","Invoiced")))))))'), False, align="center")
    # R: on the invoice picked on the Invoice tab, in date order. S: its kind of work. T: first line of each kind.
    ts.cell(r, 18, f'=IF(AND(B{r}<>"",P{r}<>"",P{r}={CHOSEN}),B{r}+MOD(N(F{r}),1)+ROW()/10000000,"")')
    ts.cell(r, 21, f'=IF(B{r}="","",IF(J{r}<>"",0,M{r}))')                                   # U: billable hours
    ts.cell(r, 19, f'=IF(R{r}="","",IF(D{r}="","Other work",D{r})&IF(AND(K{r}<>"",E{r}<>""),", "&E{r},""))')
    if r == T0:
        ts.cell(r, 20, f'=IF(R{r}="","",R{r})')
    else:
        ts.cell(r, 20, f'=IF(R{r}="","",IF(COUNTIFS(R${T0}:R{r - 1},">0",S${T0}:S{r - 1},S{r},N${T0}:N{r - 1},N{r})=0,'
                       f'R{r},""))')
for c in "RSTU":
    ts.column_dimensions[c].hidden = True
dropdown(ts, f"={CL('B')}", f"C{T0}:C{T1}")
dropdown(ts, f"={PR('B')}", f"D{T0}:D{T1}", strict=False)
dropdown(ts, f"={IV('B')}", f"L{T0}:L{T1}")
for text, colour in (("To invoice", WARN), ("Draft", WARN), ("Invoiced", INFO), ("Paid", OK)):
    ts.conditional_formatting.add(f"Q{T0}:Q{T1}", FormulaRule(formula=[f'Q{T0}="{text}"'], fill=fill(colour)))
ts.conditional_formatting.add(f"Q{T0}:Q{T1}", FormulaRule(formula=[f'LEFT(Q{T0},5)="Check"'], fill=fill(BAD)))
ts.conditional_formatting.add(f"B{T0}:Q{T1}", FormulaRule(formula=[f'$J{T0}<>""'], font=Font(name=F, italic=True,
                                                                                              color=MUTED)))
ts.freeze_panes = "C6"

# ---------------------------------------------------------------- Invoices
iv = wb.create_sheet("Invoices")
sheet_base(iv, "Invoices", "One line per invoice. The work of that client in those dates goes on it by itself. Type "
           "the day it is paid.", [3, 16, 20, 13, 12, 12, 9, 12, 10, 12, 13, 13, 16], rows=V1 + 2)
header(iv, 5, 2, ["Number", "Client", "Date", "Work from", "Work to", "Hours", "Before tax", "Tax", "Total", "Due",
                  "Paid on", "Status"])
for r in range(V0, V1 + 1):
    x = invoice_log[r - V0] if r - V0 < len(invoice_log) else (None,) * 6
    number, client, dte, frm, to, paid = x
    style(iv.cell(r, 2, number), True, bold=True)
    style(iv.cell(r, 3, client), True)
    style(iv.cell(r, 4, dte), True, DATE)
    style(iv.cell(r, 5, frm), True, DATE)
    style(iv.cell(r, 6, to), True, DATE)
    style(iv.cell(r, 7, f'=IF(B{r}="","",SUMIFS({TS("M")},{TS("P")},B{r}))'), False, "0.0", align="center")
    style(iv.cell(r, 8, f'=IF(B{r}="","",SUMIFS({TS("O")},{TS("P")},B{r}))'), False, MONEY)
    style(iv.cell(r, 9, f'=IF(B{r}="","",ROUND(H{r}*N({TAX}),2))'), False, MONEY)
    style(iv.cell(r, 10, f'=IF(B{r}="","",H{r}+I{r})'), False, MONEY, True)
    terms = (f'IFERROR(IF(INDEX({CL("H")},MATCH(C{r},{CL("B")},0))="",N({TERMS}),'
             f'INDEX({CL("H")},MATCH(C{r},{CL("B")},0))),N({TERMS}))')
    style(iv.cell(r, 11, f'=IF(OR(B{r}="",D{r}=""),"",D{r}+{terms})'), False, DATE, align="center")
    style(iv.cell(r, 12, paid), True, DATE)
    style(iv.cell(r, 13, f'=IF(B{r}="","",IF(L{r}<>"","Paid",IF(D{r}="","Draft",IF(K{r}<TODAY(),"Overdue "&(TODAY()-K{r})'
                         f'&IF(TODAY()-K{r}=1," day"," days"),"Waiting"))))'), False, align="center")
    iv.cell(r, 15, f'=IF(B{r}="","",ROW()-{V0 - 1})')                                          # O: which line
    iv.cell(r, 16, f'=IF(OR(B{r}="",L{r}<>"",D{r}=""),"",K{r}+ROW()/10000000)')                  # P: not paid yet
for c in "OP":
    iv.column_dimensions[c].hidden = True
dropdown(iv, f"={CL('B')}", f"C{V0}:C{V1}")
iv.conditional_formatting.add(f"M{V0}:M{V1}", FormulaRule(formula=[f'M{V0}="Paid"'], fill=fill(OK)))
iv.conditional_formatting.add(f"M{V0}:M{V1}", FormulaRule(formula=[f'LEFT(M{V0},7)="Overdue"'], fill=fill(BAD)))
iv.conditional_formatting.add(f"M{V0}:M{V1}", FormulaRule(formula=[f'OR(M{V0}="Waiting",M{V0}="Draft")'], fill=fill(WARN)))
iv.conditional_formatting.add(f"E{V0}:F{V1}", FormulaRule(formula=[f'AND($B{V0}<>"",E{V0}="")'], fill=fill(WARN)))
iv.conditional_formatting.add(f"B{V0}:B{V1}", FormulaRule(formula=[f'AND(B{V0}<>"",COUNTIF($B${V0}:$B${V1},B{V0})>1)'],
                                                          fill=fill(BAD)))
iv.freeze_panes = "C6"

# ---------------------------------------------------------------- Invoice (the page you send)
# The number at the top right picks the invoice; the page holds nothing else, so it prints the same from both apps.
inv = wb.create_sheet("Invoice")
tot = I1 + 2
pay = tot + 4
sheet_base(inv, "", "", [3, 12, 20, 34, 9, 11, 14, 4], rows=pay + 7)
thin = Side(style="thin", color="C9D3D2")
plain(inv, "B2", f"={S(4)}", 20, True, TEAL_D)
for k, row in enumerate((5, 6, 7, 8)):
    plain(inv, f"B{3 + k}", f'=IF({S(row)}="","",{S(row)})', colour=MUTED)
plain(inv, "B7", f'=IF({S(9)}="","","Tax number "&{S(9)})', colour=MUTED)
plain(inv, "F2", "Invoice", 24, True, TEAL, align="right")
inv.merge_cells("F2:G2")
# Hidden N: which line of Invoices, its client, and the kinds of work on it.
inv["N4"] = f'=IFERROR(MATCH({CHOSEN},{IV("B")},0),0)'
inv["N5"] = f'=IF(N4=0,"",IF(INDEX({IV("C")},N4)="","",INDEX({IV("C")},N4)))'
field = lambda col: f'IF(N4=0,"",IF(INDEX({IV(col)},N4)="","",INDEX({IV(col)},N4)))'
for r, label, formula, fmt in ((4, "Number", shown, None),
                               (5, "Date", f'=IF(N4=0,"",IF(INDEX({IV("D")},N4)="","Draft",INDEX({IV("D")},N4)))', DATE),
                               (6, "Due", f"={field('K')}", DATE),
                               (10, "Work from", f"={field('E')}", DATE),
                               (11, "Work to", f"={field('F')}", DATE)):
    plain(inv, f"F{r}", label, bold=True, colour=MUTED, align="right")
    plain(inv, f"G{r}", formula, bold=True, fmt=fmt, align="right")
pick = DataValidation(type="list", formula1=f"={IV('B')}", allow_blank=True, showErrorMessage=True,
                      showInputMessage=True, promptTitle="Which invoice",
                      prompt="Pick its number. Add the invoice on the Invoices tab first.")
inv.add_data_validation(pick)
pick.add("G4")
plain(inv, "B9", "Bill to", bold=True, colour=MUTED)
client_field = lambda col: f'IFERROR(INDEX({CL(col)},MATCH($N$5,{CL("B")},0)),"")'
plain(inv, "B10", "=N5", 11, True)
for k, col in enumerate("CEFD"):
    plain(inv, f"B{11 + k}", f'=IF({client_field(col)}="","",{client_field(col)})')
header(inv, 16, 2, ["Description", "", "", "Hours", "Rate", "Amount"])
inv.merge_cells("B16:D16")
inv["B16"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for i in range(I1 - I0 + 1):
    r = I0 + i
    key = f"N{r}"
    inv[key] = f'=IFERROR(SMALL({TS("T")},{i + 1}),"")'
    row = f'MATCH({key},{TS("T")},0)'
    kind = f'{TS("R")},">0",{TS("S")},B{r},{TS("N")},F{r}'
    empty = (f'IF({CHOSEN}="","Pick the invoice number at the top right.",IF(N4=0,"This number is not on the '
             f'Invoices tab.","No work matches: check the client and the dates on the Invoices tab."))') if i == 0 else '""'
    plain(inv, f"B{r}", f'=IF({key}="",{empty},INDEX({TS("S")},{row}))')
    plain(inv, f"E{r}", f'=IF({key}="","",SUMIFS({TS("M")},{kind}))', fmt="0.00", align="center")
    plain(inv, f"F{r}", f'=IF({key}="","",INDEX({TS("N")},{row}))', fmt=MONEY, align="right")
    plain(inv, f"G{r}", f'=IF({key}="","",SUMIFS({TS("O")},{kind}))', fmt=MONEY, align="right")
    for c in "BCDEFG":
        inv[f"{c}{r}"].border = Border(bottom=thin)
    inv.row_dimensions[r].height = 19
inv.conditional_formatting.add(f"B{I0}", FormulaRule(formula=[f'$N${I0}=""'], font=Font(name=F, bold=True, color="B4541F")))
plain(inv, f"B{I1 + 1}", f'=IF(COUNT({TS("T")})>{I1 - I0 + 1},"More than {I1 - I0 + 1} kinds of work: split them over two '
      f'invoices.","")', bold=True, colour="B4541F")
inv.column_dimensions["N"].hidden = True
plain(inv, f"D{tot}", "Total hours", bold=True, colour=MUTED, align="right")
plain(inv, f"E{tot}", f'=SUMIFS({TS("M")},{TS("R")},">0")', bold=True, fmt="0.00", align="center")
plain(inv, f"F{tot}", "Before tax", bold=True, colour=MUTED, align="right")
plain(inv, f"G{tot}", f'=SUMIFS({TS("O")},{TS("R")},">0")', fmt=MONEY, align="right")
plain(inv, f"F{tot + 1}", f'=IF(N({TAX})=0,"",{TAXNAME}&" "&ROUND(N({TAX})*100,1)&"%")', bold=True, colour=MUTED,
      align="right")
plain(inv, f"G{tot + 1}", f'=IF(N({TAX})=0,"",ROUND(G{tot}*N({TAX}),2))', fmt=MONEY, align="right")
plain(inv, f"F{tot + 2}", f'="Total ("&{CUR}&")"', 12, True, TEAL_D, align="right")
plain(inv, f"G{tot + 2}", f"=G{tot}+N(G{tot + 1})", 12, True, fmt=MONEY, align="right")
for c in "FG":
    inv[f"{c}{tot + 2}"].border = Border(top=Side(style="medium", color=TEAL))
plain(inv, f"B{pay}", "How to pay", bold=True, colour=MUTED)
plain(inv, f"B{pay + 1}", f'=IF({S(14)}="","",{S(14)})')
plain(inv, f"B{pay + 2}", f'=IF({S(15)}="","",{S(15)})')
plain(inv, f"B{pay + 3}", f'=IF(G4="","",IF(G6="","Please quote "&G4&" when you pay.","Please pay by "&{long_date("G6")}'
                          f'&", quoting "&G4&"."))')
plain(inv, f"B{pay + 5}", f'=IF({S(16)}="","",{S(16)})', italic=True, colour=MUTED)
inv.print_area = f"A1:H{pay + 6}"
inv.page_setup.orientation = "portrait"
inv.page_setup.fitToWidth = 1
inv.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Hours report (every line of that invoice)
hr = wb.create_sheet("Hours report")
sheet_base(hr, "", "", [3, 11, 22, 36, 7, 7, 8, 11, 3, 10], rows=H1 + 6)
period = (f'IF(OR(Invoice!$G$10="",Invoice!$G$11=""),"",", work from "&{short("Invoice!$G$10")}&" to "&'
          f'{long_date("Invoice!$G$11")})')
plain(hr, "B2", f'=IF({CHOSEN}="","Hours report","Hours for invoice "&{CHOSEN})', 20, True, TEAL_D)
plain(hr, "B3", f'=Invoice!$N$5&{period}', 11, colour=MUTED, italic=True)
plain(hr, "B4", f"={S(4)}", 11, colour=MUTED)
header(hr, 6, 2, ["Date", "Project", "What", "Start", "End", "Hours", "Amount"])
for i in range(H1 - H0 + 1):
    r = H0 + i
    key = f"J{r}"
    hr[key] = f'=IFERROR(SMALL({TS("R")},{i + 1}),"")'
    row = f'MATCH({key},{TS("R")},0)'
    for c, formula, fmt, align in (
            ("B", f'=IF({key}="","",INT({key}))', "d mmm", "left"),
            ("C", f'=IF({key}="","",IF(INDEX({TS("D")},{row})="","",INDEX({TS("D")},{row})))', None, "left"),
            ("D", f'=IF({key}="","",IF(INDEX({TS("E")},{row})="","",INDEX({TS("E")},{row})))', None, "left"),
            ("E", f'=IF({key}="","",IF(INDEX({TS("F")},{row})="","",INDEX({TS("F")},{row})))', "hh:mm", "center"),
            ("F", f'=IF({key}="","",IF(INDEX({TS("G")},{row})="","",INDEX({TS("G")},{row})))', "hh:mm", "center"),
            ("G", f'=IF({key}="","",INDEX({TS("M")},{row}))', "0.00", "center"),
            ("H", f'=IF({key}="","",INDEX({TS("O")},{row}))', MONEY, "right")):
        cell = plain(hr, f"{c}{r}", formula, 9, fmt=fmt, align=align)
        cell.border = Border(bottom=thin)
hr.column_dimensions["J"].hidden = True
end = H1 + 2
plain(hr, f"F{end}", "Total", bold=True, colour=MUTED, align="right")
plain(hr, f"G{end}", f'=SUMIFS({TS("M")},{TS("R")},">0")', bold=True, fmt="0.00", align="center")
plain(hr, f"H{end}", f'=SUMIFS({TS("O")},{TS("R")},">0")', bold=True, fmt=MONEY, align="right")
for c in "FGH":
    hr[f"{c}{end}"].border = Border(top=Side(style="medium", color=TEAL))
plain(hr, f"B{end + 1}", f'=IF(COUNT({TS("R")})>{H1 - H0 + 1},"Only the first {H1 - H0 + 1} lines fit on this page. '
      f'The invoice counts them all.","")', 9, colour="B4541F", italic=True)
hr.print_area = f"A1:I{end + 2}"
hr.page_setup.orientation = "portrait"
hr.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Hours and invoices", "The work on the left, the money on the right.",
           [3, 17, 15, 15, 15, 3, 17, 16, 15, 15, 3], rows=45)
monday = "TODAY()-WEEKDAY(TODAY(),3)"
in_month = f'{TS("B")},">="&{month0},{TS("B")},"<="&{month1}'
year0, year1 = "DATE(YEAR(TODAY()),1,1)", "DATE(YEAR(TODAY()),12,31)"
tile(db, "B", 4, "Hours this week", f'=SUMIFS({TS("M")},{TS("B")},">="&{monday},{TS("B")},"<="&{monday}+6)', "0.0")
tile(db, "C", 4, "Hours this month", f'=SUMIFS({TS("M")},{in_month})', "0.0")
tile(db, "D", 4, "Worth this month", f'=SUMIFS({TS("O")},{in_month})', MONEY)
tile(db, "E", 4, "Worth this year", f'=SUMIFS({TS("O")},{TS("B")},">="&{year0},{TS("B")},"<="&{year1})', MONEY)
tile(db, "G", 4, "To invoice", f'=SUMIFS({TS("O")},{TS("Q")},"To invoice")+SUMIFS({TS("O")},{TS("Q")},"Draft")', MONEY,
     "B07A00")
tile(db, "H", 4, "Waiting", f'=SUMIFS({IV("J")},{IV("M")},"Waiting")', MONEY, "B07A00")
tile(db, "I", 4, "Overdue", f'=SUMIFS({IV("J")},{IV("M")},"Overdue*")', MONEY, "B4541F")
tile(db, "J", 4, "Paid this year", f'=SUMIFS({IV("J")},{IV("L")},">="&{year0},{IV("L")},"<="&{year1})', MONEY)

header(db, 8, 2, ["This week", "Hours", "Billable", "Worth"])
for k, name in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]):
    r = 9 + k
    d = f"({monday}+{k})"
    style(db.cell(r, 2, f'="{name} "&{short(d)}'), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({TS("M")},{TS("B")},{d})'), False, "0.0;;", align="center")
    style(db.cell(r, 4, f'=SUMIFS({TS("U")},{TS("B")},{d})'), False, "0.0;;", align="center")
    style(db.cell(r, 5, f'=SUMIFS({TS("O")},{TS("B")},{d})'), False, "#,##0.00;;")
db.conditional_formatting.add("B9:E15", FormulaRule(formula=["ROW()-9=WEEKDAY(TODAY(),3)"], fill=fill(INFO)))

header(db, 19, 2, ["Month", "Hours", "Worth", "Paid in"])
for k in range(6):
    r = 20 + k
    start = f"DATE(YEAR(TODAY()),MONTH(TODAY())-{5 - k},1)"
    span = f'">="&{start},{{0}},"<="&EOMONTH({start},0)'
    style(db.cell(r, 2, f'=CHOOSE(MONTH({start}),{MON_LIST})&" "&YEAR({start})'), False, bold=True)
    style(db.cell(r, 3, f'=SUMIFS({TS("M")},{TS("B")},{span.format(TS("B"))})'), False, "0.0;;", align="center")
    style(db.cell(r, 4, f'=SUMIFS({TS("O")},{TS("B")},{span.format(TS("B"))})'), False, "#,##0.00;;")
    style(db.cell(r, 5, f'=SUMIFS({IV("J")},{IV("L")},{span.format(IV("L"))})'), False, "#,##0.00;;")

header(db, 8, 7, ["Client", "Hours this month", "To invoice", "Not paid"])
for k in range(8):
    r = 9 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL(Clients!$J${K0}:$J${K1},{k + 1}),"")'
    none = '"No work this month yet"' if k == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},INDEX({CL("B")},{key}-{K0 - 1}))'), False, bold=True)
    per = f'{TS("C")},G{r}'
    style(db.cell(r, 8, f'=IF({key}="","",SUMIFS({TS("M")},{per},{in_month}))'), False, "0.0;;", align="center")
    style(db.cell(r, 9, f'=IF({key}="","",SUMIFS({TS("O")},{per},{TS("Q")},"To invoice")+SUMIFS({TS("O")},{per},'
                        f'{TS("Q")},"Draft"))'), False, "#,##0.00;;")
    style(db.cell(r, 10, f'=IF({key}="","",SUMIFS({IV("J")},{IV("C")},G{r},{IV("M")},"Waiting")+SUMIFS({IV("J")},'
                         f'{IV("C")},G{r},{IV("M")},"Overdue*"))'), False, "#,##0.00;;")

header(db, 19, 7, ["Not paid yet", "Client", "Total", "Status"])
for i in range(8):
    r = 20 + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({IV("P")},{i + 1}),"")'
    row = f'MATCH({key},{IV("P")},0)'
    none = '"Every invoice is paid"' if i == 0 else '""'
    state = f'INDEX({IV("M")},{row})'
    style(db.cell(r, 7, f'=IF({key}="",{none},INDEX({IV("B")},{row}))'), False, bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",INDEX({IV("C")},{row}))'), False)
    style(db.cell(r, 9, f'=IF({key}="","",INDEX({IV("J")},{row}))'), False, MONEY)
    style(db.cell(r, 10, f'=IF({key}="","",IF({state}="Waiting","Due "&{short(f"INT({key})")},{state}))'), False,
          align="center")
db.conditional_formatting.add("J20:J27", FormulaRule(formula=['LEFT(J20,7)="Overdue"'], fill=fill(BAD)))
db.conditional_formatting.add("J20:J27", FormulaRule(formula=['LEFT(J20,3)="Due"'], fill=fill(WARN)))

chart = BarChart()
chart.type = "col"
chart.title = "Month by month"
chart.height, chart.width = 6.5, 24
chart.add_data(Reference(db, min_col=4, max_col=5, min_row=19, max_row=25), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=20, max_row=25))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "F2A65A"
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "B30")
for c in "OP":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: your name, address and bank details as the invoice shows them, the currency and any tax."),
    ("2", "Clients: each client with the rate per hour. Projects only if one has its own rate or a budget of hours."),
    ("3", "Timesheet: one line per piece of work, with start and end times or the hours typed in."),
    ("4", "Invoices: a line with the number, the client and the dates of the work. That work goes on it by itself."),
    ("5", "Invoice: pick the number at the top right, then print it or save it as a PDF."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example studio, its clients and their hours are made up. Delete them and add your own.",
                          "The Hours report lists every line of the same invoice, for a client who wants the detail.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Timesheet", "Invoices", "Invoice", "Hours report", "Clients", "Projects", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Timesheet-Invoice-Maker.xlsx")
wb.save(out)
print("saved", out, len(lines), "timesheet lines", len(invoice_log), "invoices, showing", shown)
