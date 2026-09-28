"""Pet Care Tracker (Excel + Google Sheets).

Every pet with its age, chip, insurance and vet; vaccines and treatments with the next dose; vet visits; weight
over time; the food bag and the day it runs out; and what each pet costs by category. The Dashboard lists what is
overdue or due soon, the food to buy and the next vet visits. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
P0, P1 = 6, 25         # pets
H0, H1 = 6, 65         # health schedule
V0, V1 = 6, 305        # vet visits
W0, W1 = 6, 505        # weights
F0, F1 = 6, 25         # food
S0, S1 = 6, 1005       # spending
MON_LIST = ",".join(f'"{m}"' for m in ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
PE = lambda col: f"Pets!${col}${P0}:${col}${P1}"
HE = lambda col: f"Health!${col}${H0}:${col}${H1}"
VE = lambda col: f"'Vet visits'!${col}${V0}:${col}${V1}"
WE = lambda col: f"Weight!${col}${W0}:${col}${W1}"
FO = lambda col: f"Food!${col}${F0}:${col}${F1}"
SP = lambda col: f"Spending!${col}${S0}:${col}${S1}"
AHEAD, BUY = "Dashboard!$C$4", "Dashboard!$G$4"
SPECIES = ["Dog", "Cat", "Rabbit", "Guinea pig", "Bird", "Fish", "Reptile", "Horse", "Other"]
CATS = ["Food", "Vet", "Medicine", "Insurance", "Grooming", "Toys", "Supplies", "Boarding", "Other"]


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


def days_text(n):
    return f'{n}&IF({n}=1," day"," days")'


def ago(years=0, months=0, days=0):
    m = today.month - 1 - months
    y = today.year - years + m // 12
    m = m % 12 + 1
    d = min(today.day, [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1])
    return dt.date(y, m, d) - dt.timedelta(days=days)


# ---------------------------------------------------------------- example data (made-up pets, vets and shops)
pets = [
    # name, species, breed, birthday, sex, neutered, chip, insurance, vet, notes
    ("Biscuit", "Dog", "Labrador", ago(4, 3, 9), "Male", "Yes", "981000000000101", "PetShield, policy 44-1021",
     "Riverside Vets", "Loves the lake, hates the hoover"),
    ("Miso", "Cat", "Tabby", ago(7, 1, 20), "Female", "Yes", "981000000000102", None, "Riverside Vets", "Indoor cat"),
    ("Clover", "Rabbit", "Mini lop", ago(2, 5, 2), "Female", "Yes", None, None, "Burrow Exotic Vets", "Free roam in the kitchen"),
]
health = [
    # pet, what, every (months), last done
    ("Biscuit", "Yearly vaccine", 12, ago(0, 11, -20)), ("Biscuit", "Rabies vaccine", 36, ago(1, 2)),
    ("Biscuit", "Flea and tick", 1, ago(0, 1, 6)), ("Biscuit", "Worming", 3, ago(0, 2, 14)),
    ("Biscuit", "Teeth check", 12, ago(0, 5)),
    ("Miso", "Cat flu vaccine", 12, ago(0, 12, 10)), ("Miso", "Flea treatment", 1, ago(0, 0, 20)),
    ("Miso", "Worming", 3, ago(0, 1, 3)),
    ("Clover", "Myxo and RHD vaccine", 12, ago(0, 10, 15)), ("Clover", "Nail trim", 2, ago(0, 1, 27)),
]
visits = [
    # date, pet, reason, vet, notes, cost
    (ago(0, 8, 3), "Miso", "Yearly check and vaccine", "Riverside Vets", "Healthy, a bit overweight", 68.00),
    (ago(0, 6, 12), "Biscuit", "Limping after a walk", "Riverside Vets", "Sprain, rest for a week", 92.50),
    (ago(0, 5), "Biscuit", "Teeth check", "Riverside Vets", "Teeth fine", 45.00),
    (ago(0, 3, 18), "Clover", "Not eating", "Burrow Exotic Vets", "Gut slowdown, gave medicine", 118.00),
    (ago(0, 1, 5), "Miso", "Ear scratching", "Riverside Vets", "Ear mites, drops for 2 weeks", 54.00),
    (today + dt.timedelta(days=9), "Biscuit", "Yearly vaccine", "Riverside Vets", "Booked at 10:30", None),
    (today + dt.timedelta(days=23), "Miso", "Yearly check", "Riverside Vets", "Booked at 17:15", None),
]
rng = random.Random(27)
weights = []
for name, start, step in (("Biscuit", 33.2, 0.12), ("Miso", 5.6, -0.04), ("Clover", 1.72, 0.01)):
    w = start
    for k in range(8, -1, -1):
        weights.append((ago(0, k, 2), name, round(w, 2 if name == "Clover" else 1)))
        w += step + rng.uniform(-0.1, 0.1) * (0.1 if name == "Clover" else 1)
weights.sort()
food = [
    # pet, food, bag size, a day, opened on, price per bag, where
    ("Biscuit", "Adult dry food", 12000, 300, today - dt.timedelta(days=31), 54.99, "Pet shop"),
    ("Miso", "Indoor cat biscuits", 4000, 60, today - dt.timedelta(days=52), 22.50, "Online"),
    ("Miso", "Wet food pouches", 48, 2, today - dt.timedelta(days=12), 19.20, "Supermarket"),
    ("Clover", "Hay", 2500, 120, today - dt.timedelta(days=18), 9.80, "Pet shop"),
    ("Clover", "Rabbit pellets", 2000, 25, today - dt.timedelta(days=40), 11.50, "Pet shop"),
]
spending = []
for k in range(today.month - 1, -1, -1):
    first = dt.date(Y, today.month - k, 1)
    if first > today:
        continue
    spending.append((first + dt.timedelta(days=2), "Biscuit", "Insurance", "Monthly premium", 24.90))
    spending.append((first + dt.timedelta(days=rng.randint(5, 20)), "Biscuit", "Food", "Adult dry food, 12 kg", 54.99))
    if k % 2 == 0:
        spending.append((first + dt.timedelta(days=rng.randint(3, 25)), "Miso", "Food", "Cat biscuits, 4 kg", 22.50))
    spending.append((first + dt.timedelta(days=rng.randint(3, 25)), "Miso", "Food", "Wet food, 48 pouches", 19.20))
    spending.append((first + dt.timedelta(days=rng.randint(3, 25)), "Clover", "Food", "Hay and pellets", 21.30))
    spending.append((first + dt.timedelta(days=rng.randint(1, 10)), "All", "Medicine", "Flea and worming", 18.60))
spending += [(ago(0, 4, 8), "Biscuit", "Grooming", "Wash and trim", 45.00), (ago(0, 2, 20), "Biscuit", "Toys", "Rope toy and ball", 12.40),
             (ago(0, 3, 2), "Miso", "Supplies", "New scratching post", 36.00), (ago(0, 5, 11), "All", "Boarding", "Pet sitter, 4 days", 120.00),
             (ago(0, 1, 16), "Clover", "Supplies", "Litter", 8.75)]
spending = sorted([s for s in spending if s[0] <= today and s[0].year == Y])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Pets
pe = wb.active
pe.title = "Pets"
sheet_base(pe, "Pets", "Each pet once, with its papers. Age and the latest weight fill in by themselves.",
           [3, 13, 11, 14, 13, 14, 9, 10, 18, 26, 20, 11, 11, 30], rows=P1 + 2)
header(pe, 5, 2, ["Name", "Species", "Breed", "Birthday", "Age", "Sex", "Neutered", "Microchip", "Insurance", "Vet",
                  "Weight", "Change", "Notes"])
last_weight = lambda r, rank: f'SUMIFS({WE("D")},{WE("C")},$B{r},{WE("F")},{rank})'
for r in range(P0, P1 + 1):
    p = pets[r - P0] if r - P0 < len(pets) else (None,) * 10
    name, species, breed, born, sex, neutered, chip, insurance, vet, note = p
    style(pe.cell(r, 2, name), True, bold=True)
    style(pe.cell(r, 3, species), True)
    style(pe.cell(r, 4, breed), True)
    style(pe.cell(r, 5, born), True, DATE)
    style(pe.cell(r, 6, f'=IF(E{r}="","",IF(DATEDIF(E{r},TODAY(),"y")=0,"",DATEDIF(E{r},TODAY(),"y")&"y ")'
                        f'&DATEDIF(E{r},TODAY(),"ym")&"m")'), False, align="center")
    style(pe.cell(r, 7, sex), True, align="center")
    style(pe.cell(r, 8, neutered), True, align="center")
    style(pe.cell(r, 9, chip), True, "@")
    style(pe.cell(r, 10, insurance), True)
    style(pe.cell(r, 11, vet), True)
    has = f'COUNTIFS({WE("C")},$B{r},{WE("F")},0)'
    style(pe.cell(r, 12, f'=IF(OR(B{r}="",{has}=0),"",{last_weight(r, 0)})'), False, "0.0#", align="center")
    style(pe.cell(r, 13, f'=IF(OR(L{r}="",COUNTIFS({WE("C")},$B{r},{WE("F")},1)=0),"",L{r}-{last_weight(r, 1)})'), False,
          '+0.0#;-0.0#;0', align="center")
    style(pe.cell(r, 14, note), True)
dropdown(pe, f'"{",".join(SPECIES)}"', f"C{P0}:C{P1}", strict=False)
dropdown(pe, '"Male,Female"', f"G{P0}:G{P1}", strict=False)
dropdown(pe, '"Yes,No"', f"H{P0}:H{P1}", strict=False)
pe.freeze_panes = "C6"

# ---------------------------------------------------------------- Health
he = wb.create_sheet("Health")
sheet_base(he, "Health", "Vaccines, flea and worm treatments, check-ups: how often, and when it was last done. The "
           "next date and the reminder follow.", [3, 13, 26, 11, 14, 14, 16, 30], rows=H1 + 2)
header(he, 5, 2, ["Pet", "What", "Every (months)", "Last done", "Next due", "Status", "Notes"])
for r in range(H0, H1 + 1):
    h = health[r - H0] if r - H0 < len(health) else (None,) * 4
    pet, what, every, last = h
    style(he.cell(r, 2, pet), True, bold=True)
    style(he.cell(r, 3, what), True)
    style(he.cell(r, 4, every), True, "0", align="center")
    style(he.cell(r, 5, last), True, DATE)
    style(he.cell(r, 6, f'=IF(OR(C{r}="",E{r}="",N(D{r})=0),"",EDATE(E{r},D{r}))'), False, DATE, align="center")
    style(he.cell(r, 7, f'=IF(F{r}="","",IF(F{r}<TODAY(),"Overdue "&{days_text(f"(TODAY()-F{r})")},'
                        f'IF(F{r}=TODAY(),"Due today",IF(F{r}-TODAY()<=N({AHEAD}),"Due in "&{days_text(f"(F{r}-TODAY())")},"OK"))))'),
          False, align="center")
    style(he.cell(r, 8), True)
    he.cell(r, 10, f'=IF(G{r}="","",IF(G{r}="OK","",F{r}+ROW()/10000000))')                          # J: order
he.column_dimensions["J"].hidden = True
dropdown(he, f"={PE('B')}", f"B{H0}:B{H1}")
he.conditional_formatting.add(f"G{H0}:G{H1}", FormulaRule(formula=[f'LEFT(G{H0},7)="Overdue"'], fill=fill(BAD)))
he.conditional_formatting.add(f"G{H0}:G{H1}", FormulaRule(formula=[f'LEFT(G{H0},3)="Due"'], fill=fill(WARN)))
he.conditional_formatting.add(f"G{H0}:G{H1}", FormulaRule(formula=[f'G{H0}="OK"'], fill=fill(OK)))
he[f"B{H1 + 1}"] = "When it is done, change Last done to that day."
he[f"B{H1 + 1}"].font = font(9, color=MUTED, italic=True)
he.freeze_panes = "C6"

# ---------------------------------------------------------------- Vet visits
ve = wb.create_sheet("Vet visits")
sheet_base(ve, "Vet visits", "Every visit, past or booked: why, what the vet said and what it cost. The cost counts "
           "under Vet in the spending.", [3, 13, 13, 30, 20, 36, 11], rows=V1 + 2)
header(ve, 5, 2, ["Date", "Pet", "Why", "Vet", "What the vet said", "Cost"])
for r in range(V0, V1 + 1):
    v = visits[r - V0] if r - V0 < len(visits) else (None,) * 6
    style(ve.cell(r, 2, v[0]), True, DATE)
    style(ve.cell(r, 3, v[1]), True, bold=True)
    style(ve.cell(r, 4, v[2]), True)
    style(ve.cell(r, 5, v[3]), True)
    style(ve.cell(r, 6, v[4]), True)
    style(ve.cell(r, 7, v[5]), True, MONEY)
    ve.cell(r, 9, f'=IF(AND(B{r}<>"",C{r}<>"",B{r}>=TODAY()),B{r}+ROW()/10000000,"")')                # I: booked
ve.column_dimensions["I"].hidden = True
dropdown(ve, f"={PE('B')}", f"C{V0}:C{V1}")
ve.conditional_formatting.add(f"B{V0}:G{V1}", FormulaRule(formula=[f'AND($B{V0}<>"",$B{V0}>=TODAY())'], fill=fill("E3E8FF")))
ve.freeze_panes = "C6"

# ---------------------------------------------------------------- Weight
we = wb.create_sheet("Weight")
sheet_base(we, "Weight", "Weigh each pet now and then, in the unit you like. The Pets tab shows the latest weight and "
           "the change since the time before.", [3, 13, 13, 10, 30], rows=W1 + 2)
header(we, 5, 2, ["Date", "Pet", "Weight", "Notes"])
for r in range(W0, W1 + 1):
    w = weights[r - W0] if r - W0 < len(weights) else (None,) * 3
    style(we.cell(r, 2, w[0]), True, DATE)
    style(we.cell(r, 3, w[1]), True, bold=True)
    style(we.cell(r, 4, w[2]), True, "0.0#", align="center")
    style(we.cell(r, 5), True)
    we.cell(r, 7, f'=IF(OR(B{r}="",C{r}="",D{r}=""),"",B{r}+ROW()/10000000)')                          # G: order
    we.cell(r, 6, f'=IF(G{r}="","",COUNTIFS({WE("C")},C{r},{WE("G")},">"&G{r}))')                      # F: 0 newest
for c in "FG":
    we.column_dimensions[c].hidden = True
dropdown(we, f"={PE('B')}", f"C{W0}:C{W1}")
we.freeze_panes = "C6"

# ---------------------------------------------------------------- Food
fo = wb.create_sheet("Food")
sheet_base(fo, "Food", "Each bag or box in use: its size, how much the pet eats a day, and the day you opened it. Use "
           "the same unit for both, like grams or pouches.", [3, 13, 24, 11, 10, 14, 11, 10, 14, 11, 14, 12, 16],
           rows=F1 + 2)
header(fo, 5, 2, ["Pet", "Food", "Bag size", "A day", "Opened on", "Price", "Lasts (days)", "Runs out", "Days left",
                  "Buy by", "A month", "Where"])
for r in range(F0, F1 + 1):
    f_ = food[r - F0] if r - F0 < len(food) else (None,) * 7
    pet, name, size, daily, opened, price, where = f_
    style(fo.cell(r, 2, pet), True, bold=True)
    style(fo.cell(r, 3, name), True)
    style(fo.cell(r, 4, size), True, "#,##0", align="center")
    style(fo.cell(r, 5, daily), True, "#,##0.#", align="center")
    style(fo.cell(r, 6, opened), True, DATE)
    style(fo.cell(r, 7, price), True, MONEY)
    style(fo.cell(r, 8, f'=IF(OR(N(D{r})=0,N(E{r})=0),"",INT(D{r}/E{r}))'), False, "0", align="center")
    style(fo.cell(r, 9, f'=IF(OR(H{r}="",F{r}=""),"",F{r}+H{r})'), False, "ddd d mmm", align="center")
    style(fo.cell(r, 10, f'=IF(I{r}="","",I{r}-TODAY())'), False, "0", True, "center")
    style(fo.cell(r, 11, f'=IF(I{r}="","",IF(I{r}-N({BUY})<=TODAY(),"Now",I{r}-N({BUY})))'), False, "ddd d mmm",
          align="center")
    style(fo.cell(r, 12, f'=IF(OR(N(D{r})=0,N(E{r})=0,G{r}=""),"",G{r}/D{r}*E{r}*30.4)'), False, MONEY)
    style(fo.cell(r, 13, where), True)
    fo.cell(r, 15, f'=IF(I{r}="","",I{r}+ROW()/10000000)')                                             # O: order
fo.column_dimensions["O"].hidden = True
dropdown(fo, f"={PE('B')}", f"B{F0}:B{F1}", strict=False)
fo.conditional_formatting.add(f"K{F0}:K{F1}", FormulaRule(formula=[f'K{F0}="Now"'], fill=fill(BAD)))
fo.conditional_formatting.add(f"J{F0}:J{F1}", FormulaRule(formula=[f'AND(ISNUMBER(J{F0}),J{F0}<=N({BUY})+7)'], fill=fill(WARN)))
fo[f"B{F1 + 1}"] = "When you open a new bag, change Opened on to that day."
fo[f"B{F1 + 1}"].font = font(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Spending
sp = wb.create_sheet("Spending")
sheet_base(sp, "Spending", "Everything you buy for your pets. Vet visits count from their own tab, so leave them out "
           "here. Use All for things they share.", [3, 12, 13, 14, 34, 11], rows=S1 + 2)
header(sp, 5, 2, ["Date", "Pet", "Category", "What", "Amount"])
for r in range(S0, S1 + 1):
    s = spending[r - S0] if r - S0 < len(spending) else (None,) * 5
    style(sp.cell(r, 2, s[0]), True, DATE)
    style(sp.cell(r, 3, s[1]), True, bold=True)
    style(sp.cell(r, 4, s[2]), True)
    style(sp.cell(r, 5, s[3]), True)
    style(sp.cell(r, 6, s[4]), True, MONEY)
dropdown(sp, f'"{",".join(CATS)}"', f"D{S0}:D{S1}", strict=False)
sp.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Pet care", "What is due for your pets, the food to buy, the next vet visits and what they cost.",
           [3, 14, 24, 13, 18, 3, 14, 20, 11, 12, 3], rows=45)
db["B4"] = "Remind me"
db["B4"].font = font(11, True)
style(db["C4"], True, '0" days ahead"', True, "center")
db["C4"] = 30
db["E4"] = "Buy food"
db["E4"].font = font(11, True)
db["E4"].alignment = Alignment(horizontal="right")
style(db["G4"], True, '0" days before"', True, "center")
db["G4"] = 5
year = f'">="&DATE(YEAR(TODAY()),1,1)'
spent_year = f'SUMIFS({SP("F")},{SP("B")},{year})+SUMIFS({VE("G")},{VE("B")},{year},{VE("B")},"<="&TODAY())'
month = f'">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1)'
spent_month = f'SUMIFS({SP("F")},{SP("B")},{month})+SUMIFS({VE("G")},{VE("B")},{month},{VE("B")},"<="&TODAY())'
tile(db, "B", 6, "Pets", f'=COUNTIF({PE("B")},"?*")', "0")
tile(db, "C", 6, "Overdue", f'=COUNTIF({HE("G")},"Overdue*")', "0", "B4541F")
tile(db, "D", 6, "Due soon", f'=COUNTIF({HE("G")},"Due*")', "0", "B07A00")
tile(db, "E", 6, "Food to buy", f'=COUNTIF({FO("K")},"Now")', "0", "B4541F")
tile(db, "G", 6, "Spent this year", f"={spent_year}", MONEY)
tile(db, "H", 6, "This month", f"={spent_month}", MONEY)

header(db, 9, 2, ["Health", "", "Due", "Status"])
db.merge_cells("B9:C9")
for i in range(8):
    r = 10 + i
    key = f"M{r}"
    db[key] = f'=IFERROR(SMALL({HE("J")},{i + 1}),"")'
    row = f'MATCH({key},{HE("J")},0)'
    none = '"Nothing due soon"' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({HE("B")},{row}))'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",INDEX({HE("C")},{row}))'), False)
    style(db.cell(r, 4, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 5, f'=IF({key}="","",INDEX({HE("G")},{row}))'), False, align="center")
db.conditional_formatting.add("E10:E17", FormulaRule(formula=['LEFT(E10,7)="Overdue"'], fill=fill(BAD)))
db.conditional_formatting.add("E10:E17", FormulaRule(formula=['LEFT(E10,3)="Due"'], fill=fill(WARN)))

header(db, 9, 7, ["Food", "", "Days left", "Buy by"])
db.merge_cells("G9:H9")
for i in range(8):
    r = 10 + i
    key = f"N{r}"
    db[key] = f'=IFERROR(SMALL({FO("O")},{i + 1}),"")'
    row = f'MATCH({key},{FO("O")},0)'
    style(db.cell(r, 7, f'=IF({key}="","",INDEX({FO("B")},{row}))'), False, bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",INDEX({FO("C")},{row}))'), False)
    style(db.cell(r, 9, f'=IF({key}="","",INDEX({FO("J")},{row}))'), False, "0", align="center")
    style(db.cell(r, 10, f'=IF({key}="","",INDEX({FO("K")},{row}))'), False, "ddd d mmm", align="center")
db.conditional_formatting.add("J10:J17", FormulaRule(formula=['J10="Now"'], fill=fill(BAD)))

header(db, 20, 2, ["Vet visits booked", "", "Date", "Why"])
db.merge_cells("B20:C20")
for i in range(5):
    r = 21 + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({VE("I")},{i + 1}),"")'
    row = f'MATCH({key},{VE("I")},0)'
    none = '"No visit booked"' if i == 0 else '""'
    style(db.cell(r, 2, f'=IF({key}="",{none},INDEX({VE("C")},{row}))'), False, bold=True)
    style(db.cell(r, 3, f'=IF({key}="","",INDEX({VE("E")},{row}))'), False)
    style(db.cell(r, 4, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
    style(db.cell(r, 5, f'=IF({key}="","",INDEX({VE("D")},{row}))'), False)

header(db, 20, 7, ["This year", "", "Spent"])
db.merge_cells("G20:H20")
for k, cat in enumerate(CATS):
    r = 21 + k
    style(db.cell(r, 7, cat), False, bold=True)
    style(db.cell(r, 8), False)
    db.merge_cells(f"G{r}:H{r}")
    formula = f'SUMIFS({SP("F")},{SP("D")},G{r},{SP("B")},{year})'
    if cat == "Vet":
        formula += f'+SUMIFS({VE("G")},{VE("B")},{year},{VE("B")},"<="&TODAY())'
    style(db.cell(r, 9, f"={formula}"), False, "#,##0.00;;")

header(db, 31, 2, ["Each pet this year", "", "Spent", "Of which vet"])
db.merge_cells("B31:C31")
for k in range(8):
    r = 32 + k
    name = f"Pets!B{P0 + k}"
    style(db.cell(r, 2, f'=IF({name}="","",{name})'), False, bold=True)
    style(db.cell(r, 3), False)
    db.merge_cells(f"B{r}:C{r}")
    vet = f'SUMIFS({VE("G")},{VE("C")},B{r},{VE("B")},{year},{VE("B")},"<="&TODAY())'
    style(db.cell(r, 4, f'=IF(B{r}="","",SUMIFS({SP("F")},{SP("C")},B{r},{SP("B")},{year})+{vet})'), False, MONEY)
    style(db.cell(r, 5, f'=IF(B{r}="","",{vet})'), False, "#,##0.00;;")
style(db.cell(40, 2, "Shared (All)"), False, bold=True)
style(db.cell(40, 3), False)
db.merge_cells("B40:C40")
style(db.cell(40, 4, f'=SUMIFS({SP("F")},{SP("C")},"All",{SP("B")},{year})'), False, MONEY)
for c in "MNO":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Pets: each pet once, with its birthday, chip, insurance and vet."),
    ("2", "Health: vaccines and treatments, how many months apart, and when each was last done."),
    ("3", "Vet visits and Weight: a line for each visit, booked or past, and each time you weigh a pet."),
    ("4", "Food: each bag in use, how much a day, and the day you opened it. It says when to buy the next one."),
    ("5", "Spending: what you buy for them. The Dashboard shows what is due, the food to buy and the costs."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["How often a vaccine or treatment is needed depends on your pet and your country: ask your vet.",
                          "The example pets, vets and shops are made up. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency and any unit."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Pets", "Health", "Vet visits", "Weight", "Food", "Spending", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Pet-Care-Tracker.xlsx")
wb.save(out)
print("saved", out, len(pets), "pets", len(spending), "spending lines", len(weights), "weights")
