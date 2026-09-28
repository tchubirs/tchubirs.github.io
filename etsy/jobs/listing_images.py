"""Etsy listing photos for the Job Application Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Job-Application-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
apps_end, int_end, ct_end = last(wb["Applications"], 305), last(wb["Interviews"], 205), last(wb["Contacts"], 205)

r = Renders(xlsx, tmp, one_page=("Dashboard",))
close = Renders(xlsx, tmp, one_page=("Dashboard",), hide={"Applications": ["E", "H", "I", "O", "P"]},
                hide_rows={"Applications": range(apps_end + 1, 306), "Interviews": range(int_end + 1, 206),
                           "Contacts": range(ct_end + 1, 206)})

img = canvas(dark=True)
title(img, "Job application tracker", "Every job from saved to the answer, the next step for each, the interviews, "
      "and each week against your goal", dark=True, width=2300)
laptop(img, r.shot("Job search"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-applications.jpg"), "Every application and where it stands",
       "Type the stage and the last contact. A line with no news asks for a follow-up by itself, and a quiet one is "
       "marked No reply.", close.shot("Applications"),
       [("#F9C9C4", "Follow up"), ("#FFE2B8", "Due soon"), ("#D8F0DC", "Offer"), ("#E3E8FF", "Closed or no reply")])

detail(os.path.join(img_dir, "03-interviews.jpg"), "Ready for each interview",
       "What to prepare, who you meet, how it went, and whether the thank-you went out",
       close.shot("Interviews"), [("#E3E8FF", "Coming up"), ("#D8F0DC", "Done"), ("#F9C9C4", "Send thank-you")])

# The Dashboard from the weeks down: the week table and the chart.
page = r.doc[r._index("Job search")]
cut = int((page.search_for("Week of")[0].y0 - 14) * 220 / 72)
full = r.page("Job search")
lower = Image.new("RGB", (full.width, full.height - cut + 60), "#FFFFFF")
lower.paste(full.crop((0, cut, full.width, full.height)), (0, 30))
detail(os.path.join(img_dir, "04-weeks.jpg"), "Each week against your goal",
       "How many you applied for each week, the interviews coming, and the chart of the last eight weeks",
       content(lower))

detail(os.path.join(img_dir, "05-contacts.jpg"), "The people who can help",
       "Recruiters, referrals and friends in the business, with the next time to get in touch",
       close.shot("Contacts"), [("#FFE2B8", "Get in touch")])

img = canvas()
y = title(img, "Inside the file", "Four tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Job search")), "Dashboard", "The whole search"),
    (content(r.page("Applications")), "Applications", "Every job and its status"),
    (content(r.page("Interviews")), "Interviews", "Prep and thank-you notes"),
    (content(r.page("Contacts")), "Contacts", "Who can help"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
