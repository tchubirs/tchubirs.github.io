"""Job Application Tracker (Excel + Google Sheets).

Every job from saved to the answer: the company, role, pay, source, the stage it has reached and the next step. A
line with no next step gets a follow-up date by itself, and one that has been quiet too long is marked No reply. An
Interviews tab with what to prepare and whether the thank-you went out, a Contacts tab, and a Dashboard with the
pipeline, the next steps, the interviews coming and the applications week by week. Only functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MUTED, OK, TEAL, WARN, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
A0, A1 = 6, 305        # applications
I0, I1 = 6, 205        # interviews
C0, C1 = 6, 205        # contacts
S0, S1 = 11, 19        # stages (Settings)
STAGES = ["Saved", "Applied", "Screening", "Interview", "Final interview", "Offer", "Accepted", "Rejected", "Withdrawn"]
CLOSED = ["Accepted", "Rejected", "Withdrawn"]
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
AP = lambda col: f"Applications!${col}${A0}:${col}${A1}"
IV = lambda col: f"Interviews!${col}${I0}:${col}${I1}"
CT = lambda col: f"Contacts!${col}${C0}:${col}${C1}"
GOAL, FOLLOW, QUIET, CUR = "Settings!$C$4", "Settings!$C$5", "Settings!$C$6", "Settings!$C$7"
RED, AMBER = "B4541F", "B07A00"


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


D = lambda n: today + dt.timedelta(days=n)

# ---------------------------------------------------------------- example data (a made-up job search)
apps = [
    # company, role, location, link, pay from, pay to, source, saved, applied, stage, last contact, next step,
    # next step on, priority, contact
    ("Fernhill Labs", "Product Designer", "Remote", "fernhill.example/jobs/41", 90000, 105000, "Job board", D(-48), D(-47),
     "Rejected", D(-30), None, None, "Medium", None),
    ("Quillmark", "UX Designer", "Hybrid, Riverton", "quillmark.example/careers", 85000, 95000, "Company site", D(-46),
     D(-45), "No reply", None, None, None, "Low", None),
    ("Tidewell Health", "Senior Product Designer", "Remote", "tidewell.example/jobs/ux", 105000, 120000, "Referral", D(-44),
     D(-42), "Final interview", D(-2), "Final round", D(3), "High", "Nina Park"),
    ("Brambleton Bank", "UI Designer", "On site, Eastbridge", "brambleton.example/jobs", 80000, 92000, "Recruiter", D(-40),
     D(-39), "Rejected", D(-20), None, None, "Low", "Owen Hart"),
    ("Larkspur Learning", "Product Designer", "Remote", "larkspur.example/join", 88000, 100000, "Job board", D(-36),
     D(-35), "Offer", D(-1), "Reply to the offer", D(4), "High", None),
    ("Oakhaven Foods", "UX Designer", "Hybrid, Riverton", "oakhaven.example/careers/12", 82000, 90000, "Job board", D(-33),
     D(-33), "Applied", None, None, None, "Medium", None),
    ("Wrenfield Travel", "Product Designer", "Remote", "wrenfield.example/jobs/7", 92000, 104000, "Company site", D(-30),
     D(-29), "Interview", D(-5), "Portfolio review", D(6), "High", None),
    ("Cloverdale Logistics", "UX Designer", "On site, Northam", "cloverdale.example/jobs", 78000, 88000, "Job board",
     D(-27), D(-26), "Withdrawn", D(-15), None, None, "Low", None),
    ("Hollowbrook Media", "Design Lead", "Remote", "hollowbrook.example/careers/3", 115000, 130000, "Recruiter", D(-24),
     D(-23), "Screening", D(-9), None, None, "Medium", "Owen Hart"),
    ("Starling Software", "Senior Product Designer", "Remote", "starling.example/jobs/22", 110000, 125000, "Referral",
     D(-20), D(-19), "Interview", D(-3), "Take-home task due", D(1), "High", "Leo Grant"),
    ("Pebblestone Games", "UI Designer", "Hybrid, Riverton", "pebblestone.example/jobs", 80000, 95000, "Job board", D(-17),
     D(-16), "Rejected", D(-4), None, None, "Medium", None),
    ("Thistle Insurance", "Product Designer", "On site, Westford", "thistle.example/careers", 86000, 98000, "Job board",
     D(-12), D(-11), "Applied", None, None, None, "Low", None),
    ("Marigold Retail", "UX Designer", "Remote", "marigold.example/jobs/9", 84000, 96000, "Company site", D(-8), D(-8),
     "Applied", None, None, None, "Medium", None),
    ("Kestrel Analytics", "Product Designer", "Remote", "kestrel.example/jobs/5", 95000, 110000, "Job board", D(-5), D(-4),
     "Applied", None, None, None, "High", None),
    ("Ashgrove Studio", "Senior Product Designer", "Hybrid, Riverton", "ashgrove.example/jobs", 100000, 115000, "Referral",
     D(-3), D(-2), "Screening", D(-1), "Call with the recruiter", D(2), "High", "Nina Park"),
    ("Glimmerglass Energy", "UX Designer", "Remote", "glimmerglass.example/careers", 90000, 100000, "Job board", D(-2),
     None, "Saved", None, "Apply", D(5), "Medium", None),
    ("Pinecrest Studio", "Product Designer", "Remote", "pinecrest.example/jobs/2", 88000, 99000, "Company site", D(-1),
     None, "Saved", None, None, None, "Medium", None),
    ("Silverfin Travel", "Design Lead", "On site, Eastbridge", "silverfin.example/jobs", 120000, 135000, "Recruiter", D(-6),
     None, "Saved", None, "Apply before it closes", D(-1), "Low", "Owen Hart"),
]
apps += [
    ("Copperleaf Books", "UX Designer", "Remote", "copperleaf.example/jobs", 80000, 90000, "Job board", D(-38), D(-37),
     "Applied", None, None, None, "Low", None),
    ("Driftwood Hotels", "Product Designer", "On site, Seaford", "driftwood.example/careers", 84000, 94000, "Job board",
     D(-31), D(-30), "Rejected", D(-18), None, None, "Medium", None),
    ("Emberly Home", "UI Designer", "Remote", "emberly.example/jobs/4", 78000, 90000, "Company site", D(-24), D(-22),
     "Rejected", D(-10), None, None, "Low", None),
    ("Foxglove Health", "Product Designer", "Hybrid, Riverton", "foxglove.example/jobs", 90000, 102000, "Job board",
     D(-14), D(-13), "Applied", None, None, None, "Medium", None),
    ("Juniper Bank", "Product Designer", "Remote", "juniper.example/careers/8", 95000, 108000, "Recruiter", D(-1), D(0),
     "Applied", None, None, None, "High", "Owen Hart"),
]
# "No reply" is not a stage: that line is just Applied long ago.
apps = [a[:9] + ("Applied",) + a[10:] if a[9] == "No reply" else a for a in apps]
apps.sort(key=lambda a: a[7])
interviews = [
    # date, time, company, role, type, with, prep, how it went, thank-you sent
    (D(-36), dt.time(10, 0), "Fernhill Labs", "Product Designer", "Phone screen", "Recruiter", None,
     "Friendly, 20 min", "x"),
    (D(-31), dt.time(14, 0), "Tidewell Health", "Senior Product Designer", "Video", "Head of design",
     "Case study on the booking flow", "Went well", "x"),
    (D(-22), dt.time(11, 30), "Brambleton Bank", "UI Designer", "On site", "Two designers", None, "Hard questions on accessibility", "x"),
    (D(-14), dt.time(9, 30), "Larkspur Learning", "Product Designer", "Video", "Design manager", "Their mobile app",
     "Good, they liked the research work", "x"),
    (D(-9), dt.time(15, 0), "Tidewell Health", "Senior Product Designer", "Task", "Design team", "Two-hour exercise",
     "Presented the task", "x"),
    (D(-5), dt.time(13, 0), "Wrenfield Travel", "Product Designer", "Phone screen", "Recruiter", None, "Short and clear",
     None),
    (D(-3), dt.time(16, 0), "Starling Software", "Senior Product Designer", "Video", "Lead designer", "Their design system",
     "Asked for a take-home task", "x"),
    (D(3), dt.time(10, 0), "Tidewell Health", "Senior Product Designer", "Final", "Design team and VP",
     "Walk through two projects", None, None),
    (D(6), dt.time(11, 0), "Wrenfield Travel", "Product Designer", "Video", "Design manager", "Portfolio review, 45 min",
     None, None),
]
contacts = [
    # name, company, title, how we met, email, last contact, next contact, notes
    ("Nina Park", "Tidewell Health", "Product manager", "Old colleague", "nina@tidewell.example", D(-2), D(12),
     "Referred me to Tidewell and Ashgrove"),
    ("Owen Hart", "Recruiting agency", "Recruiter", "Reached out on a job board", "owen@agency.example", D(-9), D(-1),
     "Sends design lead roles"),
    ("Leo Grant", "Starling Software", "Designer", "Design meetup", "leo@starling.example", D(-3), D(10), None),
    ("Maya Chen", "Freelance", "Design mentor", "Mentoring programme", "maya@studio.example", D(-20), D(2),
     "Portfolio feedback every month"),
    ("Sam Ortiz", "Kestrel Analytics", "Engineer", "University", "sam@kestrel.example", D(-4), None, None),
]

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "Your weekly goal, when to follow up, and the currency of the pay.", [3, 36, 14, 4, 60],
           rows=S1 + 3)
for r, label, value, fmt in ((4, "Applications a week (goal)", 5, "0"), (5, "Follow up after (days)", 7, "0"),
                             (6, "No reply after (days)", 30, "0"), (7, "Currency of the pay", "USD", None)):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], True, fmt, True, "center")
    se[f"C{r}"] = value
for k, text in enumerate(["The Dashboard counts each week against this goal.",
                          "With no next step, a line asks for a follow-up this many days after the last contact.",
                          "A line with no news for this many days is marked No reply."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)
header(se, S0 - 1, 2, ["Stages, in order"])
for k, s in enumerate(STAGES):
    style(se.cell(S0 + k, 2, s), False, bold=True)

# ---------------------------------------------------------------- Applications
ap = wb.create_sheet("Applications")
sheet_base(ap, "Applications", "One line per job. Keep the stage and the last contact up to date: the status and the "
           "next steps follow.", [3, 20, 22, 17, 24, 11, 11, 13, 12, 12, 15, 12, 26, 12, 9, 13, 8, 20],
           rows=A1 + 2)
header(ap, 5, 2, ["Company", "Role", "Where", "Link", "Pay from", "Pay to", "Source", "Saved on", "Applied on", "Stage",
                  "Last contact", "Next step", "Next step on", "Priority", "Contact", "Days", "Status"])
stage_list = f"Settings!$B${S0}:$B${S1}"
for r in range(A0, A1 + 1):
    x = apps[r - A0] if r - A0 < len(apps) else (None,) * 15
    fmts = (None, None, None, None, "#,##0", "#,##0", None, DATE, DATE, None, DATE, None, DATE, None, None)
    aligns = (None, None, None, None, None, None, None, "center", "center", "center", "center", None, "center", "center",
              None)
    for c, (v, fmt, align) in enumerate(zip(x, fmts, aligns)):
        style(ap.cell(r, 2 + c, v), True, fmt, c == 0, align)
    style(ap.cell(r, 17, f'=IF(OR(B{r}="",J{r}=""),"",TODAY()-J{r})'), False, "0", align="center")
    ap[f"T{r}"] = f"=MAX(N(J{r}),N(L{r}))"                                                  # T: last news
    ap[f"U{r}"] = f'=IF(N{r}<>"",N{r},IF(T{r}=0,"",T{r}+N({FOLLOW})))'                     # U: when to act
    closed = " ".join(f'K{r}="{s}",' for s in CLOSED).rstrip(",").replace(" ", "")
    status = (f'=IF(B{r}="","",IF(OR({closed}),K{r},IF(K{r}="Offer",IF(N{r}="","Offer","Offer, reply by "&'
              f'{short(f"N{r}")}),IF(OR(K{r}="",K{r}="Saved"),IF(N{r}="","To apply",IF(N{r}<TODAY(),"Closing date passed",'
              f'"Apply by "&{short(f"N{r}")})),IF(AND(N{r}="",T{r}>0,TODAY()-T{r}>=N({QUIET})),"No reply",'
              f'IF(U{r}="","Waiting",IF(U{r}<TODAY(),"Follow up",IF(U{r}-TODAY()<=2,"Due "&{short(f"U{r}")},'
              f'"Waiting"))))))))')
    style(ap.cell(r, 18, status), False, align="center")
    act = f'OR(R{r}="Follow up",LEFT(R{r},3)="Due",R{r}="Waiting",LEFT(R{r},5)="Apply",LEFT(R{r},5)="Offer")'
    ap[f"V{r}"] = f'=IF(AND(B{r}<>"",U{r}<>"",{act}),U{r}+ROW()/10000000,"")'               # V: next steps
    ap[f"W{r}"] = f'=IF(B{r}="","",IF(M{r}<>"",M{r},IF(OR(K{r}="",K{r}="Saved"),"Apply",IF(K{r}="Offer",' \
                  f'"Reply to the offer","Follow up"))))'                                    # W: what to do
for c in "TUVW":
    ap.column_dimensions[c].hidden = True
dropdown(ap, f"={stage_list}", f"K{A0}:K{A1}")
dropdown(ap, '"Job board,Company site,Referral,Recruiter,Other"', f"H{A0}:H{A1}", strict=False)
dropdown(ap, '"High,Medium,Low"', f"O{A0}:O{A1}", strict=False)
dropdown(ap, f"={CT('B')}", f"P{A0}:P{A1}", strict=False)
for text, colour, how in (("Follow up", BAD, "="), ("Closing date passed", BAD, "="), ("Due", WARN, "left"),
                          ("Apply by", WARN, "left"), ("No reply", INFO, "="), ("Offer", OK, "left"), ("Accepted", OK, "="),
                          ("Rejected", INFO, "="), ("Withdrawn", INFO, "=")):
    rule = f'R{A0}="{text}"' if how == "=" else f'LEFT(R{A0},{len(text)})="{text}"'
    ap.conditional_formatting.add(f"R{A0}:R{A1}", FormulaRule(formula=[rule], fill=fill(colour)))
ap.conditional_formatting.add(f"B{A0}:Q{A1}", FormulaRule(formula=[f'OR($K{A0}="Rejected",$K{A0}="Withdrawn")'],
                                                          font=Font(name=F, italic=True, color=MUTED)))
ap.freeze_panes = "C6"

# ---------------------------------------------------------------- Interviews
iv = wb.create_sheet("Interviews")
sheet_base(iv, "Interviews", "Each call, video, visit or task, with what to prepare and whether the thank-you went out.",
           [3, 12, 8, 20, 22, 13, 18, 28, 28, 10, 15], rows=I1 + 2)
header(iv, 5, 2, ["Date", "Time", "Company", "Role", "Type", "With", "To prepare", "How it went", "Thank-you sent",
                  "Status"])
for r in range(I0, I1 + 1):
    x = interviews[r - I0] if r - I0 < len(interviews) else (None,) * 9
    fmts = (DATE, "hh:mm", None, None, None, None, None, None, None)
    aligns = ("center", "center", None, None, None, None, None, None, "center")
    for c, (v, fmt, align) in enumerate(zip(x, fmts, aligns)):
        style(iv.cell(r, 2 + c, v), True, fmt, c == 2, align)
    style(iv.cell(r, 11, f'=IF(B{r}="","",IF(B{r}>TODAY(),"Coming up",IF(B{r}=TODAY(),"Today",IF(J{r}<>"","Done",'
                         f'"Send thank-you"))))'), False, align="center")
    iv[f"L{r}"] = f'=IF(OR(B{r}="",B{r}<TODAY()),"",B{r}+MOD(N(C{r}),1)+ROW()/10000000)'    # L: coming up
iv.column_dimensions["L"].hidden = True
dropdown(iv, f"={AP('B')}", f"D{I0}:D{I1}", strict=False)
dropdown(iv, '"Phone screen,Video,On site,Task,Final"', f"F{I0}:F{I1}", strict=False)
for text, colour in (("Coming up", INFO), ("Today", WARN), ("Done", OK), ("Send thank-you", BAD)):
    iv.conditional_formatting.add(f"K{I0}:K{I1}", FormulaRule(formula=[f'K{I0}="{text}"'], fill=fill(colour)))
iv.freeze_panes = "C6"

# ---------------------------------------------------------------- Contacts
ct = wb.create_sheet("Contacts")
sheet_base(ct, "Contacts", "Recruiters, referrals and people who can help, with the next time to get in touch.",
           [3, 18, 20, 18, 24, 28, 13, 13, 30, 16], rows=C1 + 2)
header(ct, 5, 2, ["Name", "Company", "Title", "How you met", "Email", "Last contact", "Next contact", "Notes", "Status"])
for r in range(C0, C1 + 1):
    x = contacts[r - C0] if r - C0 < len(contacts) else (None,) * 8
    fmts = (None, None, None, None, None, DATE, DATE, None)
    for c, (v, fmt) in enumerate(zip(x, fmts)):
        style(ct.cell(r, 2 + c, v), True, fmt, c == 0, "center" if c in (5, 6) else None)
    style(ct.cell(r, 10, f'=IF(OR(B{r}="",H{r}=""),"",IF(H{r}<=TODAY(),"Get in touch","On "&{short(f"H{r}")}))'),
          False, align="center")
ct.conditional_formatting.add(f"J{C0}:J{C1}", FormulaRule(formula=[f'J{C0}="Get in touch"'], fill=fill(WARN)))
ct.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Job search", "Where every application stands, and what to do next.",
           [3, 18, 14, 14, 16, 3, 20, 22, 25, 12, 3], rows=48)
monday = "(TODAY()-WEEKDAY(TODAY(),3))"
applied = f"COUNT({AP('J')})"
replied = "+".join(f'COUNTIF({AP("K")},"{s}")' for s in ("Screening", "Interview", "Final interview", "Offer", "Accepted",
                                                        "Rejected"))
tile(db, "B", 4, "Applied", f"={applied}", "0")
tile(db, "C", 4, "This week", f'=COUNTIFS({AP("J")},">="&{monday},{AP("J")},"<="&({monday}+6))&" of "&N({GOAL})',
     "General")
tile(db, "D", 4, "Replies", f'=IF({applied}=0,"",({replied})/{applied})', "0%")
tile(db, "E", 4, "Interviews", f'=COUNTIFS({IV("B")},"<="&TODAY())', "0")
tile(db, "G", 4, "Offers", f'=COUNTIF({AP("K")},"Offer")+COUNTIF({AP("K")},"Accepted")', "0")
tile(db, "H", 4, "To follow up", f'=COUNTIF({AP("R")},"Follow up")+COUNTIF({AP("R")},"Due*")', "0", RED)
tile(db, "I", 4, "No reply", f'=COUNTIF({AP("R")},"No reply")', "0", AMBER)
tile(db, "J", 4, "Days searching", f'=IF({applied}=0,"",TODAY()-MIN({AP("J")}))', "0")

header(db, 8, 2, ["Pipeline", "Now", "Share", ""])
for k, s in enumerate(STAGES):
    r = 9 + k
    style(db.cell(r, 2, f"=Settings!B{S0 + k}"), False, bold=True)
    style(db.cell(r, 3, f'=COUNTIF({AP("K")},B{r})'), False, "0;;", align="center")
    total = f'COUNTA({AP("B")})'
    style(db.cell(r, 4, f'=IF({total}=0,"",C{r}/{total})'), False, "0%;;", align="center")
    style(db.cell(r, 5, f'=IF(OR(D{r}="",N(D{r})=0),"",{bar(f"D{r}", 12)})'), False)
    db.cell(r, 5).font = Font(name=F, size=9, color=TEAL)

header(db, 8, 7, ["Next steps", "Role", "What", "When"])
for i in range(10):
    r = 9 + i
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({AP("V")},{i + 1}),"")'
    row = f'MATCH({key},{AP("V")},0)'
    at = lambda col: f"INDEX({AP(col)},{row})"
    none = '"Nothing to do yet"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},{at("B")})'), False, bold=True)
    style(db.cell(r, 8, f'=IF({key}="","",{at("C")})'), False)
    style(db.cell(r, 9, f'=IF({key}="","",{at("W")})'), False)
    style(db.cell(r, 10, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm", align="center")
db.conditional_formatting.add("J9:J18", FormulaRule(formula=['AND(ISNUMBER(J9),J9<TODAY())'], fill=fill(BAD)))
db.conditional_formatting.add("J9:J18", FormulaRule(formula=['AND(ISNUMBER(J9),J9-TODAY()<=2)'], fill=fill(WARN)))

header(db, 21, 2, ["Week of", "Applied", "Goal"])
for k in range(8):
    r = 22 + k
    wk = f"({monday}-{7 * (7 - k)})"
    style(db.cell(r, 2, f"={wk}"), False, "d mmm", True, "center")
    style(db.cell(r, 3, f'=COUNTIFS({AP("J")},">="&B{r},{AP("J")},"<="&(B{r}+6))'), False, "0", align="center")
    style(db.cell(r, 4, f"=N({GOAL})"), False, "0", align="center")
db.conditional_formatting.add("C22:C29", FormulaRule(formula=["C22>=D22"], fill=fill(OK)))

header(db, 21, 7, ["Coming interviews", "Company", "Type", "Time"])
for i in range(6):
    r = 22 + i
    key = f"P{r}"
    db[key] = f'=IFERROR(SMALL({IV("L")},{i + 1}),"")'
    row = f'MATCH({key},{IV("L")},0)'
    at = lambda col: f"INDEX({IV(col)},{row})"
    none = '"None booked yet"' if i == 0 else '""'
    style(db.cell(r, 7, f'=IF({key}="",{none},INT({key}))'), False, "ddd d mmm", True, "center")
    style(db.cell(r, 8, f'=IF({key}="","",{at("D")})'), False)
    style(db.cell(r, 9, f'=IF({key}="","",IF({at("F")}="","",{at("F")}))'), False)
    style(db.cell(r, 10, f'=IF({key}="","",IF({at("C")}="","",{at("C")}))'), False, "hh:mm", align="center")

chart = BarChart()
chart.type = "col"
chart.title = "Applications a week"
chart.height, chart.width = 6.5, 24
chart.add_data(Reference(db, min_col=3, max_col=4, min_row=21, max_row=29), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=22, max_row=29))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "B8D3CF"
chart.y_axis.number_format = "0"
db.add_chart(chart, "B32")
for c in "OP":
    db.column_dimensions[c].hidden = True

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Settings: how many applications a week you aim for, and after how many days to follow up."),
    ("2", "Applications: every job you save or apply for. Keep the stage and the last contact up to date."),
    ("3", "Next step and its date: a call, a task, a closing date. Without one, a follow-up date comes by itself."),
    ("4", "Interviews: each one with what to prepare. An x once the thank-you is sent."),
    ("5", "Contacts and Dashboard: people who can help, and where the whole search stands."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example search, its companies and people are made up. Delete them and add your own.",
                          "Rejected and withdrawn lines stay in the list, in grey, so the numbers stay honest.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "Applications", "Interviews", "Contacts", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Job-Application-Tracker.xlsx")
wb.save(out)
print("saved", out, len(apps), "applications", len(interviews), "interviews", len(contacts), "contacts")
