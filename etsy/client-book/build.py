"""Client and Appointment Tracker for people who work one to one (Excel + Google Sheets).

Hair stylists, trainers, tutors, therapists: clients, services and prices,
appointments with their status and payment, prepaid session packages, a week
view, and a dashboard with income, money owed, clients to follow up and
birthdays. Only functions both apps have.
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, INFO, MONEY, MUTED, OK, TEAL, TEAL_D, WARN, F,  # noqa: E402
                      box, fill, font, header, sheet_base, style)

today = dt.date.today()
Y = today.year
D = lambda y, m, d: dt.date(y, m, d)
S0, S1 = 6, 35        # services
C0, C1 = 6, 305       # clients
A0, A1 = 6, 2005      # appointments
K0, K1 = 6, 305       # packages
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
AP = lambda col: f"Appointments!${col}${A0}:${col}${A1}"
CL = lambda col: f"Clients!${col}${C0}:${col}${C1}"
PK = lambda col: f"Packages!${col}${K0}:${col}${K1}"
SV = lambda col: f"Services!${col}${S0}:${col}${S1}"
GREY = "EEF1F1"
TIME = "h:mm"
STATUSES = "Booked,Done,No-show,Cancelled"
METHODS = "Card,Cash,Transfer,App,Package"

# ---------------------------------------------------------------- example data
services = [("Women's haircut", 55, 60), ("Men's haircut", 30, 30), ("Kids haircut", 20, 30), ("Blow dry", 35, 45),
            ("Colour", 95, 120), ("Highlights", 140, 150), ("Toner", 40, 30), ("Treatment", 30, 30)]
price = {s: p for s, p, _ in services}
minutes = {s: m for s, _, m in services}
clients = [
    # name, phone, email, birthday, client since, notes, visit every (days), services, last visit if they stopped
    ("Anna Rossi", "555 0101", "anna.r@example.com", D(1988, 9, 14), D(2024, 3, 2), "Likes it shorter at the back", 35,
     ["Women's haircut", "Colour"], None),
    ("Ben Carter", "555 0102", None, None, D(2023, 11, 20), None, 28, ["Men's haircut"], None),
    ("Chloe Martin", "555 0103", "chloe.m@example.com", D(1992, 2, 3), D(2022, 6, 11), "Blow dry card since September",
     42, ["Highlights", "Toner"], None),
    ("Daniel Kim", "555 0104", None, D(1990, 12, 1), D(2024, 1, 15), None, 30, ["Men's haircut"], None),
    ("Emma Johnson", "555 0105", "emma.j@example.com", D(1995, 4, 22), D(2023, 5, 5), "Blow dry card", None, [], None),
    ("Farah Ali", "555 0106", None, D(1985, 7, 9), D(2021, 9, 30), "Patch test before colour", 56,
     ["Colour", "Women's haircut"], None),
    ("Grace Lee", "555 0107", "grace.l@example.com", None, D(2024, 8, 8), None, 35, ["Women's haircut", "Blow dry"], None),
    ("Hugo Silva", "555 0108", None, D(1987, 3, 17), D(2023, 2, 1), None, 28, ["Men's haircut"], None),
    ("Isla Brown", "555 0109", None, None, D(2022, 10, 4), "Sensitive scalp", 49, ["Women's haircut", "Treatment"], None),
    ("Jack Wilson", "555 0110", None, None, D(2024, 5, 21), None, 21, ["Men's haircut"], None),
    ("Katie Moore", "555 0111", "katie.m@example.com", D(1993, 11, 8), D(2023, 7, 7), None, 40, ["Women's haircut"],
     D(Y, 5, 20)),
    ("Liam Evans", "555 0112", None, None, D(2025, 1, 10), None, 30, ["Men's haircut"], None),
    ("Maya Patel", "555 0113", "maya.p@example.com", D(1991, 6, 25), D(2024, 2, 14), "Treatment card", None, [], None),
    ("Noah Taylor", "555 0114", None, None, D(2024, 9, 3), "Brings his son Leo", 35, ["Men's haircut", "Kids haircut"],
     None),
    ("Olivia Green", "555 0115", "olivia.g@example.com", D(1989, 9, 30), D(2022, 3, 12), None, 35,
     ["Women's haircut", "Blow dry"], None),
    ("Paul Wright", "555 0116", None, None, D(2023, 4, 18), None, 45, ["Men's haircut"], D(Y, 6, 12)),
    ("Quinn Harris", "555 0117", None, None, D(2025, 3, 3), None, 60, ["Highlights", "Women's haircut"], None),
    ("Rosa Diaz", "555 0118", "rosa.d@example.com", D(1984, 1, 29), D(2021, 5, 25), None, 35, ["Colour", "Women's haircut"],
     None),
    ("Sam Clarke", "555 0119", None, None, D(2024, 10, 10), None, 28, ["Men's haircut"], None),
    ("Tara Singh", "555 0120", None, D(1996, 8, 12), D(2023, 8, 19), None, 45, ["Women's haircut", "Treatment"],
     D(Y, 7, 15)),
    ("Uma Nair", "555 0121", None, None, D(2024, 4, 4), None, 40, ["Women's haircut"], None),
    ("Victor Lane", "555 0122", None, None, D(2025, 2, 2), None, 30, ["Men's haircut"], None),
    ("Wendy Scott", "555 0123", "wendy.s@example.com", None, D(2022, 12, 1), None, 50, ["Women's haircut", "Toner"], None),
    ("Yara Haddad", "555 0124", "yara.h@example.com", None, D(Y, 9, 5), "New client, found us online", None, [], None),
    ("Zoe Turner", "555 0125", None, D(1998, 5, 5), D(2024, 6, 30), None, 33, ["Women's haircut", "Blow dry"], None),
]
rng = random.Random(11)


def open_day(d):
    """The salon opens Tuesday to Saturday: move d to the next open day."""
    while d.weekday() in (0, 6):
        d += dt.timedelta(days=1)
    return d


FIRST = ["Ava", "Mia", "Lily", "Ella", "Sophie", "Ruby", "Alice", "Nora", "Hannah", "Jade", "Leah", "Megan", "Holly",
         "Amy", "Beth", "Clara", "Dina", "Eva", "Fiona", "Gina", "Helen", "Iris", "Julia", "Kara", "Lucy", "Nina",
         "Paula", "Rita", "Sara", "Tina", "Adam", "Chris", "Dan", "Eli", "Finn", "George", "Harry", "Ian", "Jon", "Kyle",
         "Luke", "Matt", "Nick", "Owen", "Pete", "Ryan", "Tom", "Will"]
LAST = ["Adams", "Baker", "Bell", "Brooks", "Cole", "Cook", "Cox", "Davies", "Ellis", "Fox", "Gray", "Hall", "Hill",
        "Hughes", "James", "Kelly", "King", "Lewis", "Mason", "Mills", "Morgan", "Murphy", "Owens", "Price", "Reed",
        "Reid", "Shaw", "Stone", "Walsh", "Ward", "Webb", "West", "Wood", "Young"]
MIXES = [["Women's haircut", "Blow dry"], ["Colour", "Women's haircut"], ["Highlights", "Toner"],
         ["Women's haircut", "Treatment"], ["Blow dry", "Women's haircut"]]
taken = {c[0] for c in clients}
while len(clients) < 145:
    first, last = rng.choice(FIRST), rng.choice(LAST)
    name = f"{first} {last}"
    if name in taken:
        continue
    taken.add(name)
    man = FIRST.index(first) >= 30
    k = len(clients)
    clients.append((name, f"555 {200 + k:04d}", f"{first.lower()}.{last[0].lower()}@example.com" if rng.random() < 0.4 else None,
                    D(rng.randint(1965, 2004), rng.randint(1, 12), rng.randint(1, 28)) if rng.random() < 0.3 else None,
                    D(rng.randint(2021, Y - 1), rng.randint(1, 12), rng.randint(1, 28)), None,
                    rng.choice([21, 28, 28, 35]) if man else rng.choice([28, 35, 35, 42, 49, 56]),
                    ["Men's haircut"] if man else rng.choice(MIXES),
                    D(Y, rng.randint(4, 6), rng.randint(1, 28)) if k % 40 == 0 else None))

visits = []   # date, client, service
load = {}


def book(d, name, service, cap=7):
    """Put a visit on the first open day from d that still has room."""
    d = open_day(d)
    while load.get(d, 0) >= cap:
        d = open_day(d + dt.timedelta(days=1))
    load[d] = load.get(d, 0) + 1
    visits.append([d, name, service])
    return d


# Package clients and a new client first, so their dates stay exact.
for d in [D(Y, 6, 30), D(Y, 7, 21), D(Y, 8, 11), D(Y, 9, 1), D(Y, 9, 22), D(Y, 10, 13)]:
    book(d, "Emma Johnson", "Blow dry", cap=99)
for d in [D(Y, 6, 9), D(Y, 7, 7), D(Y, 8, 4), D(Y, 9, 1)]:
    book(d, "Maya Patel", "Treatment", cap=99)
book(D(Y, 9, 5), "Yara Haddad", "Women's haircut", cap=99)
book(D(Y, 10, 10), "Yara Haddad", "Highlights", cap=99)
book(D(Y, 9, 8), "Chloe Martin", "Blow dry", cap=99)
end = today + dt.timedelta(days=21)
for name, *_, every, kinds, stopped in clients:
    if not every:
        continue
    d = D(Y, 1, 6) + dt.timedelta(days=rng.randrange(every))
    while d <= (stopped or end):
        d = book(d, name, kinds[0] if len(kinds) == 1 or rng.random() < 0.65 else kinds[1])
        if name == "Noah Taylor":
            visits.append([d, name, "Kids haircut"])
        d = d + dt.timedelta(days=every + rng.randint(-4, 4))
visits.sort(key=lambda v: (v[0], v[1]))

appointments = []
by_day = {}
for d, name, service in visits:
    by_day.setdefault(d, []).append((name, service))
yesterday_open = today - dt.timedelta(days=1)
while yesterday_open.weekday() in (0, 6):
    yesterday_open -= dt.timedelta(days=1)
for d in sorted(by_day):
    t = dt.datetime(2000, 1, 1, 9, 0)
    for k, (name, service) in enumerate(by_day[d]):
        status, discount, paid, method, note = "Done", None, None, None, None
        if d >= today:
            status = "Booked"
        elif d == yesterday_open and k == len(by_day[d]) - 1:
            status = "Booked"                                  # never updated after the visit
        elif rng.random() < 0.035:
            status = "No-show"
        elif rng.random() < 0.03:
            status = "Cancelled"
        if d < today and name in ("Emma Johnson", "Maya Patel", "Rosa Diaz", "Wendy Scott"):
            status = "Done"                                    # keep the package and money examples exact
        if status == "Done":
            method = rng.choices(["Card", "Cash", "Transfer", "App"], [60, 25, 10, 5])[0]
            if rng.random() < 0.04:
                discount, note = 10, "Friends and family"
            paid = price[service] - (discount or 0)
        if name == "Emma Johnson" and d >= D(Y, 7, 18) or name == "Maya Patel" or \
                name == "Chloe Martin" and service == "Blow dry":
            method, paid = ("Package", None) if status == "Done" else (None, None)
        appointments.append([d, t.time(), name, service, status, discount, paid, method, note])
        t += dt.timedelta(minutes=minutes[service] + 10)
        if t.hour >= 18:
            t = dt.datetime(2000, 1, 1, 9, 15)
# Money still owed by two clients.
for row in reversed(appointments):
    if row[2] == "Rosa Diaz" and row[4] == "Done":
        row[3], row[5], row[6], row[8] = "Colour", None, 55, "Pays the rest next time"
        break
for row in reversed(appointments):
    if row[2] == "Wendy Scott" and row[4] == "Done":
        row[6], row[7], row[8] = None, None, "Forgot her card"
        break
packages = [(D(Y, 5, 30), "Maya Patel", "Treatment card, 4 sessions", 4, 100, "Card"),
            (D(Y, 7, 18), "Emma Johnson", "Blow dry card, 5 sessions", 5, 150, "Transfer"),
            (D(Y, 9, 1), "Chloe Martin", "Blow dry card, 5 sessions", 5, 150, "Card")]

wb = Workbook()


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


def status_colours(ws, rng_, first, past_date=None):
    if past_date:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'AND({first}="Booked",{past_date}<TODAY())'],
                                                        fill=fill(WARN), font=Font(name=F, bold=True, color="8A4B00")))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="Done"'], fill=fill(OK)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="Booked"'], fill=fill(INFO)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="No-show"'], fill=fill(BAD)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="Cancelled"'], fill=fill(GREY),
                                                    font=Font(name=F, color="9AA7A9")))


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


# ---------------------------------------------------------------- Appointments
ap = wb.active
ap.title = "Appointments"
sheet_base(ap, "Appointments", "One line per appointment. Change the status to Done after the visit and type what was paid.",
           [3, 13, 8, 20, 18, 11, 10, 10, 10, 11, 11, 26], rows=A1 + 2)
header(ap, 5, 2, ["Date", "Time", "Client", "Service", "Status", "Price", "Discount", "Paid", "Paid with", "Still to pay",
                  "Note"])
for r in range(A0, A1 + 1):
    a = appointments[r - A0] if r - A0 < len(appointments) else (None,) * 9
    d_, t_, name, service, status, discount, paid, method, note = a
    style(ap.cell(r, 2, d_), True, DATE)
    style(ap.cell(r, 3, t_), True, TIME, align="center")
    style(ap.cell(r, 4, name), True)
    style(ap.cell(r, 5, service), True)
    style(ap.cell(r, 6, status), True, align="center")
    style(ap.cell(r, 7, f'=IF(E{r}="","",MAX(0,IFERROR(INDEX({SV("C")},MATCH(E{r},{SV("B")},0)),0)-N(H{r})))'), False, MONEY)
    style(ap.cell(r, 8, discount), True, MONEY)
    style(ap.cell(r, 9, paid), True, MONEY)
    style(ap.cell(r, 10, method), True, align="center")
    style(ap.cell(r, 11, f'=IF(OR(B{r}="",F{r}<>"Done",J{r}="Package"),"",IF(N(G{r})-N(I{r})>0,N(G{r})-N(I{r}),""))'),
          False, MONEY, True)
    style(ap.cell(r, 12, note), True)
    ap.cell(r, 13, f'=IF(OR(B{r}="",D{r}="",F{r}="Cancelled"),"",B{r}+N(C{r})+ROW()/10000000)')   # sort key
    ap.cell(r, 14, f'=IF(AND(F{r}="Done",M{r}<>""),IF(COUNTIFS({AP("D")},D{r},{AP("F")},"Done",{AP("M")},">"&M{r})=0,1,0),0)')
    ap.cell(r, 15, f'=IF(AND(F{r}="Booked",M{r}<>"",N(B{r})>=TODAY()),IF(COUNTIFS({AP("D")},D{r},{AP("F")},"Booked",'
                   f'{AP("B")},">="&TODAY(),{AP("M")},"<"&M{r})=0,1,0),0)')
for col in "MNO":
    ap.column_dimensions[col].hidden = True
dropdown(ap, f"={CL('B')}", f"D{A0}:D{A1}")
dropdown(ap, f"={SV('B')}", f"E{A0}:E{A1}")
dropdown(ap, f'"{STATUSES}"', f"F{A0}:F{A1}")
dropdown(ap, f'"{METHODS}"', f"J{A0}:J{A1}")
status_colours(ap, f"F{A0}:F{A1}", f"$F{A0}", f"$B{A0}")
ap.conditional_formatting.add(f"K{A0}:K{A1}", FormulaRule(formula=[f"N($K{A0})>0"], fill=fill(BAD)))
ap.freeze_panes = "B6"

# ---------------------------------------------------------------- Clients
cl = wb.create_sheet("Clients")
sheet_base(cl, "Clients", "One line per client. Visits, payments and package sessions come from the other tabs.",
           [3, 18, 11, 22, 9, 12, 26, 8, 12, 12, 11, 10, 9, 22], rows=C1 + 2)
header(cl, 5, 2, ["Client", "Phone", "Email", "Birthday", "Client since", "Notes", "Visits", "Last visit", "Next booking",
                  "Paid in total", "Owes", "Sessions left", "Follow up"])
for r in range(C0, C1 + 1):
    c_ = clients[r - C0][:6] if r - C0 < len(clients) else (None,) * 6
    for col, (v, fmt) in enumerate(zip(c_, (None, None, None, "d mmm", DATE, None)), start=2):
        style(cl.cell(r, col, v), True, fmt, bold=col == 2)
    style(cl.cell(r, 8, f'=IF(B{r}="","",COUNTIFS({AP("D")},B{r},{AP("F")},"Done"))'), False, "0", align="center")
    style(cl.cell(r, 9, f'=IF(OR(B{r}="",N(H{r})=0),"",INT(SUMIFS({AP("M")},{AP("D")},B{r},{AP("N")},1)))'), False, DATE)
    nxt = f'SUMIFS({AP("M")},{AP("D")},B{r},{AP("O")},1)'
    style(cl.cell(r, 10, f'=IF(B{r}="","",IF({nxt}=0,"",INT({nxt})))'), False, DATE)
    style(cl.cell(r, 11, f'=IF(B{r}="","",SUMIFS({AP("I")},{AP("D")},B{r})+SUMIFS({PK("F")},{PK("C")},B{r}))'), False, MONEY)
    owes = f'SUMIFS({AP("K")},{AP("D")},B{r})'
    style(cl.cell(r, 12, f'=IF(B{r}="","",IF({owes}=0,"",{owes}))'), False, MONEY, True)
    bought = f'SUMIFS({PK("E")},{PK("C")},B{r})'
    used = (f'COUNTIFS({AP("D")},B{r},{AP("J")},"Package",{AP("F")},"Done")'
            f'+IF(Services!$I$6="Yes",COUNTIFS({AP("D")},B{r},{AP("J")},"Package",{AP("F")},"No-show"),0)')
    style(cl.cell(r, 13, f'=IF(OR(B{r}="",{bought}=0),"",{bought}-{used})'), False, "0", True, "center")
    style(cl.cell(r, 14, f'=IF(OR(B{r}="",I{r}=""),"",IF(AND(J{r}="",TODAY()-I{r}>=N(Services!$I$7)),'
                         f'"Not seen for "&(TODAY()-I{r})&" days",""))'), False)
    cl.cell(r, 15, f'=IF(N(E{r})=0,"",MONTH(E{r}))')                     # birthday month
cl.column_dimensions["O"].hidden = True
cl.conditional_formatting.add(f"L{C0}:L{C1}", FormulaRule(formula=[f"N($L{C0})>0"], fill=fill(BAD),
                                                          font=Font(name=F, bold=True, color="9B1C1C")))
cl.conditional_formatting.add(f"M{C0}:M{C1}", FormulaRule(formula=[f'AND($M{C0}<>"",N($M{C0})<=1)'], fill=fill(WARN)))
cl.conditional_formatting.add(f"N{C0}:N{C1}", FormulaRule(formula=[f'$N{C0}<>""'], fill=fill(INFO)))
cl.conditional_formatting.add(f"E{C0}:E{C1}", FormulaRule(formula=[f'AND(N($E{C0})>0,$O{C0}=MONTH(TODAY()))'],
                                                          fill=fill("FCE9B8"), font=Font(name=F, bold=True)))
cl.freeze_panes = "C6"

# ---------------------------------------------------------------- Packages
pk = wb.create_sheet("Packages")
sheet_base(pk, "Packages", "Prepaid sessions. Each appointment paid with Package uses one.",
           [3, 13, 18, 28, 10, 11, 11, 12, 24], rows=K1 + 2)
header(pk, 5, 2, ["Date bought", "Client", "Package", "Sessions", "Price paid", "Paid with", "Client has left", "Note"])
for r in range(K0, K1 + 1):
    p = packages[r - K0] if r - K0 < len(packages) else (None,) * 6
    for col, (v, fmt) in enumerate(zip(p, (DATE, None, None, "0", MONEY, None)), start=2):
        style(pk.cell(r, col, v), True, fmt, align="center" if col in (5, 7) else None)
    style(pk.cell(r, 8, f'=IF(C{r}="","",IFERROR(INDEX({CL("M")},MATCH(C{r},{CL("B")},0)),""))'), False, "0", True, "center")
    style(pk.cell(r, 9), True)
dropdown(pk, f"={CL('B')}", f"C{K0}:C{K1}")
dropdown(pk, '"Card,Cash,Transfer,App"', f"G{K0}:G{K1}")
pk.conditional_formatting.add(f"H{K0}:H{K1}", FormulaRule(formula=[f'AND($H{K0}<>"",N($H{K0})<=1)'], fill=fill(WARN)))
pk.freeze_panes = "B6"

# ---------------------------------------------------------------- Services
sv = wb.create_sheet("Services")
sheet_base(sv, "Services and settings", "Your services with their price and length. The appointments take the price from here.",
           [3, 22, 11, 10, 12, 13, 4, 34, 10], rows=S1 + 2)
header(sv, 5, 2, ["Service", "Price", "Minutes", "Done this year", "Paid this year"])
year_from = '">="&DATE(YEAR(TODAY()),1,1)'
for r in range(S0, S1 + 1):
    s = services[r - S0] if r - S0 < len(services) else (None,) * 3
    for col, (v, fmt) in enumerate(zip(s, (None, MONEY, "0")), start=2):
        style(sv.cell(r, col, v), True, fmt, bold=col == 2, align="center" if col == 4 else None)
    style(sv.cell(r, 5, f'=IF(B{r}="","",COUNTIFS({AP("E")},B{r},{AP("F")},"Done",{AP("B")},{year_from}))'), False, "0",
          align="center")
    style(sv.cell(r, 6, f'=IF(B{r}="","",SUMIFS({AP("I")},{AP("E")},B{r},{AP("B")},{year_from}))'), False, MONEY)
header(sv, 5, 8, ["Setting", ""])
sv.merge_cells("H5:I5")
for r, (label, value) in ((6, ("A no-show uses a package session", "Yes")),
                          (7, ("Follow up a client after this many days", 60))):
    style(sv.cell(r, 8, label), False)
    style(sv.cell(r, 9, value), True, "0" if r == 7 else None, True, "center")
dropdown(sv, '"Yes,No"', "I6")
sv.freeze_panes = "C6"

# ---------------------------------------------------------------- Week
wk = wb.create_sheet("Week")
sheet_base(wk, "Week", "The appointments of one week, in time order. Change the first day to see another week.",
           [3, 4] + [24] * 7, rows=24)
wk["B4"] = "Week starting"
wk["B4"].font = font(11, True)
style(wk["D4"], True, DATE, True, "center")
wk["D4"] = "=TODAY()-WEEKDAY(TODAY(),3)"
wk["E4"] = "Shows this week until you type another date."
wk["E4"].font = font(10, color=MUTED, italic=True)
header(wk, 6, 3, [""] * 7)
for k in range(7):
    col = L(3 + k)
    wk[f"{col}6"] = f"=$D$4+{k}"
    wk[f"{col}6"].number_format = "dddd d mmm"
    day = f"{col}$6"
    for s in range(10):
        r = 7 + s
        key = f"{L(18 + k)}{r}"                                            # hidden keys R..X
        nth = f'SMALL({AP("M")},COUNTIF({AP("M")},"<"&{day})+{s + 1})'
        wk[key] = f'=IFERROR(IF({nth}<{day}+1,{nth},""),"")'
        row = f'MATCH({key},{AP("M")},0)'
        c = wk.cell(r, 3 + k, f'=IF({key}="","",HOUR({key})&":"&RIGHT("0"&MINUTE({key}),2)&"  "&INDEX({AP("D")},{row})'
                               f'&", "&INDEX({AP("E")},{row}))')
        c.font = font(10)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = box
        wk[f"{L(26 + k)}{r}"] = f'=IF({key}="","",INDEX({AP("F")},{row}))'  # hidden status Z..AF
    style(wk.cell(17, 3 + k, f'=COUNTIF({col}7:{col}16,"?*")'), False, "0", True, "center")
    style(wk.cell(18, 3 + k, f'=SUMIFS({AP("I")},{AP("B")},{col}$6)'), False, MONEY, True, "center")
for r in range(7, 17):
    wk.row_dimensions[r].height = 30
style(wk.cell(17, 2, "Appointments"), False, bold=True)
style(wk.cell(18, 2, "Paid"), False, bold=True)
wk.column_dimensions["B"].width = 14
for c in range(18, 33):
    wk.column_dimensions[L(c)].hidden = True
grid = "C7:I16"
wk.conditional_formatting.add(grid, FormulaRule(formula=['Z7="Done"'], fill=fill(OK)))
wk.conditional_formatting.add(grid, FormulaRule(formula=['Z7="Booked"'], fill=fill(INFO)))
wk.conditional_formatting.add(grid, FormulaRule(formula=['Z7="No-show"'], fill=fill(BAD)))
wk.conditional_formatting.add("C6:I6", FormulaRule(formula=["C6=TODAY()"], fill=fill(WARN),
                                                   font=Font(name=F, bold=True, color=TEAL_D)))
key_row = 20
wk.cell(key_row, 2, "Colours:").font = font(10, True, MUTED)
for k, (label, colour) in enumerate([("Booked", INFO), ("Done", OK), ("No-show", BAD)]):
    c = wk.cell(key_row, 3 + k, label)
    c.fill = fill(colour)
    c.font = font(10)
    c.border = box
wk[f"C{key_row + 1}"] = "Cancelled appointments are left out. A day with more than 10 appointments shows the first 10."
wk[f"C{key_row + 1}"].font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Client dashboard", "This month, what is coming up, and who needs a message.",
           [3, 14, 8, 20, 20, 13, 13, 3, 16, 14, 14, 16, 3], rows=60)
ms, me = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "EOMONTH(TODAY(),0)"
this_month = lambda rng_: f'{rng_},">="&{ms},{rng_},"<="&{me}'
tile(db, "B", 5, "Paid this month",
     f'=SUMIFS({AP("I")},{this_month(AP("B"))})+SUMIFS({PK("F")},{this_month(PK("B"))})', MONEY)
db.merge_cells("B5:C5")
db.merge_cells("B6:C6")
tile(db, "D", 5, "Done this month", f'=COUNTIFS({AP("F")},"Done",{this_month(AP("B"))})', "0")
tile(db, "E", 5, "Booked ahead", f'=COUNTIFS({AP("F")},"Booked",{AP("B")},">="&TODAY())', "0")
tile(db, "F", 5, "No-shows", f'=COUNTIFS({AP("F")},"No-show",{this_month(AP("B"))})', "0", "B4541F")
tile(db, "G", 5, "Owed to you", f'=SUM({AP("K")})', MONEY, "B4541F")
tile(db, "I", 5, "Clients", f'=COUNTIF({CL("B")},"?*")', "0")
tile(db, "J", 5, "Visits this year", f'=COUNTIFS({AP("F")},"Done",{AP("B")},">="&DATE(YEAR(TODAY()),1,1))', "#,##0")
tile(db, "K", 5, "Paid this year",
     f'=SUMIFS({AP("I")},{AP("B")},">="&DATE(YEAR(TODAY()),1,1))+SUMIFS({PK("F")},{PK("B")},">="&DATE(YEAR(TODAY()),1,1))',
     MONEY)
db.merge_cells("K5:L5")
db.merge_cells("K6:L6")
for a in ("C5", "C6", "L5", "L6"):
    db[a].border = box

db["B8"] = "Coming up"
db["B8"].font = font(13, True, TEAL_D)
header(db, 9, 2, ["Date", "Time", "Client", "Service"])
for k in range(10):
    r = 10 + k
    key = f"O{r}"
    db[key] = f'=IFERROR(SMALL({AP("M")},COUNTIF({AP("M")},"<"&NOW())+{k + 1}),"")'
    row = f'MATCH({key},{AP("M")},0)'
    style(db.cell(r, 2, f'=IF({key}="","",INT({key}))'), False, "ddd d mmm")
    style(db.cell(r, 3, f'=IF({key}="","",INDEX({AP("C")},{row}))'), False, TIME, align="center")
    style(db.cell(r, 4, f'=IF({key}="","",INDEX({AP("D")},{row}))'), False, bold=True)
    style(db.cell(r, 5, f'=IF({key}="","",INDEX({AP("E")},{row}))'), False)
db.column_dimensions["O"].hidden = True

stale = f'COUNTIFS({AP("F")},"Booked",{AP("B")},"<"&TODAY())'
owe_n, owe_sum = f'COUNTIFS({CL("L")},">0")', f'SUM({CL("L")})'
low = f'COUNTIFS({CL("M")},"<=1")'
bdays = f'COUNTIFS({CL("O")},MONTH(TODAY()))'
follow = f'COUNTIFS({CL("N")},"Not seen*")'
notes = [
    (f'=IF({stale}=0,"Every past appointment has its status",IF({stale}=1,"1 past appointment is",'
     f'{stale}&" past appointments are")&" still marked Booked")', WARN, "Every"),
    (f'=IF({owe_n}=0,"Nobody owes you money",IF({owe_n}=1,"1 client owes you ",{owe_n}&" clients owe you ")'
     f'&FIXED({owe_sum},2))', BAD, "Nobody"),
    (f'=IF({low}=0,"No package is running out",IF({low}=1,"1 client has",{low}&" clients have")'
     f'&" 1 package session left or none")', WARN, "No package"),
    (f'=IF({follow}=0,"Every regular client has been back lately",IF({follow}=1,"1 client has not",'
     f'{follow}&" clients have not")&" been back in "&Services!$I$7&" days")', INFO, "Every"),
    (f'=IF({bdays}=0,"No birthdays this month",IF({bdays}=1,"1 client has a birthday",'
     f'{bdays}&" clients have a birthday")&" this month")', "FCE9B8", "No birth"),
]
header(db, 9, 9, ["Needs attention", "", "", ""])
db.merge_cells("I9:L9")
for k, (formula, colour, calm) in enumerate(notes):
    r = 10 + k
    db.merge_cells(f"I{r}:L{r}")
    c = db.cell(r, 9, formula)
    c.fill = fill(colour)
    c.font = font(11)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    for col in "IJKL":
        db[f"{col}{r}"].border = box
    db.conditional_formatting.add(f"I{r}", FormulaRule(formula=[f'LEFT(I{r},{len(calm)})="{calm}"'], fill=fill(OK)))

db["B22"] = "Year"
db["B22"].font = font(11, True)
style(db["C22"], True, "0", True, "center")
db["C22"] = Y
header(db, 24, 2, ["Month", "Done", "No-shows", "Appointments paid", "Packages sold", "Total"])
for m in range(12):
    r = 25 + m
    a, b = f"DATE($C$22,{m + 1},1)", f"EOMONTH(DATE($C$22,{m + 1},1),0)"
    within = lambda rng_: f'{rng_},">="&{a},{rng_},"<="&{b}'
    style(db.cell(r, 2, MONTHS[m]), False, bold=True, align="center")
    style(db.cell(r, 3, f'=COUNTIFS({AP("F")},"Done",{within(AP("B"))})'), False, "0", align="center")
    style(db.cell(r, 4, f'=COUNTIFS({AP("F")},"No-show",{within(AP("B"))})'), False, "0", align="center")
    style(db.cell(r, 5, f'=SUMIFS({AP("I")},{within(AP("B"))})'), False, MONEY)
    style(db.cell(r, 6, f'=SUMIFS({PK("F")},{within(PK("B"))})'), False, MONEY)
    style(db.cell(r, 7, f"=E{r}+F{r}"), False, MONEY, True)
style(db.cell(37, 2, "Year"), False, bold=True, align="center")
for col, fmt in (("C", "#,##0"), ("D", "#,##0"), ("E", MONEY), ("F", MONEY), ("G", MONEY)):
    style(db[f"{col}37"], False, fmt, True, "center" if fmt == "#,##0" else None)
    db[f"{col}37"] = f"=SUM({col}25:{col}36)"
db.conditional_formatting.add("B25:G36", FormulaRule(formula=["AND($C$22=YEAR(TODAY()),ROW()-24=MONTH(TODAY()))"],
                                                     fill=fill("E6F2F2")))
chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.overlap = 100
chart.title = "Paid by month"
chart.height, chart.width = 8.5, 15
chart.add_data(Reference(db, min_col=5, max_col=6, min_row=24, max_row=36), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=25, max_row=36))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[1].graphicalProperties.solidFill = "E0A21B"
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "I22")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells. Everything else is calculated.", [3, 6, 106])
steps = [
    ("1", "Services: your services with their price and length, and the two settings on the right."),
    ("2", "Clients: one line per client, with the phone, birthday and notes you want to keep."),
    ("3", "Appointments: one line per booking. After the visit, set the status to Done and type what was paid."),
    ("4", "Packages: a client who prepays several sessions gets one line here. Their visits are then paid with Package."),
    ("5", "Dashboard and Week: what is coming up, what you earned, who owes you and who has not been back."),
]
for k, (n, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The price of an appointment comes from Services, minus any discount you type.",
                          "An appointment whose date has passed but is still Booked turns orange, so you remember to update it.",
                          "The example clients and appointments show how it works. Delete them and add your own.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(16 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Client-Appointment-Tracker.xlsx")
wb.save(out)
print("saved", out, len(clients), "clients", len(appointments), "appointments", len(packages), "packages")
