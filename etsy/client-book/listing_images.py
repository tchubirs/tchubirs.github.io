"""Etsy listing photos for the Client and Appointment Tracker (run build.py first)."""
import datetime as dt, os, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Client-Appointment-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
os.makedirs(tmp, exist_ok=True)
r = Renders(xlsx, tmp)
img_dir = os.path.join(HERE, "images")

# A copy whose Week tab shows the last past week with a no-show, so every colour appears.
wb = load_workbook(xlsx)
ap = wb["Appointments"]
rows = [(row[0].value, row[4].value, row[0].row) for row in ap.iter_rows(min_row=6, max_row=2005, min_col=2, max_col=6)
        if row[0].value]
today = dt.date.today()
monday = today - dt.timedelta(days=today.weekday())
shows = [d.date() for d, status, _ in rows if status == "No-show" and d.date() < monday]
wb["Week"]["D4"] = max(shows) - dt.timedelta(days=max(shows).weekday())
# Appointments close-up: the rows around today.
first = next(n for d, _, n in rows if d.date() >= today - dt.timedelta(days=5))
wb.save(os.path.join(tmp, "week-copy.xlsx"))
week = Renders(os.path.join(tmp, "week-copy.xlsx"), tmp)

img = canvas(dark=True)
title(img, "Client and\nappointment tracker", "For hair stylists, trainers, tutors and anyone who works one to one",
      dark=True, width=2300)
laptop(img, r.shot("Client dashboard"), 600, 700, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-week.jpg"), "Your week, in time order",
       "Green when done, blue when booked, red for a no-show, with what each day brought in",
       week.shot("Week", until="Paid"))

close = Renders(xlsx, tmp, hide={"Clients": ["C", "D"]},
                hide_rows={"Clients": range(31, 306), "Appointments": list(range(6, first)) + list(range(first + 26, 2006))})
detail(os.path.join(img_dir, "03-clients.jpg"), "Every client on one line",
       "Visits, last visit, next booking, what they owe, package sessions left, and who has not been back",
       close.shot("Clients"),
       [("#F9C9C4", "Owes you"), ("#FFE2B8", "Package almost used"), ("#E3E8FF", "Time to follow up")])

detail(os.path.join(img_dir, "04-appointments.jpg"), "One line per appointment",
       "The price comes from your services. After the visit, mark it done and type what was paid.",
       close.shot("Appointments"),
       [("#D8F0DC", "Done"), ("#E3E8FF", "Booked"), ("#FFE2B8", "Past but still booked")])

img = canvas()
y = title(img, "Inside the file", "Six tabs and a Start Here page. You type in the yellow cells and the rest is calculated.")
grid(img, [
    (content(r.page("Client dashboard")), "Dashboard", "This month, coming up, who needs a message"),
    (content(week.page("Week")), "Week", "Seven days of appointments in time order"),
    (content(r.page("Clients")), "Clients", "Visits, payments, packages, birthdays"),
    (content(r.page("Services and settings")), "Services", "Prices, minutes and two settings"),
], y + 80)
save(img, os.path.join(img_dir, "05-whats-inside.jpg"))
