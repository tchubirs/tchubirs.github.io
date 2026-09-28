"""Chores and Allowance Tracker for kids (Excel + Google Sheets).

Each child's chores with points and the days they are expected; a This week chart to tick and to print for the
fridge; the week's pay from a base amount and the points; a Money log where pocket money is split into Save, Spend and
Give jars by each child's own shares; and a goal for each child with how long it will take at the recent rate. Only
functions both apps have.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sheetkit import (BAD, DATE, MONEY, MUTED, OK, TEAL, TEAL_D, F,  # noqa: E402
                      bar, box, fill, font, header, sheet_base, style)

today = dt.date.today()
K0, K1 = 6, 11         # kids
C0, C1 = 6, 45         # chores (the This week chart has the same rows)
G0, G1 = 6, 505        # money
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
KD = lambda col: f"Kids!${col}${K0}:${col}${K1}"
CH = lambda col: f"Chores!${col}${C0}:${col}${C1}"
WK = lambda col: f"'This week'!${col}${C0}:${col}${C1}"
MN = lambda col: f"Money!${col}${G0}:${col}${G1}"
CUR, WEEK = "Settings!$C$4", "Settings!$C$6"
PICK = "'Fridge chart'!$B$2"
F0, F1 = 6, 17         # rows on the fridge chart
MON_LIST = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'
SOFT_TEAL, DONE = "DDEFEC", "BFE3C7"


def short(d):
    return f'DAY({d})&" "&CHOOSE(MONTH({d}),{MON_LIST})'


monday = today - dt.timedelta(days=today.weekday())
D = lambda n: today + dt.timedelta(days=n)

# ---------------------------------------------------------------- example data (a made-up family)
kids = [
    # name, age, base a week, per point, save %, give %, goal, goal price
    ("Mia", 9, 2.00, 0.25, 0.40, 0.10, "Art kit", 30),
    ("Leo", 7, 1.50, 0.20, 0.30, 0.10, "Building set", 45),
    ("Ava", 12, 3.00, 0.30, 0.50, 0.10, "Headphones", 60),
]
E7 = "xxxxxxx"
chores = [
    # chore, child, points, days (Mon..Sun as x or .), notes
    ("Make bed", "Mia", 1, E7, None), ("Feed the cat", "Mia", 2, "xxxxx..", "Before school"),
    ("Set the table", "Mia", 1, E7, None), ("Tidy room", "Mia", 3, ".....x.", None),
    ("Water the plants", "Mia", 2, "..x...x", None),
    ("Make bed", "Leo", 1, E7, None), ("Put toys away", "Leo", 1, E7, "Before bath"),
    ("Help with laundry", "Leo", 2, ".....x.", None), ("Feed the fish", "Leo", 1, "x.x.x..", "A pinch only"),
    ("Make bed", "Ava", 1, E7, None), ("Empty the dishwasher", "Ava", 2, "xxxxx..", None),
    ("Walk the dog", "Ava", 3, ".x.x.x.", None), ("Homework before screens", "Ava", 2, "xxxx...", None),
    ("Vacuum the stairs", "Ava", 4, "......x", None),
]
# This week's ticks: most expected days done, a few missed.
missed = {(1, 3), (4, 6), (6, 2), (8, 4), (11, 1), (12, 2), (13, 3)}
ticks = []
for i, (_, _, _, days, _) in enumerate(chores):
    ticks.append(["x" if days[d] == "x" and (i, d) not in missed else None for d in range(7)])
ticks[4][0] = "x"                                         # an extra, on a day it was not asked for
money = []
for w in range(6, 0, -1):
    pay_day = monday - dt.timedelta(days=7 * w - 6)     # the Sunday of each past week
    for name, base, rate, pts in (("Mia", 2.00, 0.25, 22 - w % 3), ("Leo", 1.50, 0.20, 16 - w % 2),
                                  ("Ava", 3.00, 0.30, 26 - w % 4)):
        money.append((pay_day, name, "Pocket money", round(base + rate * pts, 2), None, "Split"))
money += [
    (monday - dt.timedelta(days=30), "Leo", "Birthday money from Grandma", 20, None, "Save"),
    (monday - dt.timedelta(days=26), "Mia", "Comic", None, 4.50, "Spend"),
    (monday - dt.timedelta(days=19), "Ava", "Cinema with friends", None, 8.00, "Spend"),
    (monday - dt.timedelta(days=16), "Leo", "Sweets", None, 1.20, "Spend"),
    (monday - dt.timedelta(days=12), "Mia", "Animal shelter", None, 3.00, "Give"),
    (monday - dt.timedelta(days=9), "Ava", "Sold old books", 6.00, None, "Split"),
    (monday - dt.timedelta(days=5), "Ava", "School charity day", None, 4.00, "Give"),
    (monday - dt.timedelta(days=3), "Leo", "Stickers", None, 2.50, "Spend"),
]
money.sort(key=lambda x: x[0])

wb = Workbook()


def dropdown(ws, formula, cells, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(cells)


def tile(ws, col, row, label, formula, fmt, colour=TEAL):
    ws[f"{col}{row}"] = label
    ws[f"{col}{row}"].font = font(10, True, MUTED)
    ws[f"{col}{row + 1}"] = formula
    ws[f"{col}{row + 1}"].number_format = fmt
    ws[f"{col}{row + 1}"].font = font(16, True, colour)
    ws[f"{col}{row + 1}"].alignment = Alignment(horizontal="right")
    for r in (row, row + 1):
        ws[f"{col}{r}"].border = box
    ws.row_dimensions[row + 1].height = 34


# ---------------------------------------------------------------- Settings
se = wb.active
se.title = "Settings"
sheet_base(se, "Settings", "The currency and the week the chart shows.", [3, 26, 16, 4, 64], rows=10)
for r, label, value, fmt, typed in ((4, "Currency", "USD", None, True),
                                    (5, "Show another week", None, "ddd d mmm yyyy", True),
                                    (6, "The week starts on", "=IF(C5=\"\",TODAY(),C5)-WEEKDAY(IF(C5=\"\",TODAY(),C5),3)",
                                     "ddd d mmm yyyy", False)):
    se[f"B{r}"] = label
    se[f"B{r}"].font = font(11, True)
    style(se[f"C{r}"], typed, fmt, True, "center")
    se[f"C{r}"] = value
for k, text in enumerate(["Amounts have no currency sign, so any currency works.",
                          "Leave it empty for the week we are in. Type any day to see and print its week.",
                          "The Monday of that week."]):
    se.cell(4 + k, 5, text).font = font(10, color=MUTED, italic=True)

# ---------------------------------------------------------------- Kids
kd = wb.create_sheet("Kids")
sheet_base(kd, "Kids", "Each child's pocket money: a base amount a week and so much a point, split into three jars.",
           [3, 14, 7, 11, 11, 9, 9, 9, 18, 11, 11, 11, 11, 20, 14], rows=K1 + 3)
header(kd, 5, 2, ["Name", "Age", "Base a week", "Per point", "Save", "Give", "Spend", "Goal", "Goal price", "Save jar",
                  "Spend jar", "Give jar", "Goal saved", "Weeks to go"])
for r in range(K0, K1 + 1):
    x = kids[r - K0] if r - K0 < len(kids) else (None,) * 8
    name, age, base, rate, save, give, goal, price = x
    for c, v, fmt in ((2, name, None), (3, age, "0"), (4, base, MONEY), (5, rate, MONEY), (6, save, "0%"),
                      (7, give, "0%"), (9, goal, None), (10, price, MONEY)):
        style(kd.cell(r, c, v), True, fmt, c == 2, "center" if c in (3, 6, 7) else None)
    style(kd.cell(r, 8, f'=IF(B{r}="","",1-N(F{r})-N(G{r}))'), False, "0%", align="center")
    jar = lambda col: f'=IF(B{r}="","",SUMIFS({MN(col)},{MN("C")},B{r}))'
    style(kd.cell(r, 11, jar("H")), False, MONEY)
    style(kd.cell(r, 12, jar("I")), False, MONEY)
    style(kd.cell(r, 13, jar("J")), False, MONEY)
    style(kd.cell(r, 14, f'=IF(OR(B{r}="",N(J{r})=0),"",{bar(f"K{r}/J{r}", 12)})'), False)
    kd.cell(r, 14).font = Font(name=F, size=9, color=TEAL)
    kd.row_dimensions[r].height = 20
    for c in range(2, 16):
        kd.cell(r, c).alignment = Alignment(horizontal=kd.cell(r, c).alignment.horizontal, vertical="center")
    recent = f'SUMIFS({MN("K")},{MN("C")},B{r},{MN("B")},">"&(TODAY()-28))'
    kd[f"Q{r}"] = f'=IF(B{r}="","",{recent}/4)'                                             # Q: saved a week lately
    style(kd.cell(r, 15, f'=IF(OR(B{r}="",N(J{r})=0),"",IF(K{r}>=J{r},"Reached",IF(N(Q{r})<=0,"",'
                         f'ROUNDUP(ROUND((J{r}-K{r})/Q{r},6),0))))'), False, "0", align="center")
kd.column_dimensions["Q"].hidden = True
kd.conditional_formatting.add(f"O{K0}:O{K1}", FormulaRule(formula=[f'O{K0}="Reached"'], fill=fill(OK)))
for rng, top in ((f"H{K0}:H{K1}", f"H{K0}"), (f"K{K0}:M{K1}", f"K{K0}")):              # below zero
    kd.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({top}),{top}<0)"], fill=fill(BAD)))

# ---------------------------------------------------------------- Chores
ch = wb.create_sheet("Chores")
sheet_base(ch, "Chores", "Each chore once per child, with its points and an x on the days it is expected.",
           [3, 24, 12, 8] + [6] * 7 + [26, 10], rows=C1 + 2)
header(ch, 5, 2, ["Chore", "Child", "Points"] + DAYS + ["Notes", "A week"])
for r in range(C0, C1 + 1):
    x = chores[r - C0] if r - C0 < len(chores) else (None, None, None, ".......", None)
    chore, child, pts, days, note = x
    style(ch.cell(r, 2, chore), True, bold=True)
    style(ch.cell(r, 3, child), True)
    style(ch.cell(r, 4, pts), True, "0", align="center")
    for d in range(7):
        style(ch.cell(r, 5 + d, "x" if days[d] == "x" else None), True, align="center")
    style(ch.cell(r, 12, note), True)
    style(ch.cell(r, 13, f'=IF(B{r}="","",N(D{r})*COUNTA(E{r}:K{r}))'), False, "0", align="center")
    ch[f"N{r}"] = f'=IF(B{r}="","",COUNTA(E{r}:K{r}))'                                    # N: times a week
    ch[f"O{r}"] = f"=IF(OR(B{r}=\"\",C{r}=\"\"),\"\",IF(C{r}={PICK},ROW(),\"\"))"          # O: on the fridge chart
for c in "NO":
    ch.column_dimensions[c].hidden = True
dropdown(ch, f"={KD('B')}", f"C{C0}:C{C1}")

# ---------------------------------------------------------------- This week (the chart to tick and to print)
wk = wb.create_sheet("This week")
sheet_base(wk, "", "", [3, 12, 24, 8] + [9] * 7 + [8, 9], rows=C1 + 2)
wk["B2"] = f'="Chores for the week of "&{short(WEEK)}&" to "&{short(f"({WEEK}+6)")}'
wk["B2"].font = font(20, True, TEAL_D)
wk["B3"] = "An x in the box when it is done. Shaded boxes are the days it is expected. To print, use Fridge chart."
wk["B3"].font = font(11, color=MUTED, italic=True)
header(wk, 5, 2, ["Child", "Chore", "Points"] + DAYS + ["Done", "Earned"])
for d in range(7):
    wk.cell(5, 5 + d, f'="{DAYS[d]} "&DAY({WEEK}+{d})')
for r in range(C0, C1 + 1):
    x = ticks[r - C0] if r - C0 < len(ticks) else [None] * 7
    style(wk.cell(r, 2, f'=IF(Chores!B{r}="","",IF(Chores!C{r}="","",Chores!C{r}))'), False, bold=True, align="left")
    style(wk.cell(r, 3, f'=IF(Chores!B{r}="","",Chores!B{r})'), False, align="left")
    style(wk.cell(r, 4, f'=IF(Chores!B{r}="","",N(Chores!D{r}))'), False, "0", align="center")
    for d in range(7):
        c = style(wk.cell(r, 5 + d, x[d]), False, align="center")
        c.font = font(12, True, TEAL_D)
    style(wk.cell(r, 12, f'=IF(C{r}="","",COUNTA(E{r}:K{r}))'), False, "0", align="center")
    style(wk.cell(r, 13, f'=IF(C{r}="","",L{r}*D{r})'), False, "0", True, "center")
    wk.row_dimensions[r].height = 20
for d in range(7):
    col = wk.cell(1, 5 + d).column_letter
    rng = f"{col}{C0}:{col}{C1}"
    wk.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($C{C0}<>"",{col}{C0}<>"")'], fill=fill(DONE)))
    wk.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($C{C0}<>"",Chores!{col}{C0}<>"")'], fill=fill(SOFT_TEAL)))

# ---------------------------------------------------------------- Fridge chart (one child, to print)
fc = wb.create_sheet("Fridge chart")
sheet_base(fc, "", "", [3, 30, 8] + [11] * 7 + [8], rows=F1 + 5)
style(fc["B2"], True, bold=True)
fc["B2"] = kids[0][0]
fc["B2"].font = font(24, True, TEAL_D)
fc["B2"].alignment = Alignment(horizontal="left", vertical="center")
dropdown(fc, f"={KD('B')}", "B2")
fc["D2"] = f'="Chores for the week of "&{short(WEEK)}&" to "&{short(f"({WEEK}+6)")}'
fc["D2"].font = font(15, True, TEAL)
fc["D2"].alignment = Alignment(vertical="center")
fc.row_dimensions[2].height = 42
fc["B3"] = "An x in the box when it is done. The shaded boxes are the days to do it."
fc["B3"].font = font(12, color=MUTED, italic=True)
header(fc, 5, 2, ["Chore", "Points"] + DAYS + ["Done"])
for d in range(7):
    fc.cell(5, 4 + d, f'="{DAYS[d]} "&DAY({WEEK}+{d})')
for r in range(F0, F1 + 1):
    fc[f"N{r}"] = f"=IFERROR(SMALL({CH('O')},{r - F0 + 1}),\"\")"                          # N: the row on Chores
    at = lambda sheet, col: f"INDEX({sheet}!${col}$1:${col}${C1},$N{r})"
    style(fc.cell(r, 2, f'=IF($N{r}="","",{at("Chores", "B")})'), False, bold=True, align="left").font = font(14, True)
    style(fc.cell(r, 3, f'=IF($N{r}="","",N({at("Chores", "D")}))'), False, "0", align="center").font = font(13)
    for d in range(7):
        col = "EFGHIJK"[d]
        tick = at("'This week'", col)
        c = style(fc.cell(r, 4 + d, f'=IF($N{r}="","",IF({tick}="","",{tick}))'), False, align="center")
        c.font = font(18, True, TEAL_D)
        fc.cell(r, 16 + d, f'=IF($N{r}="",0,IF({at("Chores", col)}="",0,1))')               # P to V: expected
    done = at("'This week'", "L")
    style(fc.cell(r, 11, f'=IF($N{r}="","",{done})'), False, "0", True, "center").font = font(13, True)
    fc.row_dimensions[r].height = 30
for c in "NOPQRSTUV":
    fc.column_dimensions[c].hidden = True
fc.conditional_formatting.add(f"D{F0}:J{F1}", FormulaRule(formula=[f'AND(ISNUMBER($N{F0}),D{F0}<>"")'], fill=fill(DONE)))
fc.conditional_formatting.add(f"D{F0}:J{F1}", FormulaRule(formula=[f"AND(ISNUMBER($N{F0}),P{F0}=1)"], fill=fill(SOFT_TEAL)))
fc[f"B{F1 + 1}"] = (f'=IF(COUNT({CH("O")})>{F1 - F0 + 1},"And "&(COUNT({CH("O")})-{F1 - F0 + 1})&'
                    f'" more on This week","")')
fc[f"B{F1 + 1}"].font = font(10, color=MUTED, italic=True)
kid_at = lambda col: f"IFERROR(N(INDEX({KD(col)},MATCH($B$2,{KD('B')},0))),0)"
for k, (label, formula, fmt) in enumerate((
        ("Points this week", f'=IF($B$2="","",SUMIFS({WK("M")},{WK("B")},$B$2))', "0"),
        (f'="Pocket money this week"&IF({CUR}="",""," ("&{CUR}&")")',
         f'=IF($B$2="","",{kid_at("D")}+{kid_at("E")}*N(C{F1 + 2}))', MONEY))):
    r = F1 + 2 + k
    fc[f"B{r}"] = label
    fc[f"B{r}"].font = font(13, True, MUTED)
    fc[f"B{r}"].alignment = Alignment(horizontal="right")
    fc[f"C{r}"] = formula
    fc[f"C{r}"].number_format = fmt
    fc[f"C{r}"].font = font(14, True, TEAL_D)
    fc[f"C{r}"].alignment = Alignment(horizontal="right")
    fc.row_dimensions[r].height = 24
fc.page_setup.fitToHeight = 1

# ---------------------------------------------------------------- Money
mn = wb.create_sheet("Money")
sheet_base(mn, "Money", "Money in and out for each child. Split shares it by the child's jars; or name one jar.",
           [3, 14, 12, 28, 10, 10, 10], rows=G1 + 2)
header(mn, 5, 2, ["Date", "Child", "What", "In", "Out", "Jar"])
for r in range(G0, G1 + 1):
    x = money[r - G0] if r - G0 < len(money) else (None,) * 6
    for c, (v, fmt) in enumerate(zip(x, (DATE, None, None, MONEY, MONEY, None))):
        style(mn.cell(r, 2 + c, v), True, fmt, c == 1, "center" if c in (0, 5) else None)
    share = lambda col: f'IFERROR(N(INDEX({KD(col)},MATCH(C{r},{KD("B")},0))),0)'
    amount = f"(N(E{r})-N(F{r}))"
    split = f'OR(G{r}="",G{r}="Split")'
    mn[f"H{r}"] = f'=IF(C{r}="",0,IF({split},{amount}*{share("F")},IF(G{r}="Save",{amount},0)))'   # H: save jar
    mn[f"J{r}"] = f'=IF(C{r}="",0,IF({split},{amount}*{share("G")},IF(G{r}="Give",{amount},0)))'   # J: give jar
    mn[f"I{r}"] = f'=IF(C{r}="",0,{amount}-H{r}-J{r})'                                             # I: spend jar
    mn[f"K{r}"] = f'=IF(C{r}="",0,IF({split},N(E{r})*{share("F")},IF(G{r}="Save",N(E{r}),0)))'    # K: saved from money in
for c in "HIJK":
    mn.column_dimensions[c].hidden = True
dropdown(mn, f"={KD('B')}", f"C{G0}:C{G1}")
dropdown(mn, '"Split,Save,Spend,Give"', f"G{G0}:G{G1}")
mn.freeze_panes = "C6"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
sheet_base(db, "Chores and pocket money", "", [3, 16, 16, 16, 16, 3, 16, 14, 14, 14, 3], rows=40)
db["B3"] = (f'="The week of "&{short(WEEK)}&" to "&{short(f"({WEEK}+6)")}&", and the jars so far"&'
            f'IF({CUR}="",".",", in "&{CUR}&".")')
db["B3"].font = font(11, color=MUTED, italic=True)
tile(db, "B", 4, "Done this week", f'=SUM({WK("L")})&" of "&SUM({CH("N")})', "General")
tile(db, "C", 4, "Points", f"=SUM({WK('M')})", "0")
tile(db, "D", 4, "To pay this week", "=SUM(E9:E14)", MONEY)
tile(db, "E", 4, "In the jars", f"=SUM({KD('K')})+SUM({KD('L')})+SUM({KD('M')})", MONEY)
tile(db, "G", 4, "Save jars", f"=SUM({KD('K')})", MONEY)
tile(db, "H", 4, "Spend jars", f"=SUM({KD('L')})", MONEY)
tile(db, "I", 4, "Give jars", f"=SUM({KD('M')})", MONEY)
tile(db, "J", 4, "Goals reached", f'=COUNTIF({KD("O")},"Reached")&" of "&COUNT({KD("J")})', "General")

header(db, 8, 2, ["This week", "Done", "Points", "Earned"])
header(db, 8, 7, ["Jars", "Save", "Spend", "Give"])
header(db, 17, 2, ["Goals", "For", "Saved", "Weeks to go"])
for k in range(K1 - K0 + 1):
    r, kr = 9 + k, K0 + k
    name = f"Kids!$B${kr}"
    style(db.cell(r, 2, f'=IF({name}="","",{name})'), False, bold=True)
    style(db.cell(r, 3, f'=IF(B{r}="","",SUMIFS({WK("L")},{WK("B")},B{r}))'), False, "0", align="center")
    style(db.cell(r, 4, f'=IF(B{r}="","",SUMIFS({WK("M")},{WK("B")},B{r}))'), False, "0", align="center")
    style(db.cell(r, 5, f'=IF(B{r}="","",N(Kids!$D${kr})+N(Kids!$E${kr})*D{r})'), False, MONEY)
    style(db.cell(r, 7, f'=IF({name}="","",{name})'), False, bold=True)
    for c, col in ((8, "K"), (9, "L"), (10, "M")):
        style(db.cell(r, c, f'=IF(G{r}="","",Kids!${col}${kr})'), False, MONEY)
    g = 18 + k
    style(db.cell(g, 2, f'=IF(OR({name}="",Kids!$I${kr}=""),"",{name})'), False, bold=True)
    style(db.cell(g, 3, f'=IF(B{g}="","",Kids!$I${kr})'), False)
    style(db.cell(g, 4, f'=IF(B{g}="","",FIXED(Kids!$K${kr},2)&IF(N(Kids!$J${kr})=0,"",'
                        f'" of "&FIXED(Kids!$J${kr},2)))'), False, align="center")
    style(db.cell(g, 5, f'=IF(B{g}="","",Kids!$O${kr})'), False, "0", align="center")
db.conditional_formatting.add("E18:E23", FormulaRule(formula=['E18="Reached"'], fill=fill(OK)))

chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.overlap = 100
chart.title = "The jars"
chart.height, chart.width = 7, 12.5
chart.add_data(Reference(db, min_col=8, max_col=10, min_row=8, max_row=14), titles_from_data=True)
chart.set_categories(Reference(db, min_col=7, min_row=9, max_row=14))
for s, colour in zip(chart.series, (TEAL, "F2A65A", "B8D3CF")):
    s.graphicalProperties.solidFill = colour
chart.y_axis.number_format = "#,##0"
db.add_chart(chart, "G17")

# ---------------------------------------------------------------- Start Here
st = wb.create_sheet("Start Here")
sheet_base(st, "Start here", "Type only in the yellow cells, and tick the boxes on This week.", [3, 6, 106])
steps = [
    ("1", "Kids: each child with a base amount a week, so much a point, and the shares to save and to give."),
    ("2", "Chores: each chore for each child, its points and an x on the days it is expected."),
    ("3", "This week: an x in the box when a chore is done."),
    ("4", "Fridge chart: pick a child in the yellow box and print the chart. It shows the ticks made so far."),
    ("5", "Pay day: the Dashboard shows what each child earned. Add it on Money as Pocket money, jar Split."),
    ("6", "Money: gifts in, and what they spend or give out of a jar. Each child's goal fills from the Save jar."),
]
for k, (n_, text) in enumerate(steps):
    r = 5 + k * 2
    st.cell(r, 2, n_).font = font(18, True, TEAL)
    st.cell(r, 3, text).font = font(13)
for k, text in enumerate(["The example family and their money are made up. Delete them and add your own.",
                          "A new week: clear the boxes on This week. The money stays on the Money tab.",
                          "Works in Google Sheets and Microsoft Excel, in any currency."]):
    st.cell(18 + k * 2, 3, text).font = font(11, color=MUTED, italic=True)

order = ["Dashboard", "This week", "Fridge chart", "Chores", "Money", "Kids", "Settings", "Start Here"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Chores-Allowance-Tracker.xlsx")
wb.save(out)
print("saved", out, len(chores), "chores", len(money), "money lines")
