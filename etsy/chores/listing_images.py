"""Etsy listing photos for the Chores and Allowance Tracker (run build.py first)."""
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from listingkit import M, Renders, canvas, content, detail, grid, laptop, note, save, title  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
xlsx = os.path.join(HERE, "Chores-Allowance-Tracker.xlsx")
tmp = os.path.join(HERE, "out")
img_dir = os.path.join(HERE, "images")

# Close-ups stop under the last line typed in.
wb = load_workbook(xlsx)
last = lambda ws, end: max(r for r in range(6, end + 1) if ws.cell(r, 2).value not in (None, ""))
chores_end, money_end, kids_end = last(wb["Chores"], 45), last(wb["Money"], 505), last(wb["Kids"], 11)
fridge_end = 5 + sum(1 for r in range(6, 46) if wb["Chores"].cell(r, 3).value == wb["Fridge chart"]["B2"].value)

r = Renders(xlsx, tmp, one_page=("Dashboard",))
close = Renders(xlsx, tmp, one_page=("Dashboard",), hide={"Kids": ["C", "D", "E"]},
                hide_rows={"This week": range(chores_end + 1, 46), "Money": range(money_end + 1, 506),
                           "Kids": range(kids_end + 1, 12), "Fridge chart": range(fridge_end + 1, 18)})


def exact(renders, name):
    """The page whose title is exactly this (Chores is also the start of other titles)."""
    pix = renders.doc[renders.titles.index(name)].get_pixmap(dpi=220)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


img = canvas(dark=True)
title(img, "Chores and pocket money", "Chores with points, a chart for the fridge, and pocket money shared out into "
      "Save, Spend and Give jars", dark=True, width=2300)
laptop(img, r.shot("Chores and pocket money"), 600, 560, 1950)
note(img, ["Google Sheets", "and Excel"], M, 1720, dark=True)
save(img, os.path.join(img_dir, "01-dashboard.jpg"))

detail(os.path.join(img_dir, "02-fridge-chart.jpg"), "A chart for each child to print",
       "Pick a child and print. It shows the days to do each chore, the ticks so far, and the pocket money earned.",
       close.shot("Mia"), [("#DDEFEC", "A day to do it"), ("#BFE3C7", "Done")])

detail(os.path.join(img_dir, "03-this-week.jpg"), "The whole family's week",
       "An x in the box when a chore is done. The points and the pay add up by themselves.",
       close.shot("Chores for the week"), [("#DDEFEC", "A day to do it"), ("#BFE3C7", "Done")])

# Each child's shares and jars from Kids, over the goals and the jar chart from the Dashboard.
page = r.doc[r._index("Chores and pocket money")]
cut = int((page.search_for("Weeks to go")[0].y0 - 14) * 220 / 72)
full = r.page("Chores and pocket money")
lower = content(full.crop((0, cut, full.width, full.height)))
top = close.shot("Kids")
both = Image.new("RGB", (max(top.width, lower.width), top.height + lower.height + 40), "#FFFFFF")
both.paste(top, (0, 0))
both.paste(lower, (0, top.height + 40))
detail(os.path.join(img_dir, "04-jars.jpg"), "Save, Spend and Give jars",
       "Each child's shares, what is in each jar, and how many weeks until the goal at the rate they save",
       both, [("#D8F0DC", "Goal reached"), ("#F9C9C4", "Below zero")])

detail(os.path.join(img_dir, "05-money.jpg"), "Pocket money shared out by itself",
       "Split shares each amount by the child's own jars. Or name one jar for a gift to save or a treat to spend.",
       close.shot("Money"))

img = canvas()
y = title(img, "Inside the file", "Six tabs, settings and a Start Here page. You type in the yellow cells and the "
          "rest is calculated.")
grid(img, [
    (content(r.page("Chores and pocket money")), "Dashboard", "The week and the jars"),
    (content(r.page("Mia")), "Fridge chart", "One child, ready to print"),
    (content(exact(r, "Chores")), "Chores", "Points and the days"),
    (content(r.page("Money")), "Money", "In and out, by jar"),
], y + 80, cols=2)
save(img, os.path.join(img_dir, "06-whats-inside.jpg"))
